"""The firm keeps growing: next month's exports, updated lists, the same file once more.

Every scenario here was found by an independent review and once broke the system.
"""
import io
import json
import shutil
import tempfile
import time
import unittest
import zipfile
from pathlib import Path

from physarum import ops
from physarum.demo.generate import generate
from physarum.grow import ingest, regrow
from physarum.questions import _rid_for, answer, hints, open_questions
from physarum.runtime import Runtime
from physarum.store import Store, loads


def csv(head, rows):
    return (head + "\n" + "\n".join(";".join(r) for r in rows) + "\n").encode("utf-8")


class Firm:
    def __init__(self):
        self.dir = Path(tempfile.mkdtemp(prefix="physarum-regrow-"))
        self.store = Store(self.dir / "firma")

    def feed(self, files, actor="Test"):
        ingest(self.store, [(n, b if isinstance(b, bytes) else b.encode("utf-8")) for n, b in files], actor)
        return regrow(self.store)

    def rt(self):
        rt = Runtime(self.store)
        rt.ensure()
        return rt

    def live(self):
        return [r for r in self.store.q("SELECT id, type, data, origin FROM records WHERE deleted=0 AND type!='doc'")]

    def close(self):
        shutil.rmtree(self.dir, ignore_errors=True)


class Regrowth(unittest.TestCase):
    def setUp(self):
        self.f = Firm()

    def tearDown(self):
        self.f.close()

    def test_next_months_export_updates_instead_of_rebuilding(self):
        base = [(f"K-{1000 + i}", f"Firma {chr(65 + i % 26)}{i} GmbH", ["Hamburg", "Kiel", "Bremen"][i % 3],
                 f"info{i}@firma{i}.de") for i in range(40)]
        head = "Kundennr;Firma;Ort;E-Mail"
        self.f.feed([("kunden_2025-10.csv", csv(head, base))])
        rt = self.f.rt()
        t = next(t for t in rt.model["types"] if t["kind"] != "document")
        ort = next(a for a in t["attrs"] if a["name"] == "Ort")
        first = rt.records[t["id"]][0]
        ops.update(rt, first["id"], {ort["id"]: "Lübeck"}, "Tester")
        ops.add_note(rt, first["id"], "nur vormittags anrufen", "Tester")
        changed = list(base)
        changed[5] = (base[5][0], base[5][1], "Rostock", base[5][3])  # moved in the newer export
        self.f.feed([("kunden_2025-11.csv", csv(head, changed + [("K-1040", "Neu GmbH", "Kiel", "a@neu.de")]))])
        rt = self.f.rt()
        t2 = next(t for t in rt.model["types"] if t["kind"] != "document")
        self.assertEqual(t2["id"], t["id"])
        self.assertEqual(len(rt.records[t["id"]]), 41)
        self.assertEqual(len(self.f.live()), 41)
        rec = self.f.store.record(first["id"])
        self.assertFalse(rec["deleted"])
        self.assertEqual(rec["data"][ort["id"]], "Lübeck")  # the person's edit stands
        k5 = next(r for r in rt.records[t["id"]] if r["data"].get(t["key"]) == "K-1005")
        self.assertEqual(k5["data"][ort["id"]], "Rostock")  # the newest version speaks
        self.assertEqual(self.f.store.q1("SELECT COUNT(*) n FROM notes WHERE record_id=?", (first["id"],))["n"], 1)
        self.assertEqual(self.f.store.q1("SELECT COUNT(*) n FROM changes WHERE kind='remove'")["n"], 0)

    def test_overlapping_bank_statements_do_not_double(self):
        head = "Buchungstag;Verwendungszweck;Betrag"
        a = [(f"{d:02d}.03.2026", f"Zahlung {d}", f"-{10 + d},00") for d in range(1, 21)]
        a += [("05.03.2026", "Kaffee Automat", "-3,50"), ("05.03.2026", "Kaffee Automat", "-3,50")]
        b = a[10:] + [(f"{d:02d}.03.2026", f"Zahlung {d}", f"-{10 + d},00") for d in range(21, 29)]
        self.f.feed([("Umsaetze_1.csv", csv(head, a))])
        ids1 = sorted(r["id"] for r in self.f.live())
        self.f.feed([("Umsaetze_2.csv", csv(head, b))])
        live = self.f.live()
        self.assertEqual(len(live), 30)  # 20 + 2 coffees + 8 new, the overlap counted once
        self.assertTrue(set(ids1) <= {r["id"] for r in live})
        coffee = [r for r in live if "Kaffee" in r["data"]]
        self.assertEqual(len(coffee), 2)  # two equal bookings on one day stay two

    def test_identical_keyless_rows_keep_their_ids(self):
        rows = [(f"{(i % 28) + 1:02d}.03.2026", f"Kartenzahlung {i}", f"-{20 + i},50") for i in range(30)]
        rows += [("05.03.2026", "Kaffee Automat", "-3,50")] * 2
        self.f.feed([("kontoauszug_maerz.csv", csv("Buchungstag;Verwendungszweck;Betrag", rows))])
        before = sorted(r["id"] for r in self.f.live())
        regrow(self.f.store)
        regrow(self.f.store)
        self.assertEqual(sorted(r["id"] for r in self.f.live()), before)
        self.assertEqual(self.f.store.q1("SELECT COUNT(*) n FROM changes WHERE kind='remove'")["n"], 0)

    def test_a_changed_source_value_is_shown_next_to_the_own_edit(self):
        head = "Kundennr;Firma;Ort;Telefon"
        rows = [(f"K-{1000 + i}", f"Firma {i} GmbH", "Hamburg", f"040 {1000000 + i}") for i in range(20)]
        self.f.feed([("kunden_2026-09-15.csv", csv(head, rows))])
        rt = self.f.rt()
        t = next(t for t in rt.model["types"] if t["name"] == "Kunde")
        tel = next(a for a in t["attrs"] if a["name"] == "Telefon")
        rec = next(r for r in rt.records[t["id"]] if r["data"].get(t["key"]) == "K-1000")
        time.sleep(1.1)
        ops.update(rt, rec["id"], {tel["id"]: "0171 5555555"}, "Tester")
        newer = list(rows)
        newer[0] = (rows[0][0], rows[0][1], rows[0][2], "040 7777777")  # the CRM changed it too, before the edit
        self.f.feed([("kunden_2026-09-16.csv", csv(head, newer))])
        r = self.f.store.record(rec["id"])
        self.assertEqual(r["data"][tel["id"]], "0171 5555555")       # the own edit stands …
        self.assertEqual(r["meta"]["shadow"][tel["id"]][0], "040 7777777")  # … and the other value is shown

    def test_user_edit_survives_unrelated_newer_file(self):
        kunden = csv("Kundennr;Firma;Ort;Telefon", [(f"K-{1000 + i}", f"Firma {i} GmbH", "Hamburg", f"040 {1000000 + i}")
                                                    for i in range(20)])
        self.f.feed([("kunden.csv", kunden)])
        rt = self.f.rt()
        t = next(t for t in rt.model["types"] if t["name"] == "Kunde")
        tel = next(a for a in t["attrs"] if a["name"] == "Telefon")
        rec = next(r for r in rt.records[t["id"]] if r["data"].get(t["key"]) == "K-1000")
        time.sleep(1.1)
        ops.update(rt, rec["id"], {tel["id"]: "0171 5555555"}, "Tester")
        time.sleep(1.1)
        orders = csv("Auftragsnr;Datum;Kundennr;Firma;Ort;Betrag",
                     [(f"A-{5000 + i}", f"{(i % 28) + 1:02d}.02.2026", f"K-{1000 + i % 20}", f"Firma {i % 20} GmbH",
                       "Hamburg", f"{100 + i},00") for i in range(80)])
        self.f.feed([("auftraege.csv", orders)])
        self.assertEqual(self.f.store.record(rec["id"])["data"][tel["id"]], "0171 5555555")

    def test_same_file_inside_zip_and_reupload(self):
        kunden = csv("Kundennr;Firma;Ort", [(f"K-{1000 + i}", f"Firma {i} GmbH", "Kiel") for i in range(30)])
        self.f.feed([("kunden.csv", kunden)])
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as z:
            z.writestr("export/kunden.csv", kunden)
        self.f.feed([("export.zip", buf.getvalue())])
        self.assertEqual(len(self.f.live()), 30)
        sid = self.f.store.q1("SELECT id FROM sources WHERE name='kunden.csv' AND parent_id IS NULL")["id"]
        with self.f.store.tx() as c:
            self.f.store.set_active(sid, False, c)
        regrow(self.f.store)
        self.f.feed([("kunden.csv", kunden)])  # dropped in again
        self.assertEqual(len(self.f.live()), 30)
        self.assertEqual(self.f.store.q1("SELECT COUNT(*) n FROM raw_tables WHERE source_id=?", (sid,))["n"], 1)

    def test_nested_zip_is_removed_with_its_parent(self):
        inner = io.BytesIO()
        with zipfile.ZipFile(inner, "w") as z:
            z.writestr("liste/kunden.csv", csv("Kundennr;Firma", [(f"K-{i}", f"Firma {i}") for i in range(10)]))
        outer = io.BytesIO()
        with zipfile.ZipFile(outer, "w") as z:
            z.writestr("inner.zip", inner.getvalue())
        self.f.feed([("outer.zip", outer.getvalue())])
        sid = self.f.store.q1("SELECT id FROM sources WHERE name='outer.zip'")["id"]
        with self.f.store.tx() as c:
            self.f.store.set_active(sid, False, c)
        self.assertEqual(self.f.store.q1("SELECT COUNT(*) n FROM sources WHERE active=1")["n"], 0)

    def test_branches_with_numbers_stay_apart(self):
        rows = [(f"Bäckerei Hansen Filiale {i}", f"Hauptstraße {i}", "Hamburg", f"040 70000{i:02d}") for i in range(1, 32)]
        self.f.feed([("filialen.csv", csv("Name;Straße;Ort;Telefon", rows))])
        self.assertEqual(len(self.f.live()), 31)

    def test_impossible_and_placeholder_dates_keep_pages_working(self):
        from starlette.testclient import TestClient
        from physarum.web.app import create_app
        rows = [(f"A-{500 + i}", f"{(i % 28) + 1:02d}.01.2026", f"{(i % 28) + 1:02d}.02.2026" if i % 4 else "",
                 f"Kunde {i % 5}", f"{100 + i},00") for i in range(40)]
        rows += [("A-997", "31.02.2026", "", "Kunde 1", "50,00"),       # a typo
                 ("A-998", "10.01.2026", "01.01.0001", "Kunde 2", "80,00"),  # "no date" in an old system
                 ("A-999", "31.12.9999", "", "Kunde 3", "90,00")]           # "open-ended"
        self.f.feed([("auftraege.csv", csv("Auftragsnr;Datum;Geliefert am;Kunde;Betrag", rows))])
        c = TestClient(create_app(self.f.dir / "firma"))
        rt = c.app.state.rt
        rt.ensure()
        t = next(t for t in rt.model["types"] if t["name"] == "Auftrag")
        for path in ["/", "/ablaeufe", f"/t/{t['id']}"] + [f"/r/{r['id']}" for r in rt.records[t["id"]]
                                                          if r["data"].get(t["key"]) in ("A-997", "A-998", "A-999")]:
            self.assertEqual(c.get(path).status_code, 200, path)


class DemoDecisions(unittest.TestCase):
    """Answers are constraints for all future growth – and must not break it."""

    @classmethod
    def setUpClass(cls):
        cls.f = Firm()
        res = generate(cls.f.dir / "dump")
        ingest(cls.f.store, [(p.name, p.read_bytes()) for p in res["initial"]], "Test")
        regrow(cls.f.store)

    @classmethod
    def tearDownClass(cls):
        cls.f.close()

    def test_every_answer_then_growth(self):
        store = self.f.store
        rt = self.f.rt()
        # merge a duplicate while a number equal to the dropped id exists elsewhere
        q = next(q for q in open_questions(store) if q["kind"] == "same_record")
        c = store.conn()
        drop = max(_rid_for(c, q["type"], q["idents"][0]), _rid_for(c, q["type"], q["idents"][1]))
        art = next(t for t in rt.model["types"] if t["name"] == "Artikel")
        num = next(a for a in art["attrs"] if a["kind"] == "number" and not a.get("hidden"))
        rid = ops.create(rt, art["id"], {art["title"] or art["key"]: "Testartikel", num["id"]: str(drop)}, "Tester")
        answer(store, q["id"], "same", "Tester")
        self.assertEqual(store.record(rid)["data"][num["id"]], drop)
        # split a structural merge, then grow again
        split = [q for q in hints(store) if q["kind"] == "confirm_merge"]
        for q in open_questions(store) + hints(store):
            answer(store, q["id"], "split" if q["kind"] == "confirm_merge" else
                   ("different" if q["kind"] == "same_type" else "no"), "Tester")
        regrow(store)
        m = store.model()
        self.assertTrue(m["types"])
        for q in split:  # "keep apart" really keeps them apart – not merged again through another file
            home = {mem["label"]: t["id"] for t in m["types"] for mem in t["members"]}
            obs, into = q["decision"]["obs"], q["decision"]["into"]
            if obs in home and into in home:
                self.assertNotEqual(home[obs], home[into], q["text"])


class OrderIndependence(unittest.TestCase):
    def test_upload_order_does_not_change_the_model(self):
        def grow(reverse):
            f = Firm()
            try:
                res = generate(f.dir / "dump")
                files = sorted(res["initial"], key=lambda p: p.name, reverse=reverse)
                for p in files:
                    ingest(f.store, [(p.name, p.read_bytes())], "Test")
                regrow(f.store)
                m = f.store.model()
                return sorted((t["name"], t["kind"], t["count"]) for t in m["types"])
            finally:
                f.close()
        self.assertEqual(grow(False), grow(True))


class NewerCustomerList(unittest.TestCase):
    """The owner's own, newer customer list: one customer moved, one name has a typo."""

    def test_a_move_is_taken_as_a_whole_and_a_typo_does_not_win(self):
        import openpyxl
        from datetime import date
        f = Firm()
        try:
            res = generate(f.dir / "dump")
            f.feed([(p.name, p.read_bytes()) for p in res["initial"]])
            src = next(x for x in res["initial"] if x.name == "Kundenliste_alt.xlsx")
            wb = openpyxl.load_workbook(src)
            ws = wb.active
            ws.cell(row=2, column=1).value = f"Stand: {date.today().strftime('%d.%m.%Y')}"
            for row in ws.iter_rows(min_row=5):
                if row[0].value == "K-1003":            # moved to Hamburg, new number
                    row[3].value, row[4].value, row[5].value, row[6].value = "Elbchaussee 12", "22763", "Hamburg", "040 3041499"
                if row[0].value == "K-1001":            # a slip of the finger
                    row[1].value = "Bistro Krabenkutter GmbH & Co. KG"
            out = f.dir / "Kundenliste_2026.xlsx"
            wb.save(out)
            time.sleep(1)
            f.feed([(out.name, out.read_bytes())])
            rt = f.rt()
            kunde = next(t for t in rt.model["types"] if t["name"] == "Kunde")
            by = {r["data"].get(kunde["key"]): r for r in rt.records[kunde["id"]]}
            attr = {a["name"]: a["id"] for a in kunde["attrs"]}
            moved = by["K-1003"]["data"]
            self.assertEqual((moved[attr["Straße"]], moved[attr["PLZ"]], moved[attr["Ort"]]), ("Elbchaussee 12", "22763", "Hamburg"))
            self.assertEqual(moved[attr["Telefon"]], "040 3041499")
            self.assertEqual(by["K-1001"]["data"][kunde["title"]], "Bistro Krabbenkutter GmbH & Co. KG")
            self.assertIn("Bistro Krabenkutter GmbH & Co. KG",
                          [x[0] for x in (by["K-1001"]["meta"].get("alts") or {}).get(kunde["title"], [])])
            member = next(m for m in kunde["members"] if m.get("versions") and m["versions"].get("files", 0) > 1)
            self.assertEqual(member["asof"], date.today().isoformat())   # the list is as current as its newest version
        finally:
            f.close()


if __name__ == "__main__":
    unittest.main()
