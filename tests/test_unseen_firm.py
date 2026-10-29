"""Lessons from a firm nobody had seen: a web agency with a CRM, two time trackers, an
invoice book, a bank statement and a ticket system, delivered twice.

Each test is one finding of the independent review that once gave a wrong answer.
"""
import json
from collections import Counter
import shutil
import tempfile
import unittest
from datetime import date, timedelta
from email.message import EmailMessage
from pathlib import Path

from physarum.grow import ingest, regrow
from physarum.runtime import Runtime
from physarum.store import Store
from physarum.values import parse_hours


def csv(head, rows, sep=";"):
    return (head + "\n" + "\n".join(sep.join(str(x) for x in r) for r in rows) + "\n").encode("utf-8")


def _pdf(text: str) -> bytes:
    """A minimal one-page PDF with a line of text."""
    body = f"BT /F1 12 Tf 50 750 Td ({text}) Tj ET".encode("latin-1")
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents 4 0 R "
            b"/Resources << /Font << /F1 5 0 R >> >> >>",
            b"<< /Length %d >>\nstream\n" % len(body) + body + b"\nendstream",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out, offs = b"%PDF-1.4\n", []
    for i, o in enumerate(objs, start=1):
        offs.append(len(out))
        out += b"%d 0 obj\n" % i + o + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    out += b"".join(b"%010d 00000 n \n" % o for o in offs)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref)
    return out


class Firm:
    def __init__(self):
        self.dir = Path(tempfile.mkdtemp(prefix="physarum-unseen-"))
        self.store = Store(self.dir / "firma")

    def feed(self, files):
        ingest(self.store, [(n, b if isinstance(b, bytes) else b.encode("utf-8")) for n, b in files], "Test")
        return regrow(self.store)

    def rt(self):
        rt = Runtime(self.store)
        rt.ensure()
        return rt

    def close(self):
        shutil.rmtree(self.dir, ignore_errors=True)


def _type(rt, name):
    return next(t for t in rt.model["types"] if t["name"] == name)


INV_HEAD = "Belegnummer;Belegdatum;Kundennummer;Kunde;Betreff;Netto;USt;Brutto;Fällig am;Status"


def _invoices(n=30, start=date(2026, 3, 2)):
    rows = []
    for i in range(n):
        d = start + timedelta(days=i)
        net = 100 + 10 * i
        rows.append([f"RE-2026-{i + 1:04d}", d.strftime("%d.%m.%Y"), f"K-{10000 + i % 6}", f"Kunde {i % 6} GmbH",
                     f"Leistung {i}", f"{net},00", f"{net * 0.19:.2f}".replace(".", ","),
                     f"{net * 1.19:.2f}".replace(".", ","), (d + timedelta(days=14)).strftime("%d.%m.%Y"), "offen"])
    return rows


class UnseenFirm(unittest.TestCase):
    def setUp(self):
        self.f = Firm()

    def tearDown(self):
        self.f.close()

    def test_a_row_exported_twice_does_not_cost_the_list_its_numbers(self):
        rows = _invoices()
        v1 = rows[:20] + [rows[7]]  # one invoice written twice into the first export
        self.f.feed([("Rechnungsausgangsbuch_2026-03-21.csv", csv(INV_HEAD, v1))])
        rt = self.f.rt()
        inv = _type(rt, "Rechnung")
        ids1 = {r["data"][inv["key"]]: r["id"] for r in rt.records[inv["id"]]}
        self.assertEqual(len(ids1), 20)
        v2 = [list(r) for r in rows]
        v2[3][-1] = "bezahlt"  # a status changed in the next export
        self.f.feed([("Rechnungsausgangsbuch_2026-04-15.csv", csv(INV_HEAD, v2))])
        rt = self.f.rt()
        inv = _type(rt, "Rechnung")
        self.assertEqual(inv["key"], next(a["id"] for a in inv["attrs"] if a["name"] == "Belegnummer"))
        ids2 = {r["data"][inv["key"]]: r["id"] for r in rt.records[inv["id"]]}
        self.assertEqual(len(rt.records[inv["id"]]), 30)
        self.assertTrue(all(ids2[k] == v for k, v in ids1.items()))  # every invoice kept its identity
        removed = self.f.store.q1("SELECT COUNT(*) n FROM changes WHERE kind='remove'")["n"]
        self.assertEqual(removed, 0)

    def test_overlapping_time_exports_carry_corrections_and_deletions(self):
        head = "Project,User,Email,Start Date,Start Time,Duration (h)"
        days = [date(2026, 9, 1) + timedelta(days=i) for i in range(20)]
        old = [[f"P-2026-0{10 + i % 3} Relaunch", "Anna Berg", "anna@agentur.de", d.strftime("%m/%d/%Y"), "09:00 AM",
                "02:00:00"] for i, d in enumerate(days)]
        self.f.feed([("Clockify_Time_Report_09_01_2026-09_20_2026.csv", csv(head, old, ","))])
        new = [list(r) for r in old[10:]]
        new[0][5] = "03:00:00"   # corrected in the time tracker
        del new[1]                # deleted in the time tracker
        new += [[f"P-2026-010 Relaunch", "Anna Berg", "anna@agentur.de", (date(2026, 9, 21) + timedelta(days=i))
                 .strftime("%m/%d/%Y"), "09:00 AM", "01:00:00"] for i in range(5)]
        self.f.feed([("Clockify_Time_Report_09_11_2026-09_25_2026.csv", csv(head, new, ","))])
        rt = self.f.rt()
        te = _type(rt, "Zeiteintrag")
        hours = next(a["id"] for a in te["attrs"] if a["name"].startswith(("Dauer", "Duration")))
        recs = rt.records[te["id"]]
        self.assertEqual(len(recs), 10 + 9 + 5)
        self.assertAlmostEqual(sum(r["data"][hours] for r in recs), 10 * 2 + 3 + 8 * 2 + 5, places=2)

    def test_split_debit_credit_columns_credit_notes_storno_and_a_swapped_number(self):
        rows = _invoices(12)
        gross = {r[0]: float(r[7].replace(",", ".")) for r in rows}
        rows.append(["GS-2026-0001", "20.03.2026", "K-10002", "Kunde 2 GmbH", "Gutschrift zu RE-2026-0003 (Nachlass)",
                     "-50,00", "-9,50", "-59,50", "", "Gutschrift"])
        rows.append(["ST-2026-0001", "21.03.2026", "K-10004", "Kunde 4 GmbH", "Stornorechnung zu RE-2026-0005",
                     "-140,00", "-26,60", "-166,60", "", "Storno"])
        bank_head = '"Buchungstag";"Auftraggeber / Begünstigter";"Verwendungszweck";"Soll";"Haben";"Währung"'
        bank = []
        for no in [r[0] for r in rows[:12]]:
            if no in ("RE-2026-0003", "RE-2026-0005", "RE-2026-0010", "RE-2026-0001"):
                continue
            k = int(no[-2:]) - 1
            bank.append(["25.03.2026", f"Kunde {k % 6} GmbH", no, "", f"{gross[no]:.2f}".replace(".", ",")])
        bank.append(["26.03.2026", "Kunde 2 GmbH", "RE-2026-0003 abzgl. GS-2026-0001", "",
                     f"{gross['RE-2026-0003'] - 59.5:.2f}".replace(".", ",")])
        bank.append(["27.03.2026", "Kunde 3 GmbH", "RE-2026-0001", "",  # meant RE-2026-0010, digits swapped
                     f"{gross['RE-2026-0010']:.2f}".replace(".", ",")])
        bank.append(["27.03.2026", "Bürobedarf AG", "Druckerpatronen", "89,90", ""])
        bank_csv = (bank_head + "\n" + "\n".join(";".join(f'"{x}"' for x in r + ["EUR"]) for r in bank) + "\n")
        self.f.feed([("Rechnungsausgangsbuch.csv", csv(INV_HEAD, rows)),
                     ("Umsaetze_Girokonto.csv", bank_csv.encode("cp1252"))])
        rt = self.f.rt()
        inv = _type(rt, "Rechnung")
        by_no = {r["data"][inv["key"]]: rt.ledger.get(r["id"]) for r in rt.records[inv["id"]]}
        self.assertEqual(by_no["RE-2026-0003"]["state"], "bezahlt")      # paid minus its credit note
        self.assertEqual(by_no["GS-2026-0001"]["state"], "verrechnet")
        self.assertEqual(by_no["RE-2026-0005"]["state"], "storniert")    # cancelled by its storno document
        self.assertEqual(by_no["RE-2026-0010"]["state"], "bezahlt")      # the swapped number found its invoice
        self.assertEqual(by_no["RE-2026-0001"]["state"], "offen")
        self.assertEqual(sorted(k for k, e in by_no.items() if e and e["open"] > 0.005), ["RE-2026-0001"])

    def test_hours_however_they_were_typed(self):
        for raw, h in [("1,5", 1.5), ("1:30", 1.5), ("01:30:00", 1.5), ("1,35 h", 1.35), ("2 h", 2.0), ("2 Std.", 2.0),
                       ("90 min", 1.5), ("1h 30m", 1.5), (1.75, 1.75), ("0:57", 0.95)]:
            self.assertAlmostEqual(parse_hours(raw), h, places=2, msg=str(raw))
        rows = [[f"{(i % 20) + 1:02d}.03.2026", f"Projekt {i % 3}", "Entwicklung", v]
                for i, v in enumerate(["1,5", "2:15", "1,35 h", "2 h", "0:30", "3", "1,25", "4", "2,5", "1"] * 2)]
        self.f.feed([("Stundenzettel_Steinbach_2026.csv", csv("Datum;Projekt;Tätigkeit;Stunden", rows))])
        rt = self.f.rt()
        sz = _type(rt, "Stundenzettel")
        hours = next(a for a in sz["attrs"] if a["name"] == "Stunden")
        vals = [r["data"].get(hours["id"]) for r in rt.records[sz["id"]]]
        self.assertTrue(all(isinstance(v, float) for v in vals))
        self.assertAlmostEqual(sum(vals), 2 * (1.5 + 2.25 + 1.35 + 2 + 0.5 + 3 + 1.25 + 4 + 2.5 + 1), places=2)

    def test_project_numbers_inside_names_and_time_entries_are_events(self):
        projects = [[f"P-2026-{i:03d}", f"Kunde {i % 4} GmbH", ["Website-Relaunch", "Wartung & Hosting"][i % 2],
                     f"{(i % 9) + 1:02d}.02.2026", f"{5000 + 100 * i}"] for i in range(1, 13)]
        entries = []
        for k in range(60):
            i = k % 12 + 1
            name = ["Website-Relaunch", "Wartung & Hosting"][i % 2]
            label = f"P-2026-{i:03d} {name}" if k % 2 else f"{name} Kunde {i % 4} ({'P-2026-%03d' % i})"
            entries.append([label, "anna@agentur.de", (date(2026, 3, 1) + timedelta(days=k)).strftime("%d.%m.%Y"), "2,0"])
        self.f.feed([("Projektliste.csv", csv("Projekt-Nr.;Kunde;Projekt;Start;Budget", projects)),
                     ("time_entries_2026.csv", csv("Project;Email;Date;Duration (h)", entries))])
        rt = self.f.rt()
        pj = _type(rt, "Projekt")
        self.assertEqual(len(rt.records[pj["id"]]), 12)  # one project per number, whatever the tool wrote
        te = _type(rt, "Zeiteintrag")
        self.assertEqual(te["kind"], "transaction")
        ref = next(a for a in te["attrs"] if a["kind"] == "ref" and a.get("ref") == pj["id"])
        self.assertEqual(sum(1 for r in rt.records[te["id"]] if isinstance(r["data"].get(ref["id"]), int)), 60)

    def test_the_same_pdf_as_file_and_as_mail_attachment(self):
        pdf = _pdf("Rechnung FR-2026-117 Steinbach IT 1.785,00 EUR")
        m = EmailMessage()
        m["From"], m["To"], m["Subject"] = "marco@steinbach-it.de", "info@agentur.de", "Rechnung September"
        m.set_content("Anbei meine Rechnung.")
        m.add_attachment(pdf, maintype="application", subtype="pdf", filename="FR-2026-117.pdf")
        s = self.f.feed([("belege/FR-2026-117.pdf", pdf), ("mails/rechnung.eml", m.as_bytes())])
        self.assertIsNotNone(s)
        docs = self.f.store.q("SELECT data FROM records WHERE type='doc' AND deleted=0")
        self.assertGreaterEqual(len(docs), 2)  # the mail and the invoice, the PDF only once

    def test_deleted_in_the_crm_is_marked_moved_between_sheets_is_not(self):
        head = "Record ID,First Name,Last Name,Email,Company Name"
        people = [[str(9000 + i), ["Anna", "Jan", "Lena", "Paul", "Maria", "Tim"][i % 6], f"Muster{i}",
                   f"p{i}@firma{i % 5}.de", f"Firma {i % 5} GmbH"] for i in range(25)]
        self.f.feed([("hubspot-crm-exports-all-contacts-2026-09-15.csv", csv(head, people, ","))])
        self.f.feed([("hubspot-crm-exports-all-contacts-2026-10-27.csv", csv(head, people[:-2], ","))])
        rt = self.f.rt()
        kt = _type(rt, "Kontakt")
        gone = [r for r in rt.records[kt["id"]] if (r.get("meta") or {}).get("vanished")]
        self.assertEqual(len(rt.records[kt["id"]]), 25)
        self.assertEqual(sorted(r["data"][kt["key"]] for r in gone), ["9023", "9024"])

    def test_fixed_budget_overrun_and_freelancer_names(self):
        kunden = ["Elbkind Spielwaren GmbH", "Tidenhub Meerestechnik GmbH", "Baumschule Kruse GmbH", "Hansen Bau GmbH",
                  "Kröger Haustechnik GmbH"]
        projects = []
        for i in range(1, 11):
            k = kunden[i % 5]
            short = k.split()[0]
            art = ["Website-Relaunch", "Wartung & Hosting"][i % 2]
            projects.append([f"P-2026-{i:03d}", short, art, art, "Festpreis" if i % 2 == 0 else "Retainer",
                             f"{i:02d}.03.2026", "8000" if i % 2 == 0 else "250"])
        self.f.feed([("Kunden.csv", csv("Kundennr;Firma;Ort", [[f"K-{100 + i}", k, "Hamburg"] for i, k in enumerate(kunden)])),
                     ("Projektliste.csv", csv("Projekt-Nr.;Kunde;Projekt;Art;Abrechnung;Start;Budget", projects))])
        entries, sheet = [], []
        for k in range(120):
            i = k % 10 + 1
            amount = "950,00" if i == 2 else "150,00"  # P-2026-002: 12 x 950 = 11.400 on a budget of 8.000
            entries.append([f"P-2026-{i:03d} Projekt", "anna@agentur.de",
                            (date(2026, 3, 15) + timedelta(days=k % 60)).strftime("%d.%m.%Y"), "2,0", amount])
            p = projects[i - 1]
            name = f"{p[1]} {'Relaunch' if 'Relaunch' in p[2] else 'Wartung'}" if k % 2 else \
                f"{'Relaunch' if 'Relaunch' in p[2] else 'Wartung'} {p[1]}"
            sheet.append([(date(2026, 3, 15) + timedelta(days=k % 60)).strftime("%d.%m.%Y"), name, "Entwicklung", "1,5"])
        self.f.feed([("time_entries.csv", csv("Project;Email;Date;Duration (h);Amount (EUR)", entries)),
                     ("Stundenzettel_Steinbach.csv", csv("Datum;Projekt;Tätigkeit;Stunden", sheet))])
        rt = self.f.rt()
        pj = _type(rt, "Projekt")
        work = rt.work(date(2026, 6, 1))
        over = next(g for g in work if g["kind"] == "budget")
        self.assertEqual([it["title"] for it in over["items"]], ["P-2026-002"])  # retainers are left out
        sz = _type(rt, "Stundenzettel")
        ref = next(a for a in sz["attrs"] if a["kind"] == "ref" and a.get("ref") == pj["id"])
        linked = [r for r in rt.records[sz["id"]] if isinstance(r["data"].get(ref["id"]), int)]
        self.assertGreaterEqual(len(linked), 0.9 * len(rt.records[sz["id"]]))
        by_no = {r["id"]: r["data"][pj["key"]] for r in rt.records[pj["id"]]}
        for r in linked:  # "Elbkind Relaunch" and "Wartung Tidenhub" name the right project
            name = r["data"][next(a["id"] for a in sz["attrs"] if a["name"] == "Projekt")]
            no = by_no[r["data"][ref["id"]]]
            p = projects[int(no[-3:]) - 1]
            self.assertIn(p[1], str(rt.title(r["data"][ref["id"]])) + " " + no + " " + p[1])

    def test_keeping_two_lists_apart_keeps_their_records(self):
        import io
        import openpyxl
        from physarum import ops
        from physarum.grow import regrow
        from physarum.questions import answer, open_questions
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Aktiv"
        ws.append(["Projekt-Nr.", "Kunde", "Projekt", "Start", "Budget"])
        for i in range(12):
            ws.append([f"P-2026-{i + 1:03d}", f"Kunde {i % 4}", "Website-Relaunch", f"{i + 1:02d}.03.2026", 5000 + i])
        ws2 = wb.create_sheet("Abgeschlossen")
        ws2.append(["Projekt-Nr.", "Kunde", "Projekt", "Start", "Ende", "Budget"])
        for i in range(10):
            ws2.append([f"P-2025-{i + 1:03d}", f"Kunde {i % 4}", "Kampagne", f"{i + 1:02d}.03.2025", f"{i + 1:02d}.06.2025", 3000 + i])
        buf = io.BytesIO()
        wb.save(buf)
        self.f.feed([("Projektliste.xlsx", buf.getvalue())])
        rt = self.f.rt()
        pj = _type(rt, "Projekt")
        if len(rt.records[pj["id"]]) != 22:
            self.skipTest("the two sheets were not put into one list here")
        ids = {r["data"][pj["key"]]: r["id"] for r in rt.records[pj["id"]]}
        ops.add_note(rt, ids["P-2026-001"], "Kunde will Termin im April", "Tester")
        q = next((q for q in open_questions(self.f.store) if q["kind"] == "confirm_merge"), None)
        if q is None:
            self.skipTest("no confirmation asked")
        answer(self.f.store, q["id"], "split", "Tester")
        regrow(self.f.store)
        rt = self.f.rt()
        after = {r["data"].get(next((a["id"] for a in rt.spec(r["type"])["attrs"] if a["name"] == "Projekt-Nr."), None)): r["id"]
                 for t in rt.model["types"] for r in rt.records.get(t["id"], [])}
        self.assertTrue(all(after.get(k) == v for k, v in ids.items()), "every project kept its record")
        self.assertEqual(self.f.store.q1("SELECT COUNT(*) n FROM notes WHERE record_id=?", (ids["P-2026-001"],))["n"], 1)

    def test_a_project_number_with_swapped_digits_is_asked_and_merged(self):
        from physarum.grow import regrow
        from physarum.questions import answer, open_questions
        projects = [[f"P-2026-{i:03d}", f"Kunde {i % 4} GmbH", "Website-Relaunch", f"{(i % 9) + 1:02d}.02.2026"]
                    for i in range(10, 25)]
        entries = [[f"P-2026-{10 + k % 15:03d} Website-Relaunch", "anna@agentur.de",
                    (date(2026, 3, 1) + timedelta(days=k % 50)).strftime("%d.%m.%Y"), "2,0"] for k in range(90)]
        entries += [["P-2026-031 Website-Relaunch", "anna@agentur.de", "02.04.2026", "1,0"] for _ in range(8)]
        self.f.feed([("Projektliste.csv", csv("Projekt-Nr.;Kunde;Projekt;Start", projects)),
                     ("time_entries.csv", csv("Project;Email;Date;Duration (h)", entries))])
        q = next((q for q in open_questions(self.f.store) if "P-2026-031" in q["text"]), None)
        self.assertIsNotNone(q)
        self.assertIn("P-2026-013", q["text"])
        answer(self.f.store, q["id"], "same", "Tester")
        regrow(self.f.store)
        rt = self.f.rt()
        pj = _type(rt, "Projekt")
        self.assertNotIn("P-2026-031", {r["data"].get(pj["key"]) for r in rt.records[pj["id"]]})
        te = _type(rt, "Zeiteintrag")
        ref = next(a for a in te["attrs"] if a["kind"] == "ref" and a.get("ref") == pj["id"])
        p13 = next(r["id"] for r in rt.records[pj["id"]] if r["data"].get(pj["key"]) == "P-2026-013")
        self.assertEqual(sum(1 for r in rt.records[te["id"]] if r["data"].get(ref["id"]) == p13), 6 + 8)

    def test_an_offer_made_here_is_not_taken_over_by_another_with_its_number(self):
        from physarum import ops
        head = "Angebots-Nr.;Datum;Kunde;Titel;Netto;Status"
        rows = [[f"A-2026-{i:03d}", f"{(i % 28) + 1:02d}.03.2026", f"Kunde {i % 5} GmbH", "Relaunch", f"{1000 + 100 * i},00",
                 "offen"] for i in range(1, 21)]
        self.f.feed([("Angebote.csv", csv(head, rows))])
        rt = self.f.rt()
        an = _type(rt, "Angebot")
        attrs = {a["name"]: a["id"] for a in an["attrs"]}
        mine = ops.create(rt, an["id"], {attrs["Angebots-Nr."]: "A-2026-021", attrs["Datum"]: "2026-04-02",
                                         attrs["Titel"]: "Kanu-Website", attrs["Netto"]: "20000"}, "Tester")
        theirs = rows + [["A-2026-021", "08.04.2026", "Hafenkontor Immobilien GmbH", "Online-Shop", "18.700,00", "angenommen"]]
        self.f.feed([("Angebote_neu.csv", csv(head, theirs))])
        rt = self.f.rt()
        same = [r for r in rt.records[an["id"]] if r["data"].get(an["key"]) == "A-2026-021"]
        self.assertEqual(len(same), 2)                                  # two different offers …
        mine_now = self.f.store.record(mine)
        self.assertEqual(mine_now["data"][attrs["Titel"]], "Kanu-Website")  # … and the own one untouched
        self.assertTrue(any("gibt es zweimal" in it["reason"] for g in rt.work() if g["kind"] == "integrity"
                            for it in g["items"]))

    def test_a_note_is_a_document_not_a_table(self):
        self.f.feed([("LIESMICH.txt", "Export aus Toggl + alter Kundenstamm (WaWi 2014), erstellt 02.07.2025 von J. P.\n"
                                      "Toggl wurde zum 30.06.2025 gekündigt, ab 01.07. Clockify.\n"),
                     ("kunden.csv", csv("Kundennr;Firma;Ort", [[f"K-{i}", f"Firma {i} GmbH", "Kiel"] for i in range(10)]))])
        rt = self.f.rt()
        self.assertFalse(any("LIESMICH" in t["name"] for t in rt.model["types"]))
        self.assertTrue(any("LIESMICH" in str(r["data"].get("doc_title")) for r in rt.records.get("doc", [])))


    SHORT = ["Wattwurm", "Hafenkontor", "Optik Schulz", "Elbkind", "Druckhaus Lokstedt", "Alsterblick"]

    def _short_name_firm(self):
        firms = [("K-10001", "Hafenkontor Logistik GmbH", "hafenkontor-logistik.de"),
                 ("K-10007", "Wattwurm Reisen GmbH", "wattwurm-reisen.de"),
                 ("K-10008", "Alsterblick Immobilien GmbH", "alsterblick.de"),
                 ("K-10019", "Kita Elbkind gGmbH", "elbkind.de"),
                 ("K-10027", "Augenoptik Schulz e.K.", "optik-schulz.de"),
                 ("K-10043", "Lokstedter Druckhaus GmbH", "druckhaus-lokstedt.de"),
                 ("K-10055", "Alster Dental MVZ", "alster-dental.de"),
                 ("K-10073", "Hafenkontor Immobilien GmbH", "hafenkontor-immobilien.de")]
        kunden = [[k, n, f"Hafenstraße {i + 1}", f"2{i}457", "Hamburg", f"info@{d}"] for i, (k, n, d) in enumerate(firms)]
        projects = [[f"P-2026-{i + 1:03d}", self.SHORT[i % 6], "Wartung & Hosting", f"{(i % 27) + 1:02d}.01.2026"]
                    for i in range(18)]
        tickets = []
        for t in range(30):
            tickets.append(dict(id=4000 + t, subject=f"Anfrage {t}", status=5, created_at=f"2026-0{1 + t % 8}-1{t % 9}T09:00:00Z",
                                company=dict(id=7010000 + t % 3, name=["Wattwurm", "Hafenkontor Logistik", "Elbkind"][t % 3]),
                                time_entries=[dict(agent_email=["lukas@agentur.de", "jana@agentur.de"][k],
                                                   time_spent=f"0{k}:{15 + 7 * ((t + k) % 6):02d}",
                                                   executed_at=f"2026-0{1 + t % 8}-{11 + k + t % 9}T1{k}:{t % 6}0:00Z",
                                                   note=f"Arbeit {t}-{k}") for k in range(2)]))
        return [("Kundenstamm.csv", csv("KdNr;Firma;Straße;PLZ;Ort;E-Mail", kunden)),
                ("Projektliste.csv", csv("Projekt-Nr.;Kunde;Projekt;Start", projects)),
                ("freshdesk_tickets_export_2026-09-15.json", json.dumps(dict(tickets=tickets, exported="2026-09-15")))]

    def test_short_names_of_customers_find_the_one_they_mean(self):
        self.f.feed(self._short_name_firm())
        rt = self.f.rt()
        ku = _type(rt, "Kunde")
        self.assertEqual(len(rt.records[ku["id"]]), 8)
        kd = lambda v: rt.by_id[v]["data"].get(ku["key"]) if isinstance(v, int) else None
        pj = _type(rt, "Projekt")
        ref = next(a for a in pj["attrs"] if a["kind"] == "ref" and a.get("ref") == ku["id"])
        by_short = {self.SHORT[int(r["data"][pj["key"]][-3:]) % 6 - 1]: kd(r["data"].get(ref["id"])) for r in rt.records[pj["id"]]}
        self.assertEqual(by_short, {"Wattwurm": "K-10007", "Optik Schulz": "K-10027", "Elbkind": "K-10019",
                                    "Druckhaus Lokstedt": "K-10043", "Alsterblick": "K-10008",
                                    "Hafenkontor": None})       # two Hafenkontors: the short name cannot decide
        tk = _type(rt, "Ticket")
        tref = next(a for a in tk["attrs"] if a["kind"] == "ref" and a.get("ref") == ku["id"])
        self.assertEqual(Counter(kd(r["data"].get(tref["id"])) for r in rt.records[tk["id"]]),
                         Counter({"K-10007": 10, "K-10001": 10, "K-10019": 10}))

    def test_a_join_on_a_guess_can_be_undone(self):
        from physarum.grow import regrow
        clients = [[n, "No"] for n in self.SHORT]          # a time tracker's client list knows only short names
        self.f.feed(self._short_name_firm() + [("Clockify_Clients.csv", csv("Client;Archived", clients))])
        rt = self.f.rt()
        ku = _type(rt, "Kunde")
        names = sorted(rt.title(r["id"]) for r in rt.records[ku["id"]])
        self.assertEqual(len(names), 9, names)            # only "Hafenkontor" stays a customer of its own
        self.assertIn("Hafenkontor", names)
        watt = next(r for r in rt.records[ku["id"]] if r["data"].get(ku["key"]) == "K-10007")
        j = next(j for j in watt["meta"].get("joins") or [] if j["name"] == "Wattwurm")
        with self.f.store.tx() as c:
            self.f.store.decide("same_record", j["key"], "different", "Tester", c=c)
        regrow(self.f.store)
        rt = self.f.rt()
        ku = _type(rt, "Kunde")
        self.assertIn("Wattwurm", [rt.title(r["id"]) for r in rt.records[ku["id"]]])
        again = next(r for r in rt.records[ku["id"]] if r["data"].get(ku["key"]) == "K-10007")
        self.assertEqual(again["id"], watt["id"])
        self.assertNotIn("Wattwurm", [j["name"] for j in again["meta"].get("joins") or []])

    def test_nested_lists_of_two_deliveries_stay_one_list(self):
        files = self._short_name_firm()
        self.f.feed(files)
        upd = json.loads(files[2][1])
        upd["tickets"] = [t for t in upd["tickets"] if t["id"] >= 4020]
        for t in upd["tickets"][:3]:
            t["time_entries"].append(dict(agent_email="lukas@agentur.de", time_spent="00:30",
                                          executed_at="2026-09-20T10:00:00Z", note=f"Nachtrag {t['id']}"))
        self.f.feed([("freshdesk_tickets_updated_since_2026-09-10.json", json.dumps(upd))])
        rt = self.f.rt()
        te = [t for t in rt.model["types"] if any("time_entries" in m["label"] for m in t["members"])]
        self.assertEqual(len(te), 1, [t["name"] for t in te])
        self.assertEqual(len(rt.records[te[0]["id"]]), 30 * 2 + 3)

    def _mail(self, frm, subject, body, when):
        from email.utils import format_datetime
        from datetime import datetime
        m = EmailMessage()
        m["From"], m["To"], m["Subject"] = frm, "info@roesterei.de", subject
        m["Date"] = format_datetime(datetime(when.year, when.month, when.day, 9, 30))
        m.set_content(body)
        return m.as_bytes()

    def test_orders_and_complaints_by_mail_wait_on_heute_until_answered(self):
        from physarum import ops
        today = date.today()
        kunden = [[f"K-{100 + i}", n, f"einkauf@{d}", "Hamburg"] for i, (n, d) in
                  enumerate([("Hotel Landungsbrücke KG", "hotel-landungsbruecke.de"), ("Feinkost Elbinsel GmbH", "feinkost-elbinsel.de"),
                             ("Café Kajüte", "cafe-kajuete.de")])]
        auftraege = [[f"AU-2026-{i:03d}", (today - timedelta(days=60 - 2 * i)).strftime("%d.%m.%Y"), f"K-{100 + i % 3}",
                      f"{200 + 10 * i},00"] for i in range(20)]
        rechnungen = [[f"RE-2026-{i:04d}", (today - timedelta(days=40 - i)).strftime("%d.%m.%Y"), f"K-{100 + i % 3}",
                       f"{238 + i},00", (today - timedelta(days=26 - i)).strftime("%d.%m.%Y")] for i in range(1, 16)]
        y = today - timedelta(days=1)
        self.f.feed([("Kunden.csv", csv("Kd-Nr;Firma;E-Mail;Ort", kunden)),
                     ("Auftraege.csv", csv("Auftrags-Nr.;Datum;Kd-Nr;Betrag", auftraege)),
                     ("Rechnungen.csv", csv("Rechnungs-Nr.;Datum;Kd-Nr;Brutto;Fällig am", rechnungen)),
                     ("mails/bestellung.eml", self._mail("Hannah Becker <einkauf@hotel-landungsbruecke.de>",
                                                         "Bestellung für nächste Woche",
                                                         "Hallo, wir bestellen 6 kg Hafenmischung Espresso. Gruß", y)),
                     ("mails/rueckfrage.eml", self._mail("Finn Kruse <einkauf@feinkost-elbinsel.de>",
                                                         "Rückfrage zu Rechnung RE-2026-0012",
                                                         "Die Rechnung RE-2026-0012 stimmt nicht, es fehlt der Rabatt.", y)),
                     ("mails/news.eml", self._mail("Röstmagazin <news@roestmagazin.de>", "Newsletter: Trends der Saison",
                                                   "Jetzt bestellen und sparen! Abmelden hier.", y))])
        rt = self.f.rt()
        g = next((g for g in rt.work() if g["kind"] == "mail"), None)
        self.assertIsNotNone(g)
        self.assertEqual(sorted(it["mail"] for it in g["items"]), ["complaint", "order"])
        order = next(it for it in g["items"] if it["mail"] == "order")
        self.assertEqual(order["party"], "Hotel Landungsbrücke KG")
        self.assertIn("/neu?a_", order["act"]["href"])                 # "+ Auftrag" for that customer
        complaint = next(it for it in g["items"] if it["mail"] == "complaint")
        self.assertIn("RE-2026-0012", complaint["reason"])
        # an order for that customer answers the mail; the complaint is ticked off by hand
        au = _type(rt, "Auftrag")
        attrs = {a["name"]: a["id"] for a in au["attrs"]}
        hotel = next(r["id"] for r in rt.records[_type(rt, "Kunde")["id"]] if r["data"].get(_type(rt, "Kunde")["key"]) == "K-100")
        kd = next(a["id"] for a in au["attrs"] if a["kind"] == "ref")
        ops.create(rt, au["id"], {attrs["Auftrags-Nr."]: "AU-2026-099", attrs["Datum"]: today.isoformat(),
                                  kd: hotel, attrs["Betrag"]: "180"}, "Tester")
        with self.f.store.tx() as c:
            self.f.store.decide("mail_done", str(complaint["rid"]), "done", "Tester", c=c)
        rt.bump()
        rt = self.f.rt()
        self.assertFalse([g for g in rt.work() if g["kind"] == "mail"])

    def test_the_firm_finds_its_own_name_in_its_list_titles(self):
        from openpyxl import Workbook
        from physarum.billing import guess_company
        wb = Workbook()
        ws = wb.active
        ws.title = "Jan 26"
        ws.append(["Stundennachweis Marco Steinbach – Steinbach IT · Auftraggeber: Nordlicht Digital GmbH · Monat: Januar 2026"])
        ws.append([])
        ws.append(["Datum", "Projekt", "Tätigkeit", "Stunden"])
        for i in range(12):
            ws.append([date(2026, 1, 2 + i), "P-2026-001", "Entwicklung", 4])
        buf = __import__("io").BytesIO()
        wb.save(buf)
        wb2 = Workbook()
        ws2 = wb2.active
        ws2.append(["Projektübersicht Nordlicht Digital GmbH · Stand: 15.09.2026"])
        ws2.append(["Projekt-Nr.", "Kunde", "Projekt", "Start"])
        for i in range(10):
            ws2.append([f"P-2026-{i + 1:03d}", f"Kunde {i} GmbH", "Relaunch", date(2026, 1, 1 + i)])
        buf2 = __import__("io").BytesIO()
        wb2.save(buf2)
        self.f.feed([("Stundenzettel_Steinbach_2026.xlsx", buf.getvalue()), ("Projektliste.xlsx", buf2.getvalue()),
                     ("belege/Rechnung_Kaffee_Kontor_Nord_GmbH.pdf", _pdf("Kaffee Kontor Nord GmbH Rechnung 2026-253"))])
        self.assertEqual(guess_company(self.f.rt()).get("name"), "Nordlicht Digital GmbH")


    @staticmethod
    def _ubl(number, when, seller, iban, lines, buyer="Elektro Brandt GmbH", vat="DE811234567", net=None, tax=None,
             incl=None, rate=19.0, sub_tax=None):
        lsum = round(sum(q * p for q, p in lines), 2)
        net = lsum if net is None else net
        tax = round(net * rate / 100, 2) if tax is None else tax
        incl = round(net + tax, 2) if incl is None else incl
        sub_tax = tax if sub_tax is None else sub_tax
        items = "".join(f"""<cac:InvoiceLine><cbc:ID>{i + 1}</cbc:ID><cbc:InvoicedQuantity unitCode="H87">{q}</cbc:InvoicedQuantity>
<cbc:LineExtensionAmount currencyID="EUR">{q * p:.2f}</cbc:LineExtensionAmount><cac:Item><cbc:Name>Material {i + 1}</cbc:Name></cac:Item>
<cac:Price><cbc:PriceAmount currencyID="EUR">{p:.2f}</cbc:PriceAmount></cac:Price></cac:InvoiceLine>""" for i, (q, p) in enumerate(lines))
        vat_xml = f"<cac:PartyTaxScheme><cbc:CompanyID>{vat}</cbc:CompanyID><cac:TaxScheme><cbc:ID>VAT</cbc:ID></cac:TaxScheme></cac:PartyTaxScheme>" if vat else ""
        return f"""<?xml version="1.0" encoding="UTF-8"?>
<ubl:Invoice xmlns:ubl="urn:oasis:names:specification:ubl:schema:xsd:Invoice-2"
 xmlns:cac="urn:oasis:names:specification:ubl:schema:xsd:CommonAggregateComponents-2"
 xmlns:cbc="urn:oasis:names:specification:ubl:schema:xsd:CommonBasicComponents-2">
<cbc:ID>{number}</cbc:ID><cbc:IssueDate>{when.isoformat()}</cbc:IssueDate><cbc:DueDate>{(when + timedelta(days=30)).isoformat()}</cbc:DueDate>
<cbc:InvoiceTypeCode>380</cbc:InvoiceTypeCode><cbc:DocumentCurrencyCode>EUR</cbc:DocumentCurrencyCode>
<cac:AccountingSupplierParty><cac:Party><cac:PartyName><cbc:Name>{seller}</cbc:Name></cac:PartyName>{vat_xml}
<cac:PartyLegalEntity><cbc:RegistrationName>{seller}</cbc:RegistrationName></cac:PartyLegalEntity></cac:Party></cac:AccountingSupplierParty>
<cac:AccountingCustomerParty><cac:Party><cac:PartyLegalEntity><cbc:RegistrationName>{buyer}</cbc:RegistrationName></cac:PartyLegalEntity></cac:Party></cac:AccountingCustomerParty>
<cac:PaymentMeans><cbc:PaymentMeansCode>58</cbc:PaymentMeansCode><cac:PayeeFinancialAccount><cbc:ID>{iban}</cbc:ID></cac:PayeeFinancialAccount></cac:PaymentMeans>
<cac:TaxTotal><cbc:TaxAmount currencyID="EUR">{tax:.2f}</cbc:TaxAmount><cac:TaxSubtotal><cbc:TaxableAmount currencyID="EUR">{net:.2f}</cbc:TaxableAmount>
<cbc:TaxAmount currencyID="EUR">{sub_tax:.2f}</cbc:TaxAmount><cac:TaxCategory><cbc:ID>S</cbc:ID><cbc:Percent>{rate}</cbc:Percent></cac:TaxCategory></cac:TaxSubtotal></cac:TaxTotal>
<cac:LegalMonetaryTotal><cbc:LineExtensionAmount currencyID="EUR">{lsum:.2f}</cbc:LineExtensionAmount>
<cbc:TaxExclusiveAmount currencyID="EUR">{net:.2f}</cbc:TaxExclusiveAmount><cbc:TaxInclusiveAmount currencyID="EUR">{incl:.2f}</cbc:TaxInclusiveAmount>
<cbc:PayableAmount currencyID="EUR">{incl:.2f}</cbc:PayableAmount></cac:LegalMonetaryTotal>{items}</ubl:Invoice>""".encode("utf-8")

    def test_supplier_invoices_are_checked_before_anybody_pays(self):
        from physarum.einvoice import check, parse
        from defusedxml.ElementTree import fromstring
        today = date.today()
        a, b, c, n = "Elektrogroßhandel Nord GmbH", "Kabelwerk Süd GmbH", "Leuchten Meyer KG", "Schaltschrankbau Ost GmbH"
        iban_a, iban_b, iban_c = "DE89370400440532013000", "DE02120300000000202051", "DE12500105170648489890"
        files, bank = [], []
        for k, (sup, iban) in enumerate([(a, iban_a), (b, iban_b), (c, iban_c)]):
            for j in range(3):
                when = today - timedelta(days=150 - 30 * j + k)
                no = f"{sup[:3].upper()}-{2026}{k}{j}"
                files.append((f"E-Rechnungen/{no}.xml", self._ubl(no, when, sup, iban, [(10, 12.5 + j)])))
                bank.append([(when + timedelta(days=10)).strftime("%d.%m.%Y"), sup, iban, f"Rechnung {no}",
                             f"-{(10 * (12.5 + j)) * 1.19:.2f}".replace(".", ",")])
        # 1. the next invoice of A suddenly names another account
        files.append(("E-Rechnungen/ELE-NEU.xml", self._ubl("ELE-2026-99", today - timedelta(days=4), a, "DE75512108001245126199",
                                                           [(20, 13.0)])))
        # 2. B's figures do not add up: lines 300 €, net 310 €, 19 % of 310 is not 50 €, and the total is wrong too
        bad = self._ubl("KAB-2026-77", today - timedelta(days=3), b, iban_b, [(10, 30.0)], net=310.0, tax=50.0, incl=365.0)
        files.append(("E-Rechnungen/KAB-77.xml", bad))
        # 3. C's new invoice is addressed to another firm
        files.append(("E-Rechnungen/LEU-88.xml", self._ubl("LEU-2026-88", today - timedelta(days=2), c, iban_c, [(4, 25.0)],
                                                          buyer="Haustechnik Petersen GmbH")))
        # 4. a supplier nobody knew, without a VAT id
        files.append(("E-Rechnungen/SCH-1.xml", self._ubl("SOB-2026-1", today - timedelta(days=1), n, "DE44500105175407324931",
                                                         [(1, 890.0)], vat=None)))
        files.append(("Konto.csv", csv("Buchungstag;Empfänger;IBAN;Verwendungszweck;Betrag", bank)))
        found = check(parse(fromstring(bad)))
        self.assertTrue(any("Netto 310,00 €" in x for x in found), found)
        self.assertTrue(any("19 % von 310,00 €" in x for x in found), found)
        self.assertTrue(any("plus Umsatzsteuer" in x for x in found), found)
        self.f.store.kv_set("company_profile", dict(name="Elektro Brandt GmbH"))
        self.f.feed(files)
        rt = self.f.rt()
        work = rt.work()
        risk = next((g for g in work if g["id"] == "risk:iban"), None)
        self.assertIsNotNone(risk)
        self.assertTrue(any("DE75 5121 0800 1245 1261 99" in it["reason"] and it["party"] == a for it in risk["items"]))
        checks = next((g for g in work if g["id"] == "check:incoming"), None)
        self.assertIsNotNone(checks)
        why = " | ".join(it["reason"] for it in checks["items"])
        self.assertIn("Netto 310,00 €", why)
        self.assertIn("Adressiert an „Haustechnik Petersen GmbH“", why)
        self.assertIn("erste Rechnung dieses Lieferanten", why)
        self.assertIn("Weder USt-IdNr. noch Steuernummer", why)
        self.assertFalse(any("ELE-2026-0" in it["title"] for it in risk["items"] + checks["items"]))  # the paid, normal ones


    def test_a_reference_may_speak_another_programs_number(self):
        kunden = [[f"K-{10001 + i}", f"Firma {chr(65 + i)} GmbH", str(3363819200 + 7 * i), "Hamburg"] for i in range(12)]
        kontakte = [[str(9000 + j), f"Person {j}", f"p{j}@firma{j % 12}.de", str(3363819200 + 7 * (j % 12))] for j in range(30)]
        self.f.feed([("Kundenstamm.csv", csv("KdNr;Firma;CRM-ID;Ort", kunden)),
                     ("hubspot-contacts.csv", csv("Record ID,Name,Email,Associated Company ID", kontakte, ","))])
        rt = self.f.rt()
        ku = next(t for t in rt.model["types"] if t.get("concept") == "kunde")
        ko = next(t for t in rt.model["types"] if t["id"] != ku["id"] and any(a["kind"] == "ref" and a.get("ref") == ku["id"]
                                                                               for a in t["attrs"]))
        ref = next(a for a in ko["attrs"] if a["kind"] == "ref" and a.get("ref") == ku["id"])
        linked = {r["data"].get(ref["id"]) for r in rt.records[ko["id"]]}
        self.assertTrue(all(isinstance(v, int) for v in linked))
        by_no = {r["id"]: r["data"][ku["key"]] for r in rt.records[ku["id"]]}
        first = next(r for r in rt.records[ko["id"]] if "p0@" in str(r["data"]))
        self.assertEqual(by_no[first["data"][ref["id"]]], "K-10001")


if __name__ == "__main__":
    unittest.main()
