#!/usr/bin/env python3
"""Hollerieth Catering & Events GmbH, Regensburg – ein erfundener Caterer (Veranstaltungen, Mittagessen für
Kitas, Schule und Tagespflege, Business-Lunch-Abo, Rahmenverträge für Besprechungs- und Seminarverpflegung;
11 Leute plus Minijobber und Leihpersonal) mit zwei Lieferungen, so wie die Ablage eines solchen Betriebs
wirklich aussieht.

    python3 catering.py OUTDIR  ->  OUTDIR/lieferung1/   Stand Fr 18.09.2026 (Daten bis Do 17.09.)
                                    OUTDIR/lieferung2/   Stand Fr 30.10.2026 (Daten bis Do 29.10.), kompletter
                                                         neuer Ordnerstand
                                    OUTDIR/truth.json    Format des Prüfstands (bench/one.py)
                                    OUTDIR/answer_key.json  Veranstaltungen, Belege, Geldbewegungen, Lieferscheine

Fester Seed; nur Standardbibliothek + openpyxl + reportlab (+ pypdf für ZUGFeRD).

Die Dateien: Faktura/Rechnungsliste_Export_<Stand>.csv (Export des Rechnungsprogramms, UTF-8, alle Belege 2026,
Status und „Offener Betrag“ so, wie das Büro Zahlungen abgehakt hat: bis 31.08. bzw. 15.10.; Entwürfe ohne Nummer),
Faktura/Kontakte_Export_<Stand>.csv (cp1252), Faktura/Artikel_Export_<Stand>.csv, Veranstaltungen_2026.xlsx
(Liste von Tobias Hollerieth, Stand 11.09. bzw. 23.10., Monatszeilen, gestrichene Stornos, Blatt „Lieferscheine
RV + Abo“, in L2 Blatt „2027“), Angebote_2025-2026.xlsx, Vertraege_Verpflegung.xlsx, Kueche/Essenszahlen_2026.xlsx
(ein Blatt je Monat: Tage × Einrichtungen, „zu“, „Ferien“, Summenzeile), Kassenbuch_2026.xlsx (ein Blatt je Monat),
Bank/camt053_<IBAN>_<JJJJ-MM>.zip (Tagesauszüge camt.053.001.08, ein XML je Buchungstag, Lohn als Sammler mit
Einzelposten; KEIN CSV), Rechnungen_Ausgang/<JJJJ-MM>/*.pdf (einige eigene Rechnungen, Gutschriften, Stornos),
Rechnungen_Ausgang/XRechnung_Stadt/*.xml (eigene XRechnungen, UBL), Eingangsrechnungen/<JJJJ-MM>/ (ab August:
ZUGFeRD-PDF, PDF, XRechnung-UBL), Mails/*.eml (Ein- und Ausgang, mit PDF- und XML-Anhängen).

ANSWER KEY – die Fallen (alle Zahlen stehen in truth.json / answer_key.json; check_traps() prüft die hier genannten)
Bank und Buch
K01 Die Bank gibt es nur als camt.053: ein XML je Buchungstag, gezippt je Monat. L2 bringt Juni–August
    byte-gleich wieder, September vollständig unter demselben Namen (L1 bis 17.09.), Oktober neu – keine Buchung
    doppelt. Ohne die Bank bleibt alles offen, was das Büro noch nicht abgehakt hat.
K02 Kontowechsel zum 01.06.2026: Rechnungen bis Mai wurden aufs alte Sparkassenkonto bezahlt (nur „Bezahlt am“ im
    Export) -> bezahlt. Ausnahme R-2026-0043 Dengler (50. Geburtstag, Rest bestritten): „teilbezahlt“, Bezahlt am
    20.05. (1.000,00 aufs alte Konto), Offener Betrag 3.464,03 -> L1 und L2 offen 3.464,03, überfällig.
K03 Saldoübertrag vom alten Konto am 01.06. (61.288,17, Absender die Firma selbst) – kein Geld ohne Rechnung.
K04 Das Buch hinkt: R-2026-0088 (Eventwerk, bezahlt 10.09.), R-2026-0092/0093 (Stadt, 01.09.), R-2026-0103
    (Schmid, 11.09.) u. a. stehen in L1 „überfällig“ – bezahlt laut Bank. Fällig heute ist nicht überfällig
    (R-2026-0134 Kellner, fällig 30.10.).
Kitas, Schule, Tagespflege (MV-01…MV-06, Monatsrechnung am 2. Bankarbeitstag nach Essenszahlen)
K05 Stadt Regensburg (MV-01 + MV-02, Leitweg-ID): beide Monatsrechnungen in einer Überweisung nach ~4 Wochen
    („KASSENZ. 0542110048 RE R-2026-0104 R-2026-0105“). L1: R-2026-0104/0105 offen, nicht fällig.
K06 Eigene XRechnungen an die Stadt (Rechnungen_Ausgang/XRechnung_Stadt) sind Ausgangs-, keine Eingangsrechnungen.
    R-2026-0122 (MV-02 September) gibt es zweimal (abgewiesen wegen falscher Leitweg-ID, korrigiert am 12.10.) ->
    eine Rechnung, L2 offen 3.743,34, nicht überfällig (fällig 01.11.).
K07 Kirchenstiftung (MV-03) zahlt zwei Monate in einer Überweisung ohne Nummer („ESSENSGELD KITA ST. WOLFGANG
    AUGUST + SEPTEMBER“). L1: R-2026-0106 (976,32) offen, überfällig; L2 bezahlt (20.10. mit R-2026-0123).
K08 Elterninitiative Wirbelwind (MV-04) zahlt ab Juni per Dauerauftrag fest 900,00 € ohne Nummer, die Rechnungen
    schwanken (Schließzeit) -> älteste zuerst (§ 366 Abs. 2 BGB). L1: R-2026-0107 offen 250,34, überfällig;
    R-2026-0095 bezahlt (das Buch sagt noch „teilbezahlt 567,08“). L2: R-2026-0124 offen 584,80, überfällig.
K09 Mittagsbetreuung Grundschule (MV-05): Sommerferien 03.08.–14.09., kein Essen, keine Augustrechnung – richtig,
    nicht „nicht berechnet“; Septemberrechnung nur ab 15.09.
K10 Tagespflege (MV-06): Augustrechnung vergessen -> L1 nicht berechnet „MV-06“ (325 Essen). L2 am 06.10.
    nachberechnet (R-2026-0128, Nummer nach der Septemberrechnung R-2026-0126), beide am 21.10. bezahlt.
K11 Preisanpassung zum 01.09. (Kitas 3,95 -> 4,15, St. Wolfgang 3,85 -> 4,05, Schule 4,40 -> 4,60) – die
    Septemberrechnungen stimmen.
K12 Kündigung der Kita St. Wolfgang zum 31.12.2026 per Mail (12.10.) – wartet auf Bestätigung, ändert noch nichts.
Firmen und Rahmenverträge
K13 Lunch-Abo Nordwand (FV-01): Monatsrechnung nachträglich mit den Sonderlieferungen (Lieferscheine), Zahllauf
    jeden zweiten Freitag. L1: R-2026-0101 (fällig 15.09., Zahllauf erst 18.09.) offen, überfällig. Ab 01.10.
    höhere Pauschale (Mail 22.09.). Sonderlieferungen des laufenden Monats sind nicht „nicht berechnet“.
K14 Kanzlei Seidl (FV-02): Dauerauftrag 363,80 € zum 3. ohne Nummer, Rechnung im Voraus -> bezahlt.
K15 Rahmenvertrag Donau Energie Netz (RV-01): Abrufe per Lieferschein, quartalsweise nachträglich; SAP zahlt mit
    der Nummer ohne Bindestriche („4510022817 R20260077“). L1: Q3-Lieferscheine laufen noch – nicht „nicht
    berechnet“. L2: Q3-Rechnung R-2026-0127 offen 2.501,71, nicht fällig.
K16 Rahmenvertrag Bildungszentrum (RV-02): Q3 nie abgerechnet -> L2 nicht berechnet: LS-26-0031, -0032, -0034,
    -0042, -0045, -0047. In L1 nichts (Quartal läuft).
Veranstaltungen (VA-Nummer nach Eingang des Auftrags, Rechnung meist mit „Referenz VA-…“)
K17 Anzahlung ohne Rechnung (Privatkunden 30 %) ist Geld ohne Rechnung, bis die Schlussrechnung sie abzieht.
    L1: 4.040,00 Wenzl/Mayer (01.07., Hochzeit 10.10.), 3.060,00 Kellner (20.07., 17.10.), 1.010,00 Moser
    (08.09., 03.10.). In L2 sind alle drei abgerechnet.
K18 Neue Anzahlungen ohne Rechnung in L2: 1.830,00 Sportfreunde (50 %, Weihnachtsfeier 18.12.), 3.630,00 Zeitler
    (Hochzeit 12.06.2027, VA-27-001, Blatt „2027“).
K19 Schlussrechnung „abzgl. erhaltene Anzahlung“: es zählt der Zahlbetrag (R-2026-0109 Aigner: brutto 8.716,14,
    Zahlbetrag 6.146,14; L1 offen, fällig 19.09.; L2 bezahlt). Anzahlung aufs alte Konto (Lehner 20.05.) ->
    R-2026-0135 L2 offen 7.511,52, nicht fällig.
K20 Anzahlungsrechnung bezahlt, Schlussrechnung zieht sie ab: Trachtenverein R-2026-0084 / R-2026-0114,
    Pflaum R-2026-0120 (13.10.).
K21 Anzahlungsrechnung unbezahlt: R-2026-0131 Autohaus Lindbüchl (Weihnachtsfeier 12.12.) L2 offen 3.055,60,
    überfällig (fällig 26.10.); Mail 28.10.: 150 statt 120 Personen.
K22 Ratenzahlung ohne Absprache: R-2026-0070 Hofbauer, 4.000 + 2.000 (+ 1.000 in L2) -> L1 offen 2.548,46,
    L2 1.548,46, überfällig.
K23 Minderung akzeptiert: Wimmeder, Gutschrift GS-2026-002 (-606,69), Zahlung „RE R-2026-0085 ABZGL.
    GUTSCHRIFT GS-2026-002“ -> erledigt.
K24 Minderung einseitig: Lindbüchl zahlt 85 % von R-2026-0089, abgelehnt -> L1 und L2 offen 1.640,70, überfällig.
K25 Stornokosten 50 % (Absage 26 Tage vorher): R-2026-0108 Pongratz, abzgl. Anzahlung 1.420,00 -> 949,24; L1
    offen und überfällig (fällig 12.09.), L2 bezahlt.
K26 Kostenfreie Absage (Ebner, mehr als 60 Tage vorher): Anzahlung 1.200,00 am 19.06. zurück -> nichts.
K27 Falscher Rechnungsempfänger: R-2026-0080 an Dr. Seidl privat, Storno GS-2026-001, neu R-2026-0083 an die
    Kanzlei, bezahlt -> nichts offen; Kanzlei (10011) und Dr. Seidl (10037) bleiben zwei Kontakte.
K28 Ein Fremder zahlt: R-2026-0064 (Hochzeit Bauer/Schreiner) zahlt der Schwiegervater ohne Nummer -> bezahlt.
    Die Kundin heißt nach der Hochzeit Schreiner: 10010 und 10030 sind dieselbe Person (Mail, Telefon, Adresse).
K29 Bar im Büro (Kassenbuch): R-2026-0099 Huber, R-2026-0116 Reisinger -> bezahlt; die Bareinzahlungen
    2.800,00 (17.08.) und 3.700,00 (08.10.) sind kein Geld ohne Rechnung.
K30 Überzahlung: Moser zahlt R-2026-0129 voll, ohne die Anzahlung abzuziehen -> L2 Geld ohne Rechnung 1.010,00
    (Mail 26.10. bittet um Erstattung).
K31 Rechnung nennt nur die Angebotsnummer: R-2026-0130 Klier (Referenz AN-26-038, Liste noch „bestätigt“),
    Zahlung „Angebot AN-26-038“ -> bezahlt; VA-26-049 ist berechnet.
K32 Fertig, nicht berechnet: L1 VA-26-042 (Praxisjubiläum 11.09.), VA-26-045 (Richtfest 16.09.); L2 VA-26-050
    (Herbstempfang 22.10.: im Export nur ein Entwurf ohne Nummer – kein offener Posten) und VA-26-052
    (Firmenjubiläum 23.10.).
K33 Liste nicht nachgetragen, aber berechnet: VA-26-018 Berger (L1, R-2026-0111), VA-26-006 Lehner (L2,
    R-2026-0135); kostenlos: Probeessen VA-26-039; Trauerkaffees (Bestattungen Hirsch, telefonisch) haben
    Rechnungen ohne VA-Nummer – keiner davon „nicht berechnet“.
K34 Bestattungen Hirsch zahlt am 10. des Folgemonats alle Trauerkaffees zusammen -> R-2026-0110 (L1) und
    R-2026-0133 (L2) offen, nicht fällig.
K35 Die Eventagentur zahlt nach eigenen Einkaufsbedingungen nach 45 Tagen: R-2026-0115 L2 offen 9.702,73,
    überfällig (die Rechnung sagt 14 Tage).
K36 Elternbeirat Kita Sonnenblume (10038) ist nicht die Stadt, der Träger der Kita.
Lieferanten, Mails, Sonstiges
K37 Eingangsrechnungen: Gastro-Großmarkt wöchentlich als ZUGFeRD (Lastschrift nach 7 Tagen), Metzgerei Eibl (PDF,
    2 % Skonto 10 Tage), Pöschl (Pfand/Leergut), ServicePlus (XRechnung UBL), Wäscherei (Lastschrift), Zeltverleih.
    Offen L1: GD-26-018248 (Lastschrift 21.09.), Eibl 2026-1031 (Skonto bis 24.09.); L2: GD-26-018254
    (Lastschrift 02.11.), Eibl 2026-1036 (Skonto bis 05.11.).
K38 Eibl 2026-1029 kommt auch als Mail-Anhang (einmal bezahlen); Eibls Zahlungserinnerung (23.10.) zu 2026-1032
    ist überholt (bezahlt 21.10.).
K39 Betrug (L2): Mail von gastro-grossmarkt-donau.com (echt: .de), „keine Lastschrift mehr, bitte überweisen“,
    litauische IBAN, „korrigierte“ Rechnung GD-26-018253 (am 26.10. schon per Lastschrift bezahlt).
K40 Geld, das keine Kundenrechnung bezahlt und keine Aufgabe ist: Saldoübertrag, Versicherungsleistung (14.07.,
    2.380,00), USt-Erstattung (18.08., 1.844,30), Bareinzahlungen.
K41 Lohn als Sammler mit Einzelposten, Sozialversicherung, Finanzamt, Pacht, Leasing: Ausgaben ohne Belege im
    Ordner – keine Lieferantenrechnungen zu bezahlen.
K42 Mails, die warten: L1 Anfrage Optik Reindl (14.09.), Ratenbitte Hofbauer (21.08.); L2 Anfrage Obermeier (27.10.),
    Lindbüchl 150 Personen (28.10.), Erstattungsbitte Moser (26.10.), Kündigung St. Wolfgang (12.10.). Ausgehende
    Mails (buero@/info@hollerieth-catering.de) und Anfragen, auf die ein Angebot folgte, warten nicht.
K43 Angebote nachfassen: L1 AN-26-040 (Zeitler, 11.09.); L2 AN-26-041 (Wolf, 21.09.), AN-26-043 (Reindl, 05.10.).
    Nicht: AN-26-032 (abgelaufen), angenommene oder abgelehnte Angebote; „gültig bis“ und der Veranstaltungstermin
    eines angenommenen Angebots sind keine Fristen.
K44 Keine Festpreisprojekte: Schlussrechnungen über dem Angebot (mehr Gäste, Getränke nach Verbrauch) sind kein
    „über Budget“ (over_budget = []).
"""
from __future__ import annotations

import calendar
import hashlib
import io
import json
import random
import re
import sys
import zipfile
from datetime import date, datetime, timedelta, timezone
from email.message import EmailMessage
from email.utils import format_datetime
from pathlib import Path
from xml.sax.saxutils import escape as xesc

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas

SEED = 20261030


def rr(*key) -> random.Random:
    """A random generator that depends only on the key (stable when other parts change)."""
    h = hashlib.sha256(("|".join(map(str, key)) + f"|{SEED}").encode()).hexdigest()
    return random.Random(int(h[:16], 16))


def D(s: str) -> date:
    return date.fromisoformat(s)


# ------------------------------------------------------------------ calendar (Bayern, Regensburg)
HOL = {D(x) for x in ("2025-12-25", "2025-12-26", "2026-01-01", "2026-01-06", "2026-04-03", "2026-04-06",
                      "2026-05-01", "2026-05-14", "2026-05-25", "2026-06-04", "2026-08-15", "2026-10-03",
                      "2026-11-01", "2026-12-25", "2026-12-26", "2027-01-01", "2027-01-06")}
SCHOOL_HOL = [(D("2025-12-22"), D("2026-01-05")), (D("2026-02-16"), D("2026-02-20")), (D("2026-03-30"), D("2026-04-10")),
              (D("2026-05-26"), D("2026-06-05")), (D("2026-08-03"), D("2026-09-14")), (D("2026-11-02"), D("2026-11-06")),
              (D("2026-12-24"), D("2027-01-08"))]
MONATE = ["Januar", "Februar", "März", "April", "Mai", "Juni", "Juli", "August", "September", "Oktober",
          "November", "Dezember"]
MON3 = ["Jan", "Feb", "Mär", "Apr", "Mai", "Jun", "Jul", "Aug", "Sep", "Okt", "Nov", "Dez"]
WT = ["Mo", "Di", "Mi", "Do", "Fr", "Sa", "So"]


def is_bd(d: date) -> bool:
    return d.weekday() < 5 and d not in HOL


def next_bd(d: date) -> date:
    while not is_bd(d):
        d += timedelta(days=1)
    return d


def prev_bd(d: date) -> date:
    while not is_bd(d):
        d -= timedelta(days=1)
    return d


def nth_bd(y: int, m: int, n: int) -> date:
    d, k = date(y, m, 1), 0
    while True:
        if is_bd(d):
            k += 1
            if k == n:
                return d
        d += timedelta(days=1)


def last_bd(y: int, m: int) -> date:
    return prev_bd(date(y, m, calendar.monthrange(y, m)[1]))


def month_end(y: int, m: int) -> date:
    return date(y, m, calendar.monthrange(y, m)[1])


def in_ranges(d: date, ranges) -> bool:
    return any(a <= d <= b for a, b in ranges)


def dde(d: date) -> str:
    return d.strftime("%d.%m.%Y")


def days(n: int) -> timedelta:
    return timedelta(days=n)


# ------------------------------------------------------------------ money (cents)
def vat(net: int, rate: int) -> int:
    s = 1 if net >= 0 else -1
    return s * ((abs(net) * rate + 50) // 100)


def de(c: int) -> str:
    s = f"{abs(c) // 100:,}".replace(",", ".") + f",{abs(c) % 100:02d}"
    return ("-" if c < 0 else "") + s


def xamt(c: int) -> str:
    return f"{c / 100:.2f}"


def eur(c: int) -> float:
    return round(c / 100, 2)


def pct(c: int, p: float) -> int:
    return int(round(c * p / 100))


# ------------------------------------------------------------------ identifiers
def iban_de(blz: str, kto: str) -> str:
    bban = blz + str(kto).zfill(10)
    chk = 98 - int(bban + "131400") % 97
    return f"DE{chk:02d}{bban}"


def iban_lt(bank: str, acct: str) -> str:
    bban = bank + acct
    chk = 98 - int(bban + "212900") % 97
    return f"LT{chk:02d}{bban}"


def fmt_iban(i: str) -> str:
    return " ".join(i[k:k + 4] for k in range(0, len(i), 4))


def creditor_id(nat: str) -> str:
    chk = 98 - int(nat + "131400") % 97
    return f"DE{chk:02d}ZZZ{nat}"


def vat_id(base8: str) -> str:
    """DE + 8 digits + check digit (ISO 7064 MOD 11,10), as the Bundeszentralamt issues them."""
    p = 10
    for ch in base8:
        s = (int(ch) + p) % 10 or 10
        p = (2 * s) % 11
    c = 11 - p
    return f"DE{base8}{0 if c == 10 else c}"


BICS = {"75050000": "BYLADEM1RBG", "75090000": "GENODEF1R01", "75020073": "HYVEDEMM447", "75040062": "COBADEFFXXX",
        "75070024": "DEUTDEDB750", "12030000": "BYLADEM1001", "50010517": "INGDDEFFXXX", "75069171": "GENODEF1DTR",
        "76010085": "PBNKDEFFXXX", "75051040": "BYLADEM1SAB", "74251020": "BYLADEM1DEG", "10011001": "NTSBDEB1XXX"}


def acct(*key, blz: str | None = None) -> tuple[str, str]:
    r = rr("acct", *key)
    b = blz or r.choice(["75050000", "75050000", "75090000", "75090000", "75020073", "75040062", "75070024",
                         "12030000", "50010517", "76010085", "75051040"])
    return iban_de(b, str(r.randint(10 ** 8, 10 ** 10 - 1))), BICS[b]


# ------------------------------------------------------------------ the firm
FIRM = dict(name="Hollerieth Catering & Events GmbH", short="Hollerieth Catering", street="Im Gewerbepark D 14",
            zip="93059", city="Regensburg", phone="0941 4609 310", mail="info@hollerieth-catering.de",
            office="buero@hollerieth-catering.de", web="www.hollerieth-catering.de",
            iban=iban_de("75069171", "10471180"), bic="GENODEF1DTR", bank="Raiffeisenbank Donautal eG",
            vat=vat_id("31528647"), taxno="244/123/45678", hrb="HRB 16274, Amtsgericht Regensburg",
            gf="Katrin Hollerieth, Tobias Hollerieth", old_iban=iban_de("75050000", "760123458"))
CAMT_NAME = "HOLLERIETH CATERING UND EVENTS GMBH"
BANK_START = D("2026-06-01")          # Kontowechsel: das neue Konto bei der Raiffeisenbank läuft ab 01.06.2026
DELIV = [dict(n=1, folder="lieferung1", cut=D("2026-09-17"), today=D("2026-09-18"), check=D("2026-08-31"),
              list_upd=D("2026-09-11")),
         dict(n=2, folder="lieferung2", cut=D("2026-10-29"), today=D("2026-10-30"), check=D("2026-10-15"),
              list_upd=D("2026-10-23"))]
LAST_CUT = DELIV[-1]["cut"]

# ------------------------------------------------------------------ customers (Faktura-Kontakte)
CUSTOMERS: dict[int, dict] = {}


def C(kd, typ, name, street, plz, city, email, phone, term, created, payer=None, bankblz=None, anrede="",
      vor="", nach="", leitweg="", note="", firma=None):
    ib, bic = acct("kd", kd, blz=bankblz)
    CUSTOMERS[kd] = dict(kd=kd, typ=typ, name=name, firma=firma if firma is not None else (name if typ != "Privatperson" else ""),
                         anrede=anrede, vor=vor, nach=nach, street=street, zip=plz, city=city, email=email,
                         phone=phone, term=term, created=D(created), payer=payer or name, iban=ib, bic=bic,
                         leitweg=leitweg, note=note)


C(10001, "Firma", "Bestattungen Hirsch e.K.", "Friedhofweg 4", "93051", "Regensburg", "info@bestattungen-hirsch-rgb.de",
  "0941 79 44 21", 14, "2019-03-11", payer="BESTATTUNGEN HIRSCH E.K.")
C(10002, "Öffentliche Einrichtung", "Stadt Regensburg – Amt für Tagesbetreuung von Kindern", "Postfach 11 06 43", "93019",
  "Regensburg", "kita-verpflegung@regensburg-stadt.de", "0941 507 5512", 30, "2023-07-20",
  payer="STADT REGENSBURG STADTKASSE", bankblz="75050000", leitweg="09362000-0042-35",
  note="Rechnungen nur als XRechnung über das Portal, Leitweg-ID angeben. Kassenzeichen 0542110048.")
C(10003, "Verein", "Kath. Kirchenstiftung St. Wolfgang", "Wolfgangstraße 12", "93051", "Regensburg",
  "kita.st-wolfgang@pfarrei-stwolfgang-rgb.de", "0941 9 24 87", 14, "2022-06-30", payer="KATH.KIRCHENSTIFTUNG ST.WOLFGANG",
  bankblz="75090000", note="Kita St. Wolfgang (Träger) und Pfarrfeste")
C(10004, "Verein", "Trachtenverein D'Donautaler e.V.", "Am Sportplatz 2", "93055", "Regensburg",
  "vorstand@donautaler-trachtenverein.de", "0170 4418 230", 14, "2023-01-16", payer="TRACHTENVEREIN D DONAUTALER E.V.")
C(10005, "Firma", "Donau Energie Netz GmbH", "Netzallee 7", "93055", "Regensburg", "einkauf@donau-energie-netz.de",
  "0941 2804 0", 30, "2024-02-05", payer="DONAU ENERGIE NETZ GMBH", bankblz="75020073",
  note="Rahmenvertrag RV-01, Bestellnr. 4510022817 auf jeder Rechnung")
C(10006, "Verein", "Förderverein Mittagsbetreuung GS Kumpfmühl e.V.", "Kumpfmühler Straße 52", "93051", "Regensburg",
  "foerderverein.gs-kumpfmuehl@schulfreunde-rgb.de", "0941 2 98 64 11", 14, "2024-07-22",
  payer="FOERDERVEREIN MITTAGSBETREUUNG GS KUMPFMUEHL")
C(10007, "Verein", "Elterninitiative Kinderhaus Wirbelwind e.V.", "Prüfeninger Straße 88", "93049", "Regensburg",
  "vorstand@kinderhaus-wirbelwind.de", "0941 3 07 72 19", 14, "2025-01-13", payer="ELTERNINITIATIVE KINDERHAUS WIRBE",
  bankblz="75090000", note="zahlt ab Juni per Dauerauftrag (fester Betrag)")
C(10008, "Firma", "Sozialwerk Donau gGmbH", "Am Hohen Kreuz 12", "93055", "Regensburg", "rechnungen@sozialwerk-donau.de",
  "0941 46 39 70", 14, "2025-03-03", payer="SOZIALWERK DONAU GGMBH", note="Lieferadresse Tagespflege Lebensabend, Bajuwarenstraße 40")
C(10009, "Firma", "Bildungszentrum Ostbayern gGmbH", "Galgenbergstraße 24", "93053", "Regensburg",
  "verwaltung@bz-ostbayern.de", "0941 7 85 30 0", 30, "2025-02-10", payer="BILDUNGSZENTRUM OSTBAYERN GGMBH",
  note="Rahmenvertrag RV-02 Seminarverpflegung")
C(10010, "Privatperson", "Michaela Bauer", "Brunnweg 9", "93080", "Pentling", "michaela.bauer86@webpost.example",
  "0151 2338 4170", 10, "2025-10-02", anrede="Frau", vor="Michaela", nach="Bauer")
C(10011, "Firma", "Kanzlei Dr. Seidl & Partner mbB", "Gesandtenstraße 21", "93047", "Regensburg",
  "kanzlei@seidl-partner-rgb.de", "0941 5 99 60 0", 14, "2025-10-14", payer="KANZLEI DR. SEIDL UND PARTNER MBB",
  bankblz="75070024")
C(10012, "Privatperson", "Verena Hofbauer", "Lindenweg 3", "93138", "Lappersdorf", "verena.hofbauer@donaumail.example",
  "0176 8120 4433", 10, "2025-11-10", payer="HOFBAUER VERENA", anrede="Frau", vor="Verena", nach="Hofbauer")
C(10013, "Firma", "Nordwand Software GmbH", "Lilienthalstraße 9", "93049", "Regensburg", "buchhaltung@nordwand-software.de",
  "0941 8 30 55 0", 14, "2025-11-24", payer="NORDWAND SOFTWARE GMBH", bankblz="75040062",
  note="Business-Lunch-Abo FV-01 ab 01.03.2026, Zahllauf alle 14 Tage freitags")
C(10014, "Firma", "Eventwerk Donau GmbH", "Weißenburgstraße 15", "93055", "Regensburg", "rechnung@eventwerk-donau.de",
  "0941 6 40 21 80", 14, "2025-09-01", payer="EVENTWERK DONAU GMBH")
C(10015, "Privatperson", "Johanna Lehner", "Am Kreuzhof 7", "93152", "Nittendorf", "johanna.lehner@webpost.example",
  "0160 9187 2201", 10, "2025-12-01", anrede="Frau", vor="Johanna", nach="Lehner")
C(10016, "Privatperson", "Renate Sommer", "Ringstraße 30", "93073", "Neutraubling", "r.sommer@donaumail.example",
  "09401 5 23 18", 10, "2026-01-12", anrede="Frau", vor="Renate", nach="Sommer")
C(10017, "Verein", "Sportfreunde Ziegetsdorf e.V.", "Ziegetsdorfer Straße 101", "93051", "Regensburg",
  "kasse@sf-ziegetsdorf.de", "0941 9 42 08 65", 14, "2026-02-12", payer="SPORTFREUNDE ZIEGETSDORF E.V.")
C(10018, "Privatperson", "Lena Aigner", "Kirchweg 5", "93161", "Sinzing", "lena.aigner@webpost.example",
  "0152 0441 9087", 10, "2026-02-02", payer="AIGNER LENA", anrede="Frau", vor="Lena", nach="Aigner")
C(10019, "Firma", "Spedition Wimmeder GmbH & Co. KG", "Am Güterbahnhof 18", "93073", "Neutraubling",
  "buchhaltung@wimmeder-logistik.de", "09401 9 33 0", 14, "2026-02-20", payer="SPEDITION WIMMEDER GMBH U. CO. KG")
C(10020, "Privatperson", "Stefan Wolf", "Gartenstraße 3", "93080", "Pentling", "stefan.wolf.pentling@donaumail.example",
  "0171 5520 6031", 10, "2026-02-26", anrede="Herr", vor="Stefan", nach="Wolf")
C(10021, "Privatperson", "Markus Dengler", "Am Stadl 2", "93053", "Regensburg", "m.dengler@webpost.example",
  "0160 7782 1145", 10, "2026-03-12", anrede="Herr", vor="Markus", nach="Dengler")
C(10022, "Firma", "Autohaus Lindbüchl GmbH", "Straubinger Straße 140", "93055", "Regensburg",
  "buchhaltung@autohaus-lindbuechl.de", "0941 7 08 99 0", 14, "2026-03-16", payer="AUTOHAUS LINDBUECHL GMBH",
  bankblz="75090000")
C(10023, "Privatperson", "Julia Kraus", "Am Anger 11", "93128", "Regenstauf", "julia.kraus@donaumail.example",
  "09402 77 31 40", 10, "2026-03-30", anrede="Frau", vor="Julia", nach="Kraus")
C(10024, "Privatperson", "Thomas Ebert", "Schillerstraße 6", "93049", "Regensburg", "thomas.ebert@webpost.example",
  "0151 4467 0098", 10, "2025-10-30", anrede="Herr", vor="Thomas", nach="Ebert")
C(10025, "Privatperson", "Gertraud Pongratz", "Graßer Straße 19", "93080", "Pentling", "",
  "0941 9 13 25", 10, "2026-06-02", payer="PONGRATZ GERTRAUD", anrede="Frau", vor="Gertraud", nach="Pongratz",
  note="Kontakt über Tochter Monika (Mail monika.pongratz@webpost.example)")
C(10026, "Privatperson", "Daniel Ebner", "Weinweg 23", "93049", "Regensburg", "daniel.ebner@donaumail.example",
  "0176 2210 5546", 10, "2026-05-25", anrede="Herr", vor="Daniel", nach="Ebner")
C(10027, "Privatperson", "Sophie Wenzl", "Holzgartenstraße 44", "93059", "Regensburg", "sophie.wenzl@webpost.example",
  "0157 3390 1288", 10, "2026-06-05", payer="SOPHIE WENZL UND LUKAS MAYER", anrede="Frau", vor="Sophie", nach="Wenzl")
C(10028, "Privatperson", "Andrea Kastner", "Am Weinberg 12", "93105", "Tegernheim", "andrea.kastner@donaumail.example",
  "09403 95 12 07", 10, "2026-06-02", anrede="Frau", vor="Andrea", nach="Kastner")
C(10029, "Firma", "Zahnärzte am Domplatz MVZ GmbH", "Domplatz 3", "93047", "Regensburg", "verwaltung@zahnaerzte-domplatz.de",
  "0941 5 67 01 0", 14, "2026-07-08", payer="ZAHNAERZTE AM DOMPLATZ MVZ GMBH")
C(10030, "Privatperson", "Michaela Schreiner", "Brunnweg 9", "93080", "Pentling", "michaela.bauer86@webpost.example",
  "0151 2338 4170", 10, "2026-06-16", anrede="Frau", vor="Michaela", nach="Schreiner",
  note="neuer Name nach Hochzeit (vorher Bauer)")
C(10031, "Firma", "Wohnbau Regental eG", "Regentalstraße 40", "93059", "Regensburg", "info@wohnbau-regental.de",
  "0941 8 97 64 0", 14, "2026-08-03", payer="WOHNBAU REGENTAL EG")
C(10032, "Privatperson", "Werner Huber", "Dorfstraße 8", "93080", "Pentling-Leoprechting", "",
  "0941 9 46 03", 10, "2026-06-15", anrede="Herr", vor="Werner", nach="Huber")
C(10033, "Privatperson", "Andreas Moser", "Am Brunnenfeld 7", "93083", "Obertraubling", "andreas.moser@webpost.example",
  "0170 9935 1270", 10, "2026-08-20", payer="MOSER ANDREAS", anrede="Herr", vor="Andreas", nach="Moser")
C(10034, "Firma", "Ingenieurbüro Pflaum GmbH", "Hafenstraße 9", "93055", "Regensburg", "office@ib-pflaum.de",
  "0941 7 93 31 0", 14, "2026-09-15", payer="INGENIEURBUERO PFLAUM GMBH")
C(10035, "Privatperson", "Patrick Zeitler", "Sonnenhang 4", "93161", "Sinzing", "patrick.zeitler@donaumail.example",
  "0152 2290 6614", 10, "2026-08-28", payer="ZEITLER PATRICK", anrede="Herr", vor="Patrick", nach="Zeitler")
C(10036, "Firma", "Steuerberatung Klier & Kollegen", "Maximilianstraße 26", "93047", "Regensburg", "kanzlei@stb-klier.de",
  "0941 5 84 72 0", 14, "2026-08-24", payer="STEUERBERATUNG KLIER U. KOLLEGEN")
C(10037, "Privatperson", "Dr. Matthias Seidl", "Am Singrün 15", "93051", "Regensburg", "kanzlei@seidl-partner-rgb.de",
  "0941 5 99 60 0", 10, "2026-05-20", anrede="Herr Dr.", vor="Matthias", nach="Seidl")
C(10038, "Verein", "Elternbeirat Kita Sonnenblume", "Hemauerstraße 20", "93047", "Regensburg",
  "elternbeirat.sonnenblume@donaumail.example", "0176 4400 9812", 14, "2026-05-11", payer="ELTERNBEIRAT KITA SONNENBLUME",
  note="Sommerfest; nicht die Stadt!")
C(10039, "Privatperson", "Birgit Fischer", "Hochweg 22", "93049", "Regensburg", "birgit.fischer@webpost.example",
  "0941 3 29 81", 10, "2026-06-18", anrede="Frau", vor="Birgit", nach="Fischer")
C(10040, "Firma", "Elektro Gschwendtner GmbH", "Gewerbering 5", "93138", "Lappersdorf", "info@elektro-gschwendtner.de",
  "0941 8 20 30 0", 14, "2026-08-27", payer="ELEKTRO GSCHWENDTNER GMBH")
C(10041, "Privatperson", "Anneliese Brunner", "Graßer Straße 2", "93080", "Pentling", "", "0941 9 11 07", 10,
  "2026-08-10", anrede="Frau", vor="Anneliese", nach="Brunner", note="Rechnung zahlt Tochter Claudia Brunner")
C(10042, "Privatperson", "Kathrin Fuchs", "Am Römling 9", "93047", "Regensburg", "kathrin.fuchs@donaumail.example",
  "0171 6634 2208", 10, "2026-01-15", anrede="Frau", vor="Kathrin", nach="Fuchs")
C(10043, "Privatperson", "Simone Schmid", "Wernerwerkstraße 3", "93049", "Regensburg", "simone.schmid@webpost.example",
  "0160 2214 8873", 10, "2026-02-09", anrede="Frau", vor="Simone", nach="Schmid")
C(10044, "Privatperson", "Carina Berger", "Bergstraße 17", "93152", "Nittendorf", "carina.berger@donaumail.example",
  "0157 7781 3324", 10, "2026-03-02", payer="BERGER CARINA", anrede="Frau", vor="Carina", nach="Berger")
C(10045, "Privatperson", "Elena Kellner", "Lindenstraße 41", "93055", "Regensburg", "elena.kellner@webpost.example",
  "0176 6609 4471", 10, "2026-04-10", payer="KELLNER ELENA", anrede="Frau", vor="Elena", nach="Kellner")
C(10046, "Firma", "Optik Reindl GmbH", "Gesandtenstraße 8", "93047", "Regensburg", "info@optik-reindl-rgb.de",
  "0941 5 12 44", 14, "2026-09-21")
C(10047, "Privatperson", "Gabriele Reisinger", "Leoprechtinger Straße 3", "93080", "Pentling", "", "0941 9 40 77", 10,
  "2026-07-20", anrede="Frau", vor="Gabriele", nach="Reisinger")
C(10048, "Privatperson", "Nadine Wolf", "Am Tiergarten 6", "93049", "Regensburg", "nadine.wolf@donaumail.example",
  "0151 5530 9921", 10, "2026-09-07", anrede="Frau", vor="Nadine", nach="Wolf")

THIRD = {  # people who pay for a customer
    "schreiner": ("JOSEF SCHREINER", *acct("third", "schreiner")),
    "brunner": ("CLAUDIA BRUNNER", *acct("third", "brunner")),
    "fischer": ("FISCHER BIRGIT U. MARKUS", *acct("third", "fischer")),
}

# ------------------------------------------------------------------ contracts: Mittagsverpflegung (MV), Firmen (FV), Rahmenverträge (RV)
PRICE_CHANGE = D("2026-09-01")
XMAS = (D("2025-12-24"), D("2026-01-06"))
MV = [
    dict(id="MV-01", name="Kita Sonnenblume", kd=10002, addr="Hemauerstraße 20, 93047 Regensburg", start=D("2023-09-01"),
         base=68, p_old=395, p_new=415, wd=(0, 1, 2, 3, 4),
         closed=[XMAS, (D("2026-05-15"),) * 2, (D("2026-06-05"),) * 2, (D("2026-08-10"), D("2026-08-28")), (D("2026-10-09"),) * 2],
         note="Schließzeit 10.08.–28.08.2026, Teamtag 09.10."),
    dict(id="MV-02", name="Kita Am Stadtpark", kd=10002, addr="Prinzenweg 9, 93047 Regensburg", start=D("2024-01-08"),
         base=52, p_old=395, p_new=415, wd=(0, 1, 2, 3, 4),
         closed=[XMAS, (D("2026-03-20"),) * 2, (D("2026-05-15"),) * 2, (D("2026-06-05"),) * 2, (D("2026-08-24"), D("2026-09-04"))],
         note="Schließzeit 24.08.–04.09.2026"),
    dict(id="MV-03", name="Kita St. Wolfgang", kd=10003, addr="Wolfgangstraße 14, 93051 Regensburg", start=D("2022-09-01"),
         base=44, p_old=385, p_new=405, wd=(0, 1, 2, 3, 4),
         closed=[XMAS, (D("2026-05-15"),) * 2, (D("2026-06-05"),) * 2, (D("2026-08-03"), D("2026-08-21"))],
         note="Schließzeit 03.08.–21.08.2026; Träger zahlt meist zwei Monate zusammen"),
    dict(id="MV-04", name="Kinderhaus Wirbelwind", kd=10007, addr="Prüfeninger Straße 88, 93049 Regensburg",
         start=D("2025-02-01"), base=14, p_old=395, p_new=415, wd=(0, 1, 2, 3, 4),
         closed=[XMAS, (D("2026-05-15"),) * 2, (D("2026-08-17"), D("2026-08-28"))],
         note="ab 06/2026 Dauerauftrag 900,00 € monatlich"),
    dict(id="MV-05", name="Mittagsbetreuung GS Kumpfmühl", kd=10006, addr="Kumpfmühler Straße 52, 93051 Regensburg",
         start=D("2024-09-10"), base=36, p_old=440, p_new=460, wd=(0, 1, 2, 3), closed=SCHOOL_HOL,
         note="nur Mo–Do, nur Schultage (Ferien Bayern)"),
    dict(id="MV-06", name="Tagespflege Lebensabend", kd=10008, addr="Bajuwarenstraße 40, 93053 Regensburg",
         start=D("2025-04-01"), base=17, p_old=690, p_new=690, wd=(0, 1, 2, 3, 4),
         closed=[(D("2025-12-24"),) * 2, (D("2025-12-31"),) * 2, (D("2026-12-24"),) * 2], note="Menü inkl. Dessert"),
]
MVD = {m["id"]: m for m in MV}
FV = [dict(id="FV-01", kd=10013, name="Business-Lunch-Abo Nordwand Software", what="Mittagessen Mo–Do für ca. 25 Mitarbeitende",
           start=D("2026-03-01"), flat=[(D("2026-03-01"), 385000), (D("2026-10-01"), 430000)], billing="monatlich nachträglich",
           pay="Überweisung 14 Tage (Zahllauf)", note="Sonderlieferungen per Lieferschein, mit der Monatsrechnung"),
      dict(id="FV-02", kd=10011, name="Frühstücks- & Obstservice Kanzlei Seidl", what="Obstkorb und Frühstücksplatten 2× pro Woche",
           start=D("2025-11-01"), flat=[(D("2025-11-01"), 34000)], billing="monatlich im Voraus",
           pay="Dauerauftrag 363,80 € zum 3.", note="")]
RV = [dict(id="RV-01", kd=10005, name="Rahmenvertrag Besprechungsverpflegung", start=D("2024-03-01"), billing="quartalsweise nachträglich",
           order="4510022817", food=("Besprechungsverpflegung (Kaffee, Gebäck, Mittagssnack)", 1290), drinks=("Kalt- und Heißgetränke", 280),
           q_dates={(2025, 4): D("2026-01-07"), (2026, 1): D("2026-04-02"), (2026, 2): D("2026-07-02"), (2026, 3): D("2026-10-05")}),
      dict(id="RV-02", kd=10009, name="Rahmenvertrag Seminarverpflegung", start=D("2025-02-15"), billing="quartalsweise nachträglich",
           order="", food=("Seminarverpflegung ganztags (2 Kaffeepausen, Mittagessen)", 2350), drinks=("Kaltgetränke", 290),
           q_dates={(2025, 4): D("2026-01-07"), (2026, 1): D("2026-04-02"), (2026, 2): D("2026-07-03"), (2026, 3): None})]
RVD = {r["id"]: r for r in RV}


def mv_price(m: dict, d: date) -> int:
    return m["p_new"] if d >= PRICE_CHANGE else m["p_old"]


def fv_flat(f: dict, d: date) -> int:
    return [amt for since, amt in f["flat"] if since <= d][-1]


def meals_on(m: dict, d: date) -> int | None:
    """Portions delivered on day d (None = no delivery)."""
    if d < m["start"] or d.weekday() not in m["wd"] or d in HOL or in_ranges(d, m["closed"]):
        return None
    r = rr("meal", m["id"], d)
    f = r.uniform(0.86, 0.98) if m["id"] != "MV-06" else r.uniform(0.80, 1.0)
    if d.weekday() == 4 and m["id"] != "MV-06":
        f *= 0.93
    if (d.month, d.day) in ((5, 15), (6, 5)):  # Brückentage: half the children stay home
        f *= 0.55
    return max(0, int(round(m["base"] * f)))


MEALS: dict[tuple[str, int, int], dict[date, int]] = {}   # (MV, y, m) -> {day: portions}
for m in MV:
    d = D("2025-12-01")
    while d <= D("2026-10-31"):
        n = meals_on(m, d)
        if n is not None:
            MEALS.setdefault((m["id"], d.year, d.month), {})[d] = n
        d += days(1)

# ------------------------------------------------------------------ call-offs (Lieferscheine) under RV and FV-01 extras
LS: list[dict] = []
_RV_DATES = {
    "RV-01": ["2025-10-14", "2025-11-06", "2025-11-25", "2025-12-09", "2026-01-20", "2026-02-11", "2026-02-26", "2026-03-17",
              "2026-03-26", "2026-04-21", "2026-05-07", "2026-05-20", "2026-06-10", "2026-06-23", "2026-06-30", "2026-07-08",
              "2026-07-21", "2026-08-12", "2026-09-08", "2026-09-15", "2026-09-24", "2026-09-29", "2026-10-13", "2026-10-27"],
    "RV-02": ["2025-10-09", "2025-10-23", "2025-11-13", "2025-11-27", "2025-12-04", "2026-01-15", "2026-01-29", "2026-02-05",
              "2026-03-05", "2026-03-12", "2026-03-24", "2026-04-16", "2026-04-23", "2026-05-06", "2026-05-21", "2026-06-11",
              "2026-06-18", "2026-06-25", "2026-07-09", "2026-07-16", "2026-07-23", "2026-09-22", "2026-09-24", "2026-09-29",
              "2026-10-08", "2026-10-15", "2026-10-22"]}
_RV_WHAT = {"RV-01": ["Vorstandssitzung", "Schulung Netzleitstelle", "Aufsichtsratssitzung", "Projektmeeting Umspannwerk Ost",
                      "Betriebsratssitzung", "Kundenbeirat", "Jour fixe Netzplanung"],
            "RV-02": ["Seminar Erste Hilfe", "Lehrgang Ausbildereignung", "Seminar Führung", "Kurs Buchhaltung kompakt",
                      "Workshop Projektmanagement", "Seminar Arbeitssicherheit", "Lehrgang Fachwirt"]}
for rv_id, ds in _RV_DATES.items():
    for i, s in enumerate(ds):
        r = rr("ls", rv_id, s)
        LS.append(dict(date=D(s), kd=RVD[rv_id]["kd"], contract=rv_id, what=r.choice(_RV_WHAT[rv_id]),
                       pers=r.randint(10, 34) if rv_id == "RV-01" else r.randint(12, 26)))
_FV_EXTRAS = [("2026-03-12", "Geburtstagskuchen", 1), ("2026-04-09", "Teamfrühstück", 30), ("2026-05-21", "Obst- & Snackplatten", 2),
              ("2026-06-18", "Getränkepaket Sommerfest", 1), ("2026-07-09", "Teamfrühstück", 28), ("2026-07-30", "Geburtstagskuchen", 1),
              ("2026-08-27", "Obst- & Snackplatten", 1), ("2026-09-03", "Geburtstagskuchen", 1), ("2026-09-10", "Teamfrühstück", 32),
              ("2026-09-24", "Obst- & Snackplatten", 2), ("2026-10-08", "Teamfrühstück", 35), ("2026-10-22", "Geburtstagskuchen", 2)]
EXTRA_PRICE = {"Geburtstagskuchen": (4800, 7, "Stk."), "Teamfrühstück": (980, 7, "Pers."), "Obst- & Snackplatten": (6500, 7, "Platte"),
               "Getränkepaket Sommerfest": (18000, 19, "psch.")}
for s, what, q in _FV_EXTRAS:
    LS.append(dict(date=D(s), kd=10013, contract="FV-01", what=what, pers=q))
LS.sort(key=lambda x: (x["date"], x["contract"]))
_n = {2025: 388, 2026: 0}
for x in LS:
    _n[x["date"].year] += 1
    x["no"] = f"LS-{x['date'].year % 100:02d}-{_n[x['date'].year]:04d}"
    if x["contract"] in RVD:
        rv = RVD[x["contract"]]
        x["lines"] = [(f"{rv['food'][0]} – {x['what']}", x["pers"], "Pers.", rv["food"][1], 7),
                      (rv["drinks"][0], x["pers"], "Pers.", rv["drinks"][1], 19)]
    else:
        p, rate, unit = EXTRA_PRICE[x["what"]]
        x["lines"] = [(x["what"], x["pers"], unit, p, rate)]
    x["net"] = sum(round(q * p) for _, q, _, p, _ in x["lines"])
    x["inv"] = None   # set when billed


def quarter(d: date) -> tuple[int, int]:
    return d.year, (d.month - 1) // 3 + 1


# ------------------------------------------------------------------ events (Veranstaltungen)
MENU = {"finger": ("Fingerfood-Buffet (8 Sorten)", 2650), "bay": ("Bayerisches Buffet", 3450),
        "premium": ("Hochzeitsbuffet Premium", 4900), "menue4": ("4-Gang-Menü", 5800), "bbq": ("BBQ-Buffet", 3150),
        "kaffee": ("Kaffee & Kuchen", 1190), "brotzeit": ("Brotzeitplatten", 1680),
        "trauer": ("Trauerkaffee (Kaffee, Kuchen, belegte Semmeln)", 1450), "xmas": ("Weihnachtsbuffet", 3950),
        "flying": ("Flying Dinner", 4400), "tagung": ("Tagungspauschale ganztags", 3800), "brunch": ("Brunch-Buffet", 2450),
        "grillkids": ("Grillbuffet Kinder & Familien", 1890)}
DRINK = {"pauschale": "Getränkepauschale (Bier, Wein, Softdrinks, Kaffee)", "alkfrei": "Getränkepauschale alkoholfrei",
         "kaffee": "Kaffee-/Tee-Station", "verbrauch": "Getränke nach Verbrauch lt. Aufstellung"}
STAFF = {"S": ("Servicepersonal", 3200), "K": ("Koch/Köchin vor Ort", 4200)}
EVENTS: list[dict] = []


def EV_(key, kd, anlass, when, ort, pers, final, food, drinks=None, staff=(), log=0, rental=(), offer=None, booked=None,
        dep=None, sr=None, sr_kd=None, pays=(), cancel=None, credit=None, wrong=None, draft=None, ref=None, list_status=None,
        free=False, deceased=None, note=""):
    EVENTS.append(dict(key=key, kd=kd, anlass=anlass, date=D(when), ort=ort, pers=pers, final=final, food=list(food),
                       drinks=drinks, staff=list(staff), log=log, rental=list(rental), offer=D(offer) if offer else None,
                       booked=D(booked), dep=dep, sr=D(sr) if sr else None, sr_kd=sr_kd or kd, pays=pays, cancel=cancel,
                       credit=credit, wrong=wrong, draft=D(draft) if draft else None, ref=ref, list_status=list_status,
                       free=free, deceased=deceased, note=note))


def ev_lines(e: dict, persons: int, offer: bool = False) -> list[tuple]:
    out = []
    for f in e["food"]:
        if f[0] == "M":
            t, p = MENU[f[1]]
            out.append((t, persons, "Pers.", p, 7))
        else:
            out.append((f[1], f[2], "Stk." if f[2] > 1 else "psch.", f[3], 7))
    for role, n, h in e["staff"]:
        t, p = STAFF[role]
        out.append((f"{t} ({n} × {h} Std.)", n * h, "Std.", p, 7))
    if e["log"]:
        out.append(("Anlieferung, Aufbau & Abholung inkl. Geschirr und Besteck", 1, "psch.", e["log"], 7))
    if e["drinks"]:
        kind, val = e["drinks"]
        if kind == "verbrauch":
            v = val if not offer else int(round(val * rr("est", e["key"]).uniform(0.85, 1.08) / 1000)) * 1000
            out.append((DRINK[kind] + (" (Schätzung)" if offer else ""), 1, "psch.", v, 19))
        else:
            out.append((DRINK[kind], persons, "Pers.", val, 19))
    for t, q, p in e["rental"]:
        out.append((t, q, "Stk." if q > 1 else "psch.", p, 19))
    return out


def totals(lines) -> dict:
    n7 = sum(round(q * p) for _, q, _, p, r in lines if r == 7)
    n19 = sum(round(q * p) for _, q, _, p, r in lines if r == 19)
    v7, v19 = vat(n7, 7), vat(n19, 19)
    return dict(n7=n7, v7=v7, n19=n19, v19=v19, gross=n7 + v7 + n19 + v19)


# --- January to May (paid on the old account at the Sparkasse unless said otherwise)
EV_("den_neujahr", 10005, "Neujahrsempfang", "2026-01-16", "Foyer Netzallee 7, Regensburg", 120, 118, [("M", "flying")],
    ("verbrauch", 114000), [("S", 4, 6), ("K", 1, 7)], 12000, offer="2025-11-20", booked="2025-11-28", sr="2026-01-20",
    pays=[dict(d="2026-02-18", ch="alt")])
EV_("tracht_fasching", 10004, "Faschingsball", "2026-02-07", "Vereinsheim Am Sportplatz, Regensburg", 180, 200, [("M", "bay")],
    None, [("S", 5, 7), ("K", 1, 8)], 15000, offer="2025-11-27", booked="2025-12-03", sr="2026-02-10",
    pays=[dict(d="2026-02-25", ch="alt")])
EV_("hirsch_jan", 10001, "Trauerkaffee", "2026-01-23", "Pfarrsaal St. Anton, Regensburg", 40, 38, [("M", "trauer")], None,
    [("S", 1, 4)], 6500, booked="2026-01-19", sr="2026-01-26", pays="hirsch", deceased="Hofmeister")
EV_("nordwand_kickoff", 10013, "Kick-off 2026", "2026-03-05", "Nordwand Software, Lilienthalstraße 9", 60, 64, [("M", "finger")],
    ("alkfrei", 950), [("S", 2, 5)], 8500, offer="2026-01-20", booked="2026-01-27", sr="2026-03-09", pays="zahllauf")
EV_("sommer60", 10016, "60. Geburtstag", "2026-02-21", "Pfarrheim Graß", 45, 45, [("M", "bay")], ("pauschale", 2400),
    [("S", 2, 7)], 8500, offer="2026-01-12", booked="2026-01-30", dep=dict(kind="ohne", pct=30, paid="2026-02-06", ch="alt"),
    sr="2026-02-24", pays=[dict(d="2026-03-06", ch="alt")])
EV_("hirsch_mar", 10001, "Trauerkaffee", "2026-03-13", "Pfarrheim Königswiesen", 55, 52, [("M", "trauer")], None,
    [("S", 1, 4)], 6500, booked="2026-03-09", sr="2026-03-16", pays="hirsch", deceased="Grabmeier")
EV_("sfz_jhv", 10017, "Jahreshauptversammlung", "2026-03-20", "Sportheim Ziegetsdorf", 80, 76, [("M", "brotzeit")], None,
    [("S", 1, 5)], 6500, offer="2026-02-12", booked="2026-02-20", sr="2026-03-23", pays=[dict(d="2026-04-08", ch="alt")])
EV_("wolf_kommunion", 10020, "Erstkommunion", "2026-04-12", "privat, Gartenstraße 3, Pentling", 30, 31, [("M", "brunch")],
    ("alkfrei", 950), [("S", 1, 5)], 8500, offer="2026-02-26", booked="2026-03-05",
    dep=dict(kind="ohne", pct=30, paid="2026-03-10", ch="alt"), sr="2026-04-14", pays=[dict(d="2026-04-22", ch="alt")])
EV_("eventwerk_apr", 10014, "Fachtagung (Kunde Pharma)", "2026-04-23", "Eventhalle Donaulände", 90, 86, [("M", "tagung")],
    ("kaffee", 390), [("S", 3, 9)], 12000, offer="2026-03-05", booked="2026-03-12", sr="2026-04-27",
    pays=[dict(d="2026-06-11", vwz="Projekt 26-071 Re. {no}")])
EV_("dengler50", 10021, "50. Geburtstag", "2026-04-25", "Stadl Dengler, Oberisling", 70, 74, [("M", "bay")], ("pauschale", 2400),
    [("S", 3, 8), ("K", 1, 9)], 12000, rental=[("Stehtisch inkl. Husse", 6, 900)], offer="2026-03-12", booked="2026-03-20",
    dep=dict(kind="ohne", pct=30, paid="2026-03-27", ch="alt"), sr="2026-04-28", pays=[dict(d="2026-05-20", amt=100000, ch="alt")],
    note="Kunde reklamiert (Buffet um 21 Uhr leer) und zahlt den Rest nicht")
EV_("kraus_kommunion", 10023, "Erstkommunion", "2026-05-03", "privat, Am Anger 11, Regenstauf", 25, 27, [("M", "brunch")],
    ("alkfrei", 950), [("S", 1, 5)], 6500, offer="2026-03-30", booked="2026-04-10", sr="2026-05-05",
    pays=[dict(d="2026-05-12", ch="alt")])
EV_("ebert", 10024, "Hochzeit Ebert/Lang", "2026-05-16", "Hofgut Eichhofen (Remise)", 100, 96, [("M", "premium")],
    ("verbrauch", 172000), [("S", 6, 9), ("K", 2, 10)], 18000, offer="2025-10-30", booked="2025-11-15",
    dep=dict(kind="ohne", pct=30, paid="2025-11-25", ch="alt"), sr="2026-05-19",
    pays=[dict(d="2026-06-02", vwz="Hochzeit Ebert Lang 16.05. Rechnung {no}")])
EV_("hirsch_may", 10001, "Trauerkaffee", "2026-05-22", "Pfarrsaal St. Anton, Regensburg", 45, 47, [("M", "trauer")], None,
    [("S", 1, 4)], 6500, booked="2026-05-18", sr="2026-05-26", pays="hirsch", deceased="Wiesent")

# --- June to December 2026 (and the first booking for 2027)
EV_("bauer", 10010, "Hochzeit Bauer/Schreiner", "2026-06-13", "Gutshof Irlbach (Scheune)", 90, 95, [("M", "menue4")],
    ("pauschale", 2400), [("S", 6, 9), ("K", 2, 10)], 18000, rental=[("Stehtisch inkl. Husse", 8, 900)], offer="2025-10-02",
    booked="2025-10-12", dep=dict(kind="ohne", pct=30, paid="2026-04-20", ch="alt"), sr="2026-06-17", sr_kd=10030,
    pays=[dict(d="2026-06-29", who="schreiner", vwz="Hochzeit Michaela und Jonas 13.06.2026 Restbetrag")])
EV_("hofbauer", 10012, "Hochzeit Hofbauer/Gruber", "2026-06-27", "Eventstadel Wolkering", 120, 118, [("M", "premium")],
    ("verbrauch", 218040), [("S", 7, 9), ("K", 2, 10)], 18000, rental=[("Stehtisch inkl. Husse", 12, 900)], offer="2025-11-10",
    booked="2025-11-20", dep=dict(kind="ohne", pct=30, paid="2026-06-05", vwz="Anzahlung Hochzeit 27.06. Hofbauer/Gruber"),
    sr="2026-07-01", pays=[dict(d="2026-07-17", amt=400000, vwz="Teilzahlung {no} Hochzeit"),
                           dict(d="2026-08-18", amt=200000, vwz="2. Rate Hochzeit Hofbauer {no}"),
                           dict(d="2026-10-15", amt=100000, vwz="3. Rate {no}")])
EV_("pfarrfest", 10003, "Pfarrfest St. Wolfgang", "2026-06-28", "Pfarrgarten St. Wolfgang", 150, 162,
    [("M", "brotzeit"), ("M", "kaffee")], None, [("S", 2, 6)], 8500, offer="2026-05-04", booked="2026-05-12", sr="2026-06-30",
    pays=[dict(d="2026-07-15", vwz="Pfarrfest Rechnung {no}")])
EV_("hirsch_jun1", 10001, "Trauerkaffee", "2026-06-12", "Pfarrsaal St. Anton, Regensburg", 45, 43, [("M", "trauer")], None,
    [("S", 1, 4)], 6500, booked="2026-06-08", sr="2026-06-15", pays="hirsch", deceased="Zollner")
EV_("hirsch_jun2", 10001, "Trauerkaffee", "2026-06-26", "Pfarrheim Königswiesen", 60, 64, [("M", "trauer")], None,
    [("S", 2, 4)], 6500, booked="2026-06-22", sr="2026-06-29", pays="hirsch", deceased="Rappl")
EV_("hirsch_jul1", 10001, "Trauerkaffee", "2026-07-03", "Pfarrheim Graß", 35, 35, [("M", "trauer")], None, [("S", 1, 4)], 6500,
    booked="2026-06-30", sr="2026-07-06", pays="hirsch", deceased="Bauernfeind")
EV_("hirsch_jul2", 10001, "Trauerkaffee", "2026-07-23", "Pfarrsaal St. Anton, Regensburg", 50, 48, [("M", "trauer")], None,
    [("S", 1, 4)], 6500, booked="2026-07-20", sr="2026-07-24", pays="hirsch", deceased="Lermer")
EV_("hirsch_aug", 10001, "Trauerkaffee", "2026-08-21", "Pfarrheim Königswiesen", 40, 42, [("M", "trauer")], None,
    [("S", 1, 4)], 6500, booked="2026-08-17", sr="2026-08-24", pays="hirsch", deceased="Schindlbeck")
EV_("hirsch_sep", 10001, "Trauerkaffee", "2026-09-09", "Pfarrsaal St. Anton, Regensburg", 70, 74, [("M", "trauer")], None,
    [("S", 2, 4)], 6500, booked="2026-09-04", sr="2026-09-10", pays="hirsch", deceased="Ettl")
EV_("hirsch_oct", 10001, "Trauerkaffee", "2026-10-16", "Pfarrheim Graß", 55, 51, [("M", "trauer")], None, [("S", 1, 4)], 6500,
    booked="2026-10-12", sr="2026-10-19", pays="hirsch", deceased="Kiendl")
EV_("eventwerk_jun", 10014, "Seminar (Kunde Software)", "2026-06-17", "Eventhalle Donaulände", 60, 58, [("M", "tagung")],
    ("kaffee", 390), [("S", 2, 9)], 12000, offer="2026-05-04", booked="2026-05-13", sr="2026-06-19",
    pays=[dict(d="2026-08-03", vwz="Projekt 26-098 Re. {no}")])
EV_("seidl_sommer", 10011, "Sommerfest Kanzlei", "2026-07-10", "Kanzlei Seidl & Partner, Dachterrasse", 40, 42, [("M", "finger")],
    ("pauschale", 2400), [("S", 2, 6)], 8500, offer="2026-05-20", booked="2026-06-01", sr="2026-07-13", sr_kd=10037,
    wrong=dict(kd=10011, storno="2026-07-20", new="2026-07-20"), pays=[dict(d="2026-07-31", vwz="Rechnung {no} Sommerfest")])
EV_("elternbeirat", 10038, "Sommerfest Kita Sonnenblume", "2026-07-10", "Kita Sonnenblume, Garten", 90, 96, [("M", "grillkids")],
    ("alkfrei", 950), [("S", 2, 5)], 8500, offer="2026-05-11", booked="2026-05-22", sr="2026-07-13",
    pays=[dict(d="2026-07-22", vwz="Sommerfest Elternbeirat Sonnenblume {no}")])
EV_("fuchs", 10042, "Hochzeit Fuchs/Maier", "2026-07-11", "Eventstadel Wolkering", 80, 81,
    [("M", "bay"), ("X", "Hochzeitstorte (3-stöckig)", 1, 42000)], ("pauschale", 2400), [("S", 5, 9), ("K", 1, 10)], 15000,
    offer="2026-01-15", booked="2026-01-28", dep=dict(kind="ohne", pct=30, paid="2026-02-10", ch="alt"), sr="2026-07-14",
    pays=[dict(d="2026-07-24", vwz="{no}")])
EV_("wimmeder", 10019, "Sommerfest Spedition Wimmeder", "2026-07-17", "Betriebshof Wimmeder, Neutraubling", 160, 180,
    [("M", "bbq")], ("verbrauch", 294000), [("S", 6, 7), ("K", 2, 8)], 15000,
    rental=[("Zelt 10 × 20 m inkl. Auf- und Abbau (Fremdleistung)", 1, 145000)], offer="2026-02-20", booked="2026-03-02",
    sr="2026-07-21", credit=dict(d="2026-07-28", pct=10), pays=[dict(d="2026-08-03", vwz="RE {no} ABZGL. GUTSCHRIFT {gs}")])
EV_("wenzl_probe", 10027, "Probeessen Hochzeit Wenzl/Mayer", "2026-07-18", "Küche Gewerbepark D 14", 2, 2,
    [("X", "Probeessen 4-Gang-Menü (kostenlos)", 2, 0)], None, [], 0, booked="2026-07-01", free=True)
EV_("kastner", 10028, "Taufe", "2026-07-19", "privat, Am Weinberg 12, Tegernheim", 35, 33, [("M", "kaffee"), ("M", "brotzeit")],
    ("alkfrei", 950), [("S", 2, 5)], 6500, offer="2026-06-02", booked="2026-06-09", sr="2026-07-21",
    pays=[dict(d="2026-07-30", vwz="Rechnung {no} Taufe")])
EV_("eventwerk_jul", 10014, "Sommerempfang (Kunde Medizintechnik)", "2026-07-23", "Eventhalle Donaulände", 200, 190,
    [("M", "flying")], ("verbrauch", 312000), [("S", 8, 7), ("K", 2, 8)], 18000, offer="2026-04-28", booked="2026-05-08",
    sr="2026-07-27", pays=[dict(d="2026-09-10", vwz="Projekt 26-117 Re. {no}")])
EV_("lindbuechl_sommer", 10022, "Sommerfest Autohaus Lindbüchl", "2026-07-25", "Autohaus Lindbüchl, Ausstellungshalle", 140, 138,
    [("M", "bay")], ("pauschale", 2400), [("S", 5, 8), ("K", 1, 9)], 12000, offer="2026-03-16", booked="2026-03-25",
    sr="2026-07-28", pays=[dict(d="2026-08-10", amt="pct85", vwz="{no} abzgl. 15 % Minderung lt. Schreiben v. 10.08.2026")])
EV_("fischer18", 10039, "18. Geburtstag Laura", "2026-08-01", "privat, Hochweg 22, Regensburg", 40, 40, [("M", "finger")],
    ("alkfrei", 950), [("S", 1, 6)], 6500, offer="2026-06-18", booked="2026-06-25", sr="2026-08-04",
    pays=[dict(d="2026-08-12", who="fischer", vwz="{no} Geburtstag Laura")])
EV_("huber60", 10032, "60. Geburtstag", "2026-08-08", "Vereinsheim Leoprechting", 50, 52, [("M", "bay")], ("pauschale", 2400),
    [("S", 2, 7)], 8500, offer="2026-06-15", booked="2026-06-30",
    dep=dict(kind="ohne", pct=30, paid="2026-07-09", vwz="Anzahlung 60. Geburtstag Werner Huber 08.08."), sr="2026-08-11",
    pays=[dict(d="2026-08-14", ch="kasse")])
EV_("ebner40", 10026, "40. Geburtstag", "2026-08-22", "privat, Weinweg 23, Regensburg", 55, 0, [("M", "bbq")], ("pauschale", 2400),
    [("S", 2, 7)], 8500, offer="2026-05-25", booked="2026-06-05",
    dep=dict(kind="ohne", pct=30, paid="2026-06-08", vwz="Anzahlung Geburtstag 22.08. Ebner"),
    cancel=dict(d="2026-06-15", fee=0, refund="2026-06-19"))
EV_("schmid", 10043, "Hochzeit Schmid/Huber", "2026-08-29", "Hofgut Eichhofen (Remise)", 100, 97, [("M", "premium")],
    ("verbrauch", 176000), [("S", 6, 9), ("K", 2, 10)], 18000, offer="2026-02-09", booked="2026-02-23",
    dep=dict(kind="ohne", pct=30, paid="2026-05-28", ch="alt"), sr="2026-09-01",
    pays=[dict(d="2026-09-11", vwz="Rechnung {no} Hochzeit Schmid")])
EV_("aigner", 10018, "Hochzeit Aigner/Brandl", "2026-09-05", "Hofgut Eichhofen (Remise)", 85, 88, [("M", "premium")],
    ("verbrauch", 164000), [("S", 5, 9), ("K", 1, 10)], 15000, offer="2026-02-02", booked="2026-02-14",
    dep=dict(kind="ohne", pct=30, paid="2026-06-22", vwz="Anzahlung Hochzeit 05.09.2026 Aigner"), sr="2026-09-09",
    pays=[dict(d="2026-09-21", vwz="Hochzeit Aigner Brandl Rechnung {no}")])
EV_("mvz", 10029, "Praxisjubiläum 20 Jahre", "2026-09-11", "Zahnärzte am Domplatz, Praxisräume", 70, 66, [("M", "finger")],
    ("pauschale", 2400), [("S", 3, 6)], 8500, offer="2026-07-08", booked="2026-07-15", sr="2026-09-22",
    pays=[dict(d="2026-10-05", vwz="RG {no}")])
EV_("berger", 10044, "Hochzeit Berger/Lechner", "2026-09-12", "Gutshof Irlbach (Scheune)", 90, 92, [("M", "menue4")],
    ("pauschale", 2400), [("S", 6, 9), ("K", 2, 10)], 18000, offer="2026-03-02", booked="2026-03-14",
    dep=dict(kind="ohne", pct=30, paid="2026-06-12", vwz="Anzahlung Hochzeit Berger 12.09."), sr="2026-09-15",
    pays=[dict(d="2026-09-25", vwz="Hochzeit Berger Lechner {no}")])
EV_("wohnbau", 10031, "Richtfest Neubau Burgweinting", "2026-09-16", "Baustelle Burgweinting, Haus C", 60, 63, [("M", "brotzeit")],
    ("verbrauch", 42000), [("S", 1, 4)], 8500, offer="2026-08-03", booked="2026-08-12", sr="2026-09-21",
    pays=[dict(d="2026-10-14", vwz="Rechnung {no} Richtfest")])
EV_("tracht_jub", 10004, "100 Jahre Trachtenverein – Festabend", "2026-09-19", "Festzelt am Sportplatz", 220, 231, [("M", "bay")],
    None, [("S", 8, 8), ("K", 2, 9)], 18000, rental=[("Zelt 15 × 30 m inkl. Auf- und Abbau (Fremdleistung)", 1, 220000)],
    offer="2026-03-30", booked="2026-04-08",
    dep=dict(kind="AR", pct=30, ar="2026-07-20", paid="2026-07-31", vwz="Anzahlung {ar} Festabend 19.09."),
    sr="2026-09-24", pays=[dict(d="2026-10-07", vwz="Schlussrechnung {no} 100 Jahre")])
EV_("eventwerk_sep", 10014, "Firmenevent (Kunde Versicherung)", "2026-09-25", "Eventhalle Donaulände", 150, 146,
    [("M", "finger")], ("pauschale", 2400), [("S", 6, 6)], 15000, offer="2026-07-01", booked="2026-07-10", sr="2026-09-29",
    pays=[dict(d="2026-11-13", vwz="Projekt 26-142 Re. {no}")])
EV_("pongratz", 10025, "70. Geburtstag", "2026-09-26", "Pfarrheim Graß", 60, 0, [("M", "bay")], ("pauschale", 2400),
    [("S", 3, 7)], 8500, offer="2026-06-02", booked="2026-06-10",
    dep=dict(kind="ohne", pct=30, paid="2026-06-16", vwz="Anzahlung 70. Geburtstag G. Pongratz"),
    cancel=dict(d="2026-08-31", fee=50, inv="2026-09-02"), pays=[dict(d="2026-09-28", vwz="Rechnung {no} Storno")])
EV_("reisinger", 10047, "50. Geburtstag", "2026-09-26", "Vereinsheim Leoprechting", 50, 48, [("M", "bay")], ("pauschale", 2400),
    [("S", 2, 7)], 8500, offer="2026-07-20", booked="2026-07-30", sr="2026-09-29", pays=[dict(d="2026-10-05", ch="kasse")])
EV_("brunner75", 10041, "75. Geburtstag", "2026-09-27", "Pfarrheim Graß", 40, 38, [("M", "kaffee"), ("M", "brotzeit")],
    ("alkfrei", 950), [("S", 2, 5)], 6500, offer="2026-08-10", booked="2026-08-19", sr="2026-09-29",
    pays=[dict(d="2026-10-06", who="brunner", vwz="Rechnung {no} Geburtstag Mutter")])
EV_("klier", 10036, "Firmenfeier 30 Jahre Kanzlei", "2026-10-02", "Steuerberatung Klier, Maximilianstraße 26", 35, 36,
    [("M", "finger")], ("pauschale", 2400), [("S", 2, 5)], 6500, offer="2026-08-24", booked="2026-09-01", sr="2026-10-07",
    ref="AN", list_status="bestätigt", pays=[dict(d="2026-10-20", vwz="Angebot {an} Firmenfeier")])
EV_("moser", 10033, "Gartenfest", "2026-10-03", "privat, Am Brunnenfeld 7, Obertraubling", 45, 47, [("M", "bbq")],
    ("pauschale", 2400), [("S", 2, 7)], 8500, offer="2026-08-20", booked="2026-08-28",
    dep=dict(kind="ohne", pct=30, paid="2026-09-08", vwz="Anzahlung Gartenfest 03.10. Moser"), sr="2026-10-06",
    pays=[dict(d="2026-10-15", amt="gross", vwz="Rechnung {no} Gartenfest")])
EV_("wenzl", 10027, "Hochzeit Wenzl/Mayer", "2026-10-10", "Eventstadel Wolkering", 110, 107, [("M", "menue4")],
    ("pauschale", 2400), [("S", 7, 9), ("K", 2, 10)], 18000, rental=[("Heizpilz inkl. Gas", 6, 3500)], offer="2026-06-05",
    booked="2026-06-18", dep=dict(kind="ohne", pct=30, paid="2026-07-01", vwz="Anzahlung Hochzeit 10.10. Wenzl/Mayer"),
    sr="2026-10-14", pays=[dict(d="2026-10-27", vwz="Restzahlung Hochzeit {no}")])
EV_("kellner", 10045, "Hochzeit Kellner/Roth", "2026-10-17", "Eventstadel Wolkering", 95, 98, [("M", "premium")],
    ("pauschale", 2400), [("S", 6, 9), ("K", 1, 10)], 18000, offer="2026-04-10", booked="2026-04-24",
    dep=dict(kind="ohne", pct=30, paid="2026-07-20", vwz="Anzahlung Hochzeit 17.10. Kellner Roth"), sr="2026-10-20",
    pays=[dict(d="2026-11-02", vwz="Rechnung {no}")])
EV_("den_herbst", 10005, "Herbstempfang", "2026-10-22", "Foyer Netzallee 7, Regensburg", 90, 93, [("M", "flying")],
    ("verbrauch", 98000), [("S", 4, 6)], 12000, offer="2026-08-18", booked="2026-09-02", draft="2026-10-27")
EV_("gschwendtner", 10040, "25 Jahre Elektro Gschwendtner", "2026-10-23", "Firmenhalle Gschwendtner, Lappersdorf", 80, 84,
    [("M", "bay")], ("pauschale", 2400), [("S", 3, 7)], 12000, offer="2026-08-27", booked="2026-09-08")
EV_("lehner", 10015, "Hochzeit Lehner/Pichler", "2026-10-24", "Hofgut Eichhofen (Remise)", 100, 102, [("M", "premium")],
    ("verbrauch", 198000), [("S", 6, 9), ("K", 2, 10)], 18000, rental=[("Stehtisch inkl. Husse", 10, 900)], offer="2025-12-01",
    booked="2025-12-15", dep=dict(kind="ohne", pct=30, paid="2026-05-20", ch="alt"), sr="2026-10-28", pays=[])
EV_("pflaum_xmas", 10034, "Weihnachtsfeier", "2026-12-04", "Ingenieurbüro Pflaum, Hafenstraße 9", 45, 45, [("M", "xmas")],
    ("pauschale", 2400), [("S", 2, 6)], 8500, offer="2026-09-15", booked="2026-09-28",
    dep=dict(kind="AR", pct=30, ar="2026-10-01", paid="2026-10-13", vwz="AR {ar} Weihnachtsfeier"))
EV_("lindbuechl_xmas", 10022, "Weihnachtsfeier Autohaus Lindbüchl", "2026-12-12", "Autohaus Lindbüchl, Ausstellungshalle", 120, 120,
    [("M", "xmas")], ("pauschale", 2400), [("S", 5, 7), ("K", 1, 8)], 12000, offer="2026-06-22", booked="2026-07-06",
    dep=dict(kind="AR", pct=30, ar="2026-10-12", paid=None))
EV_("sfz_xmas", 10017, "Weihnachtsfeier", "2026-12-18", "Sportheim Ziegetsdorf", 70, 70, [("M", "xmas")], None, [("S", 3, 6)], 8500,
    offer="2026-10-05", booked="2026-10-14",
    dep=dict(kind="ohne", pct=50, paid="2026-10-22", vwz="Anzahlung Weihnachtsfeier 18.12. Sportfreunde"))
EV_("zeitler", 10035, "Hochzeit Zeitler/Wittmann", "2027-06-12", "Gutshof Irlbach (Scheune)", 120, 120, [("M", "premium")],
    ("verbrauch", 210000), [("S", 7, 9), ("K", 2, 10)], 18000, offer="2026-08-28", booked="2026-10-19",
    dep=dict(kind="ohne", pct=30, paid="2026-10-27", vwz="Anzahlung Hochzeit 12.06.2027 Zeitler"))
EVD = {e["key"]: e for e in EVENTS}

# offers that never became an event
OFFERS_LOST = [
    dict(key="reindl", kd=10046, prospect=None, anlass="Weihnachtsfeier", when=D("2026-12-11"), pers=40, food="xmas",
         drinks=("pauschale", 2400), staff=[("S", 2, 6)], log=6500, date=D("2026-09-21"), status="offen", follow=D("2026-10-05")),
    dict(key="nwolf", kd=10048, prospect=None, anlass="Hochzeit Wolf/Bachmeier", when=D("2027-07-03"), pers=130, food="premium",
         drinks=("verbrauch", 230000), staff=[("S", 7, 9), ("K", 2, 10)], log=18000, date=D("2026-09-07"), status="offen",
         follow=D("2026-09-21")),
    dict(key="haslbeck", kd=None, prospect="Josef Haslbeck (Anfrage telefonisch)", anlass="65. Geburtstag", when=D("2026-08-15"),
         pers=60, food="bay", drinks=("pauschale", 2400), staff=[("S", 3, 7)], log=8500, date=D("2026-06-25"), status="abgelehnt",
         follow=None, why="zu teuer, feiert im Gasthaus"),
    dict(key="kollmer", kd=None, prospect="Malerei Kollmer GmbH, Frau Kollmer", anlass="Betriebsfeier", when=D("2026-09-18"),
         pers=25, food="bbq", drinks=("pauschale", 2400), staff=[("S", 1, 6)], log=6500, date=D("2026-07-14"), status="abgelaufen",
         follow=D("2026-07-28"), why="keine Rückmeldung"),
    dict(key="sfz_sommer", kd=10017, prospect=None, anlass="Sommerfest 2026", when=D("2026-07-25"), pers=150, food="bbq",
         drinks=None, staff=[("S", 3, 6)], log=12000, date=D("2026-04-20"), status="abgelehnt", follow=None,
         why="Verein grillt selbst"),
]

# ------------------------------------------------------------------ customer documents (Faktura)
INV: list[dict] = []
KIND_LABEL = {"RE": "Rechnung", "SR": "Schlussrechnung", "AR": "Anzahlungsrechnung", "STK": "Rechnung", "GS": "Gutschrift",
              "STO": "Stornorechnung"}


def new_inv(kind, d, kd, subject, lines, ref="", sdate="", due_days=None, deduct=0, deduct_text="", src=None,
            applies_to=None, draft=False) -> dict:
    t = totals(lines)
    term = CUSTOMERS[kd]["term"] if due_days is None else due_days
    inv = dict(no=None, kind=kind, date=d, kd=kd, subject=subject, lines=list(lines), ref=ref, sdate=sdate, **t,
               deduct=deduct, deduct_text=deduct_text, pay=t["gross"] - deduct, due=d + days(term), src=src or {},
               applies_to=applies_to, draft=draft, seq=len(INV))
    INV.append(inv)
    return inv


def mv_inv_date(mid: str, y: int, m: int) -> date | None:
    if (mid, y, m) == ("MV-06", 2026, 8):
        return D("2026-10-06")          # vergessen, erst im Oktober nachberechnet
    ny, nm = (y + 1, 1) if m == 12 else (y, m + 1)
    return nth_bd(ny, nm, 2)


def build_contract_invoices():
    for m in MV:
        for (y, mo) in [(2025, 12)] + [(2026, k) for k in range(1, 10)]:
            cnt = MEALS.get((m["id"], y, mo))
            if not cnt:
                continue                      # Ferien: kein Essen, keine Rechnung
            meals = sum(cnt.values())
            price = mv_price(m, date(y, mo, 1))
            d = mv_inv_date(m["id"], y, mo)
            new_inv("RE", d, m["kd"], f"Mittagsverpflegung {m['name']} – Leistungsmonat {mo:02d}/{y}",
                    [(f"Mittagessen {m['name']} {MONATE[mo - 1]} {y} ({len(cnt)} Liefertage)", meals, "Essen", price, 7)],
                    ref=m["id"], sdate=f"{mo:02d}/{y}", src=dict(type="MV", contract=m["id"], y=y, m=mo, meals=meals))
    f1, f2 = FV
    for mo in range(3, 10):                   # FV-01: monthly in arrears, with the month's extra deliveries
        d = nth_bd(2026, mo + 1, 1)
        flat = fv_flat(f1, date(2026, mo, 1))
        lines = [(f"Business-Lunch-Abo Pauschale {MONATE[mo - 1]} 2026 (Mo–Do)", 1, "Monat", flat, 7)]
        extras = [x for x in LS if x["contract"] == "FV-01" and (x["date"].year, x["date"].month) == (2026, mo)]
        for x in extras:
            lines += [(f"Sonderlieferung {x['no']} vom {dde(x['date'])}: {t}", q, u, p, r) for t, q, u, p, r in x["lines"]]
        inv = new_inv("RE", d, f1["kd"], f"Business-Lunch-Abo {mo:02d}/2026" + (" inkl. Sonderlieferungen" if extras else ""),
                      lines, ref="FV-01", sdate=f"{mo:02d}/2026", src=dict(type="FV", contract="FV-01", y=2026, m=mo))
        for x in extras:
            x["inv"] = inv
    for mo in range(1, 11):                   # FV-02: monthly in advance
        d = nth_bd(2026, mo, 1)
        new_inv("RE", d, f2["kd"], f"Frühstücks- & Obstservice {mo:02d}/2026",
                [(f"Frühstücks- & Obstservice {MONATE[mo - 1]} 2026 (Pauschale)", 1, "Monat", fv_flat(f2, d), 7)],
                ref="FV-02", sdate=f"{mo:02d}/2026", src=dict(type="FV", contract="FV-02", y=2026, m=mo))
    for rv in RV:
        for (y, q), d in sorted(rv["q_dates"].items()):
            if d is None:
                continue                      # Q3 des Bildungszentrums: nie abgerechnet
            items = [x for x in LS if x["contract"] == rv["id"] and quarter(x["date"]) == (y, q)]
            lines = []
            for x in items:
                lines += [(f"{x['no']} vom {dde(x['date'])}: {t}", qq, u, p, r) for t, qq, u, p, r in x["lines"]]
            qs, qe = date(y, 3 * q - 2, 1), month_end(y, 3 * q)
            inv = new_inv("RE", d, rv["kd"], f"Quartalsabrechnung Q{q}/{y} gem. {rv['name']} {rv['id']}", lines,
                          ref=f"{rv['id']}" + (f" / Bestellnr. {rv['order']}" if rv["order"] else ""),
                          sdate=f"{qs.strftime('%d.%m.')}–{qe.strftime('%d.%m.%Y')}", due_days=30,
                          src=dict(type="RV", contract=rv["id"], y=y, q=q))
            for x in items:
                x["inv"] = inv


def build_event_invoices():
    for e in EVENTS:
        if e["free"]:
            continue
        ol = ev_lines(e, e["pers"], offer=True)
        e["offer_lines"], e["offer_tot"] = ol, totals(ol)
        e["inv"], e["ar"], e["stk"], e["gs"], e["wrong_inv"], e["sto"], e["draft_inv"] = (None,) * 7
        dep = e["dep"]
        if dep and dep["kind"] == "ohne":
            e["dep_amt"] = int(round(e["offer_tot"]["gross"] * dep["pct"] / 100 / 1000)) * 1000
        if dep and dep["kind"] == "AR":
            n7, n19 = pct(e["offer_tot"]["n7"], dep["pct"]), pct(e["offer_tot"]["n19"], dep["pct"])
            lines = [(f"Anzahlung {dep['pct']} % gem. Auftragsbestätigung – Speisen und Service", 1, "psch.", n7, 7)]
            if n19:
                lines.append((f"Anzahlung {dep['pct']} % gem. Auftragsbestätigung – Getränke und Miete", 1, "psch.", n19, 19))
            e["ar"] = new_inv("AR", D(dep["ar"]), e["kd"], f"Anzahlungsrechnung {dep['pct']} % – {e['anlass']} am {dde(e['date'])}",
                              lines, ref="{va}", sdate=dde(e["date"]), due_days=14, src=dict(type="EV", key=e["key"]))
        if e["cancel"]:
            c = e["cancel"]
            if c["fee"]:
                fee = pct(e["offer_tot"]["gross"], c["fee"])
                net = int(round(fee / 1.19))
                ded = e.get("dep_amt", 0)
                e["stk"] = new_inv("STK", D(c["inv"]), e["kd"], f"Stornierungskosten {e['anlass']} am {dde(e['date'])}",
                                   [(f"Stornierungskosten gem. Ziff. 8 AGB: {c['fee']} % des Auftragswerts "
                                     f"(Angebot {{an}}, {de(e['offer_tot']['gross'])} € brutto), Absage vom {dde(D(c['d']))}",
                                     1, "psch.", net, 19)],
                                   ref="{va}", sdate=dde(e["date"]), deduct=ded,
                                   deduct_text=f"abzgl. erhaltene Anzahlung vom {dde(D(dep['paid']))}" if ded else "",
                                   src=dict(type="EV", key=e["key"]))
            continue
        lines = ev_lines(e, e["final"])
        ded, dtext = 0, ""
        if dep and dep["kind"] == "ohne":
            ded, dtext = e["dep_amt"], f"abzgl. erhaltene Anzahlung vom {dde(D(dep['paid']))}"
        elif dep and dep["kind"] == "AR":
            ded, dtext = e["ar"]["gross"], "abzgl. Anzahlungsrechnung {ar} vom " + dde(e["ar"]["date"])
        kind = "SR" if ded else "RE"
        if e["kd"] == 10001:
            subj, ref = f"Trauerkaffee am {dde(e['date'])} (Trauerfall {e['deceased']})", ""
        else:
            subj, ref = f"{e['anlass']} am {dde(e['date'])}" + (" – Schlussrechnung" if ded else ""), "{va}"
        if e["ref"] == "AN":
            ref = "{an}"
        if e["draft"]:
            e["draft_inv"] = new_inv(kind, e["draft"], e["sr_kd"], subj, lines, ref=ref, sdate=dde(e["date"]), deduct=ded,
                                     deduct_text=dtext, src=dict(type="EV", key=e["key"]), draft=True)
            continue
        if not e["sr"]:
            continue
        inv = new_inv(kind, e["sr"], e["sr_kd"], subj, lines, ref=ref, sdate=dde(e["date"]), deduct=ded, deduct_text=dtext,
                      src=dict(type="EV", key=e["key"]))
        if e["wrong"]:
            w = e["wrong"]
            e["wrong_inv"] = inv
            e["sto"] = new_inv("STO", D(w["storno"]), e["sr_kd"], "Stornorechnung zu {wr} (falscher Rechnungsempfänger)",
                               [(t, -q, u, p, r) for t, q, u, p, r in lines], ref=ref, sdate=dde(e["date"]),
                               applies_to=inv, src=dict(type="EV", key=e["key"]))
            inv = new_inv(kind, D(w["new"]), w["kd"], subj, lines, ref=ref, sdate=dde(e["date"]),
                          src=dict(type="EV", key=e["key"]))
        e["inv"] = inv
        if e["credit"]:
            c = e["credit"]
            food = sum(round(q * p) for t, q, u, p, r in lines if r == 7 and any(t == MENU[f[1]][0] for f in e["food"] if f[0] == "M"))
            e["gs"] = new_inv("GS", D(c["d"]), e["kd"], "Gutschrift zu {inv} (Minderung Speisen)",
                              [(f"Minderung {c['pct']} % auf Speisen lt. Reklamation vom 22.07.2026 (Rechnung {{inv}})", 1,
                                "psch.", -pct(food, c["pct"]), 7)], ref=ref, sdate=dde(e["date"]), applies_to=inv,
                              src=dict(type="EV", key=e["key"]))


def number_all():
    """VA numbers by entry (booking) order per event year, AN numbers by offer date per year, invoice numbers by date."""
    n = {}
    for e in sorted(EVENTS, key=lambda e: (e["booked"], e["date"])):
        y = e["date"].year % 100
        n[y] = n.get(y, 0) + 1
        e["va"] = f"VA-{y:02d}-{n[y]:03d}"
    offers = [("ev", e) for e in EVENTS if e["offer"]] + [("lost", o) for o in OFFERS_LOST]
    n = {25: 79, 26: 0}
    for kind, o in sorted(offers, key=lambda x: (x[1]["offer"] if x[0] == "ev" else x[1]["date"], x[1]["key"])):
        d = o["offer"] if kind == "ev" else o["date"]
        n[d.year % 100] += 1
        o["an"] = f"AN-{d.year % 100:02d}-{n[d.year % 100]:03d}"
    for e in EVENTS:
        e.setdefault("an", "")
    k = 0
    for inv in sorted((i for i in INV if i["kind"] in ("RE", "SR", "AR", "STK") and not i["draft"]), key=lambda i: (i["date"], i["seq"])):
        k += 1
        inv["no"] = f"R-2026-{k:04d}"
    k = 0
    for inv in sorted((i for i in INV if i["kind"] in ("GS", "STO")), key=lambda i: (i["date"], i["seq"])):
        k += 1
        inv["no"] = f"GS-2026-{k:03d}"
    for e in EVENTS:            # placeholders that need the numbers
        for inv in (e.get("inv"), e.get("ar"), e.get("stk"), e.get("gs"), e.get("wrong_inv"), e.get("sto"), e.get("draft_inv")):
            if not inv:
                continue
            fill = dict(va=e["va"], an=e["an"], ar=e["ar"]["no"] if e.get("ar") else "",
                        inv=e["inv"]["no"] if e.get("inv") else "", wr=e["wrong_inv"]["no"] if e.get("wrong_inv") else "")
            inv["ref"] = inv["ref"].format(**fill)
            inv["subject"] = inv["subject"].format(**fill)
            inv["deduct_text"] = inv["deduct_text"].format(**fill)
            inv["lines"] = [(t.format(**fill), q, u, p, r) for t, q, u, p, r in inv["lines"]]


def inv_of(src_type: str, **kw) -> dict:
    hits = [i for i in INV if i["src"].get("type") == src_type and all(i["src"].get(k) == v for k, v in kw.items())]
    assert len(hits) == 1, (src_type, kw, len(hits))
    return hits[0]


def mv_inv(mid: str, m: int, y: int = 2026) -> dict | None:
    hits = [i for i in INV if i["src"].get("type") == "MV" and i["src"]["contract"] == mid and (i["src"]["y"], i["src"]["m"]) == (y, m)]
    return hits[0] if hits else None

# ------------------------------------------------------------------ money: every movement, on the new account (bank),
# the old Sparkasse account (alt, before 01.06.2026, not in the files) or the cash box (kasse)
MONEY: list[dict] = []


def money(d: date, amount: int, **kw) -> dict:
    m = dict(id=len(MONEY), date=d, amount=amount, ch="bank", kind="cust", kd=None, payer="", iban="", bic="", vwz="",
             alloc=[], dep_for=None, over=0, refund_of=None, gvc="166", btext="GUTSCHR. UEBERWEISUNG", e2e="NOTPROVIDED",
             mandate="", ci="", batch=None, note="", ultimate="")
    m.update(kw)
    if m["ch"] == "bank" and m["date"] < BANK_START:
        m["ch"] = "alt"
    MONEY.append(m)
    return m


INTENTS: list[tuple] = []     # (date, order, fn) – payments that need to know what is open on their day
OPEN: dict[int, int] = {}     # invoice seq -> open cents while the money is simulated
ISSUED: set[int] = set()


def intent(d: date, fn, order: int = 0):
    INTENTS.append((d, order, len(INTENTS), fn))


def nocomp(no: str) -> str:
    return no.replace("-", "")


def payer_of(kd: int, third: str | None = None) -> tuple[str, str, str]:
    if third and third != "kd":
        return THIRD[third]
    c = CUSTOMERS[kd]
    return c["payer"].upper() if c["typ"] == "Privatperson" and c["payer"] == c["name"] else c["payer"], c["iban"], c["bic"]


def open_of(kd: int, d: date, pred=lambda i: True) -> list[dict]:
    return sorted((i for i in INV if i["seq"] in ISSUED and i["kd"] == kd and OPEN.get(i["seq"], 0) > 0 and pred(i)),
                  key=lambda i: (i["date"], i["no"]))


def pay_invoices(d: date, invs: list[dict], amount: int | None = None, **kw) -> dict | None:
    """Money for these invoices; amount None = exactly what is open. The rest of a too large amount stays unallocated."""
    want = sum(OPEN.get(i["seq"], 0) for i in invs)
    amount = want if amount is None else amount
    if amount <= 0:
        return None
    rest, alloc = amount, []
    for i in invs:
        c = min(rest, OPEN.get(i["seq"], 0))
        if c > 0:
            alloc.append((i, c))
            OPEN[i["seq"]] -= c
            rest -= c
    return money(d, amount, alloc=alloc, over=rest, **kw)


def build_money():
    """Customer money: how each customer really pays. (Every closure binds its payer: late binding would hand
    all of them the last customer's name.)"""
    def who(kd, third=None):
        name, ib, bic = payer_of(kd, third)
        return dict(kd=kd, payer=name, iban=ib, bic=bic)

    # --- the city: both Kitas of a month in one transfer, ~4 weeks after the invoice, Tue/Thu
    kw = who(10002)
    for mo in [(2025, 12)] + [(2026, k) for k in range(1, 10)]:
        a, b = mv_inv("MV-01", mo[1], mo[0]), mv_inv("MV-02", mo[1], mo[0])
        d = a["date"] + days(26)
        while not (is_bd(d) and d.weekday() in (1, 3)):
            d += days(1)
        invs = [a] if mo == (2026, 9) else [a, b]   # September Am Stadtpark: XRechnung abgewiesen, später bezahlt
        intent(d, lambda d, invs=invs, kw=kw: pay_invoices(
            d, invs, vwz=f"KASSENZ. 0542110048 RE {' '.join(i['no'] for i in invs)} MITTAGSVERPFLEGUNG", **kw))
    # --- the church: two months together, no numbers
    kw = who(10003)
    for months, d, text in [([(2025, 12), (2026, 1)], "2026-02-19", "DEZEMBER/JANUAR"), ([(2026, 2), (2026, 3)], "2026-04-16", "FEBRUAR/MAERZ"),
                            ([(2026, 4)], "2026-05-21", "APRIL 2026"), ([(2026, 5), (2026, 6)], "2026-07-08", "MAI/JUNI 2026"),
                            ([(2026, 7)], "2026-08-14", "JULI 2026"), ([(2026, 8), (2026, 9)], "2026-10-20", "AUGUST + SEPTEMBER")]:
        invs = [mv_inv("MV-03", m, y) for y, m in months]
        intent(D(d), lambda d, invs=invs, text=text, kw=kw: pay_invoices(
            d, invs, vwz=f"ESSENSGELD KITA ST. WOLFGANG {text}", **kw))
    # --- Wirbelwind: exact amounts until April, from June a fixed standing order of 900,00 (oldest first)
    kw = who(10007)
    for mo in [(2025, 12)] + [(2026, k) for k in range(1, 5)]:
        i = mv_inv("MV-04", mo[1], mo[0])
        intent(next_bd(i["date"] + days(10)), lambda d, i=i, kw=kw: pay_invoices(d, [i], vwz=f"Essensgeld {i['no']}", **kw))
    for mo in range(6, 11):
        d = next_bd(date(2026, mo, 5))
        intent(d, lambda d, kw=kw: pay_invoices(d, open_of(10007, d), 90000, vwz="ESSENSGELD KINDERHAUS WIRBELWIND",
                                                gvc="152", btext="DAUERAUFTRAGSGUTSCHR.", **kw))
    # --- Förderverein: with number, ~9 days
    kw = who(10006)
    for i in [x for x in INV if x["src"].get("contract") == "MV-05"]:
        intent(next_bd(i["date"] + days(9)), lambda d, i=i, kw=kw: pay_invoices(d, [i], vwz=f"Mittagsbetreuung {i['no']}", **kw))
    # --- Sozialwerk: with number, ~14 days; the late August invoice together with September
    kw = who(10008)
    for i in [x for x in INV if x["src"].get("contract") == "MV-06"]:
        if (i["src"]["y"], i["src"]["m"]) in ((2026, 8), (2026, 9)):
            continue
        intent(next_bd(i["date"] + days(14)), lambda d, i=i, kw=kw: pay_invoices(
            d, [i], vwz=f"RG {i['no']} TAGESPFLEGE LEBENSABEND", **kw))
    a, b = mv_inv("MV-06", 8), mv_inv("MV-06", 9)
    intent(D("2026-10-21"), lambda d, kw=kw: pay_invoices(d, [a, b], vwz=f"RG {a['no']} UND {b['no']} TAGESPFLEGE", **kw))
    # --- Nordwand: payment run every second Friday (even ISO weeks), everything due within a week
    kw = who(10013)
    d = D("2026-01-02")
    while d <= D("2026-11-30"):
        if d.weekday() == 4 and d.isocalendar()[1] % 2 == 0 and is_bd(d):
            def run(d, kw=kw):
                invs = open_of(10013, d, lambda i: i["due"] <= d + days(7))
                if invs:
                    pay_invoices(d, invs, vwz=f"ZAHLLAUF {d.isoformat()} RE " + " ".join(i["no"] for i in invs), **kw)
            intent(d, run)
        d += days(1)
    # --- Kanzlei: standing order 363,80 on the 3rd for the month's service invoice
    kw = who(10011)
    for mo in range(1, 11):
        i = inv_of("FV", contract="FV-02", m=mo)
        intent(next_bd(date(2026, mo, 3)), lambda d, i=i, kw=kw: pay_invoices(
            d, [i], 36380, vwz="OBST- UND FRUEHSTUECKSSERVICE HOLLERIETH", gvc="152", btext="DAUERAUFTRAGSGUTSCHR.", **kw))
    # --- Donau Energie Netz: SAP pays on the due date, invoice number without hyphens
    kw = who(10005)
    for i in [x for x in INV if x["src"].get("contract") == "RV-01"]:
        intent(next_bd(i["due"]), lambda d, i=i, kw=kw: pay_invoices(
            d, [i], vwz=f"4510022817 {nocomp(i['no'])} HOLLERIETH CATERING", e2e=f"5100{i['date'].strftime('%y%m%d')}{i['seq']:04d}", **kw))
    # --- Bildungszentrum: with number, ~25 days
    kw = who(10009)
    for i in [x for x in INV if x["src"].get("contract") == "RV-02"]:
        intent(next_bd(i["date"] + days(25)), lambda d, i=i, kw=kw: pay_invoices(
            d, [i], vwz=f"Rechnung {i['no']} Seminarverpflegung", **kw))
    # --- Hirsch: on the 10th all invoices of the previous month
    kw = who(10001)
    for mo in range(2, 12):
        d = next_bd(date(2026, mo, 10))

        def run(d, mo=mo, kw=kw):
            invs = open_of(10001, d, lambda i: (i["date"].year, i["date"].month) == (2026, mo - 1))
            if invs:
                pay_invoices(d, invs, vwz="SAMMELZAHLUNG RG " + ", ".join(i["no"] for i in invs), **kw)
        intent(d, run)
    # --- events: deposits, refunds, final payments
    for e in EVENTS:
        if e["free"]:
            continue
        dep = e["dep"]
        kw = who(e["kd"])
        if dep and dep["kind"] == "ohne" and dep.get("paid"):
            e["dep_money"] = money(D(dep["paid"]), e["dep_amt"], vwz=dep.get("vwz", f"Anzahlung {e['anlass']}").format(va=e["va"]),
                                   dep_for=e["key"], ch=dep.get("ch", "bank"), **kw)
            if e["cancel"] and not e["cancel"]["fee"]:
                money(D(e["cancel"]["refund"]), -e["dep_amt"], kind="refund",
                      vwz=f"Rueckerstattung Anzahlung {e['anlass']} {dde(e['date'])} {CUSTOMERS[e['kd']]['nach']}",
                      refund_of=e["dep_money"]["id"], gvc="116", btext="ONLINE-UEBERWEISUNG", **kw)
        if dep and dep["kind"] == "AR" and dep.get("paid"):
            intent(D(dep["paid"]), lambda d, ar=e["ar"], vwz=e["dep"]["vwz"], kw=kw: pay_invoices(
                d, [ar], vwz=vwz.format(ar=ar["no"]), **kw))
        target = e.get("stk") or e.get("inv")
        if not target or isinstance(e["pays"], str):
            continue
        for p in e["pays"]:
            def fn(d, p=p, t=target, e=e, kw=who(e["kd"], p.get("who"))):
                amt = p.get("amt", "rest")
                if amt == "rest":
                    amount = None
                elif amt == "gross":
                    amount = t["gross"]
                elif amt == "pct85":
                    amount = int(round(t["pay"] * 0.85))
                else:
                    amount = amt
                vwz = p.get("vwz", "").format(no=t["no"], gs=e["gs"]["no"] if e.get("gs") else "", an=e["an"], va=e["va"])
                return pay_invoices(d, [t], amount, vwz=vwz, ch=p.get("ch", "bank"), **kw)
            intent(D(p["d"]), fn)


def run_money():
    timeline = [(i["date"], 0, i["seq"], ("issue", i)) for i in INV if not i["draft"]]
    timeline += [(d, 2, k, ("pay", fn)) for d, order, k, fn in INTENTS]
    for d, prio, k, (what, x) in sorted(timeline, key=lambda t: (t[0], t[1], t[2])):
        if what == "issue":
            ISSUED.add(x["seq"])
            if x["kind"] in ("GS", "STO"):
                t = x["applies_to"]
                OPEN[t["seq"]] = OPEN.get(t["seq"], 0) + x["gross"]   # gross is negative
                assert OPEN[t["seq"]] >= 0, (x["no"], t["no"])
            else:
                OPEN[x["seq"]] = x["pay"]
        else:
            x(d)

# ------------------------------------------------------------------ suppliers and the firm's own expenses
SUPPLIERS = {
    "ggm": dict(name="Gastro-Großmarkt Donau GmbH", street="Industriestraße 30", zip="93055", city="Regensburg",
                vat=vat_id("81244719"), mail="rechnung@gastro-grossmarkt-donau.de", phone="0941 6 98 00-0",
                iban=iban_de("75040062", "0384410900"), bic="COBADEFFXXX", ci=creditor_id("00000518842"), kdnr="48812",
                hrb="HRB 9921, Amtsgericht Regensburg", mandate="GGM-48812-01"),
    "eibl": dict(name="Metzgerei Eibl GmbH", street="Hauptstraße 14", zip="93087", city="Alteglofsheim",
                 vat=vat_id("12844057"), mail="buchhaltung@metzgerei-eibl.de", phone="09453 2 31 40",
                 iban=iban_de("75090000", "0102284410"), bic="GENODEF1R01", kdnr="2207", hrb="HRB 4410, Amtsgericht Regensburg"),
    "poeschl": dict(name="Getränke Pöschl KG", street="Straubinger Straße 77", zip="93055", city="Regensburg",
                    vat=vat_id("26633018"), mail="rechnung@getraenke-poeschl.de", phone="0941 7 03 88-0",
                    iban=iban_de("75050000", "0026103377"), bic="BYLADEM1RBG", kdnr="7712", hrb="HRA 2210, Amtsgericht Regensburg"),
    "serviceplus": dict(name="ServicePlus Personal GmbH", street="Bahnhofstraße 12", zip="93047", city="Regensburg",
                        vat=vat_id("29401176"), mail="rechnungen@serviceplus-personal.de", phone="0941 2 08 14 70",
                        iban=iban_de("75070024", "0412290600"), bic="DEUTDEDB750", kdnr="K-3381",
                        hrb="HRB 13307, Amtsgericht Regensburg"),
    "blank": dict(name="Wäscherei Blank e.K.", street="Am Kalkofen 3", zip="93083", city="Obertraubling",
                  vat=vat_id("17725590"), mail="info@waescherei-blank.de", phone="09401 8 12 05",
                  iban=iban_de("75090000", "0007731205"), bic="GENODEF1R01", ci=creditor_id("00000731205"), kdnr="1150",
                  hrb="HRA 1874, Amtsgericht Regensburg", mandate="WB-1150"),
    "pfeffer": dict(name="Zeltverleih Pfeffer GmbH", street="Industriestraße 4", zip="93128", city="Regenstauf",
                    vat=vat_id("30118842"), mail="buero@zeltverleih-pfeffer.de", phone="09402 50 77 0",
                    iban=iban_de("75051040", "0380114422"), bic="BYLADEM1SAB", kdnr="20-118", hrb="HRB 11042, Amtsgericht Regensburg"),
}
FRAUD = dict(mail="buchhaltung@gastro-grossmarkt-donau.com", iban=iban_lt("39990", "88710492216"), bic="XFRDLT21")  # erfunden
SUPINV: list[dict] = []
FILES_FROM = D("2026-08-01")   # seit August legt das Büro die Eingangsrechnungen digital ab


def sup_inv(sup, no, d, lines, due_days, fmt, skonto=None, how="transfer", paid=None, pay_amt=None, note=""):
    t = totals(lines)
    x = dict(sup=sup, no=no, date=d, lines=lines, **t, due=d + days(due_days), fmt=fmt, skonto=skonto,
             skonto_until=d + days(skonto[0]) if skonto else None, how=how, paid=paid, pay_amt=pay_amt, note=note)
    SUPINV.append(x)
    return x


def build_suppliers():
    # Gastro-Großmarkt: weekly collective invoice (ZUGFeRD), SEPA direct debit 7 days later
    d, k = D("2026-06-08"), 18233
    while d <= D("2026-10-26"):
        r = rr("ggm", d)
        kw = (d - days(7)).isocalendar()[1]
        lines = [(f"Lebensmittel lt. Lieferscheinen KW {kw}", 1, "psch.", r.randint(120000, 290000), 7),
                 (f"Non-Food (Reinigung, Verpackung) KW {kw}", 1, "psch.", r.randint(8000, 30000), 19)]
        if r.random() < 0.5:
            lines.append((f"Getränke (Säfte, Softdrinks) KW {kw}", 1, "psch.", r.randint(6000, 25000), 19))
        k += 1
        x = sup_inv("ggm", f"GD-26-{k:06d}", d, lines, 7, "zugferd", how="LS")
        x["paid"] = next_bd(d + days(7))
        x["pay_amt"] = x["gross"]
        d += days(7)
    # Metzgerei Eibl: about every 9 days, 2 % Skonto within 10 days
    d, k, step = D("2026-06-02"), 1018, 0
    while d <= D("2026-10-26"):
        r = rr("eibl", d)
        k += 1
        x = sup_inv("eibl", f"2026-{k}", d, [(f"Fleisch- und Wurstwaren lt. Lieferschein {r.randint(60000, 69999)}", 1, "psch.",
                                             r.randint(52000, 188000), 7)], 30, "pdf", skonto=(10, 2))
        x["paid"], x["pay_amt"] = next_bd(d + days(8)), x["gross"] - pct(x["gross"], 2)
        step += 1
        d += days(7 if step % 2 else 10)
        d = d if d.weekday() < 5 else d + days(7 - d.weekday())
    late = min((x for x in SUPINV if x["sup"] == "eibl" and x["date"] >= D("2026-09-18")), key=lambda x: x["date"])
    late["paid"], late["pay_amt"], late["note"] = D("2026-10-21"), late["gross"], "zu spät, ohne Skonto; Erinnerung kreuzt sich"
    last = max((x for x in SUPINV if x["sup"] == "eibl"), key=lambda x: x["date"])
    last["paid"], last["pay_amt"] = None, None
    # Getränke Pöschl: monthly, Pfand and Leergut on the invoice
    for mo in range(6, 11):
        d = nth_bd(2026, mo, 3)
        r = rr("poeschl", mo)
        pm = mo - 1
        lines = [(f"Getränke lt. Lieferscheinen {MONATE[pm - 1]} 2026", 1, "psch.", r.randint(150000, 420000), 19),
                 ("Pfand Kästen/Flaschen (Ausgabe)", 1, "psch.", r.randint(8000, 20000), 19),
                 ("Leergut-Rücknahme (Pfand)", 1, "psch.", -r.randint(6000, 18000), 19)]
        x = sup_inv("poeschl", f"L-48{300 + mo}", d, lines, 14, "pdf")
        x["paid"], x["pay_amt"] = next_bd(d + days(12)), x["gross"]
    # ServicePlus: monthly XRechnung for service staff
    for mo in range(6, 11):
        d = nth_bd(2026, mo, 2)
        h = rr("sp", mo).randint(140, 330)
        x = sup_inv("serviceplus", f"SP-26-{1100 + mo}", d,
                    [(f"Arbeitnehmerüberlassung Servicekräfte {MONATE[mo - 2]} 2026 lt. Stundennachweis", h, "Std.", 2850, 19)],
                    14, "xrechnung")
        x["paid"], x["pay_amt"] = next_bd(d + days(13)), x["gross"]
    # Wäscherei Blank: monthly, direct debit after 10 days
    for mo in range(6, 11):
        d = nth_bd(2026, mo, 1)
        x = sup_inv("blank", f"WB-2026-{mo + 40:03d}", d,
                    [(f"Mietwäsche und Reinigung Tischwäsche {MONATE[mo - 2]} 2026", 1, "psch.", rr("blank", mo).randint(38000, 62000), 19)],
                    10, "pdf", how="LS")
        x["paid"], x["pay_amt"] = next_bd(d + days(10)), x["gross"]
    # Zeltverleih Pfeffer: tents for two events
    for d, no, text, net, paid in [(D("2026-07-20"), "2026-118", "Zeltmiete 10 × 20 m inkl. Auf-/Abbau, Veranstaltung 17.07.2026 Neutraubling", 118000, D("2026-08-03")),
                                   (D("2026-09-21"), "2026-164", "Zeltmiete 15 × 30 m inkl. Auf-/Abbau, Festabend 19.09.2026 Sportplatz", 189000, D("2026-10-05"))]:
        x = sup_inv("pfeffer", no, d, [(text, 1, "psch.", net, 19)], 14, "pdf")
        x["paid"], x["pay_amt"] = paid, x["gross"]
    for x in SUPINV:
        s = SUPPLIERS[x["sup"]]
        if x["paid"]:
            if x["how"] == "LS":
                money(x["paid"], -x["pay_amt"], kind="sup", payer=s["name"].upper(), iban=s["iban"], bic=s["bic"],
                      vwz=f"RE {x['no']} KD {s['kdnr']}", gvc="105", btext="FOLGELASTSCHRIFT", mandate=s["mandate"], ci=s["ci"],
                      e2e=x["no"])
            else:
                sk = " ABZGL. 2% SKONTO" if x["pay_amt"] != x["gross"] else ""
                money(x["paid"], -x["pay_amt"], kind="sup", payer=s["name"].upper(), iban=s["iban"], bic=s["bic"],
                      vwz=f"RG {x['no']} KD {s['kdnr']}{sk}", gvc="116", btext="ONLINE-UEBERWEISUNG")
            x["money"] = MONEY[-1]


EMPLOYEES = [("Katrin Hollerieth", 340000), ("Tobias Hollerieth", 340000), ("Florian Wagner", 248000), ("Ayse Demir", 231000),
             ("Martin Liebl", 205000), ("Sabine Danzer", 138000), ("Olga Petrenko", 162000), ("Jana Hartl", 115000),
             ("Lisa Gruber", 60300), ("Tim Seitz", 60300), ("Anna Krinner", 54000), ("Lukas Meindl", 60300),
             ("Selin Yilmaz", 48000)]


def build_expenses():
    money(D("2026-06-01"), 6128817, kind="own", payer=CAMT_NAME, iban=FIRM["old_iban"], bic="BYLADEM1RBG",
          vwz="SALDOUEBERTRAG KONTOAUFLOESUNG SPK KTO 760123458", note="Umbuchung vom alten Konto")
    for mo in range(6, 11):
        mm = f"{mo:02d}/2026"
        money(nth_bd(2026, mo, 1), -285000, kind="exp", payer="GEWERBEHOF DONAUPARK GMBH UND CO. KG",
              iban=iban_de("75020073", "0661200345"), bic=BICS["75020073"], vwz=f"PACHT HALLE 3 PRODUKTIONSKUECHE {mm}",
              gvc="117", btext="DAUERAUFTRAG")
        for vno, amt in (("77-412883", 58940), ("77-415102", 61210)):
            money(nth_bd(2026, mo, 2), -amt, kind="exp", payer="AUTOLEASING SUED GMBH", iban=iban_de("76010085", "0904412001"),
                  bic=BICS["76010085"], vwz=f"VERTRAG {vno} RATE {mm}", gvc="105", btext="FOLGELASTSCHRIFT",
                  mandate=f"ALS-{vno}", ci=creditor_id("00000994120"))
        money(next_bd(date(2026, mo, 5)), -124000, kind="exp", payer="ENERGIEVERSORGUNG DONAUSTADT GMBH",
              iban=iban_de("75050000", "0000300441"), bic=BICS["75050000"], vwz=f"ABSCHLAG STROM/GAS {mm} VK 300441872",
              gvc="105", btext="FOLGELASTSCHRIFT", mandate="VK300441872", ci=creditor_id("00000300441"))
        money(next_bd(date(2026, mo, 20)), -8990, kind="exp", payer="NETZKOM BAYERN GMBH", iban=iban_de("75020073", "0015528193"),
              bic=BICS["75020073"], vwz=f"KD 5528193 RECHNUNG {mm}", gvc="105", btext="FOLGELASTSCHRIFT", mandate="NK5528193",
              ci=creditor_id("00000552819"))
        money(next_bd(date(2026, mo, 12)), -rr("fuel", mo).randint(52000, 71000), kind="exp", payer="TANKCARD SUED GMBH",
              iban=iban_de("75040062", "0071044000"), bic=BICS["75040062"], vwz=f"ABRECHNUNG {mo - 1:02d}/2026 KARTE 7044",
              gvc="105", btext="FOLGELASTSCHRIFT", mandate="TC-7044", ci=creditor_id("00000710440"))
        money(next_bd(date(2026, mo, 15)), -3490, kind="exp", payer="RECHNUNGSDIENST ONLINE GMBH",
              iban=iban_de("50010517", "5512340987"), bic=BICS["50010517"], vwz=f"ABO FAKTURA {mm} KD 20991", gvc="105",
              btext="FOLGELASTSCHRIFT", mandate="RDO-20991", ci=creditor_id("00000209910"))
        fa = iban_de("75000000", "0075001500")
        money(next_bd(date(2026, mo, 10)), -rr("lst", mo).randint(160000, 210000), kind="tax", payer="FINANZAMT REGENSBURG",
              iban=fa, bic="MARKDEF1750", vwz=f"LST {mo - 1:02d}/2026 STNR 244/123/45678", gvc="105", btext="FOLGELASTSCHRIFT",
              mandate="FA-24412345678", ci=creditor_id("00000000244"))
        if mo != 8:
            money(next_bd(date(2026, mo, 10)), -rr("ust", mo).randint(250000, 600000), kind="tax", payer="FINANZAMT REGENSBURG",
                  iban=fa, bic="MARKDEF1750", vwz=f"USTVA {mo - 2:02d}/2026 STNR 244/123/45678", gvc="105", btext="FOLGELASTSCHRIFT",
                  mandate="FA-24412345678", ci=creditor_id("00000000244"))
        d3 = last_bd(2026, mo)
        for _ in range(2):
            d3 = prev_bd(d3 - days(1))
        for kk, (who, lo, hi) in enumerate((("GESUNDHEITSKASSE OSTBAYERN", 540000, 570000), ("BKK DONAUTAL", 140000, 160000),
                                            ("MINIJOB-ZENTRALE", 110000, 130000))):
            money(d3, -rr("sv", mo, kk).randint(lo, hi), kind="exp", payer=who, iban=iban_de("75050000", f"0{9000 + kk}00012"),
                  bic=BICS["75050000"], vwz=f"BEITRAG {mm} BETRIEBSNR 81449120", gvc="105", btext="FOLGELASTSCHRIFT",
                  mandate=f"SV-81449120-{kk}", ci=creditor_id(f"0000081449{kk}"))
        items = []
        for name, net in EMPLOYEES:
            amt = net if net > 100000 else min(60300, int(round(net * rr("mini", name, mo).uniform(0.7, 1.0))))
            items.append(dict(name=name.upper(), iban=acct("emp", name)[0], amount=amt))
        money(last_bd(2026, mo), -sum(x["amount"] for x in items), kind="sal", payer="SAMMLER LOHN/GEHALT",
              vwz=f"LOHN/GEHALT {mm}", gvc="191", btext="SAMMEL-UEBERWEISUNG", batch=items)
        money(last_bd(2026, mo), -1290, kind="fee", payer="", vwz=f"ENTGELTABSCHLUSS {mm}", gvc="805", btext="ABSCHLUSS")
    for d in (D("2026-07-01"), D("2026-10-01")):
        money(d, -138420, kind="exp", payer="BAYERISCHE GEWERBEVERSICHERUNG AG", iban=iban_de("75020073", "0003881200"),
              bic=BICS["75020073"], vwz=f"BEITRAG Q{quarter(d)[1]}/2026 BETRIEBSHAFTPFLICHT/INHALT VS 4471-22", gvc="105",
              btext="FOLGELASTSCHRIFT", mandate="BGV-447122", ci=creditor_id("00004471220"))
    money(D("2026-06-30"), -1897410, kind="exp", payer="GROSSKUECHENTECHNIK DANNER GMBH", iban=iban_de("74251020", "0610045501"),
          bic=BICS["74251020"], vwz="RE 2026-0611 KOMBIDAEMPFER 20-1 INKL. MONTAGE", gvc="116", btext="ONLINE-UEBERWEISUNG")
    money(D("2026-07-14"), 238000, kind="other_in", payer="BAYERISCHE GEWERBEVERSICHERUNG AG", iban=iban_de("75020073", "0003881200"),
          bic=BICS["75020073"], vwz="SCHADEN 2026-04471 KUEHLZELLE ERSTATTUNG", note="Versicherungsleistung")
    money(D("2026-08-18"), 184430, kind="other_in", payer="FINANZAMT REGENSBURG", iban=iban_de("75000000", "0075001500"),
          bic="MARKDEF1750", vwz="ERSTATTUNG UST 06/2026 STNR 244/123/45678", note="Steuererstattung")


# ------------------------------------------------------------------ cash box (Kassenbuch)
CASH: list[dict] = []


def build_cash():
    bal = 41235
    rows = [(D("2026-06-03"), "Wochenmarkt: Kräuter, Blüten", -2340), (D("2026-06-18"), "Parkgebühren Veranstaltung", -800),
            (D("2026-07-14"), "Trockeneis Kühlung Sommerfest", -3600), (D("2026-08-28"), "Porto/Briefmarken", -1650),
            (D("2026-09-11"), "Wochenmarkt: Kräuter, Pilze", -1890), (D("2026-10-16"), "Getränke Kleineinkauf", -2460)]
    for m in MONEY:
        if m["ch"] == "kasse":
            i = m["alloc"][0][0]
            rows.append((m["date"], f"Barzahlung {i['no']} {CUSTOMERS[i['kd']]['nach']} ({i['subject'].split(' am ')[0]})", m["amount"]))
    rows.sort(key=lambda r: r[0])
    out, k = [], 0
    for d, text, amt in rows:
        k += 1
        bal += amt
        out.append(dict(date=d, no=f"K-26-{k + 140:03d}", text=text, amount=amt, bal=bal))
        if amt > 0:                                 # the cash goes to the bank a few days later
            dep = (bal - 30000) // 10000 * 10000
            dd = next_bd(d + days(3))
            k += 1
            bal -= dep
            out.append(dict(date=dd, no=f"K-26-{k + 140:03d}", text="Bareinzahlung Raiffeisenbank Donautal", amount=-dep, bal=bal))
            money(dd, dep, kind="own", payer="", vwz=f"BAREINZAHLUNG {dde(dd)} FIL. GEWERBEPARK", gvc="082", btext="BAREINZAHLUNG",
                  note="Bareinzahlung aus der Kasse")
    out.sort(key=lambda r: (r["date"], r["no"]))
    bal = 41235
    for r in out:
        bal += r["amount"]
        r["bal"] = bal
        assert bal >= 0
    CASH.extend(out)


# ------------------------------------------------------------------ truth
def paid_until(inv: dict, cut: date) -> int:
    return sum(c for m in MONEY if m["date"] <= cut for i, c in m["alloc"] if i is inv)


def credited_until(inv: dict, cut: date) -> int:
    return sum(-g["gross"] for g in INV if g["kind"] in ("GS", "STO") and g["applies_to"] is inv and g["date"] <= cut)


def open_at(inv: dict, cut: date) -> int:
    return inv["pay"] - paid_until(inv, cut) - credited_until(inv, cut)


def final_doc(e: dict) -> dict | None:
    """The document that settles an event: the final invoice, or the cancellation-fee invoice."""
    return e.get("stk") or e.get("inv")


def receipts_at(cut: date) -> list[dict]:
    out = []
    for m in MONEY:
        if m["ch"] != "bank" or m["date"] > cut or m["amount"] <= 0 or m["kind"] != "cust":
            continue
        if any(r["refund_of"] == m["id"] and r["date"] <= cut for r in MONEY):
            continue
        if m["dep_for"]:
            e = EVD[m["dep_for"]]
            t = final_doc(e)
            if t and t["date"] <= cut:
                continue                      # the final invoice deducts the deposit
            out.append(dict(amount=m["amount"], date=m["date"], payer=m["payer"], vwz=m["vwz"],
                            note=f"Anzahlung ohne Rechnung für {e['va']} ({e['anlass']} am {dde(e['date'])})"))
        elif m["over"] > 0:
            out.append(dict(amount=m["over"], date=m["date"], payer=m["payer"], vwz=m["vwz"],
                            note="zu viel bezahlt: " + ", ".join(i["no"] for i, _ in m["alloc"]) + " (Anzahlung nicht abgezogen)"))
    return out


def unbilled_at(cut: date) -> list[dict]:
    out = []
    for e in sorted(EVENTS, key=lambda e: e["date"]):
        if e["free"] or e["cancel"] or e["date"] > cut:
            continue
        t = e.get("wrong_inv") or e.get("inv")
        if t is None or t["date"] > cut:
            extra = " (nur ein Entwurf vom " + dde(e["draft"]) + ")" if e["draft"] and e["draft"] <= cut else ""
            out.append(dict(id=e["va"], what=f"{e['anlass']} ({CUSTOMERS[e['kd']]['name']}) am {dde(e['date'])}{extra}",
                            net=totals(ev_lines(e, e["final"]))["n7"] + totals(ev_lines(e, e["final"]))["n19"]))
    for m in MV:
        for (mid, y, mo), cnt in sorted(MEALS.items()):
            if mid != m["id"] or month_end(y, mo) > cut or (y, mo) < (2025, 12):
                continue
            i = mv_inv(mid, mo, y)
            if i is None or i["date"] > cut:
                out.append(dict(id=mid, what=f"{m['name']} – Leistungsmonat {mo:02d}/{y} nicht berechnet ({sum(cnt.values())} Essen)",
                                net=sum(cnt.values()) * mv_price(m, date(y, mo, 1))))
    for rv in RV:
        for (y, q), d in sorted(rv["q_dates"].items()):
            if month_end(y, 3 * q) > cut or (d is not None and d <= cut):
                continue
            for x in LS:
                if x["contract"] == rv["id"] and quarter(x["date"]) == (y, q):
                    out.append(dict(id=x["no"], what=f"{rv['id']} {CUSTOMERS[rv['kd']]['name']}: Lieferschein vom {dde(x['date'])} "
                                                     f"({x['what']}), Quartal Q{q}/{y} nicht abgerechnet", net=x["net"]))
    return out


def truth_for(dv: dict) -> dict:
    cut, today = dv["cut"], dv["today"]
    known = [i for i in INV if not i["draft"] and i["date"] <= cut and i["kind"] in ("RE", "SR", "AR", "STK")]
    op = {}
    for i in known:
        o = open_at(i, cut)
        assert o >= 0, (i["no"], o)
        if o > 0:
            op[i["no"]] = i
    overdue = sorted(no for no, i in op.items() if i["due"] < today)
    rec = receipts_at(cut)
    unb = unbilled_at(cut)
    payables = []
    for x in SUPINV:
        if x["date"] <= cut and (x["paid"] is None or x["paid"] > cut):
            s = SUPPLIERS[x["sup"]]
            payables.append(dict(supplier=s["name"], invoice=x["no"], date=x["date"].isoformat(), gross=eur(x["gross"]),
                                 due=x["due"].isoformat(), how="Lastschrift" if x["how"] == "LS" else "Überweisung",
                                 skonto_until=x["skonto_until"].isoformat() if x["skonto_until"] else None,
                                 skonto_amount=eur(x["gross"] - pct(x["gross"], 2)) if x["skonto"] else None,
                                 file=x["date"] >= FILES_FROM))
    detail = {}
    for no, i in sorted(op.items()):
        paid = paid_until(i, cut)
        detail[no] = dict(kunde=CUSTOMERS[i["kd"]]["name"], datum=i["date"].isoformat(), faellig=i["due"].isoformat(),
                          zahlbetrag=eur(i["pay"]), bezahlt=eur(paid), gutgeschrieben=eur(credited_until(i, cut)),
                          offen=eur(open_at(i, cut)), betreff=i["subject"])
    not_unbilled = []
    for e in EVENTS:
        if e["date"] > cut:
            continue
        if e["free"]:
            not_unbilled.append(dict(id=e["va"], why="Probeessen, kostenlos"))
        elif e["cancel"]:
            not_unbilled.append(dict(id=e["va"], why="storniert" + (" (Stornokosten berechnet)" if e["cancel"]["fee"] else " (kostenfrei)")))
        elif e.get("inv") and e["inv"]["date"] <= cut and (e["list_status"] or e["inv"]["date"] > dv["list_upd"] or e["kd"] == 10001
                                                          or e["ref"] == "AN"):
            why = "berechnet mit " + e["inv"]["no"]
            if e["ref"] == "AN":
                why += " (Rechnung nennt nur die Angebotsnummer " + e["an"] + ")"
            if e["list_status"] or e["inv"]["date"] > dv["list_upd"]:
                why += "; Veranstaltungsliste noch nicht nachgetragen"
            if e["kd"] == 10001:
                why += " (Rechnung ohne VA-Nummer, nur 'Trauerfall …')"
            not_unbilled.append(dict(id=e["va"], why=why))
    for x in LS:
        if x["date"] <= cut and (x["inv"] is None or x["inv"]["date"] > cut) and x["no"] not in {u["id"] for u in unb}:
            not_unbilled.append(dict(id=x["no"], why=f"{x['contract']}: wird {('mit der Monatsrechnung' if x['contract'] == 'FV-01' else 'quartalsweise')} abgerechnet, Zeitraum läuft noch"))
    if cut >= D("2026-08-31"):
        not_unbilled.append(dict(id="MV-05", why="August 2026: Sommerferien, kein Essen, keine Rechnung"))
    return dict(folder=dv["folder"], today=today.isoformat(), cut=cut.isoformat(),
                open={no: eur(open_at(i, cut)) for no, i in sorted(op.items())},
                open_total=eur(sum(open_at(i, cut) for i in op.values())),
                overdue=overdue,
                unbilled=[u["id"] for u in unb],
                over_budget=[],
                receipts=sorted(eur(x["amount"]) for x in rec),
                open_detail=detail,
                unbilled_detail=[dict(u, net=eur(u["net"])) for u in unb],
                not_unbilled=not_unbilled,
                receipts_detail=[dict(amount=eur(x["amount"]), date=x["date"].isoformat(), payer=x["payer"], vwz=x["vwz"],
                                      note=x["note"]) for x in rec],
                payables=payables,
                same_customer=[dict(a=10010, b=10030, why="Michaela Bauer heißt nach der Hochzeit Schreiner: gleiche Mail, "
                                                           "gleiches Telefon, gleiche Adresse – ein Kunde")],
                different_customers=[dict(a=10038, b=10002, why="Elternbeirat Kita Sonnenblume ist nicht die Stadt (Träger der Kita)"),
                                     dict(a=10011, b=10037, why="Kanzlei und Dr. Seidl privat: falscher Rechnungsempfänger, storniert"),
                                     dict(a=10020, b=10048, why="Stefan Wolf und Nadine Wolf: nicht verwandt")],
                not_receipts=[dict(what=m["vwz"], amount=eur(m["amount"]), why=m["note"]) for m in MONEY
                              if m["ch"] == "bank" and m["date"] <= cut and m["amount"] > 0 and m["kind"] in ("own", "other_in")])


# ------------------------------------------------------------------ writers: Faktura exports (CSV)
def q(v) -> str:
    return '"' + str(v).replace('"', '""') + '"'


def csv_bytes(rows: list[list], enc: str = "utf-8-sig") -> bytes:
    return ("\r\n".join(";".join(q("" if v is None else v) for v in r) for r in rows) + "\r\n").encode(enc, errors="replace")


DUNNING = {1: {"dengler50": 2, "lindbuechl_sommer": 1, "hofbauer": 1}, 2: {"dengler50": 2, "lindbuechl_sommer": 2, "hofbauer": 1}}


def book_view(i: dict, dv: dict) -> dict:
    """What the invoicing program knows: payments booked by hand up to dv['check'], its own credit notes."""
    chk, cut = dv["check"], dv["cut"]
    pays = [(m, c) for m in MONEY if m["date"] <= chk for j, c in m["alloc"] if j is i]
    rec = sum(c for _, c in pays)
    cred = credited_until(i, cut)
    left = i["pay"] - rec - cred
    how = ""
    if pays:
        last = max(pays, key=lambda x: x[0]["date"])[0]
        how = "Bar" if last["ch"] == "kasse" else ("Dauerauftrag" if last["gvc"] == "152" else "Überweisung")
    if i["kind"] in ("GS", "STO"):
        status = "verrechnet" if i["kind"] == "GS" else "Storno"
    elif i["draft"]:
        status = "Entwurf"
    elif any(g["kind"] == "STO" and g["applies_to"] is i and g["date"] <= cut for g in INV):
        status, left = "storniert", 0
    elif left <= 0:
        status = "bezahlt"
    elif rec > 0 or cred > 0:
        status = "teilbezahlt"
    else:
        status = "überfällig" if i["due"] < cut else "offen"
    paid_at = max(m["date"] for m, _ in pays) if pays and status in ("bezahlt", "teilbezahlt") else None
    return dict(status=status, left=max(left, 0) if i["kind"] not in ("GS", "STO") else 0, paid_at=paid_at, how=how)


def write_rechnungsliste(path: Path, dv: dict):
    cut = dv["cut"]
    heads = ["Belegnummer", "Belegart", "Belegdatum", "Leistungsdatum", "Kundennummer", "Kunde", "Betreff", "Referenz",
             "Netto 7 %", "USt 7 %", "Netto 19 %", "USt 19 %", "Brutto", "abzgl. Anzahlung", "Zahlbetrag", "Offener Betrag",
             "Fällig am", "Status", "Bezahlt am", "Zahlungsart", "Mahnstufe"]
    rows = [heads]
    docs = sorted((i for i in INV if i["date"] <= cut), key=lambda i: (i["draft"], i["date"], i["no"] or ""))
    dun = {EVD[k]["inv"]["seq"]: v for k, v in DUNNING[dv["n"]].items()}
    for i in docs:
        b = book_view(i, dv)
        c = CUSTOMERS[i["kd"]]
        rows.append([i["no"] or "", KIND_LABEL[i["kind"]], dde(i["date"]), i["sdate"], i["kd"], c["name"], i["subject"], i["ref"],
                     de(i["n7"]), de(i["v7"]), de(i["n19"]), de(i["v19"]), de(i["gross"]), de(i["deduct"]) if i["deduct"] else "",
                     de(i["pay"]), de(b["left"]), dde(i["due"]) if i["kind"] not in ("GS", "STO") else "", b["status"],
                     dde(b["paid_at"]) if b["paid_at"] else "", b["how"], dun.get(i["seq"], "") if b["status"] not in ("bezahlt",) else ""])
    path.write_bytes(csv_bytes(rows))


def write_kontakte(path: Path, dv: dict):
    heads = ["Kundennummer", "Typ", "Firmenname", "Anrede", "Vorname", "Nachname", "Straße", "PLZ", "Ort", "Land", "E-Mail",
             "Telefon", "Zahlungsziel (Tage)", "Leitweg-ID", "Notiz", "Angelegt am"]
    rows = [heads]
    for kd, c in sorted(CUSTOMERS.items()):
        if c["created"] > dv["cut"]:
            continue
        rows.append([kd, c["typ"], c["firma"], c["anrede"], c["vor"], c["nach"], c["street"], c["zip"], c["city"], "Deutschland",
                     c["email"], c["phone"], c["term"], c["leitweg"], c["note"], dde(c["created"])])
    path.write_bytes(csv_bytes(rows, "cp1252"))


def write_artikel(path: Path, dv: dict):
    rows = [["Artikelnummer", "Bezeichnung", "Einheit", "Nettopreis", "USt-Satz", "Kategorie"]]
    k = 100
    for key, (t, p) in MENU.items():
        k += 1
        rows.append([f"S-{k}", t, "Person", de(p), "7 %", "Speisen"])
    for key, t in DRINK.items():
        if key == "verbrauch":
            continue
        k += 1
        p = {"pauschale": 2400, "alkfrei": 950, "kaffee": 390}[key]
        rows.append([f"G-{k}", t, "Person", de(p), "19 %", "Getränke"])
    rows += [["P-301", "Servicepersonal", "Stunde", de(3200), "7 %", "Personal"],
             ["P-302", "Koch/Köchin vor Ort", "Stunde", de(4200), "7 %", "Personal"],
             ["L-401", "Anlieferung, Aufbau & Abholung inkl. Geschirr und Besteck", "pauschal", "ab 65,00", "7 %", "Logistik"],
             ["M-501", "Stehtisch inkl. Husse", "Stück", de(900), "19 %", "Miete"],
             ["M-502", "Heizpilz inkl. Gas", "Stück", de(3500), "19 %", "Miete"],
             ["V-601", "Mittagessen Kita (bis 31.08.2026)", "Essen", de(395), "7 %", "Verträge"],
             ["V-602", "Mittagessen Kita (ab 01.09.2026)", "Essen", de(415), "7 %", "Verträge"],
             ["V-603", "Mittagessen Schule (bis 31.08.2026)", "Essen", de(440), "7 %", "Verträge"],
             ["V-604", "Mittagessen Schule (ab 01.09.2026)", "Essen", de(460), "7 %", "Verträge"],
             ["V-605", "Menü Tagespflege inkl. Dessert", "Essen", de(690), "7 %", "Verträge"]]
    for key, (p, rate, unit) in EXTRA_PRICE.items():
        k += 1
        rows.append([f"X-{k}", key, unit, de(p), f"{rate} %", "Sonderlieferung"])
    path.write_bytes(csv_bytes(rows))


# ------------------------------------------------------------------ writers: hand-kept Excel lists
def fixed_props(wb, who="Tobias Hollerieth"):
    wb.properties.creator = who
    wb.properties.created = datetime(2026, 1, 7, 8, 0, 0)
    wb.properties.modified = datetime(2026, 1, 7, 8, 0, 0)


def xl_header(ws, row, heads, widths=None):
    for j, h in enumerate(heads, start=1):
        c = ws.cell(row=row, column=j, value=h)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E8E2D0")
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if widths:
            ws.column_dimensions[c.column_letter].width = widths[j - 1]


def save_wb(wb, path: Path):
    buf = io.BytesIO()
    wb.save(buf)
    path.write_bytes(_fix_zip(buf.getvalue()))


def _fix_zip(data: bytes) -> bytes:
    """Same bytes for the same content: openpyxl stamps the time of saving into the archive."""
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for info in src.infolist():
            zi = zipfile.ZipInfo(info.filename, date_time=(2026, 1, 7, 8, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            data = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                data = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)", rb"\g<1>2026-01-07T08:00:00Z\g<2>", data)
                data = re.sub(rb"(<dcterms:created[^>]*>)[^<]*(</dcterms:created>)", rb"\g<1>2026-01-07T08:00:00Z\g<2>", data)
            z.writestr(zi, data)
    return out.getvalue()


def ev_status(e: dict, upd: date) -> tuple[str, str]:
    """(status, invoice number) as the list says on the day it was last updated."""
    if e["cancel"] and D(e["cancel"]["d"]) <= upd:
        return "storniert", (e["stk"]["no"] if e.get("stk") and e["stk"]["date"] <= upd else "")
    if e["date"] > upd:
        return "bestätigt", ""
    if e["free"]:
        return "durchgeführt", ""
    if e["list_status"]:
        return e["list_status"], ""
    t = e.get("inv")
    if t and t["date"] <= upd:
        return "abgerechnet", t["no"]
    return "durchgeführt", ""


def write_veranstaltungen(path: Path, dv: dict):
    upd, cut = dv["list_upd"], dv["cut"]
    wb = openpyxl.Workbook()
    fixed_props(wb)
    ws = wb.active
    ws.title = "Veranstaltungen 2026"
    ws.cell(row=1, column=1, value="Hollerieth Catering & Events – Veranstaltungen 2026").font = Font(bold=True, size=13)
    ws.cell(row=2, column=1, value=f"Stand: {dde(upd)} (TH)")
    heads = ["VA-Nr.", "Datum", "Tag", "Anlass", "Kunde", "Kd.-Nr.", "Ort / Location", "Pers. Angebot", "Pers. final", "Angebot",
             "Status", "Anzahlung", "Anzahlung erhalten", "Rechnung", "Bemerkung"]
    xl_header(ws, 4, heads, [10, 11, 5, 34, 30, 8, 34, 9, 9, 11, 13, 14, 12, 13, 40])
    r, month = 5, None
    for e in sorted((e for e in EVENTS if e["date"].year == 2026 and e["booked"] <= upd), key=lambda e: (e["date"], e["va"])):
        if e["date"].month != month:
            month = e["date"].month
            ws.cell(row=r, column=1, value=MONATE[month - 1].upper()).font = Font(bold=True)
            r += 1
        st, rno = ev_status(e, upd)
        dep, dep_in = "", ""
        if e["dep"]:
            if e["dep"]["kind"] == "ohne":
                dep = f"{e['dep']['pct']} % = {de(e['dep_amt'])}"
                if e["dep"].get("paid") and D(e["dep"]["paid"]) <= upd:
                    dep_in = dde(D(e["dep"]["paid"]))
            else:
                dep = f"AR {e['ar']['no']}" if e["ar"]["date"] <= upd else f"{e['dep']['pct']} % per AR"
                if e["dep"].get("paid") and D(e["dep"]["paid"]) <= upd:
                    dep_in = dde(D(e["dep"]["paid"]))
        note = e["note"]
        if e["free"]:
            note = "kostenlos (Probeessen für " + EVD["wenzl"]["va"] + ")"
        if e["cancel"] and D(e["cancel"]["d"]) <= upd:
            note = f"abgesagt am {dde(D(e['cancel']['d']))}" + (", Anzahlung zurück" if not e["cancel"]["fee"] else f", Stornokosten {e['cancel']['fee']} %")
        if e["kd"] == 10001:
            note = f"Trauerfall {e['deceased']}"
        if e["key"] == "bauer":
            note = "Rechnung auf Schreiner (neuer Name)"
        kunde = CUSTOMERS[e["kd"]]["name"]
        anl = e["anlass"] if e["kd"] != 10001 else f"Trauerkaffee Fam. {e['deceased']}"
        done = e["date"] <= upd and st != "storniert"
        vals = [e["va"], e["date"], WT[e["date"].weekday()], anl, kunde, e["kd"], e["ort"], e["pers"],
                e["final"] if done else None, e["an"] or "tel.", st, dep, dep_in, rno, note]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=j, value=v)
            if isinstance(v, date):
                c.number_format = "DD.MM.YYYY"
        if st == "storniert":
            for j in range(1, len(vals) + 1):
                ws.cell(row=r, column=j).font = Font(strike=True, color="888888")
        r += 1
    ws.freeze_panes = "A5"
    # call-offs under the framework contracts and the lunch contract
    ws2 = wb.create_sheet("Lieferscheine RV + Abo")
    ws2.cell(row=1, column=1, value="Lieferscheine Rahmenverträge (RV-01, RV-02) und Sonderlieferungen Lunch-Abo (FV-01)").font = Font(bold=True, size=12)
    ws2.cell(row=2, column=1, value="Abrechnung: RV quartalsweise nachträglich, FV-01 mit der Monatsrechnung")
    xl_header(ws2, 4, ["Lieferschein", "Datum", "Kunde", "Vertrag", "Anlass / Inhalt", "Pers. / Menge", "Netto €", "abgerechnet mit"],
              [13, 11, 32, 9, 48, 10, 11, 15])
    r = 5
    for x in LS:
        if x["date"] > cut or x["date"] < D("2026-01-01"):
            continue
        billed = x["inv"]["no"] if x["inv"] and x["inv"]["date"] <= upd else ""
        vals = [x["no"], x["date"], CUSTOMERS[x["kd"]]["name"], x["contract"], x["what"], x["pers"], eur(x["net"]), billed]
        for j, v in enumerate(vals, start=1):
            c = ws2.cell(row=r, column=j, value=v)
            if isinstance(v, date):
                c.number_format = "DD.MM.YYYY"
            if j == 7:
                c.number_format = "#,##0.00"
        r += 1
    later = [e for e in EVENTS if e["date"].year == 2027 and e["booked"] <= upd]
    if later:
        ws3 = wb.create_sheet("2027")
        xl_header(ws3, 1, heads)
        for k, e in enumerate(later, start=2):
            vals = [e["va"], e["date"], WT[e["date"].weekday()], e["anlass"], CUSTOMERS[e["kd"]]["name"], e["kd"], e["ort"],
                    e["pers"], None, e["an"], "bestätigt", f"{e['dep']['pct']} % = {de(e['dep_amt'])}",
                    dde(D(e["dep"]["paid"])) if D(e["dep"]["paid"]) <= upd else "", "", ""]
            for j, v in enumerate(vals, start=1):
                c = ws3.cell(row=k, column=j, value=v)
                if isinstance(v, date):
                    c.number_format = "DD.MM.YYYY"
    save_wb(wb, path)


def write_angebote(path: Path, dv: dict):
    upd = dv["list_upd"]
    wb = openpyxl.Workbook()
    fixed_props(wb)
    ws = wb.active
    ws.title = "Angebote"
    ws.cell(row=1, column=1, value="Angebote 2025/2026").font = Font(bold=True, size=13)
    ws.cell(row=2, column=1, value=f"Stand {dde(upd)}")
    heads = ["Angebot", "Datum", "Kunde / Interessent", "Anlass", "Termin", "Pers.", "Summe netto", "Summe brutto", "gültig bis",
             "Status", "Nachfassen am", "VA-Nr.", "Bemerkung"]
    xl_header(ws, 4, heads, [11, 11, 32, 30, 11, 6, 12, 12, 11, 11, 13, 10, 30])
    items = []
    for e in EVENTS:
        if e["offer"] and e["offer"] <= upd:
            st = "angenommen" if e["booked"] <= upd else "offen"
            items.append((e["offer"], e["an"], CUSTOMERS[e["kd"]]["name"], e["anlass"], e["date"], e["pers"], e["offer_tot"],
                          st, None if st == "angenommen" else e["offer"] + days(14), e["va"] if st == "angenommen" else "", ""))
    for o in OFFERS_LOST:
        if o["date"] > upd:
            continue
        fake = dict(food=[("M", o["food"])], drinks=o["drinks"], staff=o["staff"], log=o["log"], rental=[], key=o["key"])
        tot = totals(ev_lines(fake, o["pers"], offer=True))
        who = CUSTOMERS[o["kd"]]["name"] if o["kd"] else o["prospect"]
        items.append((o["date"], o["an"], who, o["anlass"], o["when"], o["pers"], tot, o["status"], o["follow"], "", o.get("why", "")))
    r = 5
    for d, an, who, anl, when, pers, tot, st, follow, va, note in sorted(items, key=lambda x: x[1]):
        vals = [an, d, who, anl, when, pers, eur(tot["n7"] + tot["n19"]), eur(tot["gross"]), d + days(30), st, follow, va, note]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=j, value=v)
            if isinstance(v, date):
                c.number_format = "DD.MM.YYYY"
            if j in (7, 8):
                c.number_format = "#,##0.00"
        r += 1
    save_wb(wb, path)


def write_vertraege(path: Path, dv: dict):
    wb = openpyxl.Workbook()
    fixed_props(wb, "Sabine Danzer")
    ws = wb.active
    ws.title = "Kitas & Schulen"
    ws.cell(row=1, column=1, value="Verpflegungsverträge Kitas / Schulen / Tagespflege").font = Font(bold=True, size=13)
    heads = ["Vertrag", "Einrichtung", "Träger (Rechnungsempfänger)", "Kd.-Nr.", "Lieferadresse", "Beginn", "Preis/Essen netto",
             "ab 01.09.2026", "Liefertage", "Abrechnung", "Zahlung", "Leitweg-ID", "gekündigt zum", "Bemerkung"]
    xl_header(ws, 3, heads, [8, 28, 40, 8, 34, 11, 10, 10, 10, 26, 30, 18, 12, 44])
    pay = {"MV-01": "Überweisung 30 Tage", "MV-02": "Überweisung 30 Tage", "MV-03": "Überweisung 14 Tage",
           "MV-04": "Dauerauftrag 900,00 € zum 5.", "MV-05": "Überweisung 14 Tage", "MV-06": "Überweisung 14 Tage"}
    for r, m in enumerate(MV, start=4):
        note = m["note"]
        gek = None
        if m["id"] == "MV-03" and dv["cut"] >= D("2026-10-12"):
            gek, note = D("2026-12-31"), note + "; Kündigung per Mail vom 12.10.2026"
        vals = [m["id"], m["name"], CUSTOMERS[m["kd"]]["name"], m["kd"], m["addr"], m["start"], eur(m["p_old"]), eur(m["p_new"]),
                "Mo–Do" if m["wd"] == (0, 1, 2, 3) else "Mo–Fr", "monatlich nach Essenszahlen", pay[m["id"]],
                CUSTOMERS[m["kd"]]["leitweg"], gek, note]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=j, value=v)
            if isinstance(v, date):
                c.number_format = "DD.MM.YYYY"
    ws2 = wb.create_sheet("Firmen & Rahmenverträge")
    heads2 = ["Vertrag", "Kunde", "Kd.-Nr.", "Leistung", "Beginn", "Pauschale netto/Monat", "Abrechnung", "Zahlung", "Bemerkung"]
    xl_header(ws2, 1, heads2, [8, 36, 8, 48, 11, 14, 26, 32, 50])
    r = 2
    for f in FV:
        note = f["note"]
        flat = eur(f["flat"][0][1])
        if f["id"] == "FV-01" and dv["cut"] >= D("2026-09-23"):
            note += "; ab 01.10.2026 4.300,00 € (30 Personen, Mail 22.09.)"
        vals = [f["id"], CUSTOMERS[f["kd"]]["name"], f["kd"], f["what"], f["start"], flat, f["billing"], f["pay"], note]
        for j, v in enumerate(vals, start=1):
            c = ws2.cell(row=r, column=j, value=v)
            if isinstance(v, date):
                c.number_format = "DD.MM.YYYY"
        r += 1
    for rv in RV:
        vals = [rv["id"], CUSTOMERS[rv["kd"]]["name"], rv["kd"], f"{rv['name']}: {rv['food'][0]} {de(rv['food'][1])} € p. P. + "
                f"{rv['drinks'][0]} {de(rv['drinks'][1])} € p. P.", rv["start"], "nach Abruf", rv["billing"],
                "Überweisung 30 Tage", f"Bestellnummer {rv['order']} auf jeder Rechnung" if rv["order"] else "Abruf per Mail"]
        for j, v in enumerate(vals, start=1):
            c = ws2.cell(row=r, column=j, value=v)
            if isinstance(v, date):
                c.number_format = "DD.MM.YYYY"
        r += 1
    save_wb(wb, path)


def write_essenszahlen(path: Path, dv: dict):
    cut = dv["cut"]
    wb = openpyxl.Workbook()
    fixed_props(wb, "Florian Wagner")
    wb.remove(wb.active)
    for mo in range(1, cut.month + 1):
        ws = wb.create_sheet(f"{MON3[mo - 1]} 26")
        ws.cell(row=1, column=1, value=f"Essenszahlen {MONATE[mo - 1]} 2026 (Portionen)").font = Font(bold=True, size=12)
        heads = ["Datum", "Tag"] + [f"{m['id']} {m['name']}" for m in MV] + ["Summe"]
        xl_header(ws, 3, heads, [11, 5] + [14] * len(MV) + [8])
        r = 4
        d = date(2026, mo, 1)
        sums = [0] * len(MV)
        while d.month == mo and d <= cut:
            if d.weekday() < 5:
                ws.cell(row=r, column=1, value=d).number_format = "DD.MM.YYYY"
                ws.cell(row=r, column=2, value=WT[d.weekday()])
                tot = 0
                for j, m in enumerate(MV):
                    n = MEALS.get((m["id"], 2026, mo), {}).get(d)
                    if n is None:
                        v = "Feiertag" if d in HOL else ("Ferien" if m["id"] == "MV-05" and in_ranges(d, SCHOOL_HOL) else
                                                         ("–" if d.weekday() not in m["wd"] else "zu"))
                    else:
                        v, tot, sums[j] = n, tot + n, sums[j] + n
                    ws.cell(row=r, column=3 + j, value=v)
                ws.cell(row=r, column=3 + len(MV), value=tot)
                r += 1
            d += days(1)
        ws.cell(row=r, column=1, value="Summe").font = Font(bold=True)
        for j, s in enumerate(sums):
            ws.cell(row=r, column=3 + j, value=s).font = Font(bold=True)
        ws.cell(row=r, column=3 + len(MV), value=sum(sums)).font = Font(bold=True)
        ws.cell(row=r + 1, column=1, value="Preis netto")
        for j, m in enumerate(MV):
            ws.cell(row=r + 1, column=3 + j, value=eur(mv_price(m, date(2026, mo, 1)))).number_format = "0.00"
    save_wb(wb, path)


def write_kassenbuch(path: Path, dv: dict):
    cut = dv["cut"]
    wb = openpyxl.Workbook()
    fixed_props(wb, "Sabine Danzer")
    wb.remove(wb.active)
    bal = 41235
    for mo in range(6, cut.month + 1):
        ws = wb.create_sheet(MONATE[mo - 1])
        ws.cell(row=1, column=1, value=f"Kassenbuch {MONATE[mo - 1]} 2026 – Hollerieth Catering & Events GmbH").font = Font(bold=True, size=12)
        xl_header(ws, 3, ["Datum", "Beleg-Nr.", "Buchungstext", "Einnahmen", "Ausgaben", "Bestand"], [11, 10, 52, 11, 11, 11])
        ws.cell(row=4, column=3, value="Übertrag")
        ws.cell(row=4, column=6, value=eur(bal)).number_format = "#,##0.00"
        r = 5
        for x in CASH:
            if x["date"].month != mo or x["date"] > cut:
                continue
            bal += x["amount"]
            ws.cell(row=r, column=1, value=x["date"]).number_format = "DD.MM.YYYY"
            ws.cell(row=r, column=2, value=x["no"])
            ws.cell(row=r, column=3, value=x["text"])
            ws.cell(row=r, column=4 if x["amount"] > 0 else 5, value=eur(abs(x["amount"]))).number_format = "#,##0.00"
            ws.cell(row=r, column=6, value=eur(bal)).number_format = "#,##0.00"
            r += 1
    save_wb(wb, path)


# ------------------------------------------------------------------ writers: bank statements as camt.053 (one XML per booking day, a ZIP per month)
GVC = {"166": ("PMNT", "RCDT", "ESCT", "NTRF"), "152": ("PMNT", "RCDT", "ESCT", "NSTO"), "105": ("PMNT", "IDDT", "ESDD", "NDDT"),
       "116": ("PMNT", "ICDT", "ESCT", "NTRF"), "117": ("PMNT", "ICDT", "ESCT", "NSTO"), "191": ("PMNT", "ICDT", "ESCT", "NTRF"),
       "082": ("PMNT", "CNTR", "CDPT", "NCMI"), "805": ("ACMT", "MDOP", "CHRG", "NCHG")}


def tz_for(d: date) -> str:
    return "+01:00" if d >= D("2026-10-25") else "+02:00"


def camt_party(tag: str, name: str, iban: str, bic: str, ci: str = "") -> str:
    pid = (f"<Id><PrvtId><Othr><Id>{ci}</Id><SchmeNm><Prtry>SEPA</Prtry></SchmeNm></Othr></PrvtId></Id>" if ci else "")
    acc = f"<{tag}Acct><Id><IBAN>{iban}</IBAN></Id></{tag}Acct>" if iban else ""
    return f"<{tag}><Pty><Nm>{xesc(name)}</Nm>{pid}</Pty></{tag}>{acc}"


def camt_tx(m: dict, amount: int, cd: str, name: str, iban: str, bic: str, vwz: str, ref: str) -> str:
    dom, fam, sub, tc = GVC[m["gvc"]]
    prtry = f"{tc}+{m['gvc']}+{9300 + int(ref[-3:]) % 97:04d}"
    refs = f"<AcctSvcrRef>{ref}</AcctSvcrRef><EndToEndId>{xesc(m['e2e'] or 'NOTPROVIDED')}</EndToEndId>"
    if m["mandate"]:
        refs += f"<MndtId>{xesc(m['mandate'])}</MndtId>"
    own = camt_party("Cdtr" if cd == "CRDT" else "Dbtr", CAMT_NAME, FIRM["iban"], FIRM["bic"])
    if m["gvc"] in ("082", "805"):
        parties, agents = "", ""
    elif cd == "CRDT":
        parties = f"<RltdPties>{camt_party('Dbtr', name, iban, bic)}{own}</RltdPties>"
        agents = f"<RltdAgts><DbtrAgt><FinInstnId><BICFI>{bic}</BICFI></FinInstnId></DbtrAgt></RltdAgts>" if bic else ""
    else:
        parties = f"<RltdPties>{own}{camt_party('Cdtr', name, iban, bic, m['ci'])}</RltdPties>"
        agents = f"<RltdAgts><CdtrAgt><FinInstnId><BICFI>{bic}</BICFI></FinInstnId></CdtrAgt></RltdAgts>" if bic else ""
    rmt = f"<RmtInf><Ustrd>{xesc(vwz[:140])}</Ustrd></RmtInf>" if vwz else ""
    return (f"<TxDtls><Refs>{refs}</Refs><Amt Ccy=\"EUR\">{xamt(amount)}</Amt><CdtDbtInd>{cd}</CdtDbtInd>"
            f"<BkTxCd><Domn><Cd>{dom}</Cd><Fmly><Cd>{fam}</Cd><SubFmlyCd>{sub}</SubFmlyCd></Fmly></Domn>"
            f"<Prtry><Cd>{prtry}</Cd><Issr>DK</Issr></Prtry></BkTxCd>{parties}{agents}{rmt}"
            f"<AddtlTxInf>{xesc(m['btext'])}</AddtlTxInf></TxDtls>")


def camt_entry(m: dict, k: int) -> str:
    amt, cd = abs(m["amount"]), ("CRDT" if m["amount"] > 0 else "DBIT")
    ref = f"{m['date']:%Y%m%d}{k:06d}"
    dom, fam, sub, tc = GVC[m["gvc"]]
    if m["batch"]:
        txs = "".join(camt_tx(m, x["amount"], cd, x["name"], x["iban"], "", m["vwz"], f"{ref}{j:02d}") for j, x in enumerate(m["batch"], 1))
        dtls = (f"<NtryDtls><Btch><NbOfTxs>{len(m['batch'])}</NbOfTxs><TtlAmt Ccy=\"EUR\">{xamt(amt)}</TtlAmt>"
                f"<CdtDbtInd>{cd}</CdtDbtInd></Btch>{txs}</NtryDtls>")
        bb = "<BtchBookg>true</BtchBookg>"
    else:
        dtls = f"<NtryDtls>{camt_tx(m, amt, cd, m['payer'], m['iban'], m['bic'], m['vwz'], ref)}</NtryDtls>"
        bb = ""
    return (f"<Ntry><NtryRef>{k}</NtryRef><Amt Ccy=\"EUR\">{xamt(amt)}</Amt><CdtDbtInd>{cd}</CdtDbtInd>"
            f"<Sts><Cd>BOOK</Cd></Sts><BookgDt><Dt>{m['date'].isoformat()}</Dt></BookgDt><ValDt><Dt>{m['date'].isoformat()}</Dt></ValDt>"
            f"<AcctSvcrRef>{ref}</AcctSvcrRef><BkTxCd><Domn><Cd>{dom}</Cd><Fmly><Cd>{fam}</Cd><SubFmlyCd>{sub}</SubFmlyCd></Fmly></Domn>"
            f"<Prtry><Cd>{tc}+{m['gvc']}</Cd><Issr>DK</Issr></Prtry></BkTxCd>{bb}{dtls}"
            f"<AddtlNtryInf>{xesc(m['btext'])}</AddtlNtryInf></Ntry>")


def camt_statements(cut: date) -> list[tuple[date, str, bytes]]:
    """All daily statements of the account up to cut: (day, file name, xml)."""
    bank = sorted((m for m in MONEY if m["ch"] == "bank" and m["date"] <= cut), key=lambda m: (m["date"], m["id"]))
    out, bal, seq, k = [], 0, 0, 0
    for d in sorted({m["date"] for m in bank}):
        seq += 1
        day = [m for m in bank if m["date"] == d]
        opening = bal
        ents = []
        for m in day:
            k += 1
            ents.append(camt_entry(m, k))
            bal += m["amount"]
        created = datetime.combine(next_bd(d + days(1)), datetime.min.time()).replace(hour=5, minute=10 + seq % 40)
        cre = created.strftime("%Y-%m-%dT%H:%M:%S") + tz_for(created.date())
        cr = [m["amount"] for m in day if m["amount"] > 0]
        db = [-m["amount"] for m in day if m["amount"] < 0]

        def bal_xml(code, amount, when):
            return (f"<Bal><Tp><CdOrPrtry><Cd>{code}</Cd></CdOrPrtry></Tp><Amt Ccy=\"EUR\">{xamt(abs(amount))}</Amt>"
                    f"<CdtDbtInd>{'CRDT' if amount >= 0 else 'DBIT'}</CdtDbtInd><Dt><Dt>{when.isoformat()}</Dt></Dt></Bal>")
        xml = (f"<?xml version=\"1.0\" encoding=\"UTF-8\"?>\n"
               f"<Document xmlns=\"urn:iso:std:iso:20022:tech:xsd:camt.053.001.08\" xmlns:xsi=\"http://www.w3.org/2001/XMLSchema-instance\" "
               f"xsi:schemaLocation=\"urn:iso:std:iso:20022:tech:xsd:camt.053.001.08 camt.053.001.08.xsd\">\n"
               f"<BkToCstmrStmt><GrpHdr><MsgId>C53{d:%Y%m%d}{seq:05d}RBDT</MsgId><CreDtTm>{cre}</CreDtTm>"
               f"<MsgPgntn><PgNb>1</PgNb><LastPgInd>true</LastPgInd></MsgPgntn></GrpHdr>\n"
               f"<Stmt><Id>{FIRM['iban']}{d:%Y%m%d}{seq:04d}</Id><ElctrncSeqNb>{seq}</ElctrncSeqNb><LglSeqNb>{seq}</LglSeqNb>"
               f"<CreDtTm>{cre}</CreDtTm><FrToDt><FrDtTm>{d.isoformat()}T00:00:00{tz_for(d)}</FrDtTm>"
               f"<ToDtTm>{d.isoformat()}T23:59:59{tz_for(d)}</ToDtTm></FrToDt>"
               f"<Acct><Id><IBAN>{FIRM['iban']}</IBAN></Id><Ccy>EUR</Ccy><Ownr><Nm>{CAMT_NAME}</Nm></Ownr>"
               f"<Svcr><FinInstnId><BICFI>{FIRM['bic']}</BICFI><Nm>{FIRM['bank']}</Nm></FinInstnId></Svcr></Acct>\n"
               f"{bal_xml('PRCD', opening, d - days(1))}{bal_xml('CLBD', bal, d)}\n"
               f"<TxsSummry><TtlNtries><NbOfNtries>{len(day)}</NbOfNtries></TtlNtries>"
               f"<TtlCdtNtries><NbOfNtries>{len(cr)}</NbOfNtries><Sum>{xamt(sum(cr))}</Sum></TtlCdtNtries>"
               f"<TtlDbtNtries><NbOfNtries>{len(db)}</NbOfNtries><Sum>{xamt(sum(db))}</Sum></TtlDbtNtries></TxsSummry>\n"
               + "\n".join(ents) + "\n</Stmt></BkToCstmrStmt></Document>\n")
        out.append((d, f"{d:%Y%m%d}_C53_{FIRM['iban']}_EUR_{seq:05d}.xml", xml.encode("utf-8")))
    return out


def write_camt(folder: Path, dv: dict):
    by_month: dict[tuple, list] = {}
    for d, name, data in camt_statements(dv["cut"]):
        by_month.setdefault((d.year, d.month), []).append((name, data))
    for (y, mo), files in sorted(by_month.items()):
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
            for name, data in files:
                zi = zipfile.ZipInfo(name, date_time=(y, mo, 28, 6, 0, 0))
                zi.compress_type = zipfile.ZIP_DEFLATED
                z.writestr(zi, data)
        (folder / f"camt053_{FIRM['iban']}_{y}-{mo:02d}.zip").write_bytes(buf.getvalue())


# ------------------------------------------------------------------ writers: PDFs (own invoices, offers, supplier invoices)
def _pdf(draw) -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4, invariant=1)
    c.setTitle("Beleg")
    c.setAuthor("")
    draw(c)
    c.showPage()
    c.save()
    return buf.getvalue()


def _table(c, y, rows, totals_rows, foot):
    c.setFont("Helvetica-Bold", 8.5)
    cols = [20, 30, 120, 138, 160, 178]
    for x, h in zip(cols, ["Pos.", "Bezeichnung", "Menge", "Einheit", "Einzelpreis", "Gesamt"]):
        c.drawString(x * mm, y, h)
    c.line(20 * mm, y - 2, 195 * mm, y - 2)
    y -= 13
    c.setFont("Helvetica", 8.5)
    for k, (t, qq, u, p, r) in enumerate(rows, start=1):
        if y < 60 * mm:
            c.showPage()
            c.setFont("Helvetica", 8.5)
            y = 270 * mm
        c.drawString(20 * mm, y, str(k))
        c.drawString(30 * mm, y, t[:62] + (f" ({r} %)" if r else ""))
        c.drawRightString(132 * mm, y, (f"{qq:,}".replace(",", ".") if isinstance(qq, int) else str(qq)))
        c.drawString(138 * mm, y, u)
        c.drawRightString(175 * mm, y, de(p))
        c.drawRightString(195 * mm, y, de(round(qq * p)))
        y -= 12
    c.line(120 * mm, y + 4, 195 * mm, y + 4)
    y -= 6
    for label, val, bold in totals_rows:
        c.setFont("Helvetica-Bold" if bold else "Helvetica", 9)
        c.drawString(120 * mm, y, label)
        c.drawRightString(195 * mm, y, val)
        y -= 12
    y -= 8
    c.setFont("Helvetica", 8.5)
    for line in foot:
        c.drawString(20 * mm, y, line)
        y -= 11
    return y


def _footer(c, lines):
    c.setFont("Helvetica", 6.8)
    c.line(20 * mm, 22 * mm, 195 * mm, 22 * mm)
    for k, line in enumerate(lines):
        c.drawString(20 * mm, (18 - 3.2 * k) * mm, line)


def own_footer() -> list[str]:
    return [f"{FIRM['name']} · {FIRM['street']} · {FIRM['zip']} {FIRM['city']} · Tel. {FIRM['phone']} · {FIRM['mail']} · {FIRM['web']}",
            f"Geschäftsführung: {FIRM['gf']} · {FIRM['hrb']} · USt-IdNr. {FIRM['vat']} · Steuernummer {FIRM['taxno']}",
            f"{FIRM['bank']} · IBAN {fmt_iban(FIRM['iban'])} · BIC {FIRM['bic']}"]


def _addr(c, lines, x=20, y=232):
    c.setFont("Helvetica", 7)
    c.drawString(x * mm, (y + 8) * mm, f"{FIRM['name']} · {FIRM['street']} · {FIRM['zip']} {FIRM['city']}")
    c.setFont("Helvetica", 10)
    for k, line in enumerate(lines):
        c.drawString(x * mm, (y - 5 * k) * mm, line)


def cust_addr(kd: int) -> list[str]:
    c = CUSTOMERS[kd]
    if c["typ"] == "Privatperson":
        return [f"{c['anrede']} {c['vor']} {c['nach']}".strip(), c["street"], f"{c['zip']} {c['city']}"]
    return [c["name"], c["street"], f"{c['zip']} {c['city']}"]


def own_head(c):
    c.setFont("Helvetica-Bold", 20)
    c.drawString(20 * mm, 272 * mm, "Hollerieth")
    c.setFont("Helvetica", 10)
    c.drawString(20 * mm, 266 * mm, "Catering & Events")


TITLE = {"RE": "Rechnung", "SR": "Schlussrechnung", "AR": "Anzahlungsrechnung", "STK": "Rechnung", "GS": "Gutschrift",
         "STO": "Stornorechnung"}


def pdf_own(i: dict) -> bytes:
    def draw(c):
        own_head(c)
        _addr(c, cust_addr(i["kd"]))
        c.setFont("Helvetica", 9)
        info = [("Rechnungsnummer" if i["kind"] not in ("GS", "STO") else "Belegnummer", i["no"]), ("Datum", dde(i["date"])),
                ("Kundennummer", str(i["kd"])), ("Leistungsdatum", i["sdate"]), ("Ihre Referenz", i["ref"] or "–"),
                ("Ansprechpartner", "Sabine Danzer")]
        for k, (a, b) in enumerate(info):
            c.drawString(130 * mm, (240 - 5 * k) * mm, a)
            c.drawString(162 * mm, (240 - 5 * k) * mm, b)
        c.setFont("Helvetica-Bold", 13)
        c.drawString(20 * mm, 196 * mm, f"{TITLE[i['kind']]} {i['no']}")
        c.setFont("Helvetica", 10)
        c.drawString(20 * mm, 189 * mm, i["subject"])
        tr = []
        if i["n7"]:
            tr += [("Summe netto 7 %", de(i["n7"]), False), ("zzgl. 7 % USt", de(i["v7"]), False)]
        if i["n19"]:
            tr += [("Summe netto 19 %", de(i["n19"]), False), ("zzgl. 19 % USt", de(i["v19"]), False)]
        tr.append(("Rechnungsbetrag" if i["kind"] not in ("GS", "STO") else "Betrag", de(i["gross"]) + " €", True))
        if i["deduct"]:
            tr += [(i["deduct_text"][:46], "-" + de(i["deduct"]) + " €", False), ("Zahlbetrag", de(i["pay"]) + " €", True)]
        if i["kind"] == "GS":
            foot = [f"Den Betrag der Gutschrift verrechnen wir mit Rechnung {i['applies_to']['no']}.",
                    f"Offen bleibt aus Rechnung {i['applies_to']['no']}: {de(i['applies_to']['pay'] + i['gross'])} €."]
        elif i["kind"] == "STO":
            foot = [f"Diese Stornorechnung hebt Rechnung {i['applies_to']['no']} vom {dde(i['applies_to']['date'])} vollständig auf.",
                    "Bitte zahlen Sie nur die neue Rechnung."]
        elif i["kind"] == "AR":
            foot = [f"Bitte überweisen Sie die Anzahlung bis zum {dde(i['due'])} unter Angabe der Rechnungsnummer.",
                    "Die Anzahlung wird in der Schlussrechnung nach der Veranstaltung abgezogen."]
        else:
            foot = [f"Bitte überweisen Sie den Zahlbetrag bis zum {dde(i['due'])} ohne Abzug unter Angabe der Rechnungsnummer.",
                    "Speisen und Service 7 % USt, Getränke und Mietartikel 19 % USt."]
        foot += ["", "Vielen Dank für Ihren Auftrag!", "Katrin und Tobias Hollerieth"]
        _table(c, 180 * mm, i["lines"], tr, foot)
        _footer(c, own_footer())
    return _pdf(draw)


def pdf_offer(e_or_o: dict, an: str, kd: int | None, d: date, lines, when: date, anlass: str, conf: dict | None = None) -> bytes:
    t = totals(lines)

    def draw(c):
        own_head(c)
        _addr(c, cust_addr(kd) if kd else ["Interessent"])
        c.setFont("Helvetica", 9)
        rows = [("Angebot" if not conf else "Auftrag", an if not conf else conf["va"]), ("Datum", dde(d)),
                ("Veranstaltung", dde(when)), ("Kundennummer", str(kd) if kd else "–")]
        for k, (a, b) in enumerate(rows):
            c.drawString(130 * mm, (240 - 5 * k) * mm, a)
            c.drawString(162 * mm, (240 - 5 * k) * mm, b)
        c.setFont("Helvetica-Bold", 13)
        c.drawString(20 * mm, 196 * mm, ("Angebot " + an) if not conf else ("Auftragsbestätigung " + conf["va"]))
        c.setFont("Helvetica", 10)
        c.drawString(20 * mm, 189 * mm, f"{anlass} am {dde(when)}")
        tr = [("Summe netto", de(t["n7"] + t["n19"]), False), ("USt 7 % / 19 %", de(t["v7"] + t["v19"]), False),
              ("Angebotssumme brutto", de(t["gross"]) + " €", True)]
        foot = ["Dieses Angebot gilt 30 Tage. Endabrechnung nach tatsächlicher Personenzahl und Verbrauch.",
                "Bei Auftragserteilung bitten wir um eine Anzahlung von 30 % (Privatkunden), Rest nach Rechnung.",
                "Stornierung: bis 60 Tage vorher kostenfrei, 59–30 Tage 25 %, 29–14 Tage 50 %, danach 80 % (Ziff. 8 AGB)."]
        if conf:
            foot = [f"Wir bestätigen Ihren Auftrag {conf['va']} gemäß Angebot {an}.",
                    f"Bitte überweisen Sie die Anzahlung von {conf['pct']} % = {de(conf['amount'])} € bis zum {dde(conf['until'])}.",
                    "Eine Anzahlungsrechnung erhalten Sie auf Wunsch; die Anzahlung wird in der Schlussrechnung abgezogen."]
        _table(c, 180 * mm, lines, tr, foot)
        _footer(c, own_footer())
    return _pdf(draw)


def pdf_supplier(x: dict, fraud: bool = False) -> bytes:
    s = SUPPLIERS[x["sup"]]

    def draw(c):
        c.setFont("Helvetica-Bold", 15)
        c.drawString(20 * mm, 272 * mm, s["name"])
        c.setFont("Helvetica", 8)
        c.drawString(20 * mm, 266 * mm, f"{s['street']} · {s['zip']} {s['city']} · Tel. {s['phone']} · "
                                        f"{FRAUD['mail'] if fraud else s['mail']}")
        c.setFont("Helvetica", 7)
        c.drawString(20 * mm, 240 * mm, f"{s['name']} · {s['street']} · {s['zip']} {s['city']}")
        c.setFont("Helvetica", 10)
        for k, line in enumerate([FIRM["name"], "Buchhaltung", FIRM["street"], f"{FIRM['zip']} {FIRM['city']}"]):
            c.drawString(20 * mm, (232 - 5 * k) * mm, line)
        c.setFont("Helvetica", 9)
        rows = [("Rechnungs-Nr.", x["no"]), ("Rechnungsdatum", dde(x["date"])), ("Kunden-Nr.", s["kdnr"]),
                ("USt-IdNr.", s["vat"])]
        for k, (a, b) in enumerate(rows):
            c.drawString(130 * mm, (240 - 5 * k) * mm, a)
            c.drawString(162 * mm, (240 - 5 * k) * mm, b)
        c.setFont("Helvetica-Bold", 13)
        c.drawString(20 * mm, 196 * mm, ("KORRIGIERTE RECHNUNG " if fraud else "Rechnung ") + x["no"])
        tr = []
        if x["n7"]:
            tr += [("Netto 7 %", de(x["n7"]), False), ("USt 7 %", de(x["v7"]), False)]
        if x["n19"]:
            tr += [("Netto 19 %", de(x["n19"]), False), ("USt 19 %", de(x["v19"]), False)]
        tr.append(("Rechnungsbetrag", de(x["gross"]) + " €", True))
        if fraud:
            foot = ["Achtung – geänderte Bankverbindung! Bitte überweisen Sie den Betrag umgehend auf unser neues Konto:",
                    f"IBAN {fmt_iban(FRAUD['iban'])} · BIC {FRAUD['bic']}. Die Lastschrift wird nicht mehr eingezogen."]
        elif x["how"] == "LS":
            foot = [f"Der Rechnungsbetrag wird am {dde(next_bd(x['date'] + days(7 if x['sup'] == 'ggm' else 10)))} per SEPA-Lastschrift "
                    f"eingezogen (Mandat {s['mandate']}, Gläubiger-ID {s['ci']})."]
        elif x["skonto"]:
            foot = [f"Zahlbar bis {dde(x['skonto_until'])} abzüglich 2 % Skonto ({de(x['gross'] - pct(x['gross'], 2))} €)",
                    f"oder bis {dde(x['due'])} netto ohne Abzug."]
        else:
            foot = [f"Zahlbar bis {dde(x['due'])} ohne Abzug."]
        _table(c, 180 * mm, x["lines"], tr, foot)
        bank = (f"IBAN {fmt_iban(FRAUD['iban'])} · BIC {FRAUD['bic']}" if fraud else f"IBAN {fmt_iban(s['iban'])} · BIC {s['bic']}")
        _footer(c, [f"{s['name']} · {s['street']} · {s['zip']} {s['city']} · {s['hrb']}", f"USt-IdNr. {s['vat']} · {bank}"])
    data = _pdf(draw)
    if x["fmt"] == "zugferd" and not fraud:
        from pypdf import PdfReader, PdfWriter
        w = PdfWriter(clone_from=PdfReader(io.BytesIO(data)))
        w.add_attachment("factur-x.xml", cii_xml(x))
        w._ID = None
        buf = io.BytesIO()
        w.write(buf)
        data = buf.getvalue()
    return data


# ------------------------------------------------------------------ writers: e-invoices
def cii_xml(x: dict) -> bytes:
    s = SUPPLIERS[x["sup"]]
    dd = x["date"].strftime("%Y%m%d")
    items = []
    for k, (t, qq, u, p, r) in enumerate(x["lines"], start=1):
        items.append(f"""  <ram:IncludedSupplyChainTradeLineItem>
   <ram:AssociatedDocumentLineDocument><ram:LineID>{k}</ram:LineID></ram:AssociatedDocumentLineDocument>
   <ram:SpecifiedTradeProduct><ram:Name>{xesc(t)}</ram:Name></ram:SpecifiedTradeProduct>
   <ram:SpecifiedLineTradeAgreement><ram:NetPriceProductTradePrice><ram:ChargeAmount>{xamt(p)}</ram:ChargeAmount></ram:NetPriceProductTradePrice></ram:SpecifiedLineTradeAgreement>
   <ram:SpecifiedLineTradeDelivery><ram:BilledQuantity unitCode="C62">{qq}</ram:BilledQuantity></ram:SpecifiedLineTradeDelivery>
   <ram:SpecifiedLineTradeSettlement><ram:ApplicableTradeTax><ram:TypeCode>VAT</ram:TypeCode><ram:CategoryCode>S</ram:CategoryCode><ram:RateApplicablePercent>{r}</ram:RateApplicablePercent></ram:ApplicableTradeTax>
    <ram:SpecifiedTradeSettlementLineMonetarySummation><ram:LineTotalAmount>{xamt(round(qq * p))}</ram:LineTotalAmount></ram:SpecifiedTradeSettlementLineMonetarySummation></ram:SpecifiedLineTradeSettlement>
  </ram:IncludedSupplyChainTradeLineItem>""")
    taxes = []
    for rate, n, v in ((7, x["n7"], x["v7"]), (19, x["n19"], x["v19"])):
        if n:
            taxes.append(f"   <ram:ApplicableTradeTax><ram:CalculatedAmount>{xamt(v)}</ram:CalculatedAmount><ram:TypeCode>VAT</ram:TypeCode>"
                         f"<ram:BasisAmount>{xamt(n)}</ram:BasisAmount><ram:CategoryCode>S</ram:CategoryCode>"
                         f"<ram:RateApplicablePercent>{rate}</ram:RateApplicablePercent></ram:ApplicableTradeTax>")
    due = next_bd(x["date"] + days(7))
    xml = f"""<?xml version="1.0" encoding="UTF-8"?>
<rsm:CrossIndustryInvoice xmlns:rsm="urn:un:unece:uncefact:data:standard:CrossIndustryInvoice:100" xmlns:ram="urn:un:unece:uncefact:data:standard:ReusableAggregateBusinessInformationEntity:100" xmlns:udt="urn:un:unece:uncefact:data:standard:UnqualifiedDataType:100" xmlns:qdt="urn:un:unece:uncefact:data:standard:QualifiedDataType:100">
 <rsm:ExchangedDocumentContext><ram:GuidelineSpecifiedDocumentContextParameter><ram:ID>urn:cen.eu:en16931:2017</ram:ID></ram:GuidelineSpecifiedDocumentContextParameter></rsm:ExchangedDocumentContext>
 <rsm:ExchangedDocument><ram:ID>{x['no']}</ram:ID><ram:TypeCode>380</ram:TypeCode><ram:IssueDateTime><udt:DateTimeString format="102">{dd}</udt:DateTimeString></ram:IssueDateTime></rsm:ExchangedDocument>
 <rsm:SupplyChainTradeTransaction>
{chr(10).join(items)}
  <ram:ApplicableHeaderTradeAgreement>
   <ram:BuyerReference>{s['kdnr']}</ram:BuyerReference>
   <ram:SellerTradeParty><ram:Name>{xesc(s['name'])}</ram:Name>
    <ram:PostalTradeAddress><ram:PostcodeCode>{s['zip']}</ram:PostcodeCode><ram:LineOne>{xesc(s['street'])}</ram:LineOne><ram:CityName>{s['city']}</ram:CityName><ram:CountryID>DE</ram:CountryID></ram:PostalTradeAddress>
    <ram:URIUniversalCommunication><ram:URIID schemeID="EM">{s['mail']}</ram:URIID></ram:URIUniversalCommunication>
    <ram:SpecifiedTaxRegistration><ram:ID schemeID="VA">{s['vat']}</ram:ID></ram:SpecifiedTaxRegistration></ram:SellerTradeParty>
   <ram:BuyerTradeParty><ram:Name>{xesc(FIRM['name'])}</ram:Name>
    <ram:PostalTradeAddress><ram:PostcodeCode>{FIRM['zip']}</ram:PostcodeCode><ram:LineOne>{FIRM['street']}</ram:LineOne><ram:CityName>{FIRM['city']}</ram:CityName><ram:CountryID>DE</ram:CountryID></ram:PostalTradeAddress></ram:BuyerTradeParty>
  </ram:ApplicableHeaderTradeAgreement>
  <ram:ApplicableHeaderTradeDelivery><ram:ActualDeliverySupplyChainEvent><ram:OccurrenceDateTime><udt:DateTimeString format="102">{dd}</udt:DateTimeString></ram:OccurrenceDateTime></ram:ActualDeliverySupplyChainEvent></ram:ApplicableHeaderTradeDelivery>
  <ram:ApplicableHeaderTradeSettlement>
   <ram:CreditorReferenceID>{s['ci']}</ram:CreditorReferenceID>
   <ram:InvoiceCurrencyCode>EUR</ram:InvoiceCurrencyCode>
   <ram:SpecifiedTradeSettlementPaymentMeans><ram:TypeCode>59</ram:TypeCode><ram:PayerPartyDebtorFinancialAccount><ram:IBANID>{FIRM['iban']}</ram:IBANID></ram:PayerPartyDebtorFinancialAccount></ram:SpecifiedTradeSettlementPaymentMeans>
{chr(10).join(taxes)}
   <ram:SpecifiedTradePaymentTerms><ram:Description>Einzug per SEPA-Lastschrift am {dde(due)}</ram:Description><ram:DueDateDateTime><udt:DateTimeString format="102">{due:%Y%m%d}</udt:DateTimeString></ram:DueDateDateTime><ram:DirectDebitMandateID>{s['mandate']}</ram:DirectDebitMandateID></ram:SpecifiedTradePaymentTerms>
   <ram:SpecifiedTradeSettlementHeaderMonetarySummation><ram:LineTotalAmount>{xamt(x['n7'] + x['n19'])}</ram:LineTotalAmount><ram:TaxBasisTotalAmount>{xamt(x['n7'] + x['n19'])}</ram:TaxBasisTotalAmount><ram:TaxTotalAmount currencyID="EUR">{xamt(x['v7'] + x['v19'])}</ram:TaxTotalAmount><ram:GrandTotalAmount>{xamt(x['gross'])}</ram:GrandTotalAmount><ram:DuePayableAmount>{xamt(x['gross'])}</ram:DuePayableAmount></ram:SpecifiedTradeSettlementHeaderMonetarySummation>
  </ram:ApplicableHeaderTradeSettlement>
 </rsm:SupplyChainTradeTransaction>
</rsm:CrossIndustryInvoice>
"""
    return xml.encode("utf-8")


def ubl_xml(no, d, due, seller: dict, buyer: dict, lines, t: dict, buyer_ref: str, note: str, pay_iban: str, pay_bic: str,
            period: tuple | None = None) -> bytes:
    def party(p, endpoint, scheme):
        return f"""<cac:Party><cbc:EndpointID schemeID="{scheme}">{xesc(endpoint)}</cbc:EndpointID>
   <cac:PartyName><cbc:Name>{xesc(p['name'])}</cbc:Name></cac:PartyName>
   <cac:PostalAddress><cbc:StreetName>{xesc(p['street'])}</cbc:StreetName><cbc:CityName>{xesc(p['city'])}</cbc:CityName><cbc:PostalZone>{p['zip']}</cbc:PostalZone><cac:Country><cbc:IdentificationCode>DE</cbc:IdentificationCode></cac:Country></cac:PostalAddress>
   {f'<cac:PartyTaxScheme><cbc:CompanyID>{p["vat"]}</cbc:CompanyID><cac:TaxScheme><cbc:ID>VAT</cbc:ID></cac:TaxScheme></cac:PartyTaxScheme>' if p.get('vat') else ''}
   <cac:PartyLegalEntity><cbc:RegistrationName>{xesc(p['name'])}</cbc:RegistrationName></cac:PartyLegalEntity>
   {f'<cac:Contact><cbc:Name>{xesc(p.get("contact", ""))}</cbc:Name><cbc:Telephone>{p.get("phone", "")}</cbc:Telephone><cbc:ElectronicMail>{p.get("mail", "")}</cbc:ElectronicMail></cac:Contact>' if p.get('mail') else ''}</cac:Party>"""
    subs = []
    for rate, n, v in ((7, t["n7"], t["v7"]), (19, t["n19"], t["v19"])):
        if n:
            subs.append(f'<cac:TaxSubtotal><cbc:TaxableAmount currencyID="EUR">{xamt(n)}</cbc:TaxableAmount><cbc:TaxAmount currencyID="EUR">{xamt(v)}</cbc:TaxAmount>'
                        f'<cac:TaxCategory><cbc:ID>S</cbc:ID><cbc:Percent>{rate}</cbc:Percent><cac:TaxScheme><cbc:ID>VAT</cbc:ID></cac:TaxScheme></cac:TaxCategory></cac:TaxSubtotal>')
    il = []
    for k, (tx, qq, u, p, r) in enumerate(lines, start=1):
        il.append(f'<cac:InvoiceLine><cbc:ID>{k}</cbc:ID><cbc:InvoicedQuantity unitCode="{"HUR" if u == "Std." else "C62"}">{qq}</cbc:InvoicedQuantity>'
                  f'<cbc:LineExtensionAmount currencyID="EUR">{xamt(round(qq * p))}</cbc:LineExtensionAmount>'
                  f'<cac:Item><cbc:Name>{xesc(tx)}</cbc:Name><cac:ClassifiedTaxCategory><cbc:ID>S</cbc:ID><cbc:Percent>{r}</cbc:Percent>'
                  f'<cac:TaxScheme><cbc:ID>VAT</cbc:ID></cac:TaxScheme></cac:ClassifiedTaxCategory></cac:Item>'
                  f'<cac:Price><cbc:PriceAmount currencyID="EUR">{xamt(p)}</cbc:PriceAmount></cac:Price></cac:InvoiceLine>')
    per = (f"<cac:InvoicePeriod><cbc:StartDate>{period[0].isoformat()}</cbc:StartDate><cbc:EndDate>{period[1].isoformat()}</cbc:EndDate></cac:InvoicePeriod>"
           if period else "")
    xml = f"""<?xml version="1.0" encoding="UTF-8"?>
<ubl:Invoice xmlns:ubl="urn:oasis:names:specification:ubl:schema:xsd:Invoice-2" xmlns:cac="urn:oasis:names:specification:ubl:schema:xsd:CommonAggregateComponents-2" xmlns:cbc="urn:oasis:names:specification:ubl:schema:xsd:CommonBasicComponents-2">
 <cbc:CustomizationID>urn:cen.eu:en16931:2017#compliant#urn:xeinkauf.de:kosit:xrechnung_3.0</cbc:CustomizationID>
 <cbc:ProfileID>urn:fdc:peppol.eu:2017:poacc:billing:01:1.0</cbc:ProfileID>
 <cbc:ID>{no}</cbc:ID>
 <cbc:IssueDate>{d.isoformat()}</cbc:IssueDate>
 <cbc:DueDate>{due.isoformat()}</cbc:DueDate>
 <cbc:InvoiceTypeCode>380</cbc:InvoiceTypeCode>
 <cbc:Note>{xesc(note)}</cbc:Note>
 <cbc:DocumentCurrencyCode>EUR</cbc:DocumentCurrencyCode>
 <cbc:BuyerReference>{xesc(buyer_ref)}</cbc:BuyerReference>
 {per}
 <cac:AccountingSupplierParty>{party(seller, seller['mail'], 'EM')}</cac:AccountingSupplierParty>
 <cac:AccountingCustomerParty>{party(buyer, buyer_ref if buyer_ref.count('-') == 2 else buyer.get('mail', ''), '0204' if buyer_ref.count('-') == 2 else 'EM')}</cac:AccountingCustomerParty>
 <cac:PaymentMeans><cbc:PaymentMeansCode>58</cbc:PaymentMeansCode><cac:PayeeFinancialAccount><cbc:ID>{pay_iban}</cbc:ID><cbc:Name>{xesc(seller['name'])}</cbc:Name><cac:FinancialInstitutionBranch><cbc:ID>{pay_bic}</cbc:ID></cac:FinancialInstitutionBranch></cac:PayeeFinancialAccount></cac:PaymentMeans>
 <cac:PaymentTerms><cbc:Note>Zahlbar bis {dde(due)} ohne Abzug</cbc:Note></cac:PaymentTerms>
 <cac:TaxTotal><cbc:TaxAmount currencyID="EUR">{xamt(t['v7'] + t['v19'])}</cbc:TaxAmount>{''.join(subs)}</cac:TaxTotal>
 <cac:LegalMonetaryTotal><cbc:LineExtensionAmount currencyID="EUR">{xamt(t['n7'] + t['n19'])}</cbc:LineExtensionAmount><cbc:TaxExclusiveAmount currencyID="EUR">{xamt(t['n7'] + t['n19'])}</cbc:TaxExclusiveAmount><cbc:TaxInclusiveAmount currencyID="EUR">{xamt(t['gross'])}</cbc:TaxInclusiveAmount><cbc:PayableAmount currencyID="EUR">{xamt(t['gross'])}</cbc:PayableAmount></cac:LegalMonetaryTotal>
 {chr(10).join(il)}
</ubl:Invoice>
"""
    return xml.encode("utf-8")


FIRM_PARTY = dict(name=FIRM["name"], street=FIRM["street"], city=FIRM["city"], zip=FIRM["zip"], vat=FIRM["vat"],
                  mail=FIRM["office"], phone=FIRM["phone"], contact="Sabine Danzer")


def xr_own(i: dict, leitweg: str | None = None) -> bytes:
    c = CUSTOMERS[i["kd"]]
    buyer = dict(name=c["name"], street=c["street"], city=c["city"], zip=c["zip"], mail=c["email"])
    y, mo = i["src"]["y"], i["src"]["m"]
    m = MVD[i["src"]["contract"]]
    return ubl_xml(i["no"], i["date"], i["due"], FIRM_PARTY, buyer, i["lines"], i, leitweg or c["leitweg"],
                   f"{i['subject']}. Lieferadresse: {m['addr']}.", FIRM["iban"], FIRM["bic"], (date(y, mo, 1), month_end(y, mo)))


def xr_supplier(x: dict) -> bytes:
    s = SUPPLIERS[x["sup"]]
    seller = dict(name=s["name"], street=s["street"], city=s["city"], zip=s["zip"], vat=s["vat"], mail=s["mail"], phone=s["phone"],
                  contact="Debitorenbuchhaltung")
    return ubl_xml(x["no"], x["date"], x["due"], seller, dict(FIRM_PARTY, mail=FIRM["office"]), x["lines"], x, s["kdnr"],
                   "Arbeitnehmerüberlassung gem. Rahmenvertrag; Stundennachweise liegen vor.", s["iban"], s["bic"])


# ------------------------------------------------------------------ mails
OWN_INFO = "Hollerieth Catering & Events <info@hollerieth-catering.de>"
OWN_BUERO = "Sabine Danzer <buero@hollerieth-catering.de>"
OWN_TH = "Tobias Hollerieth <t.hollerieth@hollerieth-catering.de>"
SIG = ("\n\nMit freundlichen Grüßen\n{who}\n\nHollerieth Catering & Events GmbH\nIm Gewerbepark D 14 · 93059 Regensburg\n"
       "Tel. 0941 4609 310 · www.hollerieth-catering.de\n")


def sig(who="Sabine Danzer, Buchhaltung"):
    return SIG.format(who=who)


def build_mails() -> list[dict]:
    E = EVD
    inv = lambda k: E[k]["inv"]["no"]
    M = []

    def mail(key, when, frm, to, subject, body, att=(), reply=None):
        M.append(dict(key=key, when=when, frm=frm, to=to, subject=subject, body=body, att=list(att), reply=reply))

    eb, dg, lb, hb = E["ebner40"], E["dengler50"], E["lindbuechl_sommer"], E["hofbauer"]
    mail("ebner_dep", "2026-06-08T10:41", "Daniel Ebner <daniel.ebner@donaumail.example>", OWN_INFO,
         "Anzahlung Geburtstag 22.08.", f"Hallo Frau Hollerieth,\n\nich habe heute die Anzahlung von {de(eb['dep_amt'])} € "
         f"für meinen 40. am 22.08. überwiesen.\n\nViele Grüße\nDaniel Ebner")
    mail("ebner_cancel", "2026-06-15T19:22", "Daniel Ebner <daniel.ebner@donaumail.example>", OWN_INFO,
         "Absage Geburtstagsfeier 22.08.", "Hallo Frau Hollerieth,\n\nleider müssen wir die Feier absagen – ich wechsle beruflich "
         "nach Hamburg und wir ziehen schon im August um. Nach Ihren AGB ist das ja noch kostenfrei möglich.\n"
         "Bitte erstatten Sie mir die Anzahlung auf mein Konto.\n\nDanke und viele Grüße\nDaniel Ebner")
    mail("ebner_re", "2026-06-16T08:55", OWN_BUERO, "Daniel Ebner <daniel.ebner@donaumail.example>",
         "AW: Absage Geburtstagsfeier 22.08.", f"Hallo Herr Ebner,\n\nschade, aber natürlich – die Stornierung ist kostenfrei "
         f"(mehr als 60 Tage vorher). Die Anzahlung von {de(eb['dep_amt'])} € überweisen wir Ihnen bis 19.06. zurück.\n"
         f"Den Auftrag {eb['va']} haben wir storniert." + sig(), reply="ebner_cancel")
    mail("dengler_er", "2026-06-10T09:30", OWN_BUERO, "Markus Dengler <m.dengler@webpost.example>",
         f"Zahlungserinnerung Rechnung {inv('dengler50')}",
         f"Sehr geehrter Herr Dengler,\n\nzu unserer Rechnung {inv('dengler50')} vom {dde(dg['inv']['date'])} (50. Geburtstag am "
         f"25.04.2026) ist nach Ihrer Teilzahlung von 1.000,00 € noch ein Betrag von {de(dg['inv']['pay'] - 100000)} € offen.\n"
         f"Bitte überweisen Sie den Betrag bis zum 24.06.2026." + sig())
    mail("dengler_aw", "2026-06-23T07:58", "Markus Dengler <m.dengler@webpost.example>", OWN_BUERO,
         f"AW: Zahlungserinnerung Rechnung {inv('dengler50')}",
         "Guten Morgen,\n\ndas Buffet war um 21 Uhr leer, meine Gäste haben Pizza bestellt. Das habe ich Frau Hollerieth "
         "schon im Mai geschrieben. Den Rest zahle ich erst, wenn Sie mir ein vernünftiges Angebot machen.\n\nM. Dengler",
         reply="dengler_er")
    mail("dengler_m2", "2026-07-15T10:05", OWN_BUERO, "Markus Dengler <m.dengler@webpost.example>",
         f"2. Mahnung Rechnung {inv('dengler50')}",
         f"Sehr geehrter Herr Dengler,\n\nunser Angebot vom 30.06. (Nachlass 250,00 €) haben Sie nicht angenommen. Wir bitten Sie "
         f"daher, den offenen Betrag von {de(dg['inv']['pay'] - 100000)} € bis zum 29.07.2026 zu überweisen.\n"
         "Die Personenzahl lag mit 74 Gästen über den bestellten 70." + sig(), reply="dengler_aw")
    mail("bauer_name", "2026-06-16T21:03", "Michaela Bauer <michaela.bauer86@webpost.example>", OWN_INFO,
         "Rechnung Hochzeit – neuer Name", "Liebe Frau Hollerieth,\n\nnochmal tausend Dank für Samstag, es war traumhaft!\n"
         "Bitte stellen Sie die Rechnung auf meinen neuen Namen aus: Michaela Schreiner, Adresse bleibt gleich. "
         "Bezahlen wird mein Schwiegervater Josef Schreiner – das ist sein Hochzeitsgeschenk.\n\nHerzliche Grüße\nMichaela Schreiner")
    to_kitas = ", ".join(["kita-verpflegung@regensburg-stadt.de", "kita.st-wolfgang@pfarrei-stwolfgang-rgb.de",
                          "vorstand@kinderhaus-wirbelwind.de", "foerderverein.gs-kumpfmuehl@schulfreunde-rgb.de"])
    mail("preis", "2026-07-14T11:20", OWN_BUERO, to_kitas, "Preisanpassung Mittagsverpflegung zum 01.09.2026",
         "Sehr geehrte Damen und Herren,\n\ngemäß § 5 unseres Verpflegungsvertrags passen wir die Preise zum 01.09.2026 an die "
         "gestiegenen Lebensmittelkosten an:\n\nKitas: 4,15 € statt 3,95 € netto je Essen\nKita St. Wolfgang: 4,05 € statt 3,85 € netto\n"
         "Mittagsbetreuung Grundschule: 4,60 € statt 4,40 € netto\n\nAlle Preise zzgl. 7 % USt. Die Tagespflege ist nicht betroffen." + sig())
    sd = E["seidl_sommer"]
    mail("seidl", "2026-07-15T10:20", "Kanzlei Dr. Seidl & Partner <kanzlei@seidl-partner-rgb.de>", OWN_BUERO,
         f"Ihre Rechnung {sd['wrong_inv']['no']} Sommerfest",
         f"Sehr geehrte Frau Danzer,\n\ndie Rechnung {sd['wrong_inv']['no']} für unser Sommerfest ist an Dr. Seidl privat adressiert. "
         "Bitte stellen Sie sie auf die Kanzlei Dr. Seidl & Partner mbB aus, sonst können wir die Vorsteuer nicht ziehen.\n\n"
         "Freundliche Grüße\nB. Aumer, Kanzleiorganisation")
    mail("seidl_re", "2026-07-20T09:12", OWN_BUERO, "Kanzlei Dr. Seidl & Partner <kanzlei@seidl-partner-rgb.de>",
         f"AW: Ihre Rechnung {sd['wrong_inv']['no']} Sommerfest",
         f"Sehr geehrte Frau Aumer,\n\nentschuldigen Sie bitte. Anbei die Stornorechnung {sd['sto']['no']} und die neue Rechnung "
         f"{sd['inv']['no']} auf die Kanzlei. Bitte zahlen Sie nur die neue Rechnung." + sig(),
         att=[(f"Stornorechnung_{sd['sto']['no']}.pdf", ("pdf_own", sd["sto"])), (f"Rechnung_{sd['inv']['no']}.pdf", ("pdf_own", sd["inv"]))],
         reply="seidl")
    wm = E["wimmeder"]
    mail("wim", "2026-07-22T14:05", "Spedition Wimmeder – Buchhaltung <buchhaltung@wimmeder-logistik.de>", OWN_INFO,
         "Reklamation Sommerfest 17.07.", "Sehr geehrte Damen und Herren,\n\nunser Sommerfest war insgesamt schön, aber der Grill "
         "war 40 Minuten zu spät und die Nachspeisen waren um 20:30 Uhr aus – bei 180 Gästen ärgerlich. Wir bitten um eine "
         "angemessene Minderung.\n\nMit freundlichen Grüßen\nR. Wimmeder\nGeschäftsführung")
    mail("wim_re", "2026-07-27T09:30", OWN_TH, "Spedition Wimmeder – Buchhaltung <buchhaltung@wimmeder-logistik.de>",
         "AW: Reklamation Sommerfest 17.07.", "Sehr geehrter Herr Wimmeder,\n\nSie haben recht, das war nicht unser Anspruch. "
         "Wir schreiben Ihnen 10 % auf die Speisen gut, die Gutschrift kommt in den nächsten Tagen.\n\n"
         "Mit freundlichen Grüßen\nTobias Hollerieth", reply="wim")
    mail("wim_gs", "2026-07-28T10:02", OWN_BUERO, "Spedition Wimmeder – Buchhaltung <buchhaltung@wimmeder-logistik.de>",
         f"Gutschrift {wm['gs']['no']} zu Rechnung {wm['inv']['no']}",
         f"Sehr geehrte Damen und Herren,\n\nanbei die Gutschrift {wm['gs']['no']} über {de(-wm['gs']['gross'])} €. "
         f"Bitte überweisen Sie von Rechnung {wm['inv']['no']} nur noch {de(wm['inv']['pay'] + wm['gs']['gross'])} €." + sig(),
         att=[(f"Gutschrift_{wm['gs']['no']}.pdf", ("pdf_own", wm["gs"]))])
    lb_paid = int(round(lb["inv"]["pay"] * 0.85))
    mail("lind", "2026-08-10T16:44", "Autohaus Lindbüchl – Buchhaltung <buchhaltung@autohaus-lindbuechl.de>", OWN_BUERO,
         f"Ihre Rechnung {inv('lindbuechl_sommer')} – Minderung",
         f"Sehr geehrte Frau Danzer,\n\nbeim Sommerfest waren die Speisen teilweise kalt und der Service war überfordert. Wir haben "
         f"heute 85 % des Rechnungsbetrags überwiesen ({de(lb_paid)} €). Damit ist die Sache für uns erledigt.\n\n"
         "Mit freundlichen Grüßen\nAutohaus Lindbüchl GmbH\ni. A. Petra Lindbüchl")
    mail("lind_re", "2026-08-12T11:02", OWN_TH, "Autohaus Lindbüchl – Buchhaltung <buchhaltung@autohaus-lindbuechl.de>",
         f"AW: Ihre Rechnung {inv('lindbuechl_sommer')} – Minderung",
         "Sehr geehrte Frau Lindbüchl,\n\neine Minderung akzeptieren wir nicht: Während der Veranstaltung gab es keine Reklamation, "
         f"und 138 Gäste wurden vollständig versorgt. Bitte gleichen Sie den Restbetrag von {de(lb['inv']['pay'] - lb_paid)} € bis "
         "zum 26.08.2026 aus.\n\nMit freundlichen Grüßen\nTobias Hollerieth", reply="lind")
    mail("lind_er", "2026-09-01T10:30", OWN_BUERO, "Autohaus Lindbüchl – Buchhaltung <buchhaltung@autohaus-lindbuechl.de>",
         f"Zahlungserinnerung Rechnung {inv('lindbuechl_sommer')} (Restbetrag)",
         f"Sehr geehrte Damen und Herren,\n\nzu Rechnung {inv('lindbuechl_sommer')} ist noch ein Restbetrag von "
         f"{de(lb['inv']['pay'] - lb_paid)} € offen. Bitte überweisen Sie ihn bis zum 15.09.2026." + sig())
    mail("hof_er", "2026-08-20T09:40", OWN_BUERO, "Verena Hofbauer <verena.hofbauer@donaumail.example>",
         f"Zahlungserinnerung Rechnung {inv('hofbauer')}",
         f"Liebe Frau Hofbauer,\n\nvielen Dank für Ihre beiden Zahlungen. Aus Rechnung {inv('hofbauer')} sind noch "
         f"{de(hb['inv']['pay'] - 600000)} € offen. Bitte überweisen Sie den Rest bis 03.09.2026." + sig())
    mail("hof_aw", "2026-08-21T22:16", "Verena Hofbauer <verena.hofbauer@donaumail.example>", OWN_BUERO,
         f"AW: Zahlungserinnerung Rechnung {inv('hofbauer')}", "Hallo Frau Danzer,\n\nes tut mir leid, die Hochzeit ist teurer "
         "geworden als geplant. Wir zahlen den Rest in Raten, so schnell es geht. Ist das in Ordnung?\n\nLiebe Grüße\nVerena (Gruber)",
         reply="hof_er")
    mo = E["moser"]
    mail("moser_anf", "2026-08-19T15:10", "Andreas Moser <andreas.moser@webpost.example>", OWN_INFO, "Anfrage Gartenfest 03.10.",
         "Hallo,\n\nwir feiern am 03.10. ein Gartenfest bei uns in Obertraubling, ca. 45 Gäste. Können Sie ein Grillbuffet mit "
         "Getränken und 2 Servicekräften anbieten?\n\nGruß\nAndreas Moser")
    mail("moser_ang", "2026-08-20T17:45", OWN_BUERO, "Andreas Moser <andreas.moser@webpost.example>",
         f"Ihr Angebot {mo['an']} – Gartenfest 03.10.2026", "Hallo Herr Moser,\n\nanbei unser Angebot. Wir freuen uns auf Ihre "
         "Rückmeldung." + sig(), att=[(f"Angebot_{mo['an']}.pdf", ("offer", "moser"))], reply="moser_anf")
    mail("moser_ok", "2026-08-28T08:30", "Andreas Moser <andreas.moser@webpost.example>", OWN_BUERO,
         f"AW: Ihr Angebot {mo['an']} – Gartenfest 03.10.2026", "Hallo Frau Danzer,\n\npasst, wir machen das so. 45 Personen, "
         "evtl. ein paar mehr.\n\nGruß\nA. Moser", reply="moser_ang")
    mail("moser_ab", "2026-08-28T11:15", OWN_BUERO, "Andreas Moser <andreas.moser@webpost.example>",
         f"Auftragsbestätigung {mo['va']} Gartenfest 03.10.2026",
         f"Hallo Herr Moser,\n\nvielen Dank für Ihren Auftrag! Anbei die Auftragsbestätigung. Bitte überweisen Sie die Anzahlung "
         f"von {de(mo['dep_amt'])} € bis zum 11.09.2026." + sig(), att=[(f"Auftragsbestaetigung_{mo['va']}.pdf", ("conf", "moser"))],
         reply="moser_ok")
    mail("moser_dep", "2026-09-07T12:00", "Andreas Moser <andreas.moser@webpost.example>", OWN_BUERO, "Anzahlung überwiesen",
         f"Hallo,\n\nhabe heute die Anzahlung ({de(mo['dep_amt'])} €) überwiesen.\n\nGruß Moser")
    ze = E["zeitler"]
    mail("zeit_anf", "2026-08-24T20:14", "Patrick Zeitler <patrick.zeitler@donaumail.example>", OWN_INFO, "Anfrage Hochzeit 12.06.2027",
         "Hallo zusammen,\n\nwir heiraten am 12.06.2027 im Gutshof Irlbach (ca. 120 Gäste) und suchen ein Catering mit Buffet, "
         "Getränken nach Verbrauch und Service. Können Sie uns ein Angebot machen?\n\nViele Grüße\nPatrick Zeitler & Carolin Wittmann")
    mail("zeit_ang", "2026-08-28T16:20", OWN_BUERO, "Patrick Zeitler <patrick.zeitler@donaumail.example>",
         f"Angebot {ze['an']} Hochzeit 12.06.2027", "Hallo Herr Zeitler,\n\nanbei unser Angebot für Ihre Hochzeit. Gern laden wir "
         "Sie zu einem Probeessen ein." + sig(), att=[(f"Angebot_{ze['an']}.pdf", ("offer", "zeitler"))], reply="zeit_anf")
    pg = E["pongratz"]
    mail("pong", "2026-08-31T18:30", "Monika Pongratz <monika.pongratz@webpost.example>", OWN_INFO,
         "Absage 70. Geburtstag meiner Mutter am 26.09.", "Sehr geehrte Damen und Herren,\n\nmeine Mutter Gertraud Pongratz liegt "
         "im Krankenhaus, wir müssen die Feier am 26.09. leider absagen. Was kostet uns das jetzt?\n\nMit freundlichen Grüßen\nMonika Pongratz")
    mail("pong_re", "2026-09-01T09:14", OWN_BUERO, "Monika Pongratz <monika.pongratz@webpost.example>",
         "AW: Absage 70. Geburtstag meiner Mutter am 26.09.",
         f"Sehr geehrte Frau Pongratz,\n\nzunächst gute Besserung für Ihre Mutter! Nach Ziff. 8 unserer AGB berechnen wir bei einer "
         f"Absage 29–14 Tage vorher 50 % des Auftragswerts. Die Anzahlung von {de(pg['dep_amt'])} € ziehen wir ab; die Rechnung "
         "schicken wir Ihrer Mutter in den nächsten Tagen." + sig(), reply="pong")
    kl = E["klier"]
    mail("klier", "2026-09-01T17:02", "Steuerberatung Klier & Kollegen <kanzlei@stb-klier.de>", OWN_INFO,
         f"Ihr Angebot {kl['an']} – Auftrag", f"Sehr geehrte Damen und Herren,\n\nwir nehmen Ihr Angebot {kl['an']} vom "
         f"{dde(kl['offer'])} für unsere Feier am 02.10. an. Bitte geben Sie auf der Rechnung die Angebotsnummer an, das braucht "
         "unsere Buchhaltung.\n\nMit freundlichen Grüßen\nAndrea Klier")
    nw = next(o for o in OFFERS_LOST if o["key"] == "nwolf")
    mail("nwolf", "2026-09-04T19:30", "Nadine Wolf <nadine.wolf@donaumail.example>", OWN_INFO, "Anfrage Hochzeit 03.07.2027",
         "Hallo,\n\nwir planen unsere Hochzeit am 03.07.2027 mit ca. 130 Gästen. Haben Sie noch frei? Wir hätten gern ein "
         "Premium-Buffet.\n\nLG Nadine Wolf & Jakob Bachmeier")
    mail("nwolf_ang", "2026-09-07T11:05", OWN_BUERO, "Nadine Wolf <nadine.wolf@donaumail.example>", f"Angebot {nw['an']} Hochzeit 03.07.2027",
         "Hallo Frau Wolf,\n\nder Termin ist noch frei – anbei unser Angebot." + sig(),
         att=[(f"Angebot_{nw['an']}.pdf", ("offer_lost", "nwolf"))], reply="nwolf")
    mail("reindl", "2026-09-14T10:37", "Optik Reindl GmbH <info@optik-reindl-rgb.de>", OWN_INFO,
         "Anfrage Weihnachtsfeier 11.12. für ca. 40 Personen", "Guten Tag,\n\nwir suchen für unsere Weihnachtsfeier am Freitag, "
         "11.12.2026, ein warmes Buffet für ca. 40 Personen in unseren Geschäftsräumen. Bitte schicken Sie uns ein Angebot.\n\n"
         "Freundliche Grüße\nSandra Reindl")
    mail("stadt_hinweis", "2026-09-16T09:00", "Stadt Regensburg – Amt für Tagesbetreuung <kita-verpflegung@regensburg-stadt.de>",
         OWN_BUERO, "Hinweis zur Rechnungsstellung ab 01.10.2026", "Sehr geehrte Damen und Herren,\n\nab 01.10.2026 nehmen wir "
         "Rechnungen ausschließlich als XRechnung über das Rechnungseingangsportal entgegen. Bitte geben Sie immer die Leitweg-ID "
         "09362000-0042-35 an; Rechnungen ohne gültige Leitweg-ID werden automatisch abgewiesen.\n\nMit freundlichen Grüßen\n"
         "Amt für Tagesbetreuung von Kindern, Sachgebiet Verpflegung")
    eibl_dup = max((x for x in SUPINV if x["sup"] == "eibl" and x["date"] <= D("2026-09-01")), key=lambda x: x["date"])
    mail("eibl_re", "2026-09-02T07:48", "Metzgerei Eibl – Buchhaltung <buchhaltung@metzgerei-eibl.de>", OWN_BUERO,
         f"Rechnung {eibl_dup['no']}", f"Guten Morgen,\n\nanbei unsere Rechnung {eibl_dup['no']}.\n\nMit freundlichen Grüßen\n"
         "Metzgerei Eibl GmbH", att=[(f"Eibl_RE_{eibl_dup['no']}.pdf", ("sup", eibl_dup))])
    mail("ggm_news", "2026-09-10T06:00", "Gastro-Großmarkt Donau <newsletter@gastro-grossmarkt-donau.de>", OWN_INFO,
         "Herbstaktion: Wild, Kürbis & Co.", "Liebe Gastro-Profis,\n\nab sofort: Rehrücken aus heimischer Jagd, Hokkaido-Kürbis "
         "und neue Weine im Markt. Nur für Gewerbekunden.\n\nIhr Gastro-Großmarkt Donau\n\nAbmelden: newsletter@gastro-grossmarkt-donau.de")
    ro = next(o for o in OFFERS_LOST if o["key"] == "reindl")
    mail("reindl_ang", "2026-09-21T10:05", OWN_BUERO, "Optik Reindl GmbH <info@optik-reindl-rgb.de>",
         f"AW: Anfrage Weihnachtsfeier 11.12. – Angebot {ro['an']}", "Sehr geehrte Frau Reindl,\n\nvielen Dank für Ihre Anfrage, "
         "anbei unser Angebot." + sig(), att=[(f"Angebot_{ro['an']}.pdf", ("offer_lost", "reindl"))], reply="reindl")
    mail("nordwand", "2026-09-22T13:40", "Jonas Pfeiffer – Nordwand Software <office@nordwand-software.de>", OWN_INFO,
         "Lunch-Abo ab Oktober – 30 Personen", "Hallo Herr Hollerieth,\n\nwir sind ab Oktober 30 Leute im Büro. Können Sie das "
         "Lunch-Abo entsprechend erweitern?\n\nViele Grüße\nJonas Pfeiffer\nOffice Management")
    mail("nordwand_re", "2026-09-23T09:10", OWN_TH, "Jonas Pfeiffer – Nordwand Software <office@nordwand-software.de>",
         "AW: Lunch-Abo ab Oktober – 30 Personen", "Hallo Herr Pfeiffer,\n\ngern! Ab 01.10.2026 beträgt die Pauschale für "
         "30 Personen 4.300,00 € netto im Monat, sonst bleibt alles wie gehabt.\n\nViele Grüße\nTobias Hollerieth", reply="nordwand")
    sp = max((x for x in SUPINV if x["sup"] == "serviceplus" and x["date"] <= D("2026-10-01")), key=lambda x: x["date"])
    mail("sp_xr", "2026-10-01T07:12", "ServicePlus Personal – Rechnungswesen <rechnungen@serviceplus-personal.de>", OWN_BUERO,
         f"Rechnung {sp['no']} (XRechnung)", f"Sehr geehrte Damen und Herren,\n\nanbei erhalten Sie unsere Rechnung {sp['no']} "
         "als XRechnung.\n\nMit freundlichen Grüßen\nServicePlus Personal GmbH", att=[(f"{sp['no']}.xml", ("xr_sup", sp))])
    sep_b = mv_inv("MV-02", 9)
    mail("stadt_ab", "2026-10-07T11:23", "Rechnungseingang Stadt Regensburg <rechnungseingang@regensburg-stadt.de>", OWN_BUERO,
         f"Abweisung Ihrer Rechnung {sep_b['no']} – Leitweg-ID fehlerhaft",
         f"Sehr geehrte Damen und Herren,\n\nIhre XRechnung {sep_b['no']} vom {dde(sep_b['date'])} wurde abgewiesen: Die Leitweg-ID "
         "09362000-0042-53 ist keinem Rechnungsempfänger zugeordnet. Bitte reichen Sie die Rechnung mit der Leitweg-ID "
         "09362000-0042-35 erneut ein. Die Zahlungsfrist beginnt mit dem Eingang der korrekten Rechnung.\n\n"
         "Diese Nachricht wurde automatisch erstellt.")
    mail("stadt_korr", "2026-10-12T08:40", OWN_BUERO, "Rechnungseingang Stadt Regensburg <rechnungseingang@regensburg-stadt.de>",
         f"Korrigierte XRechnung {sep_b['no']}", f"Sehr geehrte Damen und Herren,\n\nanbei die korrigierte XRechnung {sep_b['no']} "
         "mit der richtigen Leitweg-ID." + sig(), att=[(f"XRechnung_{sep_b['no']}_korrigiert.xml", ("xr_own", sep_b))], reply="stadt_ab")
    mail("kuend", "2026-10-12T14:02", "Kath. Kirchenstiftung St. Wolfgang <kita.st-wolfgang@pfarrei-stwolfgang-rgb.de>", OWN_BUERO,
         "Kündigung Mittagsverpflegung Kita St. Wolfgang", "Sehr geehrte Damen und Herren,\n\nhiermit kündigen wir den "
         "Verpflegungsvertrag für die Kita St. Wolfgang fristgerecht zum 31.12.2026. Ab Januar kocht unsere neue Küche im "
         "Pfarrzentrum selbst. Für die gute Zusammenarbeit danken wir Ihnen herzlich.\n\nBitte bestätigen Sie den Eingang.\n\n"
         "Mit freundlichen Grüßen\nKirchenverwaltung St. Wolfgang\ni. A. Hildegard Brunnhuber")
    pf, lx = E["pflaum_xmas"], E["lindbuechl_xmas"]
    mail("pflaum_ar", "2026-10-01T15:15", OWN_BUERO, "Ingenieurbüro Pflaum <office@ib-pflaum.de>",
         f"Anzahlungsrechnung {pf['ar']['no']} Weihnachtsfeier 04.12.", "Sehr geehrte Damen und Herren,\n\nanbei die "
         "Anzahlungsrechnung zu Ihrer Weihnachtsfeier." + sig(), att=[(f"Anzahlungsrechnung_{pf['ar']['no']}.pdf", ("pdf_own", pf["ar"]))])
    mail("lind_ar", "2026-10-12T09:05", OWN_BUERO, "Autohaus Lindbüchl – Buchhaltung <buchhaltung@autohaus-lindbuechl.de>",
         f"Anzahlungsrechnung {lx['ar']['no']} Weihnachtsfeier 12.12.", "Sehr geehrte Damen und Herren,\n\nanbei die "
         "Anzahlungsrechnung für Ihre Weihnachtsfeier am 12.12.2026." + sig(),
         att=[(f"Anzahlungsrechnung_{lx['ar']['no']}.pdf", ("pdf_own", lx["ar"]))])
    sz = E["sfz_xmas"]
    mail("sfz_auftrag", "2026-10-14T16:30", "Sportfreunde Ziegetsdorf – Kasse <kasse@sf-ziegetsdorf.de>", OWN_INFO,
         "Weihnachtsfeier 18.12. – Auftrag", f"Hallo,\n\nwir nehmen Ihr Angebot {sz['an']} für unsere Weihnachtsfeier am 18.12. an "
         "(70 Personen). Die Anzahlung von 50 % überweisen wir nächste Woche.\n\nSportliche Grüße\nThomas Kerschl, Kassier")
    mail("lind_m2", "2026-10-15T10:00", OWN_BUERO, "Autohaus Lindbüchl – Buchhaltung <buchhaltung@autohaus-lindbuechl.de>",
         f"2. Mahnung Rechnung {inv('lindbuechl_sommer')}", f"Sehr geehrte Damen und Herren,\n\ntrotz unserer Erinnerung ist der "
         f"Restbetrag von {de(lb['inv']['pay'] - lb_paid)} € aus Rechnung {inv('lindbuechl_sommer')} weiterhin offen. Bitte "
         "überweisen Sie ihn bis zum 29.10.2026." + sig())
    mail("eventwerk_ziel", "2026-10-14T11:48", "Eventwerk Donau – Rechnungsprüfung <rechnung@eventwerk-donau.de>", OWN_BUERO,
         f"Ihre Rechnung {inv('eventwerk_sep')}", f"Sehr geehrte Frau Danzer,\n\nwir zahlen Ihre Rechnung {inv('eventwerk_sep')} "
         "wie gewohnt gemäß unseren Einkaufsbedingungen 45 Tage nach Rechnungseingang.\n\nMit freundlichen Grüßen\nEventwerk Donau GmbH")
    mail("zeit_ok", "2026-10-19T21:07", "Patrick Zeitler <patrick.zeitler@donaumail.example>", OWN_BUERO,
         f"AW: Angebot {ze['an']} Hochzeit 12.06.2027", "Hallo Frau Danzer,\n\nwir haben uns entschieden: Wir nehmen Ihr Angebot "
         "an! Die Anzahlung überweisen wir nächste Woche.\n\nViele Grüße\nPatrick & Carolin", reply="zeit_ang")
    mail("sfz_dep", "2026-10-22T18:11", "Sportfreunde Ziegetsdorf – Kasse <kasse@sf-ziegetsdorf.de>", OWN_INFO,
         "Anzahlung Weihnachtsfeier überwiesen", f"Hallo,\n\nheute sind 50 % = {de(sz['dep_amt'])} € an Sie rausgegangen.\n\n"
         "Gruß Thomas Kerschl")
    ggm_last = max((x for x in SUPINV if x["sup"] == "ggm" and x["date"] <= D("2026-10-21")), key=lambda x: x["date"])
    mail("fraud", "2026-10-21T16:12", f"Gastro-Großmarkt Donau – Buchhaltung <{FRAUD['mail']}>", OWN_BUERO,
         "Wichtig: Änderung unserer Bankverbindung", "Sehr geehrte Kundin, sehr geehrter Kunde,\n\naufgrund einer Umstellung "
         "unseres Zahlungsverkehrs ziehen wir ab sofort keine Lastschriften mehr ein. Bitte überweisen Sie die beigefügte, "
         f"korrigierte Rechnung {ggm_last['no']} umgehend auf unser neues Konto:\n\nIBAN {fmt_iban(FRAUD['iban'])}\nBIC {FRAUD['bic']}\n\n"
         "Mit freundlichen Grüßen\nBuchhaltung Gastro-Großmarkt Donau GmbH",
         att=[(f"Rechnung_{ggm_last['no']}_korrigiert.pdf", ("fraud", ggm_last))])
    late = next(x for x in SUPINV if x["sup"] == "eibl" and x["paid"] == D("2026-10-21"))
    mail("eibl_er", "2026-10-23T08:02", "Metzgerei Eibl – Buchhaltung <buchhaltung@metzgerei-eibl.de>", OWN_BUERO,
         f"Zahlungserinnerung Rechnung {late['no']}", f"Sehr geehrte Damen und Herren,\n\nunsere Rechnung {late['no']} vom "
         f"{dde(late['date'])} über {de(late['gross'])} € ist seit {dde(late['due'])} fällig. Bitte überweisen Sie den Betrag "
         "umgehend. Sollten Sie bereits gezahlt haben, betrachten Sie dieses Schreiben als gegenstandslos.\n\nMetzgerei Eibl GmbH")
    mail("moser_doppelt", "2026-10-26T19:48", "Andreas Moser <andreas.moser@webpost.example>", OWN_BUERO, "Gartenfest – doppelt bezahlt?",
         f"Hallo Frau Danzer,\n\nich glaube, ich habe bei der Rechnung {mo['inv']['no']} die Anzahlung vergessen abzuziehen und den "
         f"vollen Betrag überwiesen. Können Sie mir die {de(mo['dep_amt'])} € bitte zurücküberweisen?\n\nDanke und Gruß\nAndreas Moser")
    mail("obermeier", "2026-10-27T12:30", "Julia Obermeier <julia.obermeier@webpost.example>", OWN_INFO, "Anfrage Hochzeit 19.06.2027",
         "Hallo liebes Hollerieth-Team,\n\nwir heiraten am 19.06.2027 und hätten gern ein Angebot für ca. 90 Gäste "
         "(Menü oder Buffet). Habt ihr an dem Tag noch Kapazität?\n\nLiebe Grüße\nJulia Obermeier")
    mail("lind_150", "2026-10-28T15:55", "Autohaus Lindbüchl – Petra Lindbüchl <p.lindbuechl@autohaus-lindbuechl.de>", OWN_TH,
         "Weihnachtsfeier 12.12. – jetzt 150 Personen", "Hallo Herr Hollerieth,\n\nwir sind zur Weihnachtsfeier jetzt 150 "
         "statt 120 Personen. Geht das noch? Die Anzahlung überweisen wir, sobald das geklärt ist.\n\nGruß\nPetra Lindbüchl")
    return sorted(M, key=lambda m: m["when"])


def write_eml(path: Path, m: dict, ids: dict) -> str:
    dt = datetime.fromisoformat(m["when"])
    tzh = 1 if dt.date() >= D("2026-10-25") else 2
    dt = dt.replace(tzinfo=timezone(timedelta(hours=tzh)))
    msg = EmailMessage()
    msg["From"] = m["frm"]
    msg["To"] = m["to"]
    msg["Date"] = format_datetime(dt)
    msg["Subject"] = m["subject"]
    dom = m["frm"].split("@")[-1].strip(">")
    mid = f"<{hashlib.md5(m['key'].encode()).hexdigest()[:20]}@{dom}>"
    msg["Message-ID"] = mid
    if m["reply"] and m["reply"] in ids:
        msg["In-Reply-To"] = ids[m["reply"]]
        msg["References"] = ids[m["reply"]]
    msg["MIME-Version"] = "1.0"
    msg.set_content(m["body"])
    for name, (kind, obj) in m["att"]:
        data = attachment_bytes(kind, obj)
        if name.endswith(".xml"):
            msg.add_attachment(data, maintype="application", subtype="xml", filename=name)
        else:
            msg.add_attachment(data, maintype="application", subtype="pdf", filename=name)
    if m["att"]:
        msg.set_boundary("==" + hashlib.md5((m["key"] + "|b").encode()).hexdigest()[:24] + "==")
    path.write_bytes(bytes(msg))
    return mid


def attachment_bytes(kind: str, obj) -> bytes:
    if kind == "pdf_own":
        return pdf_own(obj)
    if kind == "sup":
        return pdf_supplier(obj)
    if kind == "fraud":
        return pdf_supplier(obj, fraud=True)
    if kind == "xr_sup":
        return xr_supplier(obj)
    if kind == "xr_own":
        return xr_own(obj)
    if kind in ("offer", "conf"):
        e = EVD[obj]
        conf = None
        if kind == "conf":
            conf = dict(va=e["va"], pct=e["dep"]["pct"], amount=e["dep_amt"], until=D("2026-09-11"))
        return pdf_offer(e, e["an"], e["kd"], e["offer"] if kind == "offer" else e["booked"], e["offer_lines"], e["date"],
                         e["anlass"], conf)
    if kind == "offer_lost":
        o = next(o for o in OFFERS_LOST if o["key"] == obj)
        fake = dict(food=[("M", o["food"])], drinks=o["drinks"], staff=o["staff"], log=o["log"], rental=[], key=o["key"])
        return pdf_offer(o, o["an"], o["kd"], o["date"], ev_lines(fake, o["pers"], offer=True), o["when"], o["anlass"])
    raise ValueError(kind)


def slug(s: str) -> str:
    tr = {"ä": "ae", "ö": "oe", "ü": "ue", "Ä": "Ae", "Ö": "Oe", "Ü": "Ue", "ß": "ss"}
    s = "".join(tr.get(ch, ch) for ch in s)
    return "".join(ch if ch.isalnum() else "_" for ch in s).strip("_")[:60]


def write_mails(folder: Path, dv: dict):
    ids = {}
    for m in MAILS:
        if datetime.fromisoformat(m["when"]).date() > dv["cut"]:
            continue
        ids[m["key"]] = write_eml(folder / f"{m['when'][:10]}_{slug(m['subject'])}.eml", m, ids)


# ------------------------------------------------------------------ the deliveries
OWN_PDF_EVENTS = ["wimmeder", "lindbuechl_sommer", "lindbuechl_xmas", "tracht_jub", "pflaum_xmas", "seidl_sommer", "pongratz",
                  "eventwerk_jul", "eventwerk_sep", "moser", "aigner", "mvz", "klier", "lehner", "kellner", "berger"]
SUP_FILE = {"ggm": "GGM_Sammelrechnung_{no}.pdf", "eibl": "Eibl_RE_{no}.pdf", "poeschl": "Poeschl_Getraenke_{no}.pdf",
            "serviceplus": "ServicePlus_{no}.xml", "blank": "Waescherei_Blank_{no}.pdf", "pfeffer": "Zeltverleih_Pfeffer_{no}.pdf"}
MAILS: list[dict] = []


def write_delivery(root: Path, dv: dict) -> list[str]:
    out, cut = root / dv["folder"], dv["cut"]
    for sub in ("Faktura", "Bank", "Kueche", "Mails", "Rechnungen_Ausgang/XRechnung_Stadt", "Eingangsrechnungen"):
        (out / sub).mkdir(parents=True, exist_ok=True)
    write_rechnungsliste(out / "Faktura" / f"Rechnungsliste_Export_{cut.isoformat()}.csv", dv)
    write_kontakte(out / "Faktura" / f"Kontakte_Export_{cut.isoformat()}.csv", dv)
    write_artikel(out / "Faktura" / f"Artikel_Export_{cut.isoformat()}.csv", dv)
    write_veranstaltungen(out / "Veranstaltungen_2026.xlsx", dv)
    write_angebote(out / "Angebote_2025-2026.xlsx", dv)
    write_vertraege(out / "Vertraege_Verpflegung.xlsx", dv)
    write_essenszahlen(out / "Kueche" / "Essenszahlen_2026.xlsx", dv)
    write_kassenbuch(out / "Kassenbuch_2026.xlsx", dv)
    write_camt(out / "Bank", dv)
    docs = []
    for k in OWN_PDF_EVENTS:
        e = EVD[k]
        docs += [x for x in (e.get("ar"), e.get("wrong_inv"), e.get("sto"), e.get("inv"), e.get("stk"), e.get("gs")) if x]
    docs += [i for i in INV if i["src"].get("contract") == "RV-01" and i["date"] >= D("2026-07-01")]
    for i in docs:
        if i["date"] > cut:
            continue
        sub = out / "Rechnungen_Ausgang" / i["date"].strftime("%Y-%m")
        sub.mkdir(exist_ok=True)
        (sub / f"{i['no']}_{slug(CUSTOMERS[i['kd']]['name'])[:30]}.pdf").write_bytes(pdf_own(i))
    sep_b = mv_inv("MV-02", 9)
    for i in INV:
        if i["src"].get("contract") in ("MV-01", "MV-02") and i["date"] <= cut:
            wrong = "09362000-0042-53" if i is sep_b else None
            (out / "Rechnungen_Ausgang" / "XRechnung_Stadt" / f"XRechnung_{i['no']}.xml").write_bytes(xr_own(i, wrong))
    if cut >= D("2026-10-12"):
        (out / "Rechnungen_Ausgang" / "XRechnung_Stadt" / f"XRechnung_{sep_b['no']}_korrigiert.xml").write_bytes(xr_own(sep_b))
    for x in SUPINV:
        if FILES_FROM <= x["date"] <= cut:
            sub = out / "Eingangsrechnungen" / x["date"].strftime("%Y-%m")
            sub.mkdir(exist_ok=True)
            data = xr_supplier(x) if x["fmt"] == "xrechnung" else pdf_supplier(x)
            (sub / SUP_FILE[x["sup"]].format(no=x["no"])).write_bytes(data)
    write_mails(out / "Mails", dv)
    return sorted(str(p.relative_to(out)) for p in out.rglob("*") if p.is_file())


def build_all():
    build_contract_invoices()
    build_event_invoices()
    number_all()
    build_money()
    run_money()
    build_suppliers()
    build_expenses()
    build_cash()
    MAILS.extend(build_mails())


def main(outdir: str) -> None:
    root = Path(outdir)
    root.mkdir(parents=True, exist_ok=True)
    build_all()
    truth = dict(firm="Hollerieth Catering & Events GmbH (Catering und Veranstaltungen, Review 9)", deliveries=[])
    files = {}
    for dv in DELIV:
        files[dv["folder"]] = write_delivery(root, dv)
        truth["deliveries"].append(truth_for(dv))
    check_traps(truth)
    (root / "truth.json").write_text(json.dumps(truth, ensure_ascii=False, indent=1), encoding="utf-8")
    key = dict(events={e["va"]: dict(key=e["key"], kd=e["kd"], anlass=e["anlass"], date=e["date"].isoformat(), offer=e["an"],
                                     booked=e["booked"].isoformat(), invoice=(final_doc(e) or {}).get("no"),
                                     deposit=eur(e.get("dep_amt", 0)) or None) for e in EVENTS},
               invoices={i["no"]: dict(kind=KIND_LABEL[i["kind"]], kd=i["kd"], date=i["date"].isoformat(), gross=eur(i["gross"]),
                                       pay=eur(i["pay"]), due=i["due"].isoformat(), ref=i["ref"], subject=i["subject"])
                         for i in INV if i["no"]},
               money=[dict(date=m["date"].isoformat(), amount=eur(m["amount"]), ch=m["ch"], kind=m["kind"], payer=m["payer"],
                           vwz=m["vwz"], alloc=[(i["no"], eur(c)) for i, c in m["alloc"]], over=eur(m["over"]),
                           deposit_for=EVD[m["dep_for"]]["va"] if m["dep_for"] else None, note=m["note"])
                      for m in MONEY if m["kind"] in ("cust", "refund", "own", "other_in")],
               lieferscheine={x["no"]: dict(date=x["date"].isoformat(), contract=x["contract"], what=x["what"], net=eur(x["net"]),
                                            invoice=x["inv"]["no"] if x["inv"] else None) for x in LS},
               files=files)
    (root / "answer_key.json").write_text(json.dumps(key, ensure_ascii=False, indent=1), encoding="utf-8")
    for d in truth["deliveries"]:
        print(d["folder"], d["today"], "offen:", len(d["open"]), d["open_total"], "überfällig:", len(d["overdue"]),
              "nicht berechnet:", d["unbilled"], "Geld ohne Rechnung:", d["receipts"])


def check_traps(truth: dict):
    """The header promises these; if a change breaks one, the generator says so."""
    l1, l2 = truth["deliveries"]
    want = {1: {"R-2026-0043": 3464.03, "R-2026-0070": 2548.46, "R-2026-0089": 1640.70, "R-2026-0101": 4189.05,
                "R-2026-0104": 1593.39, "R-2026-0105": 2983.91, "R-2026-0106": 976.32, "R-2026-0107": 250.34,
                "R-2026-0108": 949.24, "R-2026-0109": 6146.14, "R-2026-0110": 1491.58, "R-2026-0111": 7947.40},
            2: {"R-2026-0043": 3464.03, "R-2026-0070": 1548.46, "R-2026-0089": 1640.70, "R-2026-0115": 9702.73,
                "R-2026-0122": 3743.34, "R-2026-0124": 584.80, "R-2026-0127": 2501.71, "R-2026-0131": 3055.60,
                "R-2026-0133": 997.78, "R-2026-0134": 7367.98, "R-2026-0135": 7511.52}}
    for d, w in ((l1, want[1]), (l2, want[2])):
        assert d["open"] == w, (d["folder"], d["open"])
    assert l1["overdue"] == ["R-2026-0043", "R-2026-0070", "R-2026-0089", "R-2026-0101", "R-2026-0106", "R-2026-0107",
                             "R-2026-0108"], l1["overdue"]
    assert l2["overdue"] == ["R-2026-0043", "R-2026-0070", "R-2026-0089", "R-2026-0115", "R-2026-0124", "R-2026-0131"], l2["overdue"]
    assert l1["unbilled"] == ["VA-26-042", "VA-26-045", "MV-06"], l1["unbilled"]
    assert l2["unbilled"] == ["VA-26-050", "VA-26-052", "LS-26-0031", "LS-26-0032", "LS-26-0034", "LS-26-0042", "LS-26-0045",
                              "LS-26-0047"], l2["unbilled"]
    assert l1["receipts"] == [1010.0, 3060.0, 4040.0] and l2["receipts"] == [1010.0, 1830.0, 3630.0]
    assert [p["invoice"] for p in l1["payables"]] == ["GD-26-018248", "2026-1031"]
    assert [p["invoice"] for p in l2["payables"]] == ["GD-26-018254", "2026-1036"]
    e = EVD
    assert (e["bauer"]["inv"]["no"], e["wimmeder"]["gs"]["no"], e["seidl_sommer"]["wrong_inv"]["no"], e["seidl_sommer"]["sto"]["no"],
            e["seidl_sommer"]["inv"]["no"], e["tracht_jub"]["ar"]["no"], e["pflaum_xmas"]["ar"]["no"], e["klier"]["an"],
            e["den_herbst"]["va"], e["wenzl_probe"]["va"], e["zeitler"]["va"]) == (
        "R-2026-0064", "GS-2026-002", "R-2026-0080", "GS-2026-001", "R-2026-0083", "R-2026-0084", "R-2026-0120", "AN-26-038",
        "VA-26-050", "VA-26-039", "VA-27-001")
    assert (mv_inv("MV-06", 8)["no"], mv_inv("MV-06", 9)["no"]) == ("R-2026-0128", "R-2026-0126")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit("Aufruf: python3 catering.py OUTDIR")
    main(sys.argv[1])
