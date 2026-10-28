#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Testdatensatz "Brenner Haustechnik GmbH" – SHK-Handwerksbetrieb (Sanitär/Heizung), 8 Personen, Dresden.

Erzeugt unter <base>:
  daten/        Lieferung 1 (der "Datendump")
  lieferung2/   spätere Lieferung derselben Dateien mit geänderten Werten
  kaputt/       Dateien mit kaputtem Inhalt
  wahrheit.json die Wahrheit (Entitäten, Dubletten, Namensvettern, Beziehungen, Formeln, Auffälligkeiten)

Aufruf: python3 generate_brenner.py <base>
"""
import csv, io, json, random, sys, zipfile, datetime as dt
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from collections import defaultdict, Counter
from email.message import EmailMessage
from email.utils import format_datetime
from email import policy

import xlsxwriter
from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from reportlab.lib.units import mm

import os
UMGEHUNG = os.environ.get("UMGEHUNG") == "1"   # Variante ohne die beiden Absturz-Auslöser
BASE = Path(sys.argv[1]).resolve()
OUT, L2, KAPUTT = BASE / "daten", BASE / "lieferung2", BASE / "kaputt"
for p in (OUT, OUT / "Mails", L2, KAPUTT):
    p.mkdir(parents=True, exist_ok=True)

R = random.Random(20261028)
EXPORT = dt.date(2026, 10, 15)           # Stand der Exporte
BANK_FROM, BANK_TO = dt.date(2026, 1, 1), dt.date(2026, 9, 30)
TZ = dt.timezone(dt.timedelta(hours=2))

TRUTH = {"firma": "Brenner Haustechnik GmbH (SHK), Hechtstraße 41, 01097 Dresden",
         "dateien": {}, "entitaeten": {}, "dubletten": [], "namensvettern": [],
         "beziehungen": [], "formeln": [], "nummernkreise": {}, "auffaellig": {},
         "werte_fallen": [], "lieferung2": {}, "kaputt": {}}


# ---------------------------------------------------------------- Hilfen
def q2(x):
    return Decimal(str(x)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def de(x):
    s = f"{q2(x):,.2f}"
    return s.replace(",", "§").replace(".", ",").replace("§", ".")


def dd(d):
    return d.strftime("%d.%m.%Y")


HOL = {dt.date(2025, 1, 1), dt.date(2025, 4, 18), dt.date(2025, 4, 21), dt.date(2025, 5, 1), dt.date(2025, 5, 29),
       dt.date(2025, 6, 9), dt.date(2025, 10, 3), dt.date(2025, 10, 31), dt.date(2025, 11, 19),
       dt.date(2025, 12, 25), dt.date(2025, 12, 26), dt.date(2026, 1, 1), dt.date(2026, 4, 3), dt.date(2026, 4, 6),
       dt.date(2026, 5, 1), dt.date(2026, 5, 14), dt.date(2026, 5, 25), dt.date(2026, 11, 18),
       dt.date(2026, 12, 25), dt.date(2026, 12, 26), dt.date(2024, 1, 1), dt.date(2024, 3, 29), dt.date(2024, 4, 1),
       dt.date(2024, 5, 1), dt.date(2024, 5, 9), dt.date(2024, 5, 20), dt.date(2024, 10, 3), dt.date(2024, 10, 31),
       dt.date(2024, 11, 20), dt.date(2024, 12, 25), dt.date(2024, 12, 26)}


def wd(d):
    return d.weekday() < 5 and d not in HOL


def next_wd(d):
    while not wd(d):
        d += dt.timedelta(days=1)
    return d


def add_wd(d, n):
    while n > 0:
        d += dt.timedelta(days=1)
        if wd(d):
            n -= 1
    return d


def wd_range(a, b):
    out, d = [], a
    while d <= b:
        if wd(d):
            out.append(d)
        d += dt.timedelta(days=1)
    return out


def iban(blz, konto):
    bban = f"{blz}{konto:010d}"
    return f"DE{98 - int(bban + '131400') % 97:02d}{bban}"


BLZS = ["85050300", "85090000", "12030000", "50010517", "86010090", "85040000", "85020086"]


def rand_iban():
    return iban(R.choice(BLZS), R.randrange(10 ** 8, 10 ** 10))


OWN_IBAN = iban("85090000", 4471230)

# ---------------------------------------------------------------- Mitarbeiter
EMP = [
    dict(pnr=101, vorname="Thomas", nachname="Brenner", kz="TB", position="Geschäftsführer / SHK-Meister",
         dept="Management", hire=dt.date(2009, 4, 1), hours=45, mobile="+49 172 3301101", netto=3800),
    dict(pnr=102, vorname="Sabine", nachname="Brenner", kz="SB", position="Office / Accounting",
         dept="Administration", hire=dt.date(2009, 4, 1), hours=30, mobile="+49 172 3301102", netto=1650),
    dict(pnr=103, vorname="Jens", nachname="Kowalski", kz="JK", position="Kundendiensttechniker",
         dept="Service", hire=dt.date(2012, 8, 15), hours=40, mobile="+49 172 3301103", netto=2450),
    dict(pnr=104, vorname="Marco", nachname="Richter", kz="MR", position="Anlagenmechaniker SHK",
         dept="Installation", hire=dt.date(2016, 2, 1), hours=40, mobile="+49 172 3301104", netto=2280),
    dict(pnr=105, vorname="Dennis", nachname="Hoffmann", kz="DH", position="Anlagenmechaniker SHK",
         dept="Installation", hire=dt.date(2019, 9, 1), hours=40, mobile="+49 172 3301105", netto=2190),
    dict(pnr=106, vorname="Ayşe", nachname="Yılmaz", kz="AY", position="Anlagenmechanikerin SHK",
         dept="Installation", hire=dt.date(2021, 3, 15), hours=40, mobile="+49 172 3301106", netto=2190),
    dict(pnr=107, vorname="Kevin", nachname="Schmidt", kz="KS", position="Auszubildender (3. Lj.)",
         dept="Installation", hire=dt.date(2023, 8, 28), hours=40, mobile="+49 152 5507107", netto=980),
    dict(pnr=108, vorname="Lukas", nachname="Weber", kz="LW", position="Auszubildender (1. Lj.)",
         dept="Installation", hire=dt.date(2025, 9, 1), hours=40, mobile="+49 152 5507108", netto=820),
]
EMP_BY_KZ = {e["kz"]: e for e in EMP}
for e in EMP:
    e["email"] = f"{e['vorname'].lower().replace('ş', 's')}.{e['nachname'].lower().replace('ı', 'i')}@brenner-haustechnik.de"

# ---------------------------------------------------------------- Kunden
CUST = []
BY_NR = {}


def add_c(**kw):
    c = dict(objekte=[], bankname=None, variants=None, mobil=None, email=None, ansprech="", bem="",
             skonto=False, vorname=None, nachname=None, anrede="Firma", true_id=None)
    c.update(kw)
    c["iban"] = rand_iban()
    if not c["true_id"]:
        c["true_id"] = f"K{c['kdnr']}"
    CUST.append(c)
    BY_NR[c["kdnr"]] = c
    return c


# Spezialfälle (Dubletten, Namensvettern) -------------------------------------------------
WG_ELBTAL = add_c(kdnr=10005, gruppe="WG", name="Wohnungsgenossenschaft Elbtal eG", ansprech="Hr. Neumann (Technik)",
                  strasse="Kesselsdorfer Straße", hnr="161", plz="01169", ort="Dresden", tel="+49 351 4102030",
                  email="technik@wg-elbtal.de", ziel=30, seit=dt.date(2011, 1, 1),
                  bem="Rahmenvertrag Wartung, Rechnung je Objekt!", bankname="WOHNUNGSGENOSSENSCHAFT ELBTAL EG",
                  objekte=["Amalie-Dietrich-Platz 3", "Leutewitzer Ring 12", "Harthaer Straße 7",
                           "Wölfnitzer Ring 40", "Omsewitzer Ring 2", "Merianplatz 9"],
                  variants=[("Wohnungsgenossenschaft Elbtal eG", .9), ("WG Elbtal eG", .7),
                            ("Wohnungsgen. Elbtal", .4), ("WG Elbtal", .5)])
HV_SCHULZE = add_c(kdnr=10008, gruppe="HV", name="Hausverwaltung Schulze & Partner GbR",
                   ansprech="Fr. Schulze-Wendt", strasse="Schandauer Straße", hnr="34", plz="01277", ort="Dresden",
                   tel="+49 351 3104455", email="technik@hv-schulze-partner.de", ziel=21, seit=dt.date(2012, 5, 1),
                   bankname="HAUSVERWALTUNG SCHULZE UND PARTNER",
                   objekte=["Altgorbitzer Ring 21", "Gompitzer Höhe 4", "Kesselsdorfer Straße 210",
                            "Tharandter Straße 55", "Reisewitzer Straße 18"],
                   variants=[("Hausverwaltung Schulze & Partner GbR", .9), ("HV Schulze & Partner", .6),
                             ("Schulze & Partner", .4), ("HV Schulze u. Partner", .5)])
HOFMANN = add_c(kdnr=10014, gruppe="P", anrede="Herr", vorname="Klaus", nachname="Hofmann", name="Hofmann, Klaus",
                strasse="Borsbergstraße", hnr="22", plz="01309", ort="Dresden", tel="+49 351 3361208",
                email="k.hofmann@gmx.de", ziel=14, seit=dt.date(2012, 9, 12), bankname="KLAUS HOFMANN")
MUELLER_G = add_c(kdnr=10023, gruppe="G", name="Müller GmbH & Co. KG", ansprech="Hr. Wagner (Buchhaltung)",
                  strasse="Großenhainer Straße", hnr="101", plz="01127", ort="Dresden", tel="+49 351 8494420",
                  email="buchhaltung@mueller-getraenke.de", ziel=14, seit=dt.date(2014, 3, 1),
                  bem="Getränkegroßhandel, Lager + Büro", bankname="MUELLER GMBH + CO. KG",
                  variants=[("Müller GmbH & Co. KG", .9), ("Fa. Müller (Getränke)", 0.0),
                            ("Müller GmbH & Co.KG", .7), ("Mueller GmbH", 0.0)])
KRUEGER = add_c(kdnr=10031, gruppe="P", anrede="Herr", vorname="Andreas", nachname="Krüger",
                name="Krüger, Andreas (Dr.)", strasse="Bergstraße", hnr="61", plz="01069", ort="Dresden",
                tel="+49 351 4718823", email="andreas.krueger@t-online.de", ziel=14, seit=dt.date(2015, 4, 20),
                bankname="DR. ANDREAS KRUEGER",
                variants=[("Krüger, Andreas (Dr.)", .9), ("Dr. Krüger", .3), ("A. Krueger", .0),
                          ("Dr. Andreas Krüger", .6)])
LEHMANN = add_c(kdnr=10040, gruppe="G", name="Zahnarztpraxis Dr. Lehmann", ansprech="Fr. Dr. Lehmann",
                strasse="Wiener Straße", hnr="33", plz="01219", ort="Dresden", tel="+49 351 4796610",
                email="praxis@za-lehmann.de", ziel=14, skonto=True, seit=dt.date(2016, 1, 11),
                bankname="ZAHNARZTPRAXIS DR. LEHMANN",
                variants=[("Zahnarztpraxis Dr. Lehmann", .9), ("Praxis Dr. med. dent. Lehmann", .3),
                          ("ZA-Praxis Lehmann", .2)])
MUELLER_BAU = add_c(kdnr=10045, gruppe="G", name="Müller Bau GmbH", ansprech="Hr. Müller jun.",
                    strasse="Radeburger Straße", hnr="5", plz="01109", ort="Dresden", tel="+49 351 8803311",
                    email="info@mueller-bau-dd.de", ziel=30, seit=dt.date(2016, 6, 1), bem="Zahlt schleppend!",
                    bankname="MUELLER BAU GMBH",
                    variants=[("Müller Bau GmbH", .9), ("Müller-Bau", .0), ("Fa. Müller Bau", .5)])
SCHULZE_G = add_c(kdnr=10052, gruppe="P", anrede="Frau", vorname="Gerda", nachname="Schulze", name="Schulze, Gerda",
                  strasse="Hüblerstraße", hnr="8", plz="01309", ort="Dresden", tel="+49 351 3108877", ziel=14,
                  seit=dt.date(2017, 2, 2), bankname="GERDA SCHULZE",
                  variants=[("Schulze, Gerda", .9), ("Fr. Schulze", .5)])
SCHMIDT_P = add_c(kdnr=10061, gruppe="P", anrede="Herr", vorname="Peter", nachname="Schmidt", name="Schmidt, Peter",
                  strasse="Leipziger Straße", hnr="12", plz="01097", ort="Dresden", tel="+49 351 8015544",
                  email="p.schmidt@web.de", ziel=14, seit=dt.date(2017, 9, 5), bankname="PETER SCHMIDT")
ELEKTRO_M = add_c(kdnr=10067, gruppe="G", name="Elektro Müller e.K.", ansprech="Inh. Frank Müller",
                  strasse="Pirnaer Landstraße", hnr="88", plz="01237", ort="Dresden", tel="+49 351 2563340",
                  email="kontakt@elektro-mueller-dresden.de", ziel=14, seit=dt.date(2018, 3, 3),
                  bankname="ELEKTRO MUELLER E.K.",
                  variants=[("Elektro Müller e.K.", .9), ("Müller Elektro", .3), ("Elektro Müller", .6)])
BAECKER_H = add_c(kdnr=10070, gruppe="G", name="Bäckerei Hoffmann", ansprech="Fr. Hoffmann",
                  strasse="Königsbrücker Straße", hnr="56", plz="01099", ort="Dresden", tel="+49 351 8028811",
                  ziel=14, seit=dt.date(2018, 7, 1), bankname="BAECKEREI HOFFMANN")
HEIKE_1 = add_c(kdnr=10088, gruppe="P", anrede="Frau", vorname="Heike", nachname="Müller", name="Müller, Heike",
                strasse="Kesselsdorfer Straße", hnr="40", plz="01159", ort="Dresden", tel="+49 351 4210977",
                email="heike.mueller@gmx.de", ziel=14, seit=dt.date(2019, 11, 4), bankname="HEIKE MUELLER",
                variants=[("Müller, Heike", .8), ("Heike Müller", .5)])
HOFFMANN_DUP = add_c(kdnr=10119, gruppe="P", anrede="Herr", vorname="Klaus", nachname="Hoffmann",
                     name="Hoffmann, Klaus", strasse="Borsbergstr.", hnr="22", plz="01309", ort="Dresden",
                     tel="+49 351 3361208", ziel=14, seit=dt.date(2025, 6, 3), bem="Neuanlage Tel.-Auftrag",
                     bankname="KLAUS HOFMANN", true_id="K10014")
HEIKE_2 = add_c(kdnr=10131, gruppe="P", anrede="Frau", vorname="Heike", nachname="Müller", name="Müller, Heike",
                strasse="Am Hang", hnr="3", plz="01328", ort="Dresden", tel="+49 351 2689044",
                email="heike.mueller@web.de", ziel=14, seit=dt.date(2026, 2, 10), bankname="HEIKE MUELLER",
                variants=[("Müller, Heike", .6), ("Heike Müller", .0)])
SCHMIDT_PETRA = add_c(kdnr=10134, gruppe="P", anrede="Frau", vorname="Petra", nachname="Schmidt",
                      name="Schmidt, Petra", strasse="Hüblerstraße", hnr="14", plz="01309", ort="Dresden",
                      tel="+49 351 3108899", email="petra.schmidt@gmx.net", ziel=14, seit=dt.date(2026, 9, 29),
                      bankname="PETRA SCHMIDT", bem="Angebot Badumbau")

SPECIAL_NRS = {c["kdnr"] for c in CUST}

STREETS = [("Hechtstraße", "01097"), ("Bautzner Straße", "01099"), ("Königsbrücker Straße", "01099"),
           ("Leipziger Straße", "01127"), ("Großenhainer Straße", "01127"), ("Kesselsdorfer Straße", "01159"),
           ("Pirnaer Landstraße", "01237"), ("Borsbergstraße", "01309"), ("Schandauer Straße", "01277"),
           ("Wehlener Straße", "01279"), ("Augsburger Straße", "01309"), ("Bergstraße", "01069"),
           ("Nürnberger Straße", "01187"), ("Chemnitzer Straße", "01187"), ("Löbtauer Straße", "01159"),
           ("Hamburger Straße", "01157"), ("Bürgerstraße", "01127"), ("Rudolf-Leonhard-Straße", "01097"),
           ("Louisenstraße", "01099"), ("Alaunstraße", "01099"), ("Görlitzer Straße", "01099"),
           ("Martin-Luther-Straße", "01099"), ("Prießnitzstraße", "01099"), ("Bischofsweg", "01099"),
           ("Stauffenbergallee", "01099"), ("Karcherallee", "01277"), ("Tolkewitzer Straße", "01279"),
           ("Winterbergstraße", "01237"), ("Zwinglistraße", "01277"), ("Gamigstraße", "01239"),
           ("Prohliser Allee", "01239"), ("Dohnaer Straße", "01219"), ("Teplitzer Straße", "01219"),
           ("Lockwitzer Straße", "01219"), ("Altenberger Straße", "01277"), ("Tharandter Straße", "01159"),
           ("Freiberger Straße", "01067"), ("Cossebauder Straße", "01157"), ("Meißner Landstraße", "01157"),
           ("Hansastraße", "01097"), ("Grundstraße", "01326"), ("Pillnitzer Landstraße", "01326"),
           ("Hauptstraße", "01445"), ("Dresdner Straße", "01705"), ("Gartenstraße", "01796"),
           ("Bahnhofstraße", "01454")]
ORT = {"01445": "Radebeul", "01705": "Freital", "01796": "Pirna", "01454": "Radeberg"}

FIRST_M = ["Andreas", "Bernd", "Christian", "Detlef", "Frank", "Gerd", "Hans-Jürgen", "Holger", "Jan", "Jörg",
           "Karsten", "Lars", "Matthias", "Michael", "Olaf", "Rainer", "Ralf", "Steffen", "Sven", "Torsten",
           "Uwe", "Volker", "Wolfgang", "Stefan", "Tobias", "Florian", "Sebastian", "Martin", "René", "Mirko"]
FIRST_F = ["Anja", "Birgit", "Claudia", "Doreen", "Elke", "Franziska", "Gabriele", "Ines", "Jana", "Katrin",
           "Kerstin", "Manuela", "Nadine", "Peggy", "Ramona", "Sandra", "Silke", "Simone", "Susanne", "Ute",
           "Yvonne", "Anett", "Cornelia", "Ilona", "Marion", "Steffi", "Antje", "Grit", "Romy", "Dagmar"]
LAST = ["Wagner", "Becker", "Schulz", "Koch", "Bauer", "Klein", "Wolf", "Schröder", "Neumann", "Schwarz",
        "Zimmermann", "Braun", "Hartmann", "Lange", "Werner", "Krause", "Meier", "Seidel", "Fischer", "Vogel",
        "Günther", "Frank", "Berger", "Winkler", "Lorenz", "Ludwig", "Böhme", "Kühn", "Pohl", "Haase", "Franke",
        "Engel", "Horn", "Arnold", "Thiele", "Kaiser", "Fiedler", "Ullrich", "Kretschmar", "Heinrich", "Göbel",
        "Jähnig", "Lindner", "Ebert", "Uhlig", "Petzold", "Hering", "Gebhardt", "Rößler", "Grafe"]
GEWERBE = ["Friseursalon Haarmonie", "Restaurant Zum Goldenen Anker GmbH", "Physiotherapie am Alaunpark",
           "Kanzlei Seidel & Kollegen", "Autohaus Pieschen GmbH", "Café Neustadt UG", "Fitnessstudio ElbeFit GmbH",
           "Blumen Kretschmar", "Fleischerei Geißler GmbH", "Optik Haubold e.K.", "Pension Loschwitz",
           "Kindergarten Elbspatzen e.V.", "Tierarztpraxis Dr. Fiedler", "Steuerbüro Ullrich",
           "Hotel Elbblick GmbH", "Seniorenresidenz Am Großen Garten gGmbH", "Kita Sonnenblume",
           "Gaststätte Zur Linde", "Malerbetrieb Thiele GmbH", "Apotheke am Bischofsplatz",
           "Sportverein Blau-Weiß Dresden e.V.", "Druckerei Lindner GmbH", "Weinstube Radebeul",
           "Hostel Neustadt GmbH"]
HVWG_EXTRA = [("Immobilienverwaltung Weißer Hirsch GmbH", "HV"), ("Hausverwaltung Kröber", "HV"),
              ("WEG-Verwaltung Blasewitz GmbH", "HV"), ("Wohnungsbaugenossenschaft Johannstadt eG", "WG"),
              ("Städtische Wohnungsgesellschaft Freital mbH", "WG")]

all_nrs = [n for n in range(10001, 10139) if n not in (10017, 10093)]   # Lücken im Kundenstamm
free_nrs = [n for n in all_nrs if n not in SPECIAL_NRS]
assert len(free_nrs) == len(all_nrs) - len(SPECIAL_NRS)
kinds = (["G"] * len(GEWERBE)) + [k for _, k in HVWG_EXTRA]
kinds += ["P"] * (len(free_nrs) - len(kinds))
R.shuffle(kinds)
gew_iter, hv_iter = iter(GEWERBE), iter(HVWG_EXTRA)
used_names = set()
used_addr = {(c["strasse"], c["hnr"]) for c in CUST}
for nr, kind in zip(free_nrs, kinds):
    st, plz = R.choice(STREETS)
    hnr = str(R.randint(1, 180)) + R.choice(["", "", "", "", "a", "b"])
    while (st, hnr) in used_addr:
        hnr = str(R.randint(1, 180))
    used_addr.add((st, hnr))
    frac = (nr - 10001) / 140
    seit = dt.date(2010, 1, 1) + dt.timedelta(days=int(frac * 5900 + R.randint(-60, 60)))
    tel = f"+49 351 {R.randint(2000000, 8999999)}" if plz not in ORT else f"+49 {R.choice(['351', '3528', '3501'])} {R.randint(200000, 899999)}"
    if kind == "P":
        while True:
            female = R.random() < 0.5
            vn = R.choice(FIRST_F if female else FIRST_M)
            nn = R.choice(LAST)
            if (nn, vn) not in used_names:
                used_names.add((nn, vn))
                break
        em = None
        if R.random() < 0.55:
            em = f"{vn.lower()[0]}.{nn.lower()}@{R.choice(['gmx.de', 'web.de', 't-online.de', 'gmail.com', 'posteo.de'])}"
            em = em.replace("ö", "oe").replace("ü", "ue").replace("ä", "ae").replace("ß", "ss")
        add_c(kdnr=nr, gruppe="P", anrede="Frau" if female else "Herr", vorname=vn, nachname=nn,
              name=f"{nn}, {vn}", strasse=st, hnr=hnr, plz=plz, ort=ORT.get(plz, "Dresden"), tel=tel,
              mobil=(f"+49 1{R.choice(['51', '52', '60', '62', '70', '71', '72', '76'])} {R.randint(1000000, 9999999)}" if R.random() < .3 else None),
              email=em, ziel=14, seit=seit,
              bankname=R.choice([f"{vn} {nn}".upper().replace("Ä", "AE").replace("Ö", "OE").replace("Ü", "UE").replace("ß", "SS"),
                                 f"{nn}, {vn}", f"{vn} {nn}"]))
    elif kind == "G":
        nm = next(gew_iter)
        slug = nm.split()[0].lower().replace("ä", "ae").replace("ö", "oe").replace("ü", "ue").replace("é", "e")
        add_c(kdnr=nr, gruppe="G", name=nm, ansprech=R.choice(["", "Hr. " + R.choice(LAST), "Fr. " + R.choice(LAST)]),
              strasse=st, hnr=hnr, plz=plz, ort=ORT.get(plz, "Dresden"), tel=tel,
              email=f"info@{slug}-dresden.de", ziel=R.choice([14, 14, 30]), skonto=R.random() < .25, seit=seit,
              bankname=nm.upper().replace("Ä", "AE").replace("Ö", "OE").replace("Ü", "UE").replace("ß", "SS")[:27])
    else:
        nm, g = next(hv_iter)
        objs = []
        for _ in range(R.randint(3, 6)):
            s2, _p = R.choice(STREETS[:40])
            objs.append(f"{s2} {R.randint(1, 120)}")
        add_c(kdnr=nr, gruppe=g, name=nm, ansprech="Technik / " + R.choice(["Hr. ", "Fr. "]) + R.choice(LAST),
              strasse=st, hnr=hnr, plz=plz, ort=ORT.get(plz, "Dresden"), tel=tel,
              email="technik@" + nm.split()[0].lower().replace("ä", "ae").replace("ß", "ss").replace("-", "") + ".de",
              ziel=30, seit=seit, objekte=objs,
              bankname=nm.upper().replace("Ä", "AE").replace("Ö", "OE").replace("Ü", "UE").replace("ß", "SS")[:27])

CUST.sort(key=lambda c: c["kdnr"])
assert len(CUST) == 136, len(CUST)
TRUE_CUSTOMERS = sorted({c["true_id"] for c in CUST})


def short_street(s):
    return s.replace("straße", "str.").replace("Straße", "Str.")


def addr(c):
    return f"{c['strasse']} {c['hnr']}"


def name_variant(c):
    """(wie im Auftragsbuch getippt, Kd-Nr eingetragen?)"""
    if c["variants"]:
        nm, p = R.choice(c["variants"])
        return nm, R.random() < p
    if c["gruppe"] == "P":
        r = R.random()
        if r < .7:
            nm = c["name"]
        elif r < .85:
            nm = f"{c['vorname']} {c['nachname']}"
        else:
            nm = ("Hr. " if c["anrede"] == "Herr" else "Fr. ") + c["nachname"]
        return nm, R.random() < .85
    nm = c["name"]
    if R.random() < .2:
        for suf in [" GmbH", " gGmbH", " e.K.", " e.V.", " eG", " UG", " mbH"]:
            nm = nm.replace(suf, "")
    return nm, R.random() < .85


# ---------------------------------------------------------------- Artikel (Preisliste)
ART = []
import itertools
GROUPS = {   # (Vorlage, {Platzhalter: Werte}, Hersteller, ME, EK von, EK bis)
    "HEIZUNG": [("Thermostatkopf {m} {x}", {"x": ["RA 2000", "RAW-K", "K-Kopf", "Uni XH", "Uni LH", "VK M30x1,5"]}, ["Danfoss", "Heimeier", "Oventrop"], "Stk", 12, 28),
                ("Heizkörperventil DN{x} {y} {m}", {"x": [10, 15, 20], "y": ["Eck", "Durchgang", "Axial"]}, ["Heimeier", "Oventrop"], "Stk", 9, 22),
                ("Kompaktheizkörper {m} Typ {x} {y}x{z}", {"x": [11, 21, 22, 33], "y": [300, 400, 500, 600, 900], "z": [400, 600, 800, 1000, 1200, 1400]}, ["Kermi", "Purmo", "Buderus"], "Stk", 95, 420),
                ("Ausdehnungsgefäß {m} {x} l", {"x": [8, 12, 18, 25, 35, 50, 80]}, ["Reflex", "Flamco"], "Stk", 38, 160),
                ("Sicherheitsventil {m} {x} bar {y}", {"x": ["2,5", "3", "6", "10"], "y": ['1/2"', '3/4"', '1"']}, ["Syr", "Watts"], "Stk", 14, 45),
                ("Heizungsumwälzpumpe {m} {x}", {"x": ["25-40 180", "25-60 180", "15-60 130", "32-80 180"]}, ["Grundfos", "Wilo"], "Stk", 140, 520),
                ("Entlüfter automatisch {x} {m}", {"x": ['3/8"', '1/2"']}, ["Afriso", "Caleffi"], "Stk", 6, 18)],
    "SANITÄR": [("Waschtischarmatur {m} Einhebel Gr. {x}", {"x": ["S", "M", "L", "XL"]}, ["Grohe", "Hansgrohe", "Kludi"], "Stk", 45, 260),
                ("Spülkasten-Füllventil {m} {x}", {"x": ["Typ 380", "Typ 333", "UP320"]}, ["Geberit", "Viega"], "Stk", 11, 34),
                ("Wand-WC spülrandlos {m} {x}", {"x": ["weiß", "weiß CeramicPlus", "pergamon", "manhattan"]}, ["Villeroy & Boch", "Duravit", "Ideal Standard"], "Stk", 120, 380),
                ("Siphon DN{x} {m}", {"x": [32, 40, 50]}, ["Viega", "Dallmer"], "Stk", 8, 26),
                ("Duschrinne {x} mm {m}", {"x": [700, 800, 900, 1000, 1200]}, ["Kessel", "ACO", "Dallmer"], "Stk", 160, 480),
                ("Eckventil {x} {m}", {"x": ['1/2" x 3/8"', '1/2" x 1/2"', '3/8" x 10 mm']}, ["Schell", "Viega"], "Stk", 5, 16),
                ("UP-Spülkasten {m} {x} cm", {"x": [112, 98, 82]}, ["Geberit", "Viega"], "Stk", 95, 210),
                ("Brausegarnitur {m} {x} mm", {"x": [600, 900, 1100]}, ["Grohe", "Hansgrohe", "Kludi"], "Stk", 40, 190),
                ("Waschtisch {x} cm {m}", {"x": [45, 50, 55, 60, 65, 80, 100, 120]}, ["Villeroy & Boch", "Duravit", "Ideal Standard"], "Stk", 60, 520),
                ("Haltegriff {x} mm {m}", {"x": [300, 450, 600, 900]}, ["Hewi", "Normbau"], "Stk", 25, 90)],
    "INSTALLATION": [("Kupferrohr {x}x1 mm hart (5 m) {m}", {"x": [12, 15, 18, 22, 28, 35]}, ["KME", "Wieland"], "Stg", 22, 75),
                     ("Pressfitting {x} {y} mm {m}", {"x": ["Bogen 90°", "Bogen 45°", "T-Stück", "Muffe", "Übergang AG", "Übergang IG", "Reduzierstück"], "y": [12, 15, 18, 22, 28, 35]}, ["Viega", "Geberit"], "Stk", 3, 14),
                     ("Mehrschichtverbundrohr {x}x2 (Rolle 50 m) {m}", {"x": [16, 20, 26, 32]}, ["Uponor", "Rehau"], "Rol", 70, 190),
                     ("Rohrisolierung {x} mm, {y} mm Wand (2 m) {m}", {"x": [12, 15, 18, 22, 28, 35], "y": [9, 13, 20]}, ["Armacell", "Kaimann"], "Stk", 2, 9),
                     ("HT-Rohr DN{x} {y} mm {m}", {"x": [40, 50, 70, 100], "y": [250, 500, 1000, 2000]}, ["Ostendorf", "Rehau"], "Stk", 3, 12)],
    "VERBRAUCH": [("Gewindedichtung Hanf {x} g {m}", {"x": [100, 200, 500]}, ["Unipak"], "Pck", 2, 6),
                  ("Dichtpaste {x} g {m}", {"x": [65, 150, 300]}, ["Fermit", "Unipak"], "Tube", 3, 8),
                  ("Sanitär-Silikon {x} 310 ml {m}", {"x": ["weiß", "transparent", "manhattan", "pergamon", "grau"]}, ["Otto", "Ramsauer"], "Kart", 5, 12),
                  ("Rohrschelle {x} mm M{y} {m}", {"x": [12, 15, 18, 22, 28, 35], "y": [8, 10]}, ["Fischer", "Würth"], "Stk", 1, 4),
                  ("Dübel {x}x{y} {m}", {"x": [6, 8, 10], "y": [30, 40, 50]}, ["Fischer", "Würth"], "Pck", 3, 9)],
}
artnr = 3100100
for grp, templates in GROUPS.items():
    n_g = {"HEIZUNG": 130, "SANITÄR": 100, "INSTALLATION": 130, "VERBRAUCH": 55}[grp]
    cands = []
    for tpl, ph, mans, me, lo, hi in templates:
        keys = list(ph)
        for combo in itertools.product(mans, *[ph[k] for k in keys]):
            kw = dict(zip(["m"] + keys, combo))
            cands.append((tpl.format(**kw), kw["m"], me, lo, hi))
    R.shuffle(cands)
    cands = cands[:n_g]
    for bez, m, me, lo, hi in cands:
        artnr += R.randint(3, 97)
        ek = q2(R.uniform(lo, hi))
        ART.append(dict(nr=str(artnr), bez=bez, hersteller=m, me=me, ek=ek, vk=q2(ek * Decimal("1.35")),
                        gruppe=grp, rabatt=R.choice(["A", "B", "C", "D"])))
ART_BY_NR = {a["nr"]: a for a in ART}
PUMPS = []
pnr = 99160500
for i in range(60):
    pnr += R.randint(7, 61)
    model = R.choice(["Alpha2 25-40", "Alpha2 25-60", "Alpha1 L 25-40", "UPS 25-60", "Magna1 32-80",
                      "Comfort 15-14 BX", "Stratos Pico 25/1-6", "Yonos Maxo 30/0,5-7", "Para 25/6"])
    p = R.choice(["on request", "", None]) if R.random() < .12 else q2(R.uniform(160, 1450))
    PUMPS.append(dict(nr=str(pnr), bez=f"Circulator pump {model} {R.choice(['130 mm', '180 mm'])} {R.choice(['1x230V', '3x400V'])}",
                      price=p, dgroup=R.choice(["P1", "P2", "P3"]), lead=R.choice([2, 3, 5, 10, 14]),
                      ean=int("57" + str(R.randint(10 ** 10, 10 ** 11 - 1)))))

# ---------------------------------------------------------------- Wartungsverträge
DEVICES = ["Vaillant ecoTEC plus VC 206/5-5", "Viessmann Vitodens 200-W 19 kW", "Buderus Logamax plus GB172-20",
           "Wolf CGB-2-20", "Brötje WGB 20 E", "Junkers Cerapur ZWB 24", "Vaillant aroTHERM plus 7 kW (WP)",
           "Viessmann Vitocal 250-A (WP)", "Weishaupt WTC-GW 25", "Ölkessel Buderus Logano G125"]
CONTRACTS = []
contract_custs = [WG_ELBTAL, HV_SCHULZE, KRUEGER, SCHULZE_G, LEHMANN, MUELLER_G, HOFMANN, HEIKE_1, BAECKER_H]
others = [c for c in CUST if c["kdnr"] not in SPECIAL_NRS and c["gruppe"] in ("P", "G", "HV", "WG")]
contract_custs += R.sample(others, 26)
cseq = defaultdict(int)
for c in contract_custs:
    objs = c["objekte"][:R.randint(2, 3)] if c["gruppe"] in ("HV", "WG") else [addr(c)]
    for o in objs:
        y = R.choice([2019, 2020, 2021, 2022, 2023, 2024, 2025])
        cseq[y] += 1
        nr = f"WV-{y}-{cseq[y]:03d}"
        month = R.choice([3, 4, 5, 9, 10, 10, 11, 11])
        last = dt.date(2025, month, R.randint(1, 28))
        CONTRACTS.append(dict(nr=nr, cust=c, objekt=o, anlage=R.choice(DEVICES),
                              intervall=R.choice(["jährlich", "jährlich", "jährlich", "halbjährlich"]),
                              pauschale=q2(R.choice([139, 139, 149, 189, 229, 260])), last=last,
                              next=dt.date(2026, month, 1), start_year=y))
# fixer Vertrag für HV Schulze (E-Mail)
hvc = [k for k in CONTRACTS if k["cust"] is HV_SCHULZE]
hvc[0]["nr"] = "WV-2024-012"
hvc[0]["next"] = dt.date(2026, 10, 1)
hvc[0]["objekt"] = "Altgorbitzer Ring 21"
CONTRACTS.sort(key=lambda k: k["nr"])
assert len({k["nr"] for k in CONTRACTS}) == len(CONTRACTS)

# ---------------------------------------------------------------- Aufträge
ORDER_TYPES = {
    "Wartung": (lambda: R.choice([139, 139, 139, 149, 159, 189]),
                ["Jahreswartung Gastherme", "Wartung Gas-Brennwertkessel", "Wartung Ölheizung", "Wartung Wärmepumpe"], False),
    "Reparatur": (lambda: R.randrange(90, 1200),
                  ["Austausch Thermostatventil", "Spülkasten läuft nach", "Heizung ohne Funktion", "Druckverlust Heizung",
                   "Kein Warmwasser", "Armatur tropft", "Siphon undicht", "Zirkulationspumpe defekt"], False),
    "Notdienst": (lambda: R.randrange(180, 650),
                  ["Rohrbruch Küche", "Heizungsausfall Wochenende", "Wasserschaden Keller", "Verstopfung Grundleitung"], False),
    "Montage": (lambda: R.randrange(250, 1800),
                ["Montage Waschtisch", "Montage Wand-WC", "Montage Heizkörper", "Austausch Warmwasserspeicher",
                 "Einbau Enthärtungsanlage"], False),
    "Heizungstausch": (lambda: R.randrange(8500, 17500),
                       ["Austausch Gaskessel gegen Brennwert", "Einbau Wärmepumpe Luft/Wasser", "Heizungsmodernisierung"], True),
    "Badsanierung": (lambda: R.randrange(9000, 26000),
                     ["Komplettsanierung Bad", "Badumbau barrierefrei", "Gäste-WC Sanierung"], True),
    "Trinkwasser": (lambda: R.randrange(250, 480),
                    ["Legionellenprüfung", "Trinkwasseruntersuchung", "Spülung Trinkwasseranlage"], False),
    "Hydr. Abgleich": (lambda: R.randrange(600, 1600), ["Hydraulischer Abgleich", "Hydr. Abgleich + Pumpentausch"], False),
    "Strangsanierung": (lambda: R.randrange(18000, 46000), ["Strangsanierung Trinkwasser", "Strangsanierung Abwasser + TW", "Erneuerung Heizungsverteilung Keller"], True),
}
TYPE_W = {"P": {"Wartung": 30, "Reparatur": 28, "Notdienst": 6, "Montage": 14, "Heizungstausch": 12,
                "Badsanierung": 10, "Hydr. Abgleich": 4},
          "G": {"Wartung": 30, "Reparatur": 30, "Notdienst": 8, "Montage": 15, "Trinkwasser": 8, "Heizungstausch": 6},
          "HV": {"Wartung": 30, "Reparatur": 25, "Notdienst": 15, "Trinkwasser": 20, "Montage": 6, "Hydr. Abgleich": 4, "Strangsanierung": 5},
          "WG": {"Wartung": 30, "Reparatur": 25, "Notdienst": 15, "Trinkwasser": 20, "Montage": 6, "Hydr. Abgleich": 4, "Strangsanierung": 5}}

VACATION = {"JK": (dt.date(2026, 7, 20), dt.date(2026, 7, 31)), "MR": (dt.date(2026, 8, 3), dt.date(2026, 8, 14)),
            "DH": (dt.date(2026, 8, 17), dt.date(2026, 8, 28)), "AY": (dt.date(2026, 9, 14), dt.date(2026, 9, 18)),
            "KS": (dt.date(2026, 7, 27), dt.date(2026, 8, 7)), "LW": (dt.date(2026, 8, 10), dt.date(2026, 8, 21))}


def available(kz, d):
    a, b = VACATION.get(kz, (None, None))
    return not (a and a <= d <= b)


def pick_monteur(art, d):
    if art in ("Heizungstausch", "Badsanierung", "Strangsanierung"):
        opts = ["MR/KS", "DH/LW", "AY/KS", "MR/LW"] if d >= dt.date(2025, 9, 1) else ["MR/KS", "DH", "AY/KS"]
    elif art in ("Wartung", "Reparatur", "Notdienst"):
        opts = ["JK", "JK", "JK", "MR", "DH", "AY"]
    else:
        opts = ["MR", "DH", "AY", "JK"]
    for _ in range(20):
        m = R.choice(opts)
        if all(available(k, d) for k in m.split("/")):
            return m
    return "TB"


ORDERS = []


def mk_order(c, d, art=None, beschr=None, wert=None, monteur=None, status=None, done=None, tag=None, objekt=None):
    if art is None:
        w = TYPE_W[c["gruppe"]]
        art = R.choices(list(w), weights=list(w.values()))[0]
    vfn, blist, big = ORDER_TYPES[art]
    if wert is None:
        wert = vfn()
    if beschr is None:
        beschr = R.choice(blist)
    if objekt is None:
        objekt = R.choice(c["objekte"]) if c["objekte"] else addr(c)
    o = dict(cust=c, date=d, art=art, beschr=beschr, wert=q2(wert), monteur=monteur or pick_monteur(art, d),
             status=status, done=done, big=big, tag=tag, objekt=objekt, invoices=[])
    ORDERS.append(o)
    return o


weights = []
for c in CUST:
    w = {"P": 1.0, "G": 2.0, "HV": 6.0, "WG": 7.0}[c["gruppe"]]
    if c["kdnr"] in (10119, 10131, 10134):
        w = 0
    weights.append(w)


def random_date(y, until=None):
    a = dt.date(y, 1, 2)
    b = until or dt.date(y, 12, 22)
    while True:
        d = a + dt.timedelta(days=R.randint(0, (b - a).days))
        if d.month in (9, 10, 11, 12, 1) or R.random() < .7:
            return next_wd(d) if next_wd(d) <= b else d


for y, n in ((2025, 235), (2026, 245)):
    for _ in range(n):
        c = R.choices(CUST, weights=weights)[0]
        d = random_date(y, dt.date(2026, 10, 14) if y == 2026 else None)
        if c["seit"] > d:
            d = max(d, c["seit"])
            if d.year != y:
                continue
        mk_order(c, d)

# erzwungene Aufträge (Story) -----------------------------------------------------------
for d in (dt.date(2025, 3, 12), dt.date(2025, 10, 7), dt.date(2026, 2, 3), dt.date(2026, 5, 19)):
    mk_order(MUELLER_G, d)
O_MUELLER_AUG = mk_order(MUELLER_G, dt.date(2026, 8, 11), art="Montage", beschr="Austausch Warmwasserspeicher Lager",
                         wert=935, monteur="MR", status="abgerechnet", done=dt.date(2026, 8, 13), tag="mueller_rg_pdf")
for d in (dt.date(2025, 4, 2), dt.date(2025, 9, 23), dt.date(2026, 1, 20), dt.date(2026, 6, 9), dt.date(2026, 9, 1)):
    mk_order(MUELLER_BAU, d, art=R.choice(["Montage", "Reparatur", "Trinkwasser"]))
O_MUELLER_BAU_STORNO = mk_order(MUELLER_BAU, dt.date(2026, 3, 10), art="Montage", beschr="Montage Sanitärobjekte Baustelle Radebeul",
                                wert=4380, monteur="AY", status="abgerechnet", done=dt.date(2026, 3, 18),
                                tag="storno", objekt="Baustelle Radebeul, Meißner Str. 210")
for d in (dt.date(2025, 6, 17), dt.date(2026, 4, 14), dt.date(2026, 9, 21)):
    mk_order(ELEKTRO_M, d)
mk_order(HEIKE_1, dt.date(2025, 11, 4), art="Wartung")
mk_order(HEIKE_1, dt.date(2026, 6, 23), art="Reparatur", beschr="Spülkasten läuft nach")
mk_order(HEIKE_2, dt.date(2026, 3, 2), art="Reparatur", beschr="Druckverlust Heizung")
O_HEIKE2 = mk_order(HEIKE_2, dt.date(2026, 9, 8), art="Reparatur", beschr="Reparatur Waschtischarmatur Bad OG",
                    wert=286.5, monteur="DH", status="abgerechnet", done=dt.date(2026, 9, 9), tag="reklamation")
mk_order(HOFMANN, dt.date(2025, 2, 11), art="Wartung")
mk_order(HOFFMANN_DUP, dt.date(2025, 6, 3), art="Notdienst", beschr="Rohrbruch Küche")
mk_order(HOFFMANN_DUP, dt.date(2026, 2, 24), art="Wartung")
mk_order(KRUEGER, dt.date(2025, 5, 6), art="Wartung")
O_KRUEGER_BAD = mk_order(KRUEGER, dt.date(2026, 9, 28), art="Badsanierung", beschr="Badsanierung lt. Angebot AN-26-044",
                         wert=15480, monteur="MR/KS", status="geplant", tag="krueger_bad")
mk_order(SCHULZE_G, dt.date(2025, 10, 14), art="Wartung")
mk_order(SCHULZE_G, dt.date(2026, 10, 6), art="Wartung", status="geplant")
mk_order(SCHMIDT_P, dt.date(2025, 8, 19), art="Reparatur")
mk_order(SCHMIDT_P, dt.date(2026, 7, 7), art="Montage")
mk_order(SCHMIDT_PETRA, dt.date(2026, 10, 12), art="Badsanierung", beschr="Besichtigung Badumbau, Angebot erstellen",
         wert=0, monteur="TB", status="offen", tag="petra_angebot")
for d in (dt.date(2025, 9, 9), dt.date(2026, 3, 17), dt.date(2026, 9, 15)):
    mk_order(LEHMANN, d)
for d in (dt.date(2025, 1, 21), dt.date(2026, 1, 13)):
    mk_order(BAECKER_H, d)
O_HV_GORBITZ = mk_order(HV_SCHULZE, dt.date(2026, 10, 2), art="Wartung",
                        beschr="Wartung Heizungsanlagen lt. WV-2024-012 (Mail 01.10.)", wert=139 * 3, monteur="JK",
                        status="geplant", tag="hv_mail", objekt="Altgorbitzer Ring 21")

# Status / erledigt ----------------------------------------------------------------------
for o in ORDERS:
    d = o["date"]
    if o["status"] is None:
        r = R.random()
        if d < dt.date(2026, 8, 1):
            o["status"] = "storniert" if r < .03 else ("erledigt" if r < .08 else "abgerechnet")
        elif d < dt.date(2026, 10, 1):
            o["status"] = ("storniert" if r < .04 else "erledigt" if r < .30 else "in Arbeit" if r < .38 else "abgerechnet")
        else:
            o["status"] = R.choice(["offen", "offen", "geplant", "in Arbeit"])
        if o["big"] and d > dt.date(2026, 8, 15) and o["status"] == "abgerechnet":
            o["status"] = "in Arbeit"
    if o["done"] is None and o["status"] in ("erledigt", "abgerechnet"):
        o["done"] = add_wd(d, R.randint(12, 35) if o["big"] else R.randint(0, 4))
        if o["done"] > EXPORT:
            o["done"] = EXPORT - dt.timedelta(days=1)
            o["status"] = "erledigt"

ORDERS.sort(key=lambda o: (o["date"], o["cust"]["kdnr"]))
seq = Counter()
for o in ORDERS:
    yy = o["date"].year % 100
    seq[yy] += 1
    o["nr"] = f"A-{yy}-{seq[yy]:03d}"
ORD_BY_NR = {o["nr"]: o for o in ORDERS}
O_BY_TAG = {o["tag"]: o for o in ORDERS if o["tag"]}

# ---------------------------------------------------------------- Rechnungen
INVOICES = []


def mk_inv(o, d, netto, betreff, kind="normal", cust=None):
    inv = dict(order=o, cust=cust or o["cust"], date=d, netto=q2(netto), betreff=betreff, kind=kind,
               paid=None, paid_amount=None, status=None, bank_rows=[])
    inv["ust"] = q2(inv["netto"] * Decimal("0.19"))
    inv["brutto"] = inv["netto"] + inv["ust"]
    INVOICES.append(inv)
    if o:
        o["invoices"].append(inv)
    return inv


for o in ORDERS:
    if o["status"] != "abgerechnet":
        continue
    base = o["wert"]
    if o["big"]:
        ab = q2(round(float(base) * 0.4, -2))
        d1 = add_wd(o["date"], R.randint(3, 8))
        mk_inv(o, d1, ab, f"Abschlagsrechnung 1 {o['beschr']}, Auftrag {o['nr']}", kind="abschlag")
        d2 = add_wd(o["done"], R.randint(2, 9))
        mk_inv(o, d2, base - ab, f"Schlussrechnung {o['beschr']} (abzgl. Abschlag), Auftrag {o['nr']}", kind="schluss")
    else:
        d2 = add_wd(o["done"], R.randint(1, 9))
        if d2 > EXPORT:
            d2 = EXPORT
        style = R.random()
        if style < .45:
            bt = f"{o['beschr']} lt. Auftrag {o['nr']}"
        elif style < .65:
            bt = f"{o['art']} {o['objekt']} ({o['nr']})"
        elif style < .85:
            bt = f"{o['beschr']}, {o['objekt']}"
        else:
            bt = f"{o['art']} vom {dd(o['done'])}"
        mk_inv(o, d2, base, bt)
# Rechnungen ohne Auftrag (Ladenverkauf)
for _ in range(7):
    c = R.choice([c for c in CUST if c["gruppe"] in ("P", "G") and c["kdnr"] not in (10119, 10134)])
    d = next_wd(dt.date(2025, 1, 15) + dt.timedelta(days=R.randint(0, 600)))
    mk_inv(None, d, q2(R.uniform(18, 140)), R.choice(["Verkauf Ersatzteil ab Lager", "Materialverkauf (Abholung)",
                                                      "Verkauf Dichtungssatz + Kleinmaterial"]), cust=c, kind="ohne_auftrag")

INVOICES.sort(key=lambda i: (i["date"], i["cust"]["kdnr"]))
# Storno-Story: Rechnung zu O_MUELLER_BAU_STORNO wird storniert und neu gestellt
storno_orig = O_MUELLER_BAU_STORNO["invoices"][0]
storno_orig["kind"] = "storniert"
sd = add_wd(storno_orig["date"], 3)
storno = mk_inv(O_MUELLER_BAU_STORNO, sd, -storno_orig["netto"], "", kind="storno")
neu = mk_inv(O_MUELLER_BAU_STORNO, sd, Decimal("4180.00"),
             f"Montage Sanitärobjekte Baustelle Radebeul (korrigiert), Auftrag {O_MUELLER_BAU_STORNO['nr']}", kind="korrektur")
INVOICES.sort(key=lambda i: (i["date"], 0 if i["kind"] != "korrektur" else 1, i["cust"]["kdnr"]))

GAPS = {2025: {44, 130}, 2026: {17, 18, 102}}
seqi = Counter()
for inv in INVOICES:
    y = inv["date"].year
    seqi[y] += 1
    while seqi[y] in GAPS[y]:
        seqi[y] += 1
    inv["nr"] = f"RE-{y}-{seqi[y]:04d}"
INV_BY_NR = {i["nr"]: i for i in INVOICES}
storno["betreff"] = f"Storno zu {storno_orig['nr']}"
for inv in INVOICES:
    inv["faellig"] = inv["date"] + dt.timedelta(days=inv["cust"]["ziel"])

# Zahlungsverhalten ----------------------------------------------------------------------
LAG = []          # bezahlt laut Bank, aber offen im Export
PARTIAL = []
UNPAID = []
for inv in INVOICES:
    c = inv["cust"]
    if inv["kind"] in ("storniert", "storno"):
        inv["status"] = "storniert" if inv["kind"] == "storniert" else "Storno"
        continue
    r = R.random()
    if inv["nr"] == O_MUELLER_AUG["invoices"][0]["nr"]:
        inv["paid"] = dt.date(2026, 9, 18)
        inv["paid_amount"] = inv["brutto"]
        inv["status"] = "gemahnt"
        LAG.append(inv["nr"])
        continue
    if r < .055 and inv["date"] > dt.date(2025, 10, 1):
        UNPAID.append(inv["nr"])
        inv["status"] = "gemahnt" if inv["faellig"] < dt.date(2026, 10, 1) else "offen"
        continue
    if c["kdnr"] == 10045:
        delay = R.randint(38, 75)
    elif c["gruppe"] in ("HV", "WG"):
        delay = R.randint(15, 33)
    elif c["skonto"] and R.random() < .7:
        delay = R.randint(3, 8)
        inv["skonto"] = True
    else:
        delay = R.randint(4, 22)
    pd = next_wd(inv["date"] + dt.timedelta(days=delay))
    if pd > EXPORT:
        inv["status"] = "offen"
        continue
    inv["paid"] = pd
    amt = inv["brutto"]
    if inv.get("skonto"):
        amt = q2(amt * Decimal("0.98"))
    if r > .965 and inv["brutto"] > 300 and dt.date(2026, 1, 10) < pd < dt.date(2026, 9, 1):
        amt = q2(float(inv["brutto"]) / 2)
        PARTIAL.append(inv["nr"])
        inv["status"] = "teilbezahlt"
    inv["paid_amount"] = amt
    if inv["status"] is None:
        if dt.date(2026, 9, 1) <= pd <= BANK_TO and R.random() < .3:
            inv["status"] = "offen"
            LAG.append(inv["nr"])
        else:
            inv["status"] = "bezahlt"

# ---------------------------------------------------------------- Bankumsätze
BANK = []   # dict(date, valuta, name, iban, bic, text, amount (signed), kind, refs)


def bank(d, name, amount, text, kind, refs=(), ib=None):
    BANK.append(dict(date=d, valuta=d if R.random() < .85 else add_wd(d, 1), name=name, iban=ib or rand_iban(),
                     amount=q2(amount), text=text, kind=kind, refs=list(refs)))


def zweck(inv_list, c):
    nrs = [i["nr"] for i in inv_list]
    if len(nrs) > 1:
        return f"Kd {c['kdnr']} " + ", ".join(nrs), "multi"
    n = nrs[0]
    y, k = n.split("-")[1], n.split("-")[2]
    r = R.random()
    if r < .40:
        return n, "exakt"
    if r < .55:
        return f"Rechnung {n} vom {dd(inv_list[0]['date'])}", "exakt"
    if r < .67:
        return f"Rg.-Nr. {y}-{k} Kd-Nr. {c['kdnr']}", "ohne_praefix"
    if r < .74:
        return f"RE{y}{k}", "zusammengezogen"
    if r < .82:
        return f"Re {int(k)}/{y}", "kurz"
    if r < .9:
        return f"{inv_list[0]['order']['art'] if inv_list[0]['order'] else 'Material'} {c['nachname'] or c['name'].split()[0]}", "ohne_nummer"
    return f"SVWZ+{n} EREF+NOTPROVIDED", "sepa"


pay_rows = defaultdict(list)
for inv in INVOICES:
    if inv["paid"] and BANK_FROM <= inv["paid"] <= BANK_TO:
        pay_rows[(inv["cust"]["true_id"] if inv["cust"]["kdnr"] != 10119 else "K10119", inv["paid"])].append(inv)
BANK_PAY_TRUTH = []
for (tid, pd), invs in sorted(pay_rows.items(), key=lambda kv: kv[0][1]):
    c = invs[0]["cust"]
    groups = [invs] if (c["gruppe"] in ("HV", "WG") and len(invs) > 1) else [[i] for i in invs]
    for g in groups:
        amt = sum(i["paid_amount"] for i in g)
        z, style = zweck(g, c)
        if any(i.get("skonto") for i in g):
            z += " abzgl. 2% Skonto"
        txt = R.choice(["Gutschrift", "Überweisungsgutschr.", "SEPA-Gutschrift"]) + "\n" + z
        bank(pd, c["bankname"], amt, txt, "kunde", [i["nr"] for i in g], ib=c["iban"])
        BANK_PAY_TRUTH.append(dict(datum=dd(pd), name=c["bankname"], betrag=de(amt), rechnungen=[i["nr"] for i in g],
                                   kunde=c["true_id"], stil=style))
        for i in g:
            i["bank_rows"].append(len(BANK) - 1)

# Restzahlungen der Teilzahler (später, ohne Nummer)
for n in PARTIAL[:2]:
    inv = INV_BY_NR[n]
    d2 = add_wd(inv["paid"], 20)
    if d2 <= BANK_TO:
        rest = inv["brutto"] - inv["paid_amount"]
        bank(d2, inv["cust"]["bankname"], rest, "Überweisungsgutschr.\nRestzahlung " + inv["nr"].replace("RE-", ""),
             "kunde_rest", [n], ib=inv["cust"]["iban"])
        BANK_PAY_TRUTH.append(dict(datum=dd(d2), name=inv["cust"]["bankname"], betrag=de(rest), rechnungen=[n],
                                   kunde=inv["cust"]["true_id"], stil="rest"))

SUPPLIERS = [
    dict(nr=70001, name="Elbe-Haustechnik Großhandel GmbH", bank="ELBE-HAUSTECHNIK GROSSHANDEL", kd="33810",
         strasse="Washingtonstraße 30", plz="01139", ort="Dresden", tel="0351 2109900", w=10),
    dict(nr=70002, name="SHK-Fachhandel Sachsen KG", bank="SHK-FACHHANDEL SACHSEN KG", kd="D-4471",
         strasse="Am Waldschlösschen 4", plz="01099", ort="Dresden", tel="0351 8107700", w=4),
    dict(nr=70003, name="Schrauben-Wolf GmbH", bank="SCHRAUBEN-WOLF GMBH", kd="883 112",
         strasse="Gewerbegebiet Süd 12", plz="01723", ort="Wilsdruff", tel="035204 66100", w=3),
    dict(nr=70004, name="Thermo-Werk Kundendienst Ost", bank="THERMO-WERK GMBH", kd="KD-2210-77",
         strasse="Industriestraße 8", plz="09111", ort="Chemnitz", tel="0371 4412200", w=1),
    dict(nr=70005, name="Fliesen Kaiser (Nachunternehmer)", bank="ANDRE KAISER FLIESENLEGERMEISTER", kd="—",
         strasse="Dorfstraße 7", plz="01465", ort="Dresden-Langebrück", tel="035201 70815", w=1),
    dict(nr=70006, name="Elektrogroßhandel Nord GmbH", bank="ELEKTROGROSSHANDEL NORD", kd="10-5588",
         strasse="Radeburger Straße 150", plz="01109", ort="Dresden", tel="0351 8881200", w=1),
]
for s in SUPPLIERS:
    s["iban"] = rand_iban()
for d in wd_range(BANK_FROM, BANK_TO):
    if R.random() < .32:
        s = R.choices(SUPPLIERS, weights=[s["w"] for s in SUPPLIERS])[0]
        rg = f"{R.randint(4700000, 4799999)}"
        bank(d, s["bank"], -q2(R.uniform(150, 3900)), f"Überweisung\nRG {rg} KD-NR {s['kd']}", "lieferant", [s["nr"]], ib=s["iban"])
months = range(1, 10)
for m in months:
    first = next_wd(dt.date(2026, m, 1))
    last = max(wd_range(dt.date(2026, m, 20), (dt.date(2026, m + 1, 1) - dt.timedelta(days=1))))
    bank(first, "GRUNDSTUECKSVERWALTUNG LEHNERT", -Decimal("2150.00"),
         f"Dauerauftrag\nMiete Werkstatt/Lager Hechtstr. 41 {m:02d}/2026", "miete")
    for kfz in ("DD-BH 101", "DD-BH 102", "DD-BH 103"):
        bank(first, "AUTOLEASING SACHSEN GMBH", -Decimal("489.00"), f"Lastschrift\nLeasingrate {m:02d}/2026 Vertrag 77-{kfz[-3:]} {kfz}", "leasing")
    for e in EMP:
        bank(last, f"{e['vorname']} {e['nachname']}".replace("ş", "s").replace("ı", "i"), -Decimal(e["netto"]) - q2(R.choice([0, 0, 0, 37.5, 52.8])),
             f"Lohn/Gehalt\nLOHN/GEHALT {m:02d}/2026 PNR {e['pnr']}", "lohn", [e["pnr"]])
    bank(next_wd(dt.date(2026, m, 10)), "FINANZAMT DRESDEN-NORD", -q2(R.uniform(900, 3400)),
         f"Überweisung\nUSt-VZ {m - 1 if m > 1 else 12:02d}/{2026 if m > 1 else 2025} StNr 202/118/54321", "steuer")
    bank(next_wd(dt.date(2026, m, 10)), "FINANZAMT DRESDEN-NORD", -q2(R.uniform(1800, 2450)),
         f"Überweisung\nLohnsteuer {m - 1 if m > 1 else 12:02d}/{2026 if m > 1 else 2025} StNr 202/118/54321", "steuer")
    bank(last, "AOK PLUS", -q2(R.uniform(6400, 7300)),
         f"Lastschrift\nBeitragsnachweis {m:02d}/2026 BNR 12345678", "sv")
    bank(next_wd(dt.date(2026, m, 5)), "TELEKOM DEUTSCHLAND GMBH", -Decimal("89.90"), "Lastschrift\nFestnetz/Mobil Kd 4455661", "sonst")
    bank(next_wd(dt.date(2026, m, 15)), "DKV EURO SERVICE", -q2(R.uniform(850, 1350)), "Lastschrift\nTankkarte Abrechnung", "sonst")
    bank(last, "VOLKSBANK DRESDEN-BAUTZEN EG", -Decimal("34.50"), "Abschluss\nKontoführung", "sonst")
    for _ in range(R.randint(2, 5)):
        dd_ = next_wd(dt.date(2026, m, R.randint(1, 26)))
        bank(dd_, "BAUMARKT DRESDEN NICKERN", -q2(R.uniform(8, 160)), "Kartenzahlung girocard\nBaumarkt Kartenzahlung", "sonst")
for m in (1, 4, 7):
    bank(next_wd(dt.date(2026, m, 15)), "SV SACHSEN VERSICHERUNG AG", -Decimal("612.40"), "Lastschrift\nBetriebshaftpflicht Q", "sonst")
    bank(next_wd(dt.date(2026, m, 15)), "SV SACHSEN VERSICHERUNG AG", -Decimal("1284.00"), "Lastschrift\nKfz-Versicherung 3 Fzg", "sonst")
bank(dt.date(2026, 3, 16), "HANDWERKSKAMMER DRESDEN", -Decimal("380.00"), "Lastschrift\nKammerbeitrag 2026", "sonst")
bank(dt.date(2026, 2, 16), "INNUNG SHK DRESDEN", -Decimal("450.00"), "Überweisung\nInnungsbeitrag 2026", "sonst")
bank(dt.date(2026, 4, 15), "BERUFSGENOSSENSCHAFT ETEM", -Decimal("2840.00"), "Lastschrift\nBeitrag 2025", "sonst")
bank(dt.date(2026, 5, 6), "IKK CLASSIC", Decimal("1460.22"), "Gutschrift\nErstattung U1 Kowalski", "sonst")

BANK.sort(key=lambda b: (b["date"], b["amount"]))
OPEN_BAL = Decimal("48432.17")
end_bal = OPEN_BAL + sum(b["amount"] for b in BANK)

# ---------------------------------------------------------------- Stundenzettel Q3 2026
TS = defaultdict(list)    # kz -> list of rows
Q3 = wd_range(dt.date(2026, 7, 1), dt.date(2026, 9, 30))
sick = {("DH", dt.date(2026, 7, 14)), ("DH", dt.date(2026, 7, 15)), ("JK", dt.date(2026, 9, 3)),
        ("LW", dt.date(2026, 9, 22)), ("MR", dt.date(2026, 9, 7))}
plan = defaultdict(list)   # (kz, day) -> list of (order, hours)
for o in ORDERS:
    if o["status"] == "storniert":
        continue
    start = o["date"]
    end = o["done"] or (EXPORT if o["status"] == "in Arbeit" else None)
    if end is None:
        continue
    kzs = o["monteur"].split("/")
    days = [d for d in wd_range(start, end) if Q3[0] <= d <= Q3[-1]]
    if not days:
        continue
    if not o["big"]:
        days = [days[-1]]
        h = R.choice([1, 1.5, 2, 2.5, 3, 4]) if o["art"] != "Montage" else R.choice([3, 4, 5, 6])
    else:
        days = days[:12]
        h = 8
    for d in days:
        for k in kzs:
            if available(k, d) and (k, d) not in sick:
                plan[(k, d)].append((o, h))


def bs_day(kz, d):
    return (kz == "KS" and d.weekday() == 1) or (kz == "LW" and d.weekday() == 2)


TS_TRUTH = Counter()
for kz in ("JK", "MR", "DH", "AY", "KS", "LW"):
    for d in Q3:
        if not available(kz, d):
            TS[kz].append(dict(date=d, auftrag="", kunde="", beginn=None, ende=None, pause=0, std=8.0, taet="Urlaub", fzg="", bem=""))
            continue
        if (kz, d) in sick:
            TS[kz].append(dict(date=d, auftrag="", kunde="", beginn=None, ende=None, pause=0, std=8.0, taet="krank", fzg="", bem="AU liegt vor"))
            continue
        if bs_day(kz, d):
            TS[kz].append(dict(date=d, auftrag="", kunde="", beginn=None, ende=None, pause=0, std=8.0, taet="Berufsschule", fzg="", bem=""))
            continue
        items = plan.get((kz, d), [])
        t = dt.datetime.combine(d, dt.time(7, 0))
        total = 0
        rows = []
        for o, h in items:
            if total + h > 9:
                continue
            total += h
            pause = 30 if (h >= 6 or (total >= 6 and not any(r["pause"] for r in rows))) else 0
            b = t
            e = t + dt.timedelta(hours=h, minutes=pause)
            t = e + dt.timedelta(minutes=R.choice([15, 20, 30]))
            c = o["cust"]
            if c["gruppe"] in ("HV", "WG"):
                short = (c["variants"][1][0] if c["variants"] else c["name"].split()[0]) + " – " + o["objekt"]
            elif c["gruppe"] == "P":
                short = f"{c['nachname']}, {short_street(c['strasse'])}"
            else:
                short = c["name"].split(" GmbH")[0]
            ref = o["nr"] if R.random() < .82 else ""
            rows.append(dict(date=d, auftrag=ref, kunde=short, beginn=b.time(), ende=e.time(), pause=pause,
                             std=h, taet=o["beschr"] if R.random() < .7 else o["art"],
                             fzg=R.choice(["DD-BH 101", "DD-BH 102", "DD-BH 103"]), bem="", order=o["nr"]))
            TS_TRUTH[("mit_auftrag" if ref else "auftrag_nur_name")] += 1
        if total < 7:
            h = 8 - total if total else 8
            b = t
            e = t + dt.timedelta(hours=h, minutes=0 if total else 30)
            rows.append(dict(date=d, auftrag="", kunde="Werkstatt", beginn=b.time(), ende=e.time(),
                             pause=0 if total else 30, std=h,
                             taet=R.choice(["Lager aufräumen", "Fahrzeugpflege", "Material holen Großhandel", "Werkstatt / Vorbereitung"]),
                             fzg="", bem=""))
            TS_TRUTH["intern"] += 1
        TS[kz].extend(rows)
    # Feiertage ignorieren
TS_ROWS_TOTAL = sum(len(v) for v in TS.values())

# ---------------------------------------------------------------- Materialverbrauch
MAT = []
for o in ORDERS:
    if o["date"].year != 2026 or o["status"] not in ("erledigt", "abgerechnet", "in Arbeit"):
        continue
    n = R.randint(4, 9) if o["big"] else (R.randint(0, 2) if o["art"] == "Wartung" else R.randint(1, 4))
    for _ in range(n):
        md = o["done"] or o["date"]
        kz = o["monteur"].split("/")[0]
        r = R.random()
        if r < .06:
            MAT.append(dict(auftrag=o["nr"], date=md, kz=kz, art="", bez="Kleinmaterial pauschal", menge=Decimal(1), me="psch"))
        elif r < .1 and o["art"] in ("Heizungstausch", "Hydr. Abgleich", "Reparatur", "Strangsanierung"):
            p = R.choice(PUMPS)
            MAT.append(dict(auftrag=o["nr"], date=md, kz=kz, art=p["nr"], bez=p["bez"], menge=Decimal(1), me="Stk"))
        else:
            a = R.choice(ART)
            mg = Decimal(R.choice([1, 1, 1, 2, 2, 3, 4, 6, 10]) if a["ek"] < 60 else R.choice([1, 1, 2])) if a["me"] != "Stg" else q2(R.choice([0.5, 1, 1.5, 2, 2.5, 4]))
            MAT.append(dict(auftrag=o["nr"], date=md, kz=kz, art=a["nr"], bez=a["bez"], menge=mg, me=a["me"]))

# ---------------------------------------------------------------- Termine (JSON)
APPTS = []
PROSPECTS = [dict(name="Familie Bergmann", phone="+49 351 8331200", email="bergmann.radebeul@gmail.com",
                  street="Meißner Straße 288", zip="01445", city="Radebeul", title="Besichtigung Heizungstausch Neubau"),
             dict(name="Kerstin Vogt", phone="0160 9988776", email=None, street="Tolkewitzer Straße 41",
                  zip="01279", city="Dresden", title="Beratung Wärmepumpe"),
             dict(name="Praxis für Ergotherapie Lange", phone="+49 351 4440011", email="info@ergo-lange.de",
                  street="Teplitzer Straße 12", zip="01219", city="Dresden", title="Angebot Trinkwasser-Check")]
RES = [{"id": f"res_{e['pnr']}", "name": f"{e['vorname']} {e['nachname']}", "short": e["kz"]}
       for e in EMP if e["kz"] not in ("SB",)]
RES_BY_KZ = {r["short"]: r["id"] for r in RES}


def cust_json(c):
    if c["gruppe"] == "P":
        dn = f"{c['vorname']} {c['nachname']}" if c is not KRUEGER else "Dr. Andreas Krüger"
    elif c is MUELLER_G:
        dn = "Müller Getränke"
    else:
        dn = c["name"]
    ph = c["tel"]
    if R.random() < .4:
        ph = ph.replace("+49 ", "0")
    return {"display_name": dn, "phone": ph, "email": c["email"],
            "address": {"street": addr(c), "postal_code": c["plz"], "city": c["ort"]}}


aid = 0
APPT_TRUTH = []
for o in ORDERS:
    if o["status"] in ("offen", "geplant", "in Arbeit") or (o["status"] in ("erledigt", "abgerechnet") and o["done"] and o["done"] >= dt.date(2026, 8, 1)):
        start_d = o["done"] if o["status"] in ("erledigt", "abgerechnet") else next_wd(max(o["date"], dt.date(2026, 10, 16)) + dt.timedelta(days=R.randint(0, 20)))
        if start_d < dt.date(2026, 8, 1):
            continue
        aid += 1
        h0 = R.choice([7, 8, 9, 10, 13])
        dur = 8 if o["big"] else R.choice([1, 1.5, 2, 3])
        st = dt.datetime.combine(start_d, dt.time(h0, 0), TZ)
        en = st + dt.timedelta(hours=dur)
        cj = cust_json(o["cust"])
        if o["cust"]["gruppe"] in ("HV", "WG"):
            cj["address"]["street"] = o["objekt"]
            cj["address"]["postal_code"] = None
        APPTS.append({"id": f"tp_{R.randrange(16 ** 6):06x}", "title": o["beschr"], "start": st.isoformat(),
                      "end": en.isoformat(),
                      "resource_ids": [RES_BY_KZ[k] for k in o["monteur"].split("/") if k in RES_BY_KZ],
                      "status": "done" if o["status"] in ("erledigt", "abgerechnet") else R.choice(["confirmed", "confirmed", "tentative"]),
                      "customer": cj, "reference": o["nr"] if R.random() < .85 else None,
                      "notes": R.choice(["", "", "Schlüssel beim Hausmeister", "Hund im Haus", "bitte vorher anrufen", "Parken im Hof"]),
                      "created_at": dt.datetime.combine(o["date"], dt.time(9, 12), TZ).isoformat()})
        APPT_TRUTH.append(dict(id=APPTS[-1]["id"], auftrag=o["nr"], kunde=o["cust"]["true_id"], ref_im_json=APPTS[-1]["reference"]))
for k in CONTRACTS:
    if dt.date(2026, 10, 1) <= k["next"] <= dt.date(2026, 11, 30):
        d0 = next_wd(k["next"] + dt.timedelta(days=R.randint(10, 40)))
        st = dt.datetime.combine(d0, dt.time(R.choice([7, 8, 9, 10, 13, 14]), 0), TZ)
        cj = cust_json(k["cust"])
        cj["address"]["street"] = k["objekt"]
        APPTS.append({"id": f"tp_{R.randrange(16 ** 6):06x}", "title": "Jahreswartung " + k["anlage"].split(" (")[0],
                      "start": st.isoformat(), "end": (st + dt.timedelta(hours=1.5)).isoformat(),
                      "resource_ids": [RES_BY_KZ["JK"]], "status": "confirmed", "customer": cj,
                      "reference": k["nr"], "notes": "", "created_at": "2026-09-30T10:00:00+02:00",
                      "price_estimate": {"amount": float(k["pauschale"]), "currency": "EUR"}})
        APPT_TRUTH.append(dict(id=APPTS[-1]["id"], vertrag=k["nr"], kunde=k["cust"]["true_id"], ref_im_json=k["nr"]))
for p in PROSPECTS:
    d0 = next_wd(dt.date(2026, 10, R.randint(19, 30)))
    st = dt.datetime.combine(d0, dt.time(15, 0), TZ)
    APPTS.append({"id": f"tp_{R.randrange(16 ** 6):06x}", "title": p["title"], "start": st.isoformat(),
                  "end": (st + dt.timedelta(hours=1)).isoformat(), "resource_ids": [RES_BY_KZ["TB"]],
                  "status": "tentative",
                  "customer": {"display_name": p["name"], "phone": p["phone"], "email": p["email"],
                               "address": {"street": p["street"], "postal_code": p["zip"], "city": p["city"]}},
                  "reference": None, "notes": "Neukunde – über Website", "created_at": "2026-10-13T18:22:00+02:00"})
    APPT_TRUTH.append(dict(id=APPTS[-1]["id"], kunde="NEU:" + p["name"]))
APPTS.sort(key=lambda a: a["start"])

# ================================================================= Dateien schreiben
# 1) Kundenstamm.xlsx --------------------------------------------------------------------
TEL_STYLE = {}


def write_kundenstamm(path, custs, stand, suppliers=SUPPLIERS, record=True):
    wb = xlsxwriter.Workbook(str(path))
    bold = wb.add_format({"bold": True, "bg_color": "#DDEBF7", "border": 1})
    title = wb.add_format({"bold": True, "font_size": 14})
    datef = wb.add_format({"num_format": "dd.mm.yyyy"})
    plzf = wb.add_format({"num_format": "00000"})
    telf = wb.add_format({"num_format": "0"})
    ws = wb.add_worksheet("Kunden")
    ws.write(0, 0, "Kundenliste Brenner Haustechnik GmbH", title)
    ws.write(1, 0, f"Stand: {stand} – Export aus Handwerkersoftware, bitte NICHT sortieren!")
    hdr = ["Kd.-Nr.", "Anrede", "Name / Firma", "Ansprechpartner", "Straße", "PLZ", "Ort", "Telefon", "Mobil",
           "E-Mail", "Kundengruppe", "Zahlungsziel", "Kunde seit", "Bemerkung"]
    for j, h in enumerate(hdr):
        ws.write(2, j, h, bold)
    rr = random.Random(7)
    for i, c in enumerate(custs, start=3):
        ws.write_number(i, 0, c["kdnr"])
        ws.write(i, 1, c["anrede"])
        ws.write(i, 2, c["name"])
        ws.write(i, 3, c["ansprech"] or "")
        ws.write(i, 4, f"{c['strasse']} {c['hnr']}")
        plz_int = int(c["plz"])
        if rr.random() < .5:
            ws.write_number(i, 5, plz_int, plzf)
        else:
            ws.write_number(i, 5, plz_int)
        ws.write(i, 6, c["ort"])
        tel_digits = c["tel"].replace("+49 ", "").replace(" ", "")
        r = rr.random()
        if r < .45:
            ws.write_number(i, 7, int(tel_digits), telf)       # 3518494420 – führende 0 weg
            TEL_STYLE[c["kdnr"]] = "zahl" if record else TEL_STYLE.get(c["kdnr"])
        elif r < .7:
            area, num = c["tel"].replace("+49 ", "0").split(" ", 1)
            ws.write(i, 7, f"{area} {num}")
            TEL_STYLE[c["kdnr"]] = "text" if record else TEL_STYLE.get(c["kdnr"])
        elif r < .9:
            area, num = c["tel"].replace("+49 ", "0").split(" ", 1)
            ws.write(i, 7, f"{area}/{num}")
            TEL_STYLE[c["kdnr"]] = "slash" if record else TEL_STYLE.get(c["kdnr"])
        else:
            ws.write(i, 7, c["tel"])
            TEL_STYLE[c["kdnr"]] = "intl" if record else TEL_STYLE.get(c["kdnr"])
        if c["mobil"]:
            if rr.random() < .5:
                ws.write_number(i, 8, int(c["mobil"].replace("+49 ", "").replace(" ", "")), telf)
            else:
                ws.write(i, 8, c["mobil"].replace("+49 ", "0").replace(" ", "-"))
        ws.write(i, 9, c["email"] or "")
        ws.write(i, 10, c["gruppe"])
        z = c["ziel"]
        if c["skonto"] and UMGEHUNG:
            ws.write(i, 11, f"{z} Tage netto")
        elif c["skonto"]:
            ws.write(i, 11, f"8 Tage 2% Skonto, {z} Tage netto")
        else:
            ws.write(i, 11, rr.choice([f"{z} Tage", z, f"{z} Tage netto"]) if z != 21 else "21 Tage")
        s = c["seit"]
        r = rr.random()
        if r < .7:
            ws.write_datetime(i, 12, dt.datetime.combine(s, dt.time()), datef)
        elif r < .8:
            ws.write(i, 12, s.strftime("%m/%Y"))
        elif r < .9:
            ws.write_number(i, 12, s.year)
        bem = c["bem"] or ""
        if c["skonto"] and UMGEHUNG:
            bem = (bem + "; " if bem else "") + "8 Tage 2% Skonto"
        ws.write(i, 13, bem)
    ws.autofilter(2, 0, 2 + len(custs), len(hdr) - 1)
    ws.freeze_panes(3, 0)
    ws2 = wb.add_worksheet("Lieferanten")
    h2 = ["Lief.-Nr.", "Firma", "Straße", "PLZ", "Ort", "Unsere Kd-Nr. dort", "Telefon", "IBAN", "Konditionen"]
    for j, h in enumerate(h2):
        ws2.write(0, j, h, bold)
    for i, s in enumerate(suppliers, start=1):
        ws2.write_number(i, 0, s["nr"])
        ws2.write(i, 1, s["name"])
        ws2.write(i, 2, s["strasse"])
        ws2.write_number(i, 3, int(s["plz"]))
        ws2.write(i, 4, s["ort"])
        ws2.write(i, 5, s["kd"])
        ws2.write(i, 6, s["tel"])
        ws2.write(i, 7, " ".join(s["iban"][k:k + 4] for k in range(0, 22, 4)))
        ws2.write(i, 8, rr.choice(["30 Tage netto", "14 Tage 3% Skonto", "Lastschrift"]))
    ws3 = wb.add_worksheet("Hinweise")
    for i, line in enumerate(["Kundengruppen: P = Privat, G = Gewerbe, HV = Hausverwaltung, WG = Wohnungsgenossenschaft",
                              "Zahlungsziel Standard 14 Tage, HV/WG 30 Tage",
                              "Neue Kunden bitte immer mit Telefonnummer anlegen!",
                              "Doppelte Kunden NICHT löschen, sondern an Sabine melden (Zuordnung Rechnungen)"]):
        ws3.write(i, 0, line)
    wb.close()


write_kundenstamm(OUT / "Kundenstamm.xlsx", CUST, "01.10.2026")
TRUTH["dateien"]["Kundenstamm.xlsx"] = "Blatt Kunden (Kopf in Zeile 3, 136 Zeilen, PLZ als Zahl, Tel teils Zahl), Blatt Lieferanten (6), Blatt Hinweise (Freitext)"

# 2) Auftragsbuch.xlsx -------------------------------------------------------------------
MONTHS = ["Januar", "Februar", "März", "April", "Mai", "Juni", "Juli", "August", "September", "Oktober", "November", "Dezember"]
DATE_STYLE = {}
WERT_STYLE = {}


def write_auftragsbuch(path, orders, stand):
    wb = xlsxwriter.Workbook(str(path))
    bold = wb.add_format({"bold": True, "bg_color": "#FCE4D6", "border": 1})
    title = wb.add_format({"bold": True, "font_size": 14})
    datef = wb.add_format({"num_format": "dd.mm.yyyy"})
    eur = wb.add_format({"num_format": '#,##0.00 "€"'})
    sumf = wb.add_format({"bold": True, "num_format": '#,##0.00 "€"', "top": 1})
    sumt = wb.add_format({"bold": True, "top": 1})
    rr = random.Random(11)
    hdr = ["Auftr.-Nr.", "Datum", "Kd-Nr", "Kunde", "Objekt / Einsatzort", "Art", "Beschreibung", "Monteur",
           "Status", "Wert netto", "erledigt am", "Rechnung"]
    for year in (2025, 2026):
        ws = wb.add_worksheet(str(year))
        ws.write(0, 0, f"Auftragsbuch {year} – Brenner Haustechnik", title)
        ws.write(1, 0, f"Stand {stand}")
        for j, h in enumerate(hdr):
            ws.write(2, j, h, bold)
        row = 3
        yo = [o for o in orders if o["date"].year == year]
        for m in range(1, 13):
            mo = [o for o in yo if o["date"].month == m]
            if not mo:
                continue
            first = row
            for o in mo:
                ws.write(row, 0, o["nr"])
                d = o["date"]
                r = rr.random()
                if year == 2025 or r < .6:
                    ws.write_datetime(row, 1, dt.datetime.combine(d, dt.time()), datef)
                    DATE_STYLE[o["nr"]] = "excel"
                elif r < .8:
                    ws.write(row, 1, dd(d))
                    DATE_STYLE[o["nr"]] = "text_dd.mm.yyyy"
                elif r < .9:
                    ws.write(row, 1, f"{d.day}.{d.month}.{d.strftime('%y')}")
                    DATE_STYLE[o["nr"]] = "text_d.m.yy"
                else:
                    ws.write(row, 1, d.strftime("%d/%m/%y"))
                    DATE_STYLE[o["nr"]] = "text_dd/mm/yy"
                nm, with_nr = o.get("_nv") or name_variant(o["cust"])
                o["_nv"] = (nm, with_nr)
                o["typed_name"], o["typed_nr"] = nm, with_nr
                if with_nr:
                    ws.write_number(row, 2, o["cust"]["kdnr"])
                ws.write(row, 3, nm)
                ws.write(row, 4, o["objekt"])
                ws.write(row, 5, o["art"])
                ws.write(row, 6, o["beschr"])
                mon = o["monteur"]
                if rr.random() < .05 and "/" not in mon and mon in EMP_BY_KZ:
                    mon = EMP_BY_KZ[mon]["nachname"]
                ws.write(row, 7, mon)
                ws.write(row, 8, o["status"])
                w = o["wert"]
                r = rr.random()
                if o["tag"] == "petra_angebot":
                    ws.write(row, 9, "auf Anfrage")
                    WERT_STYLE[o["nr"]] = "auf Anfrage"
                elif o["status"] == "storniert" and r < .5:
                    WERT_STYLE[o["nr"]] = "leer"
                elif r < (.85 if year == 2025 else .70):
                    ws.write_number(row, 9, float(w), eur)
                    WERT_STYLE[o["nr"]] = "zahl"
                elif r < (.95 if year == 2025 else .85):
                    ws.write(row, 9, de(w) + " €")
                    WERT_STYLE[o["nr"]] = "text_eur"
                elif o["art"] == "Wartung":
                    ws.write(row, 9, f"pauschal {int(w)},-")
                    WERT_STYLE[o["nr"]] = "pauschal"
                elif r < .93:
                    ws.write(row, 9, "nach Aufwand")
                    WERT_STYLE[o["nr"]] = "nach Aufwand"
                else:
                    ws.write(row, 9, f"ca. {int(round(float(w), -1))}")
                    WERT_STYLE[o["nr"]] = "ca"
                if o["done"]:
                    if rr.random() < .03:
                        ws.write(row, 10, f"KW {o['done'].isocalendar()[1]}")
                    else:
                        ws.write_datetime(row, 10, dt.datetime.combine(o["done"], dt.time()), datef)
                if o["invoices"] and rr.random() < .75:
                    ws.write(row, 11, ", ".join(i["nr"] for i in o["invoices"] if i["kind"] not in ("storno",)))
                    o["inv_in_book"] = True
                row += 1
            ws.write(row, 3, f"Summe {MONTHS[m - 1]} {year}", sumt)
            tot = sum(float(o["wert"]) for o in mo if WERT_STYLE[o["nr"]] == "zahl")
            ws.write_formula(row, 9, f"=SUBTOTAL(9,J{first + 1}:J{row})", sumf, tot)
            row += 2
        ws.freeze_panes(3, 1)
    # Auswertung (Pivot-artig)
    ws = wb.add_worksheet("Auswertung")
    ws.write(0, 0, "Aufträge 2026 je Monteur und Monat", title)
    ws.write(2, 0, "Monteur", bold)
    for m in range(1, 11):
        ws.write(2, m, MONTHS[m - 1][:3], bold)
    ws.write(2, 11, "Gesamt", bold)
    for i, kz in enumerate(["TB", "JK", "MR", "DH", "AY"], start=3):
        ws.write(i, 0, kz)
        tot = 0
        for m in range(1, 11):
            cnt = sum(1 for o in orders if o["date"].year == 2026 and o["date"].month == m and kz in o["monteur"].split("/"))
            tot += cnt
            ws.write_formula(i, m, f'=COUNTIFS(\'2026\'!H:H,"*{kz}*",\'2026\'!B:B,">="&DATE(2026,{m},1),\'2026\'!B:B,"<"&DATE(2026,{m + 1},1))', None, cnt)
        ws.write_formula(i, 11, f"=SUM(B{i + 1}:K{i + 1})", None, tot)
    wb.close()


write_auftragsbuch(OUT / "Auftragsbuch.xlsx", ORDERS, "15.10.2026")
TRUTH["dateien"]["Auftragsbuch.xlsx"] = "Blätter 2025/2026 (Kopf Zeile 3, Monats-Summenzeilen, Datum gemischt), Blatt Auswertung (Pivot, keine Datensätze)"

# 3) Rechnungsausgangsbuch.csv (cp1252) ---------------------------------------------------
with open(OUT / "Rechnungsausgangsbuch_2025-2026.csv", "w", encoding="cp1252", newline="") as f:
    w = csv.writer(f, delimiter=";", lineterminator="\r\n")
    w.writerow(["Belegnr", "Belegdatum", "Kd-Nr", "Kunde", "Betreff", "Netto", "USt 19%", "Brutto", "Fällig am",
                "Bezahlt am", "Status"])
    for inv in INVOICES:
        if inv["date"] > EXPORT:
            continue
        paid_s = dd(inv["paid"]) if inv["paid"] and inv["status"] in ("bezahlt", "teilbezahlt") and inv["paid"] <= EXPORT else ""
        w.writerow([inv["nr"], dd(inv["date"]), inv["cust"]["kdnr"], inv["cust"]["name"], inv["betreff"],
                    de(inv["netto"]), de(inv["ust"]), de(inv["brutto"]), dd(inv["faellig"]), paid_s, inv["status"]])
TRUTH["dateien"]["Rechnungsausgangsbuch_2025-2026.csv"] = "cp1252, ';', Dezimalkomma, Nummern mit Lücken, Storno"

# 4) Volksbank-Umsätze (altes VR-NetWorld-Format, S/H-Spalte, absteigend) ------------------
with open(OUT / "Umsaetze_Volksbank_01-09_2026.csv", "w", encoding="cp1252", newline="") as f:
    w = csv.writer(f, delimiter=";", quoting=csv.QUOTE_ALL, lineterminator="\r\n")
    f.write('"Volksbank Dresden-Bautzen eG"\r\n\r\n"Umsatzanzeige"\r\n\r\n')
    f.write('"BLZ:";"85090000";;"Datum:";"15.10.2026"\r\n')
    f.write('"Konto:";"4471230";;"Uhrzeit:";"07:58:12"\r\n')
    f.write('"Abfrage von:";"Sabine Brenner";;"Kontoinhaber:";"Brenner Haustechnik GmbH"\r\n\r\n')
    f.write('"Zeitraum:";;"von:";"01.01.2026";"bis:";"30.09.2026"\r\n')
    f.write('"Betrag in EUR:";;"von:";" ";"bis:";" "\r\n')
    f.write('"Sortiert nach:";"Buchungstag";"absteigend"\r\n\r\n')
    w.writerow(["Buchungstag", "Valuta", "Auftraggeber/Zahlungsempfänger", "Empfänger/Zahlungspflichtiger", "Konto-Nr.",
                "IBAN", "BLZ", "BIC", "Vorgang/Verwendungszweck", "Kundenreferenz", "Währung", "Umsatz", " "])
    for b in sorted(BANK, key=lambda b: b["date"], reverse=True):
        w.writerow([dd(b["date"]), dd(b["valuta"]), "Brenner Haustechnik GmbH", b["name"], "", b["iban"], "",
                    "GENODEF1DRS" if b["iban"][4:12] == "85090000" else "", b["text"], "", "EUR",
                    de(abs(b["amount"])), "S" if b["amount"] < 0 else "H"])
    f.write("\r\n")
    f.write(f'"01.01.2026";;;;;;;;;"Anfangssaldo";"EUR";"{de(OPEN_BAL)}";"H"\r\n')
    f.write(f'"30.09.2026";;;;;;;;;"Endsaldo";"EUR";"{de(abs(end_bal))}";"{"H" if end_bal >= 0 else "S"}"\r\n')
TRUTH["dateien"]["Umsaetze_Volksbank_01-09_2026.csv"] = "altes VR-NetWorld-Format: 12 Zeilen Vorspann, Betrag immer positiv + S/H-Spalte, absteigend sortiert, mehrzeiliger Verwendungszweck, Saldo-Fußzeilen"

# 5) Stundenzettel_Q3_2026.xlsx ---------------------------------------------------------
wb = xlsxwriter.Workbook(str(OUT / "Stundenzettel_Q3_2026.xlsx"))
bold = wb.add_format({"bold": True, "bg_color": "#E2EFDA", "border": 1})
title = wb.add_format({"bold": True, "font_size": 13})
datef = wb.add_format({"num_format": "dd.mm.yy"})
timef = wb.add_format({"num_format": "hh:mm"})
numf = wb.add_format({"num_format": "0.00"})
sumf = wb.add_format({"bold": True, "num_format": "0.00", "top": 1})
WDN = ["Mo", "Di", "Mi", "Do", "Fr", "Sa", "So"]
rr = random.Random(5)
SHEETNAME = {"JK": "Kowalski", "MR": "Richter", "DH": "Hoffmann", "AY": "Yilmaz", "KS": "Schmidt KS", "LW": "Weber (Azubi)"}
TS_TEXT_TIMES = 0
for kz, rows in TS.items():
    e = EMP_BY_KZ[kz]
    ws = wb.add_worksheet(SHEETNAME[kz])
    ws.write(0, 0, f"Stundenzettel Juli – September 2026   {e['vorname']} {e['nachname']} (PNr {e['pnr']})", title)
    hdr = ["Datum", "Tag", "Auftrag", "Kunde / Baustelle", "Beginn", "Ende", "Pause (Min)", "Std.", "Tätigkeit", "Fahrzeug", "Bemerkung"]
    for j, h in enumerate(hdr):
        ws.write(2, j, h, bold)
    row = 3
    week = None
    wk_first = row
    for r_ in rows + [None]:
        wkn = r_["date"].isocalendar()[1] if r_ else None
        if week is not None and wkn != week:
            ws.write(row, 1, f"Summe KW {week}")
            tot = sum(x["std"] for x in rows if x["date"].isocalendar()[1] == week)
            ws.write_formula(row, 7, f"=SUM(H{wk_first + 1}:H{row})", sumf, tot)
            row += 1
            wk_first = row
        if r_ is None:
            break
        week = wkn
        ws.write_datetime(row, 0, dt.datetime.combine(r_["date"], dt.time()), datef)
        ws.write(row, 1, WDN[r_["date"].weekday()])
        ws.write(row, 2, r_["auftrag"])
        ws.write(row, 3, r_["kunde"])
        if r_["beginn"] is not None:
            if rr.random() < .03:
                ws.write(row, 4, r_["beginn"].strftime("%-H:%M"))
                TS_TEXT_TIMES += 1
            elif UMGEHUNG:
                ws.write(row, 4, r_["beginn"].strftime("%H:%M"))
            else:
                ws.write_datetime(row, 4, r_["beginn"], timef)
            if UMGEHUNG:
                ws.write(row, 5, r_["ende"].strftime("%H:%M"))
            else:
                ws.write_datetime(row, 5, r_["ende"], timef)
            ws.write_number(row, 6, r_["pause"])
            ws.write_formula(row, 7, f"=(F{row + 1}-E{row + 1})*24-G{row + 1}/60", numf, r_["std"])
        else:
            ws.write_number(row, 7, r_["std"], numf)
        ws.write(row, 8, r_["taet"])
        ws.write(row, 9, r_["fzg"])
        ws.write(row, 10, r_["bem"])
        row += 1
wb.close()
TRUTH["dateien"]["Stundenzettel_Q3_2026.xlsx"] = "6 Blätter (je Monteur, Name nur im Blattnamen/Titel), Kopf Zeile 3, Formel Std., Wochensummen"

# 6) Preisliste ------------------------------------------------------------------------
wb = xlsxwriter.Workbook(str(OUT / "Preisliste_Grosshandel_2026.xlsx"))
bold = wb.add_format({"bold": True, "border": 1, "bg_color": "#D9D9D9"})
grpf = wb.add_format({"bold": True, "font_color": "#1F4E78"})
eur = wb.add_format({"num_format": '#,##0.00 "€"'})
ws = wb.add_worksheet("Artikel")
hdr = ["Art.-Nr.", "Bezeichnung", "Hersteller", "ME", "EK netto", "VK netto", "Rabattgr.", "Bemerkung"]
for j, h in enumerate(hdr):
    ws.write(0, j, h, bold)
row = 1
rr = random.Random(3)
PRICE_STYLE = Counter()
for grp in GROUPS:
    ws.write(row, 0, grp, grpf)
    row += 1
    for a in [a for a in ART if a["gruppe"] == grp]:
        ws.write_number(row, 0, int(a["nr"]))
        ws.write(row, 1, a["bez"])
        ws.write(row, 2, a["hersteller"])
        ws.write(row, 3, a["me"])
        r = rr.random()
        if r < .8:
            ws.write_number(row, 4, float(a["ek"]), eur)
            ws.write_formula(row, 5, f"=ROUND(E{row + 1}*1.35,2)", eur, float(a["vk"]))
            PRICE_STYLE["formel"] += 1
        elif r < .9:
            ws.write(row, 4, de(a["ek"]) + " €")
            ws.write(row, 5, de(a["vk"]) + " €")
            PRICE_STYLE["text"] += 1
        elif r < .96:
            ws.write_number(row, 4, float(a["ek"]), eur)
            ws.write(row, 5, "auf Anfrage")
            a["vk_text"] = "auf Anfrage"
            PRICE_STYLE["auf_anfrage"] += 1
        else:
            ws.write(row, 4, "n. lieferbar")
            a["ek_text"] = "n. lieferbar"
            PRICE_STYLE["n_lieferbar"] += 1
        ws.write(row, 6, a["rabatt"])
        if rr.random() < .05:
            ws.write(row, 7, rr.choice(["Auslaufartikel", "Nachfolger folgt", "Lagerware"]))
        row += 1
ws2 = wb.add_worksheet("Pumps_Export")
h2 = ["Item No.", "Description", "List price (EUR)", "Discount group", "Lead time (days)", "EAN"]
for j, h in enumerate(h2):
    ws2.write(0, j, h, bold)
for i, p in enumerate(PUMPS, start=1):
    ws2.write_number(i, 0, int(p["nr"]))
    ws2.write(i, 1, p["bez"])
    if isinstance(p["price"], Decimal):
        ws2.write_number(i, 2, float(p["price"]))
    elif p["price"]:
        ws2.write(i, 2, p["price"])
    ws2.write(i, 3, p["dgroup"])
    ws2.write_number(i, 4, p["lead"])
    ws2.write_number(i, 5, p["ean"])     # 13-stellig als Zahl
wb.close()
TRUTH["dateien"]["Preisliste_Grosshandel_2026.xlsx"] = "Blatt Artikel (Gruppen-Zwischenzeilen, VK=ROUND(EK*1,35;2), 'auf Anfrage', Text-Beträge), Blatt Pumps_Export (englisch, EAN als Zahl)"

# 7) Materialverbrauch (ohne Kopfzeile) --------------------------------------------------
with open(OUT / "Materialverbrauch_2026.csv", "w", encoding="utf-8", newline="") as f:
    w = csv.writer(f, delimiter=";", lineterminator="\n")
    for m in MAT:
        mg = str(m["menge"]).replace(".", ",")
        if mg.endswith(",00"):
            mg = mg[:-3]
        w.writerow([m["auftrag"], dd(m["date"]), m["kz"], m["art"], m["bez"], mg, m["me"]])
TRUTH["dateien"]["Materialverbrauch_2026.csv"] = "OHNE Kopfzeile: Auftrag;Datum;Kürzel;Art.-Nr.;Bezeichnung;Menge;ME"

# 8) Termine JSON ----------------------------------------------------------------------
with open(OUT / "Termine_TerminPilot_Export.json", "w", encoding="utf-8") as f:
    json.dump({"export": {"source": "TerminPilot Online", "account": "brenner-haustechnik",
                          "exported_at": "2026-10-15T07:42:11Z", "timezone": "Europe/Berlin", "version": "3.2"},
               "resources": RES, "appointments": APPTS}, f, ensure_ascii=False, indent=2)
TRUTH["dateien"]["Termine_TerminPilot_Export.json"] = "verschachteltes JSON, resource_ids -> resources, customer als Objekt, reference = Auftrag oder Vertrag"

# 9) Mitarbeiter.csv (HR-Tool, englisch, US-Datum) -------------------------------------
with open(OUT / "Mitarbeiter_Personio.csv", "w", encoding="utf-8", newline="") as f:
    w = csv.writer(f, delimiter=",", lineterminator="\n")
    w.writerow(["Employee ID", "First name", "Last name", "Position", "Department", "Hire date", "Weekly hours", "Mobile phone", "Email", "Short code"])
    for e in EMP:
        w.writerow([e["pnr"], e["vorname"], e["nachname"], e["position"], e["dept"], e["hire"].strftime("%m/%d/%Y"),
                    e["hours"], e["mobile"], e["email"], e["kz"]])
TRUTH["dateien"]["Mitarbeiter_Personio.csv"] = "englische Spalten, Hire date im US-Format MM/DD/YYYY"

# 10) Wartungsverträge.docx -------------------------------------------------------------
doc = Document()
doc.add_heading("Wartungsverträge – Übersicht 2026", level=1)
doc.add_paragraph("Stand: September 2026. Pauschalen netto zzgl. MwSt. Material wird gesondert berechnet.")
t = doc.add_table(rows=1, cols=9)
t.style = "Table Grid"
for j, h in enumerate(["Vertrag", "Kd-Nr", "Kunde", "Objekt / Standort", "Anlage", "Intervall", "Pauschale (netto)", "Letzte Wartung", "Nächste Wartung"]):
    t.rows[0].cells[j].text = h
rr = random.Random(9)
for k in CONTRACTS:
    cells = t.add_row().cells
    c = k["cust"]
    cells[0].text = k["nr"]
    cells[1].text = str(c["kdnr"])
    cells[2].text = c["name"] if rr.random() < .7 else name_variant(c)[0]
    cells[3].text = k["objekt"]
    cells[4].text = k["anlage"]
    cells[5].text = k["intervall"]
    cells[6].text = de(k["pauschale"]) + " €"
    cells[7].text = dd(k["last"])
    nxt = k["next"]
    cells[8].text = rr.choice([f"{MONTHS[nxt.month - 1][:3]}. {nxt.year}", nxt.strftime("%m/%Y"), f"KW {nxt.isocalendar()[1]}/{nxt.year}"])
doc.add_paragraph("Kündigungsfrist jeweils 3 Monate zum Jahresende. Bei HV/WG: Rechnung je Objekt.")
doc.save(str(OUT / "Wartungsvertraege_2026.docx"))
TRUTH["dateien"]["Wartungsvertraege_2026.docx"] = f"Word-Tabelle, {len(CONTRACTS)} Verträge"

# 11) PDFs ------------------------------------------------------------------------------


def pdf_doc(path, kind, nr, date, c, positions, extra_lines=()):
    cv = canvas.Canvas(str(path), pagesize=A4)
    W, H = A4
    cv.setFont("Helvetica-Bold", 14)
    cv.drawString(20 * mm, H - 20 * mm, "Brenner Haustechnik GmbH")
    cv.setFont("Helvetica", 8)
    cv.drawString(20 * mm, H - 25 * mm, "Sanitär · Heizung · Klima   Hechtstraße 41 · 01097 Dresden · Tel. 0351 8023344")
    cv.setFont("Helvetica", 10)
    y = H - 45 * mm
    lines = [c["name"] if c["gruppe"] != "P" else f"{c['anrede']} {c['vorname']} {c['nachname']}"]
    if c["ansprech"]:
        lines.append(c["ansprech"])
    lines += [addr(c), f"{c['plz']} {c['ort']}"]
    for ln in lines:
        cv.drawString(20 * mm, y, ln)
        y -= 5 * mm
    cv.drawString(130 * mm, H - 45 * mm, f"{kind} Nr.: {nr}")
    cv.drawString(130 * mm, H - 50 * mm, f"Datum: {dd(date)}")
    cv.drawString(130 * mm, H - 55 * mm, f"Kunden-Nr.: {c['kdnr']}")
    y = H - 80 * mm
    cv.setFont("Helvetica-Bold", 12)
    cv.drawString(20 * mm, y, f"{kind} {nr}")
    y -= 8 * mm
    cv.setFont("Helvetica", 9)
    for ln in extra_lines:
        cv.drawString(20 * mm, y, ln)
        y -= 5 * mm
    y -= 3 * mm
    cv.setFont("Helvetica-Bold", 9)
    for x, h in ((20, "Pos"), (30, "Bezeichnung"), (125, "Menge"), (145, "EP netto"), (172, "GP netto")):
        cv.drawString(x * mm, y, h)
    y -= 6 * mm
    cv.setFont("Helvetica", 9)
    total = Decimal(0)
    for i, (bez, menge, me, ep) in enumerate(positions, start=1):
        gp = q2(Decimal(str(menge)) * Decimal(str(ep)))
        total += gp
        cv.drawString(20 * mm, y, str(i))
        cv.drawString(30 * mm, y, bez)
        cv.drawString(125 * mm, y, f"{str(menge).replace('.', ',')} {me}")
        cv.drawRightString(165 * mm, y, de(ep))
        cv.drawRightString(190 * mm, y, de(gp))
        y -= 5.5 * mm
    ust = q2(total * Decimal("0.19"))
    y -= 4 * mm
    for lab, val in (("Summe netto", total), ("zzgl. 19 % MwSt.", ust), ("Gesamtbetrag", total + ust)):
        cv.drawString(125 * mm, y, lab)
        cv.drawRightString(190 * mm, y, de(val) + " €")
        y -= 5.5 * mm
    cv.setFont("Helvetica", 7)
    cv.drawString(20 * mm, 15 * mm, f"Volksbank Dresden-Bautzen eG · IBAN {OWN_IBAN} · BIC GENODEF1DRS · USt-IdNr. DE287654321 · GF Thomas Brenner · AG Dresden HRB 27711")
    cv.showPage()
    cv.save()
    return total, ust


KRUEGER_POS = [("Demontage alte Sanitärobjekte inkl. Entsorgung", 1, "psch", "890.00"),
               ("Vorwandinstallation WC inkl. Wand-WC spülrandlos", 1, "Stk", "1240.00"),
               ("Bodengleiche Dusche inkl. Duschrinne 800 mm", 1, "Stk", "3480.00"),
               ("Doppelwaschtischanlage inkl. Armaturen", 1, "Stk", "2150.00"),
               ("Rohrinstallation Kalt-/Warmwasser, Abwasser", 1, "psch", "2870.00"),
               ("Handtuchheizkörper 1200 mm", 1, "Stk", "785.00"),
               ("Fliesenarbeiten (Nachunternehmer)", 32, "m²", "95.00"),
               ("Montage, Inbetriebnahme, Dichtheitsprüfung", 1, "psch", "1025.00")]
k_net, k_ust = pdf_doc(OUT / "Angebot_AN-26-044_Krueger.pdf", "Angebot", "AN-26-044", dt.date(2026, 9, 2), KRUEGER,
                       KRUEGER_POS, ["Badsanierung Obergeschoss, Bergstraße 61", "Gültig bis 31.10.2026"])
assert k_net == Decimal("15480.00"), k_net
K_BRUTTO = k_net + k_ust
mu_inv = O_MUELLER_AUG["invoices"][0]
pdf_doc(OUT / f"Rechnung_{mu_inv['nr']}.pdf", "Rechnung", mu_inv["nr"], mu_inv["date"], MUELLER_G,
        [("Warmwasserspeicher 300 l inkl. Anschlussset", 1, "Stk", "610.00"),
         ("Demontage Altspeicher + Entsorgung", 1, "psch", "85.00"),
         ("Montagestunden Geselle", 4, "Std", "60.00")],
        [f"Auftrag {O_MUELLER_AUG['nr']} – Austausch Warmwasserspeicher Lager", f"Leistungsdatum: {dd(O_MUELLER_AUG['done'])}",
         f"Zahlbar bis {dd(mu_inv['faellig'])} ohne Abzug"])
assert mu_inv["netto"] == Decimal("935.00")
TRUTH["dateien"]["Angebot_AN-26-044_Krueger.pdf"] = f"Angebot, netto {de(k_net)}, brutto {de(K_BRUTTO)}"
TRUTH["dateien"][f"Rechnung_{mu_inv['nr']}.pdf"] = f"Rechnungs-PDF zu {mu_inv['nr']} (auch im Rechnungsausgangsbuch)"

# 12) E-Mails ---------------------------------------------------------------------------


def eml(fname, frm, to, subj, date, text, html=None, attachments=()):
    m = EmailMessage()
    m["From"] = frm
    m["To"] = to
    m["Subject"] = subj
    m["Date"] = format_datetime(date)
    m["Message-ID"] = f"<{abs(hash(fname)) % 10 ** 12}@{frm.split('@')[-1].strip('>')}>"
    m.set_content(text)
    if html:
        m.add_alternative(html, subtype="html")
    for name, data, mt in attachments:
        maintype, subtype = mt.split("/")
        m.add_attachment(data, maintype=maintype, subtype=subtype, filename=name)
    (OUT / "Mails" / fname).write_bytes(m.as_bytes(policy=policy.SMTP))


eml("2026-09-14_Re_Angebot_AN-26-044.eml", "Dr. Andreas Krüger <andreas.krueger@t-online.de>",
    "info@brenner-haustechnik.de", "Re: Angebot AN-26-044 Badsanierung",
    dt.datetime(2026, 9, 14, 21, 7, tzinfo=TZ),
    f"Sehr geehrter Herr Brenner,\n\nvielen Dank für Ihr Angebot AN-26-044 vom 02.09.2026 über {de(K_BRUTTO)} € brutto.\n"
    "Grundsätzlich möchten wir Ihnen den Auftrag erteilen. Zwei Punkte:\n"
    "1) Statt der Doppelwaschtischanlage reicht ein einzelner Waschtisch (bitte Preis anpassen).\n"
    "2) Können Sie Ende Oktober beginnen? Wir sind ab 3.11. zwei Wochen verreist.\n\n"
    "Mit freundlichen Grüßen\nAndreas Krüger\nBergstraße 61, 01069 Dresden\nTel. 0351 4718823\n")
eml("2026-09-22_AW_Zahlungserinnerung.eml", "Mueller GmbH & Co. KG - Buchhaltung <buchhaltung@mueller-getraenke.de>",
    "buero@brenner-haustechnik.de", f"AW: Zahlungserinnerung {mu_inv['nr']}",
    dt.datetime(2026, 9, 22, 10, 31, tzinfo=TZ),
    f"Sehr geehrte Frau Brenner,\n\nwir haben die Rechnung {mu_inv['nr']} (Austausch Warmwasserspeicher, Auftrag {O_MUELLER_AUG['nr']}) "
    f"bereits am 18.09.2026 überwiesen, Betrag {de(mu_inv['brutto'])} EUR.\n"
    "Bitte prüfen Sie Ihren Zahlungseingang und betrachten Sie die Erinnerung als gegenstandslos.\n\n"
    "Unsere Kundennummer bei Ihnen: 10023\n\nMit freundlichen Grüßen\ni.A. Wagner\nMüller GmbH & Co. KG – Getränkegroßhandel\n"
    "Großenhainer Straße 101 · 01127 Dresden\n")
objs = [k["objekt"] for k in CONTRACTS if k["cust"] is HV_SCHULZE]
html = ("<html><body><p>Sehr geehrter Herr Brenner,</p><p>bitte führen Sie die jährliche Wartung gemäß Wartungsvertrag "
        "<b>WV-2024-012</b> in folgenden Objekten durch:</p><table border='1'><tr><th>Objekt</th><th>Ansprechpartner vor Ort</th><th>Telefon</th></tr>"
        + "".join(f"<tr><td>{o}, 01169 Dresden</td><td>Hausmeister {n}</td><td>0172 {t}</td></tr>"
                  for o, n, t in zip(objs, ["Kunze", "Pietsch", "Wendler"], ["5512340", "5512341", "5512342"]))
        + "</table><p>Termine bitte direkt mit den Hausmeistern abstimmen. Rechnung bitte je Objekt.</p>"
          "<p>Mit freundlichen Grüßen<br>Ines Schulze-Wendt<br>Hausverwaltung Schulze &amp; Partner GbR<br>Schandauer Straße 34, 01277 Dresden</p></body></html>")
eml("2026-10-01_Wartungsauftrag_Gorbitz.eml", "Hausverwaltung Schulze & Partner <technik@hv-schulze-partner.de>",
    "info@brenner-haustechnik.de", "Wartungsauftrag Heizungsanlagen Objekte Gorbitz (WV-2024-012)",
    dt.datetime(2026, 10, 1, 8, 15, tzinfo=TZ),
    "Sehr geehrter Herr Brenner,\nbitte führen Sie die jährliche Wartung gemäß WV-2024-012 durch. Objekte siehe HTML-Version.\n",
    html=html)
ab_rows = R.sample(ART, 6)
ab_csv = "Pos;Art.-Nr.;Bezeichnung;Menge;ME;Einzelpreis;Gesamt\n" + "".join(
    f"{i};{a['nr']};{a['bez']};{q};{a['me']};{de(a['ek'])};{de(a['ek'] * q)}\n" for i, (a, q) in enumerate(zip(ab_rows, [2, 1, 4, 10, 1, 3]), start=1))
eml("2026-10-06_Auftragsbestaetigung_4500123456.eml", "Elbe-Haustechnik Großhandel <auftrag@elbe-haustechnik.de>",
    "einkauf@brenner-haustechnik.de", "Auftragsbestätigung 4500123456 / Kommission Krüger Bad",
    dt.datetime(2026, 10, 6, 14, 2, tzinfo=TZ),
    f"Sehr geehrte Damen und Herren,\nanbei unsere Auftragsbestätigung 4500123456 zu Ihrer Bestellung vom 05.10.2026, "
    f"Kommission: Krüger Bad / {O_KRUEGER_BAD['nr']}.\nIhre Kundennummer: 33810. Liefertermin: KW 43.\n\nElbe-Haustechnik Großhandel GmbH\n",
    attachments=[("AB_4500123456.csv", ab_csv.encode("utf-8"), "text/csv")])
eml("2026-10-12_Reklamation_Waschtischarmatur.eml", "Heike Müller <heike.mueller@web.de>", "info@brenner-haustechnik.de",
    f"Reklamation Waschtischarmatur – Auftrag {O_HEIKE2['nr']}", dt.datetime(2026, 10, 12, 19, 44, tzinfo=TZ),
    f"Guten Tag,\n\nIhr Monteur (Herr Hoffmann) hat am {dd(O_HEIKE2['done'])} bei uns die Waschtischarmatur im Bad OG repariert "
    f"(Auftrag {O_HEIKE2['nr']}). Seit gestern tropft es wieder. Bitte kommen Sie nochmal vorbei – auf Ihre Kosten.\n\n"
    "Heike Müller\nAm Hang 3\n01328 Dresden\nTel. 0351 2689044\n")
eml("2026-10-15_Notdienstplan_November.eml", "Sabine Brenner <sabine.brenner@brenner-haustechnik.de>",
    "Thomas Brenner <thomas.brenner@brenner-haustechnik.de>", "Notdienstplan November",
    dt.datetime(2026, 10, 15, 16, 20, tzinfo=TZ),
    "Hallo Thomas,\n\nNotdienst-Wochenenden November:\n07./08.11. Jens Kowalski (0172 3301103)\n"
    "14./15.11. Marco Richter (0172 3301104)\n21./22.11. Dennis Hoffmann (0172 3301105)\n28./29.11. Ayşe Yılmaz (0172 3301106)\n\n"
    "Kevin Schmidt geht jeweils mit (Azubi, nur tagsüber). Frau Schulze (Hüblerstr.) hat wegen der Wartung angerufen, "
    "bitte Termin bestätigen.\n\nLG Sabine\n")
TRUTH["dateien"]["Mails/*.eml"] = "6 Mails: Angebotsrückfrage Krüger, Zahlungsavis Müller GmbH (lag), HTML-Mail HV mit Objekt-Tabelle, AB Großhandel mit CSV-Anhang, Reklamation Heike Müller (Am Hang = K10131!), interner Notdienstplan"

# 13) Notizen.txt (UTF-8 mit BOM) -------------------------------------------------------
notes = f"""Notizen Thomas – bitte nicht löschen!!

Telefon:
Großhandel Elbe Hotline 0351 2109900
Fliesen Kaiser 035201 70815 (für Krüger Bad fragen, ab KW 44)
Schornsteinfeger Hr. Dittrich 0351 4567123

Fa. Müller (BAU, nicht Getränke!) zahlt immer erst nach 2. Mahnung -> künftig Abschlag/Vorkasse?
Frau Heike Müller (Am Hang, NICHT Kesselsdorfer!) – Reklamation Wasserhahn {O_HEIKE2['nr']}, Kulanz anbieten.
Hofmann/Hoffmann Klaus Borsbergstr. ist zweimal angelegt (10014 und 10119) -> Sabine fragen welche bleibt.
Neukundin: Petra Schmidt, Hüblerstr. 14 – Angebot Badumbau bis Ende Oktober. (Nicht verwechseln mit Gerda Schulze Nr. 8!)
Krüger Bad: Beginn 26.10., Einzelwaschtisch statt Doppel -> Angebot AN-26-044 anpassen (ca. -900 €)

Preise 2027:
Stundensatz Geselle 62 €, Meister 82 €, Azubi 30 €
Anfahrtspauschale 39 €
Wartung Gastherme pauschal 149 € netto (bisher 139)

Fahrzeuge: DD-BH 101 (Kowalski), DD-BH 102 (Richter), DD-BH 103 (Hoffmann/Yilmaz)
TÜV DD-BH 102 im Dezember!
"""
(OUT / "Notizen.txt").write_bytes(b"\xef\xbb\xbf" + notes.replace("\n", "\r\n").encode("utf-8"))
TRUTH["dateien"]["Notizen.txt"] = "Freitext UTF-8 mit BOM"

# 14) Altdaten_2024.zip (Kunden_alt.csv + Rechnungen_2024.csv) -----------------------------
old_c = [c for c in CUST if c["seit"] < dt.date(2024, 1, 1) and c["kdnr"] != 10119]
overlap = R.sample(old_c, 42)
for must in (MUELLER_G, HOFMANN, KRUEGER, MUELLER_BAU, WG_ELBTAL):
    if must not in overlap:
        overlap[R.randrange(len(overlap))] = must
overlap = sorted(set(map(id, overlap)))
overlap = [c for c in CUST if id(c) in overlap]
OLD = []
moved = set()
for i, c in enumerate(overlap):
    st, hnr, plz = c["strasse"], c["hnr"], c["plz"]
    if c["gruppe"] == "P" and len(moved) < 3 and c not in (HOFMANN, KRUEGER):
        s2, plz = R.choice(STREETS[:40])
        st, hnr = s2, str(R.randint(1, 99))
        moved.add(c["kdnr"])
    if c["gruppe"] == "P":
        n1 = f"{c['vorname']} {c['nachname']}" if c is not KRUEGER else "Dr. Andreas Krüger"
        n2 = ""
    elif c is MUELLER_G:
        n1, n2 = "Müller GmbH & Co. KG", "Getränkegroßhandel"
    elif c is MUELLER_BAU:
        n1, n2 = "Bauunternehmen Müller", "Inh. R. Müller"
    else:
        n1, n2 = c["name"], ""
    OLD.append(dict(nr=f"D-{1001 + i}", n1=n1, n2=n2, strasse=f"{st} {hnr}", plz=plz, ort=ORT.get(plz, "Dresden"),
                    tel=c["tel"].replace("+49 ", "0").replace(" ", "-"), email=c["email"] or "", true_id=c["true_id"],
                    erstellt=c["seit"]))
for j in range(13):
    vn = R.choice(FIRST_M + FIRST_F)
    nn = R.choice(LAST)
    st, plz = R.choice(STREETS)
    OLD.append(dict(nr=f"D-{1001 + len(OLD)}", n1=f"{vn} {nn}", n2="", strasse=f"{st} {R.randint(1, 140)}", plz=plz,
                    ort=ORT.get(plz, "Dresden"), tel=f"0351-{R.randint(2000000, 8999999)}", email="",
                    true_id=f"ALT:{vn} {nn}", erstellt=dt.date(2015, 1, 1) + dt.timedelta(days=R.randint(0, 3000))))
old_buf = io.StringIO()
w = csv.writer(old_buf, delimiter=",", quoting=csv.QUOTE_ALL, lineterminator="\r\n")
w.writerow(["KdNr", "Name1", "Name2", "Strasse", "PLZ", "Ort", "Telefon", "Email", "Erstellt"])
for o in OLD:
    w.writerow([o["nr"], o["n1"], o["n2"], o["strasse"], o["plz"], o["ort"], o["tel"], o["email"], o["erstellt"].strftime("%d.%m.%Y")])
inv24_buf = io.StringIO()
w = csv.writer(inv24_buf, delimiter=";", lineterminator="\r\n")
w.writerow(["ReNr", "ReDatum", "KdNr", "Name", "Netto", "MwSt", "Brutto", "Bezahlt"])
INV24 = []
for k in range(1, 141):
    o = R.choice(OLD)
    d = next_wd(dt.date(2024, 1, 3) + dt.timedelta(days=int(k * 2.55)))
    net = q2(R.choice([139, 139, 149]) if R.random() < .4 else R.uniform(90, 2500))
    ust = q2(net * Decimal("0.19"))
    INV24.append(dict(nr=f"2024-{k:04d}", kd=o["nr"], true_id=o["true_id"]))
    w.writerow([f"2024-{k:04d}", d.isoformat(), o["nr"], o["n1"], f"{net:.2f}", f"{ust:.2f}", f"{net + ust:.2f}", "ja" if R.random() < .97 else "nein"])
with zipfile.ZipFile(OUT / "Altdaten_2024.zip", "w", zipfile.ZIP_DEFLATED) as z:
    z.writestr("Altdaten/Kunden_alt.csv", old_buf.getvalue().encode("cp1252"))
    z.writestr("Altdaten/Rechnungen_2024.csv", inv24_buf.getvalue().encode("cp1252"))
    z.writestr("Altdaten/LIESMICH.txt", "Export aus alter Software (bis 12/2024). Kundennummern D-xxxx sind NICHT die neuen Nummern!\r\n".encode("cp1252"))
TRUTH["dateien"]["Altdaten_2024.zip"] = f"Kunden_alt.csv ({len(OLD)} Zeilen, alte Nummern D-xxxx, Name 'Vorname Nachname', 42 = aktuelle Kunden, 13 nur alt), Rechnungen_2024.csv (140, ISO-Datum, Dezimalpunkt)"

# ================================================================= Lieferung 2
L2_CUST = [dict(c) for c in CUST if c["kdnr"] != 10119]          # Dublette bereinigt
by = {c["kdnr"]: c for c in L2_CUST}
by[10052]["tel"] = "+49 351 3109000"
by[10061]["strasse"], by[10061]["hnr"], by[10061]["plz"] = "Rudolf-Leonhard-Straße", "7", "01097"
by[10023]["email"] = "rechnung@mueller-getraenke.de"
by[10045]["bem"] = "NUR gegen Abschlag/Vorkasse!"
new_c = [dict(kdnr=10139, gruppe="P", anrede="Frau", vorname="Kerstin", nachname="Vogt", name="Vogt, Kerstin",
              strasse="Tolkewitzer Straße", hnr="41", plz="01279", ort="Dresden", tel="+49 160 9988776", mobil=None,
              email=None, ziel=14, skonto=False, seit=dt.date(2026, 10, 20), bem="WP-Beratung", ansprech=""),
         dict(kdnr=10140, gruppe="P", anrede="Familie", vorname="Jörg", nachname="Bergmann", name="Bergmann, Familie",
              strasse="Meißner Straße", hnr="288", plz="01445", ort="Radebeul", tel="+49 351 8331200", mobil=None,
              email="bergmann.radebeul@gmail.com", ziel=14, skonto=False, seit=dt.date(2026, 10, 21), bem="Neubau", ansprech=""),
         dict(kdnr=10141, gruppe="G", anrede="Firma", vorname=None, nachname=None, name="Praxis für Ergotherapie Lange",
              strasse="Teplitzer Straße", hnr="12", plz="01219", ort="Dresden", tel="+49 351 4440011", mobil=None,
              email="info@ergo-lange.de", ziel=14, skonto=False, seit=dt.date(2026, 10, 22), bem="", ansprech="Fr. Lange")]
write_kundenstamm(L2 / "Kundenstamm.xlsx", L2_CUST + new_c, "25.10.2026", record=False)
L2_ORD = [dict(o) for o in ORDERS]
changed = []
for o in L2_ORD:
    if o["status"] in ("erledigt",) and len(changed) < 6 and o["invoices"] == []:
        o["status"] = "abgerechnet"
        changed.append((o["nr"], "erledigt->abgerechnet"))
for o in L2_ORD:
    if o["status"] in ("offen", "geplant") and len([x for x in changed if "geplant" in x[1] or "offen" in x[1]]) < 4:
        changed.append((o["nr"], f"{o['status']}->in Arbeit"))
        o["status"] = "in Arbeit"
seq26 = max(int(o["nr"].split("-")[2]) for o in ORDERS if o["nr"].startswith("A-26"))
new_orders = []
for k, (c, art) in enumerate([(BY_NR[10005], "Notdienst"), (BY_NR[10008], "Reparatur"), (KRUEGER, "Wartung"),
                              (BY_NR[10052], "Reparatur"), (MUELLER_G, "Reparatur"), (HEIKE_2, "Reparatur"),
                              (BY_NR[10040], "Wartung"), (BY_NR[10070], "Montage")], start=1):
    d = next_wd(dt.date(2026, 10, 15) + dt.timedelta(days=k))
    vfn, bl, big = ORDER_TYPES[art]
    o = dict(cust=c, date=d, art=art, beschr=R.choice(bl), wert=q2(vfn()), monteur=R.choice(["JK", "MR", "DH"]), status="offen",
             done=None, big=big, tag=None, objekt=(c["objekte"][0] if c["objekte"] else addr(c)), invoices=[], nr=f"A-26-{seq26 + k:03d}")
    new_orders.append(o)
write_auftragsbuch(L2 / "Auftragsbuch.xlsx", L2_ORD + new_orders, "25.10.2026")
TRUTH["lieferung2"] = {
    "Kundenstamm.xlsx": {"zeilen": len(L2_CUST) + 3, "entfernt": ["10119 (Dublette bereinigt)"],
                          "neu": ["10139 Vogt, Kerstin", "10140 Bergmann, Familie", "10141 Praxis für Ergotherapie Lange"],
                          "geaendert": {"10052": "Telefon -> +49 351 3109000", "10061": "Adresse -> Rudolf-Leonhard-Straße 7",
                                        "10023": "E-Mail -> rechnung@mueller-getraenke.de", "10045": "Bemerkung -> NUR gegen Abschlag/Vorkasse!"},
                          "hinweis": "neue Kunden 10139-10141 = Interessenten aus Termine-JSON (Vogt, Bergmann, Ergotherapie Lange)"},
    "Auftragsbuch.xlsx": {"neu": [o["nr"] for o in new_orders], "statuswechsel": changed,
                           "erwartet_auftraege_gesamt": len(ORDERS) + len(new_orders)}}

# ================================================================= kaputt
good = (OUT / "Kundenstamm.xlsx").read_bytes()
(KAPUTT / "Rechnungen_Oktober.xlsx").write_bytes(good[: len(good) // 2])
(KAPUTT / "Kasse_Export.csv").write_bytes(b"Datum;Beleg;Betrag;Text\r\n01.10.2026;K-0001;12,50;Dichtung\r\n" +
                                          bytes(R.randrange(256) for _ in range(3000)) + b"\x00\x00\x00\r\n02.10.2026;K-0002;8,90;Hanf\r\n")
(KAPUTT / "termine_defekt.json").write_text(json.dumps({"appointments": APPTS[:5]}, ensure_ascii=False, indent=2)[:1700], encoding="utf-8")
(KAPUTT / "leer.csv").write_bytes(b"")
(KAPUTT / "Scan_Lieferschein.pdf").write_bytes(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n1 0 obj\n<< /Type /Catalog /Pages 2 0 R" + bytes(R.randrange(256) for _ in range(800)))
(KAPUTT / "Archiv_2023.zip").write_bytes(b"PK\x03\x04" + bytes(R.randrange(256) for _ in range(1500)))
(KAPUTT / "Inventur.csv").write_bytes("Artikel;Bestand;Lagerort\n3100105;12;Regal A\n3100190;x;\n;;;;;\nnur eine Spalte\n".encode("utf-8")
                                      + "Hanf 100g;5;Regal B;Überzählige Spalte;noch eine\n".encode("cp1252") + b'"offenes Anf\xfchrungszeichen;3;\n')
TRUTH["kaputt"] = {"Rechnungen_Oktober.xlsx": "abgeschnittene xlsx (halbe Datei)", "Kasse_Export.csv": "2 gültige Zeilen + 3 KB Binärmüll mit NUL",
                   "termine_defekt.json": "abgeschnittenes JSON", "leer.csv": "0 Byte", "Scan_Lieferschein.pdf": "PDF-Kopf + Müll",
                   "Archiv_2023.zip": "kaputtes ZIP", "Inventur.csv": "ungleiche Spaltenzahl, gemischte Kodierung, offenes Anführungszeichen"}

# ================================================================= Wahrheit
inv_in_csv = [i for i in INVOICES if i["date"] <= EXPORT]
TRUTH["entitaeten"] = {
    "Kunde": {"zeilen_kundenstamm": len(CUST), "eindeutig_aktuell": len(TRUE_CUSTOMERS),
              "plus_nur_in_altdaten": 13, "plus_interessenten_termine": len(PROSPECTS),
              "eindeutig_gesamt": len(TRUE_CUSTOMERS) + 13 + len(PROSPECTS),
              "gruppen": dict(Counter(c["gruppe"] for c in CUST))},
    "Lieferant": {"anzahl": len(SUPPLIERS), "quelle": "Kundenstamm.xlsx/Lieferanten + Bank-Gegenkonten"},
    "Mitarbeiter": {"anzahl": len(EMP), "quelle": "Mitarbeiter_Personio.csv; Kürzel im Auftragsbuch/Materialverbrauch; Blattnamen Stundenzettel; resources im JSON"},
    "Auftrag": {"anzahl": len(ORDERS), "2025": sum(1 for o in ORDERS if o["date"].year == 2025),
                "2026": sum(1 for o in ORDERS if o["date"].year == 2026),
                "status": dict(Counter(o["status"] for o in ORDERS))},
    "Rechnung": {"anzahl_2025_2026": len(inv_in_csv), "anzahl_2024_alt": len(INV24),
                 "status": dict(Counter(i["status"] for i in inv_in_csv)),
                 "luecken": [f"RE-{y}-{n:04d}" for y, s in GAPS.items() for n in sorted(s)],
                 "storno": {"original": storno_orig["nr"], "storno": storno["nr"], "neu": neu["nr"]}},
    "Kontobewegung": {"anzahl": len(BANK), "davon_kundenzahlungen": sum(1 for b in BANK if b["kind"].startswith("kunde")),
                      "anfangssaldo": de(OPEN_BAL), "endsaldo": de(end_bal),
                      "summe_haben": de(sum(b["amount"] for b in BANK if b["amount"] > 0)),
                      "summe_soll": de(-sum(b["amount"] for b in BANK if b["amount"] < 0))},
    "Stundeneintrag": {"anzahl": TS_ROWS_TOTAL, "je_mitarbeiter": {SHEETNAME[k]: len(v) for k, v in TS.items()},
                       "arten": dict(TS_TRUTH), "summenzeilen_KW_nicht_mitzaehlen": True, "text_uhrzeiten": TS_TEXT_TIMES},
    "Artikel": {"anzahl_artikel_blatt": len(ART), "anzahl_pumpen_blatt": len(PUMPS), "gruppen": list(GROUPS),
                "preis_stile": dict(PRICE_STYLE)},
    "Materialbuchung": {"anzahl": len(MAT), "ohne_artikelnr": sum(1 for m in MAT if not m["art"]),
                        "pumpen_artikelnr": sum(1 for m in MAT if m["art"] in {p["nr"] for p in PUMPS})},
    "Termin": {"anzahl": len(APPTS), "mit_auftrag": sum(1 for a in APPT_TRUTH if a.get("auftrag") and a.get("ref_im_json")),
               "mit_vertrag": sum(1 for a in APPT_TRUTH if a.get("vertrag")), "interessenten": len(PROSPECTS)},
    "Wartungsvertrag": {"anzahl": len(CONTRACTS)},
    "Angebot": {"anzahl": 1, "nr": "AN-26-044", "kunde": "K10031"},
    "Dokumente": {"mails": 6, "notizen": 1, "pdf": 2},
}
TRUTH["dubletten"] = [
    {"entitaet": "K10014", "fall": "im Kundenstamm doppelt: 10014 'Hofmann, Klaus' + 10119 'Hoffmann, Klaus' (gleiche Adresse/Telefon)", "erwartet": "zusammenführen oder fragen"},
    {"entitaet": "K10023", "varianten": ["Müller GmbH & Co. KG", "Fa. Müller (Getränke)", "Müller GmbH & Co.KG", "Mueller GmbH", "MUELLER GMBH + CO. KG", "Müller Getränke"], "erwartet": "eine Entität"},
    {"entitaet": "K10005", "varianten": ["Wohnungsgenossenschaft Elbtal eG", "WG Elbtal eG", "Wohnungsgen. Elbtal", "WG Elbtal", "WOHNUNGSGENOSSENSCHAFT ELBTAL EG"]},
    {"entitaet": "K10008", "varianten": ["Hausverwaltung Schulze & Partner GbR", "HV Schulze & Partner", "Schulze & Partner", "HV Schulze u. Partner"]},
    {"entitaet": "K10031", "varianten": ["Krüger, Andreas (Dr.)", "Dr. Krüger", "A. Krueger", "Dr. Andreas Krüger", "DR. ANDREAS KRUEGER"]},
    {"entitaet": "K10040", "varianten": ["Zahnarztpraxis Dr. Lehmann", "Praxis Dr. med. dent. Lehmann", "ZA-Praxis Lehmann"]},
    {"entitaet": "Altdaten", "fall": "Kunden_alt.csv: 42 Zeilen = aktuelle Kunden (Name 'Vorname Nachname', alte Nummer D-xxxx); 3 davon mit alter Adresse", "umgezogen": sorted(moved)},
    {"entitaet": "Interessenten", "fall": "Lieferung 2 legt Vogt/Bergmann/Ergotherapie Lange an = dieselben wie in Termine-JSON"},
]
TRUTH["namensvettern"] = [
    {"verschieden": ["K10023 Müller GmbH & Co. KG (Getränke)", "K10045 Müller Bau GmbH", "K10067 Elektro Müller e.K.", "K10088 Heike Müller (Kesselsdorfer Str. 40)", "K10131 Heike Müller (Am Hang 3)"]},
    {"verschieden": ["K10088 Müller, Heike", "K10131 Müller, Heike"], "hinweis": "exakt gleicher Name, versch. Adresse/Tel/Mail; Bankname bei beiden 'HEIKE MUELLER'"},
    {"verschieden": ["K10061 Schmidt, Peter", "K10134 Schmidt, Petra", "Mitarbeiter Kevin Schmidt"]},
    {"verschieden": ["K10008 HV Schulze & Partner", "K10052 Schulze, Gerda"]},
    {"verschieden": ["K10014 Hofmann, Klaus", "K10070 Bäckerei Hoffmann", "Mitarbeiter Dennis Hoffmann"]},
    {"verschieden": ["Mitarbeiter Thomas Brenner", "Mitarbeiterin Sabine Brenner"]},
    {"verschieden": "alle Zufallskunden sind verschieden, auch bei gleichem Nachnamen",
     "gleiche_nachnamen": {k: v for k, v in Counter(c["nachname"] for c in CUST if c["nachname"]).items() if v > 1}},
]
TRUTH["beziehungen"] = [
    "Auftrag -> Kunde (Kd-Nr, wenn leer: nur Name-Variante + Objekt)",
    "Auftrag -> Mitarbeiter (Monteur-Kürzel, teils 'MR/KS' zwei Personen, teils Nachname)",
    "Auftrag -> Rechnung (Spalte 'Rechnung', ~75 %; sonst Auftragsnummer im Betreff der Rechnung)",
    "Rechnung -> Kunde (Kd-Nr)", "Rechnung -> Auftrag (A-xx-nnn im Betreff, nicht immer)",
    "Kontobewegung -> Rechnung (Nummer im Verwendungszweck, verschiedene Schreibweisen; Sammelzahlungen)",
    "Kontobewegung -> Lieferant / Mitarbeiter (Lohn, PNR)",
    "Stundeneintrag -> Mitarbeiter (nur über Blattname!)", "Stundeneintrag -> Auftrag (Auftrag-Spalte, ~82 %)",
    "Materialbuchung -> Auftrag", "Materialbuchung -> Artikel (Art.-Nr. bzw. Item No.)",
    "Termin -> Auftrag bzw. Wartungsvertrag (reference)", "Termin -> Mitarbeiter (resource_ids)", "Termin -> Kunde (eingebettet)",
    "Wartungsvertrag -> Kunde (Kd-Nr)", "Rechnung_2024 -> Kunde_alt (D-xxxx) -> aktueller Kunde",
    "Angebot AN-26-044 -> Kunde K10031, -> Auftrag (Beschreibung 'lt. Angebot AN-26-044')",
]
TRUTH["formeln"] = ["Rechnung: Brutto = Netto + USt; USt = Netto × 0,19", "Rechnung: Fällig = Belegdatum + Zahlungsziel (14/21/30 Tage je Kunde)",
                    "Stundenzettel: Std. = (Ende − Beginn) × 24 − Pause/60", "Preisliste: VK netto = ROUND(EK netto × 1,35; 2)",
                    "Rechnungs-PDF/Angebot: GP = Menge × EP"]
TRUTH["nummernkreise"] = {"Auftrag": f"A-26-{seq['26'] if False else seq[26]:03d} zuletzt (nächste A-26-{seq[26] + 1:03d})",
                          "Rechnung": f"RE-2026-{seqi[2026]:04d} zuletzt", "Kunde": "10138 zuletzt (Lücken 10017, 10093)",
                          "Wartungsvertrag": "WV-JJJJ-nnn", "Angebot": "AN-26-044"}
TRUTH["auffaellig"] = {
    "bezahlt_laut_bank_aber_offen_im_export": LAG,
    "teilzahlungen": PARTIAL,
    "unbezahlt_ueberfaellig_zum_15_10": [n for n in UNPAID if INV_BY_NR[n]["faellig"] < EXPORT],
    "unbezahlt_alle": UNPAID,
    "erledigt_aber_nicht_abgerechnet": [o["nr"] for o in ORDERS if o["status"] == "erledigt"],
    "kunde_zahlt_schleppend": "K10045 Müller Bau",
}
TRUTH["werte_fallen"] = [
    "Kundenstamm PLZ als Zahl (1067 statt 01067) – teils mit Zahlenformat 00000",
    "Kundenstamm Telefon teils als Zahl (3518494420), führende 0 fehlt",
    "Auftragsbuch 2026 Datum: Excel / 'dd.mm.yyyy' / 'd.m.yy' / 'dd/mm/yy' (= TT/MM/JJ!)",
    "Auftragsbuch Wert: Zahl / '1.250,00 €' / 'pauschal 139,-' / 'nach Aufwand' / 'ca. 800' / 'auf Anfrage'",
    "Auftragsbuch 'erledigt am' teils 'KW 38'", "Mitarbeiter Hire date MM/DD/YYYY (08/15/2012 = 15.08.2012)",
    "Bank: Betrag immer positiv, Vorzeichen in S/H-Spalte; absteigend sortiert",
    "Preisliste EAN 13-stellig als Zahl", "Artikelnummern 7-stellig als Zahl, Kundennummern 5-stellig als Zahl",
    "Kunde seit: Datum / '03/2019' / 2015 (Jahr als Zahl) / leer",
]
TRUTH["details"] = {
    "auftraege": {o["nr"]: {"kunde": o["cust"]["true_id"], "kdnr_im_buch": bool(o.get("typed_nr")), "name_im_buch": o.get("typed_name"),
                            "datum": o["date"].isoformat(), "datum_stil": DATE_STYLE.get(o["nr"]), "wert": str(o["wert"]),
                            "wert_stil": WERT_STYLE.get(o["nr"]), "status": o["status"], "monteur": o["monteur"],
                            "rechnungen": [i["nr"] for i in o["invoices"]], "rechnung_im_buch": bool(o.get("inv_in_book"))}
                   for o in ORDERS},
    "rechnungen": {i["nr"]: {"kunde": i["cust"]["true_id"], "kdnr": i["cust"]["kdnr"], "auftrag": i["order"]["nr"] if i["order"] else None,
                             "netto": str(i["netto"]), "brutto": str(i["brutto"]), "datum": i["date"].isoformat(), "status": i["status"],
                             "bezahlt": i["paid"].isoformat() if i["paid"] else None}
                   for i in inv_in_csv},
    "bankzahlungen_kunden": BANK_PAY_TRUTH,
    "termine": APPT_TRUTH,
    "kunden": {c["kdnr"]: {"true_id": c["true_id"], "name": c["name"], "plz": c["plz"], "tel": c["tel"], "tel_stil": TEL_STYLE.get(c["kdnr"]),
                           "gruppe": c["gruppe"]} for c in CUST},
    "altkunden": {o["nr"]: o["true_id"] for o in OLD},
    "mitarbeiter": {e["pnr"]: f"{e['vorname']} {e['nachname']} ({e['kz']}), Eintritt {e['hire'].isoformat()}" for e in EMP},
    "vertraege": {k["nr"]: k["cust"]["true_id"] for k in CONTRACTS},
}
(BASE / "wahrheit.json").write_text(json.dumps(TRUTH, ensure_ascii=False, indent=1, default=str), encoding="utf-8")

print("Kunden", len(CUST), "eindeutig", len(TRUE_CUSTOMERS), "| Aufträge", len(ORDERS), "| Rechnungen", len(inv_in_csv),
      "| Bank", len(BANK), "| Stunden", TS_ROWS_TOTAL, "| Artikel", len(ART), "+", len(PUMPS), "| Material", len(MAT),
      "| Termine", len(APPTS), "| Verträge", len(CONTRACTS), "| Alt-Kunden", len(OLD), "| Alt-Rechnungen", len(INV24))
print("Endsaldo", de(end_bal), "| LAG", LAG, "| PARTIAL", PARTIAL, "| UNPAID", len(UNPAID))
