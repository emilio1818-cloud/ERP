import io
import json
import unittest
from datetime import date, datetime

import openpyxl

from physarum.ingest import read_file
from physarum.lexicon import concept_of, header_concept, header_role, looks_like_person
from physarum.text import name_similarity, norm_key, norm_name, plural_de, singular_de
from physarum.values import iban_ok, parse_date, parse_number, profile_column


class TestValues(unittest.TestCase):
    def test_numbers_german_and_english(self):
        self.assertEqual(parse_number("1.234,56", "de"), 1234.56)
        self.assertEqual(parse_number("1,234.56", "en"), 1234.56)
        self.assertEqual(parse_number("12,5", "de"), 12.5)
        self.assertEqual(parse_number("-2.350,00", "de"), -2350.0)
        self.assertEqual(parse_number("123,45-", "de"), -123.45)
        self.assertEqual(parse_number("(12,00)", "de"), -12.0)
        self.assertEqual(parse_number("24,90 €", "de"), 24.9)
        self.assertIsNone(parse_number("auf Anfrage", "de"))

    def test_dates(self):
        self.assertEqual(parse_date("31.12.2024"), date(2024, 12, 31))
        self.assertEqual(parse_date("31.12.24"), date(2024, 12, 31))
        self.assertEqual(parse_date("2026-10-28"), date(2026, 10, 28))
        self.assertEqual(parse_date("2025-01-26 13:24:49 +0100"), datetime(2025, 1, 26, 13, 24, 49))
        self.assertEqual(parse_date("28. Oktober 2026"), date(2026, 10, 28))
        self.assertEqual(parse_date("03/04/2026", "mdy"), date(2026, 3, 4))
        self.assertIsNone(parse_date("RE-2026-0001"))

    def test_iban(self):
        self.assertTrue(iban_ok("DE89 3704 0044 0532 0130 00"))
        self.assertFalse(iban_ok("DE89 3704 0044 0532 0130 01"))

    def test_profile_detects_kinds_and_roles(self):
        p = profile_column("Netto", ["104,00", "1.155,50", "7,28"])
        self.assertEqual((p.kind, p.role), ("money", "money"))
        p = profile_column("MwSt", ["7,28", "10,85", "46,00", "12,14", "80,23"])
        self.assertEqual(p.kind, "money")  # tax amounts, not a rate
        p = profile_column("MwSt", ["7%", "19%", "7%"])
        self.assertEqual(p.kind, "percent")
        self.assertAlmostEqual(p.values[1], 0.19)
        p = profile_column("PLZ", [20457, "01067", 22765])
        self.assertEqual(p.kind, "zip")
        self.assertEqual(p.values[1], "01067")
        p = profile_column("Name", ["#1001", "#1002", "#1003", "#1001"])
        self.assertEqual(p.kind, "id")  # values beat the header (Shopify)
        p = profile_column("Zahlungsziel", ["30 Tage", "14 Tage", "sofort"])
        self.assertEqual((p.kind, p.values), ("duration", [30, 14, 0]))
        p = profile_column("E-Mail", ["a@b.de", None, "c@d.com"])
        self.assertEqual(p.kind, "email")

    def test_text_helpers(self):
        self.assertEqual(norm_name("Café Hafenblick GmbH & Co. KG"), norm_name("Cafe Hafenblick"))
        self.assertEqual(norm_name("Müller & Söhne"), norm_name("Mueller und Soehne"))
        self.assertEqual(norm_key("01001"), norm_key(1001.0))
        self.assertLess(name_similarity("bistro krabbenkutter", "buero krabbenkutter"), 0.8)
        self.assertGreater(name_similarity("hotel elbblick", "hotel elbblik"), 0.95)
        self.assertEqual(plural_de("Auftrag"), "Aufträge")
        self.assertEqual(plural_de("Wartungsvertrag"), "Wartungsverträge")
        self.assertEqual(singular_de("Rechnungen"), "Rechnung")
        self.assertTrue(looks_like_person("Hannah Becker"))
        self.assertFalse(looks_like_person("Hotel Landungsbrücke"))

    def test_lexicon(self):
        self.assertEqual(concept_of("Kundenliste_2023.xlsx"), "kunde")
        self.assertEqual(concept_of("Aufträge_Großkunden"), "auftrag")
        self.assertEqual(concept_of("customer_orders"), "auftrag")
        self.assertEqual(concept_of("Lieferantenrechnung"), "rechnung")
        self.assertEqual(header_concept("Kd-Nr."), "kunde")
        self.assertIsNone(header_concept("Kundengruppe"))
        self.assertIsNone(header_concept("Rechnungsdatum"))
        self.assertEqual(header_role("MwSt"), "percent")
        self.assertEqual(header_role("Fällig am"), "due")


class TestIngest(unittest.TestCase):
    def test_csv_cp1252_title_and_sum_rows(self):
        text = ("Rechnungsausgangsbuch 2026\n\nRechnungsnr;Datum;Betrag\nRE-1;01.02.2026;1.234,50\n"
                "RE-2;03.02.2026;99,00\nSumme;;1.333,50\n")
        d = read_file("buch.csv", text.encode("cp1252"))
        self.assertEqual(len(d.tables), 1)
        t = d.tables[0]
        self.assertEqual(t.columns, ["Rechnungsnr", "Datum", "Betrag"])
        self.assertEqual(len(t.rows), 2)
        self.assertEqual(t.meta["skipped_totals"], 1)

    def test_xlsx_sections_become_a_column(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Preisliste"])
        ws.append(["Art.-Nr.", "Bezeichnung", "Preis"])
        ws.append(["ESPRESSO"])
        ws.append(["E1", "Espresso 1 kg", 24.9])
        ws.append(["E2", "Espresso 250 g", 7.5])
        ws.append(["FILTER"])
        ws.append(["F1", "Filter 1 kg", 28.9])
        buf = io.BytesIO()
        wb.save(buf)
        t = read_file("preise.xlsx", buf.getvalue()).tables[0]
        self.assertEqual(t.columns[-1], "Gruppe")
        self.assertEqual([r[-1] for r in t.rows], ["ESPRESSO", "ESPRESSO", "FILTER"])
        self.assertIn("Preisliste", t.meta["title"])

    def test_key_value_blocks(self):
        text = ("Lieferanten\n\nFirma: A GmbH\nTelefon: 040 1\nNotiz: liefert\n  dienstags\n\n"
                "Firma: B KG\nTelefon: 040 2\n")
        t = read_file("lief.txt", text.encode()).tables[0]
        self.assertEqual(t.columns, ["Firma", "Telefon", "Notiz"])
        self.assertEqual(t.rows[0][2], "liefert dienstags")
        self.assertEqual(t.rows[1][0], "B KG")

    def test_json_nested_children(self):
        data = [{"id": 1, "customer": {"name": "A"}, "items": [{"sku": "x", "qty": 2}, {"sku": "y", "qty": 1}]},
                {"id": 2, "customer": {"name": "B"}, "items": [{"sku": "x", "qty": 5}]}]
        d = read_file("orders.json", json.dumps(data).encode())
        names = [t.name for t in d.tables]
        self.assertEqual(names, ["orders", "orders.items"])
        self.assertIn("customer.name", d.tables[0].columns)
        self.assertEqual(len(d.tables[1].rows), 3)

    def test_eml_with_attachment(self):
        from email.message import EmailMessage

        m = EmailMessage()
        m["From"] = "Kunde <k@firma.de>"
        m["To"] = "info@wir.de"
        m["Subject"] = "Bestellung RE-1"
        m.set_content("Bitte liefern.")
        m.add_attachment(b"a;b\n1;2\n3;4\n", maintype="text", subtype="csv", filename="liste.csv")
        d = read_file("mail.eml", bytes(m))
        self.assertEqual(d.docs[0].title, "Bestellung RE-1")
        self.assertEqual(d.attachments[0][0], "liste.csv")

    def test_broken_file_never_raises(self):
        d = read_file("kaputt.xlsx", b"not a zip")
        self.assertTrue(d.notes)


class PeriodsAndPlaceholders(unittest.TestCase):
    def test_periods(self):
        from datetime import date
        from physarum.values import event_date, parse_date, parse_period
        self.assertEqual(parse_period("10/2026"), (date(2026, 10, 1), date(2026, 10, 31)))
        self.assertEqual(parse_period("Okt. 2026"), (date(2026, 10, 1), date(2026, 10, 31)))
        self.assertEqual(parse_period("KW 44/2026"), (date(2026, 10, 26), date(2026, 11, 1)))
        self.assertEqual(parse_period("Q4/2026"), (date(2026, 10, 1), date(2026, 12, 31)))
        self.assertIsNone(parse_period("13/2026"))
        self.assertEqual(event_date("KW 44", date(2025, 10, 2)), date(2025, 10, 27))   # the year of the case
        self.assertIsNone(parse_date("01.01.0001"))
        self.assertIsNone(parse_date("31.12.9999"))
        self.assertIsNone(parse_date("31.02.2026"))


if __name__ == "__main__":
    unittest.main()
