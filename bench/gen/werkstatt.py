#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
gen_firm.py -- unabhaengige Testfirma fuer Physarum

FIRMA: Autoservice Lindner GmbH, freie Kfz-Meisterwerkstatt, Erwitter Strasse 112, 59557 Lippstadt.
6 Personen: Juergen Lindner (Kfz-Meister, GF), Petra Lindner (Buero, seit 01.06.2026),
Dennis Hartmann + Marco Ostermann (Mechatroniker), Luca Krueger (Azubi), Andrea Kemper (Aushilfe Reifen, ab Okt.).
Seit 01.06.2026 fuehrt Petra ein eigenes Excel-Rechnungsbuch; die Serviceannahme fuehrt ein Excel-Auftragsbuch.
Privatkunden zahlen meist bar oder mit girocard bei Abholung, Firmen/Versicherungen per Ueberweisung.
Bank: Geschaeftskonto Deutsche Bank (CSV-Export mit Vorspann/Fusszeile, Soll/Haben getrennt, cp1252, neueste zuerst).

Aufruf:  python3 gen_firm.py OUTDIR
Schreibt OUTDIR/lieferung1/..., OUTDIR/lieferung2/..., OUTDIR/truth.json, OUTDIR/ledger.json, OUTDIR/answer_key.txt
(ledger.json und answer_key.txt liegen AUSSERHALB der Lieferungen und duerfen Physarum nicht gezeigt werden.)

Dateien je Lieferung (lieferung2 = kompletter neuer Stand des Ordners, gleiche Dateinamen, neu exportiert):
  Rechnungsbuch 2026.xlsx        Petras Rechnungsausgangsbuch (Titelzeilen, Monats-Trennzeilen, Summenzeile, Blatt "Mahnungen")
  Auftragsbuch Werkstatt.xlsx    Werkstattauftraege ab 01.06.2026 (Status, fertig am, Rg.)
  Kunden_Export.csv              Kundenkartei aus der Werkstattsoftware (cp1252, ;, alle Felder in "")
  Umsaetze_Geschaeftskonto.csv   Bankexport (L1: 01.07.-14.09.2026, L2: 29.07.-27.10.2026 -> Ueberlappung!)
  Mails/*.eml                    Kunden-, Versicherungs-, Lieferanten- und eigene Mails (teils mit PDF)
  Rechnungen PDF/*.pdf           an Firmenkunden gemailte Rechnungen, Storno, Gutschriften, ein KVA, eine Lieferantenrechnung
  (nur L2) Zahlungsavis Sonnenschein 07.10.2026.pdf, ~$Rechnungsbuch 2026.xlsx (Excel-Sperrdatei)

ANSWER KEY (geplante Situationen; Nummern so, wie sie im Rechnungsbuch / Auftragsbuch stehen)
  -- Lieferung 1 (heute 15.09.2026; Daten bis 14.09.; Petra hat Ueberweisungen nur bis 14.08. im Buch abgehakt)
  K01 Teilzahlungen: Taxi Yilmaz Rg 31076 (1.845,20, 13.07.) zahlt 600 am 03.08. und 600 am 31.08. -> offen 645,20, ueberfaellig
      (keine Stundung vereinbart); L2: Rest 645,20 am 05.10. -> bezahlt
  K02 Versicherung kuerzt: Rg 31132 Karin Vogt (2.263,05, Unfall) - NordWest zahlt 2.150,00 (Kuerzung Verbringung 113,05, Mail)
      -> offen 113,05, ueberfaellig; L2: Kundin zahlt Rest am 22.09. -> bezahlt
  K03 Zahlung ohne Rechnungsnummer: Sabine Hoffmann zahlt 489,90 "Inspektion Fiesta LP-SH 512" fuer Rg 31149 -> bezahlt (nur Bank)
  K04 Sammelueberweisung + Skonto: Pflegedienst Sonnenschein zahlt Rg 31092 + 31099 in EINER Ueberweisung abzgl. je 2 % Skonto
      (Avis-Mail 30.07.) -> beide bezahlt
  K05 Sammelueberweisung ohne Skonto, verspaetet: Elektro Kaiser zahlt Rg 31068 + 31102 am 02.09. zusammen -> bezahlt (nur Bank)
  K06 Skonto gewaehrt: Sonnenschein zahlt (fast) immer innerhalb 10 Tagen mit 2 % -> bezahlt
  K07 Skonto NICHT gewaehrt: Elektro Kaiser zieht bei Rg 31038 2 % ab, obwohl "30 Tage netto" -> Rest 23,12 offen, ueberfaellig (L1+L2)
  K08 Gutschrift: Rg 31079 Baeckerei Steinhoff 1.190,00, Gutschrift 31087 (-142,80), Zahlung 1.047,20 -> erledigt;
      eine Gutschrift ist selbst kein offener Posten
  K09 Storno + Neuausstellung: Rg 31121 (980,00) Autovermietung Hellweg storniert durch 31123, neu 31124 (1.028,30), bezahlt 01.09.
      -> nichts offen
  K10 Dublette: Pflegedienst Sonnenschein GmbH (Kd 10412) und "Ambulanter Pflegedienst Sonnenschein" (Kd 10671, gleiche
      Adresse "Cappelstr. 41", gleiches Telefon anders geschrieben) -> sollte als ein Kunde erkannt werden
  K11 Gleicher Name, zwei Kunden: Thomas Becker Kd 10133 (Lippstadt, LP-TB 412) und Kd 10541 (Erwitte, SO-TB 77) -> NICHT
      zusammenfuehren. Rg 31107 (Lippstadt, 498,00) offen + ueberfaellig, bezahlt erst 12.10. (L2);
      Rg 31136 (Erwitte, 645,30) steht im L1-Buch faelschlich unter Kd 10133, bezahlt 21.08. (Mail "falsche Adresse")
  K12 Familienmitglied zahlt: Klaus Schroeder zahlt Rg 31125 seines Sohnes Leon (mit Rg-Nr.) -> bezahlt
  K13 Familienmitglied ohne Rg-Nr.: Irina Nowak zahlt 723,60 "Rechnung Auto LP PN 88" fuer Rg 31140 (Pawel Nowak) -> bezahlt
  K14 Geld ohne Rechnung: Jan Petersen (kein Kunde) 180,00 "Abschleppen B55" am 03.08. -> receipt (L1+L2; liegt in der
      Bank-Ueberlappung beider Exporte)
  K15 Doppelzahlung: Ute Brinkmann zahlt Rg 31155 per girocard bei Abholung UND ueberweist 312,40 am 03.09. "Rg 31155"
      -> Rechnung bezahlt, 312,40 = receipt (L1+L2; L2: Mail mit Bitte um Erstattung, noch nicht erstattet)
  K16 Anzahlung ohne Rechnung: Rainer Kuhlmann 500,00 "Anzahlung Motor Passat" am 02.09. (Auftrag 26-0655 "wartet auf Teile")
      -> receipt L1; L2: Schlussrechnung 31287 (2.380,00 "abzgl. Anzahlung 500,00") -> offen 1.880,00, nicht faellig, kein receipt
  K17 Vor Bankexport bezahlt: Juni-Ueberweisungen nur im Buch abgehakt ("bez. am", teils als Text "22.06.26") -> bezahlt
  K18 Juni-Rechnung unbezahlt: Rg 31034 Dieter Pohl (1.120,00), Zahlungserinnerung per Mail, Teilzahlung 400,00 am 04.09.
      (in der Bank-Ueberlappung) -> offen 720,00, ueberfaellig (L1+L2). Blatt "Mahnungen" zeigt in L1 veraltet 1.120,00
  K19 Ruecklastschrift: Rg 31113 Kevin Schulz 456,70 girocard/ELV am 31.07., Lastschrift geplatzt 12.08. (-459,70 inkl. 3,00
      Gebuehr) -> offen 456,70, ueberfaellig (L1+L2; Gebuehr ist keine Rechnungsforderung)
  K20 Falsche Rg-Nr. in Ueberweisung: Gerda Pieper zahlt 412,30 fuer Rg 31147, schreibt aber "Rechnung 31174" (Zahlendreher;
      31174 ist eine bar/EC bezahlte Rechnung einer anderen Kundin vom 04.09.) -> 31147 bezahlt
  K21 Buch nicht nachgetragen: Ueberweisungen nach dem 14.08. stehen nur in der Bank -> bezahlt
  K22 Teilkasko-Glasschaden: Rg 31178 Heike Brandt 833,00, Selbstbeteiligung 150,00 per girocard, Rest 683,00 von der
      Versicherung -> L1 offen 683,00, faellig 21.09. (nicht ueberfaellig); L2 bezahlt 24.09.
  K23 Fertig, nicht berechnet (L1): 26-0633 Unfall Monika Lange (KVA-PDF, "Rechnung erst nach Gutachten"; L2 Rg 31213, bezahlt),
      26-0659 Bremsen Frank Albers (L2 Rg 31209), 26-0677 Klimaservice Getraenke Brune (L2 Rg 31201),
      26-0681 Raederwechsel 4 Mietwagen Hellweg (auch in L2 noch nicht berechnet)
  K24 Fallen, die NICHT "nicht berechnet" sind: 26-0649 Kulanz-Nacharbeit (Rg. "Kulanz"), 26-0579 Werkstattwagen (Rg. "intern"),
      26-0626 KVA abgelehnt (storniert), 26-0655 / 26-0685 noch in Arbeit, 26-0672 hat Rg 31183 im Rechnungsbuch
      (Spalte "Auftr."), nur im L1-Auftragsbuch fehlt die Rg-Nr.
  K25 Lieferantenrechnung Teile-Express (PDF + Mail) und KVA-PDF sind keine Kundenrechnungen
  K26 girocard-Tagesgutschriften (PayWest) und Bareinzahlungen in der Bank sind KEINE Zahlungen ohne Rechnung
      (sie gehoeren zu bar/EC-Rechnungen, die im Buch schon als bezahlt stehen)
  K27 Barverkauf-Rechnungen ohne Kundennummer; Excel mit Titelzeilen, Monatszeilen, Summenzeile; Blatt "Mahnungen" mit
      teils veralteten Betraegen; Bank: Vorspann, Fusszeile "Kontostand", Soll/Haben-Spalten, neueste Buchung zuerst
  -- Lieferung 2 (heute 28.10.2026; Daten bis 27.10.; Petra hat bis 02.10. abgehakt)
  K28 Buchkorrektur Betrag: Rg 31172 Bernd Hagemann L1 595,00 -> L2 535,50 (Nachlass; Mail + korrigiertes PDF), bezahlt 535,50
      -> erledigt (L1: offen 595,00, nicht ueberfaellig)
  K29 Storno einer L1-offenen Rechnung: Rg 31191 Sven Kraemer (1.287,58) storniert durch 31217, neu 31218 an Westfalen Leasing
      GmbH (neuer Kunde 10695), faellig 22.10., unbezahlt -> nur 31218 offen + ueberfaellig
  K30 Buchkorrektur Kunde: Rg 31136 im Buch von Kd 10133 auf 10541 (Becker Erwitte) korrigiert
  K31 Gutschrift auf unbezahlte Rechnung: Rg 31257 Getraenke Brune 773,50 - Gutschrift 31266 (59,50) -> offen 714,00, ueberfaellig
  K32 Sammelzahlung nur "lt. Avis": Sonnenschein zahlt 31237 + 31246 (Kd 10671!) + 31249 mit je 2 % Skonto; Rechnungsnummern
      stehen nur im Avis-PDF (neue Datei + Mail) -> bezahlt
  K33 Bankexport L2 = 29.07.-27.10. (gleicher Dateiname): Ueberlappung 29.07.-14.09. nicht doppelt zaehlen; Juli-Zahlungen
      stehen nur noch im L1-Export, sind aber im L2-Buch abgehakt
  K34 Anzahlung Winterreifen Cem Arslan 250,00 (20.10., neuer Kunde 10708, Mail) -> receipt L2
  K35 Fertig, nicht berechnet (L2): 26-0681 (Hellweg, alt), 26-0745 Achsvermessung Steinhoff, 26-0778 Olga Wiebe;
      26-0791 / 26-0792 sind noch in Arbeit
  K36 Kaiser zahlt weiter spaet; neue Ueberweisungsrechnungen teils ueberfaellig; Oktober = Reifensaison (viele kleine bar/EC)
  K37 Excel-Sperrdatei "~$Rechnungsbuch 2026.xlsx" liegt im L2-Ordner
"""

import sys
import os
import json
import random
import shutil
import hashlib
import datetime as dt
from email.message import EmailMessage
from email.utils import format_datetime
from email import policy

import openpyxl
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas

SEED = 260915
R = random.Random(SEED)
D = dt.date
ONE = dt.timedelta(days=1)

CUT = {1: D(2026, 9, 14), 2: D(2026, 10, 27)}          # letzter Datentag je Lieferung
TODAY = {1: D(2026, 9, 15), 2: D(2026, 10, 28)}
BOOK_START = D(2026, 6, 1)
BANK_FROM = {1: D(2026, 7, 1), 2: D(2026, 7, 29)}
CHECK = {1: D(2026, 8, 14), 2: D(2026, 10, 2)}         # bis hierhin hat Petra Ueberweisungen im Buch abgehakt
HOLIDAYS = {D(2026, 6, 4), D(2026, 10, 3)}             # Fronleichnam (NRW), Tag der Deutschen Einheit
START_BALANCE = 2174832                                # Cent, Kontostand vor dem 01.07.2026
FIRM_IBAN_BLZ = "47270024"
FIRM_KONTO = 1185093


def days(n):
    return dt.timedelta(days=n)


def is_bank_day(d):
    return d.weekday() < 5 and d not in HOLIDAYS


def is_shop_day(d):
    return d.weekday() < 6 and d not in HOLIDAYS


def next_bank_day(d):
    while not is_bank_day(d):
        d += ONE
    return d


def prev_bank_day(d):
    while not is_bank_day(d):
        d -= ONE
    return d


def bank_day_after(d):
    d += ONE
    return next_bank_day(d)


def sub_shop_days(d, n):
    while n > 0:
        d -= ONE
        if is_shop_day(d):
            n -= 1
    return d


def vat_of(net):
    if net < 0:
        return -vat_of(-net)
    return (net * 19 + 50) // 100


def de_amt(c):
    s = "-" if c < 0 else ""
    c = abs(c)
    e, ct = divmod(c, 100)
    return s + f"{e:,}".replace(",", ".") + f",{ct:02d}"


def eur(c):
    return round(c / 100.0, 2)


def dde(d):
    return d.strftime("%d.%m.%Y")


def make_iban(blz, konto):
    bban = f"{blz}{konto:010d}"
    num = int(bban + "131400")
    return f"DE{98 - num % 97:02d}{bban}"


BANKS = [("41650001", "WELADED1LIP"), ("41660124", "GENODEM1LPS"), ("41450075", "WELADED1SOS"),
         ("41662557", "GENODEM1SGE"), ("50010517", "INGDDEFFXXX"), ("12030000", "BYLADEM1001"),
         ("44010046", "PBNKDEFFXXX"), ("47240047", "COBADEFFXXX"), ("41662465", "GENODEM1ERW")]


def new_iban():
    blz, bic = R.choice(BANKS)
    return make_iban(blz, R.randrange(10 ** 8, 10 ** 10)), bic


FIRM_IBAN = make_iban(FIRM_IBAN_BLZ, FIRM_KONTO * 100)

# ---------------------------------------------------------------------------------------------
# Kunden
# ---------------------------------------------------------------------------------------------
CUSTOMERS = {}


def cust(kd, **kw):
    c = dict(kd=kd, anrede="", name="", vorname="", zusatz="", strasse="", plz="", ort="", tel="", mobil="",
             mail="", zb="sofort", angelegt=None, bem="", vehicles=[], kind="privat", payer=None, iban=None,
             bic=None, book_name=None, job_name=None, profile="privat", special=True, ovr={})
    c.update(kw)
    if c["iban"] is None:
        c["iban"], c["bic"] = new_iban()
    if c["payer"] is None:
        c["payer"] = (c["vorname"] + " " + c["name"]).strip()
    if c["book_name"] is None:
        c["book_name"] = f"{c['name']}, {c['vorname']}" if c["vorname"] else c["name"]
    if c["job_name"] is None:
        c["job_name"] = f"{c['name']}, {c['vorname'][:2]}." if c["vorname"] else c["name"]
    CUSTOMERS[kd] = c
    return c


FLEET = {}
FLEET["sonn"] = cust(10412, anrede="Firma", name="Pflegedienst Sonnenschein GmbH", zusatz="Fuhrpark / Frau Overhoff",
                     strasse="Cappelstraße 41", plz="59555", ort="Lippstadt", tel="02941 77880",
                     mail="fuhrpark@pflegedienst-sonnenschein.de", zb="10 Tage 2% Skonto, 30 Tage netto",
                     kind="firma", profile="sonn",
                     payer="Ambulanter Pflegedienst Sonnenschein GmbH",
                     vehicles=[("LP-PS 101", "VW up!"), ("LP-PS 102", "VW up!"), ("LP-PS 103", "Renault Clio"),
                               ("LP-PS 104", "Renault Clio"), ("LP-PS 105", "VW Polo")],
                     book_name="Pflegedienst Sonnenschein", job_name="Sonnenschein Pflegedienst")
FLEET["sonn2"] = cust(10671, anrede="Firma", name="Ambulanter Pflegedienst Sonnenschein", zusatz="",
                      strasse="Cappelstr. 41", plz="59555", ort="Lippstadt", tel="02941/77880", mail="",
                      zb="sofort", angelegt=D(2026, 7, 7), kind="firma", profile="sonn",
                      bem="neu angelegt fuer LP-PS 106",
                      payer="Ambulanter Pflegedienst Sonnenschein GmbH", iban=None,
                      vehicles=[("LP-PS 106", "Skoda Fabia")],
                      book_name="Amb. Pflegedienst Sonnenschein", job_name="Pflegedienst Sonnenschein (Fabia)")
FLEET["sonn2"]["iban"], FLEET["sonn2"]["bic"] = FLEET["sonn"]["iban"], FLEET["sonn"]["bic"]
FLEET["kaiser"] = cust(10201, anrede="Firma", name="Elektro Kaiser GmbH & Co. KG", zusatz="Buchhaltung",
                       strasse="Im Hagen 12", plz="59558", ort="Lippstadt", tel="02941 66050",
                       mail="info@elektro-kaiser-lp.de", zb="30 Tage netto", kind="firma",
                       profile="kaiser", payer="ELEKTRO KAISER GMBH U. CO. KG",
                       vehicles=[("LP-EK 310", "VW Crafter"), ("LP-EK 311", "VW Caddy"),
                                 ("LP-EK 312", "Ford Transit Custom")],
                       book_name="Elektro Kaiser", job_name="Fa. Kaiser")
FLEET["yilmaz"] = cust(10307, anrede="Herr", name="Yilmaz", vorname="Mehmet", zusatz="Taxi Yilmaz",
                       strasse="Beckumer Straße 77", plz="59555", ort="Lippstadt", tel="02941 23023",
                       mobil="0171 4433291", mail="mehmet.yilmaz@taxi-yilmaz.de", zb="14 Tage netto",
                       kind="firma", profile="yilmaz", payer="TAXI YILMAZ INH. M. YILMAZ",
                       vehicles=[("LP-TY 300", "Mercedes E 220 d"), ("LP-TY 301", "Toyota Corolla Hybrid")],
                       book_name="Taxi Yilmaz", job_name="Taxi Yilmaz")
FLEET["hellweg"] = cust(10255, anrede="Firma", name="Autovermietung Hellweg GmbH", zusatz="",
                        strasse="Bahnhofstraße 9", plz="59597", ort="Erwitte", tel="02943 4870",
                        mail="buchhaltung@av-hellweg.de", zb="30 Tage netto", kind="firma", profile="hellweg", payer="Autovermietung Hellweg GmbH",
                        vehicles=[("SO-AH 21", "Opel Corsa"), ("SO-AH 22", "Opel Corsa"),
                                  ("SO-AH 23", "Skoda Octavia"), ("SO-AH 24", "VW Golf"), ("SO-AH 25", "Ford Transit")],
                        book_name="AV Hellweg", job_name="Hellweg Autovermietung")
FLEET["steinhoff"] = cust(10180, anrede="Firma", name="Bäckerei Steinhoff GmbH", strasse="Lange Straße 55",
                          plz="59555", ort="Lippstadt", tel="02941 3317", mail="info@baeckerei-steinhoff.de",
                          zb="14 Tage netto", kind="firma", profile="firma",
                          payer="BAECKEREI STEINHOFF GMBH",
                          vehicles=[("LP-BS 12", "Mercedes Sprinter 314"), ("LP-BS 14", "Fiat Ducato")],
                          book_name="Bäckerei Steinhoff", job_name="Steinhoff Bäckerei")
FLEET["brune"] = cust(10390, anrede="Firma", name="Getränke Brune e.K.", zusatz="Inh. Ralf Brune",
                      strasse="Westernkötter Straße 120", plz="59558", ort="Lippstadt", tel="02941 961144",
                      mail="ralf@getraenke-brune.de", zb="14 Tage netto", kind="firma",
                      profile="firma", payer="Getraenke Brune e.K.",
                      vehicles=[("LP-GB 77", "Mercedes Sprinter 316"), ("LP-GB 78", "VW Crafter")],
                      book_name="Getränke Brune", job_name="Brune Getränke")
FLEET["rinsche"] = cust(10466, anrede="Firma", name="Malerbetrieb Rinsche", zusatz="Inh. Uwe Rinsche",
                        strasse="Overhagener Straße 31", plz="59557", ort="Lippstadt", tel="02941 78812",
                        mail="malerbetrieb.rinsche@t-online.de", zb="14 Tage netto", kind="firma", profile="firma", payer="Uwe Rinsche",
                        vehicles=[("LP-MR 5", "Renault Trafic")], book_name="Maler Rinsche", job_name="Rinsche")
FLEET["rinsche"]["ovr"] = {2: dict(strasse="Gewerbestraße 4", plz="59558", bem="Umzug 09/2026")}

SP = {}
SP["becker_lp"] = cust(10133, anrede="Herr", name="Becker", vorname="Thomas", strasse="Mastholter Straße 18",
                       plz="59558", ort="Lippstadt", tel="02941 5512", mail="thomas.becker64@web.de",
                       vehicles=[("LP-TB 412", "Opel Astra K")], payer="Thomas Becker",
                       job_name="Becker, Th.")
SP["becker_erw"] = cust(10541, anrede="Herr", name="Becker", vorname="Thomas", strasse="Am Stadtgraben 3",
                        plz="59597", ort="Erwitte", tel="02943 98127", mobil="0160 99812345",
                        mail="tbecker.erwitte@gmail.com", vehicles=[("SO-TB 77", "VW Touran")], payer="BECKER, THOMAS", job_name="Becker, Th.")
SP["schroeder"] = cust(10577, anrede="Herr", name="Schröder", vorname="Leon", strasse="Weihenstraße 9",
                       plz="59555", ort="Lippstadt", mobil="0157 30044711", mail="leon.schroeder07@gmx.de",
                       vehicles=[("LP-LS 2007", "VW Polo")])
SP["nowak"] = cust(10350, anrede="Herr", name="Nowak", vorname="Pawel", strasse="Rixbecker Straße 140",
                   plz="59557", ort="Lippstadt", tel="02941 270311", vehicles=[("LP-PN 88", "Opel Insignia")])
SP["hoffmann"] = cust(10612, anrede="Frau", name="Hoffmann", vorname="Sabine", strasse="Südstraße 22",
                      plz="59557", ort="Lippstadt", mobil="0151 22870146", mail="sabine.hoffmann@t-online.de",
                      vehicles=[("LP-SH 512", "Ford Fiesta")], payer="SABINE HOFFMANN")
SP["pieper"] = cust(10220, anrede="Frau", name="Pieper", vorname="Gerda", strasse="Hospitalstraße 4",
                    plz="59555", ort="Lippstadt", tel="02941 4418", vehicles=[("LP-GP 33", "Toyota Yaris")], payer="Gerda Pieper")
SP["pohl"] = cust(10091, anrede="Herr", name="Pohl", vorname="Dieter", strasse="Am Markt 6", plz="59597",
                  ort="Erwitte", tel="02943 2290", vehicles=[("SO-DP 640", "Ford Focus")],
                  payer="Dieter Pohl")
SP["schulz"] = cust(10655, anrede="Herr", name="Schulz", vorname="Kevin", strasse="Dusternweg 17", plz="59557",
                    ort="Lippstadt", mobil="0176 84412093", mail="kevin.s1999@gmail.com", vehicles=[("LP-KS 1999", "BMW 318i")])
SP["brinkmann"] = cust(10498, anrede="Frau", name="Brinkmann", vorname="Ute", strasse="Bökenförder Straße 63",
                       plz="59557", ort="Lippstadt", tel="02941 15530", mail="ute.brinkmann@web.de",
                       vehicles=[("LP-UB 26", "Hyundai i20")], payer="Ute Brinkmann")
SP["kuhlmann"] = cust(10312, anrede="Herr", name="Kuhlmann", vorname="Rainer", strasse="Lipperoder Straße 88",
                      plz="59555", ort="Lippstadt", tel="02941 61904", vehicles=[("LP-RK 66", "VW Passat Variant B7")], payer="KUHLMANN, RAINER")
SP["vogt"] = cust(10470, anrede="Frau", name="Vogt", vorname="Karin", strasse="Barbarossastraße 12", plz="59555",
                  ort="Lippstadt", mobil="0162 5503218", mail="karin.vogt@gmx.de", vehicles=[("LP-KV 803", "Mazda 3")], payer="Karin Vogt")
SP["brandt"] = cust(10599, anrede="Frau", name="Brandt", vorname="Heike", strasse="Konrad-Adenauer-Ring 40",
                    plz="59557", ort="Lippstadt", tel="02941 920077", vehicles=[("LP-HB 415", "Skoda Octavia Combi")])
SP["lange"] = cust(10533, anrede="Frau", name="Lange", vorname="Monika", strasse="Bachstraße 15", plz="59590",
                   ort="Geseke", tel="02942 57731", vehicles=[("GT-ML 22", "Audi A3")])
SP["albers"] = cust(10421, anrede="Herr", name="Albers", vorname="Frank", strasse="Woldemei 7", plz="59555",
                    ort="Lippstadt", mobil="0171 9930562", vehicles=[("LP-FA 71", "Seat Leon")],
                    payer="Frank Albers")
SP["hagemann"] = cust(10377, anrede="Herr", name="Hagemann", vorname="Bernd", strasse="Mühlenstraße 3",
                      plz="59609", ort="Anröchte", tel="02947 3611", mail="b.hagemann@web.de",
                      vehicles=[("LP-BH 905", "VW Golf Variant")], payer="Bernd Hagemann")
SP["kraemer"] = cust(10640, anrede="Herr", name="Krämer", vorname="Sven", strasse="Klusetor 2", plz="59555",
                     ort="Lippstadt", mobil="0152 01188442", mail="sven.kraemer@outlook.de", vehicles=[("LP-SK 912", "BMW 520d Touring")], bem="Leasingfahrzeug")
SP["leasing"] = cust(10695, anrede="Firma", name="Westfalen Leasing GmbH", zusatz="Abt. Service/Wartung",
                     strasse="Warendorfer Straße 210", plz="48145", ort="Münster", tel="0251 4900120",
                     mail="service@westfalen-leasing.de", zb="30 Tage netto", angelegt=D(2026, 9, 22), kind="firma",
                     profile="firma", payer="Westfalen Leasing GmbH", vehicles=[], book_name="Westfalen Leasing",
                     job_name="Westfalen Leasing")
SP["arslan"] = cust(10708, anrede="Herr", name="Arslan", vorname="Cem", strasse="Beckumer Straße 12", plz="59555",
                    ort="Lippstadt", mobil="0176 22219980", mail="cem.arslan@gmail.com", angelegt=D(2026, 10, 19),
                    vehicles=[("LP-CA 344", "Seat Ateca")], payer="CEM ARSLAN")
SP["wiebe"] = cust(10690, anrede="Frau", name="Wiebe", vorname="Olga", strasse="Am Weinberg 5", plz="59556",
                   ort="Lippstadt", tel="02941 58204", vehicles=[("LP-OW 48", "Kia Ceed")])

SPECIAL_KD = set(CUSTOMERS.keys())

FIRST_M = ["Thomas", "Michael", "Andreas", "Stefan", "Frank", "Markus", "Klaus", "Bernd", "Uwe", "Ralf", "Jens",
           "Dirk", "Martin", "Christian", "Daniel", "Tobias", "Jan", "Florian", "Lukas", "Kevin", "Sebastian",
           "Matthias", "Holger", "Heinz", "Werner", "Günter", "Peter", "Wolfgang", "Ali", "Sergej", "Alexander",
           "Dimitri", "Norbert", "Hubert", "Jonas", "Niklas", "Timo", "Björn", "Carsten", "Hans-Jürgen"]
FIRST_F = ["Sabine", "Monika", "Andrea", "Claudia", "Susanne", "Birgit", "Heike", "Gabriele", "Anja", "Nicole",
           "Stefanie", "Julia", "Laura", "Lena", "Sarah", "Katharina", "Christina", "Melanie", "Tanja", "Elke",
           "Renate", "Irina", "Elena", "Ayse", "Fatma", "Marion", "Doris", "Jessica", "Vanessa", "Jutta"]
LAST = ["Müller", "Schmidt", "Schneider", "Fischer", "Weber", "Meyer", "Wagner", "Hoffmann", "Koch", "Richter",
        "Klein", "Wolf", "Neumann", "Schwarz", "Braun", "Zimmermann", "Krüger", "Hartmann", "Werner", "Krause",
        "Lehmann", "Köhler", "Kemper", "Schulte", "Hölscher", "Wiemann", "Tillmann", "Rüther", "Bröckelmann",
        "Westermann", "Overhoff", "Linnemann", "Peitz", "Henke", "Stratmann", "Böddeker", "Grothe", "Penner",
        "Friesen", "Demir", "Kaya", "Öztürk", "Kowalczyk", "Wiens", "Brockmeier", "Hense", "Sander", "Rose",
        "Beckmann", "Wessel", "Göke", "Jasper", "Kleine", "Mertens", "Spiekermann", "Ahlke", "Brökelmann"]
TOWNS = [("59555", "Lippstadt", ["Lange Straße", "Cappelstraße", "Beckumer Straße", "Lipperoder Straße",
                                 "Wiedenbrücker Straße", "Spielplatzstraße", "Poststraße"]),
         ("59556", "Lippstadt", ["Am Weinberg", "Quellenstraße", "Hellinghauser Straße", "Lipperoder Weg"]),
         ("59557", "Lippstadt", ["Erwitter Straße", "Bökenförder Straße", "Südstraße", "Overhagener Straße",
                                 "Rixbecker Straße", "Dusternweg"]),
         ("59558", "Lippstadt", ["Mastholter Straße", "Westernkötter Straße", "Im Hagen", "Eickelborner Straße"]),
         ("59597", "Erwitte", ["Hellweg", "Bahnhofstraße", "Soester Straße", "Am Stadtgraben", "Glasmerweg"]),
         ("59590", "Geseke", ["Bachstraße", "Lüdischer Weg", "Ostenfeldmark"]),
         ("59609", "Anröchte", ["Hauptstraße", "Kreuzstraße"]),
         ("59494", "Soest", ["Brüderstraße", "Ulricher Straße"])]
AREA = {"Lippstadt": "02941", "Erwitte": "02943", "Geseke": "02942", "Anröchte": "02947", "Soest": "02921"}
MODELS = ["VW Golf VII", "VW Golf VIII", "VW Polo", "Opel Astra", "Opel Corsa", "Ford Fiesta", "Ford Focus",
          "Skoda Octavia", "Skoda Fabia", "BMW 118i", "BMW 320d", "Mercedes A 180", "Mercedes C 200",
          "Toyota Yaris", "Toyota Auris", "Hyundai i30", "Kia Ceed", "Kia Sportage", "Seat Ibiza", "Seat Leon",
          "Audi A4 Avant", "Renault Clio", "Dacia Sandero", "Dacia Duster", "Fiat 500", "Mazda CX-5", "VW Tiguan",
          "VW Touran", "Opel Zafira", "Citroen C3", "Peugeot 208", "Nissan Qashqai", "Suzuki Swift", "VW T6"]
MAILDOM = ["web.de", "gmx.de", "t-online.de", "gmail.com", "outlook.de", "freenet.de"]
USED_PLATES = set()
for c in CUSTOMERS.values():
    for p, _ in c["vehicles"]:
        USED_PLATES.add(p)
USED_PLATES.add("LP-AL 112")


def make_plate(town, ini):
    while True:
        pre = "LP" if (town == "Lippstadt" and R.random() < 0.65) else R.choice(["SO", "SO", "LP", "PB", "GT", "WAR"])
        letters = ini if R.random() < 0.45 else "".join(R.choice("ABCDEFGHJKLMNPRSTUVWZ") for _ in range(R.choice([1, 2, 2])))
        num = str(R.randrange(1, 10 ** R.choice([2, 3, 3, 4])))
        p = f"{pre}-{letters} {num}"
        if p not in USED_PLATES:
            USED_PLATES.add(p)
            return p


def kd_date(kd):
    if kd <= 10600:
        base = D(2008, 3, 1)
        span = (D(2025, 12, 31) - base).days
        return base + days(int(span * (kd - 10001) / 599))
    if kd <= 10690:
        return D(2026, 1, 2) + days(int(235 * (kd - 10601) / 89))
    return D(2026, 9, 15) + days(int(42 * (kd - 10691) / 21))


PRIVATE_POOL = []
NEW_POOL = []
kds = sorted(R.sample([k for k in range(10002, 10690) if k not in SPECIAL_KD], 135))
kds_new = [k for k in range(10691, 10713) if k not in SPECIAL_KD]
for kd in kds + kds_new:
    female = R.random() < 0.42
    vn = R.choice(FIRST_F if female else FIRST_M)
    nn = R.choice(LAST)
    plz, ort, streets = R.choice(TOWNS[:5] * 3 + TOWNS[5:])
    st = f"{R.choice(streets)} {R.randrange(1, 140)}{R.choice(['', '', '', '', 'a', 'b'])}"
    tel = f"{AREA[ort]} {R.randrange(1000, 999999)}" if R.random() < 0.7 else ""
    mobil = f"01{R.choice(['51', '52', '57', '60', '62', '70', '71', '73', '76'])} {R.randrange(1000000, 99999999)}" if R.random() < 0.55 else ""
    mail = ""
    if R.random() < 0.55:
        base = R.choice([f"{vn}.{nn}", f"{vn[0]}.{nn}", f"{nn}{R.randrange(10, 99)}", f"{vn}{nn[:3]}"])
        base = base.lower().replace("ä", "ae").replace("ö", "oe").replace("ü", "ue").replace("ß", "ss").replace(" ", "")
        mail = f"{base}@{R.choice(MAILDOM)}"
    nveh = 1 if R.random() < 0.85 else 2
    vehs = [(make_plate(ort, (vn[0] + nn[0]).upper().replace("Ö", "O").replace("Ü", "U").replace("Ä", "A")),
             R.choice(MODELS)) for _ in range(nveh)]
    style = R.random()
    if style < 0.45:
        payer = f"{vn} {nn}"
    elif style < 0.70:
        payer = f"{nn.upper()}, {vn.upper()}"
    elif style < 0.85:
        payer = f"{nn} {vn}"
    else:
        partner = R.choice(FIRST_M if female else FIRST_F)
        payer = f"{vn} und {partner} {nn}" if R.random() < 0.5 else f"{partner} {nn}"
    c = cust(kd, anrede="Frau" if female else "Herr", name=nn, vorname=vn, strasse=st, plz=plz, ort=ort, tel=tel,
             mobil=mobil, mail=mail, angelegt=kd_date(kd), vehicles=vehs, payer=payer, special=False)
    (NEW_POOL if kd > 10690 else PRIVATE_POOL).append(c)

for c in CUSTOMERS.values():
    if c["angelegt"] is None:
        c["angelegt"] = kd_date(c["kd"])

INSURERS = {
    "nordwest": dict(name="NordWest Assekuranz AG", iban=make_iban("30020900", 4471100990), bic="CMCIDEDDXXX"),
    "soester": dict(name="Soester Kreis-Versicherung VVaG", iban=make_iban("41450075", 33012877), bic="WELADED1SOS"),
    "rhein": dict(name="Rhein-Ruhr Direkt Versicherung AG", iban=make_iban("37040044", 580112300), bic="COBADEFFXXX"),
}

# ---------------------------------------------------------------------------------------------
# Arbeiten
# ---------------------------------------------------------------------------------------------
WORK = {
    "insp": ("Inspektion nach Herstellervorgabe inkl. Öl und Filter", "Inspektion", 22000, 52000, 18, 20),
    "oel": ("Ölwechsel inkl. Ölfilter", "Ölwechsel", 7500, 14500, 9, 8),
    "hu": ("HU/AU-Vorführung inkl. Prüfgebühr", "HU/AU", 12500, 17500, 10, 8),
    "brv": ("Bremsscheiben und -beläge vorne erneuert", "Bremsen VA", 28000, 52000, 8, 10),
    "brh": ("Bremsbeläge hinten, Bremsflüssigkeit gewechselt", "Bremsen HA + Bremsfl.", 18000, 38000, 6, 6),
    "rad": ("Räderwechsel Sommer/Winter", "Räderwechsel", 2500, 4500, 9, 4),
    "radein": ("Räderwechsel inkl. Einlagerung", "Räderwechsel + Einlag.", 5500, 8500, 6, 2),
    "reifen": ("4 Reifen neu inkl. Montage, Wuchten, Altreifen", "4 Reifen neu", 32000, 74000, 5, 5),
    "klima": ("Klimaservice (Desinfektion, Kältemittel)", "Klimaservice", 8900, 14900, 5, 4),
    "zahn": ("Zahnriemen und Wasserpumpe erneuert", "Zahnriemen + WaPu", 52000, 98000, 3, 2),
    "kupp": ("Kupplung erneuert", "Kupplung", 85000, 145000, 2, 2),
    "batt": ("Batterie geprüft und erneuert", "Batterie", 14000, 26000, 4, 3),
    "diag": ("Fehlerspeicher ausgelesen, Diagnose", "Diagnose", 4900, 11900, 5, 3),
    "ausp": ("Mittelschalldämpfer/Endtopf erneuert", "Auspuff", 18000, 42000, 3, 1),
    "stossd": ("Stoßdämpfer vorne erneuert, Achsvermessung", "Stoßdämpfer + Vermessung", 38000, 69000, 2, 2),
    "licht": ("Lichtmaschine erneuert", "Lichtmaschine", 38000, 72000, 1, 1),
    "achs": ("Achsvermessung, Spur eingestellt", "Achsvermessung", 7900, 11900, 2, 2),
    "delle": ("Smart Repair Delle Tür", "Smart Repair", 12000, 28000, 2, 1),
    "wischer": ("Wischerblätter und Glühlampen erneuert", "Wischer/Lampen", 2900, 6900, 3, 2),
}
SAT_WORK = ["rad", "radein", "oel", "batt", "diag", "wischer"]
TIRE = {"rad", "radein", "reifen"}
MECHS = ["DH", "MO", "LK", "JL", "DH", "MO"]
THEKE = [("Motoröl 5W-30 5 l", 3990, 5490), ("AdBlue 10 l", 1490, 1990), ("Scheibenwischer Satz", 2290, 4290),
         ("Glühlampen H7 Satz", 1290, 2490), ("Frostschutz 5 l", 1590, 2490), ("Scheibenfrostschutz 5 l", 790, 1290),
         ("Batterieladegerät Ausleihe", 1000, 1500)]


def pick_work(d, private, saturday):
    if saturday:
        keys = SAT_WORK
    else:
        keys = list(WORK.keys())
    w = []
    for k in keys:
        wt = WORK[k][4] if private else WORK[k][5]
        if d.month == 10 and k in TIRE:
            wt *= 6
        if d.month in (6, 7, 8) and k == "klima":
            wt *= 2
        w.append(wt)
    first = R.choices(keys, weights=w)[0]
    items = [first]
    if not saturday and R.random() < 0.25:
        second = R.choices(keys, weights=w)[0]
        if second != first:
            items.append(second)
    return items


def price(items):
    return sum(R.randrange(WORK[k][2], WORK[k][3]) for k in items)


def arbeiten_text(items):
    return ", ".join(WORK[k][0] for k in items)


def leistung_text(items):
    return " + ".join(WORK[k][1] for k in items)


# ---------------------------------------------------------------------------------------------
# Auftraege, Rechnungen, Zahlungen
# ---------------------------------------------------------------------------------------------
JOBS, INVS, PAYS = [], [], []
KEY = {}       # Schluessel -> Rechnung/Auftrag (fuer answer key)


def mk_job(annahme, fertig, c, veh, arbeiten, special=None, key=None, kunde=None):
    j = dict(nr=None, key=key, annahme=annahme, fertig=fertig, cust=c, kd=c["kd"] if c else None,
             kunde=kunde or (c["job_name"] if c else ""), plate=veh[0], model=veh[1],
             km=R.randrange(18000, 240000), arbeiten=arbeiten, mech=R.choice(MECHS), special=special, invs=[],
             wait=R.choice(["in Arbeit", "wartet auf Teile", "in Arbeit", "Fzg. da"]), bem={1: "", 2: ""},
             rg_ovr={}, status_ovr={}, unbilled_status="fertig", sort=R.random())
    JOBS.append(j)
    if key:
        KEY[key] = j
    return j


def mk_inv(date, c, net, leistung, zahlart, job=None, due=None, kind="R", ref=None, key=None, plate=None,
           pdf=None, terms=None, kunde=None, t=None, auftr=None):
    vat = vat_of(net)
    inv = dict(nr=None, key=key, date=date, t=R.random() if t is None else t, cust=c, kd=c["kd"] if c else None,
               kunde=kunde or (c["book_name"] if c else "Barverkauf"), job=job,
               plate=plate if plate is not None else (job["plate"] if job else ""), leistung=leistung, net=net,
               vat=vat, gross=net + vat, zahlart=zahlart, due=due, kind=kind, ref=ref, pdf=pdf, terms=terms,
               ovr={1: {}, 2: {}}, bem={1: "", 2: ""}, cancelled_by=None, credits=[], allocs=[], skontos=[],
               auftr=(R.random() < 0.75) if auftr is None else auftr, items=None, first_dl=1 if date <= CUT[1] else 2)
    if job is not None:
        job["invs"].append(inv)
    INVS.append(inv)
    if key:
        KEY[key] = inv
    if kind == "G" and ref is not None:
        ref["credits"].append(inv)
    if kind == "S" and ref is not None:
        ref["cancelled_by"] = inv
    return inv


def mk_pay(date, amount, channel, payer="", iban="", bic="", tpl="", refs=(), allocs=(), skontos=(), kind="kunde",
           cref="", umsatzart="SEPA-Gutschrift"):
    p = dict(date=date, amount=amount, channel=channel, payer=payer, iban=iban, bic=bic, tpl=tpl, refs=list(refs),
             text="", allocs=[], skontos=[], kind=kind, cref=cref, umsatzart=umsatzart)
    for a in allocs:
        inv, amt = a[0], a[1]
        eff = a[2] if len(a) > 2 else date
        p["allocs"].append((inv, amt, eff))
        inv["allocs"].append((p, amt, eff))
    for inv, amt in skontos:
        p["skontos"].append((inv, amt))
        inv["skontos"].append((p, amt))
    PAYS.append(p)
    return p


def due_for(c, d):
    if c is None:
        return None
    z = c["zb"]
    if "30 Tage" in z:
        return d + days(30)
    return d + days(14)


TERMS = {"sonn": "Zahlbar innerhalb von 10 Tagen abzgl. 2 % Skonto oder innerhalb von 30 Tagen netto.",
         "net30": "Zahlbar innerhalb von 30 Tagen ohne Abzug bis {due}.",
         "net14": "Zahlbar innerhalb von 14 Tagen ohne Abzug bis {due}.",
         "bar": "Betrag bar erhalten. Vielen Dank!",
         "ec": "Bezahlt mit girocard. Vielen Dank!",
         "vers": "Die Rechnung wird gemäß Abtretungserklärung direkt mit der Versicherung abgerechnet."}

PRIV_TPL_NUM = ["Rechnung {0}", "RG-NR. {0}", "Re.Nr. {0} vom @D", "{0}", "Rechnung {0} @P", "Rg {0} @N",
                "Rechnungsnummer {0} Kundennummer @K", "Reparatur @M Rg. {0}", "RE {0}", "Rechnungs-Nr. {0}",
                "RECHNUNG NR {0} AUTOSERVICE LINDNER", "Rg.Nr.{0}"]
PRIV_TPL_NONUM = ["Rechnung vom @D @P", "Reparatur @M @P", "@P Werkstatt", "Autoservice Lindner Rechnung vom @D"]


def fill(tpl, inv):
    c = inv["cust"]
    return (tpl.replace("@D", dde(inv["date"])).replace("@P", inv["plate"]).replace("@N", c["name"] if c else "")
            .replace("@K", str(inv["kd"] or "")).replace("@M", (inv["job"]["model"] if inv["job"] else "")))


def pay_bar_ec(inv, how, amount=None, d=None):
    amt = inv["gross"] if amount is None else amount
    return mk_pay(d or inv["date"], amt, "ec" if how == "EC" else "bar", allocs=[(inv, amt)], kind="kasse")


def private_transfer(inv, delay=None):
    c = inv["cust"]
    if delay is None:
        r = R.random()
        if r < 0.70:
            delay = R.randint(2, 13)
        elif r < 0.88:
            delay = R.randint(14, 27)
        elif r < 0.97:
            delay = R.randint(28, 50)
        else:
            delay = R.randint(60, 90)
    d = next_bank_day(inv["date"] + days(delay))
    if R.random() < 0.9:
        tpl = fill(R.choice(PRIV_TPL_NUM), inv)
        refs = [inv]
    else:
        tpl = fill(R.choice(PRIV_TPL_NONUM), inv)
        refs = []
    return mk_pay(d, inv["gross"], "bank", c["payer"], c["iban"], c["bic"], tpl, refs, allocs=[(inv, inv["gross"])])


def fleet_payment(inv):
    c = inv["cust"]
    prof = c["profile"]
    g = inv["gross"]
    if prof == "sonn":
        if R.random() < 0.85:
            d = prev_bank_day(inv["date"] + days(R.randint(5, 9)))
            if d <= inv["date"]:
                d = bank_day_after(inv["date"])
            sk = (g * 2 + 50) // 100
            tpl = R.choice(["RE {0} KD 10412 ABZGL. 2% SKONTO", "Rechnung {0} @P abzgl. Skonto",
                            "RG {0} SKONTO 2 PROZENT"])
            return mk_pay(d, g - sk, "bank", c["payer"], c["iban"], c["bic"], fill(tpl, inv), [inv],
                          allocs=[(inv, g - sk)], skontos=[(inv, sk)], cref=f"SONN-{inv['date']:%y%m}-{R.randrange(100, 999)}")
        d = next_bank_day(inv["date"] + days(R.randint(25, 30)))
        return mk_pay(d, g, "bank", c["payer"], c["iban"], c["bic"], fill("RE {0} KD 10412", inv), [inv],
                      allocs=[(inv, g)])
    if prof == "kaiser":
        d = next_bank_day(inv["date"] + days(R.randint(35, 70)))
        tpl = R.choice(["RG {0} @P", "Rechnung {0} vom @D", "RE-NR {0}"])
        return mk_pay(d, g, "bank", c["payer"], c["iban"], c["bic"], fill(tpl, inv), [inv], allocs=[(inv, g)])
    if prof == "hellweg":
        d = next_bank_day(inv["date"] + days(R.randint(18, 29)))
        return mk_pay(d, g, "bank", c["payer"], c["iban"], c["bic"], fill("Re-Nr. {0} / Fzg @P", inv), [inv],
                      allocs=[(inv, g)], cref=f"AVH{R.randrange(10000, 99999)}")
    if prof == "yilmaz":
        d = next_bank_day(inv["date"] + days(R.randint(8, 20)))
        return mk_pay(d, g, "bank", c["payer"], c["iban"], c["bic"], fill("Rechnung {0} Taxi @P", inv), [inv],
                      allocs=[(inv, g)])
    d = next_bank_day(inv["date"] + days(R.randint(6, 20)))
    return mk_pay(d, g, "bank", c["payer"], c["iban"], c["bic"], fill(R.choice(["Rechnung {0}", "RG {0} KD @K"]), inv),
                  [inv], allocs=[(inv, g)])


# ---------- geplante Faelle (vor den Zufallsauftraegen anlegen, Nummern ergeben sich aus Datum/Uhrzeit) ----------
def planted():
    S = SP
    F = FLEET
    # K18 Pohl
    j = mk_job(D(2026, 6, 16), D(2026, 6, 18), S["pohl"], S["pohl"]["vehicles"][0], "Kupplung erneuert, Ausrücklager", key="J_POHL")
    inv = mk_inv(D(2026, 6, 18), S["pohl"], 94118, "Kupplung", "Überw.", job=j, due=D(2026, 7, 2), key="POHL", terms="net14")
    mk_pay(D(2026, 9, 4), 40000, "bank", "Dieter Pohl", S["pohl"]["iban"], S["pohl"]["bic"], "Teilzahlung Rg {0} Pohl", [inv],
           allocs=[(inv, 40000)])
    # K07 Kaiser unberechtigter Skonto
    j = mk_job(D(2026, 6, 19), D(2026, 6, 22), F["kaiser"], F["kaiser"]["vehicles"][0],
               "Bremsen VA+HA erneuert, Bremsflüssigkeit", key="J_KSK")
    inv = mk_inv(D(2026, 6, 22), F["kaiser"], 97161, "Bremsen VA+HA + Bremsfl.", "Überw.", job=j, due=D(2026, 7, 22),
                 key="KAISER_SK", pdf="Rechnung {nr} Kaiser.pdf", terms="net30")
    mk_pay(D(2026, 7, 20), 113310, "bank", F["kaiser"]["payer"], F["kaiser"]["iban"], F["kaiser"]["bic"],
           "RG {0} ABZGL 2% SKONTO", [inv], allocs=[(inv, 113310)])
    # K01 Yilmaz Raten
    j = mk_job(D(2026, 7, 10), D(2026, 7, 13), F["yilmaz"], F["yilmaz"]["vehicles"][0],
               "Zahnriemen + Wasserpumpe erneuert, Bremsen vorne", key="J_YIL")
    inv = mk_inv(D(2026, 7, 13), F["yilmaz"], 155059, "Zahnriemen + WaPu, Bremsen VA", "Überw.", job=j,
                 due=D(2026, 7, 27), key="YILMAZ", pdf="RE_{nr}.pdf", terms="net14")
    yc = F["yilmaz"]
    mk_pay(D(2026, 8, 3), 60000, "bank", yc["payer"], yc["iban"], yc["bic"], "Rate 1 Rg {0}", [inv], allocs=[(inv, 60000)])
    mk_pay(D(2026, 8, 31), 60000, "bank", yc["payer"], yc["iban"], yc["bic"], "Rg {0} 2. Rate", [inv], allocs=[(inv, 60000)])
    mk_pay(D(2026, 10, 5), 64520, "bank", yc["payer"], yc["iban"], yc["bic"], "Rest Rechnung {0}", [inv], allocs=[(inv, 64520)])
    # K04 Sonnenschein Sammelzahlung mit Skonto
    j1 = mk_job(D(2026, 7, 21), D(2026, 7, 21), F["sonn"], F["sonn"]["vehicles"][2], "Inspektion, Bremsbeläge vorne", key="J_SC1")
    a = mk_inv(D(2026, 7, 21), F["sonn"], 41000, "Inspektion + Bremsen VA", "Überw.", job=j1, due=D(2026, 8, 20), key="SONN_A", terms="sonn")
    j2 = mk_job(D(2026, 7, 23), D(2026, 7, 23), F["sonn2"], F["sonn2"]["vehicles"][0], "Ölwechsel, HU/AU", key="J_SC2")
    b = mk_inv(D(2026, 7, 23), F["sonn2"], 19900, "Ölwechsel + HU/AU", "Überw.", job=j2, due=D(2026, 8, 22), key="SONN_B", terms="sonn")
    sa, sb = (a["gross"] * 2 + 50) // 100, (b["gross"] * 2 + 50) // 100
    sc = F["sonn"]
    mk_pay(D(2026, 7, 30), a["gross"] - sa + b["gross"] - sb, "bank", sc["payer"], sc["iban"], sc["bic"],
           "RE {0} UND {1} ABZGL. 2% SKONTO KD 10412", [a, b], allocs=[(a, a["gross"] - sa), (b, b["gross"] - sb)],
           skontos=[(a, sa), (b, sb)], cref="SONN-2607-AVIS")
    # K05 Kaiser Sammelzahlung verspaetet
    j1 = mk_job(D(2026, 7, 7), D(2026, 7, 8), F["kaiser"], F["kaiser"]["vehicles"][1], "Inspektion, Klimaservice", key="J_KC1")
    a = mk_inv(D(2026, 7, 8), F["kaiser"], 38500, "Inspektion + Klimaservice", "Überw.", job=j1, due=D(2026, 8, 7),
               key="KAISER_A", pdf="Rechnung {nr} Kaiser.pdf", terms="net30")
    j2 = mk_job(D(2026, 7, 22), D(2026, 7, 24), F["kaiser"], F["kaiser"]["vehicles"][2], "Kupplung geprüft, Stoßdämpfer vorne", key="J_KC2")
    b = mk_inv(D(2026, 7, 24), F["kaiser"], 62650, "Stoßdämpfer VA + Vermessung", "Überw.", job=j2, due=D(2026, 8, 23),
               key="KAISER_B", pdf="Rechnung {nr} Kaiser.pdf", terms="net30")
    kc = F["kaiser"]
    mk_pay(D(2026, 9, 2), a["gross"] + b["gross"], "bank", kc["payer"], kc["iban"], kc["bic"], "RG {0} + {1}", [a, b],
           allocs=[(a, a["gross"]), (b, b["gross"])])
    # K08 Steinhoff Gutschrift
    j = mk_job(D(2026, 7, 14), D(2026, 7, 15), F["steinhoff"], F["steinhoff"]["vehicles"][0],
               "Bremsscheiben + Beläge VA und HA erneuert", key="J_STG")
    a = mk_inv(D(2026, 7, 15), F["steinhoff"], 100000, "Bremsen VA + HA", "Überw.", job=j, due=D(2026, 7, 29),
               key="STEIN_R", pdf="RE_{nr}.pdf", terms="net14")
    g = mk_inv(D(2026, 7, 20), F["steinhoff"], -12000, "Gutschrift: Bremsscheiben HA doppelt berechnet", "", kind="G",
               ref=a, key="STEIN_G", pdf="Gutschrift_{nr}.pdf", plate=F["steinhoff"]["vehicles"][0][0], auftr=False)
    stc = F["steinhoff"]
    mk_pay(D(2026, 7, 29), a["gross"] + g["gross"], "bank", stc["payer"], stc["iban"], stc["bic"],
           "Rechnung {0} abzgl. Gutschrift {1}", [a, g], allocs=[(a, a["gross"] + g["gross"])])
    # K09 Hellweg Storno + Neuausstellung
    j = mk_job(D(2026, 8, 3), D(2026, 8, 4), F["hellweg"], F["hellweg"]["vehicles"][2], "Inspektion, Bremsen hinten, Ersatzwagen 2 Tage", key="J_HST")
    a = mk_inv(D(2026, 8, 4), F["hellweg"], 82353, "Inspektion + Bremsen HA", "Überw.", job=j, due=D(2026, 9, 3),
               key="HELL_A", pdf="RE_{nr}.pdf", terms="net30", t=0.2)
    s = mk_inv(D(2026, 8, 5), F["hellweg"], -82353, "Storno zu Rg. {0} (Ersatzwagen vergessen)", "", kind="S", ref=a,
               key="HELL_S", pdf="Stornorechnung_{nr}.pdf", plate=a["plate"], t=0.30, auftr=False)
    n = mk_inv(D(2026, 8, 5), F["hellweg"], 86412, "Inspektion + Bremsen HA + Ersatzwagen", "Überw.", job=j,
               due=D(2026, 9, 4), key="HELL_N", pdf="RE_{nr}.pdf", terms="net30", t=0.31)
    hc = F["hellweg"]
    mk_pay(D(2026, 9, 1), n["gross"], "bank", hc["payer"], hc["iban"], hc["bic"], "Re-Nr. {0} / Fzg SO-AH 23", [n],
           allocs=[(n, n["gross"])], cref="AVH55102")
    # K11 Becker Lippstadt (offen bis L2) und Becker Erwitte (falsche Kd-Nr. im L1-Buch)
    bl, be = S["becker_lp"], S["becker_erw"]
    j = mk_job(D(2026, 7, 27), D(2026, 7, 27), bl, bl["vehicles"][0], "Bremsscheiben und -beläge vorne, Bremsflüssigkeit", key="J_BLP")
    a = mk_inv(D(2026, 7, 27), bl, 41849, "Bremsen VA + Bremsfl.", "Überw.", job=j, due=D(2026, 8, 10), key="BECKER_LP", terms="net14")
    mk_pay(D(2026, 10, 12), a["gross"], "bank", bl["payer"], bl["iban"], bl["bic"], "RG {0} Becker", [a], allocs=[(a, a["gross"])])
    j = mk_job(D(2026, 8, 11), D(2026, 8, 12), be, be["vehicles"][0], "Inspektion nach Herstellervorgabe, Klimaservice", key="J_BER")
    b = mk_inv(D(2026, 8, 12), be, 54227, "Inspektion + Klimaservice", "Überw.", job=j, due=D(2026, 8, 26), key="BECKER_ERW", terms="net14")
    b["ovr"][1] = dict(kd=10133, kunde="Becker, Thomas")
    b["ovr"][2] = dict(kunde="Becker, Thomas (Erwitte)")
    mk_pay(D(2026, 8, 21), b["gross"], "bank", be["payer"], be["iban"], be["bic"], "Rechnung {0}", [b], allocs=[(b, b["gross"])])
    # K12 Schroeder (Vater zahlt)
    c = S["schroeder"]
    j = mk_job(D(2026, 8, 6), D(2026, 8, 6), c, c["vehicles"][0], "Inspektion, Bremsbeläge hinten", key="J_SCHR")
    a = mk_inv(D(2026, 8, 6), c, 32731, "Inspektion + Bremsen HA", "Überw.", job=j, due=D(2026, 8, 20), key="SCHROEDER", terms="net14")
    ib, bic = new_iban()
    mk_pay(D(2026, 8, 18), a["gross"], "bank", "Klaus Schröder", ib, bic, "Rg. {0} Leon Schröder Polo", [a], allocs=[(a, a["gross"])])
    # K13 Nowak (Ehefrau, ohne Rg-Nr.)
    c = S["nowak"]
    j = mk_job(D(2026, 8, 14), D(2026, 8, 17), c, c["vehicles"][0], "Stoßdämpfer vorne erneuert, Achsvermessung", key="J_NOW")
    a = mk_inv(D(2026, 8, 17), c, 60807, "Stoßdämpfer + Vermessung", "Überw.", job=j, due=D(2026, 8, 31), key="NOWAK", terms="net14")
    ib, bic = new_iban()
    mk_pay(D(2026, 8, 26), a["gross"], "bank", "Irina Nowak", ib, bic, "Rechnung Auto LP PN 88", [], allocs=[(a, a["gross"])])
    # K03 Hoffmann ohne Rg-Nr.
    c = S["hoffmann"]
    j = mk_job(D(2026, 8, 20), D(2026, 8, 20), c, c["vehicles"][0], "Inspektion nach Herstellervorgabe", key="J_HOF")
    a = mk_inv(D(2026, 8, 20), c, 41168, "Inspektion", "Überw.", job=j, due=D(2026, 9, 3), key="HOFFMANN", terms="net14")
    mk_pay(D(2026, 9, 1), a["gross"], "bank", c["payer"], c["iban"], c["bic"], "Inspektion Fiesta LP-SH 512", [], allocs=[(a, a["gross"])])
    # K14 Petersen (Geld ohne Rechnung)
    ib, bic = new_iban()
    mk_pay(D(2026, 8, 3), 18000, "bank", "Jan Petersen", ib, bic, "Abschleppen B55 Samstag 01.08.", [], kind="ohne")
    # K15 Brinkmann Doppelzahlung
    c = S["brinkmann"]
    j = mk_job(D(2026, 8, 24), D(2026, 8, 24), c, c["vehicles"][0], "HU/AU-Vorführung, Bremsbeläge vorne", key="J_BRI")
    a = mk_inv(D(2026, 8, 24), c, 26252, "HU/AU + Bremsbeläge VA", "EC", job=j, key="BRINKMANN", terms="ec")
    pay_bar_ec(a, "EC")
    mk_pay(D(2026, 9, 3), a["gross"], "bank", c["payer"], c["iban"], c["bic"], "Rg {0}", [a], kind="doppelt")
    # K16 Kuhlmann Anzahlung -> Schlussrechnung L2
    c = S["kuhlmann"]
    j = mk_job(D(2026, 8, 31), D(2026, 10, 15), c, c["vehicles"][0], "Motorinstandsetzung: Steuerkette, Zylinderkopf planen",
               key="J_KUH")
    j["wait"] = "wartet auf Teile"
    a = mk_inv(D(2026, 10, 16), c, 200000, "Motorinstandsetzung (Steuerkette, ZK)", "Überw.", job=j, due=D(2026, 10, 30),
               key="KUHLMANN", terms="net14")
    mk_pay(D(2026, 9, 2), 50000, "bank", c["payer"], c["iban"], c["bic"], "Anzahlung Motor Passat LP-RK 66", [],
           allocs=[(a, 50000, D(2026, 10, 16))], kind="anzahlung")
    # K20 Pieper falsche Rg-Nr. (Zahlendreher wird nach Nummerierung eingesetzt)
    c = S["pieper"]
    j = mk_job(D(2026, 8, 18), D(2026, 8, 19), c, c["vehicles"][0], "Bremsscheiben und -beläge vorne erneuert", key="J_PIE")
    a = mk_inv(D(2026, 8, 19), c, 34647, "Bremsen VA", "Überw.", job=j, due=D(2026, 9, 2), key="PIEPER", terms="net14")
    mk_pay(D(2026, 8, 28), a["gross"], "bank", c["payer"], c["iban"], c["bic"], "Rechnung @TYPO", [], allocs=[(a, a["gross"])], kind="typo")
    # K19 Schulz Ruecklastschrift
    c = S["schulz"]
    j = mk_job(D(2026, 7, 31), D(2026, 7, 31), c, c["vehicles"][0], "4 Reifen neu inkl. Montage, Wuchten", key="J_SCHU")
    a = mk_inv(D(2026, 7, 31), c, 38378, "4 Reifen neu", "EC (ELV)", job=j, key="SCHULZ", terms="ec")
    a["ovr"][1] = dict(due=D(2026, 8, 26))
    a["ovr"][2] = dict(due=D(2026, 8, 26))
    pay_bar_ec(a, "EC")
    mk_pay(D(2026, 8, 12), -a["gross"], "bounce", c["payer"], c["iban"], c["bic"],
           "RUECKLASTSCHRIFT ELV TID 65781234 VOM 31.07.2026 KONTO NICHT GEDECKT ENTGELT 3,00 EUR",
           allocs=[(a, -a["gross"])], kind="ruecklastschrift", umsatzart="Rücklastschrift")
    # K22 Brandt Glasschaden
    c = S["brandt"]
    j = mk_job(D(2026, 9, 7), D(2026, 9, 7), c, c["vehicles"][0], "Frontscheibe erneuert, Kamera kalibriert (Teilkasko)", key="J_GLAS")
    a = mk_inv(D(2026, 9, 7), c, 70000, "Frontscheibe + Kalibrierung (Teilkasko)", "EC/Vers.", job=j, due=D(2026, 9, 21),
               key="GLAS", terms="vers")
    pay_bar_ec(a, "EC", amount=15000)
    ins = INSURERS["soester"]
    mk_pay(D(2026, 9, 24), a["gross"] - 15000, "bank", ins["name"], ins["iban"], ins["bic"],
           "Glasschaden Vers.-Nr. 44-1187266 Rg {0} abzgl. SB 150,00", [a], allocs=[(a, a["gross"] - 15000)], kind="versicherung")
    # K02 Vogt Versicherung kuerzt
    c = S["vogt"]
    j = mk_job(D(2026, 7, 30), D(2026, 8, 10), c, c["vehicles"][0], "Unfallinstandsetzung Heck lt. Gutachten, Verbringung Lackiererei",
               key="J_VOGT")
    a = mk_inv(D(2026, 8, 10), c, 190172, "Unfallinstandsetzung Heck (NordWest)", "Vers.", job=j, due=D(2026, 8, 24),
               key="VOGT", terms="vers")
    ins = INSURERS["nordwest"]
    mk_pay(D(2026, 8, 27), 215000, "bank", ins["name"], ins["iban"], ins["bic"],
           "SCHADEN KH 26-118274 VN KLEIN ANSPR. VOGT RG {0} KUERZUNG VERBRINGUNG", [a], allocs=[(a, 215000)], kind="versicherung")
    mk_pay(D(2026, 9, 22), a["gross"] - 215000, "bank", c["payer"], c["iban"], c["bic"], "Restbetrag Rg {0} Versicherung", [a],
           allocs=[(a, a["gross"] - 215000)])
    # K23 Lange Unfall: fertig, Rechnung erst L2
    c = S["lange"]
    j = mk_job(D(2026, 8, 17), D(2026, 8, 28), c, c["vehicles"][0], "Unfallinstandsetzung Front lt. KVA (Stoßfänger, Scheinwerfer, Kühler)",
               key="J_LANGE")
    j["unbilled_status"] = "abgeholt"
    j["bem"][1] = "Fzg. am 29.08. abgeholt, Rechnung erst nach Gutachten Versicherung"
    a = mk_inv(D(2026, 9, 21), c, 286765, "Unfallinstandsetzung Front (NordWest)", "Vers.", job=j, due=D(2026, 10, 5),
               key="LANGE", terms="vers")
    mk_pay(D(2026, 10, 14), a["gross"], "bank", ins["name"], ins["iban"], ins["bic"], "SCHADEN KH 26-120551 RG {0} LANGE",
           [a], allocs=[(a, a["gross"])], kind="versicherung")
    # K23 Brune Klimaservice (L1 unbilled, L2 berechnet)
    c = F["brune"]
    j = mk_job(D(2026, 9, 9), D(2026, 9, 9), c, c["vehicles"][0], "Klimaservice, Kältemittel ergänzt", key="J_BRUNE_KL")
    j["unbilled_status"] = "abgeholt"
    a = mk_inv(D(2026, 9, 16), c, 16400, "Klimaservice", "Überw.", job=j, due=D(2026, 9, 30), key="BRUNE_KL", terms="net14")
    mk_pay(D(2026, 9, 29), a["gross"], "bank", c["payer"], c["iban"], c["bic"], "Rechnung {0}", [a], allocs=[(a, a["gross"])])
    # K23 Hellweg Raederwechsel (nie berechnet)
    c = F["hellweg"]
    j = mk_job(D(2026, 9, 10), D(2026, 9, 11), c, ("SO-AH 21", "Opel Corsa"),
               "Räderwechsel 4 Mietwagen (SO-AH 21, 22, 23, 24) inkl. Einlagerung", key="J_HELL_RAD")
    j["bem"][1] = j["bem"][2] = "Sammelrechnung Ende Monat"
    # K23 Albers (L1 unbilled, L2 berechnet)
    c = S["albers"]
    j = mk_job(D(2026, 9, 2), D(2026, 9, 3), c, c["vehicles"][0], "Bremsscheiben und -beläge vorne erneuert", key="J_ALB")
    j["unbilled_status"] = "abgeholt"
    a = mk_inv(D(2026, 9, 18), c, 37800, "Bremsen VA", "Überw.", job=j, due=D(2026, 10, 2), key="ALBERS", terms="net14")
    mk_pay(D(2026, 10, 6), a["gross"], "bank", c["payer"], c["iban"], c["bic"], "Rechnung {0}", [a], allocs=[(a, a["gross"])])
    # K24 Kulanz, intern, KVA abgelehnt
    j = mk_job(D(2026, 8, 24), D(2026, 8, 24), F["kaiser"], F["kaiser"]["vehicles"][0], "Nacharbeit Bremse (Geräusch) zu Rg. @KSK",
               special="kulanz", key="J_KULANZ")
    j["bem"][1] = j["bem"][2] = "ohne Berechnung (Kulanz, Chef)"
    j = mk_job(D(2026, 7, 16), D(2026, 7, 17), None, ("LP-AL 112", "VW Caddy"), "Werkstattwagen: Inspektion, HU/AU",
               special="intern", key="J_INTERN", kunde="eigenes Fzg. (Werkstattwagen)")
    cc = R.choice(PRIVATE_POOL)
    j = mk_job(D(2026, 8, 11), None, cc, cc["vehicles"][0], "KVA Kupplung + ZMS – Kunde lässt woanders machen",
               special="storniert", key="J_KVA")
    # K28 Hagemann Betragskorrektur in L2
    c = S["hagemann"]
    j = mk_job(D(2026, 9, 2), D(2026, 9, 3), c, c["vehicles"][0], "Inspektion, Zahnriemen geprüft", key="J_HAG")
    a = mk_inv(D(2026, 9, 3), c, 50000, "Inspektion + Zahnriemen", "Überw.", job=j, due=D(2026, 9, 17), key="HAGEMANN", terms="net14")
    a["ovr"][2] = dict(net=45000, vat=8550, gross=53550)
    mk_pay(D(2026, 9, 25), 53550, "bank", c["payer"], c["iban"], c["bic"], "Rechnung {0} korrigiert", [a], allocs=[(a, 53550)])
    # K29 Kraemer -> Storno -> Westfalen Leasing
    c = S["kraemer"]
    j = mk_job(D(2026, 9, 9), D(2026, 9, 10), c, c["vehicles"][0], "Inspektion, Bremsen VA+HA", key="J_KRAE")
    a = mk_inv(D(2026, 9, 10), c, 108200, "Inspektion + Bremsen VA/HA", "Überw.", job=j, due=D(2026, 9, 24), key="KRAEMER", terms="net14")
    s = mk_inv(D(2026, 9, 22), c, -108200, "Storno zu Rg. {0} (Rechnung an Leasinggeber)", "", kind="S", ref=a, key="KRAEMER_S",
               pdf="Stornorechnung_{nr}.pdf", plate=a["plate"], t=0.40, auftr=False)
    w = mk_inv(D(2026, 9, 22), S["leasing"], 108200, "Inspektion + Bremsen VA/HA (Fzg. LP-SK 912)", "Überw.", job=j,
               due=D(2026, 10, 22), key="LEASING", pdf="RE_{nr}.pdf", plate="LP-SK 912", terms="net30", t=0.41)
    j["bem"][2] = "Rg. an Westfalen Leasing (Vertrag WL-55123)"
    # K31 Brune Gutschrift auf unbezahlte Rechnung
    c = F["brune"]
    j = mk_job(D(2026, 10, 5), D(2026, 10, 6), c, c["vehicles"][1], "Bremsen VA+HA erneuert, Bremsflüssigkeit", key="J_BRUNE_B")
    a = mk_inv(D(2026, 10, 6), c, 65000, "Bremsen VA + HA + Bremsfl.", "Überw.", job=j, due=D(2026, 10, 20), key="BRUNE_B",
               pdf="RE_{nr}.pdf", terms="net14")
    mk_inv(D(2026, 10, 9), c, -5000, "Gutschrift zu Rg. {0}: Bremsflüssigkeit doppelt berechnet", "", kind="G", ref=a,
           key="BRUNE_G", pdf="Gutschrift_{nr}.pdf", plate="LP-GB 78", auftr=False)
    # K32 Sonnenschein Avis-Sammelzahlung
    sc, s2 = F["sonn"], F["sonn2"]
    j1 = mk_job(D(2026, 9, 29), D(2026, 9, 30), sc, sc["vehicles"][0], "Inspektion", key="J_SAV1")
    x1 = mk_inv(D(2026, 9, 30), sc, 23500, "Inspektion", "Überw.", job=j1, due=D(2026, 10, 30), key="SAV1", terms="sonn")
    j2 = mk_job(D(2026, 10, 1), D(2026, 10, 1), s2, s2["vehicles"][0], "Bremsen vorne, Wischerblätter", key="J_SAV2")
    x2 = mk_inv(D(2026, 10, 1), s2, 31800, "Bremsen VA + Wischer", "Überw.", job=j2, due=D(2026, 10, 31), key="SAV2", terms="sonn")
    j3 = mk_job(D(2026, 10, 2), D(2026, 10, 2), sc, sc["vehicles"][4], "Räderwechsel inkl. Einlagerung, Batterie", key="J_SAV3")
    x3 = mk_inv(D(2026, 10, 2), sc, 12900, "Räderwechsel + Batterie", "Überw.", job=j3, due=D(2026, 11, 1), key="SAV3", terms="sonn")
    sks = [(x["gross"] * 2 + 50) // 100 for x in (x1, x2, x3)]
    mk_pay(D(2026, 10, 8), sum(x["gross"] for x in (x1, x2, x3)) - sum(sks), "bank", sc["payer"], sc["iban"], sc["bic"],
           "ZAHLUNG LT. AVIS VOM 07.10.2026 KD-NR. 10412", [], allocs=[(x, x["gross"] - s_) for x, s_ in zip((x1, x2, x3), sks)],
           skontos=list(zip((x1, x2, x3), sks)), cref="SONN-2610-AVIS")
    # K34 Arslan Anzahlung
    c = S["arslan"]
    mk_pay(D(2026, 10, 20), 25000, "bank", c["payer"], c["iban"], c["bic"], "Anzahlung Winterreifen Arslan", [], kind="ohne")
    # K35 Steinhoff Achsvermessung, Wiebe Inspektion (fertig, nicht berechnet)
    c = F["steinhoff"]
    j = mk_job(D(2026, 10, 8), D(2026, 10, 9), c, c["vehicles"][1], "Achsvermessung, Spur eingestellt", key="J_STEIN_ACHS")
    j["unbilled_status"] = "abgeholt"
    c = S["wiebe"]
    j = mk_job(D(2026, 10, 22), D(2026, 10, 23), c, c["vehicles"][0], "Inspektion, Räderwechsel", key="J_WIEBE")
    # Auftraege, die am Stichtag noch in Arbeit sind (keine "nicht berechnet"-Faelle)
    c = R.choice(PRIVATE_POOL)
    j = mk_job(D(2026, 9, 14), D(2026, 9, 16), c, c["vehicles"][0], "Zahnriemen und Wasserpumpe erneuert", key="J_INARB1")
    a = mk_inv(D(2026, 9, 16), c, 71450, "Zahnriemen + WaPu", "EC", job=j, terms="ec", key="INARB1")
    pay_bar_ec(a, "EC")
    c = R.choice(PRIVATE_POOL)
    j = mk_job(D(2026, 10, 26), None, c, c["vehicles"][0], "Kupplung erneuert", key="J_INARB2")
    j["wait"] = "wartet auf Teile"
    c = R.choice(PRIVATE_POOL)
    j = mk_job(D(2026, 10, 27), None, c, c["vehicles"][0], "Inspektion, HU/AU-Vorführung", key="J_INARB3")
    j["wait"] = "in Arbeit"


planted()

# ---------- normale Auftraege ----------
LAMBDA = {6: 2.3, 7: 1.9, 8: 1.8, 9: 2.4, 10: 3.3}
FLEET_W = [("sonn", 3.0), ("hellweg", 2.0), ("kaiser", 1.4), ("yilmaz", 0.8), ("steinhoff", 0.6), ("brune", 0.6),
           ("rinsche", 0.4), ("sonn2", 0.5)]


def poisson(lam):
    L, k, p = pow(2.718281828459045, -lam), 0, 1.0
    while True:
        p *= R.random()
        if p <= L:
            return k
        k += 1


def choose_private(d):
    pool = [c for c in PRIVATE_POOL if c["angelegt"] <= d] + [c for c in NEW_POOL if c["angelegt"] <= d]
    # Stammkunden kommen oefter
    w = [3 if c["kd"] % 3 == 0 else 1 for c in pool]
    return R.choices(pool, weights=w)[0]


d = BOOK_START
while d <= CUT[2]:
    if is_shop_day(d):
        sat = d.weekday() == 5
        lam = LAMBDA[d.month] * (0.45 if sat else 1.0)
        if sat and d.month == 10:
            lam = 2.2
        n = poisson(lam)
        for _ in range(n):
            fleet = (not sat) and R.random() < 0.17
            if fleet:
                fk = R.choices([k for k, _ in FLEET_W], weights=[w for _, w in FLEET_W])[0]
                if fk == "sonn2" and d < D(2026, 7, 7):
                    fk = "sonn"
                c = FLEET[fk]
            else:
                c = choose_private(d)
            veh = R.choice(c["vehicles"])
            items = pick_work(d, not fleet, sat)
            net = price(items)
            big = net > 40000
            k = 0 if (sat or not big) else R.choice([0, 1, 1, 2, 3])
            ann = sub_shop_days(d, k)
            if ann < BOOK_START:
                ann = BOOK_START
            j = mk_job(ann, d, c, veh, arbeiten_text(items))
            # Unfaelle ueber Versicherung (selten)
            if not fleet and not sat and R.random() < 0.035:
                net = R.randrange(140000, 420000)
                j["arbeiten"] = "Unfallinstandsetzung lt. Gutachten"
                inv = mk_inv(d, c, net, "Unfallinstandsetzung (Vers.)", "Vers.", job=j, due=d + days(14), terms="vers")
                insk = R.choice(["nordwest", "soester", "rhein"])
                ins = INSURERS[insk]
                pd_ = next_bank_day(d + days(R.randint(12, 35)))
                mk_pay(pd_, inv["gross"], "bank", ins["name"], ins["iban"], ins["bic"],
                       f"SCHADEN {R.randrange(10, 99)}-{R.randrange(100000, 999999)} RG {{0}} {veh[0]}", [inv],
                       allocs=[(inv, inv["gross"])], kind="versicherung")
                continue
            if fleet:
                inv = mk_inv(d, c, net, leistung_text(items), R.choice(["Überw.", "Überw.", "Überweisung"]), job=j,
                             due=due_for(c, d) if c["profile"] != "sonn" else d + days(30),
                             terms="sonn" if c["profile"] == "sonn" else ("net30" if "30" in c["zb"] else "net14"),
                             pdf=("RE_{nr}.pdf" if c["profile"] in ("kaiser", "hellweg") else None))
                fleet_payment(inv)
                continue
            gross = net + vat_of(net)
            r = R.random()
            if gross < 10000:
                how = "bar" if r < 0.55 else "EC"
            elif gross <= 60000:
                how = "bar" if r < 0.25 else ("EC" if r < 0.75 else "Überw.")
            else:
                how = "bar" if r < 0.10 else ("EC" if r < 0.55 else "Überw.")
            if how == "Überw.":
                inv = mk_inv(d, c, net, leistung_text(items), R.choice(["Überw.", "Überweisung", "ÜW"]), job=j,
                             due=d + days(14), terms="net14")
                private_transfer(inv)
            else:
                zl = R.choice(["bar", "Bar", "bar"]) if how == "bar" else R.choice(["EC", "EC", "EC-Karte"])
                inv = mk_inv(d, c, net, leistung_text(items), zl, job=j, terms="bar" if how == "bar" else "ec")
                pay_bar_ec(inv, "bar" if how == "bar" else "EC")
        # Thekenverkauf ca. 1x pro Woche
        if R.random() < 0.18:
            art, lo, hi = R.choice(THEKE)
            net = R.randrange(lo, hi)
            how = "bar" if R.random() < 0.7 else "EC"
            inv = mk_inv(d, None, net, art, "bar" if how == "bar" else "EC", terms="bar" if how == "bar" else "ec",
                         auftr=False, kunde="Barverkauf")
            pay_bar_ec(inv, how)
    d += ONE

# Falle K24: Rechnung existiert, im L1-Auftragsbuch fehlt die Rg-Nr. -> eine private EC-Rechnung Anfang September
cands = [v for v in INVS if v["kind"] == "R" and v["key"] is None and v["zahlart"] in ("EC", "EC-Karte")
         and D(2026, 9, 7) <= v["date"] <= D(2026, 9, 10) and v["job"] is not None and v["cust"] is not None]
trap = cands[0]
trap["key"] = "TRAP_RG"
KEY["TRAP_RG"] = trap
trap["auftr"] = True
trap["job"]["key"] = "J_TRAP_RG"
KEY["J_TRAP_RG"] = trap["job"]
trap["job"]["rg_ovr"][1] = ""
trap["job"]["status_ovr"][1] = "fertig"

# ---------- Nummern vergeben ----------
for i, j in enumerate(sorted(JOBS, key=lambda x: (x["annahme"], x["sort"]))):
    j["nr"] = f"26-{500 + i:04d}"
for i, v in enumerate(sorted(INVS, key=lambda x: (x["date"], x["t"]))):
    v["nr"] = 31001 + i
INV_BY_NR = {v["nr"]: v for v in INVS}

# Zahlendreher fuer Pieper: Nummer einer bezahlten bar/EC-Rechnung eines anderen Kunden
pie = KEY["PIEPER"]
s = str(pie["nr"])
typo = None
for strict in (True, False):
    for i in range(len(s) - 1, 0, -1):
        t_ = s[:i - 1] + s[i] + s[i - 1] + s[i + 1:]
        tv = INV_BY_NR.get(int(t_))
        if tv and tv["date"] <= CUT[1] and tv["cust"] is not pie["cust"] and tv["kind"] == "R" and \
                (not strict or tv["zahlart"] in ("EC", "bar", "Bar", "EC-Karte")):
            typo = int(t_)
            break
    if typo:
        break
assert typo, "kein passender Zahlendreher"
KEY["PIEPER_TYPO"] = INV_BY_NR[typo]
for p in PAYS:
    if p["kind"] == "typo":
        p["tpl"] = p["tpl"].replace("@TYPO", str(typo))
KEY["J_KULANZ"]["arbeiten"] = KEY["J_KULANZ"]["arbeiten"].replace("@KSK", str(KEY["KAISER_SK"]["nr"]))

for p in PAYS:
    if p["channel"] in ("bank", "bounce"):
        p["text"] = p["tpl"].format(*[r["nr"] for r in p["refs"]]) if p["refs"] else p["tpl"]
for v in INVS:
    if v["kind"] in ("S", "G") and "{0}" in v["leistung"]:
        v["leistung"] = v["leistung"].format(v["ref"]["nr"])

# ---------------------------------------------------------------------------------------------
# Sichten je Lieferung
# ---------------------------------------------------------------------------------------------


def g(v, field, dl):
    return v["ovr"].get(dl, {}).get(field, v[field])


def inv_due(v, dl):
    du = g(v, "due", dl)
    return du if du is not None else v["date"]


def visible(v, dl):
    return v["date"] <= CUT[dl]


def cancelled(v, dl):
    s_ = v["cancelled_by"]
    return s_ is not None and s_["date"] <= CUT[dl]


def open_amount(v, dl):
    C = CUT[dl]
    rest = g(v, "gross", dl)
    for p, a, eff in v["allocs"]:
        if p["date"] <= C and eff <= C:
            rest -= a
    for p, a in v["skontos"]:
        if p["date"] <= C:
            rest -= a
    for cr in v["credits"]:
        if visible(cr, dl):
            rest += g(cr, "gross", dl)
    return rest


def job_state(j, dl):
    C = CUT[dl]
    if j["annahme"] > C:
        return None
    invs = [v for v in j["invs"] if v["kind"] == "R" and visible(v, dl) and not cancelled(v, dl)]
    fertig = j["fertig"] if (j["fertig"] and j["fertig"] <= C) else None
    if j["special"] == "storniert":
        st, rg, fertig = "storniert", "", None
    elif fertig is None:
        st, rg = j["wait"], ""
    elif invs:
        st, rg = "abgerechnet", invs[-1]["nr"]
    elif j["special"] == "kulanz":
        st, rg = "fertig", "Kulanz"
    elif j["special"] == "intern":
        st, rg = "fertig", "intern"
    else:
        st, rg = j["unbilled_status"], ""
    if dl in j["status_ovr"]:
        st = j["status_ovr"][dl]
    if dl in j["rg_ovr"]:
        rg = j["rg_ovr"][dl]
    return dict(status=st, rg=rg, fertig=fertig)


def compute_truth(dl):
    C = CUT[dl]
    T = TODAY[dl]
    op = {}
    overdue = []
    for v in sorted(INVS, key=lambda x: x["nr"]):
        if v["kind"] != "R" or not visible(v, dl) or cancelled(v, dl):
            continue
        rest = open_amount(v, dl)
        if rest > 0:
            op[str(v["nr"])] = eur(rest)
            if inv_due(v, dl) < T:
                overdue.append(str(v["nr"]))
    unb = []
    for j in sorted(JOBS, key=lambda x: x["nr"]):
        st = job_state(j, dl)
        if st is None or st["fertig"] is None or j["special"] in ("kulanz", "intern", "storniert"):
            continue
        has = any(v["kind"] == "R" and visible(v, dl) for v in j["invs"])
        if not has:
            unb.append(j["nr"])
    rec = []
    for p in sorted(PAYS, key=lambda x: (x["date"], x["amount"])):
        if p["channel"] != "bank" or p["date"] > C or p["date"] < BANK_FROM[1]:
            continue
        assigned = sum(a for _, a, eff in p["allocs"] if eff <= C)
        if p["amount"] - assigned > 0:
            rec.append(eur(p["amount"] - assigned))
    return {"folder": f"lieferung{dl}", "today": TODAY[dl].isoformat(), "open": op, "overdue": overdue,
            "unbilled": unb, "receipts": rec}


# ---------------------------------------------------------------------------------------------
# Bankzeilen
# ---------------------------------------------------------------------------------------------
BANKLINES = []


def bl(date, umsatzart, name, text, amount, iban="", bic="", cref="", mref="", gid="", src=None):
    BANKLINES.append(dict(date=date, umsatzart=umsatzart, name=name, text=text, amount=amount, iban=iban, bic=bic,
                          cref=cref, mref=mref, gid=gid, src=src))


PAYWEST = dict(name="PayWest Terminal GmbH", iban=make_iban("30050000", 7712040), bic="WELADEDDXXX")
for p in PAYS:
    if p["channel"] == "bank":
        bl(p["date"], p["umsatzart"], p["payer"], p["text"], p["amount"], p["iban"], p["bic"], p["cref"], src=p)
    elif p["channel"] == "bounce":
        bl(p["date"], "Rücklastschrift", p["payer"], p["text"], p["amount"] - 300, p["iban"], p["bic"], src=p)

# girocard-Tagesumsaetze -> Gutschrift am naechsten Bankarbeitstag
ec_days = {}
for p in PAYS:
    if p["channel"] == "ec":
        ec_days.setdefault(p["date"], []).append(p)
for day in sorted(ec_days):
    ps = ec_days[day]
    amt = sum(p["amount"] for p in ps)
    bl(bank_day_after(day), "SEPA-Gutschrift", PAYWEST["name"],
       f"KARTENUMSAETZE {dde(day)} TID 65781234 GIROCARD {len(ps)} TRX", amt, PAYWEST["iban"], PAYWEST["bic"],
       cref=f"PW{day:%Y%m%d}65781234", src="ec")

# Bareinzahlungen freitags
cash = 30000
d = BOOK_START
cash_by_day = {}
for p in PAYS:
    if p["channel"] == "bar":
        cash_by_day[p["date"]] = cash_by_day.get(p["date"], 0) + p["amount"]
while d <= CUT[2]:
    cash += cash_by_day.get(d, 0)
    if d.weekday() == 4 and is_bank_day(d):
        cash -= R.randrange(0, 9000)  # kleine Barausgaben
        if cash > 40000:
            dep = ((cash - 20000) // 5000) * 5000
            cash -= dep
            bl(d, "Bareinzahlung", "Autoservice Lindner GmbH", "EINZAHLUNG EINZAHLAUTOMAT FIL. LIPPSTADT", dep, src="bar")
    d += ONE

# Ausgaben
EMPL = [("Dennis Hartmann", 278044), ("Marco Ostermann", 261210), ("Luca Krüger", 81233), ("Petra Lindner", 145000),
        ("Jürgen Lindner", 390000), ("Andrea Kemper", 53800)]
EMPL_IBAN = {name: new_iban() for name, _ in EMPL}


def month_last_bank_day(y, m):
    d_ = D(y, m + 1, 1) - ONE if m < 12 else D(y, 12, 31)
    return prev_bank_day(d_)


rev_by_month = {}
ec_by_month = {}
for v in INVS:
    rev_by_month[v["date"].month] = rev_by_month.get(v["date"].month, 0) + v["net"]
for p in PAYS:
    if p["channel"] == "ec":
        ec_by_month[p["date"].month] = ec_by_month.get(p["date"].month, 0) + p["amount"]
MON = {6: "06/2026", 7: "07/2026", 8: "08/2026", 9: "09/2026", 10: "10/2026"}
for m in (7, 8, 9, 10):
    bl(next_bank_day(D(2026, m, 3)), "Dauerauftrag", "Grundstücksgemeinschaft Brockmann",
       f"Miete Werkstatthalle Erwitter Str. 112 {MON[m]}", -285000, make_iban("41650001", 1200448812), "WELADED1LIP")
    bl(next_bank_day(D(2026, m, 5)), "SEPA-Lastschrift", "Kreis-Leasing GmbH",
       f"Leasingrate Vertrag 77-40512 Hebebuehne/Diagnose {MON[m]}", -61880, make_iban("37050198", 1900223344),
       "COLSDE33XXX", mref="KL-77-40512", gid="DE91ZZZ00000188411")
    bl(next_bank_day(D(2026, m, 2)), "SEPA-Lastschrift", PAYWEST["name"],
       f"ENTGELTE KARTENAKZEPTANZ {MON[m - 1]} TID 65781234", -((ec_by_month.get(m - 1, 0) * 95 + 5000) // 10000),
       PAYWEST["iban"], PAYWEST["bic"], mref="PW-65781234", gid="DE55ZZZ00001234567")
    bl(next_bank_day(D(2026, m, 10)), "SEPA-Überweisung", "Finanzamt Lippstadt",
       f"UST-VA {MON[m - 1]} STNR 338/5712/0815", -int(rev_by_month.get(m - 1, 0) * 0.19 * 0.55),
       make_iban("44000000", 4400152), "MARKDEF1440")
    bl(next_bank_day(D(2026, m, 15)), "SEPA-Lastschrift", "Energieversorgung Hellweg GmbH",
       f"Abschlag Strom/Gas {MON[m]} VK 400128833", -41200, make_iban("41650001", 900011223), "WELADED1LIP",
       mref="EVH-400128833", gid="DE12ZZZ00000077123")
    bl(next_bank_day(D(2026, m, 20)), "SEPA-Lastschrift", "TelKom Westfalen GmbH", f"Kd-Nr 8812044 Rechnung {MON[m]}",
       -8990, make_iban("20041111", 300112290), "COBADEHDXXX", mref="TKW8812044", gid="DE33ZZZ00000999111")
    bl(next_bank_day(D(2026, m, 12)), "SEPA-Lastschrift", "Tank & Service Card GmbH", f"Tankkarte Abrechnung {MON[m - 1]}",
       -R.randrange(25000, 48000), make_iban("50070010", 912345600), "DEUTDEFFXXX", mref="TSC-55120", gid="DE77ZZZ00000555222")
    last = month_last_bank_day(2026, m)
    if last <= CUT[2]:
        for name, amt in EMPL:
            if name == "Andrea Kemper" and m != 10:
                continue
            bl(last, "SEPA-Überweisung", name, f"Lohn/Gehalt {MON[m]}", -amt, *EMPL_IBAN[name])
        kk = prev_bank_day(last - days(2))
        bl(kk, "SEPA-Lastschrift", "BKK Hellweg", f"Beitraege {MON[m]} Betriebsnr 44587712", -R.randrange(380000, 420000),
           make_iban("41650001", 700334455), "WELADED1LIP", mref="BKK-44587712", gid="DE44ZZZ00000333444")
# Teile-Express Halbmonatsrechnungen, Lastschrift 10 Tage nach Rechnungsdatum
TEILE = dict(name="Teile-Express Westfalen GmbH", iban=make_iban("47260121", 8801455200), bic="DGPBDE3MXXX")
TEILE_INV = []
nr_te = 884301
for m in (6, 7, 8, 9, 10):
    for half in (1, 2):
        start = D(2026, m, 1) if half == 1 else D(2026, m, 16)
        end = D(2026, m, 15) if half == 1 else (D(2026, m + 1, 1) - ONE)
        inv_date = end + ONE
        net_rev = sum(v["net"] for v in INVS if start <= v["date"] <= end and v["kind"] == "R")
        net_cost = int(net_rev * R.uniform(0.27, 0.33))
        gross_cost = net_cost + vat_of(net_cost)
        deb = next_bank_day(inv_date + days(9))
        nr_te += R.randrange(90, 140)
        TEILE_INV.append(dict(nr=nr_te, date=inv_date, start=start, end=end, net=net_cost, gross=gross_cost, debit=deb))
        bl(deb, "SEPA-Lastschrift", TEILE["name"], f"RE {nr_te} VOM {dde(inv_date)} KD-NR 47110", -gross_cost,
           TEILE["iban"], TEILE["bic"], mref="TE-47110", gid="DE68ZZZ00000412345")
REIFEN = dict(name="Reifen Wolf GmbH & Co. KG", iban=make_iban("41662465", 3300441), bic="GENODEM1ERW")
bl(D(2026, 9, 28), "SEPA-Überweisung", REIFEN["name"], "RE 2026-44871 Winterreifen Lager", -485000, REIFEN["iban"], REIFEN["bic"])
bl(D(2026, 10, 19), "SEPA-Überweisung", REIFEN["name"], "RE 2026-46102", -312040, REIFEN["iban"], REIFEN["bic"])
bl(D(2026, 9, 30), "Entgelt", "Deutsche Bank", "ENTGELT KONTOFUEHRUNG 01.07.2026-30.09.2026", -3870)
bl(D(2026, 8, 12), "Entgelt", "Deutsche Bank", "ENTGELT RUECKLASTSCHRIFT", 0)  # Platzhalter, wird entfernt
BANKLINES[:] = [b for b in BANKLINES if b["amount"] != 0 and b["date"] >= BANK_FROM[1] and b["date"] <= CUT[2]]

# Reihenfolge und Saldo
BANKLINES.sort(key=lambda b: (b["date"], 0 if b["amount"] > 0 else 1, b["name"], b["text"], b["amount"]))
bal = START_BALANCE
for b in BANKLINES:
    bal += b["amount"]
    b["saldo"] = bal

# ---------------------------------------------------------------------------------------------
# Schreiben
# ---------------------------------------------------------------------------------------------


def xl_date(d_):
    return dt.datetime(d_.year, d_.month, d_.day)


MONTHS_DE = {6: "Juni", 7: "Juli", 8: "August", 9: "September", 10: "Oktober"}


def book_payment_view(v, dl):
    """Was Petra zum Stichtag im Buch stehen hat: (bez_am, betrag_erh, notiz)"""
    C, CK = CUT[dl], CHECK[dl]
    if v["kind"] != "R":
        return None, None
    known = 0
    last = None
    for p, a, eff in v["allocs"]:
        if p["channel"] in ("bar", "ec"):
            ok = p["date"] <= C
        elif p["channel"] == "bounce":
            ok = p["date"] <= CK
        elif p["kind"] == "anzahlung":
            ok = eff <= C
        else:
            ok = p["date"] <= CK
        if ok:
            known += a
            if a > 0:
                last = max(last, p["date"]) if last else p["date"]
    sk = sum(a for p, a in v["skontos"] if p["date"] <= CK)
    cr = sum(-g(c_, "gross", dl) for c_ in v["credits"] if visible(c_, dl))
    gross = g(v, "gross", dl)
    erh = known if known != 0 else None
    if known + sk + cr >= gross and known > 0:
        return last, erh
    return None, erh


def write_book(path, dl):
    C = CUT[dl]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rechnungen 2026"
    ws["A1"] = "Rechnungsausgangsbuch 2026"
    ws["A1"].font = Font(bold=True, size=14)
    ws["F1"] = "Autoservice Lindner GmbH"
    ws["A2"] = f"Stand: {dde(C)}"
    hdr = ["Rg-Nr.", "Datum", "Kd-Nr.", "Kunde", "Kennz.", "Auftr.", "Leistung", "Netto", "USt 19%", "Brutto",
           "Zahlart", "fällig", "bez. am", "Betrag erh.", "Bemerkung"]
    for i, h in enumerate(hdr, 1):
        ws.cell(row=4, column=i, value=h).font = Font(bold=True)
    row = 5
    cur_m = None
    tot = [0, 0, 0]
    for v in sorted([x for x in INVS if visible(x, dl)], key=lambda x: x["nr"]):
        if v["date"].month != cur_m:
            cur_m = v["date"].month
            ws.cell(row=row, column=1, value=f"{MONTHS_DE[cur_m]} 2026").font = Font(bold=True, italic=True)
            row += 1
        net, vat, gross = g(v, "net", dl), g(v, "vat", dl), g(v, "gross", dl)
        tot[0] += net
        tot[1] += vat
        tot[2] += gross
        bez, erh = book_payment_view(v, dl)
        kd = g(v, "kd", dl)
        vals = [v["nr"], xl_date(v["date"]), kd, g(v, "kunde", dl), v["plate"] or None,
                (v["job"]["nr"] if (v["job"] and v["auftr"]) else None), v["leistung"], eur(net), eur(vat), eur(gross),
                v["zahlart"] or None, None, None, eur(erh) if erh is not None else None, None]
        if v["kind"] == "R" and v["zahlart"] not in ("bar", "Bar", "EC", "EC-Karte"):
            du = g(v, "due", dl)
            vals[11] = xl_date(du) if du else None
        if v["kind"] == "R" and v["key"] == "SCHULZ":
            vals[11] = xl_date(g(v, "due", dl))
        if bez is not None:
            vals[12] = xl_date(bez)
            if v["date"].month == 6 and v["zahlart"] not in ("bar", "Bar", "EC", "EC-Karte") and v["nr"] % 4 == 1:
                vals[12] = bez.strftime("%d.%m.%y")      # als Text eingetippt
        vals[14] = book_note(v, dl, bez, erh) or None
        for i, val in enumerate(vals, 1):
            cell = ws.cell(row=row, column=i, value=val)
            if i in (2, 12, 13) and isinstance(val, dt.datetime):
                cell.number_format = "DD.MM.YYYY"
            if i in (8, 9, 10, 14) and val is not None:
                cell.number_format = '#,##0.00 "€"'
        row += 1
    row += 1
    ws.cell(row=row, column=4, value="Summe").font = Font(bold=True)
    for i, t_ in zip((8, 9, 10), tot):
        cc_ = ws.cell(row=row, column=i, value=eur(t_))
        cc_.number_format = '#,##0.00 "€"'
        cc_.font = Font(bold=True)
    widths = [9, 11, 8, 30, 12, 9, 34, 11, 10, 11, 10, 11, 11, 11, 60]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A5"
    # Blatt Mahnungen
    wm = wb.create_sheet("Mahnungen")
    mh = ["Rg-Nr.", "Kunde", "offen €", "Erinnerung", "1. Mahnung", "2. Mahnung", "Bemerkung"]
    for i, h in enumerate(mh, 1):
        wm.cell(row=1, column=i, value=h).font = Font(bold=True)
    for r_i, rowv in enumerate(mahn_rows(dl), 2):
        for i, val in enumerate(rowv, 1):
            cell = wm.cell(row=r_i, column=i, value=val)
            if isinstance(val, dt.datetime):
                cell.number_format = "DD.MM.YYYY"
            if i == 3:
                cell.number_format = '#,##0.00'
    for i, w in enumerate([9, 28, 10, 11, 11, 11, 50], 1):
        wm.column_dimensions[get_column_letter(i)].width = w
    wb.properties.creator = "Petra Lindner"
    wb.properties.lastModifiedBy = "Petra Lindner"
    wb.properties.created = dt.datetime(2026, 6, 1, 8, 0, 0)
    wb.properties.modified = dt.datetime(C.year, C.month, C.day, 17, 45, 0)
    wb.save(path)


def book_note(v, dl, bez, erh):
    k = v["key"]
    nr = lambda key: KEY[key]["nr"]
    if v["kind"] == "S":
        return f"Storno zu Rg. {v['ref']['nr']}"
    if v["kind"] == "G":
        return f"Gutschrift zu Rg. {v['ref']['nr']}"
    if v["cancelled_by"] is not None and cancelled(v, dl):
        return f"storniert durch {v['cancelled_by']['nr']}" + (", neu an Westfalen Leasing" if k == "KRAEMER" else ", neu: " + str(nr("HELL_N")) if k == "HELL_A" else "")
    if k == "KAISER_SK":
        return "2% Skonto abgezogen – nicht vereinbart (30 Tage netto)! 23,12 nachgefordert am 27.07."
    if k == "YILMAZ":
        return "zahlt in Raten (Mail 02.08.)" if dl == 1 else "zahlt in Raten (Mail 02.08.), 2 Raten erhalten"
    if k == "SCHULZ":
        return "Lastschrift (ELV) geplatzt 12.08. (+3,00 Geb.) – Kd. angeschrieben, Zahlung bis 26.08." + (" 2. Mahnung 01.10." if dl == 2 else "")
    if k == "STEIN_R":
        return f"abzgl. Gutschrift {nr('STEIN_G')}"
    if k == "GLAS":
        return "SB 150,- per EC bezahlt, Rest Teilkasko (Soester Kreis-Vers.)"
    if k == "VOGT":
        if dl == 1:
            return "Vers. hat gekürzt (Verbringung), Rest bei Kundin anfordern"
        return "Vers. hat gekürzt (Verbringung), Rest 113,05 von Kundin am 22.09."
    if k == "LANGE":
        return "NordWest, Schaden KH 26-120551"
    if k == "HAGEMANN" and dl == 2:
        return "korrigiert 18.09. (10% Nachlass lt. Chef, s. Mail)"
    if k == "KUHLMANN":
        return "abzgl. Anzahlung 500,00 v. 02.09."
    if k == "LEASING":
        return f"Fzg. LP-SK 912, Fahrer S. Krämer, Vertrag WL-55123 (ersetzt {nr('KRAEMER')})"
    if k == "BRUNE_B" and dl == 2:
        return f"Gutschrift {nr('BRUNE_G')} abziehen"
    if k == "BRINKMANN" and dl == 2:
        return "Kundin hat am 03.09. zusätzlich überwiesen (doppelt) – Erstattung klären"
    if k == "PIEPER" and dl == 2:
        return f"Kd. hat bei Überweisung Rg-Nr. {KEY['PIEPER_TYPO']['nr']} angegeben"
    if k == "POHL":
        return "1. Mahnung 20.08." if dl == 1 else "1. Mahnung 20.08., 400,- am 04.09., 2. Mahnung 28.09."
    if k in ("SONN_A", "SONN_B") and bez:
        return "2% Skonto"
    if v["skontos"] and bez:
        return "2% Skonto"
    if k == "SAV2":
        return "Fzg. LP-PS 106"
    if v["zahlart"] == "Vers." and k is None:
        return "Abtretung an Werkstatt"
    if v["nr"] % 17 == 0 and v["kind"] == "R":
        return random.Random(v["nr"]).choice(["Kd. holt Fzg. Samstag", "Stammkunde", "Ersatzteil nachbestellt", "Altteile mitgegeben",
                          "Fzg. stand 2 Tage"])
    return ""


def mahn_rows(dl):
    rows = []
    P, Sch, K, B = KEY["POHL"], KEY["SCHULZ"], KEY["KAISER_SK"], KEY["BECKER_LP"]
    rows.append([P["nr"], "Pohl, Dieter", 1120.00 if dl == 1 else 720.00, xl_date(D(2026, 8, 5)), xl_date(D(2026, 8, 20)),
                 xl_date(D(2026, 9, 28)) if dl == 2 else None, "" if dl == 1 else "400,- am 04.09. erhalten"])
    rows.append([Sch["nr"], "Schulz, Kevin", 456.70, xl_date(D(2026, 8, 14)), xl_date(D(2026, 8, 28)) if True else None,
                 xl_date(D(2026, 10, 1)) if dl == 2 else None, "Rücklastschrift, zzgl. 3,00 Gebühr"])
    rows.append([K["nr"], "Elektro Kaiser", 23.12, xl_date(D(2026, 7, 27)), None, None, "Skonto unberechtigt"])
    rows.append([B["nr"], "Becker, Thomas", 498.00, xl_date(D(2026, 9, 1)), xl_date(D(2026, 9, 22)) if dl == 2 else None, None,
                 "" if dl == 1 else "Lippstadt, NICHT Erwitte!"])
    if dl == 2:
        y = KEY["YILMAZ"]
        rows.append([y["nr"], "Taxi Yilmaz", 645.20, xl_date(D(2026, 9, 15)), None, None, "Rest nach 2 Raten"])
    return rows


def write_jobs(path, dl):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Aufträge"
    hdr = ["Auftrag", "Annahme", "Kd.-Nr.", "Kunde", "Kennzeichen", "Fahrzeug", "km", "Arbeiten", "Mech.", "fertig am",
           "Status", "Rg.", "Bemerkung"]
    for i, h in enumerate(hdr, 1):
        ws.cell(row=1, column=i, value=h).font = Font(bold=True)
    row = 2
    for j in sorted(JOBS, key=lambda x: x["nr"]):
        st = job_state(j, dl)
        if st is None:
            continue
        vals = [j["nr"], xl_date(j["annahme"]), j["kd"], j["kunde"], j["plate"], j["model"], j["km"], j["arbeiten"],
                j["mech"], xl_date(st["fertig"]) if st["fertig"] else None, st["status"],
                st["rg"] if st["rg"] != "" else None, j["bem"].get(dl) or None]
        for i, val in enumerate(vals, 1):
            cell = ws.cell(row=row, column=i, value=val)
            if isinstance(val, dt.datetime):
                cell.number_format = "DD.MM.YYYY"
        row += 1
    for i, w in enumerate([9, 11, 8, 26, 12, 20, 8, 60, 6, 11, 16, 9, 45], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:M{row - 1}"
    wb.properties.creator = "Serviceannahme"
    wb.properties.created = dt.datetime(2026, 6, 1, 7, 30, 0)
    wb.properties.modified = dt.datetime(CUT[dl].year, CUT[dl].month, CUT[dl].day, 18, 5, 0)
    wb.save(path)


def write_customers(path, dl):
    C = CUT[dl]
    hdr = ["KdNr", "Anrede", "Name", "Vorname", "Zusatz", "Strasse", "PLZ", "Ort", "Telefon", "Mobil", "E-Mail",
           "Zahlungsbedingung", "Angelegt", "Letzter Auftrag", "Bemerkung"]
    last_job = {}
    for j in JOBS:
        if j["kd"] and j["annahme"] <= C:
            last_job[j["kd"]] = max(last_job.get(j["kd"], j["annahme"]), j["annahme"])
    lines = [";".join(f'"{h}"' for h in hdr)]
    for kd in sorted(CUSTOMERS):
        c = CUSTOMERS[kd]
        if c["angelegt"] > C:
            continue
        cc_ = dict(c)
        cc_.update(c["ovr"].get(dl, {}))
        lj = last_job.get(kd)
        vals = [str(kd), cc_["anrede"], cc_["name"], cc_["vorname"], cc_["zusatz"], cc_["strasse"], cc_["plz"],
                cc_["ort"], cc_["tel"], cc_["mobil"], cc_["mail"], cc_["zb"], dde(cc_["angelegt"]),
                dde(lj) if lj else "", cc_["bem"]]
        lines.append(";".join('"' + str(x).replace('"', "'") + '"' for x in vals))
    with open(path, "wb") as f:
        f.write(("\r\n".join(lines) + "\r\n").encode("cp1252"))


def write_bank(path, dl):
    frm, to = BANK_FROM[dl], CUT[dl]
    rows = [b for b in BANKLINES if frm <= b["date"] <= to]
    end_bal = rows[-1]["saldo"] if rows else START_BALANCE
    hdr = ["Buchungstag", "Wert", "Umsatzart", "Begünstigter / Auftraggeber", "Verwendungszweck", "IBAN", "BIC",
           "Kundenreferenz", "Mandatsreferenz ", "Gläubiger ID", "Fremde Gebühren", "Betrag", "Abweichender Empfänger",
           "Anzahl der Aufträge", "Anzahl der Schecks", "Soll", "Haben", "Währung"]
    n = len(hdr)
    pad = lambda *xs: ";".join(list(xs) + [""] * (n - len(xs)))
    out = [pad("Umsätze Girokonto"),
           pad(f"Kundennummer: 247 1185093 00 / Autoservice Lindner GmbH"),
           pad(f"{dde(frm)} - {dde(to)}"),
           pad("Letzter Kontostand", "", "", "", de_amt(end_bal), "EUR"),
           pad("Vorgemerkte und noch nicht gebuchte Umsätze sind nicht Bestandteil dieser Übersicht."),
           ";".join(hdr)]
    for b in sorted(rows, key=lambda x: (x["date"], BANKLINES.index(x)), reverse=True):
        soll = de_amt(b["amount"]) if b["amount"] < 0 else ""
        haben = de_amt(b["amount"]) if b["amount"] > 0 else ""
        vals = [dde(b["date"]), dde(b["date"]), b["umsatzart"], b["name"], b["text"], b["iban"], b["bic"], b["cref"],
                b["mref"], b["gid"], "", "", "", "", "", soll, haben, "EUR"]
        out.append(";".join(v_.replace(";", ",") for v_ in vals))
    out.append(pad("Kontostand", dde(to), "", "", de_amt(end_bal), "EUR"))
    with open(path, "wb") as f:
        f.write(("\r\n".join(out) + "\r\n").encode("cp1252"))


# ---------- PDFs ----------
FIRM_LINES = ["Autoservice Lindner GmbH", "Kfz-Meisterbetrieb · Erwitter Straße 112 · 59557 Lippstadt",
              "Tel. 02941 18220 · buero@autoservice-lindner.de"]
FOOT = ["Deutsche Bank Lippstadt · IBAN " + " ".join(FIRM_IBAN[i:i + 4] for i in range(0, 22, 4)) + " · BIC DEUTDEDB472",
        "Geschäftsführer: Jürgen Lindner · Amtsgericht Paderborn HRB 9821 · USt-IdNr. DE 287 551 903"]


def positions_for(v, dl):
    net = g(v, "net", dl)
    sign = -1 if net < 0 else 1
    a = abs(net)
    if v["kind"] in ("G", "S"):
        return [(v["leistung"], "1", sign * a, sign * a)]
    aw = max(1, int(a * 0.45 / 950))
    labor = min(aw * 950, a)
    parts = a - labor
    rows = [(f"Arbeitslohn: {v['leistung']}", f"{aw} AW", 950, labor)]
    if parts > 0:
        rows.append(("Ersatzteile/Material lt. Aufstellung", "1", parts, parts))
    if v["key"] == "HAGEMANN" and dl == 2:
        rows = [(f"Arbeitslohn: {v['leistung']}", "26 AW", 950, 24700), ("Ersatzteile/Material lt. Aufstellung", "1", 25300, 25300),
                ("Nachlass 10 % lt. Absprache", "1", -5000, -5000)]
    return rows


def pdf_invoice(path, v, dl, title=None):
    c_ = canvas.Canvas(path, pagesize=A4, invariant=1)
    kind_title = {"R": "RECHNUNG", "G": "GUTSCHRIFT", "S": "STORNORECHNUNG"}[v["kind"]]
    if title:
        kind_title = title
    c_.setTitle(f"{kind_title.title()} {v['nr']}")
    c_.setAuthor("Autoservice Lindner GmbH")
    c_.setFont("Helvetica-Bold", 15)
    c_.drawString(20 * mm, 280 * mm, FIRM_LINES[0])
    c_.setFont("Helvetica", 9)
    c_.drawString(20 * mm, 275 * mm, FIRM_LINES[1])
    c_.drawString(20 * mm, 271 * mm, FIRM_LINES[2])
    cu = v["cust"]
    if v["key"] == "BECKER_ERW" and dl == 1:
        cu = SP["becker_lp"]
    addr = [cu["name"] if cu["kind"] == "firma" else f"{cu['anrede']} {cu['vorname']} {cu['name']}"]
    if cu["zusatz"] and cu["kind"] == "firma":
        addr.append(cu["zusatz"])
    cco = dict(cu)
    cco.update(cu["ovr"].get(dl, {}))
    addr += [cco["strasse"], f"{cco['plz']} {cco['ort']}"]
    c_.setFont("Helvetica", 7)
    c_.drawString(20 * mm, 252 * mm, "Autoservice Lindner GmbH · Erwitter Str. 112 · 59557 Lippstadt")
    c_.setFont("Helvetica", 10)
    y = 246
    for a_ in addr:
        c_.drawString(20 * mm, y * mm, a_)
        y -= 5
    c_.setFont("Helvetica-Bold", 13)
    c_.drawString(20 * mm, 210 * mm, f"{kind_title} Nr. {v['nr']}")
    c_.setFont("Helvetica", 9.5)
    meta = [("Datum", dde(v["date"])), ("Kunden-Nr.", str(g(v, "kd", dl) or "")),
            ("Auftrags-Nr.", v["job"]["nr"] if v["job"] else "-"), ("Kennzeichen", v["plate"] or "-")]
    if v["job"]:
        meta.append(("Fahrzeug", v["job"]["model"]))
        meta.append(("km-Stand", f"{v['job']['km']:,}".replace(",", ".")))
    if v["kind"] in ("G", "S"):
        meta.append(("Bezug", f"Rechnung Nr. {v['ref']['nr']} vom {dde(v['ref']['date'])}"))
    yy = 236
    for k_, val in meta:
        c_.drawString(130 * mm, yy * mm, f"{k_}:")
        c_.drawRightString(190 * mm, yy * mm, val)
        yy -= 4.6
    y = 196
    c_.setFont("Helvetica-Bold", 9)
    c_.drawString(20 * mm, y * mm, "Pos.")
    c_.drawString(32 * mm, y * mm, "Bezeichnung")
    c_.drawRightString(135 * mm, y * mm, "Menge")
    c_.drawRightString(160 * mm, y * mm, "Einzelpreis")
    c_.drawRightString(190 * mm, y * mm, "Gesamt €")
    c_.line(20 * mm, (y - 1.5) * mm, 190 * mm, (y - 1.5) * mm)
    c_.setFont("Helvetica", 9)
    y -= 7
    for i, (txt, qty, ep, tot) in enumerate(positions_for(v, dl), 1):
        c_.drawString(20 * mm, y * mm, str(i))
        c_.drawString(32 * mm, y * mm, txt[:62])
        c_.drawRightString(135 * mm, y * mm, qty)
        c_.drawRightString(160 * mm, y * mm, de_amt(ep))
        c_.drawRightString(190 * mm, y * mm, de_amt(tot))
        y -= 6
    y -= 4
    c_.line(120 * mm, (y + 3) * mm, 190 * mm, (y + 3) * mm)
    for lab, val, bold in (("Summe netto", g(v, "net", dl), False), ("zzgl. 19 % USt.", g(v, "vat", dl), False),
                           ("Rechnungsbetrag" if v["kind"] == "R" else "Betrag", g(v, "gross", dl), True)):
        c_.setFont("Helvetica-Bold" if bold else "Helvetica", 10 if bold else 9)
        c_.drawString(120 * mm, y * mm, lab)
        c_.drawRightString(190 * mm, y * mm, de_amt(val) + " €")
        y -= 5.5
    y -= 6
    c_.setFont("Helvetica", 9)
    if v["kind"] == "R":
        t_ = TERMS.get(v["terms"] or "net14", "").replace("{due}", dde(g(v, "due", dl) or v["date"]))
        c_.drawString(20 * mm, y * mm, t_)
        y -= 5
        if v["key"] == "HAGEMANN" and dl == 2:
            c_.drawString(20 * mm, y * mm, "Diese korrigierte Rechnung ersetzt die Rechnung gleicher Nummer vom 03.09.2026.")
            y -= 5
        if v["key"] == "LEASING":
            c_.drawString(20 * mm, y * mm, "Leasingvertrag WL-55123 · Fahrer: Sven Krämer · ersetzt Rechnung "
                          f"{KEY['KRAEMER']['nr']} (storniert)")
    elif v["kind"] == "G":
        c_.drawString(20 * mm, y * mm, "Der Betrag wird mit der oben genannten Rechnung verrechnet.")
    else:
        c_.drawString(20 * mm, y * mm, "Die oben genannte Rechnung wird hiermit vollständig storniert.")
    c_.setFont("Helvetica", 7.5)
    c_.drawString(20 * mm, 18 * mm, FOOT[0])
    c_.drawString(20 * mm, 14 * mm, FOOT[1])
    c_.showPage()
    c_.save()


def pdf_kva(path):
    j = KEY["J_LANGE"]
    c_ = canvas.Canvas(path, pagesize=A4, invariant=1)
    c_.setTitle("Kostenvoranschlag KV-0892")
    c_.setFont("Helvetica-Bold", 15)
    c_.drawString(20 * mm, 280 * mm, FIRM_LINES[0])
    c_.setFont("Helvetica", 9)
    c_.drawString(20 * mm, 275 * mm, FIRM_LINES[1])
    c_.setFont("Helvetica", 10)
    for i, a_ in enumerate(["Frau Monika Lange", "Bachstraße 15", "59590 Geseke"]):
        c_.drawString(20 * mm, (246 - 5 * i) * mm, a_)
    c_.setFont("Helvetica-Bold", 13)
    c_.drawString(20 * mm, 210 * mm, "KOSTENVORANSCHLAG Nr. KV-0892")
    c_.setFont("Helvetica", 9.5)
    for i, (k_, val) in enumerate([("Datum", "18.08.2026"), ("Auftrags-Nr.", j["nr"]), ("Kennzeichen", "GT-ML 22"),
                                   ("Fahrzeug", "Audi A3 Sportback"), ("Schaden-Nr.", "KH 26-120551 (NordWest)")]):
        c_.drawString(120 * mm, (236 - 4.6 * i) * mm, f"{k_}:")
        c_.drawRightString(190 * mm, (236 - 4.6 * i) * mm, val)
    pos = [("Stoßfänger vorne inkl. Lackierung", 98450), ("Scheinwerfer rechts (LED)", 71230), ("Kühler + Kältemittel", 48890),
           ("Arbeitslohn Karosserie 58 AW", 55100), ("Verbringung Lackiererei", 9500), ("Kleinteile/Clips", 3595)]
    y = 196
    c_.setFont("Helvetica", 9)
    for i, (t_, a_) in enumerate(pos, 1):
        c_.drawString(20 * mm, y * mm, f"{i}  {t_}")
        c_.drawRightString(190 * mm, y * mm, de_amt(a_))
        y -= 6
    net = sum(a_ for _, a_ in pos)
    for lab, val in (("Summe netto", net), ("zzgl. 19 % USt.", vat_of(net)), ("Gesamt (voraussichtlich)", net + vat_of(net))):
        c_.drawString(120 * mm, y * mm, lab)
        c_.drawRightString(190 * mm, y * mm, de_amt(val) + " €")
        y -= 5.5
    c_.drawString(20 * mm, (y - 8) * mm, "Dies ist keine Rechnung. Kostenvoranschlag unverbindlich, gültig 4 Wochen.")
    c_.showPage()
    c_.save()
    return net + vat_of(net)


def pdf_supplier(path, ti):
    c_ = canvas.Canvas(path, pagesize=A4, invariant=1)
    c_.setTitle(f"Rechnung {ti['nr']}")
    c_.setFont("Helvetica-Bold", 16)
    c_.drawString(20 * mm, 280 * mm, "TEILE-EXPRESS Westfalen GmbH")
    c_.setFont("Helvetica", 8.5)
    c_.drawString(20 * mm, 275 * mm, "Kfz-Teile-Großhandel · Am Hafen 31 · 33102 Paderborn · www.teile-express-westfalen.de")
    c_.setFont("Helvetica", 10)
    for i, a_ in enumerate(["Autoservice Lindner GmbH", "Erwitter Straße 112", "59557 Lippstadt"]):
        c_.drawString(20 * mm, (246 - 5 * i) * mm, a_)
    c_.setFont("Helvetica-Bold", 13)
    c_.drawString(20 * mm, 210 * mm, f"Sammelrechnung Nr. {ti['nr']}")
    c_.setFont("Helvetica", 9.5)
    for i, (k_, val) in enumerate([("Rechnungsdatum", dde(ti["date"])), ("Kunden-Nr.", "47110"),
                                   ("Lieferzeitraum", f"{dde(ti['start'])} - {dde(ti['end'])}")]):
        c_.drawString(120 * mm, (236 - 4.6 * i) * mm, f"{k_}:")
        c_.drawRightString(190 * mm, (236 - 4.6 * i) * mm, val)
    rr = random.Random(ti["nr"])
    arts = ["Bremsscheibe VA", "Bremsbelagsatz VA", "Ölfilter", "Luftfilter", "Innenraumfilter", "Motoröl 5W-30 60 l Fass",
            "Zahnriemensatz m. WaPu", "Batterie 70Ah AGM", "Stoßdämpfer VA", "Wischerblatt-Satz", "Kältemittel R1234yf"]
    remaining = ti["net"]
    y = 196
    c_.setFont("Helvetica", 9)
    for i, a_ in enumerate(arts, 1):
        part = remaining if i == len(arts) else int(remaining * rr.uniform(0.06, 0.16))
        remaining -= part
        c_.drawString(20 * mm, y * mm, f"{i:>2}  {a_}")
        c_.drawRightString(190 * mm, y * mm, de_amt(part))
        y -= 5.5
    for lab, val in (("Warenwert netto", ti["net"]), ("19 % USt.", ti["gross"] - ti["net"]), ("Rechnungsbetrag", ti["gross"])):
        c_.drawString(120 * mm, y * mm, lab)
        c_.drawRightString(190 * mm, y * mm, de_amt(val) + " €")
        y -= 5.5
    c_.drawString(20 * mm, (y - 8) * mm, f"Der Rechnungsbetrag wird am {dde(ti['debit'])} per SEPA-Lastschrift von Ihrem Konto eingezogen.")
    c_.showPage()
    c_.save()


def pdf_avis(path):
    xs = [KEY["SAV1"], KEY["SAV2"], KEY["SAV3"]]
    c_ = canvas.Canvas(path, pagesize=A4, invariant=1)
    c_.setTitle("Zahlungsavis")
    c_.setFont("Helvetica-Bold", 13)
    c_.drawString(20 * mm, 280 * mm, "Ambulanter Pflegedienst Sonnenschein GmbH")
    c_.setFont("Helvetica", 9)
    c_.drawString(20 * mm, 275 * mm, "Cappelstraße 41 · 59555 Lippstadt · Buchhaltung")
    c_.setFont("Helvetica", 10)
    for i, a_ in enumerate(["Autoservice Lindner GmbH", "Erwitter Straße 112", "59557 Lippstadt"]):
        c_.drawString(20 * mm, (250 - 5 * i) * mm, a_)
    c_.setFont("Helvetica-Bold", 13)
    c_.drawString(20 * mm, 220 * mm, "ZAHLUNGSAVIS vom 07.10.2026")
    c_.setFont("Helvetica", 9.5)
    c_.drawString(20 * mm, 212 * mm, "Unsere Lieferanten-Nr. bei Ihnen: 10412     Zahlweg: Überweisung, Ausführung 08.10.2026")
    y = 200
    c_.setFont("Helvetica-Bold", 9)
    for x_, lab in ((20, "Beleg-Nr."), (50, "Belegdatum"), (85, "Kennzeichen")):
        c_.drawString(x_ * mm, y * mm, lab)
    for x_, lab in ((135, "Betrag"), (160, "Skonto"), (190, "Zahlbetrag")):
        c_.drawRightString(x_ * mm, y * mm, lab)
    c_.setFont("Helvetica", 9)
    tot = 0
    for v in xs:
        y -= 6
        sk = (v["gross"] * 2 + 50) // 100
        tot += v["gross"] - sk
        c_.drawString(20 * mm, y * mm, str(v["nr"]))
        c_.drawString(50 * mm, y * mm, dde(v["date"]))
        c_.drawString(85 * mm, y * mm, v["plate"])
        c_.drawRightString(135 * mm, y * mm, de_amt(v["gross"]))
        c_.drawRightString(160 * mm, y * mm, de_amt(sk))
        c_.drawRightString(190 * mm, y * mm, de_amt(v["gross"] - sk))
    y -= 8
    c_.setFont("Helvetica-Bold", 10)
    c_.drawString(20 * mm, y * mm, "Gesamt")
    c_.drawRightString(190 * mm, y * mm, de_amt(tot) + " EUR")
    c_.showPage()
    c_.save()
    return tot


# ---------- Mails ----------
TZ_S = dt.timezone(dt.timedelta(hours=2))
TZ_W = dt.timezone(dt.timedelta(hours=1))
OWN = "Autoservice Lindner – Petra Lindner <buero@autoservice-lindner.de>"


def tstamp(y, m, d_, hh, mi):
    tz = TZ_W if D(y, m, d_) >= D(2026, 10, 25) else TZ_S
    return dt.datetime(y, m, d_, hh, mi, 0, tzinfo=tz)


def eml(path, frm, to, subject, when, body, attachments=(), reply_to_id=None):
    m = EmailMessage(policy=policy.SMTP)
    m["From"] = frm
    m["To"] = to
    m["Subject"] = subject
    m["Date"] = format_datetime(when)
    mid = hashlib.sha1((subject + str(when)).encode()).hexdigest()[:20]
    m["Message-ID"] = f"<{mid}@{frm.split('@')[-1].rstrip('>')}>"
    if reply_to_id:
        m["In-Reply-To"] = reply_to_id
    m["MIME-Version"] = "1.0"
    m.set_content(body, subtype="plain", charset="utf-8", cte="quoted-printable")
    for fn, data in attachments:
        m.add_attachment(data, maintype="application", subtype="pdf", filename=fn)
    if m.is_multipart():
        m.set_boundary("----=_Part_" + mid[:12])
    with open(path, "wb") as f:
        f.write(m.as_bytes())


SIG = "\n\nMit freundlichen Grüßen\nPetra Lindner\n\nAutoservice Lindner GmbH\nKfz-Meisterbetrieb\nErwitter Straße 112 · 59557 Lippstadt\nTel. 02941 18220"


def write_mails_l1(md, pdfs):
    k = KEY
    a, b = k["SONN_A"], k["SONN_B"]
    sa, sb = (a["gross"] * 2 + 50) // 100, (b["gross"] * 2 + 50) // 100
    eml(os.path.join(md, "Zahlungsavis 30.07.2026.eml"),
        "Buchhaltung Pflegedienst Sonnenschein <buchhaltung@pflegedienst-sonnenschein.de>",
        "buero@autoservice-lindner.de", "Zahlungsavis 30.07.2026", tstamp(2026, 7, 30, 9, 12),
        "Sehr geehrte Damen und Herren,\n\nhiermit avisieren wir Ihnen folgende Zahlung:\n\n"
        f"Rechnung {a['nr']} vom {dde(a['date'])}   {de_amt(a['gross'])} EUR   abzgl. 2 % Skonto {de_amt(sa)} EUR\n"
        f"Rechnung {b['nr']} vom {dde(b['date'])}   {de_amt(b['gross'])} EUR   abzgl. 2 % Skonto {de_amt(sb)} EUR\n\n"
        f"Zahlbetrag gesamt: {de_amt(a['gross'] - sa + b['gross'] - sb)} EUR\n\nDie Überweisung wird heute ausgeführt.\n\n"
        "Mit freundlichen Grüßen\ni.A. Sandra Overhoff\nBuchhaltung\nAmbulanter Pflegedienst Sonnenschein GmbH\n"
        "Cappelstraße 41 · 59555 Lippstadt\nTel. 02941 77880")
    y = k["YILMAZ"]
    eml(os.path.join(md, f"AW Ihre Rechnung {y['nr']}.eml"), "Mehmet Yilmaz <mehmet.yilmaz@taxi-yilmaz.de>",
        "buero@autoservice-lindner.de", f"AW: Ihre Rechnung {y['nr']}", tstamp(2026, 8, 2, 21, 47),
        "Hallo Frau Lindner,\n\nich kann die Rechnung über 1.845,20 gerade nicht auf einmal bezahlen, die Sommerferien "
        "sind ganz schlecht fürs Taxi. Ich überweise erstmal 600 Euro und den Rest in den nächsten Wochen, "
        "versprochen.\n\nGruß Mehmet Yilmaz\nTaxi Yilmaz Lippstadt\n\n"
        f"> Am 13.07.2026 um 16:40 schrieb Autoservice Lindner <buero@autoservice-lindner.de>:\n> Sehr geehrter Herr Yilmaz,\n"
        f"> anbei erhalten Sie unsere Rechnung {y['nr']} für Ihr Fahrzeug LP-TY 300.\n> Mit freundlichen Grüßen\n> Petra Lindner",
        attachments=[(f"RE_{y['nr']}.pdf", pdfs[y["nr"]])])
    v = k["VOGT"]
    eml(os.path.join(md, "Schaden KH 26-118274 Vogt.eml"), "NordWest Assekuranz Schaden Kraftfahrt <kfz-schaden@nordwest-assekuranz.de>",
        "buero@autoservice-lindner.de", f"Schaden-Nr. KH 26-118274 / Ihre Rechnung {v['nr']} / Anspruchstellerin Karin Vogt",
        tstamp(2026, 8, 26, 11, 3),
        "Sehr geehrte Damen und Herren,\n\nzu o.g. Schaden haben wir Ihre Reparaturrechnung Nr. "
        f"{v['nr']} vom {dde(v['date'])} über {de_amt(v['gross'])} EUR geprüft.\n\n"
        "Wir regulieren einen Betrag von 2.150,00 EUR. Die berechneten Verbringungskosten zur Lackiererei (95,00 EUR netto) "
        "waren nach Prüfung durch unseren Sachverständigen nicht erforderlich und werden nicht übernommen.\n\n"
        "Die Zahlung erfolgt in den nächsten Tagen auf das in der Rechnung genannte Konto.\n\n"
        "Mit freundlichen Grüßen\nNordWest Assekuranz AG\nSchadenservice Kraftfahrt\n\nDiese E-Mail wurde maschinell erstellt.")
    p_ = k["POHL"]
    eml(os.path.join(md, f"Zahlungserinnerung Rechnung {p_['nr']}.eml"), OWN, "Dieter Pohl <d.pohl@t-online.de>",
        f"Zahlungserinnerung Rechnung {p_['nr']}", tstamp(2026, 8, 20, 10, 15),
        "Sehr geehrter Herr Pohl,\n\nbei Durchsicht unserer Buchhaltung haben wir festgestellt, dass unsere Rechnung "
        f"Nr. {p_['nr']} vom {dde(p_['date'])} über {de_amt(p_['gross'])} EUR (Kupplung Ford Focus SO-DP 640) noch nicht "
        "beglichen ist.\n\nWir bitten Sie, den Betrag bis zum 03.09.2026 auf unser Konto bei der Deutschen Bank zu überweisen. "
        "Sollten Sie die Zahlung inzwischen veranlasst haben, betrachten Sie dieses Schreiben bitte als gegenstandslos." + SIG)
    ti = [t for t in TEILE_INV if t["date"] == D(2026, 9, 1)][0]
    eml(os.path.join(md, f"Ihre Rechnung {ti['nr']} Teile-Express.eml"), "Teile-Express Westfalen <rechnung@teile-express-westfalen.de>",
        "buero@autoservice-lindner.de", f"Ihre Rechnung {ti['nr']} / Kd-Nr. 47110", tstamp(2026, 9, 1, 6, 2),
        f"Sehr geehrte Kundin, sehr geehrter Kunde,\n\nanbei erhalten Sie Ihre Sammelrechnung Nr. {ti['nr']} für den "
        f"Lieferzeitraum {dde(ti['start'])} bis {dde(ti['end'])}.\nDer Betrag von {de_amt(ti['gross'])} EUR wird am "
        f"{dde(ti['debit'])} per SEPA-Lastschrift eingezogen.\n\nIhr Teile-Express-Team\n\n-- \nTeile-Express Westfalen GmbH · "
        "Am Hafen 31 · 33102 Paderborn · HRB 4410 Paderborn", attachments=[(f"RE{ti['nr']}.pdf", pdfs["teile"])])
    kb = k["KAISER_B"]
    eml(os.path.join(md, f"Rechnung {kb['nr']} Elektro Kaiser.eml"), OWN, "Elektro Kaiser Buchhaltung <info@elektro-kaiser-lp.de>",
        f"Rechnung {kb['nr']} – Fahrzeug {kb['plate']}", tstamp(2026, 7, 24, 15, 31),
        f"Sehr geehrte Damen und Herren,\n\nanbei erhalten Sie unsere Rechnung {kb['nr']} für Ihr Fahrzeug {kb['plate']}."
        f"\n\nDie Rechnung {k['KAISER_A']['nr']} vom {dde(k['KAISER_A']['date'])} haben Sie bereits per Post erhalten." + SIG,
        attachments=[(f"Rechnung {kb['nr']} Kaiser.pdf", pdfs[kb["nr"]])])
    eml(os.path.join(md, "Termin Klimaservice Sprinter.eml"), "Ralf Brune <ralf@getraenke-brune.de>", "buero@autoservice-lindner.de",
        "Termin Klimaservice Sprinter", tstamp(2026, 9, 8, 7, 55),
        "Moin Jürgen,\n\nkönnen wir den Sprinter LP-GB 77 morgen früh um 7 bringen? Die Klima kühlt nicht mehr richtig, "
        "die Fahrer meckern. Mittags brauchen wir ihn wieder.\n\nRechnung wie immer an die Firma.\n\nGruß Ralf\n\n"
        "Getränke Brune e.K.\nWesternkötter Straße 120 · 59558 Lippstadt")
    be = k["BECKER_ERW"]
    eml(os.path.join(md, "Ihre Rechnung - falsche Adresse.eml"), "Thomas Becker <tbecker.erwitte@gmail.com>",
        "buero@autoservice-lindner.de", "Ihre Rechnung - falsche Adresse", tstamp(2026, 8, 21, 19, 20),
        f"Guten Abend,\n\nich habe heute die Rechnung {be['nr']} für meinen Touran (SO-TB 77) überwiesen. Auf der Rechnung "
        "steht allerdings die Adresse Mastholter Straße in Lippstadt – das bin ich nicht, ich wohne Am Stadtgraben 3 in "
        "Erwitte. Ich glaube, Sie haben da noch einen anderen Thomas Becker in der Kartei. Bitte korrigieren Sie das, ich "
        "brauche die Rechnung für meinen Arbeitgeber (Dienstwagenregelung).\n\nViele Grüße\nThomas Becker")


def write_mails_l2(md, pdfs):
    k = KEY
    h = k["HAGEMANN"]
    eml(os.path.join(md, f"Rechnung {h['nr']} zu hoch.eml"), "Bernd Hagemann <b.hagemann@web.de>", "buero@autoservice-lindner.de",
        f"Rechnung {h['nr']}", tstamp(2026, 9, 16, 18, 2),
        "Hallo Frau Lindner,\n\nIhr Mann hatte mir bei der Annahme ca. 450 Euro netto gesagt, jetzt sind es 500 netto bzw. "
        "595,00 Euro. Ich bitte um Prüfung, bevor ich überweise.\n\nGruß\nBernd Hagemann")
    eml(os.path.join(md, f"AW Rechnung {h['nr']} - korrigierte Rechnung.eml"), OWN, "Bernd Hagemann <b.hagemann@web.de>",
        f"AW: Rechnung {h['nr']} - korrigierte Rechnung", tstamp(2026, 9, 18, 9, 40),
        "Sehr geehrter Herr Hagemann,\n\nSie haben recht, mein Mann hat Ihnen 10 % Nachlass zugesagt. Anbei die korrigierte "
        f"Rechnung {h['nr']} über 535,50 EUR. Die alte Rechnung über 595,00 EUR ist damit hinfällig." + SIG,
        attachments=[(f"Rechnung_{h['nr']}_korrigiert.pdf", pdfs[("hag", h["nr"])])])
    kr = k["KRAEMER"]
    eml(os.path.join(md, "Rechnungsadresse LP-SK 912.eml"), "Westfalen Leasing Service <service@westfalen-leasing.de>",
        "buero@autoservice-lindner.de", f"Ihre Rechnung {kr['nr']} - Fahrzeug LP-SK 912 - Vertrag WL-55123",
        tstamp(2026, 9, 21, 14, 12),
        "Sehr geehrte Damen und Herren,\n\nunser Leasingnehmer Herr Sven Krämer hat uns Ihre Rechnung "
        f"{kr['nr']} vom {dde(kr['date'])} weitergeleitet. Wartungs- und Verschleißreparaturen sind bei diesem Vertrag "
        "Bestandteil des Servicepakets. Bitte stornieren Sie die Rechnung an den Fahrer und stellen Sie sie neu auf uns aus:\n\n"
        "Westfalen Leasing GmbH\nAbt. Service/Wartung\nWarendorfer Straße 210\n48145 Münster\n\n"
        "Bitte geben Sie die Vertragsnummer WL-55123 an. Zahlungsziel 30 Tage.\n\nMit freundlichen Grüßen\n"
        "Westfalen Leasing GmbH – Team Service")
    br = k["BRINKMANN"]
    eml(os.path.join(md, "doppelt bezahlt.eml"), "Ute Brinkmann <ute.brinkmann@web.de>", "buero@autoservice-lindner.de",
        "doppelt bezahlt", tstamp(2026, 9, 29, 20, 31),
        f"Hallo,\n\nich habe gemerkt, dass ich die Rechnung {br['nr']} (TÜV + Bremsen, 312,40) zweimal bezahlt habe: einmal "
        "mit Karte bei Abholung und dann nochmal überwiesen, weil die Rechnung mit der Post kam. Können Sie mir das "
        "Geld bitte zurücküberweisen? IBAN wie bei der Überweisung.\n\nDanke und Gruß\nUte Brinkmann")
    s = k["SCHULZ"]
    eml(os.path.join(md, f"2. Mahnung Rechnung {s['nr']}.eml"), OWN, "Kevin Schulz <kevin.s1999@gmail.com>",
        f"2. Mahnung Rechnung {s['nr']}", tstamp(2026, 10, 1, 11, 5),
        "Sehr geehrter Herr Schulz,\n\ntrotz unserer Erinnerung vom 14.08. und der Mahnung vom 28.08. ist der Betrag aus "
        f"Rechnung {s['nr']} vom {dde(s['date'])} (4 Reifen, 456,70 EUR) nach der Rücklastschrift weiterhin offen. "
        "Zuzüglich Rücklastschriftgebühr von 3,00 EUR bitten wir um Zahlung von 459,70 EUR bis 12.10.2026. "
        "Danach müssen wir die Forderung an unser Inkassobüro abgeben." + SIG)
    tot = pdfs["avis_total"]
    eml(os.path.join(md, "Zahlungsavis 07.10.2026.eml"),
        "Buchhaltung Pflegedienst Sonnenschein <buchhaltung@pflegedienst-sonnenschein.de>", "buero@autoservice-lindner.de",
        "Zahlungsavis", tstamp(2026, 10, 7, 16, 44),
        f"Sehr geehrte Damen und Herren,\n\nim Anhang erhalten Sie unser Zahlungsavis. Der Gesamtbetrag von {de_amt(tot)} EUR "
        "wird morgen überwiesen.\n\nMit freundlichen Grüßen\ni.A. Sandra Overhoff\nBuchhaltung",
        attachments=[("Avis_10412_20261007.pdf", pdfs["avis"])])
    eml(os.path.join(md, "Winterreifen Seat Ateca.eml"), "Cem Arslan <cem.arslan@gmail.com>", "buero@autoservice-lindner.de",
        "Winterreifen Seat Ateca", tstamp(2026, 10, 19, 12, 9),
        "Hallo Herr Lindner,\n\nwie am Telefon besprochen bestelle ich 4 Winterreifen Continental 215/55 R17 auf Stahlfelgen "
        "für meinen Seat Ateca LP-CA 344. Ich überweise Ihnen morgen 250 Euro Anzahlung. Termin zum Montieren dann "
        "Anfang November?\n\nMit freundlichen Grüßen\nCem Arslan\nBeckumer Straße 12, Lippstadt")
    la = k["LANGE"]
    eml(os.path.join(md, "Schaden KH 26-120551 Lange.eml"), "NordWest Assekuranz Schaden Kraftfahrt <kfz-schaden@nordwest-assekuranz.de>",
        "buero@autoservice-lindner.de", f"Schaden-Nr. KH 26-120551 / Ihre Rechnung {la['nr']}", tstamp(2026, 10, 13, 10, 30),
        f"Sehr geehrte Damen und Herren,\n\nIhre Rechnung {la['nr']} vom {dde(la['date'])} über {de_amt(la['gross'])} EUR "
        "wird in voller Höhe reguliert. Die Zahlung ist angewiesen.\n\nMit freundlichen Grüßen\nNordWest Assekuranz AG\n"
        "Schadenservice Kraftfahrt")


def generate(outdir):
    if os.path.exists(outdir):
        for sub in ("lieferung1", "lieferung2"):
            p = os.path.join(outdir, sub)
            if os.path.exists(p):
                shutil.rmtree(p)
    os.makedirs(outdir, exist_ok=True)
    tmp = os.path.join(outdir, "_pdf_cache")
    os.makedirs(tmp, exist_ok=True)
    pdfs = {}
    pdf_files = {1: [], 2: []}

    def mkpdf(v, dl, name=None):
        fn = (name or v["pdf"]).replace("{nr}", str(v["nr"]))
        path = os.path.join(tmp, f"{dl}_{fn}")
        pdf_invoice(path, v, dl)
        with open(path, "rb") as f:
            data = f.read()
        return fn, data

    for v in sorted(INVS, key=lambda x: x["nr"]):
        if not v["pdf"]:
            continue
        dl = v["first_dl"]
        fn, data = mkpdf(v, dl)
        pdfs[v["nr"]] = data
        pdf_files[dl].append((fn, data))
    h = KEY["HAGEMANN"]
    fn, data = mkpdf(h, 2, "Rechnung_{nr}_korrigiert.pdf")
    pdfs[("hag", h["nr"])] = data
    pdf_files[2].append((fn, data))
    kva_path = os.path.join(tmp, "kva.pdf")
    kva_total = pdf_kva(kva_path)
    assert kva_total == KEY["LANGE"]["gross"], (kva_total, KEY["LANGE"]["gross"])
    pdf_files[1].append(("KVA_0892_Lange_GT-ML22.pdf", open(kva_path, "rb").read()))
    ti = [t for t in TEILE_INV if t["date"] == D(2026, 9, 1)][0]
    sp = os.path.join(tmp, "teile.pdf")
    pdf_supplier(sp, ti)
    pdfs["teile"] = open(sp, "rb").read()
    pdf_files[1].append((f"Teile-Express RE{ti['nr']}.pdf", pdfs["teile"]))
    ap = os.path.join(tmp, "avis.pdf")
    pdfs["avis_total"] = pdf_avis(ap)
    pdfs["avis"] = open(ap, "rb").read()

    for dl in (1, 2):
        base = os.path.join(outdir, f"lieferung{dl}")
        os.makedirs(os.path.join(base, "Mails"), exist_ok=True)
        os.makedirs(os.path.join(base, "Rechnungen PDF"), exist_ok=True)
        write_book(os.path.join(base, "Rechnungsbuch 2026.xlsx"), dl)
        write_jobs(os.path.join(base, "Auftragsbuch Werkstatt.xlsx"), dl)
        write_customers(os.path.join(base, "Kunden_Export.csv"), dl)
        write_bank(os.path.join(base, "Umsaetze_Geschaeftskonto.csv"), dl)
        for d_ in range(1, dl + 1):
            for fn, data in pdf_files[d_]:
                with open(os.path.join(base, "Rechnungen PDF", fn), "wb") as f:
                    f.write(data)
        write_mails_l1(os.path.join(base, "Mails"), pdfs)
        if dl == 2:
            write_mails_l2(os.path.join(base, "Mails"), pdfs)
            with open(os.path.join(base, "Zahlungsavis Sonnenschein 07.10.2026.pdf"), "wb") as f:
                f.write(pdfs["avis"])
            with open(os.path.join(base, "~$Rechnungsbuch 2026.xlsx"), "wb") as f:
                f.write(b"\x0dPetra Lindner" + b"\x20" * 40 + "Petra Lindner".encode("utf-16-le") + b"\x00" * 60)
    shutil.rmtree(tmp)

    truth = {"deliveries": [compute_truth(1), compute_truth(2)]}
    with open(os.path.join(outdir, "truth.json"), "w", encoding="utf-8") as f:
        json.dump(truth, f, ensure_ascii=False, indent=1)
    write_ledger(os.path.join(outdir, "ledger.json"))
    write_key(os.path.join(outdir, "answer_key.txt"))
    return truth


def write_ledger(path):
    def dd(x):
        return x.isoformat() if x else None
    led = {"cut": {k: dd(v) for k, v in CUT.items()}, "today": {k: dd(v) for k, v in TODAY.items()},
           "bank_from": {k: dd(v) for k, v in BANK_FROM.items()}, "check": {k: dd(v) for k, v in CHECK.items()},
           "invoices": [], "jobs": [], "payments": [], "bank": []}
    for v in sorted(INVS, key=lambda x: x["nr"]):
        led["invoices"].append(dict(nr=v["nr"], key=v["key"], date=dd(v["date"]), kind=v["kind"],
                                    ref=v["ref"]["nr"] if v["ref"] else None, job=v["job"]["nr"] if v["job"] else None,
                                    zahlart=v["zahlart"], cust=v["kd"],
                                    gross={dl: g(v, "gross", dl) for dl in (1, 2)},
                                    due={dl: dd(g(v, "due", dl)) for dl in (1, 2)},
                                    book_kd={dl: g(v, "kd", dl) for dl in (1, 2)}))
    for j in sorted(JOBS, key=lambda x: x["nr"]):
        led["jobs"].append(dict(nr=j["nr"], key=j["key"], annahme=dd(j["annahme"]), fertig=dd(j["fertig"]),
                                special=j["special"], kd=j["kd"], invs=[v["nr"] for v in j["invs"]],
                                state={dl: (lambda s_: None if s_ is None else dict(status=s_["status"], rg=s_["rg"],
                                                                                     fertig=dd(s_["fertig"])))(job_state(j, dl))
                                       for dl in (1, 2)}))
    for p in sorted(PAYS, key=lambda x: (x["date"], x["amount"])):
        led["payments"].append(dict(date=dd(p["date"]), amount=p["amount"], channel=p["channel"], kind=p["kind"],
                                    payer=p["payer"], text=p["text"],
                                    allocs=[[v["nr"], a, dd(e)] for v, a, e in p["allocs"]],
                                    skontos=[[v["nr"], a] for v, a in p["skontos"]]))
    for b in BANKLINES:
        led["bank"].append(dict(date=dd(b["date"]), amount=b["amount"], name=b["name"], text=b["text"],
                                src=("pay" if isinstance(b["src"], dict) else b["src"])))
    with open(path, "w", encoding="utf-8") as f:
        json.dump(led, f, ensure_ascii=False, indent=0)


def write_key(path):
    lines = []
    for k in sorted(KEY):
        o = KEY[k]
        if "invs" in o:
            lines.append(f"{k:14s} Auftrag {o['nr']}  {o['kunde']}  {o['arbeiten'][:50]}")
        else:
            lines.append(f"{k:14s} Rg {o['nr']}  {dde(o['date'])}  {o['kunde']}  {de_amt(o['gross'])}")
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines) + "\n")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("Aufruf: python3 gen_firm.py OUTDIR")
        sys.exit(2)
    t = generate(sys.argv[1])
    for dlv in t["deliveries"]:
        print(dlv["folder"], "offen:", len(dlv["open"]), "Summe", round(sum(dlv["open"].values()), 2), "ueberfaellig:",
              len(dlv["overdue"]), "unbilled:", dlv["unbilled"], "receipts:", dlv["receipts"])
