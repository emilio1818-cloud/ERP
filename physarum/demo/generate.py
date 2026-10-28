"""Generate a deliberately messy, realistic data dump of a fictional small company.

Röstwerk Hafenbohne GmbH is a (fictional) coffee roastery in Hamburg with B2B customers
(cafés, hotels, offices), an online shop, green-coffee suppliers and a bank account.
The files mimic what such a company really has lying around: an old Excel customer
list with a title row, a CRM export in English, a price list with section rows, a flat
order workbook split by year, a Shopify-style shop export, an accounting journal in
cp1252 with decimal commas and a sum row, a Sparkasse bank statement, a supplier text
file, e-mails and a PDF supplier invoice.

Usage:  python -m physarum.demo.generate OUT_DIR [--growth]
"""
from __future__ import annotations

import csv
import io
import random
import sys
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from email.message import EmailMessage
from email.utils import format_datetime
from pathlib import Path

import openpyxl

SEED = 1818

FIRST = ["Anna", "Jonas", "Lea", "Finn", "Mia", "Paul", "Emma", "Ben", "Lina", "Tom", "Clara", "Jan",
         "Sophie", "Lukas", "Marie", "Felix", "Hannah", "Moritz", "Lena", "Nils", "Frieda", "Ole", "Greta", "Hauke"]
LAST = ["Brandt", "Petersen", "Hansen", "Jensen", "Müller", "Schröder", "Möller", "Koch", "Wagner", "Becker",
        "Lorenzen", "Thomsen", "Carstens", "Rickmers", "Behrens", "Albers", "Kruse", "Witt", "Martens", "Stüve"]
PLACE = ["Hafenblick", "Elbstrand", "Möwennest", "Alsterperle", "Fleetenkieker", "Landungsbrücke", "Kranhaus",
         "Deichgraf", "Ankerplatz", "Leuchtturm", "Kehrwieder", "Kompass", "Seemannsruh", "Nordlicht", "Backbord",
         "Steuerbord", "Kombüse", "Hafenkante", "Fährhaus", "Lotsenhaus", "Wellenreiter", "Krabbenkutter",
         "Elbinsel", "Pier Sieben", "Schanzenecke", "Ottensener Hof", "Kaispeicher", "Brückenwirt", "Zollhaus",
         "Tidenhub", "Seewind", "Fischerhaus", "Elbdeich", "Hafenmole", "Klabautermann", "Strandkorb"]
KIND = [("Café", "Café"), ("Kaffeebar", "Café"), ("Bistro", "Restaurant"), ("Restaurant", "Restaurant"),
        ("Hotel", "Hotel"), ("Bäckerei", "Einzelhandel"), ("Feinkost", "Einzelhandel"), ("Büro", "Büro"),
        ("Kanzlei", "Büro"), ("Coworking", "Büro")]
LEGAL = ["", "", "", " GmbH", " GmbH", " UG", " KG", " e.K.", " GmbH & Co. KG"]
CITIES = [("20457", "Hamburg"), ("20459", "Hamburg"), ("22765", "Hamburg"), ("22767", "Hamburg"),
          ("20354", "Hamburg"), ("22085", "Hamburg"), ("20249", "Hamburg"), ("21073", "Hamburg"),
          ("22880", "Wedel"), ("21465", "Reinbek"), ("25421", "Pinneberg"), ("21614", "Buxtehude"),
          ("22926", "Ahrensburg"), ("21335", "Lüneburg")]
STREETS = ["Hafenstraße", "Elbchaussee", "Große Elbstraße", "Kehrwieder", "Am Sandtorkai", "Ottenser Hauptstraße",
           "Schulterblatt", "Lange Reihe", "Mühlenkamp", "Bahrenfelder Straße", "Holstenstraße", "Neuer Wall",
           "Bergstraße", "Lindenallee", "Am Markt", "Rathausstraße", "Bahnhofstraße", "Deichstraße"]
PRIVATE_FIRST = FIRST + ["Karl", "Ida", "Theo", "Ella", "Henri", "Mila", "Anton", "Romy", "Fiete", "Smilla"]
PRIVATE_LAST = LAST + ["Schulz", "Hoffmann", "Fischer", "Weber", "Meyer", "Schmidt", "Neumann", "Krüger",
                       "Lange", "Vogt", "Busch", "Otto", "Sommer", "Winter", "Kühl", "Engel"]
SHOP_CITIES = ["Hamburg", "Berlin", "München", "Köln", "Bremen", "Kiel", "Hannover", "Leipzig", "Lübeck",
               "Frankfurt am Main", "Münster", "Rostock", "Dresden", "Freiburg"]


@dataclass
class Article:
    sku: str
    name: str
    group: str
    variety: str
    roast: str
    pack: str
    price: float  # net
    vat: float


ARTICLES = [
    Article("HB-E01", "Hafenmischung Espresso 1 kg", "ESPRESSO", "Brasil/Indien", "dunkel", "1 kg", 24.90, 0.07),
    Article("HB-E02", "Hafenmischung Espresso 250 g", "ESPRESSO", "Brasil/Indien", "dunkel", "250 g", 7.50, 0.07),
    Article("HB-E03", "Kaispeicher Crema 1 kg", "ESPRESSO", "Brasil/Äthiopien", "mittel", "1 kg", 26.50, 0.07),
    Article("HB-E04", "Kaispeicher Crema 250 g", "ESPRESSO", "Brasil/Äthiopien", "mittel", "250 g", 7.90, 0.07),
    Article("HB-E05", "Klabautermann Decaf 1 kg", "ESPRESSO", "Kolumbien (entkoff.)", "mittel", "1 kg", 27.90, 0.07),
    Article("HB-E06", "Klabautermann Decaf 250 g", "ESPRESSO", "Kolumbien (entkoff.)", "mittel", "250 g", 8.40, 0.07),
    Article("HB-E07", "Lotsen-Espresso Single Origin 1 kg", "ESPRESSO", "Guatemala", "mittel", "1 kg", 31.00, 0.07),
    Article("HB-E08", "Lotsen-Espresso Single Origin 250 g", "ESPRESSO", "Guatemala", "mittel", "250 g", 9.20, 0.07),
    Article("HB-F01", "Elbsegler Filter 250 g", "FILTERKAFFEE", "Äthiopien Sidamo", "hell", "250 g", 8.50, 0.07),
    Article("HB-F02", "Elbsegler Filter 1 kg", "FILTERKAFFEE", "Äthiopien Sidamo", "hell", "1 kg", 28.90, 0.07),
    Article("HB-F03", "Nordlicht Filter 250 g", "FILTERKAFFEE", "Kenia AA", "hell", "250 g", 9.90, 0.07),
    Article("HB-F04", "Nordlicht Filter 1 kg", "FILTERKAFFEE", "Kenia AA", "hell", "1 kg", 33.50, 0.07),
    Article("HB-F05", "Fährhaus Filter Hausmarke 1 kg", "FILTERKAFFEE", "Brasil/Peru", "mittel", "1 kg", 21.90, 0.07),
    Article("HB-F06", "Fährhaus Filter Hausmarke 500 g", "FILTERKAFFEE", "Brasil/Peru", "mittel", "500 g", 11.90, 0.07),
    Article("HB-F07", "Tidenhub Kolumbien 250 g", "FILTERKAFFEE", "Kolumbien Huila", "hell", "250 g", 8.90, 0.07),
    Article("HB-F08", "Tidenhub Kolumbien 1 kg", "FILTERKAFFEE", "Kolumbien Huila", "hell", "1 kg", 30.50, 0.07),
    Article("HB-F09", "Seewind Peru Bio 250 g", "FILTERKAFFEE", "Peru (Bio)", "mittel", "250 g", 8.70, 0.07),
    Article("HB-F10", "Seewind Peru Bio 1 kg", "FILTERKAFFEE", "Peru (Bio)", "mittel", "1 kg", 29.90, 0.07),
    Article("HB-B01", "Büro-Bohne 1 kg", "FILTERKAFFEE", "Brasil", "mittel", "1 kg", 18.90, 0.07),
    Article("HB-B02", "Büro-Bohne 3 kg Sack", "FILTERKAFFEE", "Brasil", "mittel", "3 kg", 52.00, 0.07),
    Article("HB-Z01", "Filtertüten Gr. 4 (100 Stk.)", "ZUBEHÖR", "", "", "100 Stk.", 2.90, 0.19),
    Article("HB-Z02", "Handmühle Edelstahl", "ZUBEHÖR", "", "", "Stk.", 39.00, 0.19),
    Article("HB-Z03", "Glaskanne 1 l", "ZUBEHÖR", "", "", "Stk.", 24.50, 0.19),
    Article("HB-Z04", "Porzellanfilter Gr. 4", "ZUBEHÖR", "", "", "Stk.", 16.90, 0.19),
    Article("HB-Z05", "Geschenkbox Hafenbohne", "ZUBEHÖR", "", "", "Stk.", 29.00, 0.19),
    Article("HB-Z06", "Barista-Tamper 58 mm", "ZUBEHÖR", "", "", "Stk.", 34.90, 0.19),
    Article("HB-Z07", "Milchkännchen 0,6 l", "ZUBEHÖR", "", "", "Stk.", 14.50, 0.19),
    Article("HB-Z08", "Thermobecher Hafenbohne", "ZUBEHÖR", "", "", "Stk.", 19.90, 0.19),
]
B2B_SKUS = ["HB-E01", "HB-E03", "HB-E05", "HB-E07", "HB-F02", "HB-F04", "HB-F05", "HB-F08", "HB-F10", "HB-B01", "HB-B02"]
SHOP_SKUS = [a.sku for a in ARTICLES if a.sku not in ("HB-B02",)]


@dataclass
class Customer:
    no: str
    name: str
    contact: str
    street: str
    plz: str
    city: str
    phone: str
    email: str
    terms: int
    group: str
    note: str = ""
    weight: float = 1.0
    since: date | None = None
    in_old_list: bool = True
    churned: date | None = None


@dataclass
class Order:
    no: str
    day: date
    customer: Customer
    lines: list = field(default_factory=list)  # (Article, qty, unit_price)
    due: date | None = None
    delivered: date | None = None


def _phone(rng: random.Random) -> str:
    n = f"{rng.randint(200, 899)}{rng.randint(1000, 9999)}"
    return rng.choice([f"040 {n}", f"040/{n}", f"+49 40 {n[:3]} {n[3:]}", f"040-{n}", f"(040) {n}"])


def _slug(s: str) -> str:
    t = s.lower()
    for a, b in (("ä", "ae"), ("ö", "oe"), ("ü", "ue"), ("ß", "ss"), ("é", "e"), ("&", ""), (".", ""), (" ", "-")):
        t = t.replace(a, b)
    return "".join(ch for ch in t if ch.isalnum() or ch == "-").strip("-").replace("--", "-")


def _round2(x: float) -> float:
    return float(f"{x + 1e-9:.2f}")


def make_iban(rng: random.Random, blz: str | None = None) -> str:
    bban = (blz or f"{rng.randint(10000000, 99999999)}") + f"{rng.randint(0, 10**10 - 1):010d}"
    num = int("".join(str(int(c, 36)) for c in bban + "DE00"))
    return f"DE{98 - num % 97:02d}{bban}"


def de_money(x: float) -> str:
    s = f"{abs(x):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return ("-" if x < 0 else "") + s


def build_world(today: date, rng: random.Random):
    start = today - timedelta(days=660)

    # --- customers ---------------------------------------------------------
    customers: list[Customer] = []
    used = set()
    k = 1001
    while len(customers) < 58:
        kind, group = rng.choice(KIND)
        place = rng.choice(PLACE)
        name = f"{kind} {place}{rng.choice(LEGAL)}"
        if (kind, place) in used:
            continue
        used.add((kind, place))
        plz, city = rng.choice(CITIES)
        contact = f"{rng.choice(FIRST)} {rng.choice(LAST)}"
        domain = _slug(f"{kind}-{place}") + ".de"
        email = rng.choice(["info", "kontakt", "einkauf", contact.split()[0].lower()]) + "@" + domain
        if rng.random() < 0.12:
            email = ""
        terms = 30 if group in ("Hotel", "Büro") and rng.random() < 0.7 else 14
        c = Customer(f"K-{k}", name, contact, f"{rng.choice(STREETS)} {rng.randint(1, 140)}", plz, city,
                     _phone(rng), email, terms, group, weight=rng.paretovariate(1.3))
        c.since = start - timedelta(days=rng.randint(30, 2000))
        customers.append(c)
        k += 1
    # a customer named Müller for umlaut matching
    customers[7].name = "Kaffeerösterei-Shop Müller & Söhne"
    customers[7].group = "Einzelhandel"
    # customers acquired after the old list was printed (only in CRM)
    for c in customers[-9:]:
        c.in_old_list = False
        c.since = start + timedelta(days=rng.randint(10, 600))
    # two regular customers that quietly stopped ordering (the "aha" for the owner)
    for c in (customers[2], customers[9]):
        c.weight = max(c.weight, 3.0)
        c.churned = today - timedelta(days=rng.randint(62, 80))
    customers[3].note = "Liefert nur dienstags"
    customers[11].note = "Rechnung per Post!"
    customers[19].note = "Großkunde, Sonderpreis Espresso"
    return start, customers


def make_orders(today: date, start: date, customers: list[Customer], rng: random.Random) -> list[Order]:
    orders: list[Order] = []
    active = [c for c in customers]
    weights = [c.weight for c in active]
    day = start
    seq = {2025: 1, 2026: 1}
    by_year_no = {}
    while day <= today:
        if day.weekday() < 5:
            for _ in range(rng.choice([0, 1, 1, 1, 2])):
                cust = rng.choices(active, weights=weights)[0]
                for _ in range(6):
                    if not (cust.churned and day >= cust.churned):
                        break
                    cust = rng.choices(active, weights=weights)[0]
                if not cust.in_old_list and cust.since and day < cust.since:
                    continue
                if cust.churned and day >= cust.churned:
                    continue
                y = day.year
                if y not in seq:
                    seq[y] = 1
                no = f"A-{y}-{seq[y]:04d}"
                seq[y] += 1
                o = Order(no, day, cust)
                for sku in rng.sample(B2B_SKUS, k=rng.choice([1, 1, 2, 2, 3, 4])):
                    art = next(a for a in ARTICLES if a.sku == sku)
                    qty = rng.choice([2, 3, 4, 5, 6, 8, 10, 12]) if art.pack != "3 kg" else rng.choice([1, 2, 3])
                    price = art.price
                    if cust.note.startswith("Großkunde") and sku.startswith("HB-E"):
                        price = _round2(price * 0.9)
                    o.lines.append((art, qty, price))
                o.due = day + timedelta(days=rng.choice([2, 3, 3, 4, 5]))
                lag = rng.choice([1, 1, 2, 2, 2, 3, 3, 4, 5]) + (rng.choice([0, 0, 0, 0, 0, 0, 3, 6]))
                d = day + timedelta(days=lag)
                if d <= today - timedelta(days=0) and not (day > today - timedelta(days=9) and rng.random() < 0.6):
                    o.delivered = d
                orders.append(o)
                by_year_no[no] = o
        day += timedelta(days=1)
    return orders


def make_invoices(today: date, orders: list[Order], rng: random.Random):
    """Returns list of invoice dicts. Bookkeeping maintained 'Bezahlt am' only until end of last month."""
    book_cutoff = today.replace(day=1) - timedelta(days=1)
    invoices = []
    seq = {}
    not_invoiced = set()
    delivered = [o for o in orders if o.delivered]
    recent = [o for o in delivered if 8 <= (today - o.delivered).days < 25]
    for o in rng.sample(recent, k=min(4, len(recent))):
        not_invoiced.add(o.no)
    for o in sorted(delivered, key=lambda o: o.delivered):
        if o.no in not_invoiced:
            continue
        inv_day = o.delivered + timedelta(days=rng.choice([0, 1, 1, 2, 3]))
        if inv_day > today - timedelta(days=1):
            continue
        y = inv_day.year
        seq[y] = seq.get(y, 0) + 1
        no = f"RE-{y}-{seq[y]:04d}"
        net = _round2(sum(q * p for _, q, p in o.lines))
        vat = _round2(net * 0.07)
        gross = _round2(net + vat)
        due = inv_day + timedelta(days=o.customer.terms)
        r = rng.random()
        status = "offen"
        paid = None
        if r < 0.02:
            status = "storniert"
        else:
            late = rng.random() < 0.12
            pay_lag = rng.randint(o.customer.terms + 8, o.customer.terms + 45) if late else rng.randint(3, o.customer.terms + 2)
            pay_day = inv_day + timedelta(days=pay_lag)
            if pay_day <= today - timedelta(days=1):
                paid = pay_day
        invoices.append(dict(no=no, day=inv_day, order=o, net=net, vat=vat, gross=gross, due=due,
                             paid=paid, status=status))
    for inv in invoices:
        if inv["status"] == "storniert":
            inv["paid_book"] = None
            continue
        if inv["paid"] and inv["paid"] <= book_cutoff:
            inv["paid_book"] = inv["paid"]
            inv["status"] = "bezahlt"
        else:
            inv["paid_book"] = None
            overdue = (min(today, book_cutoff) - inv["due"]).days
            inv["status"] = "gemahnt" if overdue > 14 else "offen"
    return invoices, book_cutoff


def write_kundenliste(path: Path, customers: list[Customer], start: date, rng: random.Random):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Kunden"
    ws.append(["Röstwerk Hafenbohne GmbH – Kundenliste"])
    ws.append([f"Stand: {(start - timedelta(days=14)).strftime('%d.%m.%Y')}"])
    ws.append([])
    ws.append(["Kd-Nr.", "Firma", "Ansprechpartner", "Straße", "PLZ", "Ort", "Telefon", "E-Mail",
               "Zahlungsziel (Tage)", "Kundengruppe", "Bemerkung"])
    rows = [c for c in customers if c.in_old_list]
    for i, c in enumerate(rows):
        name = c.name + ("  " if rng.random() < 0.1 else "")
        ws.append([c.no, name, c.contact, c.street, int(c.plz) if rng.random() < 0.5 else c.plz, c.city, c.phone,
                   c.email, c.terms, c.group, c.note or None])
        if i == 20:
            ws.append([])  # stray empty row in the middle
    # duplicate customer created twice by accident (different number, same company)
    dup = rows[4]
    ws.append(["K-1099", dup.name.replace("Café", "Cafe").replace(" GmbH", "") + " GmbH", dup.contact, dup.street,
               dup.plz, dup.city, dup.phone.replace(" ", ""), dup.email, dup.terms, dup.group, "doppelt?"])
    wb.save(path)


def write_crm(path: Path, customers: list[Customer], rng: random.Random):
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["Company", "Contact", "Email", "Phone", "City", "Customer Since", "Segment", "Notes"])
        for c in customers:
            if c.in_old_list and rng.random() < 0.35:
                continue  # CRM was introduced later: not every old customer is in it
            name = c.name
            if rng.random() < 0.25:
                name = name.replace(" GmbH & Co. KG", "").replace(" GmbH", "").replace(" e.K.", "")
            if "Müller" in name:
                name = name.replace("Müller", "Mueller").replace("Söhne", "Soehne")
            seg = {"Café": "Cafe", "Restaurant": "Gastro", "Hotel": "Hotel", "Büro": "Office",
                   "Einzelhandel": "Retail"}[c.group]
            w.writerow([name, c.contact, c.email.upper() if rng.random() < 0.1 else c.email, c.phone, c.city,
                        c.since.isoformat() if c.since else "", seg, ""])


def write_preisliste(path: Path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Preisliste 2026"
    ws.append(["Preisliste gültig ab 01.01.2026 – alle Preise netto"])
    ws.append(["Art.-Nr.", "Bezeichnung", "Sorte / Herkunft", "Röstung", "Gebinde", "VK netto €", "MwSt"])
    group = None
    for a in ARTICLES:
        if a.group != group:
            group = a.group
            ws.append([group])
        ws.append([a.sku, a.name, a.variety or None, a.roast or None, a.pack, a.price, "7%" if a.vat == 0.07 else "19%"])
    for row in ws.iter_rows(min_row=3, min_col=6, max_col=6):
        for cell in row:
            cell.number_format = '#,##0.00 "€"'
    wb.save(path)


def write_auftraege(path: Path, orders: list[Order], rng: random.Random):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    header = ["Auftrag", "Datum", "Kd-Nr", "Kunde", "Art.-Nr", "Artikel", "Menge", "Einheit", "Einzelpreis",
              "Gesamt", "Liefertermin", "Geliefert am"]
    for year in sorted({o.day.year for o in orders}):
        ws = wb.create_sheet(str(year))
        ws.append(header)
        for o in [o for o in orders if o.day.year == year]:
            cname = o.customer.name
            if rng.random() < 0.08:
                cname = cname.replace("Café", "Cafe").replace("ä", "ae")
            if rng.random() < 0.05:
                cname = cname.split(" GmbH")[0]
            for art, qty, price in o.lines:
                ws.append([o.no, o.day, o.customer.no, cname, art.sku, art.name, qty, "Stk", price,
                           _round2(qty * price), o.due, o.delivered])
        for row in ws.iter_rows(min_row=2):
            for idx in (1, 10, 11):
                row[idx].number_format = "DD.MM.YYYY"
    wb.save(path)


def write_shop(path: Path, today: date, start: date, rng: random.Random):
    shop_customers = []
    for i in range(240):
        fn, ln = rng.choice(PRIVATE_FIRST), rng.choice(PRIVATE_LAST)
        email = f"{_slug(fn)}.{_slug(ln)}{rng.randint(1, 99)}@{rng.choice(['web.de', 'gmx.de', 'posteo.de', 'gmail.com', 'mailbox.org'])}"
        shop_customers.append((f"{fn} {ln}", email, rng.choice(SHOP_CITIES), rng.paretovariate(1.6)))
    weights = [c[3] for c in shop_customers]
    rows = []
    n = 1001
    day = start + timedelta(days=20)
    while day <= today:
        for _ in range(rng.choice([0, 1, 1, 1, 2])):
            name, email, city, _w = rng.choices(shop_customers, weights=weights)[0]
            ts = datetime(day.year, day.month, day.day, rng.randint(7, 22), rng.randint(0, 59), rng.randint(0, 59))
            lines = []
            for sku in rng.sample(SHOP_SKUS, k=rng.choice([1, 1, 1, 2, 2, 3])):
                art = next(a for a in ARTICLES if a.sku == sku)
                lines.append((art, rng.choice([1, 1, 1, 2, 2, 3])))
            subtotal = _round2(sum(a.price * (1 + a.vat) * q for a, q in lines))
            shipping = 0.0 if subtotal >= 39 else 4.90
            total = _round2(subtotal + shipping)
            fin = "paid"
            r = rng.random()
            if r < 0.03:
                fin = "refunded"
            elif day > today - timedelta(days=3) and r < 0.3:
                fin = "pending"
            fulfilled_at = None
            if fin == "paid":
                fd = ts + timedelta(days=rng.choice([0, 1, 1, 1, 2, 3]), hours=rng.randint(1, 8))
                if fd.date() <= today and not (day > today - timedelta(days=4) and rng.random() < 0.5):
                    fulfilled_at = fd
            for i, (art, q) in enumerate(lines):
                first = i == 0
                rows.append([
                    f"#{n}",
                    ts.strftime("%Y-%m-%d %H:%M:%S +0100") if first else "",
                    name if first else "",
                    email,
                    city if first else "",
                    art.name, art.sku, q, f"{art.price * (1 + art.vat):.2f}",
                    f"{subtotal:.2f}" if first else "",
                    f"{shipping:.2f}" if first else "",
                    f"{total:.2f}" if first else "",
                    fin if first else "",
                    ("fulfilled" if fulfilled_at else "unfulfilled") if first else "",
                    fulfilled_at.strftime("%Y-%m-%d %H:%M:%S +0100") if (first and fulfilled_at) else "",
                ])
            n += 1
        day += timedelta(days=1)
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["Name", "Created at", "Billing Name", "Email", "Shipping City", "Lineitem name", "Lineitem sku",
                    "Lineitem quantity", "Lineitem price", "Subtotal", "Shipping", "Total", "Financial Status",
                    "Fulfillment Status", "Fulfilled at"])
        w.writerows(rows)


def write_rechnungsbuch(path: Path, invoices: list[dict], rng: random.Random):
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";", lineterminator="\r\n")
    w.writerow(["Rechnungsnr", "Rechnungsdatum", "Auftrag", "Kd-Nr", "Kunde", "Netto", "MwSt", "Brutto",
                "Fällig am", "Bezahlt am", "Status"])
    tn = tv = tg = 0.0
    for inv in invoices:
        o = inv["order"]
        kd = o.customer.no if rng.random() > 0.03 else ""
        w.writerow([inv["no"], inv["day"].strftime("%d.%m.%Y"), o.no, kd, o.customer.name, de_money(inv["net"]),
                    de_money(inv["vat"]), de_money(inv["gross"]), inv["due"].strftime("%d.%m.%Y"),
                    inv["paid_book"].strftime("%d.%m.%Y") if inv["paid_book"] else "", inv["status"]])
        tn += inv["net"]
        tv += inv["vat"]
        tg += inv["gross"]
    w.writerow(["Summe", "", "", "", "", de_money(tn), de_money(tv), de_money(tg), "", "", ""])
    path.write_bytes(buf.getvalue().encode("cp1252"))


def write_bank(path: Path, today: date, invoices: list[dict], purchases: list[dict], rng: random.Random):
    since = today - timedelta(days=58)
    tx = []
    ibans: dict[str, str] = {}
    supplier_iban = {s["name"]: s["iban"] for s in SUPPLIERS}
    paid_here = [inv for inv in invoices if inv.get("paid") and since <= inv["paid"] <= today]
    # some customers pay without naming the invoice (the amount has to tell), one pays two at once
    no_ref = {inv["no"] for inv in paid_here[2::9]}
    by_cust: dict[str, list] = {}
    for inv in paid_here:
        by_cust.setdefault(inv["order"].customer.no, []).append(inv)
    combined = next((v[:2] for v in by_cust.values() if len(v) >= 2 and not ({x["no"] for x in v[:2]} & no_ref)), [])
    for inv in paid_here:
        o = inv["order"]
        iban = ibans.setdefault(o.customer.no, make_iban(rng))
        if inv in combined:
            if inv is combined[0]:
                total = _round2(sum(x["gross"] for x in combined))
                tx.append((max(x["paid"] for x in combined), "GUTSCHR. UEBERWEISUNG",
                           f"Sammelzahlung {o.customer.name[:20]}", o.customer.name.upper()[:27], iban, total))
            continue
        if inv["no"] in no_ref:
            ref = rng.choice([f"Zahlung {o.customer.name[:22]}", "Danke für die Lieferung", f"Kd {o.customer.no}",
                              "Kaffee"])
        else:
            ref = rng.choice([inv["no"], inv["no"], f"Rechnung {inv['no'][3:]}", f"{inv['no']} Danke",
                              inv["no"].replace("-", "", 1), f"Re.Nr. {inv['no']} Kd {o.customer.no}"])
        tx.append((inv["paid"], "GUTSCHR. UEBERWEISUNG", ref, o.customer.name.upper()[:27], iban, inv["gross"]))
    for p in purchases:
        if p.get("paid") and since <= p["paid"] <= today:
            tx.append((p["paid"], "UEBERWEISUNG", f"Rechnung {p['supplier_invoice']} Bestellung {p['no']}",
                       p["supplier"].upper()[:27], supplier_iban[p["supplier"]], -p["sum"]))
    landlord, power, shop_iban, own = make_iban(rng), make_iban(rng), make_iban(rng), make_iban(rng, "20050550")
    d = since
    while d <= today:
        if d.day == 1:
            tx.append((d, "DAUERAUFTRAG", "Miete Lager Hafenstraße 12", "HAFENIMMOBILIEN KG", landlord, -2350.00))
        if d.day == 28:
            tx.append((d, "LOHN GEHALT", "Gehälter " + d.strftime("%m/%Y"), "SAMMLER LOHN", "", -18640.00))
        if d.day == 15:
            tx.append((d, "LASTSCHRIFT", "Strom Abschlag", "STADTWERKE HAMBURG", power, -412.00))
        if d.weekday() == 0:
            tx.append((d, "GUTSCHR. UEBERWEISUNG", f"Auszahlung Shop {d.isocalendar()[1]}/{d.year}", "SHOP PAYMENTS EUROPE",
                       shop_iban, round(rng.uniform(380, 1450), 2)))
        d += timedelta(days=1)
    tx.sort(key=lambda t: t[0])
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";", quoting=csv.QUOTE_ALL, lineterminator="\r\n")
    w.writerow(["Auftragskonto", "Buchungstag", "Valutadatum", "Buchungstext", "Verwendungszweck",
                "Beguenstigter/Zahlungspflichtiger", "Kontonummer/IBAN", "Betrag", "Waehrung", "Info"])
    for day, text, ref, party, iban, amount in tx:
        w.writerow([own, day.strftime("%d.%m.%y"), day.strftime("%d.%m.%y"), text, ref, party,
                    iban, de_money(amount), "EUR", "Umsatz gebucht"])
    path.write_bytes(buf.getvalue().encode("cp1252"))


SUPPLIERS = [
    dict(iban="DE34200400000123456789", name="Kaffee Kontor Nord GmbH", short="Kaffee Kontor Nord", contact="Jens Behrmann", phone="040 / 33 44 55 0",
         email="einkauf@kaffeekontor-nord.de", supplies="Rohkaffee Brasil, Kolumbien, Peru", terms="30 Tage",
         note="Mindestmenge 10 Sack, Lieferung frei Haus ab 20 Sack"),
    dict(iban="DE39290500000077123456", name="Origin Green Coffee Traders", short="Origin Green Coffee", contact="Maria Okafor", phone="+49 421 556677",
         email="orders@origin-green.example", supplies="Rohkaffee Äthiopien, Kenia, Guatemala", terms="14 Tage",
         note="Spezialitäten, Preise in EUR/kg, Qualität top"),
    dict(iban="DE13207500000045678901", name="Elbe Verpackungen GmbH", short="Elbe Verpackungen", contact="Sven Kröger", phone="04101 778899",
         email="service@elbe-verpackungen.example", supplies="Ventilbeutel 250 g / 1 kg, Kartons", terms="30 Tage",
         note=""),
    dict(iban="DE52200505501234098765", name="Barista Technik Hansen", short="Barista Technik Hansen", contact="Kai Hansen", phone="0171 2233445",
         email="kai@barista-hansen.example", supplies="Wartung Siebträger, Mühlen", terms="sofort",
         note="Notfallnummer auch am Wochenende"),
    dict(iban="DE86700500000002345678", name="Bio-Rohkaffee Import Süd", short="Bio-Rohkaffee Süd", contact="Lisa Ammann", phone="089 445566",
         email="info@bio-rohkaffee-sued.example", supplies="Bio-Rohkaffee Peru, Honduras", terms="30 Tage",
         note="Bio-Zertifikat jährlich anfordern!"),
]


def write_lieferanten(path: Path):
    lines = ["LIEFERANTEN – Stand Sommer (bitte aktuell halten!)", ""]
    for s in SUPPLIERS:
        lines += [f"Firma: {s['name']}", f"Ansprechpartner: {s['contact']}", f"Telefon: {s['phone']}",
                  f"E-Mail: {s['email']}", f"Liefert: {s['supplies']}", f"Zahlungsziel: {s['terms']}"]
        if s["note"]:
            lines.append(f"Notiz: {s['note']}")
        lines.append("")
    path.write_text("\n".join(lines), encoding="utf-8")


def make_purchases(today: date, start: date, rng: random.Random) -> list[dict]:
    purchases = []
    d = start + timedelta(days=5)
    seq = {}
    origin = [("Kaffee Kontor Nord GmbH", ["Brasil Santos", "Kolumbien Huila", "Peru Cajamarca"]),
              ("Origin Green Coffee Traders", ["Äthiopien Sidamo", "Kenia AA", "Guatemala Antigua"]),
              ("Bio-Rohkaffee Import Süd", ["Peru Bio", "Honduras Bio"])]
    inv_no = 40
    while d <= today:
        sup, origins = rng.choice(origin)
        y = d.year
        seq[y] = seq.get(y, 0) + 1
        no = f"EK-{y}-{seq[y]:03d}"
        bags = rng.choice([5, 8, 10, 10, 12, 15, 20])
        kg = bags * 60
        price = round(rng.uniform(5.2, 9.8), 2)
        total = _round2(kg * price)
        received = d + timedelta(days=rng.randint(7, 21))
        if received > today:
            received = None
        inv_no += rng.randint(3, 12)
        paid = (received + timedelta(days=rng.randint(5, 25))) if received else None
        if paid and paid > today:
            paid = None
        short = next(s["short"] for s in SUPPLIERS if s["name"] == sup)
        purchases.append(dict(no=no, day=d, supplier=sup, supplier_short=short, origin=rng.choice(origins), bags=bags,
                              kg=kg, price=price, sum=total, received=received,
                              supplier_invoice=f"{y}-{inv_no:03d}", paid=paid))
        d += timedelta(days=rng.randint(14, 30))
    return purchases


def write_einkauf(path: Path, purchases: list[dict], rng: random.Random):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Einkauf"
    ws.append(["Bestell-Nr.", "Datum", "Lieferant", "Herkunft", "Säcke", "kg", "Preis/kg €", "Summe €", "Eingang am",
               "Rechnung Lieferant"])
    for p in purchases:
        sup = p["supplier"] if rng.random() < 0.6 else p["supplier_short"]
        ws.append([p["no"], p["day"], sup, p["origin"], p["bags"], p["kg"], p["price"], p["sum"], p["received"],
                   p["supplier_invoice"] if p["received"] else None])
    for row in ws.iter_rows(min_row=2):
        for idx in (1, 8):
            row[idx].number_format = "DD.MM.YYYY"
    wb.save(path)


def write_lager(path: Path, today: date, rng: random.Random):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"KW{today.isocalendar()[1]}"
    ws.append([f"Lagerbestand KW {today.isocalendar()[1]}/{today.year}"])
    ws.append(["Art.-Nr.", "Artikel", "Bestand", "Mindestbestand", "Lagerplatz"])
    low = {"HB-E01", "HB-F05", "HB-Z01"}
    for a in ARTICLES:
        minimum = 20 if a.pack == "1 kg" else 15 if a.pack in ("250 g", "500 g") else 5
        stock = rng.randint(2, minimum - 1) if a.sku in low else rng.randint(minimum + 2, minimum * 4)
        ws.append([a.sku, a.name, stock, minimum, rng.choice(["R1-A", "R1-B", "R2-A", "R2-C", "Keller"])])
    wb.save(path)


def write_mitarbeiter(path: Path, today: date, rng: random.Random):
    rows = [("Hauke Brandt", "Geschäftsführung", "0171 1112223", "01.03.2016", 40),
            ("Frieda Brandt", "Geschäftsführung / Einkauf", "0171 1112224", "01.03.2016", 40),
            ("Ole Thomsen", "Röstmeister", "0160 5556667", "15.08.2018", 40),
            ("Greta Lorenzen", "Rösterin", "0152 3334445", "01.02.2021", 32),
            ("Jan Witt", "Versand / Lager", "0157 7778889", "01.04.2022", 38),
            ("Mia Carstens", "Vertrieb B2B", "0176 9990001", "01.09.2023", 40),
            ("Lina Albers", "Onlineshop / Social Media", "0151 2223334", "15.01.2024", 25),
            ("Theo Kruse", "Aushilfe Versand", "0159 4445556", "01.06.2025", 12)]
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["Name", "Funktion", "Mobil", "Eintritt", "Wochenstunden"])
        w.writerows(rows)


def write_mails(folder: Path, today: date, customers: list[Customer], invoices: list[dict], purchases: list[dict]):
    folder.mkdir(parents=True, exist_ok=True)
    hotel = next((c for c in customers if c.name.startswith("Hotel") and c.email), customers[0])

    def mail(fname, frm, to, subject, body, when):
        m = EmailMessage()
        m["From"] = frm
        m["To"] = to
        m["Subject"] = subject
        m["Date"] = format_datetime(datetime(when.year, when.month, when.day, 9, 14))
        m.set_content(body)
        (folder / fname).write_bytes(bytes(m))

    mail("2026-bestellung-hotel.eml", f"{hotel.contact} <{hotel.email}>", "bestellung@hafenbohne.example",
         "Bestellung für nächste Woche",
         f"Hallo Frau Carstens,\n\nbitte liefern Sie uns bis Freitag:\n- 6 x Hafenmischung Espresso 1 kg (HB-E01)\n"
         f"- 4 x Elbsegler Filter 1 kg (HB-F02)\n\nUnsere Kundennummer ist {hotel.no}.\n\n"
         f"Viele Grüße\n{hotel.contact}\n{hotel.name}", today - timedelta(days=1))
    open_inv = [i for i in invoices if i["status"] in ("offen", "gemahnt") and i["order"].customer.email]
    if open_inv:
        inv = open_inv[0]
        c = inv["order"].customer
        mail("rueckfrage-rechnung.eml", f"{c.contact} <{c.email}>", "buchhaltung@hafenbohne.example",
             f"Rückfrage zu Rechnung {inv['no']}",
             f"Guten Tag,\n\nzur Rechnung {inv['no']} vom {inv['day'].strftime('%d.%m.%Y')} über "
             f"{de_money(inv['gross'])} EUR: Es wurden nur zwei statt drei Kartons geliefert.\n"
             f"Bitte prüfen Sie das und schicken Sie uns eine korrigierte Rechnung.\n\nBeste Grüße\n{c.contact}\n{c.name}",
             today - timedelta(days=3))
    open_po = [p for p in purchases if not p["received"]] or purchases[-1:]
    p = open_po[0]
    mail("versandavis-kontor.eml", "Kaffee Kontor Nord GmbH <versand@kaffeekontor-nord.de>", "einkauf@hafenbohne.example",
         f"Versandavis zu Ihrer Bestellung {p['no']}",
         f"Sehr geehrte Damen und Herren,\n\nIhre Bestellung {p['no']} ({p['bags']} Sack {p['origin']}) "
         f"wurde heute verladen. Voraussichtliche Anlieferung in 5 Werktagen.\n\nMit freundlichen Grüßen\n"
         f"Kaffee Kontor Nord GmbH – Versand", today - timedelta(days=2))
    mail("neue-bankverbindung.eml", "Kaffee Kontor Nord – Buchhaltung <buchhaltung@kaffeekontor-nord-gmbh.de>",
         "einkauf@hafenbohne.example", "Wichtig: Neue Bankverbindung ab sofort",
         f"Sehr geehrte Damen und Herren,\n\naufgrund eines Bankwechsels bitten wir Sie, die offene Rechnung "
         f"{p['supplier_invoice']} zu Ihrer Bestellung {p['no']} ab sofort auf unser neues Konto zu überweisen:\n"
         f"IBAN DE33 5001 0517 5408 4126 93\nBIC INGDDEFFXXX\n\nMit freundlichen Grüßen\nKaffee Kontor Nord GmbH",
         today)
    mail("newsletter.eml", "Röstmagazin <news@roestmagazin.example>", "info@hafenbohne.example",
         "Die Trends der Saison: Fermentation, Anaerobic & Co.",
         "Liebe Röstfreunde,\n\nin dieser Ausgabe: Warum anaerobe Fermentation polarisiert ...\n\nAbmelden: hier klicken",
         today - timedelta(days=5))


def write_supplier_pdf(path: Path, today: date, purchases: list[dict]):
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    done = [p for p in purchases if p["received"] and p["supplier"].startswith("Kaffee Kontor")] or purchases[:1]
    p = done[-1]
    c = canvas.Canvas(str(path), pagesize=A4)
    y = 800
    for line in ["Kaffee Kontor Nord GmbH · Speicherring 4 · 20457 Hamburg", "", "Röstwerk Hafenbohne GmbH",
                 "Hafenstraße 12", "20459 Hamburg", "", f"RECHNUNG Nr. {p['supplier_invoice']}",
                 f"Rechnungsdatum: {p['received'].strftime('%d.%m.%Y')}", f"Ihre Bestellung: {p['no']}", "",
                 f"{p['bags']} Sack Rohkaffee {p['origin']} à 60 kg = {p['kg']} kg",
                 f"Preis pro kg: {de_money(p['price'])} EUR", f"Gesamtbetrag: {de_money(p['sum'])} EUR",
                 "", "Zahlbar innerhalb 30 Tagen ohne Abzug.", "IBAN DE34 2004 0000 0123 4567 89 · BIC COBADEFFXXX"]:
        c.drawString(60, y, line)
        y -= 20
    c.save()


def write_wartung(path: Path, today: date, customers: list[Customer], rng: random.Random):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Verträge"
    ws.append(["Vertrag", "Kd-Nr", "Kunde", "Maschine", "Seriennummer", "Beginn", "Monatspauschale €",
               "Letzte Wartung", "Nächste Wartung"])
    pool = [c for c in customers if c.group in ("Café", "Restaurant", "Hotel")]
    for i, c in enumerate(rng.sample(pool, k=min(18, len(pool)))):
        begin = today - timedelta(days=rng.randint(60, 900))
        last = today - timedelta(days=rng.randint(10, 200))
        nxt = last + timedelta(days=180)
        ws.append([f"WV-{101 + i}", c.no, c.name, rng.choice(["La Marzocco Linea", "Rancilio Classe 5", "Sanremo Zoe",
                                                               "Mahlkönig E65S"]),
                   f"SN{rng.randint(100000, 999999)}", begin, rng.choice([39.0, 49.0, 59.0, 79.0]), last, nxt])
    for row in ws.iter_rows(min_row=2):
        for idx in (5, 7, 8):
            row[idx].number_format = "DD.MM.YYYY"
    wb.save(path)


def generate(out: Path, today: date | None = None, growth: bool = True) -> dict:
    """Write the demo dump into `out`. Returns {'initial': [paths], 'growth': [paths]}."""
    today = today or date.today()
    rng = random.Random(SEED)
    out.mkdir(parents=True, exist_ok=True)
    start, customers = build_world(today, rng)
    orders = make_orders(today, start, customers, rng)
    invoices, _cutoff = make_invoices(today, orders, rng)
    purchases = make_purchases(today, start, rng)

    files = {
        "Kundenliste_alt.xlsx": lambda p: write_kundenliste(p, customers, start, rng),
        "CRM_Export.csv": lambda p: write_crm(p, customers, rng),
        "Artikel_Preisliste_2026.xlsx": lambda p: write_preisliste(p),
        "Aufträge_Großkunden.xlsx": lambda p: write_auftraege(p, orders, rng),
        "Shop_Bestellungen_Export.csv": lambda p: write_shop(p, today, start, rng),
        "Rechnungsausgangsbuch.csv": lambda p: write_rechnungsbuch(p, invoices, rng),
        "Kontoauszug_Sparkasse.csv": lambda p: write_bank(p, today, invoices, purchases, rng),
        "Lieferanten.txt": lambda p: write_lieferanten(p),
        "Einkauf_Rohkaffee.xlsx": lambda p: write_einkauf(p, purchases, rng),
        "Lagerbestand.xlsx": lambda p: write_lager(p, today, rng),
        "Mitarbeiter.csv": lambda p: write_mitarbeiter(p, today, rng),
        "Lieferantenrechnung_Kontor.pdf": lambda p: write_supplier_pdf(p, today, purchases),
    }
    written = []
    for name, fn in files.items():
        p = out / name
        fn(p)
        written.append(p)
    mails = out / "Mails"
    write_mails(mails, today, customers, invoices, purchases)
    written += sorted(mails.glob("*.eml"))
    result = {"initial": written, "growth": []}
    if growth:
        gdir = out.parent / (out.name + "_spaeter")
        gdir.mkdir(parents=True, exist_ok=True)
        gp = gdir / "Wartungsverträge.xlsx"
        write_wartung(gp, today, customers, rng)
        result["growth"].append(gp)
    return result


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    res = generate(Path(sys.argv[1]))
    for k, v in res.items():
        for p in v:
            print(k, p)
