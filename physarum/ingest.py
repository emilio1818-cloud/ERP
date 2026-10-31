"""Digestion: turn any file into tables (grids with a header) and documents (prose).

Handles the usual mess: title rows above the header, blank rows, section rows
("ESPRESSO" between articles), sum rows at the bottom, repeated headers, cp1252 CSVs
with semicolons, key/value text blocks, e-mails with attachments, PDFs, Word files,
JSON with nested arrays and ZIP archives.
"""
from __future__ import annotations

import csv
import io
import itertools
import json
import re
import zipfile
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime
from email import policy
from email.parser import BytesParser
from email.utils import parsedate_to_datetime
from pathlib import PurePath

from . import lexicon
from .text import fold
from .values import parse_date, parse_number

MAX_ZIP_TOTAL = 300 * 1024 * 1024
MAX_ZIP_MEMBERS = 3000
MAX_ZIP_RATIO = 150          # a member that expands more than this is a bomb, not an export
MAX_ROWS = 400_000           # per table: beyond that, a file is a database dump, not a list
csv.field_size_limit(16 * 1024 * 1024)


@dataclass
class RawTable:
    name: str
    columns: list[str]
    rows: list[list]
    lines: list[int]
    sheet: str | None = None
    meta: dict = field(default_factory=dict)


@dataclass
class RawDoc:
    kind: str
    title: str
    text: str
    sender: str | None = None
    recipients: str | None = None
    date: datetime | None = None
    meta: dict = field(default_factory=dict)


@dataclass
class Digest:
    tables: list[RawTable] = field(default_factory=list)
    docs: list[RawDoc] = field(default_factory=list)
    attachments: list[tuple[str, bytes]] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)


SHEET_EXT = {".xlsx", ".xlsm", ".xltx", ".xltm"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".gif", ".tif", ".tiff", ".bmp", ".heic", ".webp"}


def decode_text(data: bytes) -> tuple[str, str]:
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        try:
            return data.decode("utf-16"), "utf-16"
        except UnicodeDecodeError:
            pass
    head = data[:4000]
    if len(head) >= 20 and head[1::2].count(0) > 0.4 * len(head[1::2]) and head[0::2].count(0) < 0.05 * len(head[0::2]):
        try:
            return data.decode("utf-16-le"), "utf-16-le"  # Windows exports without a BOM
        except UnicodeDecodeError:
            pass
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(enc), enc
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace"), "utf-8"


def read_file(name: str, data: bytes, sender: str | None = None) -> Digest:
    """`sender`: the address of the mail the file was attached to."""
    ext = PurePath(name).suffix.lower()
    stem = PurePath(name).stem
    d = Digest()
    if stem.startswith("~$"):
        d.notes.append("Sperrdatei von Office (die Mappe war beim Kopieren geöffnet) – nichts darin, übersprungen.")
        return d
    try:
        if ext in SHEET_EXT:
            _read_xlsx(stem, data, d)
        elif ext == ".xls":
            d.notes.append("Altes Excel-Format (.xls): bitte als .xlsx speichern – Datei wurde aufbewahrt.")
        elif ext == ".ods":
            d.notes.append("OpenDocument-Tabellen (.ods) werden noch nicht gelesen – bitte als .xlsx speichern.")
        elif ext == ".json":
            _read_json(stem, data, d)
        elif ext == ".xml":
            _read_xml(stem, data, d)
        elif ext == ".eml":
            _read_eml(stem, data, d)
        elif ext == ".pdf":
            _read_pdf(stem, data, d, sender)
        elif ext == ".docx":
            _read_docx(stem, data, d)
        elif ext == ".zip":
            _read_zip(data, d)
        elif ext in (".html", ".htm"):
            text, _ = decode_text(data)
            _read_html(stem, text, d)
        elif ext in IMAGE_EXT:
            d.docs.append(RawDoc("bild", stem, "", meta={"note": "Bild ohne Texterkennung"}))
        else:
            text, enc = decode_text(data)
            if "\x00" in text[:2000]:
                d.notes.append(f"Unbekanntes Binärformat ({ext or 'ohne Endung'}) – Datei wurde aufbewahrt.")
            elif _MT940.search(text[:5000]):
                mt = _mt940_table(text, stem)
                if mt is not None:
                    d.tables.append(mt)
                else:
                    d.notes.append(f"{stem}: ein Kontoauszug (MT940) ohne Buchungen")
            else:
                _read_text(stem, text, enc, ext, d)
    except zipfile.BadZipFile:
        d.notes.append(f"Konnte {name} nicht lesen: beschädigt oder keine echte {ext[1:].upper()}-Datei. Sie ist aufbewahrt.")
    except Exception as e:  # a broken file must never break the whole dump
        d.notes.append(f"Konnte {name} nicht lesen: {type(e).__name__}: {e}")
    return d


# --- text / csv ---------------------------------------------------------------------------------

def _sniff_delimiter(text: str) -> str | None:
    """The delimiter that gives the longest run of equally wide records. Records, not lines:
    bank exports put line breaks inside quoted purposes, and many files start with a preamble
    (bank name, account, period) that has fewer fields than the table below it."""
    best, best_score = None, 0.0
    head = text[:300000]
    for delim in (";", "\t", ",", "|"):
        try:
            recs = [r for r in itertools.islice(csv.reader(io.StringIO(head), delimiter=delim), 600)
                    if any(c.strip() for c in r)]
        except csv.Error:
            continue
        if len(recs) < 2:
            continue
        counts = Counter(len(r) for r in recs)
        cands = [(w, f) for w, f in counts.items() if w >= 2]
        if not cands:
            continue
        width, freq = max(cands, key=lambda x: (x[1] * min(x[0], 12), x[0]))
        run = longest = 0
        for r in recs:
            run = run + 1 if width - 1 <= len(r) <= width + 1 else 0
            longest = max(longest, run)
        if longest < 2 or (longest < 5 and freq / len(recs) < 0.6):
            continue
        score = longest * min(width, 12)
        if score > best_score:
            best, best_score = delim, score
    return best


_DEBIT_CREDIT = {"s": -1, "soll": -1, "h": 1, "haben": 1}


def _apply_debit_credit(t: "RawTable"):
    """Bank exports with positive amounts and a separate S/H column: Soll is money going out."""
    if not t.rows:
        return
    width = len(t.columns)
    for j in range(width):
        vals = [str(r[j]).strip().casefold() for r in t.rows if j < len(r) and r[j] not in (None, "")]
        if len(vals) < max(3, 0.9 * len(t.rows)) or not all(v in _DEBIT_CREDIT for v in vals):
            continue
        amount = None
        for k in list(range(j - 1, -1, -1)) + list(range(j + 1, width)):
            nums = [r[k] for r in t.rows if k < len(r) and r[k] not in (None, "")]
            ok = sum(1 for v in nums if (isinstance(v, (int, float)) and not isinstance(v, bool))
                     or parse_number(str(v)) is not None)
            if nums and ok >= 0.9 * len(nums):
                amount = k
                break
        if amount is None:
            return
        for r in t.rows:
            if j < len(r) and amount < len(r) and _DEBIT_CREDIT.get(str(r[j]).strip().casefold()) == -1:
                v = r[amount]
                if isinstance(v, (int, float)) and not isinstance(v, bool):
                    r[amount] = -abs(v)
                elif isinstance(v, str) and v.strip() and not v.strip().startswith("-"):
                    r[amount] = "-" + v.strip()
        # the sign is in the amount now: the S/H column says nothing more
        t.columns.pop(j)
        for r in t.rows:
            if j < len(r):
                r.pop(j)
        return


_SIGNATURE = re.compile(r"^(unterschrift|datum\s*/?,?\s*unterschrift|gez\.|gepr(ü|ue)ft|freigabe|abgezeichnet)\b.{0,30}$", re.I)
_ABOUT = re.compile(r"^\s*(mitarbeiter(?:in)?|monteur(?:in)?|arbeitnehmer(?:in)?|name|pers\.?\s*-?\s*nr\.?|personalnummer|"
                    r"personal-nr\.?)\s*[:.]?\s*(\S.{0,60}?)\s*$", re.I)
_ABOUT_NAMES = {"mitarbeiter": "Mitarbeiter", "mitarbeiterin": "Mitarbeiter", "monteur": "Mitarbeiter",
                "monteurin": "Mitarbeiter", "arbeitnehmer": "Mitarbeiter", "arbeitnehmerin": "Mitarbeiter",
                "name": "Mitarbeiter"}


_CONTEXT_KEYS = {"auftrag": "Auftrag", "auftragsnr": "Auftrag", "auftragsnummer": "Auftrag", "kommission": "Auftrag",
                 "projekt": "Projekt", "projektnr": "Projekt", "bauvorhaben": "Objekt", "bv": "Objekt",
                 "objekt": "Objekt", "baustelle": "Objekt", "kunde": "Kunde", "bauherr": "Kunde", "auftraggeber": "Kunde"}
_CONTEXT = re.compile(r"^\s*([A-Za-zÄÖÜäöüß.\- ]{2,20}?)\s*:\s*(\S.{0,80}?)\s*$")
_REF_TOKEN = re.compile(r"(?<![A-Za-z0-9])([A-Za-z]{1,4}\d{2}-\d{2,5}|[A-Za-z]{1,4}-\d{2,4}-\d{2,5})(?![A-Za-z0-9])")


def _context(rows, pending_title, name: str) -> dict[str, str]:
    """What a sheet is about, written above it ("Auftrag: A26-028", "Bauherr: Lindemann Wohnbau GmbH") – a
    measurement or a material list of one job. Without it, the number in the file name ("Aufmass_A26-003_…")."""
    out: dict[str, str] = {}
    for _, r in rows:
        cells = [v for v in r if v is not None and str(v).strip()]
        pairs = []
        for i, v in enumerate(cells):
            sv = str(v).strip()
            m = _CONTEXT.match(sv)
            if m:
                pairs.append((m.group(1), m.group(2)))
            elif sv.endswith(":") and i + 1 < len(cells) and not str(cells[i + 1]).strip().endswith(":"):
                pairs.append((sv[:-1], str(cells[i + 1]).strip()))
        for k, v in pairs:
            label = _CONTEXT_KEYS.get(fold(k).replace(" ", "").replace(".", "").replace("-", ""))
            if label and v:
                out.setdefault(label, v)
    if not out.get("Auftrag") and not out.get("Projekt"):
        m = _REF_TOKEN.search(name.replace("_", " "))
        if m:
            out["Bezug"] = m.group(1)
    return out


def _about(rows, pending_title) -> dict[str, str]:
    """Cells above a table that say whom it is about ("Mitarbeiter: Jens Kortmann", "Pers.-Nr. 103")."""
    out: dict[str, str] = {}
    cells = [v for _, r in rows for v in r if isinstance(v, str)] + [x for x in pending_title if isinstance(x, str)]
    for v in cells:
        m = _ABOUT.match(v)
        if not m:
            continue
        key = fold(m.group(1)).replace(" ", "").rstrip(".")
        label = _ABOUT_NAMES.get(key) or "Pers.-Nr."
        if label == "Mitarbeiter" and (key == "name" and len(m.group(2).split()) > 4):
            continue
        out.setdefault(label, m.group(2).strip())
    return out if "Mitarbeiter" in out else {}


def _datev_header(grid: list[list]) -> dict | None:
    """The first line of a DATEV export ("EXTF";700;21;"Buchungsstapel";…): advisor and client numbers,
    the financial year and the account length – what the next export has to say the same way."""
    head = [str(c).strip().strip('"') if c is not None else "" for c in (grid[0] if grid else [])]
    if len(head) < 14 or head[0] not in ("EXTF", "DTVF"):
        return None
    def num(i):
        return head[i] if i < len(head) and head[i].isdigit() else None
    return dict(kind=head[3] if len(head) > 3 else "", berater=num(10), mandant=num(11), wj_beginn=num(12),
                sachkonten=num(13), von=num(14), bis=num(15))


_OPEN_ITEMS = ("offene posten", "op-liste", "oplist", "offene-posten")


def _open_items(t: "RawTable") -> bool:
    """An open-items list from the tax advisor ("Offene Posten Debitoren, Stichtag 30.09.2026"): per
    invoice what was still open on one day – a snapshot of the books, not a list of customers."""
    title = fold(" ".join(str(x) for x in (t.meta.get("title"), t.name) if x))
    cols = {fold(str(c or "")).strip() for c in t.columns}
    beleg = any(c.startswith(("belegfeld", "beleg-nr", "belegnr", "rechnung")) for c in cols)
    return beleg and any(c.startswith(("offen", "op-betrag", "saldo")) for c in cols) and \
        (any(w in title for w in _OPEN_ITEMS) or "konto" in cols and "faellig" in cols)


def _report_accounts(t: "RawTable") -> dict[str, str]:
    """{invoice number: debtor account} from a report of the tax advisor."""
    cols = [fold(str(c or "")).strip() for c in t.columns]
    k = next((j for j, c in enumerate(cols) if c == "konto"), None)
    b = next((j for j, c in enumerate(cols) if c.startswith(("belegfeld 1", "belegfeld1", "beleg-nr", "belegnr"))), None)
    if k is None or b is None:
        return {}
    out = {}
    for r in t.rows:
        acc = str(r[k]).strip() if k < len(r) and r[k] is not None else ""
        no = str(r[b]).strip() if b < len(r) and r[b] is not None else ""
        if acc.isdigit() and no:
            out[no] = acc
    return out


_TS_NAMES = {"tag": "Datum", "kommission": "Auftrag", "auftrag/kunde": "Auftrag", "auftrags-nr.": "Auftrag",
             "auftragsnr.": "Auftrag", "auftrag-nr.": "Auftrag", "objekt": "Baustelle", "kunde/baustelle": "Baustelle",
             "baustelle/kunde": "Baustelle", "einsatzort": "Baustelle", "arbeiten": "Tätigkeit", "was gemacht": "Tätigkeit",
             "taetigkeit": "Tätigkeit", "beschreibung": "Tätigkeit", "std.": "Stunden", "std": "Stunden",
             "stunden": "Stunden", "arbeitszeit": "Stunden", "std. gesamt": "Stunden", "von": "Von", "bis": "Bis"}
_TS_WORDS = ("stundenzettel", "zeiterfassung", "stundennachweis", "arbeitszeit", "stundenbericht", "rapport")


def _timesheet_layout(t: "RawTable"):
    """Each fitter keeps his sheet his own way ("Kommission", "Objekt", "Arbeiten", "Stunden" here,
    "Auftrag", "Baustelle", "Tätigkeit", "Std." there): one time sheet for the firm."""
    cols = [fold(str(c or "")).strip() for c in t.columns]
    if not (any(w in fold(t.name) for w in _TS_WORDS) or "mitarbeiter" in cols):
        return
    if not any(c in ("datum", "tag") for c in cols) or not any(_TS_NAMES.get(c) == "Stunden" for c in cols):
        return
    taken = {fold(c) for c in t.columns}
    for j, c in enumerate(cols):
        new = _TS_NAMES.get(c)
        if new and fold(new) not in taken:
            taken.add(fold(new))
            t.columns[j] = new


_BALANCE = {"anfangssaldo", "endsaldo", "saldo", "kontostand", "alter kontostand", "neuer kontostand", "schlusssaldo"}
_PARTY_WORDS = ("auftraggeber", "zahlungsempf", "empfaenger", "zahlungspflichtig", "beguenstigt", "glaeubiger")
_TX_KINDS = ("ueberweisung", "gutschr", "lastschrift", "abschluss", "dauerauftrag", "kartenzahlung", "bargeld",
             "sepa", "entgelt", "zinsen", "storno", "rueckbuchung", "scheck", "einzahlung", "auszahlung")


def _bank_layout(t: "RawTable"):
    """Every bank exports its statement differently. The VR banks' "Umsatzanzeige" names the amount
    "Umsatz", puts the kind of booking on the first line of "Vorgang/Verwendungszweck", and writes the
    account holder into one of two name columns depending on the direction. Here it becomes the
    statement every other bank delivers: Betrag, Buchungstext, Verwendungszweck, Zahlungspartner."""
    cols = [fold(str(c or "")).strip() for c in t.columns]
    if not any(c.startswith(("buchungstag", "buchungsdatum", "valuta")) for c in cols) or len(t.rows) < 2:
        return
    # the balance lines below the bookings are no bookings
    t_rows = [(r, l) for r, l in zip(t.rows, t.lines)
              if not any(fold(str(v or "")).strip() in _BALANCE for v in r)]
    if len(t_rows) != len(t.rows):
        t.rows, t.lines = [r for r, _ in t_rows], [l for _, l in t_rows]
    if "umsatz" in cols and "betrag" not in cols:
        t.columns[cols.index("umsatz")] = "Betrag"
    mixed = next((j for j, c in enumerate(cols) if "vorgang" in c and "verwendungszweck" in c), None)
    if mixed is not None and "verwendungszweck" not in cols:
        vals = [str(r[mixed]) for r in t.rows if mixed < len(r) and r[mixed] not in (None, "")]
        heads = [v.replace("\r", "").split("\n", 1)[0] for v in vals if "\n" in v]
        if vals and len(heads) >= 0.8 * len(vals) and sum(1 for h in heads if any(w in fold(h) for w in _TX_KINDS)) >= 0.8 * len(heads):
            t.columns[mixed] = "Verwendungszweck"
            t.columns.insert(mixed, "Buchungstext")
            for r in t.rows:
                v = r[mixed] if mixed < len(r) else None
                parts = str(v).replace("\r", "").split("\n", 1) if v not in (None, "") else ["", ""]
                head, rest = (parts[0].strip(), " ".join(parts[1].split())) if len(parts) == 2 else ("", parts[0].strip())
                while len(r) <= mixed:
                    r.append(None)
                r[mixed:mixed + 1] = [head or None, rest or None]
            cols = [fold(str(c or "")).strip() for c in t.columns]
        elif vals:
            t.columns[mixed] = "Verwendungszweck"
    party = [j for j, c in enumerate(cols) if any(w in c for w in _PARTY_WORDS)]
    if len(party) == 2:
        a, b = party
        names = Counter(fold(str(r[j])).strip() for r in t.rows for j in (a, b) if j < len(r) and r[j] not in (None, ""))
        if names:
            holder, n = names.most_common(1)[0]
            if n >= 0.5 * len(t.rows):
                for r in t.rows:
                    va = r[a] if a < len(r) else None
                    vb = r[b] if b < len(r) else None
                    r[a] = vb if fold(str(va or "")).strip() == holder else va
                t.columns[a] = "Zahlungspartner"
                t.columns.pop(b)
                for r in t.rows:
                    if b < len(r):
                        r.pop(b)
                t.meta["account_holder"] = holder


_KV = re.compile(r"^\s*([A-Za-zÄÖÜäöüß][\wÄÖÜäöüß .\-/()&]{0,40}?)\s*:\s+(.*\S)?\s*$")


def _kv_blocks(text: str) -> RawTable | None:
    blocks, cur, line_no, starts = [], [], 0, []
    for i, raw in enumerate(text.splitlines(), start=1):
        if not raw.strip():
            if cur:
                blocks.append(cur)
                starts.append(line_no)
                cur = []
            continue
        if not cur:
            line_no = i
        cur.append(raw)
    if cur:
        blocks.append(cur)
        starts.append(line_no)
    recs, rec_lines, titles = [], [], []
    for b, start in zip(blocks, starts):
        kv = []
        for l in b:
            m = _KV.match(l)
            if m and not re.match(r"^\s*(https?|www)", l):
                kv.append([m.group(1).strip(), (m.group(2) or "").strip()])
            elif kv:
                kv[-1][1] = (kv[-1][1] + " " + l.strip()).strip()
        if len(kv) >= 2:
            recs.append(kv)
            rec_lines.append(start)
        else:
            titles.append(" ".join(x.strip() for x in b))
    if len(recs) < 2:
        return None
    key_freq = Counter(k for r in recs for k, _ in {(k, 0) for k, _ in r})
    recurring = [k for k, c in key_freq.items() if c >= 2]
    if len(recurring) < 2:
        return None
    columns = []
    for r in recs:
        for k, _ in r:
            if k not in columns:
                columns.append(k)
    rows = []
    for r in recs:
        d = {}
        for k, v in r:
            d[k] = v if k not in d else d[k] + "; " + v
        rows.append([d.get(c) for c in columns])
    return RawTable(name="", columns=columns, rows=rows, lines=rec_lines,
                    meta={"format": "Schlüssel: Wert", "title": " · ".join(titles)[:300], "locale": "de"})


def _prose(text: str) -> bool:
    """A note or a read-me: sentences, not rows."""
    lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
    if not lines:
        return False
    sentences = sum(1 for ln in lines if len(ln.split()) >= 6 and ln[-1] in ".!?:")
    return len(lines) <= 2 and sentences >= 1 or sentences >= 0.5 * len(lines)


def _read_text(stem: str, text: str, enc: str, ext: str, d: Digest):
    if _prose(text):
        d.docs.append(RawDoc("text", stem, text.strip()))
        return
    delim = _sniff_delimiter(text) if ext != ".md" else None
    if delim:
        grid = list(csv.reader(io.StringIO(text), delimiter=delim))
        locale = "de" if delim == ";" else None
        tabs = detect_tables(grid, stem, {"delimiter": delim, "encoding": enc, "locale": locale})
        if tabs:
            d.tables.extend(tabs)
            return
    kv = _kv_blocks(text)
    if kv:
        kv.name = stem
        d.tables.append(kv)
        return
    if text.strip():
        d.docs.append(RawDoc("text", stem, text.strip()))


def _read_html(stem: str, html: str, d: Digest):
    try:
        from bs4 import BeautifulSoup
    except ImportError:  # pragma: no cover
        d.docs.append(RawDoc("text", stem, re.sub(r"<[^>]+>", " ", html)))
        return
    soup = BeautifulSoup(html, "html.parser")
    for i, t in enumerate(soup.find_all("table")):
        grid = [[c.get_text(" ", strip=True) for c in tr.find_all(["td", "th"])] for tr in t.find_all("tr")]
        d.tables.extend(detect_tables(grid, f"{stem} Tabelle {i + 1}" if i else stem, {}))
    text = soup.get_text("\n", strip=True)
    if text:
        d.docs.append(RawDoc("text", soup.title.get_text(strip=True) if soup.title else stem, text))


# --- spreadsheets ---------------------------------------------------------------------------

def _archive_ok(data: bytes, d: Digest, what: str) -> bool:
    """An .xlsx or .docx is a ZIP too: the same limits as for a ZIP before a library unpacks it
    (a 124 KB workbook with one 120 MB string is a bomb, not a list)."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            total = 0
            for info in z.infolist()[:MAX_ZIP_MEMBERS * 2]:
                total += info.file_size
                if info.file_size > 20 * 1024 * 1024 and info.file_size / max(info.compress_size, 1) > MAX_ZIP_RATIO:
                    d.notes.append(f"{what} entpackt sich auf das {info.file_size // max(info.compress_size, 1)}-Fache "
                                   "– nicht gelesen (sieht nicht nach einem Export aus). Die Datei wurde aufbewahrt.")
                    return False
                if total > MAX_ZIP_TOTAL:
                    d.notes.append(f"{what} ist entpackt größer als {MAX_ZIP_TOTAL // (1024 * 1024)} MB – nicht gelesen.")
                    return False
    except zipfile.BadZipFile:
        return True  # not a ZIP at all: the reader says what is wrong with it
    return True


def _read_xlsx(stem: str, data: bytes, d: Digest):
    import openpyxl

    if not _archive_ok(data, d, f"Die Mappe „{stem}“"):
        return
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheets = [ws for ws in wb.worksheets if getattr(ws, "sheet_state", "visible") == "visible"] or wb.worksheets
    for ws in sheets:
        grid, formats = [], {}
        for r_idx, row in enumerate(ws.iter_rows()):
            vals = []
            for c_idx, cell in enumerate(row):
                v = getattr(cell, "value", None)
                vals.append(v)
                if v is not None and r_idx < 300 and c_idx not in formats:
                    nf = getattr(cell, "number_format", None)
                    if nf and nf != "General" and isinstance(v, (int, float)):
                        formats[c_idx] = nf
            grid.append(vals)
            if r_idx > 200000:
                d.notes.append(f"Blatt {ws.title}: nur die ersten 200.000 Zeilen gelesen.")
                break
        name = stem if len(sheets) == 1 else f"{stem} › {ws.title}"
        for t in detect_tables(grid, name, {"number_formats_by_col": formats, "locale": "de"}):
            t.sheet = ws.title
            d.tables.append(t)
    wb.close()


# --- table detection -----------------------------------------------------------------------

def _is_label(v) -> bool:
    if not isinstance(v, str):
        return False
    s = v.strip()
    if not s or len(s) > 60:
        return False
    if parse_number(s, "de") is not None or parse_number(s, "en") is not None:
        return False
    if parse_date(s) is not None:
        return False
    if "@" in s:
        return False
    return True


def _clean_grid(grid: list[list]) -> list[list]:
    out = []
    for row in grid:
        r = []
        for v in row:
            if isinstance(v, str):
                v = v.replace(" ", " ").strip()
                if v == "":
                    v = None
            r.append(v)
        out.append(r)
    width = 0
    for r in out:
        for j in range(len(r) - 1, -1, -1):
            if r[j] is not None:
                width = max(width, j + 1)
                break
    return [r[:width] + [None] * (width - len(r[:width])) for r in out]


def _nonempty(row) -> int:
    return sum(1 for v in row if v is not None)


def detect_tables(grid: list[list], name: str, meta: dict) -> list[RawTable]:
    if len(grid) > MAX_ROWS:
        meta = dict(meta, truncated=len(grid))
        grid = grid[:MAX_ROWS]
    datev = _datev_header(grid)
    rows = _clean_grid(grid)
    if not rows or not rows[0]:
        return []
    # split into blocks at runs of >= 2 empty rows
    blocks, cur, empty_run = [], [], 0
    for i, r in enumerate(rows):
        if _nonempty(r) == 0:
            empty_run += 1
            if empty_run >= 2 and cur:
                blocks.append(cur)
                cur = []
            continue
        empty_run = 0
        cur.append((i + 1, r))
    if cur:
        blocks.append(cur)
    tables = []
    pending_title: list[str] = []
    for bi, block in enumerate(blocks):
        t = _block_to_table(block, name, meta, pending_title)
        if t is None:
            # a block without a table (e.g. a title block) gives its text to the next table
            pending_title = [" ".join(str(v) for v in r if v is not None) for _, r in block][:5]
            continue
        pending_title = []
        tables.append(t)
    if len(tables) > 1:
        for k, t in enumerate(tables):
            if t.name == name:
                t.name = f"{name} ({k + 1})"
    for t in tables:
        if datev:
            t.meta.update(report="DATEV-Buchungsstapel", datev=datev, accounts=_report_accounts(t))
            continue  # the tax advisor's bookings: evidence about invoices, not things of their own
        if _open_items(t):
            t.meta.update(report="Offene-Posten-Liste", accounts=_report_accounts(t))
            continue
        _apply_debit_credit(t)
        _bank_layout(t)
        _timesheet_layout(t)
        if _is_crosstab(t.columns):
            t.meta["report"] = "Auswertung"  # months as columns: a pivot of other data, not records
    people = [t.meta["about"] for t in tables if t.meta.get("about") and "Pers.-Nr." in t.meta["about"]]
    if people:
        stem = name.split(" › ")[0]
        tables.append(RawTable(name=f"{stem} › Mitarbeiter", columns=["Pers.-Nr.", "Mitarbeiter"],
                               rows=[[a["Pers.-Nr."], a["Mitarbeiter"]] for a in people],
                               lines=[t.lines[0] if t.lines else 1 for t in tables if t.meta.get("about")
                                      and "Pers.-Nr." in t.meta["about"]],
                               sheet="Mitarbeiter", meta=dict(meta, title="Mitarbeiter laut Stundenzettel", derived="about")))
    return tables


_PERIOD_HEAD = re.compile(r"^(jan\w*|feb\w*|m[aä]r\w*|apr\w*|mai|may|jun\w*|jul\w*|aug\w*|sep\w*|okt\w*|oct\w*|"
                          r"nov\w*|dez\w*|dec\w*|q[1-4]|kw\s?\d{1,2}|\d{4}|\d{1,2}/\d{2,4})\.?$", re.I)


def _is_crosstab(columns: list) -> bool:
    heads = [str(c).strip() for c in columns if c is not None and str(c).strip()]
    periods = sum(1 for h in heads if _PERIOD_HEAD.match(h))
    return periods >= 3 and periods >= 0.5 * len(heads)


def _block_to_table(block, name, meta, pending_title) -> RawTable | None:
    width = max(len(r) for _, r in block)
    if width < 2:
        return None
    best_i, best_score = None, 0.0
    for i, (_, r) in enumerate(block[:15]):
        ne = _nonempty(r)
        if ne < 2:
            continue
        labels = sum(1 for v in r if _is_label(v))
        label_share = labels / ne
        fill = ne / width
        nxt = [x for _, x in block[i + 1:i + 6]]
        if not nxt:
            continue
        nxt_fill = sum(_nonempty(x) for x in nxt) / (len(nxt) * width)
        # rows below a header are usually *not* all labels
        nxt_label = sum(sum(1 for v in x if _is_label(v)) / max(1, _nonempty(x)) for x in nxt) / len(nxt)
        score = label_share * (0.4 + 0.6 * fill) * (0.5 + 0.5 * min(1.0, nxt_fill / max(fill, 0.01)))
        score *= 1.0 if nxt_label < 0.95 else 0.85
        if label_share >= 0.7 and fill >= 0.34 and score > best_score + 1e-9:
            best_i, best_score = i, score
    title_lines = list(pending_title)
    if best_i is None:
        # no header: invent column names, everything is data
        first = block[0][1]
        if _nonempty(first) < 2:
            return None
        header = [None] * width
        data_rows = block
    else:
        title_lines += [" ".join(str(v) for v in r if v is not None) for _, r in block[:best_i]]
        header = block[best_i][1]
        data_rows = block[best_i + 1:]
    # "Mitarbeiter: Jens Kortmann · Pers.-Nr. 103" above a time sheet: who every row is about
    about = _about(block[:best_i] if best_i else [], pending_title)
    # "Auftrag: A26-028" above a measurement: what job every position belongs to
    context = _context(block[:best_i] if best_i else [], pending_title, name) \
        if lexicon.concept_of(f"{name} {' '.join(str(x) for x in title_lines)}") == "aufmass" else {}
    # data rows: sections, sums, repeated headers
    norm_header = [str(h).strip().lower() if h is not None else None for h in header]
    rows, lines, sections, totals = [], [], [], 0
    current_section = None
    for line, r in data_rows:
        ne = _nonempty(r)
        if ne == 0:
            continue
        if [str(v).strip().lower() if v is not None else None for v in r] == norm_header:
            continue
        first_val = next((v for v in r if v is not None), None)
        if lexicon.is_total_label(first_val) and ne <= max(2, int(0.75 * width)):
            totals += 1
            continue
        if ne <= max(3, int(0.5 * width)) and any(lexicon.is_balance_label(v) for v in r):
            totals += 1  # "Anfangssaldo", "Endsaldo", "Übertrag": a statement's balance, not a booking
            continue
        if ne <= 2 and any(isinstance(v, str) and _SIGNATURE.match(v.strip()) for v in r):
            totals += 1  # "Unterschrift:" below a time sheet is a line to sign, not an entry
            continue
        if ne == 1 and width >= 3 and isinstance(r[0], str) and r[0] is first_val and _is_label(r[0]):
            current_section = r[0].strip()
            sections.append(current_section)
            continue
        rows.append(list(r) + ([current_section] if sections else []))
        lines.append(line)
    if sections:
        header = list(header) + ["Gruppe"]
        width += 1
        # rows collected before the first section: pad
    if not rows:
        return None
    rows = [r + [None] * (width - len(r)) for r in rows]
    if about:
        # with a personnel number the rows name the person by it; the person goes into a list of their own
        have = {fold(str(h or "")).strip() for h in header}
        ref = "Pers.-Nr." if "Pers.-Nr." in about else "Mitarbeiter"
        if fold(ref) not in have:
            header = list(header) + [ref]
            width += 1
            rows = [r + [about[ref]] for r in rows]
    if context:
        have = {fold(str(h or "")).strip() for h in header}
        for label, v in context.items():
            if fold(label) not in have:
                header = list(header) + [label]
                width += 1
                rows = [r + [v] for r in rows]
    # column names
    names, seen = [], Counter()
    keep = []
    for j in range(width):
        has_data = any(r[j] is not None for r in rows)
        h = header[j] if j < len(header) else None
        h = str(h).strip() if h is not None else ""
        h = re.sub(r"\s+", " ", h).rstrip(":").strip()
        if isinstance(header[j] if j < len(header) else None, (datetime, date)):
            h = header[j].strftime("%d.%m.%Y")
        if not has_data:
            continue
        if not h:
            h = f"Spalte {j + 1}"
        base = h
        while h in seen:  # "Datum", "Datum", "Datum (2)" -> three different names
            seen[base] += 1
            h = f"{base} ({seen[base] + 1})"
        seen[h] += 1
        names.append(h)
        keep.append(j)
    if len(keep) < 1:
        return None
    rows = [[r[j] for j in keep] for r in rows]
    nf_by_col = meta.get("number_formats_by_col") or {}
    number_formats = {str(k): nf_by_col[j] for k, j in enumerate(keep) if j in nf_by_col}
    tmeta = {k: v for k, v in meta.items() if k != "number_formats_by_col"}
    if about:
        tmeta["about"] = about
    tmeta.update({"title": " · ".join(x for x in title_lines if x)[:300], "number_formats": number_formats,
                  "sections": len(sections), "skipped_totals": totals, "header_row": block[best_i][0] if best_i is not None else None})
    return RawTable(name=name, columns=names, rows=rows, lines=lines, meta=tmeta)


# --- json ------------------------------------------------------------------------------------

def _flatten(obj, prefix="", depth=0) -> tuple[dict, dict]:
    flat, nested = {}, {}
    for k, v in obj.items():
        key = f"{prefix}{k}"
        if isinstance(v, dict) and depth < 3:
            f2, n2 = _flatten(v, key + ".", depth + 1)
            flat.update(f2)
            nested.update(n2)
        elif isinstance(v, list):
            if v and all(isinstance(x, dict) for x in v):
                nested[key] = v
            else:
                flat[key] = ", ".join(str(x) for x in v if x is not None)
        else:
            flat[key] = v
    return flat, nested


_ID_FIELDS = ("id", "number", "nummer", "nr", "no", "key", "code", "uuid")


_EXPORT_WORDS = {"export", "exports", "exported", "updated", "update", "since", "all", "to", "from", "bis", "von", "ab",
                 "seit", "data", "daten", "dump", "backup", "full", "delta", "neu", "new", "latest", "stand"}


def _parent_noun(name: str) -> str:
    """What a nested list's rows belong to, without the file's date or export words: the time entries of
    "tickets_export_2026-09-15.tickets" and of "tickets_updated_since_2026-09-10.tickets" both point at a
    "ticket_id", so both deliveries stay one list."""
    last = name.rsplit(".", 1)[-1]
    words = [w for w in re.split(r"[\s_\-]+", last) if w and not any(ch.isdigit() for ch in w)
             and w.lower() not in _EXPORT_WORDS]
    noun = "_".join(words) or last
    if noun.endswith("ies") and len(noun) > 4:
        return noun[:-3] + "y"
    return noun[:-1] if noun.endswith("s") and len(noun) > 3 and noun[-2] not in "sui" else noun


def _json_rows(items: list[dict], name: str, d: Digest, parent: tuple[str, list] | None = None):
    rows, cols, children = [], [], {}
    for idx, it in enumerate(items):
        flat, nested = _flatten(it)
        if parent:
            flat = {parent[0]: parent[1][idx], **flat}
        rows.append(flat)
        for k in flat:
            if k not in cols:
                cols.append(k)
        for k, v in nested.items():
            children.setdefault(k, []).append((idx, v))
    if not rows:
        return
    grid_rows = [[r.get(c) for c in cols] for r in rows]
    d.tables.append(RawTable(name=name, columns=cols, rows=grid_rows, lines=list(range(1, len(rows) + 1)),
                             meta={"format": "JSON"}))
    if children:
        id_col = next((c for c in cols if c.lower() in _ID_FIELDS or c.lower().endswith(("_id", "id", "number"))),
                      cols[0])
        parent_vals = [r.get(id_col) for r in rows]
        sg = _parent_noun(name)
        for k, lst in children.items():
            sub_items, keys = [], []
            for idx, v in lst:
                for x in v:
                    sub_items.append(x)
                    keys.append(parent_vals[idx])
            _json_rows(sub_items, f"{name}.{k}", d, parent=(f"{sg}_{id_col}", keys))


def _read_json(stem: str, data: bytes, d: Digest):
    text, _ = decode_text(data)
    try:
        obj = json.loads(text)
    except json.JSONDecodeError:
        # JSON lines
        objs = [json.loads(l) for l in text.splitlines() if l.strip().startswith("{")]
        obj = objs
    if isinstance(obj, list):
        if obj and all(isinstance(x, dict) for x in obj):
            _json_rows(obj, stem, d)
        elif obj:
            d.tables.append(RawTable(stem, ["Wert"], [[x] for x in obj], list(range(1, len(obj) + 1)), meta={}))
    elif isinstance(obj, dict):
        found = False
        for k, v in obj.items():
            if isinstance(v, list) and v and all(isinstance(x, dict) for x in v):
                _json_rows(v, f"{stem}.{k}" if len(obj) > 1 else stem, d)
                found = True
            elif isinstance(v, dict) and v and all(isinstance(x, dict) for x in v.values()):
                items = [{"id": kk, **vv} for kk, vv in v.items()]
                _json_rows(items, f"{stem}.{k}", d)
                found = True
        if not found:
            flat, _ = _flatten(obj)
            d.tables.append(RawTable(stem, list(flat), [list(flat.values())], [1], meta={"format": "JSON"}))


# --- email / pdf / docx / zip ---------------------------------------------------------------

def _read_eml(stem: str, data: bytes, d: Digest):
    msg = BytesParser(policy=policy.default).parsebytes(data)
    subject = str(msg.get("subject") or stem)
    when = None
    try:
        if msg.get("date"):
            when = parsedate_to_datetime(str(msg.get("date")))
            when = when.replace(tzinfo=None)
    except (TypeError, ValueError):
        when = None
    body = msg.get_body(preferencelist=("plain", "html"))
    text = ""
    if body is not None:
        try:
            text = body.get_content()
        except (LookupError, KeyError):
            text = body.get_payload(decode=True).decode("utf-8", errors="replace")
        if body.get_content_type() == "text/html":
            try:
                from bs4 import BeautifulSoup
                text = BeautifulSoup(text, "html.parser").get_text("\n", strip=True)
            except ImportError:  # pragma: no cover
                text = re.sub(r"<[^>]+>", " ", text)
    for part in msg.iter_attachments():
        fn = part.get_filename()
        payload = part.get_payload(decode=True)
        if fn and payload:
            d.attachments.append((fn, payload))
    d.docs.append(RawDoc("email", subject, text.strip(), sender=str(msg.get("from") or "") or None,
                         recipients=str(msg.get("to") or "") or None, date=when,
                         meta={"attachments": [a for a, _ in d.attachments]}))


def _read_xml(stem: str, data: bytes, d: Digest):
    try:
        from defusedxml.ElementTree import fromstring  # no entity expansion, no external entities
    except ImportError:  # pragma: no cover
        from xml.etree.ElementTree import fromstring
    from .einvoice import parse, to_tables

    root = fromstring(data)
    if _camt_body(root) is not None:
        camt = _camt_table(root, stem)
        if camt is not None:
            d.tables.append(camt)
        else:
            d.notes.append(f"{stem}: ein Kontoauszug ohne Buchungen")  # a day's statement with balances only
        return
    inv = parse(root)
    if inv and inv.get("number"):
        head, lines, doc = to_tables(inv, stem)
        d.tables.append(head)
        if lines is not None:
            d.tables.append(lines)
        d.docs.append(doc)
        return
    text = " ".join(t.strip() for t in root.itertext() if t and t.strip())
    if text:
        d.docs.append(RawDoc("xml", stem, text[:200000]))


_CAMT_BODY = {"BkToCstmrStmt", "BkToCstmrAcctRpt", "BkToCstmrDbtCdtNtfctn"}


def _camt_body(root):
    def loc(el):
        return el.tag.rsplit("}", 1)[-1]
    return next((x for x in root if loc(x) in _CAMT_BODY), None) if loc(root) == "Document" else None


def _camt_table(root, stem: str) -> RawTable | None:
    """ISO 20022 bank statements (camt.053 statements, camt.052 reports, camt.054 notices), as banking
    software and many online banks export them: the bookings become the table a CSV export would give –
    one row per transaction (a batch splits into its transactions), the amount signed, the other party
    named as payer or payee."""
    def loc(el):
        return el.tag.rsplit("}", 1)[-1]

    def first(el, *path):
        cur = [el]
        for name in path:
            cur = [x for c in cur for x in c if loc(x) == name]
            if not cur:
                return None
        return cur[0]

    def txt(el, *path):
        x = first(el, *path) if el is not None else None
        return x.text.strip() if x is not None and x.text and x.text.strip() else None

    def day(s):
        try:
            return date.fromisoformat(s[:10]) if s else None
        except ValueError:
            return None

    body = _camt_body(root)
    if body is None:
        return None
    cols = ["Buchungstag", "Valutadatum", "Buchungstext", "Verwendungszweck", "Zahlungspartner", "IBAN", "BIC", "Betrag",
            "Währung", "Kundenreferenz (End-to-End)", "Mandatsreferenz", "Gläubiger-ID"]
    rows, owner, iban, accounts = [], None, None, []
    for st in (x for x in body if loc(x) in ("Stmt", "Rpt", "Ntfctn")):
        acct = txt(st, "Acct", "Id", "IBAN") or txt(st, "Acct", "Id", "Othr", "Id")
        iban = iban or acct
        owner = owner or txt(st, "Acct", "Ownr", "Nm")
        for e in (x for x in st if loc(x) == "Ntry"):
            if (txt(e, "Sts", "Cd") or txt(e, "Sts") or "BOOK") != "BOOK":
                continue  # pending or informational entries are no bookings yet
            sign = -1 if txt(e, "CdtDbtInd") == "DBIT" else 1
            booked = day(txt(e, "BookgDt", "Dt") or txt(e, "BookgDt", "DtTm"))
            value = day(txt(e, "ValDt", "Dt") or txt(e, "ValDt", "DtTm"))
            cur = first(e, "Amt")
            txs = [t for nd in e if loc(nd) == "NtryDtls" for t in nd if loc(t) == "TxDtls"] or [None]
            # a batch that names no amount per transaction is one booking of its total (never n times the total)
            lump = len(txs) > 1 and not any(txt(t, "Amt") or txt(t, "AmtDtls", "TxAmt", "Amt") for t in txs)
            mine = []
            for tx in txs:
                amt = (txt(tx, "Amt") or txt(tx, "AmtDtls", "TxAmt", "Amt")) if tx is not None else None
                if amt is None or len(txs) == 1:
                    amt = txt(e, "Amt") or amt  # one transaction: what was booked (after charges) is what moved
                ind = txt(tx, "CdtDbtInd") if tx is not None else None
                s = (-1 if ind == "DBIT" else 1) if ind in ("DBIT", "CRDT") else sign
                side = "Dbtr" if s > 0 else "Cdtr"  # the other party: who paid us, whom we paid
                name = (txt(tx, "RltdPties", "Ultmt" + side, "Pty", "Nm") or txt(tx, "RltdPties", side, "Pty", "Nm")
                        or txt(tx, "RltdPties", side, "Nm"))
                ib = txt(tx, "RltdPties", side + "Acct", "Id", "IBAN")
                bic = txt(tx, "RltdAgts", side + "Agt", "FinInstnId", "BICFI") or \
                    txt(tx, "RltdAgts", side + "Agt", "FinInstnId", "BIC")
                purpose = " ".join(x.text.strip() for x in (tx.iter() if tx is not None else [])
                                   if loc(x) == "Ustrd" and x.text and x.text.strip())
                ref = txt(tx, "RmtInf", "Strd", "CdtrRefInf", "Ref")
                if ref and ref not in purpose:
                    purpose = (purpose + " " + ref).strip()
                creditor = txt(tx, "RltdPties", "Cdtr", "Pty", "Id", "PrvtId", "Othr", "Id") or \
                    txt(tx, "RltdPties", "Cdtr", "Id", "PrvtId", "Othr", "Id")
                try:
                    money = round(s * float(amt), 2) if amt else None
                except ValueError:
                    money = None
                mine.append([booked, value, txt(tx, "AddtlTxInf") or txt(e, "AddtlNtryInf"), purpose or None, name, ib,
                             bic, money, cur.get("Ccy") if cur is not None else None, txt(tx, "Refs", "EndToEndId"),
                             txt(tx, "Refs", "MndtId"), creditor])
            if lump:  # its parties and purposes all together – not the first one's, as if the total were theirs
                same = [vs.pop() if len(vs := {r[i] for r in mine}) == 1 else None for i in range(len(cols))]
                same[3] = " / ".join(dict.fromkeys(r[3] for r in mine if r[3])) or None
                same[7] = mine[0][7]
                mine = [same]
            rows += mine
            accounts += [acct] * len(mine)
    if not rows:
        return None
    if len(set(accounts)) > 1:  # the statements of several accounts in one file: which account a booking is on
        cols, rows = ["Auftragskonto"] + cols, [[a] + r for a, r in zip(accounts, rows)]
        iban = None
    return RawTable(name=f"Kontoauszug {iban or stem}", columns=cols, rows=rows, lines=list(range(1, len(rows) + 1)),
                    meta={"format": "camt", "title": f"Kontoauszug {owner or ''} · IBAN {iban or ''}".strip()})


_MT940 = re.compile(r"^:20:.*\n(?:.*\n){0,6}?:25:", re.M)
_MT61 = re.compile(r"^(\d{6})(\d{4})?(R?[CD])[A-Z]?(\d+,\d{0,2})(?:[NF][A-Z0-9]{3})?(.*)$")
_SEPA_TAG = re.compile(r"\b(EREF|MREF|CRED|SVWZ|KREF|ABWA|ABWE|DEBT|COAM|OAMT|IBAN|BIC)\+")


def _mt940_table(text: str, stem: str) -> RawTable | None:
    """SWIFT MT940 statements, as banking programs (StarMoney, SFirm, Hibiscus) and many online banks export
    them (.sta): the bookings become the table a CSV export gives, columns as for camt – one row per :61:
    line, the amount signed, the :86: field split into booking text, purpose, the other side's name, IBAN
    and BIC, and the SEPA references of the purpose (EREF, MREF, CRED)."""
    cols = ["Buchungstag", "Valutadatum", "Buchungstext", "Verwendungszweck", "Zahlungspartner", "IBAN", "BIC", "Betrag",
            "Währung", "Kundenreferenz (End-to-End)", "Mandatsreferenz", "Gläubiger-ID"]
    fields, cur = [], None
    for line in text.splitlines():
        m = re.match(r"^:(\d{2}[A-Z]?):(.*)$", line)
        if m:
            cur = [m.group(1), m.group(2)]
            fields.append(cur)
        elif line.strip() == "-" or line.startswith("-}"):
            cur = None
        elif cur is not None:
            cur[1] += line if cur[0] == "86" else "\n" + line
    rows, account, currency, last, acct, accounts = [], None, "EUR", None, None, []
    for tag, val in fields:
        if tag == "25":
            acct = val.strip()
            account = account or acct
        elif tag in ("60F", "60M"):
            currency = val[7:10] or currency
        elif tag == "61":
            m = _MT61.match(val.split("\n")[0].strip())
            if not m:
                last = None
                continue
            try:
                value = date(2000 + int(m.group(1)[:2]), int(m.group(1)[2:4]), int(m.group(1)[4:6]))
                booked = value
                if m.group(2):
                    booked = date(value.year, int(m.group(2)[:2]), int(m.group(2)[2:]))
                    if (booked - value).days > 180:  # booked in December for a value in January, or the reverse
                        booked = booked.replace(year=booked.year - 1)
                    elif (value - booked).days > 180:
                        booked = booked.replace(year=booked.year + 1)
            except ValueError:
                last = None
                continue
            sign = -1 if m.group(3) in ("D", "RC") else 1
            amount = round(sign * float(m.group(4).replace(",", ".")), 2)
            ref = m.group(5).split("//")[0].strip()
            last = [booked, value, None, None, None, None, None, amount, currency,
                    ref if ref and ref != "NONREF" else None, None, None]
            rows.append(last)
            accounts.append(acct)
        elif tag == "86" and last is not None:
            sub = dict.fromkeys(("00", "30", "31", "32", "33"), "")
            purpose = []
            if re.match(r"^\d{3}\?", val):  # structured: GVC, then ?NN subfields
                for code, part in re.findall(r"\?(\d{2})([^?]*)", val):
                    if code in sub:
                        sub[code] += part
                    elif "20" <= code <= "29" or "60" <= code <= "63":  # a SEPA tag opens its own subfield
                        purpose.append((" " if _SEPA_TAG.match(part) else "") + part)
            else:
                purpose.append(val)
            text86 = "".join(purpose)
            tags = {}
            pieces = _SEPA_TAG.split(text86)
            if len(pieces) > 1:
                for k in range(1, len(pieces) - 1, 2):
                    tags[pieces[k]] = pieces[k + 1].strip()
                text86 = tags.get("SVWZ") or " ".join(f"{k}+{v}" for k, v in tags.items())
            last[2] = sub["00"].strip() or None
            last[3] = text86.strip() or None
            last[4] = (sub["32"] + sub["33"]).strip() or tags.get("ABWA") or None
            last[5] = sub["31"].strip() if len(sub["31"].strip()) >= 15 else None
            last[6] = sub["30"].strip() if re.match(r"^[A-Z]{6}[A-Z0-9]{2,5}$", sub["30"].strip()) else None
            last[9] = tags.get("EREF") or last[9]
            last[10] = tags.get("MREF")
            last[11] = tags.get("CRED")
            last = None
    if not rows:
        return None
    if len(set(accounts)) > 1:  # several accounts in one file: which account a booking is on (as for camt)
        cols, rows, account = ["Auftragskonto"] + cols, [[a] + r for a, r in zip(accounts, rows)], None
    return RawTable(name=f"Kontoauszug {account or stem}", columns=cols, rows=rows, lines=list(range(1, len(rows) + 1)),
                    meta={"format": "mt940", "title": f"Kontoauszug · Konto {account or ''}".strip()})


def _pdf_attachments(data: bytes, d: Digest):
    """ZUGFeRD/Factur-X: the structured invoice travels inside the PDF."""
    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        for name, contents in (reader.attachments or {}).items():
            for content in contents:
                if name.lower().endswith(".xml") and content:
                    d.attachments.append((name, content))
    except Exception:
        pass


def _read_pdf(stem: str, data: bytes, d: Digest, sender: str | None = None):
    _pdf_attachments(data, d)
    texts, grids = [], []
    try:
        import pdfplumber

        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for page in pdf.pages[:200]:
                texts.append(page.extract_text() or "")
                for t in page.extract_tables() or []:
                    if t and len(t) >= 3 and max(len(r) for r in t) >= 2:
                        grids.append(t)
    except ImportError:  # pragma: no cover
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        texts = [p.extract_text() or "" for p in reader.pages[:200]]
    for i, g in enumerate(grids):
        d.tables.extend(detect_tables(g, f"{stem} Tabelle {i + 1}", {"locale": "de"}))
    text = "\n".join(texts).strip()
    first = next((l.strip() for l in text.splitlines() if l.strip()), stem)
    meta = {"first_line": first[:120]}
    if not any(n.lower().endswith(".xml") for n, _ in d.attachments):  # an e-invoice inside says it better
        from . import pdfinvoice
        inv = pdfinvoice.read(text, sender)
        if inv:
            head, lines = pdfinvoice.to_tables(inv, stem)
            d.tables.append(head)
            if lines is not None:
                d.tables.append(lines)
            meta["invoice"] = dict(number=inv["number"], gross=inv["gross"], seller=inv["seller"])
    d.docs.append(RawDoc("pdf", stem, text, meta=meta))
    if not text:
        d.notes.append(f"{stem}.pdf enthält keinen Text (gescannt?) – aufbewahrt, aber nicht ausgelesen.")


def _read_docx(stem: str, data: bytes, d: Digest):
    import docx

    if not _archive_ok(data, d, f"Das Dokument „{stem}“"):
        return
    doc = docx.Document(io.BytesIO(data))
    paras = [p.text for p in doc.paragraphs if p.text.strip()]
    for i, t in enumerate(doc.tables):
        grid = [[c.text for c in row.cells] for row in t.rows]
        d.tables.extend(detect_tables(grid, f"{stem} Tabelle {i + 1}" if i else stem, {"locale": "de"}))
    if paras:
        d.docs.append(RawDoc("docx", stem, "\n".join(paras)))


def _read_zip(data: bytes, d: Digest):
    total = 0
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        for info in z.infolist()[:MAX_ZIP_MEMBERS]:
            if info.is_dir() or "__MACOSX" in info.filename or PurePath(info.filename).name.startswith("."):
                continue
            if info.compress_size and info.file_size > 20 * 1024 * 1024 and \
                    info.file_size / max(info.compress_size, 1) > MAX_ZIP_RATIO:
                d.notes.append(f"„{info.filename}“ entpackt sich auf das {info.file_size // max(info.compress_size, 1)}-Fache "
                               "– übersprungen (sieht nicht nach einem Export aus).")
                continue
            total += info.file_size
            if total > MAX_ZIP_TOTAL:
                d.notes.append("ZIP zu groß – nur ein Teil wurde gelesen.")
                break
            d.attachments.append((info.filename, z.read(info)))
