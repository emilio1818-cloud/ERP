#!/usr/bin/env python3
"""Kerkhoff Hausverwaltung GmbH, Brakenfeld – eine erfundene Hausverwaltung (WEG- und Mietverwaltung, 5 Leute,
14 Objekte, 226 Einheiten plus Stellplätze: 8 Miethäuser für 7 Eigentümer, 6 WEG, 5 vermietete Eigentumswohnungen
in Sondereigentumsverwaltung) mit zwei Lieferungen, so wie die Ablage eines solchen Büros wirklich aussieht.
Stadt, Firmen, Menschen, Banken, IBANs und Domains sind erfunden (private Mails unter .example).

    python3 hausverwaltung.py OUTDIR  ->  OUTDIR/lieferung1/   Stand Fr 18.09.2026 (Daten bis Do 17.09.)
                                         OUTDIR/lieferung2/   Stand Fr 30.10.2026 (Daten bis Do 29.10.), kompletter
                                                              neuer Ordnerstand (alte Exporte liegen weiter darin)
                                         OUTDIR/truth.json    Format des Prüfstands (bench/one.py)
                                         OUTDIR/answer_key.json  Sollstellungen, Geldbewegungen, Mietverhältnisse

Fester Seed; nur Standardbibliothek + openpyxl + reportlab (+ pypdf für ZUGFeRD).

Die Dateien: Objekte/Objektliste.xlsx (Blatt „beendet“), Objekte/Verwaltervertraege.xlsx, Mieter/Mieterliste_Export_<Tag>.csv
(Export des Verwaltungsprogramms, cp1252, nur laufende Mietverhältnisse), Mieter/Sollmieten_2026.xlsx (Arbeitsliste von
Frau Okafor, ein Blatt je Objekt, jedes Blatt anders, Summenzeilen; L1 Stand 01.08., L2 Stand 01.10.) und
Mieter/Sollmieten_2026_alt.xlsx (Stand Januar, vergessen), Mieter/Mietanpassungen_2026.xlsx, Mieter/Kautionen.xlsx,
Mieter/OP-Liste_Mieter_Stand_<Tag>.xlsx, Mieter/Mahnungen/*.pdf, WEG/Eigentuemerliste_WEG.xlsx (ein Blatt je WEG),
WEG/Jahresabrechnung_2025_Ergebnisse.xlsx, WEG/Sonderumlage_Brunnenstrasse6.xlsx, WEG/OP-Liste_Eigentuemer_Stand_<Tag>.xlsx,
WEG/Uebergabe_Vorverwalter_Feldkamp.xlsx, WEG/Protokolle/*.pdf, Abrechnung/BK-Abrechnung_2025_Ergebnisse.xlsx (+ zwei
Abrechnungen als PDF), Buchhaltung/Rechnungsausgang_2026_<Tag>.csv (eigene Rechnungen, Status „laut Datei“ bis 31.08.
bzw. 15.10.), Buchhaltung/Sonderleistungen_2026.xlsx, Buchhaltung/Rechnungen/*.pdf (einige eigene Rechnungen),
Bank/ – neun Konten in drei Formaten: Geschäftskonto (Volksbank-„Umsatzanzeige“-CSV, Name mit Stand), Treuhandkonto
Mietverwaltung und die WEG-Konten Am Wall, Brunnenstraße, Feldkamp, Mühlenhof (Sparkassen-CSV-CAMT, Name mit Stand),
Mühlenteich GbR, WEG Rosenhof, WEG Parkresidenz (camt.053 je abgeschlossenem Monat, camt.052 für den laufenden),
Bank/Lastschriften/*.xml (pain.008 der Mieteinzüge und des Hausgeldeinzugs Rosenhof), Eingangsrechnungen/<Monat>/
(Handwerker-PDFs, ZUGFeRD, XRechnung, an „WEG … c/o Kerkhoff“ adressiert), Mails/*.eml.

DAS FORMAT DER WAHRHEIT. Mieten, Hausgeld, Umlagen und Abrechnungsergebnisse haben keine Rechnungsnummer: geschuldet
wird nach Vertrag, Wirtschaftsplan oder Beschluss. Ein offener Posten ist hier eine Sollstellung, benannt wie im
Verwaltungsprogramm: „<Personenkonto> <Sollart> <Zeitraum>“, z. B. „10612 Miete 07/2026“, „20404 Hausgeld 05/2026“,
„20404 Sonderumlage 2026“, „20404 Abrechnung 2025“, „10304 BK 2025“, „10115 Kaution Rate 2“; die eigenen Rechnungen
(Verwaltervergütung, Sonderleistungen) mit ihrer Nummer. Fällig sind Miete und Hausgeld am 3. Werktag (Samstag zählt
nicht), Kautionsraten mit der Miete, BK-Nachzahlungen 30 Tage nach Zugang, Abrechnungsspitzen vier Wochen nach der
Versammlung. Geschuldet ist, was der Vertrag sagt – auch eine Staffel, die das Büro nicht eingetragen hat (H05).
„nicht berechnet“ kennt nur die eigenen Leistungen (Verwaltervertrag ohne Rechnung, Sonderleistung ohne Rechnung);
eine Miete wird nie berechnet. „Geld ohne Rechnung“ ist Geld, das keine Sollstellung und keine Rechnung tilgt und
geklärt werden muss. Eine Zahlung vor dem Monat, für den sie ist (Dauerauftrag zum 28., das Jobcenter zahlt im
Voraus), ist eine Vorauszahlung, kein Geld ohne Rechnung. Was keine Sollstellung ist, aber Arbeit (Mängel, Fristen,
zu zahlende Handwerkerrechnungen, Betrug), steht in truth.json unter payables, deadlines, mails_waiting, risks.

ANSWER KEY – die Fallen (Zahlen in truth.json / answer_key.json; check_traps() prüft die hier genannten)
Mieten (Treuhandkonto Mietverwaltung = Sparkassen-CSV; Mühlenteich GbR = camt)
H01 Sollmiete = Kaltmiete + BK-Vorauszahlung + HK-Vorauszahlung + Stellplatz; im Programm-Export als Spalten, in der
    Arbeitsliste je Objekt anders überschrieben, teils ohne Gesamtspalte, teils mit Summenzeile.
H02 Totalausfall: 10612 Brandl (Gartenstraße 31) zahlt ab Juli nichts (745,00). L1 offen Juli, August, September,
    überfällig (2 Monatsmieten Rückstand: Kündigungsgrund § 543 BGB). Mahnung 20.08. (PDF). L2: Zahlungszusage (Mail
    24.09.), 1.000,00 am 01.10. ohne Monatsangabe -> älteste zuerst: Juli bezahlt, August 490,00, September und
    Oktober je 745,00 offen.
H03 Teilzahlung jeden Monat: 10305 Böhm (Mühlenteich, camt) zahlt seit Juni 600,00 statt 812,00 mit „Miete Juni
    Teilzahlung“: je Monat 212,00 offen (L1 Juni–September, L2 bis Oktober); seine BK-Nachzahlung 2025 (91,20) bleibt
    ebenfalls offen. Mail 05.10.: Ratenvorschlag. Zahlungserinnerung 10.09. (PDF).
H04 Dauerauftrag mit altem Betrag nach Staffel: 10104 Hartwig, Staffel ab 01.07. 610 -> 645 kalt (Sollmieten und
    Mietanpassungen kennen sie), zahlt weiter 860,00: je Monat 35,00 offen ab Juli.
H05 Staffel, die das Büro vergessen hat: 10206 (Lindenallee 16) Staffel ab 01.09. +40,00 laut Mietanpassungen_2026.xlsx,
    nicht im Programm-Export, nicht in der Arbeitsliste: geschuldet nach Vertrag -> 40,00 je Monat ab September offen.
H06 BK-Vorauszahlung nach der Abrechnung erhöht (§ 560 BGB), Dauerauftrag nicht angepasst: 10108 +20,00 und 10210
    +15,00 ab September -> je Monat offen.
H07 Jobcenter zahlt direkt (§ 22 Abs. 7 SGB II), im Voraus um den 28., ein Zahler für mehrere Mieter, Zweck mit
    BG-Nummer und Namen: 10603 voll; 10609 Sadiku: Jobcenter 560,00, Rest 38,50 zahlt der Mieter – August und
    September fehlen (L1 offen je 38,50), am 05.10. „Restmiete Aug Sep Okt“ 115,50 -> L2 bezahlt; 10614 Wagner: das
    Jobcenter zahlt ab September nicht (Weiterbewilligung) -> L1 September 612,00 offen; am 12.10. 1.224,00 „KDU
    09+10/2026“ -> L2 bezahlt; die November-Zahlung vom 28.10. ist Vorauszahlung.
H08 Jobcenter zahlt für einen Ausgezogenen: 10611 Krüger zum 31.08. ausgezogen (Kündigung per Mail 28.05.), das
    Jobcenter überweist am 28.08. die Septembermiete 534,00 -> L1 Geld ohne Rechnung (zurückzahlen); am 22.09.
    zurücküberwiesen -> L2 nichts. Wohnung 11 steht leer, Nachmieter 10619 ab 01.11. (Kaution Rate 1 am 27.10. vorab).
H09 Rücklastschrift: 10802 (Heideweg, Mieteinzug per Sammellastschrift, Einzelposten nur in der pain.008) September
    AM04 zurück am 08.09., Gebühr 3,00 in derselben Zeile -> L1 offen 689,00; am 21.09. überwiesen -> L2 bezahlt.
H10 Doppelt gezahlt: 10107 Hartmann, Dauerauftrag 01.09. und Überweisung 07.09. je 690,00 -> L1 Geld ohne Rechnung
    690,00 (Mail 15.09.); verrechnet mit Oktober (Mail 22.09., Dauerauftrag Oktober ausgesetzt) -> L2 nichts offen,
    kein Geld ohne Rechnung.
H11 Mieterwechsel: 10105 Voigt zum 31.07. ausgezogen (Leerstand bis 14.08.), 10115 Demir ab 15.08.: August anteilig
    496,29, zusammen mit Kaution Rate 1 in einer Überweisung (1.146,29); Kautionsrate 2 fehlt -> L1 offen „10115 Kaution
    Rate 2“ 650,00; am 12.10. Rate 2+3 -> L2 nichts. Kaution Voigt: Kautionskonto am 14.09. aufgelöst (+1.833,85, kein
    Geld ohne Rechnung), 15.09. 1.533,85 an Voigt (300,00 Einbehalt für die BK 2026).
H12 Leerstand 10615 (Gartenstraße, Sanierung seit 01.06.): keine Sollstellung, kein Rückstand.
H13 BK-Abrechnung 2025: Nachzahlungen fällig 30 Tage nach Zugang. 10304 318,40 unbezahlt (L1 überfällig, 02.10. bezahlt);
    10205 412,90 in Raten 150 + 150 (L1 offen 112,90), Rest 21.09.; 10703 264,10 bestritten (Widerspruch 03.06., erst
    Belege) – L1 und L2 offen und überfällig, aber nicht mahnen. Guthaben werden ausgezahlt (Ausgang) oder verrechnet:
    10502 zahlt Oktober „abzgl. Guthaben BK 2025“ 86,40 weniger -> nichts offen. Gartenstraße 31 (106) ist noch nicht
    abgerechnet: Frist 31.12.2026 (§ 556 Abs. 3 BGB) – eine Frist, kein offener Posten.
H14 Gewerbe mit Umsatzsteuer: 10407 Café Kornblume 1.450,00 + 280,00 NK je zzgl. 19 % = 2.058,70; September erst am
    12.10. mit Oktober (4.117,40) -> L1 offen 2.058,70, L2 bezahlt. 10408 Physiotherapie (steuerfrei, keine Option):
    Indexmiete ab 01.10. 980,00 -> 1.000,58, zahlt alt -> L2 offen 20,58.
H15 Mieterhöhung § 558 BGB zum 01.09. mit Zustimmung (10303, 540 -> 580): ab September bezahlt; vorher ist der alte
    Betrag richtig – keine Teilzahlung.
H16 Stellplatz an einen Nachbarn (10325 Kowalczyk, kein Mieter im Haus): 55,00 im Monat, zahlt vierteljährlich im
    Voraus 165,00 – kein Rückstand im August/September.
H17 Zahler heißt anders: 10601 Yılmaz zahlt vom Konto „YILMAZ MEHMET“; 10507 schreibt „Mieternr 10570“ (Zahlendreher).
H18 Gleichnamig, nicht dieselben: „Neumann, Thomas“ 10209 (Lindenallee 16) und 10804 (Heideweg); „Schulz, Andrea“ 10106
    und „Schulz, Andreas“ 10319.
H19 Mietverwaltung Heideweg (108) endet am 30.09. (Verkauf): keine Sollstellungen und keine Rechnung ab Oktober (nicht
    „nicht berechnet“). 10817 zahlt am 01.10. trotzdem 655,00 aufs Treuhandkonto -> L2 Geld ohne Rechnung (weiterleiten).
H20 Dauerauftrag zum 28. bezahlt den Folgemonat: Zahlungen vom 28.10. sind Vorauszahlungen für November.
WEG (Hausgeld nach Wirtschaftsplan 2026; Rosenhof und Parkresidenz per Sammellastschrift)
H21 Hausgeld-Totalausfall: 20404 Hollmann (Brunnenstraße 6) zahlt seit Mai nichts (312,00): L1 offen Mai–September,
    dazu Sonderumlage 3.120,00 und Abrechnung 2025 188,40; L2 zusätzlich Oktober. Beschluss Klage am 08.10.
H22 Sonderumlage Brunnenstraße 6 (Dach, Beschluss 23.07., 24.000,00 nach MEA, fällig 01.09.): 20402 Kranz zahlt nach
    Absprache 1.560,00 am 01.09., Rest 01.10. -> L1 offen 1.560,00, L2 bezahlt.
H23 Rücklastschrift mit Widerspruch (MD06) nach Wochen: 20127 Dr. Riedel widerspricht am 14.09. dem Juli-Einzug ->
    L1 Juli offen 287,00; am 02.10. überwiesen -> L2 bezahlt.
H24 Eigentümerwechsel 203-11 zum 01.09.: Kowalski (20311) zahlt September trotzdem (268,00 am 01.09.) -> L1 Geld ohne
    Rechnung; am 25.09. erstattet. Hoffmeister (20361) zahlt ab September per Lastschrift.
H25 Ein Eigentümer, zwei Wohnungen, eine Überweisung: 20603 + 20604 Petersen („Hausgeld Whg 3 + 4“).
H26 Abrechnung 2025: 20211 zahlt 200,00 von 342,75 -> L1 offen 142,75, Rest 15.10.
H27 Neuer Verwaltungsbeirat, neue WEG: Feldkamp 17–19 (205) seit 01.07., Konto erst ab 01.07. im Export; Rückstand aus
    der Zeit des Vorverwalters laut Übergabeliste: 20514 Hausgeld 05/2026 und 06/2026 je 276,00 -> L1 offen, am 14.10.
    bezahlt.
H28 Hausgeld für die vermieteten Eigentumswohnungen zahlt das Büro selbst vom Treuhandkonto auf die WEG-Konten
    (Zahler „Kerkhoff Hausverwaltung GmbH Treuhandkonto“): bezahlt, kein Geld ohne Rechnung, keine Lieferantenzahlung.
H29 Versicherungsleistung Wasserschaden Rosenhof (3.480,00, 02.09.) und Zuführungen zur Erhaltungsrücklage: kein Geld
    ohne Rechnung, keine Aufgabe.
Eigene Rechnungen (Verwaltervergütung monatlich je Mandat, Entnahme vom Objektkonto)
H30 Die Vergütung wird vom Objektkonto aufs Geschäftskonto überwiesen: dasselbe Geld steht zweimal im Ordner (Ausgang
    beim Objekt, Eingang beim Büro) – die Rechnung ist einmal bezahlt.
H31 Sammelentnahme vom Treuhandkonto mit Bereich: „Verwalterverguetung 09/2026 RE-2026-0xxx bis RE-2026-0yyy“ bezahlt
    alle zehn Rechnungen dazwischen.
H32 Brunnenstraße 6 hat kein Geld: Vergütung August und September nicht entnommen -> L1 offen und überfällig; am
    05.10. mit Oktober entnommen -> L2 bezahlt.
H33 Dr. Albers zahlt selbst, vierteljährlich und spät: L1 Juli–September offen und überfällig, L2 (08.10. bezahlt)
    Oktober offen und überfällig.
H34 Feldkamp: Verwaltervertrag ab 01.07., im Buchhaltungsprogramm vergessen -> L1 nicht berechnet „VV-205“; am 05.10.
    vier Rechnungen Juli–Oktober, am 06.10. entnommen -> L2 nichts.
H35 Sonderleistungen: SL-26-011 Verwalterzustimmung 203-11 (20.08.) nie berechnet -> L1 und L2 nicht berechnet;
    SL-26-013 außerordentliche Versammlung Brunnenstraße (08.10.) -> L2 nicht berechnet; SL-26-009 berechnet, Liste
    nicht nachgetragen; SL-26-012 Kulanz (0,00); SL-26-014 Baubetreuung läuft noch – keines davon „nicht berechnet“.
    SL-26-010 Zustimmung 201-25 an die Verkäuferin berechnet, unbezahlt -> offen, überfällig.
Eingangsrechnungen, Mails
H36 Handwerkerrechnungen sind an „WEG … c/o Kerkhoff Hausverwaltung“ oder den Eigentümer „vertreten durch“ adressiert
    – richtiger Empfänger; bezahlt wird vom Objektkonto.
H37 Falk Bedachungen (Dach Brunnenstraße): Abschlag 9.200,00 (L1 zu zahlen, nach der Umlage), Schlussrechnung (L2 zu
    zahlen). Betrug L2: Mail von falk-bedachung.example (echt: falk-bedachungen.example), neue litauische IBAN,
    „korrigierte Schlussrechnung“.
H38 Rehm Heizungsbau (ZUGFeRD) kommt als Datei und als Mail-Anhang: einmal bezahlen.
H39 Neue Bankverbindung eines Eigentümers für die Auszahlung (Meinhardt, Mail 16.09.): echt, aber vor der nächsten
    Auszahlung bestätigen.
H40 Mails, die warten: L1 Schimmel 10606 (03.09.), Doppelzahlung 10107 (15.09.), Widerspruch Lastschrift Riedel (14.09.),
    Jobcenter-Lücke 10614 (10.09.), neue Bankverbindung Meinhardt (16.09.); L2 Kündigung 10313 zum 31.12. (02.10.),
    Ratenvorschlag 10305 (05.10.), Frage Feldkamp (12.10.), Wasserschaden Heideweg (20.10., nicht mehr zuständig),
    Anfrage Verwaltungsangebot (28.10.).
"""
from __future__ import annotations

import calendar
import hashlib
import io
import json
import random
import sys
from datetime import date, datetime, timedelta, timezone
from email.message import EmailMessage
from email.utils import format_datetime
from pathlib import Path
from xml.sax.saxutils import escape as xesc

import openpyxl
from openpyxl.styles import Font
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas

SEED = 20261030


def rr(*key) -> random.Random:
    h = hashlib.sha256(("|".join(map(str, key)) + f"|{SEED}").encode()).hexdigest()
    return random.Random(int(h[:16], 16))


def D(s: str) -> date:
    return date.fromisoformat(s)


# ------------------------------------------------------------------ calendar (Niedersachsen 2026)
HOL = {D(x) for x in ("2025-12-25", "2025-12-26", "2026-01-01", "2026-04-03", "2026-04-06", "2026-05-01", "2026-05-14",
                      "2026-05-25", "2026-10-03", "2026-10-31", "2026-12-25", "2026-12-26", "2027-01-01")}
MONATE = ["Januar", "Februar", "März", "April", "Mai", "Juni", "Juli", "August", "September", "Oktober", "November", "Dezember"]
MON3 = ["Jan", "Feb", "Mär", "Apr", "Mai", "Jun", "Jul", "Aug", "Sep", "Okt", "Nov", "Dez"]


def is_bd(d: date) -> bool:
    return d.weekday() < 5 and d not in HOL


def next_bd(d: date) -> date:
    while not is_bd(d):
        d += timedelta(days=1)
    return d


def prev_bd(d: date) -> date:
    while not is_bd(d):
        d -= timedelta(days=1)
    return d


def nth_bd(y: int, m: int, n: int) -> date:
    d, k = date(y, m, 1), 0
    while True:
        if is_bd(d):
            k += 1
            if k == n:
                return d
        d += timedelta(days=1)


def last_bd(y: int, m: int) -> date:
    return prev_bd(date(y, m, calendar.monthrange(y, m)[1]))


def add_bd(d: date, n: int) -> date:
    while n > 0:
        d += timedelta(days=1)
        if is_bd(d):
            n -= 1
    return d


def month_end(y: int, m: int) -> date:
    return date(y, m, calendar.monthrange(y, m)[1])


def werktag3(y: int, m: int) -> date:
    """§ 556b BGB: the rent is due by the third working day; Saturday is no working day here (BGH VIII ZR 222/15)."""
    return nth_bd(y, m, 3)


def ym_prev(y: int, m: int) -> tuple[int, int]:
    return (y - 1, 12) if m == 1 else (y, m - 1)


def ym_next(y: int, m: int) -> tuple[int, int]:
    return (y + 1, 1) if m == 12 else (y, m + 1)


def dde(d: date) -> str:
    return d.strftime("%d.%m.%Y")


def dshort(d: date) -> str:
    return d.strftime("%d.%m.%y")


# ------------------------------------------------------------------ money (cents)
def de(c: int) -> str:
    s = f"{abs(c) // 100:,}".replace(",", ".") + f",{abs(c) % 100:02d}"
    return ("-" if c < 0 else "") + s


def de_plain(c: int) -> str:
    return ("-" if c < 0 else "") + f"{abs(c) // 100},{abs(c) % 100:02d}"


def xamt(c: int) -> str:
    return f"{abs(c) / 100:.2f}"


def eur(c: int) -> float:
    return round(c / 100, 2)


def vat19(net: int) -> int:
    return (abs(net) * 19 + 50) // 100 * (1 if net >= 0 else -1)


# ------------------------------------------------------------------ identifiers (invented banks, valid check digits)
def iban_de(blz: str, kto: str) -> str:
    bban = blz + str(kto).zfill(10)
    chk = 98 - int(bban + "131400") % 97
    return f"DE{chk:02d}{bban}"


def iban_lt(bank: str, acct: str) -> str:
    bban = bank + acct
    chk = 98 - int(bban + "212900") % 97
    return f"LT{chk:02d}{bban}"


def fmt_iban(i: str) -> str:
    return " ".join(i[k:k + 4] for k in range(0, len(i), 4))


def creditor_id(nat: str) -> str:
    chk = 98 - int(nat + "131400") % 97
    return f"DE{chk:02d}ZZZ{nat}"


def vat_id(base8: str) -> str:
    p = 10
    for ch in base8:
        s = (int(ch) + p) % 10 or 10
        p = (2 * s) % 11
    c = 11 - p
    return f"DE{base8}{0 if c == 10 else c}"


BANKS = {"SPK": ("29950010", "BRAKDE21XXX", "Sparkasse Brakenfeld"),
         "VB": ("29961800", "GENODEF1VBH", "Volksbank Brakenfeld-Heide eG"),
         "NB": ("25040090", "NORBDE2HXXX", "Norddeutsche Bürgerbank AG"),
         "HB": ("20730040", "HANSDE2HXXX", "Hansebank Direkt AG"),
         "WB": ("25660011", "WESEDE2WXXX", "Weser-Kreditbank eG"),
         "DB": ("12030011", "DIRKDEB1XXX", "Direktkasse Berlin AG"),
         "PB": ("25010055", "POSTDE2HXXX", "Postkonto Bank AG"),
         "LT": ("32500", "REVOLT21XXX", "UAB Baltic Pay Bank")}


def acct(*key, bank: str | None = None) -> tuple[str, str]:
    r = rr("acct", *key)
    b = bank or r.choice(["SPK", "SPK", "SPK", "VB", "VB", "NB", "HB", "HB", "WB", "DB", "PB"])
    blz, bic, _ = BANKS[b]
    return iban_de(blz, str(r.randint(10 ** 8, 10 ** 10 - 1))), bic


# ------------------------------------------------------------------ the firm and its accounts
FIRM = dict(name="Kerkhoff Hausverwaltung GmbH", short="Kerkhoff Hausverwaltung", street="Wallstraße 7", zip="29599",
            city="Brakenfeld", phone="05861 44 71-0", mail="info@kerkhoff-hv.example", buch="buchhaltung@kerkhoff-hv.example",
            web="www.kerkhoff-hv.example", gf="Antje Kerkhoff", hrb="HRB 20714, Amtsgericht Brakenfeld",
            vat=vat_id("29417305"), taxno="52/201/04417")
OWN_INFO = "Kerkhoff Hausverwaltung <info@kerkhoff-hv.example>"
OWN_MO = "Melanie Okafor <m.okafor@kerkhoff-hv.example>"
OWN_JP = "Jens Pohlmann <j.pohlmann@kerkhoff-hv.example>"
OWN_AK = "Antje Kerkhoff <a.kerkhoff@kerkhoff-hv.example>"
STAFF = [("Antje Kerkhoff", "Geschäftsführerin", 520000), ("Jens Pohlmann", "WEG-Verwalter", 398000),
         ("Melanie Okafor", "Mietenbuchhaltung", 331000), ("Sarah Lindemann", "Assistenz", 262000),
         ("Tim Kerkhoff", "Technik (Teilzeit)", 168000)]

ACCTS: dict[str, dict] = {
    "K0": dict(holder="Kerkhoff Hausverwaltung GmbH", what="Geschäftskonto", bank="VB", fmt="vr", kto="4407120"),
    "K1": dict(holder="Kerkhoff Hausverwaltung GmbH Treuhandkonto Mietverwaltung", what="Treuhandkonto Mietverwaltung",
               bank="SPK", fmt="spk", kto="1014477"),
    "K2": dict(holder="Grundstücksgemeinschaft Mühlenteich GbR", what="Objektkonto Mühlenteich", bank="VB", fmt="camt",
               kto="4407236"),
    "K3": dict(holder="WEG Rosenhof Rosenstraße 40-44", what="WEG-Konto Rosenhof", bank="VB", fmt="camt", kto="4410201"),
    "K4": dict(holder="WEG Am Wall 9", what="WEG-Konto Am Wall 9", bank="SPK", fmt="spk", kto="1020202"),
    "K5": dict(holder="WEG Parkresidenz Ostend", what="WEG-Konto Parkresidenz", bank="VB", fmt="camt", kto="4410203"),
    "K6": dict(holder="WEG Brunnenstraße 6", what="WEG-Konto Brunnenstraße 6", bank="SPK", fmt="spk", kto="1020204"),
    "K7": dict(holder="WEG Feldkamp 17-19", what="WEG-Konto Feldkamp", bank="SPK", fmt="spk", kto="1020205",
               since=D("2026-07-01")),
    "K8": dict(holder="WEG Mühlenhof", what="WEG-Konto Mühlenhof", bank="SPK", fmt="spk", kto="1020206"),
}
for _k, _a in ACCTS.items():
    _a["iban"] = iban_de(BANKS[_a["bank"]][0], _a["kto"])
    _a["bic"] = BANKS[_a["bank"]][1]
    _a.setdefault("since", D("2026-01-01"))
    _a["ci"] = creditor_id(f"{int(hashlib.md5(_k.encode()).hexdigest()[:8], 16) % 10 ** 11:011d}")
RESERVE = {k: iban_de(BANKS[ACCTS[k]["bank"]][0], str(int(ACCTS[k]["kto"]) + 500000)) for k in ("K3", "K4", "K5", "K6", "K7", "K8")}
OWN_HOLDERS = {"KERKHOFF HAUSVERWALTUNG GMBH", "KERKHOFF HV GMBH"}

# ------------------------------------------------------------------ people (invented)
SURN = ["Becker", "Hoffmann", "Schröder", "Koch", "Bauer", "Richter", "Klein", "Wolf", "Schwarz", "Zimmermann", "Braun",
        "Hofmann", "Lange", "Schmitt", "Werner", "Meier", "Lehmann", "Schulze", "Köhler", "Herrmann", "König", "Walter",
        "Huber", "Kaiser", "Fuchs", "Peters", "Scholz", "Möller", "Weiß", "Jung", "Hahn", "Schubert", "Vogel",
        "Friedrich", "Keller", "Günther", "Frank", "Berger", "Winkler", "Roth", "Beck", "Lorenz", "Baumann", "Franke",
        "Albrecht", "Schuster", "Simon", "Ludwig", "Winter", "Kraus", "Martin", "Schumacher", "Krämer", "Vogt", "Stein",
        "Jäger", "Otto", "Groß", "Seidel", "Heinrich", "Brandt", "Haas", "Schreiber", "Graf", "Schulte", "Dietrich",
        "Ziegler", "Kuhn", "Kühn", "Pohl", "Horn", "Busch", "Bergmann", "Sauer", "Arnold", "Wolff", "Pfeiffer", "Kaya",
        "Şahin", "Çelik", "Nowak", "Wiśniewska", "Popescu", "Ivanova", "Petrov", "Haddad", "Al-Masri", "Tran", "Rossi",
        "Ferreira", "Mensah", "Jovanović", "Hodžić", "Kovač", "Berisha", "Gashi", "Ebeling", "Wübbena", "Heeren",
        "Janssen", "Ahlers", "Bruns", "Gerdes", "Harms", "Thiele", "Wendt", "Gehrke", "Behrendt", "Oltmanns", "Rieken",
        "Hinrichs", "Eilers", "Cordes", "Mertens", "Lüdemann", "Brockhaus", "Steffens", "Wilkens", "Böttcher", "Marquardt",
        "Pape", "Rohde", "Siebert", "Timmermann", "Wessels", "Ahrens", "Kruse", "Evers", "Hanke", "Fricke"]
FIRST = ["Thomas", "Michael", "Andreas", "Stefan", "Frank", "Jürgen", "Uwe", "Matthias", "Markus", "Sven", "Jan",
         "Tobias", "Florian", "Daniel", "Kevin", "Dennis", "Lukas", "Jonas", "Niklas", "Felix", "Paul", "Leon",
         "Ahmet", "Emre", "Piotr", "Andrei", "Dragan", "Omar", "Karim", "Minh", "Hauke", "Rainer", "Dieter", "Klaus",
         "Wolfgang", "Bernd", "Sabine", "Petra", "Claudia", "Susanne", "Nicole", "Stefanie", "Julia", "Katharina",
         "Anna", "Laura", "Sarah", "Jana", "Lisa", "Marie", "Sophie", "Emma", "Hannah", "Monika", "Renate", "Ursula",
         "Heike", "Birgit", "Elif", "Fatma", "Agnieszka", "Ioana", "Olga", "Layla", "Linh", "Imke", "Wiebke", "Anke",
         "Silke", "Merle", "Frauke", "Gesa", "Tanja", "Kerstin", "Doris", "Gisela", "Helga", "Jutta", "Ole", "Malte"]
MAILDOM = ["webpost.example", "fernmail.example", "heidenetz.example", "postfach.example"]
USED_NAMES: set = set()


def person(key, nach=None, vor=None, bank=None, mail=True) -> dict:
    r = rr("person", key)
    if nach is None:
        while True:
            n, v = r.choice(SURN), r.choice(FIRST)
            if (n, v) not in USED_NAMES and n not in ("Neumann", "Schulz"):
                break
    else:
        n, v = nach, vor
    USED_NAMES.add((n, v))
    iban, bic = acct("p", key, bank=bank)
    tr = str.maketrans({"ä": "ae", "ö": "oe", "ü": "ue", "ß": "ss", "ş": "s", "ç": "c", "ć": "c", "ž": "z", "ı": "i",
                        "ś": "s", "Ş": "S", "Ç": "C", "-": "", " ": ""})
    em = f"{(v or 'x')[0].lower()}.{n.lower().translate(tr)}@{r.choice(MAILDOM)}" if mail else ""
    ph = f"0{r.choice(['151', '160', '170', '176', '5861', '5862'])} {r.randint(1000000, 9999999)}"
    return dict(nach=n, vor=v, iban=iban, bic=bic, mail=em, phone=ph,
                bankname=f"{(v or '')} {n}".strip(), list=f"{n}, {v}" if v else n)


def sepa_name(s: str) -> str:
    tr = {"ä": "ae", "ö": "oe", "ü": "ue", "Ä": "Ae", "Ö": "Oe", "Ü": "Ue", "ß": "ss", "ı": "i", "İ": "I", "ş": "s",
          "Ş": "S", "ç": "c", "ć": "c", "ś": "s", "ž": "z", "–": "-", "&": "+"}
    return "".join(tr.get(ch, ch) for ch in s)[:70]


# ------------------------------------------------------------------ properties
LAGE2 = ["EG links", "EG rechts", "1. OG links", "1. OG rechts", "2. OG links", "2. OG rechts", "3. OG links",
         "3. OG rechts", "DG links", "DG rechts"]
LAGE3 = ["EG links", "EG Mitte", "EG rechts", "1. OG links", "1. OG Mitte", "1. OG rechts", "2. OG links", "2. OG Mitte",
         "2. OG rechts", "3. OG links", "3. OG Mitte", "3. OG rechts", "DG links", "DG Mitte", "DG rechts"]

OWNERS = {  # the landlords of the rental houses (clients of the Mietverwaltung)
    "ov": dict(name="Erbengemeinschaft Voßkamp", addr=["Erbengemeinschaft Voßkamp", "z. Hd. Herrn Dirk Voßkamp",
                                                       "Am Brink 4", "29599 Brakenfeld"], kd="1001",
               payout=acct("owner", "ov", bank="SPK")),
    "mt": dict(name="Grundstücksgemeinschaft Mühlenteich GbR", addr=["Grundstücksgemeinschaft Mühlenteich GbR",
                                                                     "Ralf Ebeling und Martina Kühn", "Teichweg 18",
                                                                     "29599 Brakenfeld"], kd="1002",
               payout=acct("owner", "mt", bank="VB")),
    "al": dict(name="Dr. Henrike Albers", addr=["Frau Dr. Henrike Albers", "Hasenwinkel 9", "21335 Lüneburg"], kd="1003",
               payout=acct("owner", "al", bank="HB")),
    "br": dict(name="Eheleute Brinkmann", addr=["Karl-Heinz und Renate Brinkmann", "Kastanienweg 10", "29599 Brakenfeld"],
               kd="1004", payout=acct("owner", "br", bank="SPK")),
    "wb": dict(name="Wohnbau Brakenfeld GmbH & Co. KG", addr=["Wohnbau Brakenfeld GmbH & Co. KG", "Geschäftsführung",
                                                              "Industriestraße 40", "29599 Brakenfeld"], kd="1005",
               payout=acct("owner", "wb", bank="NB")),
    "mh": dict(name="Jürgen und Petra Meinhardt", addr=["Jürgen und Petra Meinhardt", "Schillerstraße 7",
                                                        "29599 Brakenfeld"], kd="1006",
               payout=acct("owner", "mh", bank="VB"), new_payout=acct("owner", "mh2", bank="HB")),
    "hw": dict(name="Heideweg Wohnen GmbH", addr=["Heideweg Wohnen GmbH", "Moorstraße 3", "29525 Uhlenbrück"], kd="1007",
               payout=acct("owner", "hw", bank="WB")),
}
PROPS = [
    dict(no=101, kind="MV", name="Lindenallee 14", street="Lindenallee 14", owner="ov", units=12, garages=2, built=1962,
         acct="K1", since=D("2014-01-01"), heat=True),
    dict(no=102, kind="MV", name="Lindenallee 16", street="Lindenallee 16", owner="ov", units=10, built=1962, acct="K1",
         since=D("2014-01-01"), heat=True),
    dict(no=103, kind="MV", name="Am Mühlenteich 3-5", street="Am Mühlenteich 3-5", owner="mt", units=24, tg=10,
         built=1994, acct="K2", since=D("2017-07-01"), heat=True),
    dict(no=104, kind="MV", name="Bahnhofstraße 22", street="Bahnhofstraße 22", owner="al", units=6, shops=2, built=1908,
         acct="K1", since=D("2019-01-01"), heat=False),
    dict(no=105, kind="MV", name="Kastanienweg 8", street="Kastanienweg 8", owner="br", units=8, built=1971, acct="K1",
         since=D("2012-04-01"), heat=False),
    dict(no=106, kind="MV", name="Gartenstraße 31", street="Gartenstraße 31", owner="wb", units=18, built=1958, acct="K1",
         since=D("2020-01-01"), heat=True),
    dict(no=107, kind="MV", name="Schillerstraße 5", street="Schillerstraße 5", owner="mh", units=6, built=1929, acct="K1",
         since=D("2016-01-01"), heat=False),
    dict(no=108, kind="MV", name="Heideweg 2a-2c", street="Heideweg 2a-2c", owner="hw", units=20, built=1999, acct="K1",
         since=D("2021-01-01"), until=D("2026-09-30"), heat=True),
    dict(no=201, kind="WEG", name="WEG Rosenhof", street="Rosenstraße 40-44", units=36, built=1985, acct="K3",
         since=D("2015-01-01"), etv=D("2026-05-20"), ls=True),
    dict(no=202, kind="WEG", name="WEG Am Wall 9", street="Am Wall 9", units=14, built=1978, acct="K4",
         since=D("2018-01-01"), etv=D("2026-06-03"), bestellt_bis=D("2026-12-31")),
    dict(no=203, kind="WEG", name="WEG Parkresidenz Ostend", street="Ostendstraße 101", units=28, built=2004, acct="K5",
         since=D("2016-01-01"), etv=D("2026-06-10"), ls=True),
    dict(no=204, kind="WEG", name="WEG Brunnenstraße 6", street="Brunnenstraße 6", units=8, built=1912, acct="K6",
         since=D("2019-01-01"), etv=D("2026-07-23")),
    dict(no=205, kind="WEG", name="WEG Feldkamp 17-19", street="Feldkamp 17-19", units=22, built=1992, acct="K7",
         since=D("2026-07-01"), etv=None),
    dict(no=206, kind="WEG", name="WEG Mühlenhof", street="Mühlenstraße 2", units=12, built=1998, acct="K8",
         since=D("2017-01-01"), etv=D("2026-06-24")),
]
PD = {p["no"]: p for p in PROPS}
ENDED = [dict(no=209, kind="WEG", name="WEG Kirchplatz 3", street="Kirchplatz 3", units=9, built=1964,
              since=D("2016-01-01"), until=D("2025-12-31"), why="Verwalterwechsel (Beschluss ETV 2025)")]
SE_UNITS = [  # rented flats in the WEGs, managed for their owners (Sondereigentumsverwaltung)
    dict(se="901", konto="90101", weg=201, unit=12, owner="Ingrid Haverkamp", ok="20112"),
    dict(se="902", konto="90102", weg=201, unit=30, owner="Dr. Frank Ostendorf", ok="20130"),
    dict(se="903", konto="90103", weg=203, unit=5, owner="Thi Lan Nguyen", ok="20305"),
    dict(se="904", konto="90104", weg=203, unit=22, owner="Kemal Aydın", ok="20322"),
    dict(se="905", konto="90105", weg=206, unit=8, owner="Brigitte Sommer", ok="20608"),
]


def lage(p: dict, i: int) -> str:
    seq = LAGE3 if p["units"] >= 18 else LAGE2
    if p["no"] == 103:  # two entrances
        return ("Nr. 3 " if i <= 12 else "Nr. 5 ") + seq[(i - 1) % 12]
    if p["no"] == 108:
        return ["2a", "2b", "2c"][(i - 1) // 7] + " " + LAGE2[(i - 1) % 7]
    if p["no"] == 201:
        return f"Haus {['40', '42', '44'][(i - 1) // 12]} " + LAGE2[(i - 1) % 12 % 10]
    if p["no"] == 205:
        return f"Haus {'17' if i <= 11 else '19'} " + LAGE3[(i - 1) % 11]
    return seq[(i - 1) % len(seq)]


# ------------------------------------------------------------------ leases (Mietverhältnisse)
LEASES: list[dict] = []
LD: dict[str, dict] = {}


def L_(konto, obj, unit, who, start, kalt, bk, hk, stp=0, end=None, how="DA_1", jc=None, area=None, vat=False,
       changes=(), deposit=None, style=None, shop=None, unitlabel=None, in_export=True):
    r = rr("lease", konto)
    p = PD.get(obj)
    L = dict(konto=konto, obj=obj, unit=unit, who=who, start=start, end=end, how=how, jc=jc, vat=vat, shop=shop,
             area=area or round(r.uniform(42, 96), 1), rooms=None, sched=[(start, kalt, bk, hk, stp)],
             deposit=deposit if deposit is not None else 3 * kalt, style=style or r.choice(["mon", "mon", "name", "static",
                                                                                            "konto", "short"]),
             lage=unitlabel or (lage(p, unit) if p else ""), in_export=in_export, kd_changes=[])
    L["rooms"] = 1 if L["area"] < 45 else 2 if L["area"] < 62 else 3 if L["area"] < 85 else 4
    for ch in changes:  # (from, kind, kalt, bk, hk, stp, noted): noted False = the office forgot it
        frm, kind, k2, b2, h2, s2, noted = ch
        last = L["sched"][-1]
        L["sched"].append((frm, k2 if k2 is not None else last[1], b2 if b2 is not None else last[2],
                           h2 if h2 is not None else last[3], s2 if s2 is not None else last[4]))
        L["kd_changes"].append(dict(frm=frm, kind=kind, noted=noted, old=last, new=L["sched"][-1]))
    LEASES.append(L)
    LD[konto] = L
    return L


def rent_at(L: dict, d: date, noted_only=False) -> tuple[int, int, int, int]:
    cur = L["sched"][0][1:]
    for k, s in enumerate(L["sched"]):
        if s[0] <= d:
            if k and noted_only and not L["kd_changes"][k - 1]["noted"]:
                continue
            cur = s[1:]
    return cur


def total_of(L: dict, parts) -> int:
    k, b, h, s = parts
    t = k + b + h + s
    return t + vat19(t) if L["vat"] else t


def build_leases():
    # the special ones first (numbers fixed by the answer key)
    P = person
    L_("10104", 101, 4, P("10104", "Hartwig", "Sabine"), D("2019-07-01"), 61000, 16000, 9000, how="DA_1", style="static",
       changes=[(D("2026-07-01"), "Staffel", 64500, None, None, None, True)])
    L_("10105", 101, 5, P("10105", "Voigt", "Rainer"), D("2017-03-01"), 61000, 15000, 7000, end=D("2026-07-31"),
       deposit=183000)
    L_("10115", 101, 5, P("10115", "Demir", "Selin"), D("2026-08-15"), 65000, 17000, 8500, style="name")
    L_("10106", 101, 6, P("10106", "Schulz", "Andrea"), D("2021-02-01"), 49500, 13500, 6500, how="LS")
    L_("10107", 101, 7, P("10107", "Hartmann", "Lea"), D("2023-09-01"), 48000, 14000, 7000, style="static")
    L_("10108", 101, 8, P("10108", "Ahlers", "Gesa"), D("2018-05-01"), 52000, 14500, 7500,
       changes=[(D("2026-09-01"), "BK-Vorauszahlung", None, 16500, None, None, True)], style="static")
    L_("10205", 102, 5, P("10205", "Gerdes", "Malte"), D("2015-10-01"), 53000, 15500, 8000, how="UE", style="mon")
    L_("10206", 102, 6, P("10206", "Oltmanns", "Frauke"), D("2022-09-01"), 56000, 15000, 8000, style="static",
       changes=[(D("2025-09-01"), "Staffel", 56000, None, None, None, True),
                (D("2026-09-01"), "Staffel", 60000, None, None, None, False)])
    L_("10209", 102, 9, P("10209", "Neumann", "Thomas"), D("2011-04-01"), 44000, 13000, 7000)
    L_("10210", 102, 10, P("10210", "Rieken", "Imke"), D("2020-06-01"), 47000, 12500, 7000,
       changes=[(D("2026-09-01"), "BK-Vorauszahlung", None, 14000, None, None, True)], style="static")
    L_("10303", 103, 3, P("10303", "Hinrichs", "Kerstin"), D("2016-08-01"), 54000, 15000, 9000, stp=5500,
       changes=[(D("2026-09-01"), "Mieterhöhung § 558", 58000, None, None, None, True)], how="UE", style="mon")
    L_("10304", 103, 4, P("10304", "Eilers", "Bernd"), D("2018-01-01"), 51000, 14000, 8500)
    L_("10305", 103, 5, P("10305", "Böhm", "Kevin"), D("2021-11-01"), 56000, 16200, 9000, how="UE", style="mon")
    L_("10313", 103, 13, P("10313", "Cordes", "Merle"), D("2019-03-01"), 50500, 14000, 8000, stp=5500)
    L_("10319", 103, 19, P("10319", "Schulz", "Andreas"), D("2024-02-01"), 57000, 15500, 9000, stp=5500, how="UE")
    L_("10325", 103, 30, P("10325", "Kowalczyk", "Marek"), D("2020-01-01"), 0, 0, 0, stp=5500, how="Q", deposit=0,
       unitlabel="TG-Stellplatz S07")
    L_("10407", 104, 7, dict(nach="Café Kornblume", vor=None, list="Café Kornblume Inh. Svenja Rademacher",
                             bankname="Cafe Kornblume Svenja Rademacher", iban=acct("p", "10407", bank="VB")[0],
                             bic=BANKS["VB"][1], mail="info@cafe-kornblume.example", phone="05861 30 12 44"),
       D("2019-06-01"), 145000, 28000, 0, vat=True, how="UE3", area=118.0, unitlabel="EG Laden links",
       shop="Café/Bäckerei", deposit=435000, style="gew")
    L_("10408", 104, 8, dict(nach="Physiotherapie Lindner & Saß GbR", vor=None, list="Physiotherapie Lindner & Saß GbR",
                             bankname="Physiotherapie Lindner u. Sass GbR", iban=acct("p", "10408", bank="SPK")[0],
                             bic=BANKS["SPK"][1], mail="praxis@physio-lindner-sass.example", phone="05861 30 55 10"),
       D("2021-01-01"), 98000, 21000, 0, area=96.5, unitlabel="EG Laden rechts", shop="Praxis",
       changes=[(D("2026-10-01"), "Indexmiete", 100058, None, None, None, True)], deposit=294000, style="gew")
    L_("10502", 105, 2, P("10502", "Wübbena", "Hauke"), D("2014-05-01"), 50000, 18400, 10000, how="UE", style="mon")
    L_("10507", 105, 7, P("10507", "Mertens", "Doris"), D("2012-04-01"), 41000, 15000, 9000, how="UE", style="konto")
    L_("10601", 106, 1, dict(P("10601", "Yılmaz", "Ayşe"), list="Yılmaz, Ayşe und Mehmet", bankname="YILMAZ MEHMET"),
       D("2016-02-01"), 39000, 12000, 7000)
    L_("10603", 106, 3, P("10603", "Petrović", "Dragana"), D("2022-04-01"), 39800, 12500, 7500, how="JC",
       jc=dict(bg="31207BG0004711", share=59800))
    L_("10606", 106, 6, P("10606", "Kaya", "Elif"), D("2020-08-01"), 42000, 12500, 7500)
    L_("10609", 106, 9, P("10609", "Sadiku", "Arben"), D("2023-01-01"), 40550, 12000, 7300, how="JC",
       jc=dict(bg="31207BG0005288", share=56000))
    L_("10611", 106, 11, P("10611", "Krüger", "Dennis"), D("2021-05-01"), 35400, 11000, 7000, how="JC",
       jc=dict(bg="31207BG0003920", share=53400), end=D("2026-08-31"))
    L_("10612", 106, 12, P("10612", "Brandl", "Marco"), D("2022-10-01"), 52000, 14500, 8000, how="UE", style="mon")
    L_("10614", 106, 14, P("10614", "Wagner", "Jacqueline"), D("2024-03-01"), 41200, 12500, 7500, how="JC",
       jc=dict(bg="31207BG0006105", share=61200))
    L_("10615", 106, 15, P("10615", "Seidel", "Klaus"), D("2009-01-01"), 38000, 12000, 7000, end=D("2026-05-31"))
    L_("10619", 106, 11, P("10619", "Popescu", "Ioana"), D("2026-11-01"), 44500, 12000, 7000, deposit=133500)
    L_("10703", 107, 3, P("10703", "Wessels", "Gisela"), D("2010-09-01"), 39000, 14000, 8500)
    L_("10802", 108, 2, P("10802", "Hanke", "Silke"), D("2019-09-01"), 45000, 15000, 8900, how="LS")
    L_("10804", 108, 4, P("10804", "Neumann", "Thomas"), D("2023-05-01"), 46000, 14000, 8500)
    L_("10817", 108, 17, P("10817", "Evers", "Jana"), D("2020-07-01"), 43500, 14000, 8000, style="static")
    # everybody else
    for p in PROPS:
        if p["kind"] != "MV":
            continue
        n = p["units"] + p.get("shops", 0)
        for i in range(1, n + 1):
            konto = f"{p['no']}{i:02d}"
            if any(L["obj"] == p["no"] and L["unit"] == i for L in LEASES):
                continue
            r = rr("std", konto)
            area = round(r.uniform(44, 94), 1)
            rate = r.uniform(6.4, 8.9) if p["no"] != 104 else r.uniform(7.5, 9.6)
            kalt = int(round(area * rate * 100))
            bk = int(round(area * r.uniform(1.9, 2.5))) * 100
            hk = int(round(area * r.uniform(0.95, 1.3))) * 100 if p["heat"] or r.random() < 0.6 else 0
            stp = 0
            if p["no"] == 103 and i in (1, 2, 6, 8, 10, 14, 20):
                stp = 5500
            if p["no"] == 101 and i in (2, 9):
                stp = 4500  # the two garages
            how = r.choices(["DA_1", "DA_28", "LS", "UE", "DA_3"], weights=[42, 14, 20, 16, 8])[0]
            if p["acct"] != "K1" and how == "LS":
                how = "DA_1"
            start = D(f"{r.randint(2004, 2025)}-{r.randint(1, 12):02d}-01")
            L_(konto, p["no"], i, person(konto), start, kalt, bk, hk, stp=stp, how=how, area=area)
    for s in SE_UNITS:  # tenants of the rented flats in the WEGs
        r = rr("se", s["konto"])
        area = round(r.uniform(55, 88), 1)
        L_(s["konto"], int(s["se"]), 1, person(s["konto"]), D(f"{r.randint(2015, 2024)}-{r.randint(1, 12):02d}-01"),
           int(round(area * r.uniform(7.2, 8.6) * 100)), int(round(area * 2.3)) * 100, int(round(area * 1.1)) * 100,
           how=r.choice(["DA_1", "DA_1", "UE"]), area=area,
           unitlabel=f"{PD[s['weg']]['street']}, Whg {s['unit']}")
    for L in LEASES:
        L["deposit"] = L["deposit"] if L["konto"] not in ("10105",) else 183000


# ------------------------------------------------------------------ WEG owners (Eigentümer)
OWN: list[dict] = []
OD: dict[str, dict] = {}
MEA204 = [125, 130, 118, 130, 125, 122, 125, 125]


def O_(konto, weg, unit, who, hg, mea, how="DA_1", start=None, end=None, units=None):
    o = dict(konto=konto, weg=weg, unit=unit, who=who, hg=hg, mea=mea, how=how, start=start or D("2000-01-01"), end=end,
             lage=lage(PD[weg], unit), mandate=f"{PD[weg]['no']}-HG-{konto}", area=None)
    OWN.append(o)
    OD[konto] = o
    return o


def build_owners():
    se = {s["ok"]: s for s in SE_UNITS}
    for p in PROPS:
        if p["kind"] != "WEG":
            continue
        meas = MEA204 if p["no"] == 204 else None
        for i in range(1, p["units"] + 1):
            konto = f"{p['no']}{i:02d}"
            r = rr("own", konto)
            area = round(r.uniform(48, 112), 1)
            hg = int(round(area * r.uniform(2.7, 3.5))) * 100
            mea = meas[i - 1] if meas else round(1000 * area / (p["units"] * 78), 1)
            how = "LS" if p.get("ls") and r.random() < 0.85 else r.choices(["DA_1", "DA_28", "UE"], weights=[60, 20, 20])[0]
            if konto in se:
                who = dict(person(konto, *se[konto]["owner"].split()[-1:], " ".join(se[konto]["owner"].split()[:-1])),
                           list=f"{se[konto]['owner'].split()[-1]}, {' '.join(se[konto]['owner'].split()[:-1])}")
                how = "SE"
            elif konto == "20404":
                who, hg = person(konto, "Hollmann", "Gerd"), 31200
            elif konto == "20402":
                who = dict(person(konto, "Kranz", "Uwe"), list="Kranz, Uwe und Heike (Eheleute)", bankname="UWE KRANZ")
            elif konto == "20127":
                who, hg, how = person(konto, "Riedel", "Dr. Hannes"), 28700, "LS"
            elif konto == "20311":
                who, hg, how = person(konto, "Kowalski", "Stefan"), 26800, "LS"
            elif konto == "20211":
                who = person(konto, "Bergmann", "Ursula")
            elif konto in ("20603", "20604"):
                who, how = person("20603", "Petersen", "Uwe"), "PAIR"
            elif konto == "20514":
                who, hg = person(konto, "Ivanova", "Olga"), 27600
            elif konto == "20125":
                who = person(konto, "Brockmann", "Ute")
            else:
                who = person(konto)
            o = O_(konto, p["no"], i, who, hg, mea, how=how, end=D("2026-08-31") if konto == "20311" else None)
            o["area"] = area
    O_("20361", 203, 11, person("20361", "Hoffmeister", "Bianca"), 26800, OD["20311"]["mea"], how="LS",
       start=D("2026-09-01"))
    OD["20361"]["area"] = OD["20311"]["area"]
    OD["20604"]["who"] = OD["20603"]["who"]


# ------------------------------------------------------------------ claims (Sollstellungen) and money
CLAIMS: list[dict] = []
CK: dict[str, dict] = {}
MONEY: list[dict] = []


def claim(key, kind, konto, amount, due, issued, label, acct, period=None, name="", obj=None, unit=None, note=""):
    assert key not in CK, key
    c = dict(key=key, kind=kind, konto=konto, amount=amount, due=due, issued=issued, label=label, acct=acct,
             period=period, name=name, obj=obj, unit=unit, note=note, gen_paid=0)
    CLAIMS.append(c)
    CK[key] = c
    return c


def money(acct_, d, amount, name="", iban="", bic="", vwz="", btext="", alloc=(), kind="", **kw) -> dict:
    assert is_bd(d) or kind in ("prev",), (d, vwz)
    m = dict(id=len(MONEY) + 1, acct=acct_, date=d, amount=amount, name=name, iban=iban, bic=bic, vwz=vwz, btext=btext,
             alloc=[], kind=kind, e2e=kw.get("e2e", ""), mandate=kw.get("mandate", ""), ci=kw.get("ci", ""),
             sammler=kw.get("sammler", ""), urspr=kw.get("urspr"), auslagen=kw.get("auslagen"), batch=kw.get("batch"),
             refund_of=kw.get("refund_of"), note=kw.get("note", ""), over=0, settled=[], gvc=kw.get("gvc", ""),
             details=kw.get("details", True), rtr=kw.get("rtr"))
    for c, cents, *eff in alloc:
        m["alloc"].append((c, cents, eff[0] if eff else d))
        c["gen_paid"] += cents
    if amount > 0:
        m["over"] = amount - sum(x[1] for x in m["alloc"])
    MONEY.append(m)
    return m


def oldest(konto: str, amount: int, d: date, kinds=("miete",), upto=None) -> list:
    """§ 366 Abs. 2 BGB as the office books it: a payment without a month pays the oldest due rent first."""
    out, left = [], amount
    for c in sorted((c for c in CLAIMS if c["konto"] == konto and c["kind"] in kinds and c["issued"] <= (upto or d)),
                    key=lambda c: (c["due"], c["key"])):
        need = c["amount"] - c["gen_paid"]
        if need <= 0 or left <= 0:
            continue
        take = min(need, left)
        out.append((c, take))
        left -= take
    return out


def mkey(konto, kind, y, m):
    return f"{konto} {kind} {m:02d}/{y}"


MONTHS = [(2026, m) for m in range(1, 13)]
CUTS = [D("2026-09-17"), D("2026-10-29")]
LAST = CUTS[-1]


def build_rent_claims():
    for L in LEASES:
        if L["obj"] == 108:
            stop = D("2026-09-30")
        else:
            stop = date(2026, 12, 31)
        for (y, m) in MONTHS:
            a, b = date(y, m, 1), month_end(y, m)
            if L["start"] > b or (L["end"] and L["end"] < a) or a > stop:
                continue
            parts = rent_at(L, a)
            amt = total_of(L, parts)
            if L["start"] > a:  # moved in during the month: pro rata
                amt = round(amt * ((b - L["start"]).days + 1) / b.day)
            claim(mkey(L["konto"], "Miete", y, m), "miete", L["konto"], amt, werktag3(y, m), a,
                  f"Miete {m:02d}/{y}", acct_of_lease(L), period=(y, m), name=L["who"]["list"], obj=L["obj"],
                  unit=L["lage"])
        if L["start"] >= D("2026-01-01") and L["deposit"]:  # deposit in three instalments (§ 551 Abs. 2 BGB)
            rate = L["deposit"] // 3
            for k in range(3):
                y, m = L["start"].year, L["start"].month
                for _ in range(k):
                    y, m = ym_next(y, m)
                due = L["start"] if k == 0 else werktag3(y, m)
                claim(f"{L['konto']} Kaution Rate {k + 1}", "kaution", L["konto"], rate if k < 2 else L["deposit"] - 2 * rate,
                      due, L["start"] if k == 0 else date(y, m, 1), f"Kaution Rate {k + 1}", acct_of_lease(L),
                      name=L["who"]["list"], obj=L["obj"], unit=L["lage"])


def acct_of_lease(L) -> str:
    return PD[L["obj"]]["acct"] if L["obj"] in PD else "K1"


def build_owner_claims():
    for o in OWN:
        p = PD[o["weg"]]
        first = max(o["start"], D("2026-07-01") if p["no"] == 205 else D("2026-01-01"))
        for (y, m) in MONTHS:
            a, b = date(y, m, 1), month_end(y, m)
            if first > b or (o["end"] and o["end"] < a):
                continue
            claim(mkey(o["konto"], "Hausgeld", y, m), "hausgeld", o["konto"], o["hg"], werktag3(y, m), a,
                  f"Hausgeld {m:02d}/{y}", p["acct"], period=(y, m), name=o["who"]["list"], obj=p["no"], unit=o["lage"])
    # the previous administrator's arrears at Feldkamp (handover list)
    for m in (5, 6):
        claim(mkey("20514", "Hausgeld", 2026, m), "hausgeld", "20514", 27600, werktag3(2026, m), date(2026, m, 1),
              f"Hausgeld {m:02d}/2026 (Vorverwalter)", "K7", period=(2026, m), name=OD["20514"]["who"]["list"], obj=205,
              unit=OD["20514"]["lage"], note="laut Übergabeliste des Vorverwalters")


# ------------------------------------------------------------------ annual results (BK 2025, Jahresabrechnung 2025, Umlage)
BK_SENT = {101: D("2026-06-15"), 102: D("2026-06-15"), 103: D("2026-07-10"), 104: D("2026-06-22"), 105: D("2026-07-01"),
           107: D("2026-05-28"), 108: D("2026-08-14"), 901: D("2026-07-20"), 902: D("2026-07-20"), 903: D("2026-07-20"),
           904: D("2026-07-20"), 905: D("2026-07-20")}
BKRES: dict[str, dict] = {}
JARES: dict[str, dict] = {}
SU204: dict[str, int] = {}


def build_results():
    fixed = {"10304": 31840, "10205": 41290, "10703": 26410, "10502": -8640, "10105": 14260, "10612": 18830,
             "10305": 9120, "10104": -4210}
    for L in LEASES:
        sent = BK_SENT.get(L["obj"])
        if not sent or L["start"] > D("2025-06-30") or (L["end"] and L["end"] < D("2025-01-01")) or L["konto"] == "10325":
            continue
        r = rr("bk", L["konto"])
        res = fixed.get(L["konto"], r.choice([-1, 1, 1]) * r.randint(1200, 39000))
        due = sent + timedelta(days=33)
        BKRES[L["konto"]] = dict(res=res, sent=sent, due=due)
        if res > 0:
            claim(f"{L['konto']} BK 2025", "bk", L["konto"], res, due, sent, "Betriebskostenabrechnung 2025",
                  acct_of_lease(L), name=L["who"]["list"], obj=L["obj"], unit=L["lage"])
    fixed_ja = {"20404": 18840, "20211": 34275, "20127": 9410}
    for o in OWN:
        p = PD[o["weg"]]
        if not p.get("etv") or o["start"] > D("2025-12-31") or o["konto"] in ("20604",):
            continue
        r = rr("ja", o["konto"])
        res = fixed_ja.get(o["konto"], r.choice([-1, 1, 1]) * r.randint(1500, 32000))
        if o["konto"] == "20603":
            res = 2 * res  # Petersen: both flats in one statement, one result
        due = p["etv"] + timedelta(days=28)
        JARES[o["konto"]] = dict(res=res, etv=p["etv"], due=due)
        if res > 0:
            claim(f"{o['konto']} Abrechnung 2025", "ja", o["konto"], res, due, p["etv"], "Abrechnungsspitze 2025",
                  p["acct"], name=o["who"]["list"], obj=p["no"], unit=o["lage"])
    for i, o in enumerate(x for x in OWN if x["weg"] == 204):
        share = 2400000 * o["mea"] // 1000
        SU204[o["konto"]] = share
        claim(f"{o['konto']} Sonderumlage 2026", "su", o["konto"], share, D("2026-09-01"), D("2026-07-23"),
              "Sonderumlage Dachsanierung (Beschluss 23.07.2026)", "K6", name=o["who"]["list"], obj=204, unit=o["lage"])


# ------------------------------------------------------------------ own invoices (Verwaltervergütung, Sonderleistungen)
INV: list[dict] = []
IK: dict[str, dict] = {}
MANDATES = [  # (contract, objects, client, net per month, account the fee is taken from); the Treuhandkonto's first
    dict(vv="VV-101", objs=[101, 102], client="ov", net=65400, acct="K1", units="22 WE, 2 Garagen"),
    dict(vv="VV-105", objs=[105], client="br", net=24000, acct="K1", units="8 WE"),
    dict(vv="VV-106", objs=[106], client="wb", net=50400, acct="K1", units="18 WE"),
    dict(vv="VV-107", objs=[107], client="mh", net=18600, acct="K1", units="6 WE"),
    dict(vv="VV-108", objs=[108], client="hw", net=56000, acct="K1", units="20 WE", until=D("2026-09-30")),
] + [dict(vv=f"VV-{s['se']}", objs=[int(s["se"])], client=f"se{s['se']}", net=3600, acct="K1", units="1 WE (SE)")
     for s in SE_UNITS] + [
    dict(vv="VV-103", objs=[103], client="mt", net=72800, acct="K2", units="24 WE, 10 TG-Stellplätze"),
    dict(vv="VV-104", objs=[104], client="al", net=27000, acct=None, units="6 WE, 2 Gewerbe"),
    dict(vv="VV-201", objs=[201], client="w201", net=86400, acct="K3", units="36 Einheiten"),
    dict(vv="VV-202", objs=[202], client="w202", net=37800, acct="K4", units="14 Einheiten"),
    dict(vv="VV-203", objs=[203], client="w203", net=70000, acct="K5", units="28 Einheiten"),
    dict(vv="VV-204", objs=[204], client="w204", net=25600, acct="K6", units="8 Einheiten"),
    dict(vv="VV-205", objs=[205], client="w205", net=57200, acct="K7", units="22 Einheiten", since=D("2026-07-01")),
    dict(vv="VV-206", objs=[206], client="w206", net=33600, acct="K8", units="12 Einheiten"),
]
CLIENTS: dict[str, dict] = {k: dict(o, key=k) for k, o in OWNERS.items()}
for _s in SE_UNITS:
    _nm = _s["owner"]
    CLIENTS[f"se{_s['se']}"] = dict(name=_nm, addr=[_nm, f"{PD[_s['weg']]['street']}", "29599 Brakenfeld"],
                                    kd=f"10{10 + int(_s['se']) - 900}", payout=acct("owner", _s["se"]), key=f"se{_s['se']}")
for _i, _p in enumerate(p for p in PROPS if p["kind"] == "WEG"):
    CLIENTS[f"w{_p['no']}"] = dict(name=_p["name"], addr=[_p["name"], "vertreten durch Kerkhoff Hausverwaltung GmbH",
                                                          _p["street"], "29599 Brakenfeld"], kd=f"10{21 + _i}",
                                   key=f"w{_p['no']}")
CLIENTS["brockmann"] = dict(name="Ute Brockmann", addr=["Frau Ute Brockmann", "Rosenstraße 44", "29599 Brakenfeld"],
                            kd="1040", key="brockmann")
SL: list[dict] = []


def inv_(d, client, subject, net, period=None, ref="", kind="VV", sl=None, order=0):
    i = dict(no=None, date=d, client=client, subject=subject, net=net, vat=vat19(net), gross=net + vat19(net),
             period=period, ref=ref, kind=kind, sl=sl, order=order, due=d + timedelta(days=10), claim=None)
    INV.append(i)
    return i


def build_invoices():
    for (y, m) in MONTHS[:10]:  # January .. October
        d = nth_bd(y, m, 1)
        for k, md in enumerate(MANDATES):
            if md.get("until") and date(y, m, 1) > md["until"]:
                continue
            if md["vv"] == "VV-205":
                continue  # set up only on 05.10. (H34)
            objs = " / ".join(PD[o]["name"] if o in PD else f"SE-Verwaltung {o}" for o in md["objs"]) if md["objs"][0] in PD \
                else next(f"SE-Verwaltung {s['owner']}, {PD[s['weg']]['street']} Whg {s['unit']}" for s in SE_UNITS
                          if int(s["se"]) == md["objs"][0])
            inv_(d, md["client"], f"Verwaltervergütung {m:02d}/{y} – {objs}", md["net"],
                 period=(date(y, m, 1), month_end(y, m)), ref=md["vv"], order=k)
    for m in (7, 8, 9, 10):
        md = next(x for x in MANDATES if x["vv"] == "VV-205")
        inv_(D("2026-10-05"), "w205", f"Verwaltervergütung {m:02d}/2026 – WEG Feldkamp 17-19", md["net"],
             period=(date(2026, m, 1), month_end(2026, m)), ref="VV-205", order=m)
    # special services (Sonderleistungen)
    sls = [("SL-26-001", D("2026-01-20"), 202, "Bescheinigungen § 35a EStG 2025 (14 Eigentümer)", "w202", 35000, D("2026-01-26")),
           ("SL-26-002", D("2026-02-11"), 201, "Verwalterzustimmung § 12 WEG – Veräußerung Einheit 201-07", "w201", 20000, D("2026-02-13")),
           ("SL-26-003", D("2026-03-05"), 106, "Begleitung Wohnungsübergaben Gartenstraße 31 (3 Termine)", "wb", 22500, D("2026-03-09")),
           ("SL-26-004", D("2026-03-24"), 203, "Bescheinigungen § 35a EStG 2025 (28 Eigentümer)", "w203", 70000, D("2026-03-30")),
           ("SL-26-005", D("2026-04-16"), 103, "Mieterhöhungsverlangen § 558 BGB (6 Mietverhältnisse)", "mt", 36000, D("2026-04-20")),
           ("SL-26-006", D("2026-05-07"), 204, "Einholung von drei Angeboten Dachsanierung, Ortstermine", "w204", 28000, D("2026-05-12")),
           ("SL-26-007", D("2026-06-18"), 206, "Verwalterzustimmung § 12 WEG – Veräußerung Einheit 206-03", "w206", 20000, D("2026-06-22")),
           ("SL-26-008", D("2026-07-08"), 101, "Durchführung Instandsetzung Fassade Lindenallee 14 (Koordination)", "ov", 45000, D("2026-07-13")),
           ("SL-26-009", D("2026-07-20"), 202, "Bescheinigungen § 35a EStG 2025 (6 Nachzügler)", "w202", 15000, D("2026-07-23")),
           ("SL-26-010", D("2026-08-05"), 201, "Verwalterzustimmung § 12 WEG – Veräußerung Einheit 201-25 (Brockmann)", "brockmann", 20000, D("2026-08-10")),
           ("SL-26-011", D("2026-08-20"), 203, "Verwalterzustimmung § 12 WEG – Veräußerung Einheit 203-11 (Kowalski/Hoffmeister)", "w203", 20000, None),
           ("SL-26-012", D("2026-09-02"), 201, "Ortstermin Wasserschaden Rosenstraße 42 mit Versicherung", "w201", 0, None),
           ("SL-26-013", D("2026-10-08"), 204, "Außerordentliche Eigentümerversammlung Brunnenstraße 6 (Klage Hollmann, Dach)", "w204", 35000, None),
           ("SL-26-014", D("2026-10-01"), 204, "Baubetreuung Dachsanierung Brunnenstraße 6 (2 % der Bausumme)", "w204", 36800, None)]
    for no, d, obj, what, cl, net, billed in sls:
        s = dict(no=no, date=d, obj=obj, what=what, client=cl, net=net, billed_on=billed, inv=None,
                 status="laufend" if no == "SL-26-014" else "erledigt", free=net == 0)
        if billed:
            s["inv"] = inv_(billed, cl, what, net, ref=no, kind="SL", sl=s, period=None)
        SL.append(s)
    INV.sort(key=lambda i: (i["date"], i["kind"] != "VV", i["order"], i["ref"]))
    for k, i in enumerate(INV, start=1):
        i["no"] = f"RE-2026-{k:04d}"
        IK[i["no"]] = i
        i["claim"] = claim(i["no"], "re", CLIENTS[i["client"]]["kd"], i["gross"], i["due"], i["date"], i["subject"], "K0",
                           name=CLIENTS[i["client"]]["name"])


def inv_of(vv: str, y: int, m: int) -> dict | None:
    return next((i for i in INV if i["ref"] == vv and i["period"] and i["period"][0] == date(y, m, 1)), None)


# ------------------------------------------------------------------ payments: rents
JOBCENTER = dict(name="Jobcenter Landkreis Brakenfeld", bank="BUNDESKASSE BRAKENFELD",
                 iban=iban_de("29950010", "8800310450"), bic=BANKS["SPK"][1])
RENT_LS_BATCHES: list[dict] = []


def rent_vwz(L, y, m, d=None) -> str:
    st = L["style"]
    obj = PD[L["obj"]]["street"] if L["obj"] in PD else L["lage"]
    if st == "static":
        return f"Miete {obj} {L['who']['nach']}"
    if st == "name":
        return f"{L['who']['bankname']} Miete"
    if st == "konto":
        return f"Mieternr {L['konto'] if L['konto'] != '10507' else '10570'} Miete {MONATE[m - 1]}"
    if st == "short":
        return f"Mietzahlung {MON3[m - 1]}. {y % 100}"
    if st == "gew":
        return f"Miete {MONATE[m - 1]} {y} {obj} zzgl. NK"
    return f"Miete {m:02d}/{y} {obj} Whg {L['lage']}"


def payer_of(L):
    return L["who"]["bankname"].upper() if L["konto"] == "10601" else L["who"]["bankname"]


def pay_rent(L, d, amount, vwz, claims, kind="rent", **kw):
    if d > LAST:
        return None
    a = acct_of_lease(L)
    return money(a, d, amount, payer_of(L), L["who"]["iban"], L["who"]["bic"], vwz,
                 kw.pop("btext", "GUTSCHR. UEBERW. DAUERAUFTR" if L["how"].startswith("DA") else "GUTSCHR. UEBERWEISUNG"),
                 alloc=claims, kind=kind, **kw)


def build_rent_money():
    ls_items: dict[tuple, list] = {}
    for L in LEASES:
        k = L["konto"]
        for (y, m) in MONTHS:
            c = CK.get(mkey(k, "Miete", y, m))
            if c is None:
                continue
            if k in SPECIAL_RENT:
                continue
            amt_paid = c["amount"]
            if L["kd_changes"]:  # a standing order keeps the amount it was set up with (H04/H05/H06)
                fixed = total_of(L, L["sched"][0][1:])
                for ch in L["kd_changes"]:
                    if ch["frm"] <= date(y, m, 1) and ch["noted"] and ch["kind"] in ("Mieterhöhung § 558",):
                        fixed = total_of(L, ch["new"][1:])  # agreed: the tenant changed the order
                    elif ch["frm"] <= date(y, m, 1) and ch["frm"] < D("2026-01-01"):
                        fixed = total_of(L, ch["new"][1:])
                amt_paid = min(fixed, c["amount"]) if L["how"].startswith(("DA", "UE")) else c["amount"]
            if L["start"] > date(y, m, 1):  # first month pro rata: paid by transfer
                amt_paid = c["amount"]
            how = L["how"]
            if how == "LS":
                ls_items.setdefault((y, m), []).append((L, c))
                continue
            if how == "JC":
                jd = prev_bd(date(*ym_prev(y, m), 28)) if (y, m) != (2026, 1) else D("2025-12-29")
                share = min(L["jc"]["share"], c["amount"])
                if jd <= LAST:
                    money("K1", jd, share, JOBCENTER["bank"], JOBCENTER["iban"], JOBCENTER["bic"],
                          f"BG-NR {L['jc']['bg']} KDU {m:02d}.{y % 100} {sepa_name(L['who']['nach']).upper()} "
                          f"{sepa_name(L['who']['vor']).upper()} {sepa_name(PD[L['obj']]['street']).upper()}",
                          "GUTSCHR. UEBERWEISUNG", alloc=[(c, share)], kind="rent_jc")
                if c["amount"] > share:
                    pay_rent(L, nth_bd(y, m, 1), c["amount"] - share, f"Restmiete {MONATE[m - 1]} {L['who']['nach']}",
                             [(c, c["amount"] - share)])
                continue
            if how == "Q":
                if m in (1, 4, 7, 10):
                    cs = [CK[mkey(k, "Miete", y, mm)] for mm in (m, m + 1, m + 2) if mkey(k, "Miete", y, mm) in CK]
                    pay_rent(L, nth_bd(y, m, 1), sum(x["amount"] for x in cs), f"Stellplatz S07 Quartal {(m + 2) // 3}/{y}",
                             [(x, x["amount"]) for x in cs])
                continue
            r = rr("pay", k, y, m)
            if how == "DA_1":
                d = nth_bd(y, m, 1)
            elif how == "DA_3":
                d = nth_bd(y, m, 3)
            elif how == "DA_28":
                py, pm = ym_prev(y, m)
                d = next_bd(date(py, pm, 28)) if (py, pm) != (2025, 12) else D("2025-12-29")
            elif how == "UE3":
                d = nth_bd(y, m, 3)
            else:  # manual transfer: the first days of the month, now and then a little late
                d = nth_bd(y, m, r.choices([1, 2, 3, 4, 5, 7], weights=[20, 25, 25, 12, 10, 8])[0])
            pay_rent(L, d, amt_paid, rent_vwz(L, y, m, d), [(c, amt_paid)], kind="prev" if d < D("2026-01-01") else "rent")
    # the office collects these rents by direct debit (one collective booking, the items in the pain.008)
    for (y, m), items in sorted(ls_items.items()):
        d = werktag3(y, m)
        if d > LAST:
            continue
        good = []
        for L, c in items:
            good.append(dict(L=L, c=c, amount=c["amount"], e2e=f"M{L['konto']}-{y}{m:02d}", mandate=f"KHV-M{L['konto']}",
                             vwz=f"Miete {m:02d}/{y} Mieternr. {L['konto']} {PD[L['obj']]['street']} {L['lage']}"))
        total = sum(x["amount"] for x in good)
        b = money("K1", d, total, "", "", "", f"SEPA-LASTSCHRIFTEN {len(good)} POSTEN EINREICHUNG VOM {dde(prev_bd(d - timedelta(days=2)))}",
                  "SEPA-LS EINREICHUNG", alloc=[(x["c"], x["amount"]) for x in good], kind="rent_ls",
                  ci=ACCTS["K1"]["ci"], sammler=f"MIETE-{y}-{m:02d}", batch=good)
        RENT_LS_BATCHES.append(dict(date=d, items=good, money=b, acct="K1", what="Mieteinzug"))
        # H09: the September debit of 10802 comes back (AM04), the fee in the same line
        for x in good:
            if x["L"]["konto"] == "10802" and (y, m) == (2026, 9):
                rd = D("2026-09-08")
                ret = money("K1", rd, -(x["amount"] + 300), payer_of(x["L"]), x["L"]["who"]["iban"], x["L"]["who"]["bic"],
                            f"RUECKLASTSCHRIFT {x['vwz']} AM04 KONTO UNGEDECKT", "RUECKLASTSCHRIFT", kind="return",
                            e2e=x["e2e"], mandate=x["mandate"], ci=ACCTS["K1"]["ci"], urspr=x["amount"], auslagen=300)
                ret["undo"] = (x["c"], x["amount"])
                pay_rent(x["L"], D("2026-09-21"), x["amount"], "Miete September nach Ruecklastschrift Hanke Heideweg 2a",
                         [(x["c"], x["amount"])], btext="GUTSCHR. UEBERWEISUNG")
    special_rent()


SPECIAL_RENT = {"10104", "10107", "10115", "10305", "10407", "10408", "10502", "10609", "10611", "10612", "10614", "10619",
                "10817"}


def std_pay(L, y, m, amount=None, d=None, vwz=None):
    c = CK[mkey(L["konto"], "Miete", y, m)]
    amount = c["amount"] if amount is None else amount
    if d is None:
        d = nth_bd(y, m, 1) if L["how"] != "DA_28" else next_bd(date(*ym_prev(y, m), 28))
    return pay_rent(L, d, amount, vwz or rent_vwz(L, y, m), [(c, min(amount, c["amount"]))])


def special_rent():
    G = LD
    # H04 10104: standing order at the old amount after the Staffel
    for y, m in MONTHS:
        if (y, m) <= (2026, 10):
            std_pay(G["10104"], y, m, amount=86000)
    # H10 10107: double September payment, offset with October
    for y, m in MONTHS:
        if (y, m) <= (2026, 9):
            std_pay(G["10107"], y, m)
    dbl = pay_rent(G["10107"], D("2026-09-07"), 69000, "Miete September Hartmann", [], kind="rent")
    dbl["note"] = "doppelt: Dauerauftrag 01.09. und Überweisung 07.09."
    dbl["alloc"].append((CK[mkey("10107", "Miete", 2026, 10)], 69000, D("2026-09-22")))
    CK[mkey("10107", "Miete", 2026, 10)]["gen_paid"] += 69000
    dbl["over_until"] = D("2026-09-22")
    # H11 10115: move-in on 15.08., August pro rata + deposit 1 in one transfer, deposit 2 missing until 12.10.
    L = G["10115"]
    aug, r1 = CK[mkey("10115", "Miete", 2026, 8)], CK["10115 Kaution Rate 1"]
    pay_rent(L, D("2026-08-14"), aug["amount"] + r1["amount"], "Miete August anteilig und Kaution 1. Rate Demir",
             [(aug, aug["amount"]), (r1, r1["amount"])], btext="GUTSCHR. UEBERWEISUNG")
    for m in (9, 10):
        std_pay(L, 2026, m, vwz=f"{L['who']['bankname']} Miete Lindenallee 14 Whg 5")
    r2, r3 = CK["10115 Kaution Rate 2"], CK["10115 Kaution Rate 3"]
    pay_rent(L, D("2026-10-12"), r2["amount"] + r3["amount"], "Kaution Rate 2 und 3 Demir Lindenallee 14",
             [(r2, r2["amount"]), (r3, r3["amount"])], btext="GUTSCHR. UEBERWEISUNG")
    # H03 10305: 600 instead of 812 since June
    L = G["10305"]
    for y, m in MONTHS:
        if (y, m) > (2026, 10):
            continue
        c = CK[mkey("10305", "Miete", y, m)]
        d = nth_bd(y, m, 2)
        if m >= 6:
            pay_rent(L, d, 60000, f"Miete {MONATE[m - 1]} Teilzahlung Böhm", [(c, 60000)])
        else:
            pay_rent(L, d, c["amount"], rent_vwz(L, y, m), [(c, c["amount"])])
    # H14 10407 café with VAT: September late, together with October
    L = G["10407"]
    for y, m in MONTHS:
        if (y, m) <= (2026, 8):
            c = CK[mkey("10407", "Miete", y, m)]
            pay_rent(L, nth_bd(y, m, 3), c["amount"], rent_vwz(L, y, m), [(c, c["amount"])])
    s, o = CK[mkey("10407", "Miete", 2026, 9)], CK[mkey("10407", "Miete", 2026, 10)]
    pay_rent(L, D("2026-10-12"), s["amount"] + o["amount"], "Miete Sept. + Okt. 2026 Bahnhofstr. 22 Cafe Kornblume",
             [(s, s["amount"]), (o, o["amount"])])
    # H14 10408 index rent from October, standing order old
    L = G["10408"]
    for y, m in MONTHS:
        if (y, m) <= (2026, 10):
            std_pay(L, y, m, amount=119000, vwz="Miete Praxis Bahnhofstr. 22 Lindner Sass")
    # H13 10502 offsets the BK credit in October
    L = G["10502"]
    for y, m in MONTHS:
        if (y, m) <= (2026, 9):
            std_pay(L, y, m, d=nth_bd(y, m, 2))
    oc = CK[mkey("10502", "Miete", 2026, 10)]
    pay_rent(L, nth_bd(2026, 10, 1), oc["amount"] - 8640, "Miete Okt 2026 abzgl. Guthaben BK 2025 Wübbena",
             [(oc, oc["amount"])])  # the credit pays the rest
    BKRES["10502"]["offset"] = (oc, 8640, nth_bd(2026, 10, 1))
    # H07 10609: the tenant's 38.50 missing in August and September, paid on 05.10. with October
    L = G["10609"]
    for y, m in MONTHS:
        if (y, m) > (2026, 11):
            continue
        c = CK[mkey("10609", "Miete", y, m)]
        jd = prev_bd(date(*ym_prev(y, m), 28)) if (y, m) != (2026, 1) else D("2025-12-29")
        if jd <= LAST:
            money("K1", jd, 56000, JOBCENTER["bank"], JOBCENTER["iban"], JOBCENTER["bic"],
                  f"BG-NR {L['jc']['bg']} KDU {m:02d}.{y % 100} SADIKU ARBEN GARTENSTRASSE 31", "GUTSCHR. UEBERWEISUNG",
                  alloc=[(c, 56000)], kind="rent_jc", kind_jc=True)
        if (y, m) <= (2026, 7):
            pay_rent(L, nth_bd(y, m, 1), c["amount"] - 56000, f"Restmiete {MONATE[m - 1]} Sadiku", [(c, c["amount"] - 56000)])
    rest = [CK[mkey("10609", "Miete", 2026, m)] for m in (8, 9, 10)]
    pay_rent(L, D("2026-10-05"), sum(c["amount"] - 56000 for c in rest), "Restmiete Aug Sep Okt Sadiku",
             [(c, c["amount"] - 56000) for c in rest], btext="GUTSCHR. UEBERWEISUNG")
    # H07 10614: Jobcenter stops in September, pays September and October on 12.10.
    L = G["10614"]
    for y, m in MONTHS:
        c = CK.get(mkey("10614", "Miete", y, m))
        if not c or (y, m) in ((2026, 9), (2026, 10)):
            continue
        jd = prev_bd(date(*ym_prev(y, m), 28)) if (y, m) != (2026, 1) else D("2025-12-29")
        if jd <= LAST:
            money("K1", jd, c["amount"], JOBCENTER["bank"], JOBCENTER["iban"], JOBCENTER["bic"],
                  f"BG-NR {L['jc']['bg']} KDU {m:02d}.{y % 100} WAGNER JACQUELINE GARTENSTRASSE 31", "GUTSCHR. UEBERWEISUNG",
                  alloc=[(c, c["amount"])], kind="rent_jc" if jd >= D("2026-01-01") else "prev")
    s, o = CK[mkey("10614", "Miete", 2026, 9)], CK[mkey("10614", "Miete", 2026, 10)]
    money("K1", D("2026-10-12"), s["amount"] + o["amount"], JOBCENTER["bank"], JOBCENTER["iban"], JOBCENTER["bic"],
          f"BG-NR {L['jc']['bg']} KDU 09+10/2026 NACHZAHLUNG WAGNER JACQUELINE", "GUTSCHR. UEBERWEISUNG",
          alloc=[(s, s["amount"]), (o, o["amount"])], kind="rent_jc")
    # H08 10611: moved out 31.08., the Jobcenter pays September anyway; refunded on 22.09.
    L = G["10611"]
    for y, m in MONTHS:
        c = CK.get(mkey("10611", "Miete", y, m))
        jd = prev_bd(date(*ym_prev(y, m), 28)) if (y, m) != (2026, 1) else D("2025-12-29")
        if c:
            money("K1", jd, c["amount"], JOBCENTER["bank"], JOBCENTER["iban"], JOBCENTER["bic"],
                  f"BG-NR {L['jc']['bg']} KDU {m:02d}.{y % 100} KRUEGER DENNIS GARTENSTRASSE 31", "GUTSCHR. UEBERWEISUNG",
                  alloc=[(c, c["amount"])], kind="rent_jc" if jd >= D("2026-01-01") else "prev")
    wrong = money("K1", D("2026-08-28"), 53400, JOBCENTER["bank"], JOBCENTER["iban"], JOBCENTER["bic"],
                  f"BG-NR {L['jc']['bg']} KDU 09.26 KRUEGER DENNIS GARTENSTRASSE 31", "GUTSCHR. UEBERWEISUNG",
                  alloc=[], kind="rent_jc", note="Jobcenter zahlt September für den zum 31.08. ausgezogenen Mieter")
    back = money("K1", D("2026-09-22"), -53400, JOBCENTER["name"], JOBCENTER["iban"], JOBCENTER["bic"],
                 f"Rueckzahlung KdU 09/2026 BG {L['jc']['bg']} Krueger Dennis - Mieter zum 31.08.2026 ausgezogen",
                 "ONLINE-UEBERWEISUNG", kind="refund", refund_of=wrong["id"])
    wrong["settled"].append((back["date"], 53400, "an das Jobcenter zurückgezahlt"))
    # H02 10612: nothing from July, 1.000 on 01.10. without a month
    L = G["10612"]
    for y, m in MONTHS:
        if (y, m) <= (2026, 6):
            c = CK[mkey("10612", "Miete", y, m)]
            pay_rent(L, nth_bd(y, m, rr("p", m).choice([1, 2, 3])), c["amount"], rent_vwz(L, y, m), [(c, c["amount"])])
    pay_rent(L, D("2026-10-01"), 100000, "Rueckstand Teilzahlung Brandl", oldest("10612", 100000, D("2026-10-01")))
    # H19 10817 (Heideweg) still pays October after the mandate ended
    L = G["10817"]
    for y, m in MONTHS:
        if (y, m) <= (2026, 9):
            std_pay(L, y, m)
    late = pay_rent(L, D("2026-10-01"), total_of(L, rent_at(L, D("2026-10-01"))), rent_vwz(L, 2026, 10),
                    [], kind="rent", note="Heideweg wird seit 01.10. nicht mehr von uns verwaltet – an den Nachverwalter weiterleiten")
    # H08 the new tenant of 106-11 pays the first deposit instalment before moving in
    L = G["10619"]
    r1 = CK["10619 Kaution Rate 1"]
    pay_rent(L, D("2026-10-27"), r1["amount"], "Kaution Rate 1 Gartenstr. 31 Whg 11 Popescu", [(r1, r1["amount"])],
             btext="GUTSCHR. UEBERWEISUNG")


# ------------------------------------------------------------------ payments: results, deposits, Hausgeld, fees, costs
def build_results_money():
    # BK 2025: Nachzahlungen paid within the month, Guthaben paid out about three weeks after the statement
    for k, res in sorted(BKRES.items()):
        L = LD[k]
        c = CK.get(f"{k} BK 2025")
        if res["res"] > 0:
            if k == "10304":
                pay_rent(L, D("2026-10-02"), c["amount"], "Nachzahlung Betriebskosten 2025 Eilers", [(c, c["amount"])],
                         btext="GUTSCHR. UEBERWEISUNG")
            elif k == "10205":
                for d, a in ((D("2026-07-20"), 15000), (D("2026-08-20"), 15000), (D("2026-09-21"), 11290)):
                    pay_rent(L, d, a, "Rate BK-Abrechnung 2025 Gerdes", [(c, a)], btext="GUTSCHR. UEBERWEISUNG")
            elif k in ("10703", "10612"):
                continue  # disputed; the defaulter does not pay
            elif k == "10305":
                continue  # Böhm cannot pay
            else:
                d = add_bd(res["sent"], rr("bkp", k).randint(5, 24))
                if d <= LAST:
                    pay_rent(L, d, c["amount"], rr("bkt", k).choice(["Nachzahlung NK 2025", "BK-Abrechnung 2025",
                                                                     f"Betriebskosten 2025 Mieternr {k}",
                                                                     "Nebenkostennachzahlung 2025"]),
                             [(c, c["amount"])], btext="GUTSCHR. UEBERWEISUNG")
        elif k != "10502":
            d = add_bd(res["sent"], 15)
            if d <= LAST:
                money(acct_of_lease(L), d, res["res"], L["who"]["bankname"], L["who"]["iban"], L["who"]["bic"],
                      f"Guthaben Betriebskostenabrechnung 2025 Mieternr. {k}", "ONLINE-UEBERWEISUNG", kind="bk_credit")
    # Jahresabrechnung 2025: Nachschüsse; Guthaben paid out
    for k, res in sorted(JARES.items()):
        o = OD[k]
        p = PD[o["weg"]]
        c = CK.get(f"{k} Abrechnung 2025")
        if res["res"] > 0:
            if k == "20404":
                continue
            if k == "20211":
                money(p["acct"], D("2026-07-20"), 20000, o["who"]["bankname"], o["who"]["iban"], o["who"]["bic"],
                      "Abrechnung 2025 Teilbetrag Bergmann Whg 11", "GUTSCHR. UEBERWEISUNG", alloc=[(c, 20000)], kind="ja")
                money(p["acct"], D("2026-10-15"), 14275, o["who"]["bankname"], o["who"]["iban"], o["who"]["bic"],
                      "Rest Abrechnung 2025 Bergmann", "GUTSCHR. UEBERWEISUNG", alloc=[(c, 14275)], kind="ja")
                continue
            if o["how"] == "SE":
                d = add_bd(res["due"], -3)
                money("K1", d, -c["amount"], p["name"], ACCTS[p["acct"]]["iban"], ACCTS[p["acct"]]["bic"],
                      f"Abrechnungsspitze 2025 Whg {o['unit']} {o['who']['nach']} (SE-Verwaltung)", "ONLINE-UEBERWEISUNG",
                      kind="internal_out")
                money(p["acct"], d, c["amount"], ACCTS["K1"]["holder"], ACCTS["K1"]["iban"], ACCTS["K1"]["bic"],
                      f"Abrechnungsspitze 2025 Whg {o['unit']} {o['who']['nach']} (SE-Verwaltung)", "GUTSCHR. UEBERWEISUNG",
                      alloc=[(c, c["amount"])], kind="ja_se")
                continue
            d = add_bd(res["etv"], rr("jap", k).randint(6, 22))
            money(p["acct"], d, c["amount"], o["who"]["bankname"], o["who"]["iban"], o["who"]["bic"],
                  rr("jat", k).choice(["Abrechnungsspitze 2025", "Nachzahlung Hausgeldabrechnung 2025",
                                       f"Jahresabrechnung 2025 Whg {o['unit']}", f"HG-Abrechnung 2025 {o['who']['nach']}"]),
                  "GUTSCHR. UEBERWEISUNG", alloc=[(c, c["amount"])], kind="ja")
        else:
            d = add_bd(res["etv"], 12)
            to = (ACCTS["K1"] if o["how"] == "SE" else None)
            money(p["acct"], d, res["res"], to["holder"] if to else o["who"]["bankname"], to["iban"] if to else o["who"]["iban"],
                  to["bic"] if to else o["who"]["bic"], f"Guthaben Jahresabrechnung 2025 Whg {o['unit']}",
                  "ONLINE-UEBERWEISUNG", kind="ja_credit")
            if to:
                money("K1", d, -res["res"], p["name"], ACCTS[p["acct"]]["iban"], ACCTS[p["acct"]]["bic"],
                      f"Guthaben Jahresabrechnung 2025 Whg {o['unit']}", "GUTSCHR. UEBERWEISUNG", kind="internal_in")
    # Sonderumlage Brunnenstraße 6, due 01.09.
    for k, share in sorted(SU204.items()):
        o = OD[k]
        c = CK[f"{k} Sonderumlage 2026"]
        if k == "20404":
            continue
        if k == "20402":
            for d in (D("2026-09-01"), D("2026-10-01")):
                money("K6", d, share // 2, o["who"]["bankname"], o["who"]["iban"], o["who"]["bic"],
                      f"Sonderumlage Dach {'1. Haelfte' if d.month == 9 else 'Rest'} Kranz Whg 2", "GUTSCHR. UEBERWEISUNG",
                      alloc=[(c, share // 2)], kind="su")
            continue
        d = next_bd(D("2026-08-24") + timedelta(days=rr("su", k).randint(0, 9)))
        money("K6", d, share, o["who"]["bankname"], o["who"]["iban"], o["who"]["bic"],
              rr("sut", k).choice(["Sonderumlage Dachsanierung", f"Sonderumlage 2026 Whg {o['unit']}",
                                   "Sonderumlage Dach Brunnenstr. 6"]), "GUTSCHR. UEBERWEISUNG", alloc=[(c, share)], kind="su")
    # deposits of tenants before 2026 sit on deposit accounts; Voigt's comes back (H11)
    L = LD["10105"]
    money("K1", D("2026-09-14"), 183385, "Kerkhoff Hausverwaltung GmbH Mietkaution Voigt", acct("dep", "10105", bank="SPK")[0],
          BANKS["SPK"][1], "Aufloesung Mietkautionskonto Voigt Rainer Lindenallee 14 Whg 5 inkl. Zinsen 3,85",
          "GUTSCHR. UEBERWEISUNG", kind="own_in", note="Kautionskonto aufgelöst")
    money("K1", D("2026-09-15"), -153385, L["who"]["bankname"], L["who"]["iban"], L["who"]["bic"],
          "Rueckzahlung Mietkaution Lindenallee 14 Whg 5 abzgl. Einbehalt 300,00 fuer BK 2026", "ONLINE-UEBERWEISUNG",
          kind="deposit_back")
    for k in ("10115",):  # the new deposit goes to its own account
        for d in (D("2026-08-18"), D("2026-10-14")):
            amt = CK[f"{k} Kaution Rate 1"]["amount"] if d.month == 8 else 2 * CK[f"{k} Kaution Rate 2"]["amount"]
            money("K1", d, -amt, "Kerkhoff Hausverwaltung GmbH Mietkaution Demir", acct("dep", k, bank="SPK")[0],
                  BANKS["SPK"][1], f"Anlage Mietkaution Demir Selin Lindenallee 14 Whg 5", "ONLINE-UEBERWEISUNG", kind="own_out")


HG_LS_BATCHES: list[dict] = []


def build_hausgeld_money():
    batches: dict[tuple, list] = {}
    for o in OWN:
        k, p = o["konto"], PD[o["weg"]]
        for (y, m) in MONTHS:
            c = CK.get(mkey(k, "Hausgeld", y, m))
            if not c or (k == "20404" and m >= 5) or k == "20514":
                continue  # H21 pays nothing since May; 20514 pays below (H27)
            how = o["how"]
            if how == "LS":
                batches.setdefault((p["no"], y, m), []).append((o, c))
                continue
            if how == "SE":
                d = nth_bd(y, m, 4)
                if d <= LAST:
                    money("K1", d, -c["amount"], p["name"], ACCTS[p["acct"]]["iban"], ACCTS[p["acct"]]["bic"],
                          f"Hausgeld {m:02d}/{y} Whg {o['unit']} {o['who']['nach']} (SE-Verwaltung)", "ONLINE-UEBERWEISUNG",
                          kind="internal_out")
                    money(p["acct"], d, c["amount"], ACCTS["K1"]["holder"], ACCTS["K1"]["iban"], ACCTS["K1"]["bic"],
                          f"Hausgeld {m:02d}/{y} Whg {o['unit']} {o['who']['nach']} (SE-Verwaltung)", "GUTSCHR. UEBERWEISUNG",
                          alloc=[(c, c["amount"])], kind="hg_se")
                continue
            if how == "PAIR":
                if k == "20604":
                    continue
                c2 = CK[mkey("20604", "Hausgeld", y, m)]
                d = nth_bd(y, m, 1)
                if d <= LAST:
                    money(p["acct"], d, c["amount"] + c2["amount"], o["who"]["bankname"], o["who"]["iban"], o["who"]["bic"],
                          "Hausgeld Whg 3 + 4 Petersen", "GUTSCHR. UEBERW. DAUERAUFTR",
                          alloc=[(c, c["amount"]), (c2, c2["amount"])], kind="hg")
                continue
            r = rr("hg", k, y, m)
            if how == "DA_1":
                d = nth_bd(y, m, 1)
            elif how == "DA_28":
                d = next_bd(date(*ym_prev(y, m), 28)) if (y, m) != (2026, 1) else D("2025-12-29")
            else:
                d = nth_bd(y, m, r.choice([1, 2, 2, 3, 3, 4, 6]))
            if p["no"] == 205 and d < D("2026-07-01"):
                d = D("2026-07-01")
            if d > LAST:
                continue
            vwz = r.choice([f"Hausgeld {m:02d}/{y}", f"Hausgeld {MONATE[m - 1]} Whg {o['unit']}",
                            f"WEG {p['street']} Hausgeld {o['who']['nach']}", f"HG {MON3[m - 1]} {y} ET {k}"]) \
                if how == "UE" else f"Hausgeld {p['street']} Whg {o['unit']} {o['who']['nach']}"
            money(p["acct"], d, c["amount"], o["who"]["bankname"], o["who"]["iban"], o["who"]["bic"], vwz,
                  "GUTSCHR. UEBERW. DAUERAUFTR" if how.startswith("DA") else "GUTSCHR. UEBERWEISUNG",
                  alloc=[(c, c["amount"])], kind="hg" if d >= D("2026-01-01") else "prev")
    for (wno, y, m), items in sorted(batches.items()):
        p = PD[wno]
        d = werktag3(y, m)
        if d > LAST:
            continue
        good = [dict(o=o, c=c, amount=c["amount"], e2e=f"HG{o['konto']}{y}{m:02d}", mandate=o["mandate"],
                     vwz=f"Hausgeld {m:02d}/{y} {p['name']} Einheit {o['unit']} {o['lage']}") for o, c in items]
        b = money(p["acct"], d, sum(x["amount"] for x in good), p["name"], "", "",
                  f"SEPA-Lastschrifteinzug Hausgeld {m:02d}/{y} {len(good)} Posten", "SAMMEL-LS-GUTSCHRIFT",
                  alloc=[(x["c"], x["amount"]) for x in good], kind="hg_ls", ci=ACCTS[p["acct"]]["ci"],
                  sammler=f"HG-{wno}-{y}{m:02d}", batch=good, details=wno == 203)
        HG_LS_BATCHES.append(dict(date=d, items=good, money=b, acct=p["acct"], what=f"Hausgeld {p['name']}", wno=wno))
        for x in good:  # H23: Dr. Riedel objects to the July debit on 14.09. (MD06)
            if x["o"]["konto"] == "20127" and (y, m) == (2026, 7):
                ret = money(p["acct"], D("2026-09-14"), -x["amount"], x["o"]["who"]["bankname"], x["o"]["who"]["iban"],
                            x["o"]["who"]["bic"], f"Ruecklastschrift {x['vwz']} MD06 Widerspruch des Zahlungspflichtigen",
                            "RUECKLASTSCHRIFT", kind="return", e2e=x["e2e"], mandate=x["mandate"], ci=ACCTS[p["acct"]]["ci"],
                            rtr="MD06")
                ret["undo"] = (x["c"], x["amount"])
                money(p["acct"], D("2026-09-14"), -250, "", "", "", "Entgelt Ruecklastschrift", "ENTGELT", kind="fee")
    # H23 Riedel pays July by transfer on 02.10.; H24 Kowalski pays September although he sold; H27 20514 pays the old arrears
    o = OD["20127"]
    money("K3", D("2026-10-02"), 28700, o["who"]["bankname"], o["who"]["iban"], o["who"]["bic"],
          "Hausgeld Juli nach Ruecklastschrift Riedel Whg 27", "GUTSCHR. UEBERWEISUNG",
          alloc=[(CK[mkey("20127", "Hausgeld", 2026, 7)], 28700)], kind="hg")
    o = OD["20311"]
    k = money("K5", D("2026-09-01"), 26800, o["who"]["bankname"], o["who"]["iban"], o["who"]["bic"],
              "Hausgeld September Kowalski Whg 11", "GUTSCHR. UEBERWEISUNG", kind="hg",
              note="Verkäufer zahlt nach dem Besitzübergang (01.09.) – erstatten")
    back = money("K5", D("2026-09-25"), -26800, o["who"]["bankname"], o["who"]["iban"], o["who"]["bic"],
                 "Erstattung Hausgeld 09/2026 Whg 11 - Besitzuebergang 01.09.2026", "ONLINE-UEBERWEISUNG", kind="refund",
                 refund_of=k["id"])
    k["settled"].append((back["date"], 26800, "erstattet"))
    o = OD["20514"]
    cs = [CK[mkey("20514", "Hausgeld", 2026, m)] for m in (5, 6)]
    money("K7", D("2026-10-14"), 55200, o["who"]["bankname"], o["who"]["iban"], o["who"]["bic"],
          "Rueckstand Hausgeld Mai Juni Ivanova", "GUTSCHR. UEBERWEISUNG", alloc=[(c, c["amount"]) for c in cs], kind="hg")
    for y, m in MONTHS:  # 20514 pays the new administrator from July on
        c = CK.get(mkey("20514", "Hausgeld", y, m))
        if c and m >= 7 and nth_bd(y, m, 2) <= LAST:
            money("K7", nth_bd(y, m, 2), c["amount"], o["who"]["bankname"], o["who"]["iban"], o["who"]["bic"],
                  f"Hausgeld {MONATE[m - 1]} Feldkamp 19 Whg 14", "GUTSCHR. UEBERWEISUNG", alloc=[(c, c["amount"])], kind="hg")


def build_fee_money():
    """The fee is taken from the object's account to the office's account – the same money twice in the folder (H30)."""
    for (y, m) in MONTHS[:10]:
        d = nth_bd(y, m, 4)
        if d > LAST:
            continue
        k1 = [inv_of(md["vv"], y, m) for md in MANDATES if md["acct"] == "K1"]
        k1 = [i for i in k1 if i]
        if k1:
            total = sum(i["gross"] for i in k1)
            vwz = f"Verwalterverguetung {m:02d}/{y} lt. {k1[0]['no']} bis {k1[-1]['no']}"
            money("K1", d, -total, FIRM["name"], ACCTS["K0"]["iban"], ACCTS["K0"]["bic"], vwz, "ONLINE-UEBERWEISUNG",
                  kind="fee_out")
            money("K0", d, total, ACCTS["K1"]["holder"], ACCTS["K1"]["iban"], ACCTS["K1"]["bic"], vwz, "Gutschrift",
                  alloc=[(i["claim"], i["gross"]) for i in k1], kind="fee")
        for md in MANDATES:
            if md["acct"] in (None, "K1"):
                continue
            i = inv_of(md["vv"], y, m)
            if not i:
                continue
            if md["vv"] == "VV-204" and m in (8, 9, 10) or md["vv"] == "VV-205":
                continue  # no money on the account (H32); Feldkamp below (H34)
            a = ACCTS[md["acct"]]
            vwz = f"{i['no']} Verwalterverguetung {m:02d}/{y}"
            money(md["acct"], d, -i["gross"], FIRM["name"], ACCTS["K0"]["iban"], ACCTS["K0"]["bic"], vwz,
                  "ONLINE-UEBERWEISUNG", kind="fee_out")
            money("K0", d, i["gross"], a["holder"], a["iban"], a["bic"], vwz, "Gutschrift", alloc=[(i["claim"], i["gross"])],
                  kind="fee")
    # H32 Brunnenstraße: August, September and October on 05.10.
    invs = [inv_of("VV-204", 2026, m) for m in (8, 9, 10)]
    vwz = "Verwalterverguetung " + " ".join(i["no"] for i in invs)
    money("K6", D("2026-10-05"), -sum(i["gross"] for i in invs), FIRM["name"], ACCTS["K0"]["iban"], ACCTS["K0"]["bic"], vwz,
          "ONLINE-UEBERWEISUNG", kind="fee_out")
    money("K0", D("2026-10-05"), sum(i["gross"] for i in invs), ACCTS["K6"]["holder"], ACCTS["K6"]["iban"], ACCTS["K6"]["bic"],
          vwz, "Gutschrift", alloc=[(i["claim"], i["gross"]) for i in invs], kind="fee")
    # H34 Feldkamp: four invoices of 05.10., taken on 06.10.
    invs = [i for i in INV if i["ref"] == "VV-205"]
    vwz = "Verwalterverguetung 07-10/2026 " + ", ".join(i["no"] for i in invs)
    money("K7", D("2026-10-06"), -sum(i["gross"] for i in invs), FIRM["name"], ACCTS["K0"]["iban"], ACCTS["K0"]["bic"], vwz,
          "ONLINE-UEBERWEISUNG", kind="fee_out")
    money("K0", D("2026-10-06"), sum(i["gross"] for i in invs), ACCTS["K7"]["holder"], ACCTS["K7"]["iban"], ACCTS["K7"]["bic"],
          vwz, "Gutschrift", alloc=[(i["claim"], i["gross"]) for i in invs], kind="fee")
    # H33 Dr. Albers pays herself, quarterly and late
    al = CLIENTS["al"]
    for d, months in ((D("2026-04-14"), (1, 2, 3)), (D("2026-07-10"), (4, 5, 6)), (D("2026-10-08"), (7, 8, 9))):
        invs = [inv_of("VV-104", 2026, m) for m in months]
        money("K0", d, sum(i["gross"] for i in invs), "Dr. Henrike Albers", al["payout"][0], al["payout"][1],
              "Rg. " + ", ".join(i["no"] for i in invs), "Gutschrift", alloc=[(i["claim"], i["gross"]) for i in invs], kind="fee")
    # special services: paid by the client within two weeks; Brockmann does not pay (SL-26-010)
    for s in SL:
        i = s["inv"]
        if not i or s["no"] == "SL-26-010":
            continue
        cl = CLIENTS[i["client"]]
        src = next((md["acct"] for md in MANDATES if md["client"] == i["client"] and md["acct"]), None)
        d = add_bd(i["date"], 6)
        if src:
            a = ACCTS[src]
            money(src, d, -i["gross"], FIRM["name"], ACCTS["K0"]["iban"], ACCTS["K0"]["bic"], f"{i['no']} {i['subject'][:60]}",
                  "ONLINE-UEBERWEISUNG", kind="fee_out")
            money("K0", d, i["gross"], a["holder"], a["iban"], a["bic"], f"{i['no']} {i['subject'][:60]}", "Gutschrift",
                  alloc=[(i["claim"], i["gross"])], kind="fee")
        else:
            money("K0", d, i["gross"], cl["name"], cl["payout"][0], cl["payout"][1], f"Rechnung {i['no']}", "Gutschrift",
                  alloc=[(i["claim"], i["gross"])], kind="fee")


SUPPLIERS = {
    "falk": dict(name="Falk Bedachungen GmbH", street="Ziegeleiweg 12", zip="29599", city="Brakenfeld",
                 mail="rechnung@falk-bedachungen.example", phone="05861 77 40-0", vat=vat_id("31407726"),
                 hrb="HRB 18833, Amtsgericht Brakenfeld", bank="VB"),
    "rehm": dict(name="Heizungsbau Rehm GmbH", street="Kesselstraße 3", zip="29599", city="Brakenfeld",
                 mail="buchhaltung@rehm-heizung.example", phone="05861 52 10", vat=vat_id("28890415"),
                 hrb="HRB 17102, Amtsgericht Brakenfeld", bank="SPK"),
    "akin": dict(name="Sanitär Aksoy", street="Hafenstraße 21", zip="29599", city="Brakenfeld",
                 mail="info@sanitaer-aksoy.example", phone="0176 4410 2290", vat="", taxno="52/144/30812",
                 hrb="Inhaber Murat Aksoy", bank="NB"),
    "steinkamp": dict(name="Elektro Steinkamp GmbH", street="Am Umspannwerk 5", zip="29599", city="Brakenfeld",
                      mail="rechnungen@elektro-steinkamp.example", phone="05861 60 31-0", vat=vat_id("30155892"),
                      hrb="HRB 19005, Amtsgericht Brakenfeld", bank="SPK"),
    "gruen": dict(name="Grünwerk Tönjes GbR", street="Gärtnerweg 2", zip="29599", city="Brakenfeld",
                  mail="buero@gruenwerk-toenjes.example", phone="05861 41 99 7", vat=vat_id("33091264"),
                  hrb="Gesellschafter Ole und Hinnerk Tönjes", bank="VB"),
    "maler": dict(name="Malerbetrieb Hinrichsen", street="Farbenweg 9", zip="29599", city="Brakenfeld",
                  mail="info@maler-hinrichsen.example", phone="05861 21 08 3", vat=vat_id("27718340"),
                  hrb="Inhaber Jörn Hinrichsen", bank="SPK"),
}
for _k, _s in SUPPLIERS.items():
    _s["iban"], _s["bic"] = acct("sup", _k, bank=_s["bank"])
FRAUD = dict(mail="rechnung@falk-bedachung.example", iban=iban_lt("32500", "00177483922641"), bic="REVOLT21XXX")
SUPINV: list[dict] = []


def sup_inv(sup, no, d, obj, lines, due_days, fmt="pdf", paid=None, note="", to_mail=False, recipient=None):
    net = sum(round(q * p) for _, q, _, p in lines)
    x = dict(sup=sup, no=no, date=d, obj=obj, lines=lines, net=net, vat=vat19(net), gross=net + vat19(net),
             due=d + timedelta(days=due_days), fmt=fmt, paid=paid, note=note, to_mail=to_mail, recipient=recipient)
    SUPINV.append(x)
    return x


def recipient_of(obj: int) -> list[str]:
    p = PD[obj]
    if p["kind"] == "WEG":
        return [p["name"], "c/o Kerkhoff Hausverwaltung GmbH", FIRM["street"], f"{FIRM['zip']} {FIRM['city']}"]
    o = OWNERS[p["owner"]]
    return [o["name"], "vertreten durch Kerkhoff Hausverwaltung GmbH", FIRM["street"], f"{FIRM['zip']} {FIRM['city']}"]


def build_suppliers():
    sup_inv("akin", "2026-0815", D("2026-08-12"), 201, [("Rohrbruch Steigleitung Haus 42, Notdienst 10.08.2026", 1, "pausch.", 28000),
                                                        ("Austausch Kupferrohr DN 22, 3,5 m, Pressfittinge", 1, "pausch.", 162000),
                                                        ("Monteurstunden", 14, "Std.", 6200), ("Trocknungsgeräte 5 Tage", 5, "Tag", 4200)],
            14, paid=D("2026-08-20"))
    sup_inv("rehm", "R-26-10418", D("2026-09-08"), 103, [("Wartung Gas-Brennwertkessel 2 Anlagen", 2, "Stk.", 18900),
                                                         ("Abgasmessung", 2, "Stk.", 4500),
                                                         ("Ersatzteil Zündelektrode", 1, "Stk.", 3880)],
            14, fmt="zugferd", paid=D("2026-09-21"), to_mail=True)
    sup_inv("falk", "AR-2026-0112", D("2026-09-14"), 204, [("1. Abschlagsrechnung Dachsanierung Brunnenstraße 6 gem. Auftrag vom 30.07.2026",
                                                            1, "pausch.", 773109)], 14, paid=D("2026-09-30"))
    sup_inv("gruen", "GT-2026-311", D("2026-07-01"), 202, [("Gartenpflege 2. Quartal 2026 Am Wall 9", 1, "Quartal", 64500)],
            14, paid=D("2026-07-09"))
    sup_inv("gruen", "GT-2026-412", D("2026-10-01"), 202, [("Gartenpflege 3. Quartal 2026 Am Wall 9", 1, "Quartal", 64500)],
            14, paid=None)
    sup_inv("gruen", "GT-2026-413", D("2026-10-01"), 206, [("Gartenpflege 3. Quartal 2026 Mühlenstraße 2", 1, "Quartal", 51800)],
            14, paid=D("2026-10-13"))
    sup_inv("steinkamp", "ST-26-5521", D("2026-10-05"), 106, [("Austausch Treppenhausbeleuchtung LED mit Bewegungsmelder", 9, "Stk.", 8950),
                                                             ("Elektrikerstunden", 6, "Std.", 6400)], 14, fmt="xrechnung")
    sup_inv("maler", "1088", D("2026-09-11"), 106, [("Renovierung Wohnung 15 (Leerstand): Wände und Decken streichen", 1, "pausch.", 218000),
                                                    ("Lackierung Türen und Zargen", 6, "Stk.", 8500)], 21, paid=D("2026-09-29"))
    sup_inv("falk", "SR-2026-0131", D("2026-10-20"), 204, [("Schlussrechnung Dachsanierung Brunnenstraße 6", 1, "pausch.", 1546218),
                                                           ("abzgl. 1. Abschlag AR-2026-0112 netto", 1, "pausch.", -773109)],
            14, paid=None)
    sup_inv("rehm", "R-26-10655", D("2026-10-16"), 106, [("Störungsbeseitigung Heizung, Pumpe getauscht", 1, "pausch.", 42000)],
            14, fmt="zugferd")
    for x in SUPINV:
        p = PD[x["obj"]]
        a = p["acct"]
        if x["paid"] and x["paid"] <= LAST:
            s = SUPPLIERS[x["sup"]]
            money(a, x["paid"], -x["gross"], s["name"], s["iban"], s["bic"], f"Rechnung {x['no']} vom {dde(x['date'])}",
                  "ONLINE-UEBERWEISUNG", kind="sup")
    # H29: the insurer pays the water damage of Rosenhof
    money("K3", D("2026-09-02"), 348000, "Lindgrün Versicherung AG", iban_de("20730040", "0071004455"), BANKS["HB"][1],
          "Schaden-Nr. 2026-117834 Leitungswasser Rosenstr. 42 Regulierung", "GUTSCHR. UEBERWEISUNG", kind="insurance")


UTIL = [("Stadtwerke Brakenfeld GmbH", "Abschlag Wasser/Abwasser/Allgemeinstrom", creditor_id("00000104477"), 15),
        ("Stadtwerke Brakenfeld GmbH", "Abschlag Erdgas", creditor_id("00000104477"), 20)]


def build_costs():
    sw_iban, sw_bic = iban_de("29950010", "0000104477"), BANKS["SPK"][1]
    hm_iban, hm_bic = SUPPLIERS["gruen"]["iban"], SUPPLIERS["gruen"]["bic"]
    for p in PROPS:
        a = p["acct"]
        n = p["units"] + p.get("shops", 0)
        for (y, m) in MONTHS:
            for k, (who, what, ci, day) in enumerate(UTIL):
                if k == 1 and not p.get("heat", p["kind"] == "WEG"):
                    continue
                if p["no"] == 205 and m < 7:
                    continue
                if p.get("until") and date(y, m, 1) > p["until"]:
                    continue
                d = next_bd(date(y, m, day))
                if d > LAST:
                    continue
                amt = n * (2900 if k == 0 else 6100) + rr("u", p["no"], k).randint(0, 9) * 100
                money(a, d, -amt, who, sw_iban, sw_bic, f"{what} {m:02d}/{y} Vertragskonto {p['no']}04477-{k}",
                      "FOLGELASTSCHRIFT", kind="cost", ci=ci, mandate=f"SWB-{p['no']}-{k}")
            d = next_bd(date(y, m, 2))
            if d <= LAST and not (p["no"] == 205 and m < 7) and not (p.get("until") and date(y, m, 1) > p["until"]):
                money(a, d, -(n * 1150 + 4000), "Grünwerk Tönjes GbR", hm_iban, hm_bic,
                      f"Hausmeisterpauschale {m:02d}/{y} {p['street']}", "DAUERAUFTRAG", kind="cost")
            if p["kind"] == "WEG" and not (p["no"] == 205 and m < 7):
                d = nth_bd(y, m, 2)
                if d <= LAST:
                    money(a, d, -(n * 2200), p["name"] + " Rücklage", RESERVE[a], ACCTS[a]["bic"],
                          f"Zufuehrung Erhaltungsruecklage {m:02d}/{y}", "DAUERAUFTRAG", kind="own_out")
            if p["kind"] == "MV" and m in (2, 5, 8, 11):
                d = next_bd(date(y, m, 15))
                if d <= LAST and not (p.get("until") and date(y, m, 1) > p["until"]):
                    money(a, d, -(n * 3800), "Stadtkasse Brakenfeld", iban_de("29950010", "0000500100"), sw_bic,
                          f"Grundbesitzabgaben {(m + 1) // 3}. Quartal {y} Kassenzeichen 4402{p['no']}", "FOLGELASTSCHRIFT",
                          kind="cost", ci=creditor_id("00000500100"), mandate=f"GBA-{p['no']}")
            if m == 1:
                d = next_bd(date(y, 1, 15))
                money(a, d, -(n * 19000), "Lindgrün Versicherung AG", iban_de("20730040", "0071004455"), BANKS["HB"][1],
                      f"Verbundene Gebäudeversicherung 2026 Vers.-Nr. VG-{p['no']}-88", "FOLGELASTSCHRIFT", kind="cost",
                      ci=creditor_id("00000990011"), mandate=f"LV-{p['no']}") if p["no"] != 205 else None
    # payouts to the landlords (rent received less costs and fee), monthly on the 10th bank day
    for md in MANDATES:
        if md["vv"] in ("VV-201", "VV-202", "VV-203", "VV-204", "VV-205", "VV-206"):
            continue
        cl = CLIENTS[md["client"]]
        a = md["acct"] or "K1"
        for (y, m) in MONTHS[:10]:
            d = nth_bd(y, m, 10)
            if d > LAST or (md.get("until") and date(y, m, 1) > md["until"]):
                continue
            rent = sum(c["amount"] for c in CLAIMS if c["kind"] == "miete" and c["period"] == (y, m)
                       and (c["obj"] in md["objs"]))
            amt = int(rent * 0.62) // 100 * 100
            iban, bic = cl["payout"]
            if md["client"] == "mh" and d >= D("2026-10-01"):
                iban, bic = OWNERS["mh"]["new_payout"]
            money(a, d, -amt, cl["name"], iban, bic, f"Auszahlung Mieteinnahmen {m:02d}/{y} {', '.join(str(o) for o in md['objs'])}",
                  "ONLINE-UEBERWEISUNG", kind="payout")
    money("K1", D("2026-10-15"), -2131400, "Heideweg Wohnen GmbH", OWNERS["hw"]["payout"][0], OWNERS["hw"]["payout"][1],
          "Schlussauszahlung Mietverwaltung Heideweg 2a-2c zum 30.09.2026 lt. Uebergabeprotokoll", "ONLINE-UEBERWEISUNG",
          kind="payout")
    # the office's own account: salaries, rent, software, taxes
    for (y, m) in MONTHS[:10]:
        for name, role, gross in STAFF:
            d = last_bd(y, m) if m != 10 else D("2026-10-28")
            if d <= LAST:
                money("K0", d, -int(gross * 0.64), name, acct("staff", name)[0], acct("staff", name)[1],
                      f"Lohn/Gehalt {m:02d}/{y}", "Lohn/Gehalt", kind="cost")
        for d0, who, what, amt, ci in ((1, "Grundstücksverwaltung Wallstraße GbR", "Miete Büro Wallstraße 7", 168000, ""),
                                       (3, "Verwalterwerk Software GmbH", "Lizenz Verwaltungsprogramm", 41650, creditor_id("00001775020")),
                                       (10, "Finanzamt Brakenfeld", "USt-VZ und Lohnsteuer", 412000, ""),
                                       (15, "BKK Heideland", "Beitragsnachweis", 298400, creditor_id("00000301200")),
                                       (20, "Heidenetz Telekommunikation GmbH", "Telefon/Internet", 8930, creditor_id("00000220033"))):
            d = next_bd(date(y, m, d0))
            if d <= LAST:
                money("K0", d, -amt, who, acct("k0", who)[0], acct("k0", who)[1], f"{what} {m:02d}/{y}",
                      "Lastschrift" if ci else "Überweisung", kind="cost", ci=ci)


# ------------------------------------------------------------------ mails
MAILS: list[dict] = []


def mail(key, when, frm, to, subject, body, att=(), reply=None, waits=None):
    MAILS.append(dict(key=key, when=when, frm=frm, to=to, subject=subject, body=body, att=list(att), reply=reply, waits=waits))


SIG = ("\n\nMit freundlichen Grüßen\n{who}\n\nKerkhoff Hausverwaltung GmbH\nWallstraße 7 · 29599 Brakenfeld\n"
       "Tel. 05861 44 71-0 · www.kerkhoff-hv.example\n")


def sig(who="Melanie Okafor, Mietenbuchhaltung"):
    return SIG.format(who=who)


def addr(who: dict) -> str:
    return f"{who['bankname']} <{who['mail']}>"


def build_mails():
    G, O = LD, OD
    kr = G["10611"]["who"]
    mail("kuend_kr", "2026-05-28T18:40", addr(kr), OWN_INFO, "Kündigung meiner Wohnung Gartenstraße 31",
         "Sehr geehrte Damen und Herren,\n\nhiermit kündige ich meine Wohnung Gartenstraße 31, 3. OG links, fristgerecht zum "
         "31.08.2026. Das Jobcenter weiß Bescheid.\nBitte nennen Sie mir einen Termin für die Übergabe.\n\nMit freundlichen Grüßen\n"
         "Dennis Krüger")
    mail("kuend_kr_re", "2026-05-29T09:12", OWN_MO, addr(kr), "AW: Kündigung meiner Wohnung Gartenstraße 31",
         "Sehr geehrter Herr Krüger,\n\nwir bestätigen Ihre Kündigung zum 31.08.2026. Die Wohnungsübergabe ist am 31.08.2026 um "
         "10:00 Uhr. Bitte geben Sie dem Jobcenter Ihre neue Anschrift." + sig(), reply="kuend_kr")
    wg = G["10703"]["who"]
    mail("bk_wid", "2026-06-03T20:05", addr(wg), OWN_INFO, "Widerspruch Betriebskostenabrechnung 2025",
         "Guten Tag,\n\nich widerspreche der Betriebskostenabrechnung 2025 für Schillerstraße 5. Die Hausmeisterkosten sind "
         "doppelt so hoch wie im Vorjahr. Bevor ich die Nachzahlung von 264,10 € leiste, möchte ich die Belege einsehen.\n\n"
         "Mit freundlichen Grüßen\nGisela Wessels")
    mail("bk_wid_re", "2026-06-05T10:30", OWN_MO, addr(wg), "AW: Widerspruch Betriebskostenabrechnung 2025",
         "Sehr geehrte Frau Wessels,\n\ngern können Sie die Belege in unserem Büro einsehen. Bitte vereinbaren Sie telefonisch "
         "einen Termin. Bis dahin stellen wir die Nachzahlung zurück." + sig(), reply="bk_wid")
    mail("notar", "2026-08-18T11:02", "Notariat Dr. Hesse <kanzlei@notariat-hesse.example>", OWN_JP,
         "Verwalterzustimmung § 12 WEG – Ostendstraße 101, Einheit 11",
         "Sehr geehrter Herr Pohlmann,\n\nin der Kaufsache Kowalski / Hoffmeister (UR-Nr. 1288/2026) bitten wir um die "
         "Verwalterzustimmung nach § 12 WEG in notariell beglaubigter Form sowie um Ihre Bestellungsurkunde. "
         "Der Besitzübergang ist für den 01.09.2026 vorgesehen.\n\nMit freundlichen Grüßen\nNotariat Dr. Hesse")
    mail("notar_re", "2026-08-20T15:44", OWN_JP, "Notariat Dr. Hesse <kanzlei@notariat-hesse.example>",
         "AW: Verwalterzustimmung § 12 WEG – Ostendstraße 101, Einheit 11",
         "Sehr geehrte Damen und Herren,\n\ndie beglaubigte Zustimmung und die Bestellungsurkunde gehen Ihnen heute per Post zu."
         + sig("Jens Pohlmann, WEG-Verwaltung"), reply="notar")
    kz = O["20402"]["who"]
    mail("kranz", "2026-08-28T19:15", addr(kz), OWN_JP, "Sonderumlage Dach – Bitte um Teilzahlung",
         "Hallo Herr Pohlmann,\n\nwir können die Sonderumlage von 3.120 € nicht auf einmal zahlen. Wir überweisen am 01.09. die "
         "Hälfte und am 01.10. den Rest. Ist das in Ordnung?\n\nViele Grüße\nUwe und Heike Kranz")
    mail("kranz_re", "2026-08-31T08:20", OWN_JP, addr(kz), "AW: Sonderumlage Dach – Bitte um Teilzahlung",
         "Sehr geehrte Eheleute Kranz,\n\ndas ist in Ordnung: 1.560,00 € zum 01.09. und 1.560,00 € zum 01.10.2026."
         + sig("Jens Pohlmann, WEG-Verwaltung"), reply="kranz")
    ky = G["10606"]["who"]
    mail("schimmel", "2026-09-03T07:48", addr(ky), OWN_INFO, "Schimmel im Badezimmer – Gartenstraße 31, Whg 6",
         "Sehr geehrte Damen und Herren,\n\nim Badezimmer meiner Wohnung bildet sich seit Wochen Schimmel an der Außenwand, "
         "obwohl ich täglich lüfte. Bitte schicken Sie jemanden vorbei. Ich mindere die Miete vorerst nicht.\n\n"
         "Mit freundlichen Grüßen\nElif Kaya", waits="L1")
    jw = G["10614"]["who"]
    mail("jc_luecke", "2026-09-10T16:22", addr(jw), OWN_MO, "Miete September – Jobcenter",
         "Hallo Frau Okafor,\n\ndas Jobcenter hat für September nicht gezahlt, weil mein Weiterbewilligungsantrag zu spät war. "
         "Der Antrag ist jetzt gestellt, das Jobcenter zahlt nach. Bitte keine Mahnung.\n\nDanke\nJacqueline Wagner", waits="L1")
    rd = O["20127"]["who"]
    mail("riedel", "2026-09-14T09:05", addr(rd), OWN_JP, "Rücklastschrift Hausgeld Juli",
         "Sehr geehrter Herr Pohlmann,\n\nich habe der Lastschrift für Juli widersprochen. Der Wirtschaftsplan 2026 enthält die "
         "Kosten für die Tiefgarage, obwohl ich keinen Stellplatz habe. Bitte erklären Sie mir das, bevor weiter abgebucht wird.\n\n"
         "Dr. Hannes Riedel", waits="L1")
    ht = G["10107"]["who"]
    mail("dbl", "2026-09-15T21:10", addr(ht), OWN_MO, "Miete September doppelt bezahlt",
         "Hallo,\n\nich habe die Septembermiete aus Versehen doppelt überwiesen (Dauerauftrag und noch einmal per Hand am 07.09.). "
         "Was machen wir jetzt?\n\nLea Hartmann", waits="L1")
    mail("meinhardt", "2026-09-16T10:37", "Jürgen Meinhardt <j.meinhardt@fernmail.example>", OWN_AK,
         "Neue Bankverbindung für die Auszahlungen",
         f"Liebe Frau Kerkhoff,\n\nwir haben die Bank gewechselt. Bitte überweisen Sie die Mieteinnahmen ab Oktober auf unser neues "
         f"Konto: IBAN {fmt_iban(OWNERS['mh']['new_payout'][0])}, Kontoinhaber Jürgen und Petra Meinhardt.\n\n"
         f"Viele Grüße\nJürgen Meinhardt", waits="L1")
    mail("rehm_inv", "2026-09-08T14:02", f"Heizungsbau Rehm GmbH <{SUPPLIERS['rehm']['mail']}>", FIRM["buch"],
         "Rechnung R-26-10418 Wartung Am Mühlenteich 3-5",
         "Sehr geehrte Damen und Herren,\n\nanbei unsere Rechnung R-26-10418 für die Wartung der Heizungsanlagen.\n\n"
         "Mit freundlichen Grüßen\nHeizungsbau Rehm GmbH", att=[("R-26-10418.pdf", ("sup", "R-26-10418"))])
    # delivery 2
    mail("dbl_re", "2026-09-22T09:15", OWN_MO, addr(ht), "AW: Miete September doppelt bezahlt",
         "Hallo Frau Hartmann,\n\nwir verrechnen die doppelte Zahlung mit der Oktobermiete. Bitte setzen Sie den Dauerauftrag "
         "für Oktober aus; ab November läuft er wie gewohnt." + sig(), reply="dbl")
    br = G["10612"]["who"]
    mail("zusage", "2026-09-24T18:02", addr(br), OWN_MO, "Mietrückstand – Zahlungszusage",
         "Sehr geehrte Frau Okafor,\n\nich war lange krank und hatte kein Einkommen. Ich zahle am 01.10. 1.000 € und den Rest "
         "in Raten ab November. Bitte kündigen Sie mir nicht.\n\nMarco Brandl")
    mail("zusage_re", "2026-09-25T11:30", OWN_MO, addr(br), "AW: Mietrückstand – Zahlungszusage",
         "Sehr geehrter Herr Brandl,\n\nwir nehmen Ihre Zusage zur Kenntnis. Bitte senden Sie uns bis 15.10. einen Ratenplan für "
         "den verbleibenden Rückstand." + sig(), reply="zusage")
    mail("heizung", "2026-09-29T06:55", addr(O["20118"]["who"]), OWN_JP, "Heizung Haus 42 kalt",
         "Guten Morgen,\n\nin Haus 42 ist die Heizung seit Samstag kalt. Bitte dringend jemanden schicken!\n\n"
         + O["20118"]["who"]["bankname"])
    mail("heizung_re", "2026-09-29T08:10", OWN_JP, addr(O["20118"]["who"]), "AW: Heizung Haus 42 kalt",
         "Guten Morgen,\n\nHeizungsbau Rehm ist beauftragt und kommt heute bis 12 Uhr." + sig("Jens Pohlmann, WEG-Verwaltung"),
         reply="heizung")
    cd = G["10313"]["who"]
    mail("kuend_cd", "2026-10-02T17:31", addr(cd), OWN_INFO, "Kündigung Mietvertrag Am Mühlenteich",
         "Sehr geehrte Damen und Herren,\n\nhiermit kündige ich meinen Mietvertrag für die Wohnung Am Mühlenteich 5 einschließlich "
         "Stellplatz fristgerecht zum 31.12.2026. Bitte bestätigen Sie mir den Eingang.\n\nMerle Cordes", waits="L2")
    bo = G["10305"]["who"]
    mail("raten_bo", "2026-10-05T12:12", addr(bo), OWN_MO, "Ratenzahlung Rückstand",
         "Hallo,\n\nich kann im Moment nur 600 € im Monat zahlen. Ab November kann ich zusätzlich 100 € für den Rückstand zahlen. "
         "Geht das?\n\nKevin Böhm", waits="L2")
    mail("feldkamp", "2026-10-12T19:40", addr(O["20503"]["who"]), OWN_JP, "Abrechnung 2025 vom Vorverwalter?",
         "Sehr geehrter Herr Pohlmann,\n\nwann bekommen wir die Jahresabrechnung 2025? Der Vorverwalter hat sie nicht mehr gemacht.\n\n"
         + O["20503"]["who"]["bankname"], waits="L2")
    mail("wasser108", "2026-10-20T07:15", addr(G["10804"]["who"]), OWN_INFO, "Wasserschaden Heideweg 2a",
         "Hallo,\n\nbei mir tropft Wasser aus der Decke im Flur. Bitte schnell jemanden schicken!\n\nThomas Neumann", waits="L2")
    mail("anfrage", "2026-10-28T10:20", "Beirat WEG Lindenhof <beirat.lindenhof@postfach.example>", OWN_AK,
         "Anfrage Verwaltungsangebot WEG Lindenhof (16 Einheiten)",
         "Sehr geehrte Frau Kerkhoff,\n\nunsere Gemeinschaft (16 Einheiten, Baujahr 1996) sucht ab 2027 eine neue Verwaltung. "
         "Bitte senden Sie uns ein Angebot.\n\nMit freundlichen Grüßen\nDer Verwaltungsbeirat", waits="L2")
    mail("fraud", "2026-10-26T13:44", f"Falk Bedachungen GmbH <{FRAUD['mail']}>", FIRM["buch"],
         "Korrigierte Schlussrechnung SR-2026-0131 – neue Bankverbindung",
         f"Sehr geehrte Damen und Herren,\n\naufgrund einer Umstellung unseres Bankinstituts bitten wir Sie, die Schlussrechnung "
         f"SR-2026-0131 ausschließlich auf unser neues Konto zu überweisen:\nIBAN {fmt_iban(FRAUD['iban'])}, BIC {FRAUD['bic']}.\n"
         f"Anbei die korrigierte Rechnung. Bitte überweisen Sie bis 30.10.\n\nMit freundlichen Grüßen\nBuchhaltung Falk Bedachungen",
         att=[("SR-2026-0131_korrigiert.pdf", ("fraud", "SR-2026-0131"))])
    mail("jc_danke", "2026-10-27T15:00", addr(jw), OWN_MO, "AW: Miete September – Jobcenter",
         "Hallo Frau Okafor,\n\ndas Jobcenter hat am 12.10. September und Oktober nachgezahlt. Danke für Ihre Geduld!\n\n"
         "Jacqueline Wagner", reply="jc_luecke")


# ------------------------------------------------------------------ truth
DELIV = [dict(folder="lieferung1", cut=CUTS[0], today=D("2026-09-18"), list_upd=D("2026-08-01"), book_upd=D("2026-08-31"),
              op_day=D("2026-08-31")),
         dict(folder="lieferung2", cut=CUTS[1], today=D("2026-10-30"), list_upd=D("2026-10-01"), book_upd=D("2026-10-15"),
              op_day=D("2026-09-30"))]


def paid_at(c: dict, cut: date) -> int:
    got = 0
    for m in MONEY:
        if m["date"] > cut:
            continue
        for cc, cents, eff in m["alloc"]:
            if cc is c and eff <= cut:
                got += cents
        if m.get("undo") and m["undo"][0] is c:
            got -= m["undo"][1]
    off = BKRES.get(c["konto"], {}).get("offset")
    if off and off[0] is c and off[2] <= cut:
        pass  # the offset is inside the rent's allocation already
    return got


def open_at(c: dict, cut: date) -> int:
    return c["amount"] - paid_at(c, cut)


def receipts_at(cut: date) -> list[dict]:
    out = []
    for m in MONEY:
        if m["date"] > cut or m["amount"] <= 0 or m["date"] < ACCTS[m["acct"]]["since"]:
            continue
        if m["kind"] not in ("rent", "rent_jc", "hg", "ja", "su", "fee"):
            continue
        used = sum(cents for c, cents, eff in m["alloc"] if eff <= cut)  # a prepayment is used on the day it came
        back = sum(a for d, a, _ in m["settled"] if d <= cut)
        rest = m["amount"] - used - back
        if rest > 0:
            out.append(dict(amount=rest, date=m["date"], payer=m["name"], vwz=m["vwz"], acct=m["acct"], note=m["note"]))
    return out


def truth_for(dv: dict) -> dict:
    cut, today = dv["cut"], dv["today"]
    op = {}
    for c in CLAIMS:
        if c["issued"] > cut:
            continue
        o = open_at(c, cut)
        assert o >= 0, (c["key"], o)
        if o > 0:
            op[c["key"]] = c
    overdue = sorted(k for k, c in op.items() if c["due"] < today)
    rec = receipts_at(cut)
    unb = []
    if cut < D("2026-10-05"):
        unb.append(dict(id="VV-205", what="Verwaltervertrag WEG Feldkamp 17-19 ab 01.07.2026: Juli–" +
                        ("August" if cut < D("2026-09-01") else "September") + " ohne Rechnung", net=eur(57200 * 3)))
    for s in SL:
        if s["date"] <= cut and not s["free"] and s["status"] == "erledigt" and (not s["inv"] or s["inv"]["date"] > cut):
            unb.append(dict(id=s["no"], what=s["what"], net=eur(s["net"])))
    payables = []
    for x in SUPINV:
        if x["date"] <= cut and (x["paid"] is None or x["paid"] > cut):
            payables.append(dict(supplier=SUPPLIERS[x["sup"]]["name"], invoice=x["no"], date=x["date"].isoformat(),
                                 gross=eur(x["gross"]), due=x["due"].isoformat(), objekt=x["obj"],
                                 pay_from=ACCTS[PD[x["obj"]]["acct"]]["what"]))
    detail = {k: dict(art=c["kind"], konto=c["konto"], name=c["name"], objekt=c["obj"], lage=c["unit"], soll=eur(c["amount"]),
                      bezahlt=eur(paid_at(c, cut)), offen=eur(open_at(c, cut)), faellig=c["due"].isoformat(), text=c["label"],
                      note=c["note"]) for k, c in sorted(op.items())}
    by_kind, by_konto = {}, {}
    for k, c in op.items():
        by_kind[c["kind"]] = round(by_kind.get(c["kind"], 0) + eur(open_at(c, cut)), 2)
        by_konto[c["konto"]] = round(by_konto.get(c["konto"], 0) + eur(open_at(c, cut)), 2)
    not_unbilled = [dict(id="SL-26-009", why="berechnet mit " + next(s["inv"]["no"] for s in SL if s["no"] == "SL-26-009")
                                             + ", Liste nicht nachgetragen"),
                    dict(id="SL-26-012", why="Kulanz, 0,00 €"), dict(id="VV-108", why="Mietverwaltung endet am 30.09.2026")]
    if cut >= D("2026-10-01"):
        not_unbilled.append(dict(id="SL-26-014", why="Baubetreuung läuft noch"))
    not_unbilled.append(dict(id="Mietverträge", why="Mieten, Hausgeld und Umlagen werden nie berechnet: geschuldet nach Vertrag/Beschluss"))
    deadlines = [dict(what="Betriebskostenabrechnung 2025 Gartenstraße 31 (106)", until="2026-12-31",
                      why="§ 556 Abs. 3 BGB: danach keine Nachforderung mehr"),
                 dict(what="Verwalterbestellung WEG Am Wall 9 endet 31.12.2026", until="2026-12-31",
                      why="Wiederbestellung in einer Versammlung nötig")]
    if cut >= D("2026-10-02"):
        deadlines.append(dict(what="Kündigung 10313 Cordes zum 31.12.2026 bestätigen, Übergabe planen", until="2026-12-31", why=""))
    waiting = [dict(when=m["when"][:10], subject=m["subject"]) for m in MAILS
               if m["waits"] and datetime.fromisoformat(m["when"]).date() <= cut]
    risks = []
    if cut >= D("2026-10-26"):
        risks.append(dict(what="Mail von falk-bedachung.example (echt: falk-bedachungen.example) mit litauischer IBAN und "
                               "„korrigierter“ Schlussrechnung SR-2026-0131", invoice="SR-2026-0131"))
    if cut >= D("2026-09-16"):
        risks.append(dict(what="Neue Bankverbindung Meinhardt für die Auszahlung (echt, aber bestätigen)", invoice=None))
    return dict(folder=dv["folder"], today=today.isoformat(), cut=cut.isoformat(),
                open={k: eur(open_at(c, cut)) for k, c in sorted(op.items())},
                open_total=eur(sum(open_at(c, cut) for c in op.values())),
                overdue=overdue,
                unbilled=[u["id"] for u in unb],
                over_budget=[],
                receipts=sorted(eur(x["amount"]) for x in rec),
                open_detail=detail, open_by_kind=by_kind, open_by_account=dict(sorted(by_konto.items())),
                unbilled_detail=unb, not_unbilled=not_unbilled,
                receipts_detail=[dict(amount=eur(x["amount"]), date=x["date"].isoformat(), payer=x["payer"], vwz=x["vwz"],
                                      konto=ACCTS[x["acct"]]["what"], note=x["note"]) for x in rec],
                payables=payables, deadlines=deadlines, mails_waiting=waiting, risks=risks,
                underbilled=[dict(konto="10206", what="Staffel ab 01.09.2026 (+40,00) nicht im Programm – geschuldet nach Vertrag, "
                                                      "steht unter open")],
                same_person=[dict(a="20603", b="20604", why="Uwe Petersen besitzt zwei Wohnungen im Mühlenhof"),
                             dict(a="10601", b="YILMAZ MEHMET", why="zahlt vom Konto des Ehemanns")],
                different_persons=[dict(a="10209", b="10804", why="zwei Thomas Neumann"),
                                   dict(a="10106", b="10319", why="Andrea und Andreas Schulz"),
                                   dict(a="20311", b="20361", why="Verkäufer und Käuferin derselben Wohnung")],
                not_receipts=[dict(what=m["vwz"], amount=eur(m["amount"]), why=m["kind"]) for m in MONEY
                              if m["date"] <= cut and m["amount"] > 0 and m["kind"] in ("insurance", "own_in", "hg_se", "ja_se",
                                                                                         "internal_in", "rent_ls", "hg_ls")
                              and m["date"] >= ACCTS[m["acct"]]["since"]][:40])


# ------------------------------------------------------------------ writers: spreadsheets and CSV
def csv_line(vals, sep=";", quote=True) -> str:
    out = []
    for v in vals:
        s = "" if v is None else str(v)
        out.append('"' + s.replace('"', '""') + '"' if quote else s)
    return sep.join(out)


def save_wb(wb, path: Path, who="Melanie Okafor"):
    wb.properties.creator = who
    wb.properties.lastModifiedBy = who
    wb.properties.created = datetime(2026, 1, 5, 9, 0)
    wb.properties.modified = datetime(2026, 1, 5, 9, 0)
    buf = io.BytesIO()
    wb.save(buf)
    path.write_bytes(_fix_zip(buf.getvalue()))


def _fix_zip(data: bytes) -> bytes:
    """Same bytes for the same content: openpyxl stamps the time of saving into the archive and into core.xml."""
    import re
    import zipfile
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for info in sorted(src.infolist(), key=lambda i: i.filename):
            zi = zipfile.ZipInfo(info.filename, date_time=(2026, 1, 5, 9, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            body = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                body = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)", rb"\g<1>2026-01-05T09:00:00Z\g<2>", body)
            z.writestr(zi, body)
    return out.getvalue()


def head(ws, row, heads, bold=True):
    for j, h in enumerate(heads, start=1):
        c = ws.cell(row=row, column=j, value=h)
        if bold:
            c.font = Font(bold=True)


def leases_at(d: date, export=False) -> list[dict]:
    return [L for L in LEASES if L["start"] <= d + timedelta(days=30 if export else 0) and (not L["end"] or L["end"] >= d)
            and not (L["obj"] == 108 and d > D("2026-09-30"))]


def write_mieterliste(path: Path, dv: dict):
    d = dv["cut"]
    heads = ["Mieternr", "Objekt", "Whg-Nr", "Lage", "Name", "Vorname", "Mietbeginn", "Mietende", "Wohnfl. m2",
             "Nettokaltmiete", "BK-Vorausz.", "HK-Vorausz.", "Stellplatz/Garage", "USt", "Gesamtmiete", "Zahlweise",
             "IBAN", "Telefon", "E-Mail", "Kaution Soll"]
    lines = [csv_line(heads)]
    for L in sorted(leases_at(d, export=True), key=lambda L: L["konto"]):
        k, b, h, s = rent_at(L, d, noted_only=True)
        t = k + b + h + s
        u = vat19(t) if L["vat"] else 0
        zw = {"LS": "Lastschrift", "JC": "Jobcenter/Überweisung", "Q": "Überweisung vierteljährlich"}.get(L["how"], "Überweisung"
                                                                                                         if L["how"] in ("UE", "UE3") else "Dauerauftrag")
        obj = PD[L["obj"]]["street"] if L["obj"] in PD else f"SE {L['lage']}"
        lines.append(csv_line([L["konto"], L["obj"], f"{L['unit']:02d}", L["lage"], L["who"]["nach"], L["who"]["vor"] or "",
                               dde(L["start"]), dde(L["end"]) if L["end"] else "", f"{L['area']:.1f}".replace(".", ","),
                               de_plain(k), de_plain(b), de_plain(h), de_plain(s), de_plain(u), de_plain(t + u), zw,
                               L["who"]["iban"], L["who"]["phone"], L["who"]["mail"], de_plain(L["deposit"])]))
    path.write_bytes(("\r\n".join(lines) + "\r\n").encode("cp1252", errors="replace"))


def write_sollmieten(path: Path, at: date, dv: dict, old=False):
    """Frau Okafor's own list: one sheet per object, each kept a little differently."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    groups = [(p["no"], p["street"]) for p in PROPS if p["kind"] == "MV"] + [(900, "SE-Verwaltung")]
    for no, street in groups:
        if no == 108 and at > D("2026-09-30"):
            continue
        ws = wb.create_sheet(str(no) if no != 900 else "SE")
        ws.cell(row=1, column=1, value=f"Sollmieten {street} – Stand {dde(at)}").font = Font(bold=True, size=12)
        style = no % 3
        if style == 0:
            heads = ["Mieter-Nr.", "Mieter", "Whg", "Kaltmiete", "NK-VZ", "Heizk.-VZ", "Garage/Stpl.", "Gesamt"]
        elif style == 1:
            heads = ["Nr", "Name", "Lage", "Kalt", "Betriebskosten", "Heizkosten", "Stellplatz"]
        else:
            heads = ["Konto", "Mieter/in", "Wohnung", "Grundmiete", "Vorauszahlung BK", "Vorauszahlung HK", "Garage", "Summe Soll", "Bem."]
        head(ws, 3, heads)
        r = 4
        tot = 0
        for L in sorted([L for L in leases_at(at) if (L["obj"] == no or (no == 900 and L["obj"] > 900))],
                        key=lambda L: L["unit"]):
            if old and L["start"] > D("2026-01-01"):
                continue
            if not old and L["konto"] == "10115" and at < D("2026-08-15"):
                continue
            k, b, h, s = rent_at(L, at, noted_only=True)
            t = k + b + h + s
            vals = [L["konto"], L["who"]["list"] if style != 1 else f"{L['who']['vor'] or ''} {L['who']['nach']}".strip(), L["lage"],
                    eur(k), eur(b), eur(h), eur(s) if s else None]
            if style == 0:
                vals.append(eur(t))
            if style == 2:
                vals += [eur(t), "Gewerbe zzgl. USt" if L["vat"] else ("Jobcenter" if L["how"] == "JC" else None)]
            for j, v in enumerate(vals, start=1):
                ws.cell(row=r, column=j, value=v)
            tot += t
            r += 1
        if no == 106:  # the vacant flat stays on the list with a note
            ws.cell(row=r, column=1, value="10615")
            ws.cell(row=r, column=2, value="LEERSTAND (Sanierung)")
            ws.cell(row=r, column=3, value=lage(PD[106], 15))
            r += 1
        if style != 1:
            ws.cell(row=r + 1, column=2, value="Summe").font = Font(bold=True)
            ws.cell(row=r + 1, column=8 if style == 0 else 8, value=eur(tot)).font = Font(bold=True)
    save_wb(wb, path)


def write_mietanpassungen(path: Path, dv: dict):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "2026"
    ws.cell(row=1, column=1, value="Mietanpassungen 2026 (Staffel, Index, § 558, Vorauszahlungen)").font = Font(bold=True)
    head(ws, 3, ["Mieter-Nr.", "Mieter", "Objekt", "Art", "wirksam ab", "bisher", "neu", "Schreiben vom", "Zustimmung",
                 "im Programm erfasst"])
    rows = [("10104", "Staffel", D("2026-07-01"), 61000, 64500, None, None, "ja"),
            ("10206", "Staffel", D("2026-09-01"), 56000, 60000, None, None, ""),
            ("10303", "Mieterhöhung § 558", D("2026-09-01"), 54000, 58000, D("2026-06-10"), D("2026-07-28"), "ja"),
            ("10108", "BK-Vorauszahlung", D("2026-09-01"), 14500, 16500, D("2026-06-15"), None, "ja"),
            ("10210", "BK-Vorauszahlung", D("2026-09-01"), 12500, 14000, D("2026-06-15"), None, "ja")]
    if dv["cut"] >= D("2026-08-20"):
        rows.append(("10408", "Indexmiete (VPI +2,1 %)", D("2026-10-01"), 98000, 100058, D("2026-08-18"), None,
                     "ja" if dv["cut"] >= D("2026-10-01") else ""))
    for k, (kt, art, frm, old, new, letter, ok, noted) in enumerate(rows, start=4):
        L = LD[kt]
        for j, v in enumerate([kt, L["who"]["list"], PD[L["obj"]]["street"], art, frm, eur(old), eur(new), letter, ok, noted],
                              start=1):
            c = ws.cell(row=k, column=j, value=v)
            if isinstance(v, date):
                c.number_format = "DD.MM.YYYY"
    ws.cell(row=len(rows) + 6, column=1, value="Staffel 10206: lt. Mietvertrag § 3 Abs. 2 – Staffel 3 von 4")
    save_wb(wb, path)


def write_kautionen(path: Path, dv: dict):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Kautionen"
    head(ws, 1, ["Mieter-Nr.", "Name", "Objekt", "Whg", "Kaution Soll", "Art", "gezahlt", "Kautionskonto", "angelegt am",
                 "zurückgezahlt am", "Bemerkung"])
    r = 2
    for L in sorted(LEASES, key=lambda L: L["konto"]):
        if not L["deposit"] or L["start"] > dv["cut"] + timedelta(days=40) or (L["obj"] == 108 and dv["cut"] > D("2026-09-30")):
            continue
        paid = L["deposit"]
        note = None
        if L["start"] >= D("2026-01-01"):
            paid = sum(paid_at(CK[f"{L['konto']} Kaution Rate {k}"], dv["book_upd"]) for k in (1, 2, 3))
            note = "Ratenzahlung"
        back = D("2026-09-15") if L["konto"] == "10105" and dv["cut"] >= D("2026-09-15") else None
        if L["konto"] == "10105":
            note = "Einbehalt 300,00 € bis BK 2026" if back else "Auszug 31.07., Abrechnung offen"
        vals = [L["konto"], L["who"]["list"], L["obj"], L["lage"], eur(L["deposit"]),
                "Barkaution" if L["konto"] not in ("10407",) else "Bankbürgschaft", eur(paid),
                acct("dep", L["konto"], bank="SPK")[0] if L["konto"] not in ("10407",) else "Bürgschaft Hansebank Nr. 77104",
                L["start"], back, note]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=j, value=v)
            if isinstance(v, date):
                c.number_format = "DD.MM.YYYY"
        r += 1
    save_wb(wb, path)


def write_op_mieter(path: Path, day: date):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "OP Mieter"
    ws.cell(row=1, column=1, value=f"Offene Posten Mieter – Stand {dde(day)}").font = Font(bold=True, size=12)
    ws.cell(row=2, column=1, value="Verwaltungsprogramm, Auswertung OP-Liste Personenkonten (Mieter), sortiert nach Konto")
    head(ws, 4, ["Personenkonto", "Name", "Objekt", "Lage", "Sollart", "Monat", "Soll", "Ist", "Offen", "Fällig", "Mahnstufe"])
    r = 5
    for c in sorted((c for c in CLAIMS if c["kind"] in ("miete", "kaution", "bk") and c["issued"] <= day), key=lambda c: (c["konto"], c["due"])):
        o = open_at(c, day)
        if o <= 0:
            continue
        stufe = 2 if c["konto"] == "10612" and day >= D("2026-08-20") else 1 if (day - c["due"]).days > 20 else 0
        vals = [c["konto"], c["name"], c["obj"], c["unit"], {"miete": "Miete", "kaution": "Kaution", "bk": "BK-Abrechnung"}[c["kind"]],
                f"{c['period'][1]:02d}/{c['period'][0]}" if c["period"] else "", eur(c["amount"]), eur(paid_at(c, day)), eur(o),
                c["due"], stufe or None]
        for j, v in enumerate(vals, start=1):
            cell = ws.cell(row=r, column=j, value=v)
            if isinstance(v, date):
                cell.number_format = "DD.MM.YYYY"
        r += 1
    ws.cell(row=r + 1, column=2, value="Summe offene Posten").font = Font(bold=True)
    save_wb(wb, path)


def write_op_eigentuemer(path: Path, day: date):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "OP Eigentümer"
    ws.cell(row=1, column=1, value=f"Offene Posten Eigentümer – Stand {dde(day)}").font = Font(bold=True, size=12)
    head(ws, 3, ["Konto", "Eigentümer", "WEG", "Einheit", "Sollart", "Zeitraum", "Soll", "gezahlt", "Rückstand", "fällig seit"])
    r = 4
    for c in sorted((c for c in CLAIMS if c["kind"] in ("hausgeld", "su", "ja") and c["issued"] <= day), key=lambda c: (c["konto"], c["due"])):
        o = open_at(c, day)
        if o <= 0:
            continue
        vals = [c["konto"], c["name"], PD[c["obj"]]["name"], c["unit"],
                {"hausgeld": "Hausgeld", "su": "Sonderumlage", "ja": "Abrechnungsspitze 2025"}[c["kind"]],
                f"{c['period'][1]:02d}/{c['period'][0]}" if c["period"] else "2026" if c["kind"] == "su" else "2025",
                eur(c["amount"]), eur(paid_at(c, day)), eur(o), c["due"]]
        for j, v in enumerate(vals, start=1):
            cell = ws.cell(row=r, column=j, value=v)
            if isinstance(v, date):
                cell.number_format = "DD.MM.YYYY"
        r += 1
    save_wb(wb, path, who="Jens Pohlmann")


def write_eigentuemer(path: Path, dv: dict):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for p in PROPS:
        if p["kind"] != "WEG":
            continue
        ws = wb.create_sheet(p["name"].replace("WEG ", "")[:28])
        ws.cell(row=1, column=1, value=f"{p['name']}, {p['street']} – Eigentümerliste / Hausgeld lt. Wirtschaftsplan 2026").font = Font(bold=True)
        heads = ["Konto", "Einheit", "Lage", "MEA", "Eigentümer", "Anschrift", "E-Mail", "Hausgeld mtl.", "Zahlweise", "IBAN",
                 "Mandatsreferenz"] if p["no"] != 204 else ["Einh.", "Eigentümer", "MEA/1000", "HG", "Konto-Nr", "Bemerkung"]
        head(ws, 3, heads)
        r = 4
        for o in sorted((o for o in OWN if o["weg"] == p["no"]), key=lambda o: (o["unit"], o["start"])):
            if o["start"] > dv["cut"] + timedelta(days=10) or (o["end"] and o["end"] < dv["cut"] - timedelta(days=60)):
                continue
            zw = {"LS": "Lastschrift", "SE": "über SE-Verwaltung (Treuhandkonto)", "PAIR": "Dauerauftrag (mit Whg 4)",
                  "DA_1": "Dauerauftrag", "DA_28": "Dauerauftrag", "UE": "Überweisung"}[o["how"]]
            who = o["who"]["list"]
            if o["end"]:
                who += f" (bis {dde(o['end'])})"
            if p["no"] == 204:
                vals = [o["unit"], who, o["mea"], eur(o["hg"]), o["konto"],
                        "Rückstand seit 05/2026 – Klage?" if o["konto"] == "20404" else None]
            else:
                vals = [o["konto"], o["unit"], o["lage"], o["mea"], who,
                        None if rr("res", o["konto"]).random() < 0.6 else f"{rr('st', o['konto']).choice(['Birkenweg', 'Am Hang', 'Lerchenstraße', 'Ulmenallee'])} {rr('hn', o['konto']).randint(1, 40)}, 29599 Brakenfeld",
                        o["who"]["mail"] or None, eur(o["hg"]), zw, o["who"]["iban"], o["mandate"] if o["how"] == "LS" else None]
            for j, v in enumerate(vals, start=1):
                ws.cell(row=r, column=j, value=v)
            r += 1
    save_wb(wb, path, who="Jens Pohlmann")


def write_ja(path: Path, dv: dict):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for p in PROPS:
        if p["kind"] != "WEG" or not p.get("etv"):
            continue
        ws = wb.create_sheet(str(p["no"]))
        ws.cell(row=1, column=1, value=f"Jahresabrechnung 2025 {p['name']} – Beschluss ETV {dde(p['etv'])}").font = Font(bold=True)
        head(ws, 3, ["Eigentümer-Konto", "Eigentümer", "Einheit", "Abrechnungsspitze", "Nachzahlung/Guthaben", "fällig"])
        r = 4
        for k, res in sorted(JARES.items()):
            o = OD[k]
            if o["weg"] != p["no"]:
                continue
            vals = [k, o["who"]["list"], o["unit"] if k != "20603" else "3 + 4", eur(res["res"]),
                    "Nachzahlung" if res["res"] > 0 else "Guthaben", res["due"]]
            for j, v in enumerate(vals, start=1):
                c = ws.cell(row=r, column=j, value=v)
                if isinstance(v, date):
                    c.number_format = "DD.MM.YYYY"
            r += 1
    save_wb(wb, path, who="Jens Pohlmann")


def write_su(path: Path, dv: dict):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sonderumlage"
    ws.cell(row=1, column=1, value="WEG Brunnenstraße 6 – Sonderumlage Dachsanierung, Beschluss TOP 4 ETV 23.07.2026").font = Font(bold=True)
    ws.cell(row=2, column=1, value="Gesamtbetrag 24.000,00 €, Verteilung nach MEA, fällig zum 01.09.2026")
    head(ws, 4, ["Einheit", "Konto", "Eigentümer", "MEA", "Anteil", "eingegangen", "Bemerkung"])
    for r, (k, share) in enumerate(sorted(SU204.items()), start=5):
        o = OD[k]
        got = paid_at(CK[f"{k} Sonderumlage 2026"], min(dv["cut"], dv["book_upd"]))
        vals = [o["unit"], k, o["who"]["list"], o["mea"], eur(share), eur(got) if got else None,
                "Ratenzahlung 2 x 1.560,00 vereinbart" if k == "20402" else None]
        for j, v in enumerate(vals, start=1):
            ws.cell(row=r, column=j, value=v)
    save_wb(wb, path, who="Jens Pohlmann")


def write_uebergabe(path: Path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "OP 30.06.2026"
    ws.cell(row=1, column=1, value="WEG Feldkamp 17-19 – Übergabe an Kerkhoff Hausverwaltung, offene Posten zum 30.06.2026").font = Font(bold=True)
    ws.cell(row=2, column=1, value="Vorverwalter: Immobilien Dreyer, Uhlenbrück")
    head(ws, 4, ["Whg", "Eigentümer", "Monat", "Hausgeld", "offen"])
    o = OD["20514"]
    for r, m in enumerate((5, 6), start=5):
        for j, v in enumerate([14, o["who"]["list"], f"{MONATE[m - 1]} 2026", 276.0, 276.0], start=1):
            ws.cell(row=r, column=j, value=v)
    ws.cell(row=8, column=1, value="Kontostand Girokonto 30.06.2026: 18.412,77 €; Erhaltungsrücklage: 96.300,00 €")
    save_wb(wb, path, who="Jens Pohlmann")


def write_bk(path: Path, dv: dict):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    by = {}
    for k, res in BKRES.items():
        if res["sent"] <= dv["cut"]:
            by.setdefault(LD[k]["obj"] if LD[k]["obj"] < 900 else 900, []).append((k, res))
    for obj, items in sorted(by.items()):
        ws = wb.create_sheet(str(obj) if obj != 900 else "SE")
        ws.cell(row=1, column=1, value=f"Betriebskostenabrechnung 2025 – {PD[obj]['street'] if obj in PD else 'SE-Verwaltung'}").font = Font(bold=True)
        head(ws, 3, ["Mieter-Nr", "Mieter", "Zeitraum", "Kosten", "Vorauszahlungen", "Ergebnis", "versandt", "fällig", "bezahlt/erstattet"])
        r = 4
        for k, res in sorted(items):
            L = LD[k]
            vz = (rent_at(L, D("2025-06-01"))[1] + rent_at(L, D("2025-06-01"))[2]) * 12
            c = CK.get(f"{k} BK 2025")
            paid = None
            if c and open_at(c, dv["book_upd"]) == 0:
                paid = "bezahlt"
            elif res["res"] < 0:
                paid = "verrechnet" if k == "10502" else "erstattet" if add_bd(res["sent"], 15) <= dv["book_upd"] else None
            vals = [k, L["who"]["list"], "01.01.–31.12.2025", eur(vz + res["res"]), eur(vz), eur(res["res"]), res["sent"],
                    res["due"], paid]
            for j, v in enumerate(vals, start=1):
                cell = ws.cell(row=r, column=j, value=v)
                if isinstance(v, date):
                    cell.number_format = "DD.MM.YYYY"
            r += 1
    save_wb(wb, path)


def write_objekte(path: Path, dv: dict):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Objekte"
    head(ws, 1, ["Objekt-Nr", "Bezeichnung", "Straße", "PLZ", "Ort", "Art", "Eigentümer / Gemeinschaft", "Einheiten",
                 "Stellplätze", "Baujahr", "Objektkonto", "zuständig", "verwaltet seit", "verwaltet bis"])
    for r, p in enumerate(PROPS, start=2):
        owner = OWNERS[p["owner"]]["name"] if p["kind"] == "MV" else "Gemeinschaft der Wohnungseigentümer"
        vals = [p["no"], p["name"], p["street"], "29599", "Brakenfeld", "Mietverwaltung" if p["kind"] == "MV" else "WEG-Verwaltung",
                owner, p["units"] + p.get("shops", 0), p.get("tg", 0) + p.get("garages", 0) or None, p["built"],
                fmt_iban(ACCTS[p["acct"]]["iban"]), "M. Okafor" if p["kind"] == "MV" else "J. Pohlmann", p["since"],
                p.get("until")]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=j, value=v)
            if isinstance(v, date):
                c.number_format = "DD.MM.YYYY"
    ws2 = wb.create_sheet("beendet")
    head(ws2, 1, ["Objekt-Nr", "Bezeichnung", "Straße", "Art", "Einheiten", "verwaltet seit", "verwaltet bis", "Grund"])
    for r, p in enumerate(ENDED, start=2):
        for j, v in enumerate([p["no"], p["name"], p["street"], "WEG-Verwaltung", p["units"], p["since"], p["until"], p["why"]], start=1):
            c = ws2.cell(row=r, column=j, value=v)
            if isinstance(v, date):
                c.number_format = "DD.MM.YYYY"
    save_wb(wb, path, who="Antje Kerkhoff")


def write_vertraege(path: Path, dv: dict):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Verwalterverträge"
    head(ws, 1, ["Vertrag", "Objekt(e)", "Auftraggeber", "Art", "Vertragsbeginn", "Vertragsende / bestellt bis", "Einheiten",
                 "Vergütung mtl. netto", "zzgl. USt", "Abrechnung", "Zahlung", "Bemerkung"])
    for r, md in enumerate(MANDATES, start=2):
        cl = CLIENTS[md["client"]]
        p = PD.get(md["objs"][0])
        start = p["since"] if p else D("2020-01-01")
        end = md.get("until") or (p.get("bestellt_bis") if p else None) or (date(start.year + 5, 12, 31) if p and p["kind"] == "WEG" else None)
        if md["vv"] == "VV-205":
            start, end = D("2026-07-01"), D("2028-12-31")
        if p and p["kind"] == "WEG" and md["vv"] != "VV-205":
            end = p.get("bestellt_bis") or D("2028-12-31")
        vals = [md["vv"], ", ".join(str(o) for o in md["objs"]), cl["name"],
                "WEG-Verwaltung" if md["vv"] >= "VV-201" and md["vv"] < "VV-900" else "SE-Verwaltung" if md["vv"] >= "VV-900" else "Mietverwaltung",
                start, end, md["units"], eur(md["net"]), "19 %", "monatlich",
                "Überweisung durch Eigentümerin" if md["acct"] is None else "Entnahme vom Objektkonto",
                "gekündigt zum 30.09.2026 (Verkauf)" if md["vv"] == "VV-108" else
                "Wiederbestellung ETV 2026 nötig" if md["vv"] == "VV-202" else None]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=j, value=v)
            if isinstance(v, date):
                c.number_format = "DD.MM.YYYY"
    save_wb(wb, path, who="Antje Kerkhoff")


def write_rechnungsausgang(path: Path, dv: dict):
    heads = ["Rechnungsnummer", "Rechnungsdatum", "Kundennummer", "Kunde", "Betreff", "Leistungszeitraum", "Netto", "USt",
             "Brutto", "Fällig am", "Status", "Bezahlt am"]
    lines = [csv_line(heads)]
    for i in sorted(INV, key=lambda i: i["no"]):
        if i["date"] > dv["cut"]:
            continue
        cl = CLIENTS[i["client"]]
        o = open_at(i["claim"], dv["book_upd"])
        paid_day = None
        if o == 0:
            paid_day = max((m["date"] for m in MONEY for c, _, _ in m["alloc"] if c is i["claim"] and m["date"] <= dv["book_upd"]),
                           default=None)
        status = "bezahlt" if paid_day else "offen"
        per = f"{dde(i['period'][0])} - {dde(i['period'][1])}" if i["period"] else ""
        lines.append(csv_line([i["no"], dde(i["date"]), cl["kd"], cl["name"], i["subject"] + (f" ({i['ref']})" if i["ref"].startswith("VV") else ""),
                               per, de_plain(i["net"]), de_plain(i["vat"]), de_plain(i["gross"]), dde(i["due"]), status,
                               dde(paid_day) if paid_day else ""]))
    path.write_bytes(("\r\n".join(lines) + "\r\n").encode("utf-8-sig"))


def write_sonderleistungen(path: Path, dv: dict):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "2026"
    head(ws, 1, ["Nr.", "Datum", "Objekt", "Leistung", "Auftraggeber", "Honorar netto", "Status", "berechnet (RE-Nr.)", "Bemerkung"])
    r = 2
    for s in SL:
        if s["date"] > dv["cut"]:
            continue
        billed = s["inv"]["no"] if s["inv"] and s["inv"]["date"] <= dv["list_upd"] and s["no"] != "SL-26-009" else None
        vals = [s["no"], s["date"], s["obj"], s["what"], CLIENTS[s["client"]]["name"], eur(s["net"]), s["status"], billed,
                "Kulanz – keine Berechnung" if s["free"] else None]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=j, value=v)
            if isinstance(v, date):
                c.number_format = "DD.MM.YYYY"
        r += 1
    save_wb(wb, path, who="Antje Kerkhoff")


# ------------------------------------------------------------------ writers: bank
def visible(acct_: str, cut: date) -> list[dict]:
    a = ACCTS[acct_]
    return sorted((m for m in MONEY if m["acct"] == acct_ and a["since"] <= m["date"] <= cut and m["kind"] != "prev"),
                  key=lambda m: (m["date"], m["id"]))


def opening(acct_: str) -> int:
    return {"K0": 4810233, "K1": 18733120, "K2": 6120455, "K3": 4422070, "K4": 1812300, "K5": 3610090, "K6": 288010,
            "K7": 1841277, "K8": 2299310}[acct_]


def write_spk(path: Path, acct_: str, cut: date):
    heads = ["Auftragskonto", "Buchungstag", "Valutadatum", "Buchungstext", "Verwendungszweck", "Glaeubiger ID",
             "Mandatsreferenz", "Kundenreferenz (End-to-End)", "Sammlerreferenz", "Lastschrift Ursprungsbetrag",
             "Auslagenersatz Ruecklastschrift", "Beguenstigter/Zahlungspflichtiger", "Kontonummer/IBAN", "BIC (SWIFT-Code)",
             "Betrag", "Waehrung", "Info"]
    a = ACCTS[acct_]
    lines = [csv_line(heads)]
    for m in reversed(visible(acct_, cut)):
        lines.append(csv_line([a["iban"], dshort(m["date"]), dshort(m["date"]), m["btext"], m["vwz"], m["ci"], m["mandate"], m["e2e"],
                               m["sammler"], de_plain(m["urspr"]) if m["urspr"] else "", de_plain(m["auslagen"]) if m["auslagen"] else "",
                               sepa_name(m["name"]), m["iban"], m["bic"], de_plain(m["amount"]), "EUR", "Umsatz gebucht"]))
    path.write_bytes(("\r\n".join(lines) + "\r\n").encode("cp1252", errors="replace"))


def write_vr(path: Path, acct_: str, cut: date):
    a = ACCTS[acct_]
    ms = visible(acct_, cut)
    bal = opening(acct_) + sum(m["amount"] for m in ms)
    out = ['"Volksbank Brakenfeld-Heide eG"', "", '"Umsatzanzeige"', "",
           csv_line(["BLZ:", BANKS["VB"][0], None, "Datum:", dde(next_bd(cut + timedelta(days=1)))]),
           csv_line(["Konto:", a["kto"], None, "Uhrzeit:", "07:42:10"]),
           csv_line(["Abfrage von:", "Antje Kerkhoff", None, "Kontoinhaber:", a["holder"]]), "",
           csv_line(["Zeitraum:", None, "von:", "01.01.2026", "bis:", dde(cut)]),
           csv_line(["Betrag in EUR:", None, "von:", " ", "bis:", " "]),
           csv_line(["Sortiert nach:", "Buchungstag", "absteigend"]), "",
           csv_line(["Buchungstag", "Valuta", "Auftraggeber/Zahlungsempfänger", "Empfänger/Zahlungspflichtiger", "Konto-Nr.", "IBAN",
                     "BLZ", "BIC", "Vorgang/Verwendungszweck", "Kundenreferenz", "Währung", "Umsatz", " "])]
    for m in reversed(ms):
        inc = m["amount"] > 0
        out.append(csv_line([dde(m["date"]), dde(m["date"]), sepa_name(m["name"]) if inc else a["holder"],
                             a["holder"] if inc else sepa_name(m["name"]), "", m["iban"], "", m["bic"],
                             f"{m['btext']}\n{m['vwz']}", "", "EUR", de(abs(m["amount"])), "H" if inc else "S"]))
    out += ["", csv_line([dde(cut), None, None, None, None, None, None, None, "Endsaldo", None, "EUR", de(abs(bal)),
                          "H" if bal >= 0 else "S"])]
    path.write_bytes(("\r\n".join(out) + "\r\n").encode("cp1252", errors="replace"))


GVC = {"CRDT": ("PMNT", "RCDT", "ESCT", "166"), "DBIT": ("PMNT", "ICDT", "ESCT", "116"), "LS": ("PMNT", "IDDT", "ESDD", "171"),
       "RET": ("PMNT", "IDDT", "UPDD", "109"), "FEE": ("ACMT", "MDOP", "CHRG", "805"), "DD": ("PMNT", "IDDT", "ESDD", "105")}


def camt_party(tag, name, iban, ci=""):
    pid = f"<Id><PrvtId><Othr><Id>{ci}</Id><SchmeNm><Prtry>SEPA</Prtry></SchmeNm></Othr></PrvtId></Id>" if ci else ""
    acc = f"<{tag}Acct><Id><IBAN>{iban}</IBAN></Id></{tag}Acct>" if iban else ""
    return f"<{tag}><Pty><Nm>{xesc(name)}</Nm>{pid}</Pty></{tag}>{acc}"


def camt_tx(acct_, m, amount, cd, name, iban, bic, vwz, e2e, mandate, kind, ref):
    a = ACCTS[acct_]
    dom, fam, sub, gvc = GVC[kind]
    own = camt_party("Cdtr" if cd == "CRDT" else "Dbtr", a["holder"], a["iban"])
    if kind == "FEE":
        parties = ""
    elif cd == "CRDT":
        parties = f"<RltdPties>{camt_party('Dbtr', name, iban)}{own}</RltdPties>"
    else:
        parties = f"<RltdPties>{own}{camt_party('Cdtr', name, iban, m.get('ci', '') if kind in ('DD',) else '')}</RltdPties>"
    agents = (f"<RltdAgts><{'Dbtr' if cd == 'CRDT' else 'Cdtr'}Agt><FinInstnId><BICFI>{bic}</BICFI></FinInstnId>"
              f"</{'Dbtr' if cd == 'CRDT' else 'Cdtr'}Agt></RltdAgts>") if bic else ""
    refs = f"<AcctSvcrRef>{ref}</AcctSvcrRef><EndToEndId>{xesc(e2e or 'NOTPROVIDED')}</EndToEndId>"
    if mandate:
        refs += f"<MndtId>{xesc(mandate)}</MndtId>"
    rtr = (f"<RtrInf><Rsn><Cd>{m['rtr']}</Cd></Rsn></RtrInf>" if m.get("rtr") else "")
    rmt = f"<RmtInf><Ustrd>{xesc(vwz[:140])}</Ustrd></RmtInf>" if vwz else ""
    return (f"<TxDtls><Refs>{refs}</Refs><Amt Ccy=\"EUR\">{xamt(amount)}</Amt><CdtDbtInd>{cd}</CdtDbtInd>"
            f"<BkTxCd><Domn><Cd>{dom}</Cd><Fmly><Cd>{fam}</Cd><SubFmlyCd>{sub}</SubFmlyCd></Fmly></Domn>"
            f"<Prtry><Cd>NTRF+{gvc}</Cd><Issr>DK</Issr></Prtry></BkTxCd>{parties}{agents}{rmt}{rtr}"
            f"<AddtlTxInf>{xesc(m['btext'])}</AddtlTxInf></TxDtls>")


def camt_entry(acct_, m, k):
    cd = "CRDT" if m["amount"] > 0 else "DBIT"
    kind = ("LS" if m["batch"] else "RET" if m["kind"] == "return" else "FEE" if m["kind"] == "fee" else
            "DD" if m["ci"] and cd == "DBIT" else cd)
    dom, fam, sub, gvc = GVC[kind]
    ref = f"{m['date']:%Y%m%d}{k:06d}"
    if m["batch"]:
        if m["details"]:
            txs = "".join(camt_tx(acct_, m, x["amount"], cd, sepa_name(x["o"]["who"]["bankname"]), x["o"]["who"]["iban"],
                                  x["o"]["who"]["bic"], x["vwz"], x["e2e"], x["mandate"], "LS", f"{ref}{j:03d}")
                          for j, x in enumerate(m["batch"], 1))
        else:
            txs = (f"<TxDtls><Refs><AcctSvcrRef>{ref}</AcctSvcrRef><PmtInfId>{m['sammler']}</PmtInfId></Refs>"
                   f"<Amt Ccy=\"EUR\">{xamt(m['amount'])}</Amt><CdtDbtInd>CRDT</CdtDbtInd>"
                   f"<RmtInf><Ustrd>{xesc(m['vwz'])}</Ustrd></RmtInf><AddtlTxInf>{xesc(m['btext'])}</AddtlTxInf></TxDtls>")
        dtls = (f"<NtryDtls><Btch><PmtInfId>{m['sammler']}</PmtInfId><NbOfTxs>{len(m['batch'])}</NbOfTxs>"
                f"<TtlAmt Ccy=\"EUR\">{xamt(m['amount'])}</TtlAmt><CdtDbtInd>{cd}</CdtDbtInd></Btch>{txs}</NtryDtls>")
        bb = "<BtchBookg>true</BtchBookg>"
    else:
        dtls = (f"<NtryDtls>{camt_tx(acct_, m, m['amount'], cd, sepa_name(m['name']), m['iban'], m['bic'], m['vwz'], m['e2e'], m['mandate'], kind, ref)}"
                f"</NtryDtls>")
        bb = ""
    return (f"<Ntry><NtryRef>{k}</NtryRef><Amt Ccy=\"EUR\">{xamt(m['amount'])}</Amt><CdtDbtInd>{cd}</CdtDbtInd>"
            f"{'<RvslInd>true</RvslInd>' if m['kind'] == 'return' else ''}"
            f"<Sts><Cd>BOOK</Cd></Sts><BookgDt><Dt>{m['date'].isoformat()}</Dt></BookgDt><ValDt><Dt>{m['date'].isoformat()}</Dt></ValDt>"
            f"<AcctSvcrRef>{ref}</AcctSvcrRef><BkTxCd><Domn><Cd>{dom}</Cd><Fmly><Cd>{fam}</Cd><SubFmlyCd>{sub}</SubFmlyCd></Fmly></Domn>"
            f"<Prtry><Cd>NTRF+{gvc}</Cd><Issr>DK</Issr></Prtry></BkTxCd>{bb}{dtls}"
            f"<AddtlNtryInf>{xesc(m['btext'])}</AddtlNtryInf></Ntry>")


def camt_doc(acct_, ms, frm: date, to: date, seq: int, opening_bal: int, kind: str) -> bytes:
    a = ACCTS[acct_]
    close = opening_bal + sum(m["amount"] for m in ms)
    tz = "+01:00" if to >= D("2026-10-25") else "+02:00"
    created = f"{next_bd(to + timedelta(days=1)).isoformat()}T05:{10 + seq:02d}:00{tz}"
    tag, body, ns = ("Stmt", "BkToCstmrStmt", "camt.053.001.08") if kind == "053" else ("Rpt", "BkToCstmrAcctRpt", "camt.052.001.08")

    def bal(code, amount, when):
        return (f"<Bal><Tp><CdOrPrtry><Cd>{code}</Cd></CdOrPrtry></Tp><Amt Ccy=\"EUR\">{xamt(amount)}</Amt>"
                f"<CdtDbtInd>{'CRDT' if amount >= 0 else 'DBIT'}</CdtDbtInd><Dt><Dt>{when.isoformat()}</Dt></Dt></Bal>")
    ents = "\n".join(camt_entry(acct_, m, k) for k, m in enumerate(ms, start=1))
    cr = [m["amount"] for m in ms if m["amount"] > 0]
    db = [-m["amount"] for m in ms if m["amount"] < 0]
    xml = (f"<?xml version=\"1.0\" encoding=\"UTF-8\"?>\n<Document xmlns=\"urn:iso:std:iso:20022:tech:xsd:{ns}\" "
           f"xmlns:xsi=\"http://www.w3.org/2001/XMLSchema-instance\">\n<{body}><GrpHdr><MsgId>C{kind}{acct_}{to:%Y%m%d}{seq:03d}</MsgId>"
           f"<CreDtTm>{created}</CreDtTm></GrpHdr>\n<{tag}><Id>{a['iban']}{to:%Y%m%d}</Id><ElctrncSeqNb>{seq}</ElctrncSeqNb>"
           f"<CreDtTm>{created}</CreDtTm><FrToDt><FrDtTm>{frm.isoformat()}T00:00:00{tz}</FrDtTm><ToDtTm>{to.isoformat()}T23:59:59{tz}</ToDtTm></FrToDt>"
           f"<Acct><Id><IBAN>{a['iban']}</IBAN></Id><Ccy>EUR</Ccy><Ownr><Nm>{xesc(a['holder'])}</Nm></Ownr>"
           f"<Svcr><FinInstnId><BICFI>{a['bic']}</BICFI><Nm>{BANKS[a['bank']][2]}</Nm></FinInstnId></Svcr></Acct>\n"
           f"{bal('PRCD' if kind == '053' else 'OPBD', opening_bal, frm - timedelta(days=1))}{bal('CLBD' if kind == '053' else 'ITBD', close, to)}\n"
           f"<TxsSummry><TtlNtries><NbOfNtries>{len(ms)}</NbOfNtries></TtlNtries><TtlCdtNtries><NbOfNtries>{len(cr)}</NbOfNtries>"
           f"<Sum>{xamt(sum(cr))}</Sum></TtlCdtNtries><TtlDbtNtries><NbOfNtries>{len(db)}</NbOfNtries><Sum>{xamt(sum(db))}</Sum>"
           f"</TtlDbtNtries></TxsSummry>\n{ents}\n</{tag}></{body}></Document>\n")
    return xml.encode("utf-8")


def write_camt(folder: Path, acct_: str, cut: date):
    ms = visible(acct_, cut)
    bal = opening(acct_)
    a = ACCTS[acct_]
    y, mo = a["since"].year, a["since"].month
    seq = 0
    while date(y, mo, 1) <= cut:
        seq += 1
        frm, to = date(y, mo, 1), month_end(y, mo)
        part = [m for m in ms if frm <= m["date"] <= min(to, cut)]
        if to <= cut:
            (folder / f"{a['iban']}_camt053_{y}-{mo:02d}.xml").write_bytes(camt_doc(acct_, part, frm, to, seq, bal, "053"))
        else:
            (folder / f"{a['iban']}_camt052_{cut.isoformat()}.xml").write_bytes(camt_doc(acct_, part, frm, cut, seq, bal, "052"))
        bal += sum(m["amount"] for m in part)
        y, mo = ym_next(y, mo)


def write_pain008(path: Path, b: dict):
    a = ACCTS[b["acct"]]
    d = b["date"]
    created = datetime.combine(prev_bd(d - timedelta(days=2)), datetime.min.time()).replace(hour=15, minute=12)
    tx = []
    for x in b["items"]:
        who = x["L"]["who"] if "L" in x else x["o"]["who"]
        sign = (x["L"]["start"] if "L" in x else D("2016-01-01")).isoformat()
        tx.append(f"""      <DrctDbtTxInf>
        <PmtId><EndToEndId>{x['e2e']}</EndToEndId></PmtId>
        <InstdAmt Ccy="EUR">{xamt(x['amount'])}</InstdAmt>
        <DrctDbtTx><MndtRltdInf><MndtId>{x['mandate']}</MndtId><DtOfSgntr>{sign}</DtOfSgntr></MndtRltdInf></DrctDbtTx>
        <DbtrAgt><FinInstnId><BICFI>{who['bic']}</BICFI></FinInstnId></DbtrAgt>
        <Dbtr><Nm>{xesc(sepa_name(who['bankname']))}</Nm></Dbtr>
        <DbtrAcct><Id><IBAN>{who['iban']}</IBAN></Id></DbtrAcct>
        <RmtInf><Ustrd>{xesc(sepa_name(x['vwz'])[:140])}</Ustrd></RmtInf>
      </DrctDbtTxInf>""")
    total = sum(x["amount"] for x in b["items"])
    xml = f"""<?xml version="1.0" encoding="UTF-8"?>
<Document xmlns="urn:iso:std:iso:20022:tech:xsd:pain.008.001.08" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
  <CstmrDrctDbtInitn>
    <GrpHdr><MsgId>{b['money']['sammler']}</MsgId><CreDtTm>{created.isoformat()}</CreDtTm><NbOfTxs>{len(b['items'])}</NbOfTxs>
      <CtrlSum>{xamt(total)}</CtrlSum><InitgPty><Nm>{xesc(sepa_name(a['holder']))}</Nm></InitgPty></GrpHdr>
    <PmtInf>
      <PmtInfId>{b['money']['sammler']}</PmtInfId><PmtMtd>DD</PmtMtd><NbOfTxs>{len(b['items'])}</NbOfTxs><CtrlSum>{xamt(total)}</CtrlSum>
      <PmtTpInf><SvcLvl><Cd>SEPA</Cd></SvcLvl><LclInstrm><Cd>CORE</Cd></LclInstrm><SeqTp>RCUR</SeqTp></PmtTpInf>
      <ReqdColltnDt>{d.isoformat()}</ReqdColltnDt>
      <Cdtr><Nm>{xesc(sepa_name(a['holder']))}</Nm></Cdtr>
      <CdtrAcct><Id><IBAN>{a['iban']}</IBAN></Id></CdtrAcct>
      <CdtrAgt><FinInstnId><BICFI>{a['bic']}</BICFI></FinInstnId></CdtrAgt>
      <CdtrSchmeId><Id><PrvtId><Othr><Id>{a['ci']}</Id><SchmeNm><Prtry>SEPA</Prtry></SchmeNm></Othr></PrvtId></Id></CdtrSchmeId>
{chr(10).join(tx)}
    </PmtInf>
  </CstmrDrctDbtInitn>
</Document>
"""
    path.write_bytes(xml.encode("utf-8"))


# ------------------------------------------------------------------ writers: PDFs
def _pdf(draw) -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4, invariant=1)
    c.setTitle("Dokument")
    c.setAuthor("")
    draw(c)
    c.showPage()
    c.save()
    return buf.getvalue()


def _letterhead(c, sender: list[str], to: list[str], info: list[tuple[str, str]], title: str):
    c.setFont("Helvetica-Bold", 15)
    c.drawString(20 * mm, 272 * mm, sender[0])
    c.setFont("Helvetica", 8)
    c.drawString(20 * mm, 266 * mm, " · ".join(sender[1:]))
    c.setFont("Helvetica", 7)
    c.drawString(20 * mm, 240 * mm, f"{sender[0]} · {sender[1]}")
    c.setFont("Helvetica", 10)
    for k, line in enumerate(to):
        c.drawString(20 * mm, (232 - 5 * k) * mm, line)
    c.setFont("Helvetica", 9)
    for k, (a, b) in enumerate(info):
        c.drawString(130 * mm, (240 - 5 * k) * mm, a)
        c.drawString(165 * mm, (240 - 5 * k) * mm, b)
    c.setFont("Helvetica-Bold", 13)
    c.drawString(20 * mm, 196 * mm, title)


def _rows(c, y, heads, rows, widths):
    c.setFont("Helvetica-Bold", 8.5)
    x = 20
    for h, w in zip(heads, widths):
        c.drawString(x * mm, y, h)
        x += w
    c.line(20 * mm, y - 2, 195 * mm, y - 2)
    y -= 13
    c.setFont("Helvetica", 8.5)
    for row in rows:
        x = 20
        for v, w in zip(row, widths):
            c.drawString(x * mm, y, str(v))
            x += w
        y -= 12
    return y


def _text(c, y, lines, size=9):
    c.setFont("Helvetica", size)
    for line in lines:
        c.drawString(20 * mm, y, line)
        y -= 12
    return y


OWN_SENDER = [FIRM["name"], f"{FIRM['street']} · {FIRM['zip']} {FIRM['city']}", f"Tel. {FIRM['phone']}", FIRM["mail"]]


def own_footer(c):
    c.setFont("Helvetica", 6.8)
    c.line(20 * mm, 22 * mm, 195 * mm, 22 * mm)
    for k, line in enumerate([f"{FIRM['name']} · {FIRM['street']} · {FIRM['zip']} {FIRM['city']} · {FIRM['web']}",
                              f"Geschäftsführerin {FIRM['gf']} · {FIRM['hrb']} · USt-IdNr. {FIRM['vat']}",
                              f"Volksbank Brakenfeld-Heide eG · IBAN {fmt_iban(ACCTS['K0']['iban'])} · BIC {ACCTS['K0']['bic']}"]):
        c.drawString(20 * mm, (18 - 3.2 * k) * mm, line)


def pdf_own(i: dict) -> bytes:
    cl = CLIENTS[i["client"]]

    def draw(c):
        _letterhead(c, OWN_SENDER, cl["addr"], [("Rechnungs-Nr.", i["no"]), ("Datum", dde(i["date"])), ("Kunden-Nr.", cl["kd"]),
                                               ("Leistungszeitraum", f"{i['period'][0]:%m/%Y}" if i["period"] else dde(i["date"]))],
                    f"Rechnung {i['no']}")
        y = _rows(c, 180 * mm, ["Pos.", "Leistung", "Betrag"], [("1", i["subject"][:80], de(i["net"]) + " €")], [12, 140, 30])
        y = _text(c, y - 6, [f"Nettobetrag {de(i['net'])} €", f"Umsatzsteuer 19 % {de(i['vat'])} €",
                             f"Rechnungsbetrag {de(i['gross'])} €", "",
                             (f"Der Betrag wird gemäß Verwaltervertrag {i['ref']} vom Objektkonto entnommen."
                              if i["kind"] == "VV" and i["client"] != "al" else f"Zahlbar bis {dde(i['due'])} ohne Abzug.")])
        own_footer(c)
    return _pdf(draw)


def pdf_mahnung(konto: str, day: date, stufe: int) -> bytes:
    L = LD.get(konto)
    o = OD.get(konto)
    who = (L or o)["who"]
    items = [c for c in CLAIMS if c["konto"] == konto and c["issued"] <= day and open_at(c, day) > 0 and c["due"] < day]
    total = sum(open_at(c, day) for c in items)
    to = [f"{'Frau' if who['vor'] in ('Sabine', 'Jacqueline', 'Olga') else 'Herrn'} {who['vor']} {who['nach']}",
          PD[(L or {}).get("obj") or o["weg"]]["street"], "29599 Brakenfeld"]

    def draw(c):
        _letterhead(c, OWN_SENDER, to, [("Personenkonto", konto), ("Datum", dde(day)), ("Sachbearbeitung", "M. Okafor")],
                    f"{stufe}. Mahnung" if stufe > 1 else "Zahlungserinnerung")
        y = _text(c, 182 * mm, ["Sehr geehrte Damen und Herren,", "",
                                "auf Ihrem Konto sind folgende Beträge trotz Fälligkeit noch offen:"])
        y = _rows(c, y - 4, ["Sollart", "Zeitraum", "fällig", "offen"],
                  [(c_["label"].split()[0], f"{c_['period'][1]:02d}/{c_['period'][0]}" if c_["period"] else "", dde(c_["due"]),
                    de(open_at(c_, day)) + " €") for c_ in items], [40, 40, 40, 40])
        y = _text(c, y - 6, [f"Gesamt: {de(total)} €" + (" zzgl. Mahngebühr 5,00 €" if stufe > 1 else ""), "",
                             f"Bitte überweisen Sie den Betrag bis zum {dde(day + timedelta(days=10))} auf das Konto",
                             f"IBAN {fmt_iban(ACCTS[(L and acct_of_lease(L)) or PD[o['weg']]['acct']]['iban'])} unter Angabe Ihres Personenkontos.",
                             "Bei Mietrückständen von zwei Monatsmieten behalten wir uns die fristlose Kündigung vor (§ 543 BGB)."
                             if L else "Der Verwaltungsbeirat ist informiert."])
        own_footer(c)
    return _pdf(draw)


def pdf_bk(konto: str) -> bytes:
    L = LD[konto]
    res = BKRES[konto]
    vz = (rent_at(L, D("2025-06-01"))[1] + rent_at(L, D("2025-06-01"))[2]) * 12

    def draw(c):
        _letterhead(c, OWN_SENDER, [L["who"]["bankname"], PD[L["obj"]]["street"], "29599 Brakenfeld"],
                    [("Mieter-Nr.", konto), ("Datum", dde(res["sent"])), ("Zeitraum", "01.01.–31.12.2025")],
                    "Betriebskostenabrechnung 2025")
        rows = [("Grundsteuer", "Wohnfläche", de(int((vz + res['res']) * 0.12)) + " €"), ("Wasser/Abwasser", "Verbrauch", de(int((vz + res['res']) * 0.21)) + " €"),
                ("Heizung/Warmwasser", "HeizkostenV", de(int((vz + res['res']) * 0.38)) + " €"), ("Hausmeister/Garten", "Wohnfläche", de(int((vz + res['res']) * 0.16)) + " €"),
                ("Versicherung, Müll, Strom", "Wohnfläche", de((vz + res['res']) - int((vz + res['res']) * 0.12) - int((vz + res['res']) * 0.21) - int((vz + res['res']) * 0.38) - int((vz + res['res']) * 0.16)) + " €")]
        y = _rows(c, 182 * mm, ["Kostenart", "Schlüssel", "Ihr Anteil"], rows, [60, 50, 40])
        y = _text(c, y - 6, [f"Ihre Kosten {de(vz + res['res'])} €, Ihre Vorauszahlungen {de(vz)} €",
                             (f"Nachzahlung {de(res['res'])} € – bitte bis {dde(res['due'])} überweisen."
                              if res["res"] > 0 else f"Guthaben {de(-res['res'])} € – wir überweisen es Ihnen."),
                             "Belege können nach Terminabsprache in unserem Büro eingesehen werden."])
        own_footer(c)
    return _pdf(draw)


def pdf_protokoll() -> bytes:
    def draw(c):
        _letterhead(c, OWN_SENDER, ["An alle Eigentümer", "WEG Brunnenstraße 6", "29599 Brakenfeld"],
                    [("Versammlung", "23.07.2026"), ("Ort", "Büro Wallstraße 7")], "Niederschrift der Eigentümerversammlung")
        _text(c, 182 * mm, ["TOP 1 Jahresabrechnung 2025: Die Abrechnungsspitzen werden beschlossen; fällig vier Wochen nach Beschluss.",
                            "TOP 3 Dachsanierung: Auftrag an Falk Bedachungen GmbH, Brakenfeld, 18.400,00 € brutto (Angebot 30.06.).",
                            "TOP 4 Sonderumlage: 24.000,00 € nach Miteigentumsanteilen, fällig zum 01.09.2026.",
                            "TOP 5 Hausgeldrückstand Einheit 4: Die Verwaltung mahnt; bei weiterem Verzug Beschluss über Klage.",
                            "Abstimmung jeweils einstimmig bei 7 von 8 anwesenden bzw. vertretenen Einheiten."])
        own_footer(c)
    return _pdf(draw)


def pdf_supplier(x: dict, fraud=False) -> bytes:
    s = SUPPLIERS[x["sup"]]

    def draw(c):
        _letterhead(c, [s["name"], f"{s['street']} · {s['zip']} {s['city']}", f"Tel. {s['phone']}",
                        FRAUD["mail"] if fraud else s["mail"]], recipient_of(x["obj"]),
                    [("Rechnungs-Nr.", x["no"]), ("Rechnungsdatum", dde(x["date"])), ("Objekt", PD[x["obj"]]["street"]),
                     ("USt-IdNr.", s["vat"] or s.get("taxno", ""))],
                    ("KORRIGIERTE " if fraud else "") + ("Schlussrechnung " if x["no"].startswith("SR") else
                                                         "Abschlagsrechnung " if x["no"].startswith("AR") else "Rechnung ") + x["no"])
        rows = [(str(k), t[:60], f"{q} {u}", de(p), de(round(q * p))) for k, (t, q, u, p) in enumerate(x["lines"], start=1)]
        y = _rows(c, 180 * mm, ["Pos.", "Bezeichnung", "Menge", "Einzelpreis", "Gesamt"], rows, [12, 100, 25, 25, 25])
        y = _text(c, y - 6, [f"Nettobetrag {de(x['net'])} €", f"zzgl. 19 % USt {de(x['vat'])} €", f"Rechnungsbetrag {de(x['gross'])} €", "",
                             f"Zahlbar bis {dde(x['due'])} ohne Abzug." if not fraud else
                             "Achtung – neue Bankverbindung! Bitte ausschließlich auf das unten genannte Konto überweisen."])
        c.setFont("Helvetica", 6.8)
        bank = f"IBAN {fmt_iban(FRAUD['iban'])} · BIC {FRAUD['bic']}" if fraud else f"IBAN {fmt_iban(s['iban'])} · BIC {s['bic']}"
        c.drawString(20 * mm, 18 * mm, f"{s['name']} · {s['street']} · {s['zip']} {s['city']} · {s['hrb']}")
        c.drawString(20 * mm, 14.8 * mm, f"USt-IdNr. {s['vat']} · {bank}" if s["vat"] else f"Steuernummer {s.get('taxno', '')} · {bank}")
    data = _pdf(draw)
    if x["fmt"] == "zugferd" and not fraud:
        from pypdf import PdfReader, PdfWriter
        w = PdfWriter(clone_from=PdfReader(io.BytesIO(data)))
        w.add_attachment("factur-x.xml", cii_xml(x))
        w._ID = None
        buf = io.BytesIO()
        w.write(buf)
        data = buf.getvalue()
    return data


def cii_xml(x: dict) -> bytes:
    s = SUPPLIERS[x["sup"]]
    dd = x["date"].strftime("%Y%m%d")
    rec = recipient_of(x["obj"])
    items = "\n".join(f"""  <ram:IncludedSupplyChainTradeLineItem>
   <ram:AssociatedDocumentLineDocument><ram:LineID>{k}</ram:LineID></ram:AssociatedDocumentLineDocument>
   <ram:SpecifiedTradeProduct><ram:Name>{xesc(t)}</ram:Name></ram:SpecifiedTradeProduct>
   <ram:SpecifiedLineTradeAgreement><ram:NetPriceProductTradePrice><ram:ChargeAmount>{p / 100:.2f}</ram:ChargeAmount></ram:NetPriceProductTradePrice></ram:SpecifiedLineTradeAgreement>
   <ram:SpecifiedLineTradeDelivery><ram:BilledQuantity unitCode="C62">{q}</ram:BilledQuantity></ram:SpecifiedLineTradeDelivery>
   <ram:SpecifiedLineTradeSettlement><ram:ApplicableTradeTax><ram:TypeCode>VAT</ram:TypeCode><ram:CategoryCode>S</ram:CategoryCode><ram:RateApplicablePercent>19</ram:RateApplicablePercent></ram:ApplicableTradeTax>
    <ram:SpecifiedTradeSettlementLineMonetarySummation><ram:LineTotalAmount>{round(q * p) / 100:.2f}</ram:LineTotalAmount></ram:SpecifiedTradeSettlementLineMonetarySummation></ram:SpecifiedLineTradeSettlement>
  </ram:IncludedSupplyChainTradeLineItem>""" for k, (t, q, u, p) in enumerate(x["lines"], start=1))
    xml = f"""<?xml version="1.0" encoding="UTF-8"?>
<rsm:CrossIndustryInvoice xmlns:rsm="urn:un:unece:uncefact:data:standard:CrossIndustryInvoice:100" xmlns:ram="urn:un:unece:uncefact:data:standard:ReusableAggregateBusinessInformationEntity:100" xmlns:udt="urn:un:unece:uncefact:data:standard:UnqualifiedDataType:100">
 <rsm:ExchangedDocumentContext><ram:GuidelineSpecifiedDocumentContextParameter><ram:ID>urn:cen.eu:en16931:2017</ram:ID></ram:GuidelineSpecifiedDocumentContextParameter></rsm:ExchangedDocumentContext>
 <rsm:ExchangedDocument><ram:ID>{x['no']}</ram:ID><ram:TypeCode>380</ram:TypeCode><ram:IssueDateTime><udt:DateTimeString format="102">{dd}</udt:DateTimeString></ram:IssueDateTime></rsm:ExchangedDocument>
 <rsm:SupplyChainTradeTransaction>
{items}
  <ram:ApplicableHeaderTradeAgreement>
   <ram:BuyerReference>{x['obj']}</ram:BuyerReference>
   <ram:SellerTradeParty><ram:Name>{xesc(s['name'])}</ram:Name>
    <ram:PostalTradeAddress><ram:PostcodeCode>{s['zip']}</ram:PostcodeCode><ram:LineOne>{xesc(s['street'])}</ram:LineOne><ram:CityName>{s['city']}</ram:CityName><ram:CountryID>DE</ram:CountryID></ram:PostalTradeAddress>
    <ram:URIUniversalCommunication><ram:URIID schemeID="EM">{s['mail']}</ram:URIID></ram:URIUniversalCommunication>
    <ram:SpecifiedTaxRegistration><ram:ID schemeID="VA">{s['vat']}</ram:ID></ram:SpecifiedTaxRegistration></ram:SellerTradeParty>
   <ram:BuyerTradeParty><ram:Name>{xesc(rec[0])}</ram:Name>
    <ram:PostalTradeAddress><ram:PostcodeCode>{FIRM['zip']}</ram:PostcodeCode><ram:LineOne>{xesc(rec[1])}</ram:LineOne><ram:LineTwo>{FIRM['street']}</ram:LineTwo><ram:CityName>{FIRM['city']}</ram:CityName><ram:CountryID>DE</ram:CountryID></ram:PostalTradeAddress></ram:BuyerTradeParty>
  </ram:ApplicableHeaderTradeAgreement>
  <ram:ApplicableHeaderTradeDelivery><ram:ActualDeliverySupplyChainEvent><ram:OccurrenceDateTime><udt:DateTimeString format="102">{dd}</udt:DateTimeString></ram:OccurrenceDateTime></ram:ActualDeliverySupplyChainEvent></ram:ApplicableHeaderTradeDelivery>
  <ram:ApplicableHeaderTradeSettlement>
   <ram:InvoiceCurrencyCode>EUR</ram:InvoiceCurrencyCode>
   <ram:SpecifiedTradeSettlementPaymentMeans><ram:TypeCode>58</ram:TypeCode><ram:PayeePartyCreditorFinancialAccount><ram:IBANID>{s['iban']}</ram:IBANID></ram:PayeePartyCreditorFinancialAccount></ram:SpecifiedTradeSettlementPaymentMeans>
   <ram:ApplicableTradeTax><ram:CalculatedAmount>{x['vat'] / 100:.2f}</ram:CalculatedAmount><ram:TypeCode>VAT</ram:TypeCode><ram:BasisAmount>{x['net'] / 100:.2f}</ram:BasisAmount><ram:CategoryCode>S</ram:CategoryCode><ram:RateApplicablePercent>19</ram:RateApplicablePercent></ram:ApplicableTradeTax>
   <ram:SpecifiedTradePaymentTerms><ram:Description>Zahlbar bis {dde(x['due'])} ohne Abzug</ram:Description><ram:DueDateDateTime><udt:DateTimeString format="102">{x['due']:%Y%m%d}</udt:DateTimeString></ram:DueDateDateTime></ram:SpecifiedTradePaymentTerms>
   <ram:SpecifiedTradeSettlementHeaderMonetarySummation><ram:LineTotalAmount>{x['net'] / 100:.2f}</ram:LineTotalAmount><ram:TaxBasisTotalAmount>{x['net'] / 100:.2f}</ram:TaxBasisTotalAmount><ram:TaxTotalAmount currencyID="EUR">{x['vat'] / 100:.2f}</ram:TaxTotalAmount><ram:GrandTotalAmount>{x['gross'] / 100:.2f}</ram:GrandTotalAmount><ram:DuePayableAmount>{x['gross'] / 100:.2f}</ram:DuePayableAmount></ram:SpecifiedTradeSettlementHeaderMonetarySummation>
  </ram:ApplicableHeaderTradeSettlement>
 </rsm:SupplyChainTradeTransaction>
</rsm:CrossIndustryInvoice>
"""
    return xml.encode("utf-8")


def ubl_xml(x: dict) -> bytes:
    s = SUPPLIERS[x["sup"]]
    rec = recipient_of(x["obj"])
    lines = "\n".join(f"""  <cac:InvoiceLine><cbc:ID>{k}</cbc:ID><cbc:InvoicedQuantity unitCode="{'HUR' if u == 'Std.' else 'C62'}">{q}</cbc:InvoicedQuantity>
   <cbc:LineExtensionAmount currencyID="EUR">{round(q * p) / 100:.2f}</cbc:LineExtensionAmount>
   <cac:Item><cbc:Name>{xesc(t)}</cbc:Name><cac:ClassifiedTaxCategory><cbc:ID>S</cbc:ID><cbc:Percent>19</cbc:Percent><cac:TaxScheme><cbc:ID>VAT</cbc:ID></cac:TaxScheme></cac:ClassifiedTaxCategory></cac:Item>
   <cac:Price><cbc:PriceAmount currencyID="EUR">{p / 100:.2f}</cbc:PriceAmount></cac:Price></cac:InvoiceLine>""" for k, (t, q, u, p) in enumerate(x["lines"], start=1))
    xml = f"""<?xml version="1.0" encoding="UTF-8"?>
<ubl:Invoice xmlns:ubl="urn:oasis:names:specification:ubl:schema:xsd:Invoice-2" xmlns:cac="urn:oasis:names:specification:ubl:schema:xsd:CommonAggregateComponents-2" xmlns:cbc="urn:oasis:names:specification:ubl:schema:xsd:CommonBasicComponents-2">
 <cbc:CustomizationID>urn:cen.eu:en16931:2017#compliant#urn:xeinkauf.de:kosit:xrechnung_3.0</cbc:CustomizationID>
 <cbc:ProfileID>urn:fdc:peppol.eu:2017:poacc:billing:01:1.0</cbc:ProfileID>
 <cbc:ID>{x['no']}</cbc:ID>
 <cbc:IssueDate>{x['date'].isoformat()}</cbc:IssueDate>
 <cbc:DueDate>{x['due'].isoformat()}</cbc:DueDate>
 <cbc:InvoiceTypeCode>380</cbc:InvoiceTypeCode>
 <cbc:Note>Objekt {xesc(PD[x['obj']]['street'])}</cbc:Note>
 <cbc:DocumentCurrencyCode>EUR</cbc:DocumentCurrencyCode>
 <cbc:BuyerReference>{x['obj']}</cbc:BuyerReference>
 <cac:AccountingSupplierParty><cac:Party><cbc:EndpointID schemeID="EM">{s['mail']}</cbc:EndpointID>
   <cac:PartyName><cbc:Name>{xesc(s['name'])}</cbc:Name></cac:PartyName>
   <cac:PostalAddress><cbc:StreetName>{xesc(s['street'])}</cbc:StreetName><cbc:CityName>{s['city']}</cbc:CityName><cbc:PostalZone>{s['zip']}</cbc:PostalZone><cac:Country><cbc:IdentificationCode>DE</cbc:IdentificationCode></cac:Country></cac:PostalAddress>
   <cac:PartyTaxScheme><cbc:CompanyID>{s['vat']}</cbc:CompanyID><cac:TaxScheme><cbc:ID>VAT</cbc:ID></cac:TaxScheme></cac:PartyTaxScheme>
   <cac:PartyLegalEntity><cbc:RegistrationName>{xesc(s['name'])}</cbc:RegistrationName></cac:PartyLegalEntity>
   <cac:Contact><cbc:Name>Buchhaltung</cbc:Name><cbc:Telephone>{s['phone']}</cbc:Telephone><cbc:ElectronicMail>{s['mail']}</cbc:ElectronicMail></cac:Contact></cac:Party></cac:AccountingSupplierParty>
 <cac:AccountingCustomerParty><cac:Party><cbc:EndpointID schemeID="EM">{FIRM['buch']}</cbc:EndpointID>
   <cac:PartyName><cbc:Name>{xesc(rec[0])}</cbc:Name></cac:PartyName>
   <cac:PostalAddress><cbc:StreetName>{xesc(FIRM['street'])}</cbc:StreetName><cbc:AdditionalStreetName>{xesc(rec[1])}</cbc:AdditionalStreetName><cbc:CityName>{FIRM['city']}</cbc:CityName><cbc:PostalZone>{FIRM['zip']}</cbc:PostalZone><cac:Country><cbc:IdentificationCode>DE</cbc:IdentificationCode></cac:Country></cac:PostalAddress>
   <cac:PartyLegalEntity><cbc:RegistrationName>{xesc(rec[0])}</cbc:RegistrationName></cac:PartyLegalEntity></cac:Party></cac:AccountingCustomerParty>
 <cac:PaymentMeans><cbc:PaymentMeansCode>58</cbc:PaymentMeansCode><cac:PayeeFinancialAccount><cbc:ID>{s['iban']}</cbc:ID><cbc:Name>{xesc(s['name'])}</cbc:Name></cac:PayeeFinancialAccount></cac:PaymentMeans>
 <cac:PaymentTerms><cbc:Note>Zahlbar bis {dde(x['due'])} ohne Abzug</cbc:Note></cac:PaymentTerms>
 <cac:TaxTotal><cbc:TaxAmount currencyID="EUR">{x['vat'] / 100:.2f}</cbc:TaxAmount><cac:TaxSubtotal><cbc:TaxableAmount currencyID="EUR">{x['net'] / 100:.2f}</cbc:TaxableAmount><cbc:TaxAmount currencyID="EUR">{x['vat'] / 100:.2f}</cbc:TaxAmount><cac:TaxCategory><cbc:ID>S</cbc:ID><cbc:Percent>19</cbc:Percent><cac:TaxScheme><cbc:ID>VAT</cbc:ID></cac:TaxScheme></cac:TaxCategory></cac:TaxSubtotal></cac:TaxTotal>
 <cac:LegalMonetaryTotal><cbc:LineExtensionAmount currencyID="EUR">{x['net'] / 100:.2f}</cbc:LineExtensionAmount><cbc:TaxExclusiveAmount currencyID="EUR">{x['net'] / 100:.2f}</cbc:TaxExclusiveAmount><cbc:TaxInclusiveAmount currencyID="EUR">{x['gross'] / 100:.2f}</cbc:TaxInclusiveAmount><cbc:PayableAmount currencyID="EUR">{x['gross'] / 100:.2f}</cbc:PayableAmount></cac:LegalMonetaryTotal>
{lines}
</ubl:Invoice>
"""
    return xml.encode("utf-8")


# ------------------------------------------------------------------ mails out
def write_eml(path: Path, m: dict, ids: dict) -> str:
    dt = datetime.fromisoformat(m["when"])
    dt = dt.replace(tzinfo=timezone(timedelta(hours=1 if dt.date() >= D("2026-10-25") else 2)))
    msg = EmailMessage()
    msg["From"] = m["frm"]
    msg["To"] = m["to"]
    msg["Date"] = format_datetime(dt)
    msg["Subject"] = m["subject"]
    dom = m["frm"].split("@")[-1].strip(">")
    mid = f"<{hashlib.md5(m['key'].encode()).hexdigest()[:20]}@{dom}>"
    msg["Message-ID"] = mid
    if m["reply"] and m["reply"] in ids:
        msg["In-Reply-To"] = ids[m["reply"]]
        msg["References"] = ids[m["reply"]]
    msg.set_content(m["body"])
    for name, (kind, ref) in m["att"]:
        x = next(s for s in SUPINV if s["no"] == ref)
        msg.add_attachment(pdf_supplier(x, fraud=kind == "fraud"), maintype="application", subtype="pdf", filename=name)
    if m["att"]:
        msg.set_boundary("==" + hashlib.md5((m["key"] + "|b").encode()).hexdigest()[:24] + "==")
    path.write_bytes(bytes(msg))
    return mid


def slug(s: str) -> str:
    tr = {"ä": "ae", "ö": "oe", "ü": "ue", "Ä": "Ae", "Ö": "Oe", "Ü": "Ue", "ß": "ss", "–": "-", "§": "Par"}
    s = "".join(tr.get(ch, ch) for ch in s)
    return "".join(ch if ch.isalnum() else "_" for ch in s).strip("_")[:60]


# ------------------------------------------------------------------ the deliveries
def write_delivery(root: Path, dv: dict) -> list[str]:
    out, cut = root / dv["folder"], dv["cut"]
    for sub in ("Objekte", "Mieter/Mahnungen", "WEG/Protokolle", "Abrechnung", "Buchhaltung/Rechnungen", "Bank/Lastschriften",
                "Bank/camt", "Eingangsrechnungen", "Mails"):
        (out / sub).mkdir(parents=True, exist_ok=True)
    write_objekte(out / "Objekte" / "Objektliste.xlsx", dv)
    write_vertraege(out / "Objekte" / "Verwaltervertraege.xlsx", dv)
    for d0 in [x["cut"] for x in DELIV if x["cut"] <= cut]:  # the older export stays in the folder
        write_mieterliste(out / "Mieter" / f"Mieterliste_Export_{d0.isoformat()}.csv", dict(dv, cut=d0))
    write_sollmieten(out / "Mieter" / "Sollmieten_2026.xlsx", dv["list_upd"], dv)
    write_sollmieten(out / "Mieter" / "Sollmieten_2026_alt.xlsx", D("2026-01-02"), dv, old=True)
    write_mietanpassungen(out / "Mieter" / "Mietanpassungen_2026.xlsx", dv)
    write_kautionen(out / "Mieter" / "Kautionen.xlsx", dv)
    for x in DELIV:
        if x["cut"] <= cut:
            write_op_mieter(out / "Mieter" / f"OP-Liste_Mieter_Stand_{x['op_day']:%d.%m.%Y}.xlsx", x["op_day"])
            write_op_eigentuemer(out / "WEG" / f"OP-Liste_Eigentuemer_Stand_{x['op_day']:%d.%m.%Y}.xlsx", x["op_day"])
    dun = [("10612", D("2026-08-20"), 1), ("20404", D("2026-08-20"), 1), ("10305", D("2026-09-10"), 1), ("10304", D("2026-09-01"), 1),
           ("10612", D("2026-10-15"), 2), ("20404", D("2026-10-15"), 2)]
    for konto, day, st in dun:
        if day <= cut:
            (out / "Mieter" / "Mahnungen" / f"{day.isoformat()}_{'Mahnung' if st > 1 else 'Zahlungserinnerung'}_{konto}.pdf").write_bytes(
                pdf_mahnung(konto, day, st))
    write_eigentuemer(out / "WEG" / "Eigentuemerliste_WEG.xlsx", dv)
    write_ja(out / "WEG" / "Jahresabrechnung_2025_Ergebnisse.xlsx", dv)
    write_su(out / "WEG" / "Sonderumlage_Brunnenstrasse6.xlsx", dv)
    write_uebergabe(out / "WEG" / "Uebergabe_Vorverwalter_Feldkamp.xlsx")
    (out / "WEG" / "Protokolle" / "Protokoll_ETV_Brunnenstrasse6_2026-07-23.pdf").write_bytes(pdf_protokoll())
    write_bk(out / "Abrechnung" / "BK-Abrechnung_2025_Ergebnisse.xlsx", dv)
    for k in ("10304", "10502"):
        (out / "Abrechnung" / f"BK-Abrechnung_2025_{k}.pdf").write_bytes(pdf_bk(k))
    write_rechnungsausgang(out / "Buchhaltung" / f"Rechnungsausgang_2026_{dv['book_upd'].isoformat()}.csv", dv)
    write_sonderleistungen(out / "Buchhaltung" / "Sonderleistungen_2026.xlsx", dv)
    for i in INV:
        if i["date"] <= cut and (i["kind"] == "SL" or i["ref"] in ("VV-204", "VV-104") and i["date"] >= D("2026-07-01")):
            (out / "Buchhaltung" / "Rechnungen" / f"{i['no']}.pdf").write_bytes(pdf_own(i))
    # bank
    write_vr(out / "Bank" / f"Umsaetze_Geschaeftskonto_01.01.2026-{dde(cut)}.csv", "K0", cut)
    for k in ("K1", "K4", "K6", "K7", "K8"):
        write_spk(out / "Bank" / f"{next_bd(cut + timedelta(days=1)):%Y%m%d}-{ACCTS[k]['kto']}-umsatz.CSV", k, cut)
    for k in ("K2", "K3", "K5"):
        write_camt(out / "Bank" / "camt", k, cut)
    for b in RENT_LS_BATCHES + [b for b in HG_LS_BATCHES if b["wno"] == 201]:
        if b["date"] <= cut and b["date"] >= D("2026-06-01"):
            name = f"Mieteinzug_{b['date']:%Y-%m}.xml" if b["acct"] == "K1" else f"Hausgeld_Rosenhof_{b['date']:%Y-%m}.xml"
            write_pain008(out / "Bank" / "Lastschriften" / name, b)
    # incoming invoices
    for x in SUPINV:
        if x["date"] <= cut:
            sub = out / "Eingangsrechnungen" / x["date"].strftime("%Y-%m")
            sub.mkdir(exist_ok=True)
            name = f"{slug(SUPPLIERS[x['sup']]['name'])}_{x['no']}"
            (sub / (name + (".xml" if x["fmt"] == "xrechnung" else ".pdf"))).write_bytes(
                ubl_xml(x) if x["fmt"] == "xrechnung" else pdf_supplier(x))
    ids = {}
    for m in sorted(MAILS, key=lambda m: m["when"]):
        if datetime.fromisoformat(m["when"]).date() <= cut:
            ids[m["key"]] = write_eml(out / "Mails" / f"{m['when'][:10]}_{slug(m['subject'])}.eml", m, ids)
    return sorted(str(p.relative_to(out)) for p in out.rglob("*") if p.is_file())


def build_all():
    build_leases()
    build_owners()
    build_rent_claims()
    build_owner_claims()
    build_results()
    build_invoices()
    build_rent_money()
    build_results_money()
    build_hausgeld_money()
    build_fee_money()
    build_suppliers()
    build_costs()
    build_mails()


def main(outdir: str) -> None:
    root = Path(outdir)
    root.mkdir(parents=True, exist_ok=True)
    build_all()
    truth = dict(firm="Kerkhoff Hausverwaltung GmbH (WEG- und Mietverwaltung, Review 10)", deliveries=[])
    files = {}
    for dv in DELIV:
        files[dv["folder"]] = write_delivery(root, dv)
        truth["deliveries"].append(truth_for(dv))
    check_traps(truth)
    (root / "truth.json").write_text(json.dumps(truth, ensure_ascii=False, indent=1), encoding="utf-8")
    key = dict(leases={L["konto"]: dict(obj=L["obj"], lage=L["lage"], name=L["who"]["list"], start=L["start"].isoformat(),
                                        end=L["end"].isoformat() if L["end"] else None, how=L["how"],
                                        soll=[(s[0].isoformat(), eur(total_of(L, s[1:]))) for s in L["sched"]]) for L in LEASES},
               owners={o["konto"]: dict(weg=o["weg"], unit=o["unit"], name=o["who"]["list"], hausgeld=eur(o["hg"]), how=o["how"])
                       for o in OWN},
               invoices={i["no"]: dict(date=i["date"].isoformat(), client=CLIENTS[i["client"]]["name"], subject=i["subject"],
                                       gross=eur(i["gross"]), due=i["due"].isoformat(), ref=i["ref"]) for i in INV},
               money=[dict(acct=ACCTS[m["acct"]]["what"], date=m["date"].isoformat(), amount=eur(m["amount"]), name=m["name"],
                           vwz=m["vwz"], kind=m["kind"], alloc=[(c["key"], eur(x), e.isoformat()) for c, x, e in m["alloc"]])
                      for m in MONEY if m["kind"] not in ("cost",)],
               files=files)
    (root / "answer_key.json").write_text(json.dumps(key, ensure_ascii=False, indent=1), encoding="utf-8")
    for d in truth["deliveries"]:
        print(d["folder"], d["today"], "offen:", len(d["open"]), d["open_total"], "überfällig:", len(d["overdue"]),
              "nicht berechnet:", d["unbilled"], "Geld ohne Rechnung:", d["receipts"], "nach Art:", d["open_by_kind"])


def check_traps(truth: dict):
    """The header promises these; if a change breaks one, the generator says so."""
    l1, l2 = truth["deliveries"]
    o1, o2 = l1["open"], l2["open"]
    assert [k for k in o1 if k.startswith("10612 Miete")] == ["10612 Miete 07/2026", "10612 Miete 08/2026", "10612 Miete 09/2026"]
    assert (o2["10612 Miete 08/2026"], o2["10612 Miete 09/2026"], o2["10612 Miete 10/2026"]) == (490.0, 745.0, 745.0)
    assert "10612 Miete 07/2026" not in o2
    assert [o1[f"10305 Miete {m:02d}/2026"] for m in (6, 7, 8, 9)] == [212.0] * 4 and o2["10305 Miete 10/2026"] == 212.0
    assert [o1.get(f"10104 Miete {m:02d}/2026") for m in (7, 8, 9)] == [35.0] * 3 and o2["10104 Miete 10/2026"] == 35.0
    assert o1["10206 Miete 09/2026"] == 40.0 and o2["10206 Miete 10/2026"] == 40.0
    assert o1["10108 Miete 09/2026"] == 20.0 and o1["10210 Miete 09/2026"] == 15.0
    assert o1["10609 Miete 08/2026"] == 38.5 and o1["10609 Miete 09/2026"] == 38.5 and "10609 Miete 08/2026" not in o2
    assert o1["10614 Miete 09/2026"] == 612.0 and "10614 Miete 09/2026" not in o2
    assert o1["10802 Miete 09/2026"] == 689.0 and "10802 Miete 09/2026" not in o2
    assert o1["10115 Kaution Rate 2"] == 650.0 and not [k for k in o2 if k.startswith("10115")]
    assert o1["10304 BK 2025"] == 318.4 and "10304 BK 2025" not in o2
    assert o1["10305 BK 2025"] == 91.2 and o2["10305 BK 2025"] == 91.2
    assert o1["10205 BK 2025"] == 112.9 and "10205 BK 2025" not in o2
    assert o1["10703 BK 2025"] == 264.1 and o2["10703 BK 2025"] == 264.1
    assert o1["10407 Miete 09/2026"] == 2058.7 and not [k for k in o2 if k.startswith("10407")]
    assert o2["10408 Miete 10/2026"] == 20.58
    assert not [k for k in o1 if k.startswith(("10303", "10325", "10107", "10502", "10615", "10611", "10603"))]
    assert not [k for k in o2 if k.startswith(("10303", "10325", "10107", "10502", "10615", "10611", "10603", "108"))]
    assert [k for k in o1 if k.startswith("20404")] == ["20404 Abrechnung 2025"] + [f"20404 Hausgeld {m:02d}/2026" for m in (5, 6, 7, 8, 9)] + \
        ["20404 Sonderumlage 2026"]
    assert o1["20402 Sonderumlage 2026"] == 1560.0 and "20402 Sonderumlage 2026" not in o2
    assert o1["20127 Hausgeld 07/2026"] == 287.0 and "20127 Hausgeld 07/2026" not in o2
    assert o1["20211 Abrechnung 2025"] == 142.75 and "20211 Abrechnung 2025" not in o2
    assert o1["20514 Hausgeld 05/2026"] == 276.0 and "20514 Hausgeld 06/2026" in o1 and "20514 Hausgeld 05/2026" not in o2
    assert sorted(l1["receipts"]) == [268.0, 534.0, 690.0], l1["receipts"]
    assert sorted(l2["receipts"]) == [sorted(l2["receipts"])[0]] and l2["receipts"][0] == eur(total_of(LD["10817"], rent_at(LD["10817"], D("2026-10-01")))), l2["receipts"]
    assert l1["unbilled"] == ["VV-205", "SL-26-011"], l1["unbilled"]
    assert l2["unbilled"] == ["SL-26-011", "SL-26-013"], l2["unbilled"]
    re1 = sorted(k for k in o1 if k.startswith("RE-"))
    re2 = sorted(k for k in o2 if k.startswith("RE-"))
    assert {IK[k]["ref"] for k in re1} == {"VV-104", "VV-204", "SL-26-010"}, re1
    assert {IK[k]["ref"] for k in re2} == {"VV-104", "SL-26-010"} and len(re2) == 2, re2
    assert [p["invoice"] for p in l1["payables"]] == ["R-26-10418", "AR-2026-0112", "1088"], l1["payables"]
    assert [p["invoice"] for p in l2["payables"]] == ["GT-2026-412", "ST-26-5521", "SR-2026-0131", "R-26-10655"], l2["payables"]


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit("Aufruf: python3 hausverwaltung.py OUTDIR")
    main(sys.argv[1])
